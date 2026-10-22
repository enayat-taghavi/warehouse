from rest_framework.test import APITestCase
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
import io
import openpyxl
from personnel.models import (
    Company,
    Project,
    ProjectSection,
    SectionAssignment,
    PersonnelProfile,
    PersonnelBankAccount,
    InsuranceJobTitle
)

User = get_user_model()


class Phase5ExcelAndEndpointsTests(APITestCase):
    def setUp(self):
        self.company = Company.objects.create(
            name="شرکت پارسیان تست",
            national_id="10100000001",
            economic_code="111111111111"
        )
        self.project = Project.objects.create(
            name="پروژه تست",
            company=self.company
        )
        self.section = ProjectSection.objects.create(
            name="بخش تست",
            project=self.project
        )
        self.manager = User.objects.create_superuser(
            username="manager_p5",
            email="mgr_p5@test.com",
            password="pass"
        )
        self.operator = User.objects.create_user(
            username="operator_p5",
            email="op_p5@test.com",
            password="pass"
        )
        SectionAssignment.objects.create(
            user=self.operator,
            section=self.section,
            is_active=True
        )
        # پیش‌بارگذاری سمت شغلی استاندارد
        InsuranceJobTitle.objects.get_or_create(
            title="کارشناس فنی",
            job_code="027212",
            is_active=True
        )

    def test_job_titles_extended_endpoint(self):
        self.client.force_authenticate(user=self.operator)
        res = self.client.get('/api/personnel/profiles/job-titles/')
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn('job_titles', data)
        self.assertIn('items', data)
        self.assertIn('کارشناس فنی', data['job_titles'])
        
        # بررسی وجود items با کدهای ۶ رقمی بیمه
        tech_item = next((item for item in data['items'] if item['title'] == 'کارشناس فنی'), None)
        self.assertIsNotNone(tech_item)
        self.assertIn('027212', tech_item['job_codes'])

        # تست جستجو
        res_search = self.client.get('/api/personnel/profiles/job-titles/?q=027212')
        self.assertEqual(res_search.status_code, 200)
        search_data = res_search.json()
        self.assertIn('کارشناس فنی', search_data['job_titles'])

    def test_download_template_and_export_excel_46_columns(self):
        self.client.force_authenticate(user=self.manager)
        
        # ۱. تست دانلود قالب ۴۶ ستونی
        res_tpl = self.client.get('/api/personnel/profiles/download-template/')
        self.assertEqual(res_tpl.status_code, 200)
        wb_tpl = openpyxl.load_workbook(io.BytesIO(res_tpl.content), data_only=True)
        ws_tpl = wb_tpl['Emp_info']
        self.assertEqual(ws_tpl.max_column, 46)
        self.assertEqual(ws_tpl.freeze_panes, 'A3')
        self.assertEqual(ws_tpl.cell(row=1, column=4).value, 'کد ملی')
        self.assertEqual(ws_tpl.cell(row=2, column=4).value, 'national_code')
        self.assertEqual(ws_tpl.cell(row=3, column=4).value, '0010376488')

        # ۲. ایجاد یک پرسنل برای اکسپورت
        profile = PersonnelProfile.objects.create(
            company=self.company,
            section=self.section,
            national_code="0010376488",
            first_name="علی",
            last_name="محمدی",
            job_title="کارشناس فنی",
            job_code="027212",
            birth_date="13700101",
            daily_base_wage=3500000,
            account_number="1234567890",
            sheba_number="IR120120000000001234567890",
            approval_status="approved",
            is_active=True
        )

        # ۳. تست خروجی اکسل
        res_exp = self.client.get('/api/personnel/profiles/export-excel/')
        self.assertEqual(res_exp.status_code, 200)
        wb_exp = openpyxl.load_workbook(io.BytesIO(res_exp.content), data_only=True)
        ws_exp = wb_exp['Emp_info']
        self.assertEqual(ws_exp.max_column, 46)
        self.assertEqual(ws_exp.freeze_panes, 'A3')
        self.assertEqual(ws_exp.cell(row=1, column=4).value, 'کد ملی')
        self.assertEqual(ws_exp.cell(row=2, column=4).value, 'national_code')
        # بررسی فرمت نرمال‌شده تاریخ در سطر داده‌ها
        self.assertEqual(ws_exp.cell(row=3, column=8).value, '1370/01/01')
        # بررسی کد شغل ۶ رقمی
        self.assertEqual(ws_exp.cell(row=3, column=26).value, '027212')

    def test_bank_accounts_nested_endpoint_and_sod(self):
        # ایجاد پرسنل در وضعیت draft
        profile = PersonnelProfile.objects.create(
            company=self.company,
            section=self.section,
            national_code="0010376488",
            first_name="حساب",
            last_name="بانکی",
            approval_status="draft",
            is_active=True
        )

        self.client.force_authenticate(user=self.operator)

        # ۱. افزودن حساب بانکی جدید
        from personnel.sheba_utils import generate_sheba_from_account
        valid_sheba = generate_sheba_from_account("017", "101111111001")
        payload = {
            "bank_name": "بانک ملی ایران",
            "account_number": "101111111001",
            "sheba_number": valid_sheba,
            "is_primary": True
        }
        res_post = self.client.post(f"/api/personnel/profiles/{profile.id}/bank-accounts/", payload, format="json")
        self.assertEqual(res_post.status_code, 201, res_post.data)
        acc_id = res_post.data['id']
        self.assertTrue(PersonnelBankAccount.objects.filter(id=acc_id, is_primary=True).exists())

        # ۲. دریافت فهرست حساب‌ها
        res_get = self.client.get(f"/api/personnel/profiles/{profile.id}/bank-accounts/")
        self.assertEqual(res_get.status_code, 200)
        self.assertEqual(len(res_get.data), 1)

        # ۳. ویرایش حساب
        res_put = self.client.put(f"/api/personnel/profiles/{profile.id}/bank-accounts/", {
            "id": acc_id,
            "bank_name": "بانک تجارت"
        }, format="json")
        self.assertEqual(res_put.status_code, 200)
        self.assertEqual(PersonnelBankAccount.objects.get(id=acc_id).bank_name, "بانک تجارت")

        # ۴. تست مسدودسازی SoD روی پرونده approved
        profile.approval_status = "approved"
        profile.save()

        res_blocked_post = self.client.post(f"/api/personnel/profiles/{profile.id}/bank-accounts/", payload, format="json")
        self.assertEqual(res_blocked_post.status_code, 400)
        self.assertIn("ویرایش یا تغییر حساب‌های بانکی برای پرونده مصوب مجاز نیست", res_blocked_post.json().get('error', ''))

        res_blocked_del = self.client.delete(f"/api/personnel/profiles/{profile.id}/bank-accounts/", {"id": acc_id}, format="json")
        self.assertEqual(res_blocked_del.status_code, 400)
