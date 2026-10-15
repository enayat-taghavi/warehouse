from rest_framework import viewsets, permissions, status, parsers, serializers
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied
from django.http import HttpResponse
from django.db import transaction
from django.apps import apps
from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.contrib.auth import get_user_model
import io
import zipfile
import openpyxl
from decimal import Decimal
from common.date_utils import format_to_shamsi_str, normalize_digits
import jdatetime

User = get_user_model()

from accounts.permissions import (
    CanApprovePersonnelManager,
    CanApprovePersonnelFinance,
    CanApproveFleetManager,
    CanApproveFleetFinance
)
from accounts.views import OptionalPageNumberPagination

from .models import (
    Company,
    CompanyDocument,
    CompanyBankAccount,
    CompanyBoardMember,
    CompanyFiscalPeriod,
    UserCompanyAccess,
    FinancialProject,
    ProjectSection,
    UserSectionAssignment,
    Counterparty,
    ExpenseInvoice,
    PettyCashAccount,
    PettyCashTransaction,
    PersonnelProfile,
    VehicleDriverProfile,
    PersonnelChangeRequest,
    VehicleChangeRequest,
    MonthlyWorkPeriod,
    DailyAttendance,
    AttendanceAuditLog,
    VehicleTripLog,
    VehicleTripAuditLog,
    PayrollYearlySettings,
    JobGradeTier,
    WorkshopInsuranceSettings,
    TaxRuleSettings,
    BankExportSettings,
    MonthlyPayrollRecord,
    WorkflowAuditLog,
    PersonnelSectionAssignment,
    PersonnelBankAccount,
    InsuranceJobTitle,
    PersonnelDocument
)
from .serializers import (
    PersonnelBankAccountSerializer,
    PersonnelDocumentSerializer,
    normalize_plate,
    CompanySerializer,
    CompanyDocumentSerializer,
    CompanyBankAccountSerializer,
    CompanyBoardMemberSerializer,
    CompanyFiscalPeriodSerializer,
    UserCompanyAccessSerializer,
    FinancialProjectSerializer,
    ProjectSectionSerializer,
    UserSectionAssignmentSerializer,
    CounterpartySerializer,
    ExpenseInvoiceSerializer,
    PettyCashAccountSerializer,
    PettyCashTransactionSerializer,
    PersonnelProfileSerializer,
    VehicleDriverProfileSerializer,
    PersonnelChangeRequestSerializer,
    VehicleChangeRequestSerializer,
    DailyAttendanceSerializer,
    BulkAttendanceMatrixSerializer,
    BulkMonthlyAttendanceGridSerializer,
    MonthlyGridItemInputSerializer,
    AttendanceAuditLogSerializer,
    VehicleTripLogSerializer,
    VehicleTripAuditLogSerializer,
    BulkVehicleTripMatrixSerializer,
    VehicleDayTripUpdateSerializer,
    VehicleMonthlyGridItemInputSerializer,
    BulkVehicleMonthlyGridSerializer,
    MonthlyWorkPeriodSerializer,
    PayrollYearlySettingsSerializer,
    JobGradeTierSerializer,
    WorkshopInsuranceSettingsSerializer,
    TaxRuleSettingsSerializer,
    BankExportSettingsSerializer,
    MonthlyPayrollRecordSerializer,
    WorkflowAuditLogSerializer
)
from .permissions import (
    IsOrgStructureManagerOrReadOnly,
    IsProjectStructureManagerOrReadOnly,
    IsCounterpartyManagerOrReadOnly,
    IsPayrollSettingsManagerOrReadOnly,
    requires_company_scope,
    has_structure_write_permission,
)
from .payroll_engine import calculate_monthly_payroll_for_period, get_effective_payroll_settings
from .dbf_generator import generate_dskkar_bytes, generate_dskwor_bytes
from .tax_bank_exporter import generate_wh_tax_content, generate_wp_tax_content, generate_bank_meli_excel
from .payroll_excel_exporter import generate_monthly_payroll_excel
from .fleet_excel_engine import export_fleet_monthly_excel, import_fleet_monthly_excel
from .fleet_settlement_engine import calculate_monthly_fleet_settlement, generate_fleet_bank_meli_excel


# قواعد دسترسی شرکت در هستهٔ companies هستند و این‌جا دوباره صادر می‌شوند؛ دسترسی مشتق از عضویت بخش
# در personnel.apps ثبت می‌شود (personnel.company_access).
from companies.realtime import broadcast_org_structure_updated as core_broadcast_org_structure_updated  # noqa: E402
from companies.views import (  # noqa: E402,F401
    CompanyViewSet as CoreCompanyViewSet,
    UserCompanyAccessViewSet,
)
from companies.access import (  # noqa: E402,F401
    get_user_allowed_companies,
    get_request_company_id,
    validate_user_company_access,
    ensure_scoped_write,
)


def user_belongs_to_company(user_id, company_id):
    """
    آیا کاربر عضو شرکت است یا بوده است؟ (دسترسی صریح workspace_full، یا هر انتساب فعال یا غیرفعال در بخش‌های همان شرکت)
    انتساب بخش هرگز نباید برای کاربر دسترسی شرکت تازه بسازد، چون عضویت بخش خودش دسترسی شرکت می‌دهد.
    انتساب غیرفعالِ قبلی حساب می‌شود تا فعال‌سازی دوباره یا جابه‌جایی درون همان شرکت رد نشود.
    """
    if not user_id or not company_id:
        return False
    return (
        UserCompanyAccess.objects.filter(user_id=user_id, company_id=company_id, access_level='workspace_full').exists()
        or UserSectionAssignment.objects.filter(user_id=user_id, section__project__company_id=company_id).exists()
    )


def ensure_scoped_assignment(actor, target_user_id, section, final_active=True):
    """
    انتساب بخش: بخش باید در شرکت‌های مجاز نویسنده باشد و کاربر هدف عضو همان شرکت.
    راهبر تشکیلاتی که مجوز قدیمی نوشتن ساختار ندارد، نمی‌تواند خودش را منتسب کند (خط قرمز نقش راهبر).
    غیرفعال‌سازی انتساب این دو شرط عضویت را لازم ندارد.
    """
    company_id = section.project.company_id if section is not None and section.project_id else None
    ensure_scoped_write(actor, company_id, "انتساب فقط در بخش‌های شرکت‌های مجاز شما ممکن است (بخش پروژهٔ سراسری مجاز نیست).")
    if not final_active:
        return
    if target_user_id == actor.id and not has_structure_write_permission(actor):
        raise PermissionDenied("انتساب خودتان به بخش را مدیر ارشد یا مدیر انجام می‌دهد، نه راهبر تشکیلات.")
    if not user_belongs_to_company(target_user_id, company_id):
        raise PermissionDenied("این کاربر عضو شرکتِ این بخش نیست. دسترسی شرکت را مدیر ارشد می‌دهد، نه انتساب بخش.")


def resolve_write_company(request):
    """
    شرکت مقصد درون‌ریزی اکسل برای کاربر غیرسوپریوزر: شرکت فعال درخواست (پارامتر company_id یا هدر X-Company-ID)
    که باید جزو شرکت‌های مجاز او باشد. اگر شرکتی مشخص نشده و کاربر فقط یک شرکت مجاز دارد، همان برگردانده می‌شود.
    """
    allowed_ids = set(get_user_allowed_companies(request.user).values_list('id', flat=True))
    raw = request.data.get('company_id') if hasattr(request.data, 'get') else None
    raw = raw or get_request_company_id(request)
    if raw:
        try:
            cid = int(raw)
        except (ValueError, TypeError):
            raise serializers.ValidationError({'error': 'شناسهٔ شرکت نامعتبر است.'})
        if cid not in allowed_ids:
            raise PermissionDenied("شما به این شرکت دسترسی ندارید.")
        return cid
    if len(allowed_ids) == 1:
        return next(iter(allowed_ids))
    raise serializers.ValidationError({'error': 'شرکت فعال مشخص نیست؛ ابتدا شرکت را در سربرگ انتخاب کنید.'})


def resolve_import_company(request):
    """
    شرکت مقصد درون‌ریزی اکسل: برای غیرسوپریوزر resolve_write_company (اجباری و مجاز).
    برای سوپریوزر شرکت فعال درخواست اگر آمده باشد، وگرنه تهی (رفتار قبلی: تطبیق روی کل جدول).
    """
    if requires_company_scope(request.user):
        return resolve_write_company(request)
    raw = request.data.get('company_id') if hasattr(request.data, 'get') else None
    raw = raw or get_request_company_id(request)
    try:
        return int(raw) if raw else None
    except (ValueError, TypeError):
        return None


def readable_company_ids(user):
    """شرکت‌هایی که خروجی اکسل کاربر غیرسوپریوزر به آن‌ها محدود می‌شود (به‌علاوهٔ رکوردهای سراسری). تهی یعنی همه."""
    if user.is_superuser:
        return None
    return set(get_user_allowed_companies(user).values_list('id', flat=True))


def is_personnel_global_auditor(user):
    """نقش‌های مدیر و حسابدار کل که به بخش‌های منتسب محدود نیستند (ولی همچنان به شرکت‌های مجاز خود محدودند)."""
    return bool(
        user and (
            user.is_superuser
            or user.has_perm('accounts.perm_approve_personnel_manager')
            or user.has_perm('accounts.can_act_as_manager')
            or user.has_perm('accounts.perm_approve_personnel_finance')
        )
    )


def scope_personnel_queryset(request, qs):
    """
    قلمرو امنیتی پرونده‌های پرسنل (Guardian G1 / BOLA): کاربر عادی فقط بخش‌های منتسب فعال یا پرونده‌های ثبت خودش،
    و هر کاربر غیرسوپریوزر فقط شرکت‌های مجاز خودش؛ چه سربرگ شرکت بیاید چه نیاید.
    پیش‌تر نبودِ سربرگ X-Company-ID قلمرو شرکت را برای مدیر و حسابدار کل برمی‌داشت.
    """
    user = request.user
    if not user or not user.is_authenticated:
        return qs.none()
    cid = validate_user_company_access(user, get_request_company_id(request))
    return scope_personnel_for_user(user, qs, company_id=cid)


def scope_personnel_for_user(user, qs, company_id=None):
    """
    همان قاعدهٔ scope_personnel_queryset بدون request (برای گزارش‌ساز): company_id تهی یعنی همهٔ شرکت‌های مجاز.
    """
    if not user or not user.is_authenticated:
        return qs.none()
    if not is_personnel_global_auditor(user):
        user_section_ids = list(
            UserSectionAssignment.objects.filter(user=user, is_active=True).values_list('section_id', flat=True)
        )
        qs = qs.filter(Q(section_id__in=user_section_ids) | Q(created_by=user))

    cid = company_id
    if cid:
        return qs.filter(Q(company_id=cid) | Q(section__project__company_id=cid))
    if user.is_superuser:
        return qs
    allowed_ids = list(get_user_allowed_companies(user).values_list('id', flat=True))
    return qs.filter(
        Q(company_id__in=allowed_ids)
        | Q(section__project__company_id__in=allowed_ids)
        # پیش‌نویسِ بی‌شرکت و بی‌بخش فقط برای ثبت‌کننده‌اش
        | Q(company__isnull=True, section__isnull=True, created_by=user)
    )


def ensure_record_write_scope(user, company=None, section=None, project=None, label='رکورد'):
    """
    نوشتن رکورد فقط درون شرکت‌های مجاز کاربر (تصمیم کاربر، ۹ مهر ۱۴۰۵؛ requires_company_scope):
    شرکت، بخش و پروژهٔ پرونده نباید به شرکتی بیرون از قلمرو کاربر اشاره کنند. در صورت نقض PermissionDenied.
    """
    if not requires_company_scope(user):
        return
    allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
    company_ids = set()
    if company is not None:
        company_ids.add(getattr(company, 'pk', company))
    if section is not None:
        section_company_id = section.project.company_id if section.project_id else None
        if section_company_id is None:
            raise PermissionDenied(f"بخش انتخاب‌شده به شرکتی تعلق ندارد و ثبت {label} در آن مجاز نیست.")
        company_ids.add(section_company_id)
    if project is not None:
        company_ids.add(project.company_id)
    if any(cid is None or cid not in allowed_ids for cid in company_ids):
        raise PermissionDenied(f"شرکت، بخش یا پروژهٔ انتخاب‌شده برای {label} در قلمرو شرکت‌های مجاز شما نیست.")
    if len(company_ids) > 1:
        raise PermissionDenied(f"شرکت، بخش و پروژهٔ {label} باید به یک شرکت تعلق داشته باشند.")


def ensure_personnel_write_scope(user, company=None, section=None, project=None):
    return ensure_record_write_scope(user, company=company, section=section, project=project, label='پرسنل')


def is_vehicle_global_auditor(user):
    return bool(
        user and (
            user.is_superuser
            or user.has_perm('accounts.perm_approve_fleet_manager')
            or user.has_perm('accounts.can_act_as_manager')
            or user.has_perm('accounts.perm_approve_fleet_finance')
        )
    )


def scope_vehicle_queryset(request, qs, prefix=''):
    """
    قلمرو امنیتی پرونده‌های ناوگان و خودروها (ب-۲): کاربر عادی فقط بخش‌های منتسب فعال یا پرونده‌های ثبت خودش،
    و هر کاربر غیرسوپریوزر فقط شرکت‌های مجاز خودش؛ چه سربرگ شرکت بیاید چه نیاید.
    """
    user = request.user
    if not user or not user.is_authenticated:
        return qs.none()

    if not is_vehicle_global_auditor(user):
        user_section_ids = list(
            UserSectionAssignment.objects.filter(user=user, is_active=True).values_list('section_id', flat=True)
        )
        sec_field = f"{prefix}section_id__in"
        creator_field = f"{prefix}created_by" if prefix == '' else "requested_by"
        qs = qs.filter(Q(**{sec_field: user_section_ids}) | Q(**{creator_field: user}))

    cid = validate_user_company_access(user, get_request_company_id(request))
    comp_field = f"{prefix}company_id"
    sec_comp_field = f"{prefix}section__project__company_id"

    if cid:
        return qs.filter(Q(**{comp_field: cid}) | Q(**{sec_comp_field: cid}))
    if user.is_superuser:
        return qs
    allowed_ids = list(get_user_allowed_companies(user).values_list('id', flat=True))
    creator_field = f"{prefix}created_by" if prefix == '' else "requested_by"
    return qs.filter(
        Q(**{f"{comp_field}__in": allowed_ids})
        | Q(**{f"{sec_comp_field}__in": allowed_ids})
        | Q(**{f"{prefix}company__isnull": True, f"{prefix}section__isnull": True, creator_field: user})
    )


def scope_trips_for_user(user, qs, company_id=None):
    """
    قلمرو امنیتی سفرهای ناوگان:
    - کاربر ناشناس: خالی
    - با company_id: سفرهایی که بخش، پروژه، خودرو یا بخشِ خودرو متعلق به آن شرکت باشد
    - بی company_id و سوپریوزر: بدون محدودیت
    - بی company_id و غیرسوپریوزر: یکی از همان چهار مسیر در شرکت‌های مجاز یا قاعدهٔ خودروی یتیم سازنده (بی‌شرکت و بی‌بخش)
    """
    if not user or not user.is_authenticated:
        return qs.none()

    if company_id:
        return qs.filter(
            Q(section__project__company_id=company_id)
            | Q(project__company_id=company_id)
            | Q(vehicle__company_id=company_id)
            | Q(vehicle__section__project__company_id=company_id)
        ).distinct()

    if user.is_superuser:
        return qs

    try:
        from companies.access import get_user_allowed_companies
        allowed_ids = list(get_user_allowed_companies(user).values_list('id', flat=True))
    except Exception:
        allowed_ids = []

    orphan_q = Q(vehicle__company__isnull=True, vehicle__section__isnull=True, vehicle__created_by=user)
    if allowed_ids:
        return qs.filter(
            Q(section__project__company_id__in=allowed_ids)
            | Q(project__company_id__in=allowed_ids)
            | Q(vehicle__company_id__in=allowed_ids)
            | Q(vehicle__section__project__company_id__in=allowed_ids)
            | orphan_q
        ).distinct()

    return qs.filter(orphan_q).distinct()


def broadcast_personnel_update(personnel, action_type='updated', message=None, sender_id=None, client_tab_id=None):
    """
    تسک ۴.۳: ارسال بلادرنگ رویدادهای تغییر وضعیت یا ویرایش پرونده پرسنل به کانال اختصاصی شرکت مربوطه.
    توقف پخش سراسری رویداد پرسنل بی‌شرکت به global_notifications، و غنی‌سازی پی‌لود جهت به‌روزرسانی درجای جدول و شمارنده‌ها.
    """
    try:
        from channels.layers import get_channel_layer
        from asgiref.sync import async_to_sync
        channel_layer = get_channel_layer()
        if channel_layer is not None:
            full_name = getattr(personnel, 'full_name', '') or f"{getattr(personnel, 'first_name', '')} {getattr(personnel, 'last_name', '')}".strip()
            cid = getattr(personnel, 'company_id', None)
            if not cid and hasattr(personnel, 'section') and personnel.section and personnel.section.project:
                cid = personnel.section.project.company_id

            if not cid:
                # طبق الزامات تسک ۴.۳: رویداد بدون شرکت دیگر نباید سراسری پخش شود تا از نشت دیتا جلوگیری گردد
                import logging
                logging.getLogger(__name__).debug(f"[WebSocket] Skipping global broadcast for personnel without company: {full_name}")
                return

            payload = {
                'type': 'send_notification',
                'type_str': 'personnel_updated',
                'action': action_type,
                'personnel_id': getattr(personnel, 'id', None),
                'company_id': cid,
                'section_id': getattr(personnel, 'section_id', None),
                'project_id': getattr(personnel, 'project_id', None),
                'full_name': full_name,
                # کد ملی در پخش گروهی شرکت فرستاده نمی‌شود (حریم خصوصی؛ آزمون test_websocket_broadcast_isolation_and_privacy)
                'job_title': getattr(personnel, 'job_title', ''),
                'approval_status': getattr(personnel, 'approval_status', ''),
                'is_active': getattr(personnel, 'is_active', True),
                'has_pending_changes': getattr(personnel, 'has_pending_changes', False),
                'valid_until_date': getattr(personnel, 'valid_until_date', '') or '',
                'rejection_reason': getattr(personnel, 'rejection_reason', '') or '',
                'message': message or f'پرونده پرسنل «{full_name}» به‌روزرسانی شد.',
                'sender_id': sender_id,
                'client_tab_id': client_tab_id,
            }
            target_group = f'company_{cid}_notifications'
            async_to_sync(channel_layer.group_send)(
                target_group,
                payload
            )
    except Exception as e:
        import logging
        logging.getLogger(__name__).warning(f"[WebSocket] Error broadcasting personnel_updated: {e}")


def broadcast_vehicle_update(vehicle, action_type='updated', message=None, sender_id=None, client_tab_id=None):
    """
    تسک ۴.۳ و ب-۶: ارسال بلادرنگ رویدادهای تغییر وضعیت، ثبت یا ویرایش پرونده ناوگان با پی‌لود غنی اختصاصی کانال شرکت
    """
    try:
        from channels.layers import get_channel_layer
        from asgiref.sync import async_to_sync
        channel_layer = get_channel_layer()
        if channel_layer is not None:
            plate_number = getattr(vehicle, 'plate_number', '') or ''
            driver_name = getattr(vehicle, 'driver_name', '') or ''
            cid = getattr(vehicle, 'company_id', None)
            if not cid and hasattr(vehicle, 'section') and vehicle.section and vehicle.section.project:
                cid = vehicle.section.project.company_id

            if not cid:
                import logging
                logging.getLogger(__name__).debug(f"[WebSocket] Skipping global broadcast for vehicle without company: {plate_number}")
                return

            proj_id = getattr(vehicle, 'project_id', None)
            if not proj_id and hasattr(vehicle, 'section') and vehicle.section:
                proj_id = getattr(vehicle.section, 'project_id', None)

            payload = {
                'type': 'send_notification',
                'type_str': 'vehicle_updated',
                'action': action_type,
                'vehicle_id': getattr(vehicle, 'id', None),
                'company_id': cid,
                'section_id': getattr(vehicle, 'section_id', None),
                'project_id': proj_id,
                'plate_number': plate_number,
                'driver_name': driver_name,
                # کدهای ملی راننده و مالک جهت رعایت حریم خصوصی در پخش عمومی ارسال نمی‌شوند
                'vehicle_type': getattr(vehicle, 'vehicle_type', ''),
                'ownership_type': getattr(vehicle, 'ownership_type', ''),
                'approval_status': getattr(vehicle, 'approval_status', ''),
                'is_active': getattr(vehicle, 'is_active', True),
                'has_pending_changes': getattr(vehicle, 'has_pending_changes', False),
                'valid_until_date': getattr(vehicle, 'valid_until_date', '') or '',
                'insurance_expiry_date': getattr(vehicle, 'insurance_expiry_date', '') or '',
                'inspection_expiry_date': getattr(vehicle, 'inspection_expiry_date', '') or '',
                'end_date': getattr(vehicle, 'end_date', '') or '',
                'rejection_reason': getattr(vehicle, 'rejection_reason', '') or '',
                'message': message or f'پرونده خودرو «{plate_number}» با راننده «{driver_name}» به‌روزرسانی شد.',
                'sender_id': sender_id,
                'client_tab_id': client_tab_id,
            }
            target_group = f'company_{cid}_notifications'
            async_to_sync(channel_layer.group_send)(
                target_group,
                payload
            )
    except Exception as e:
        import logging
        logging.getLogger(__name__).warning(f"[WebSocket] Error broadcasting vehicle_updated: {e}")


class PersonnelProfileViewSet(viewsets.ModelViewSet):
    queryset = PersonnelProfile.objects.all().select_related('user').prefetch_related('bank_accounts', 'documents')
    serializer_class = PersonnelProfileSerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    pagination_class = None

    def _is_global_auditor(self, user):
        return is_personnel_global_auditor(user)

    def _scoped_base_queryset(self):
        """قلمرو امنیتی کاربر (بخش‌ها و شرکت‌های مجاز) بدون فیلترهای نمایشی مثل جستجو یا بخش انتخابی."""
        return scope_personnel_queryset(self.request, super().get_queryset())

    def get_queryset(self):
        qs = self._scoped_base_queryset()
        user = self.request.user
        is_global_auditor = self._is_global_auditor(user)

        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True))
        
        section_id = self.request.query_params.get('section_id')
        if section_id:
            if not is_global_auditor and user and user.is_authenticated:
                user_section_ids = list(
                    UserSectionAssignment.objects.filter(user=user, is_active=True).values_list('section_id', flat=True)
                )
                if str(section_id).isdigit() and int(section_id) not in user_section_ids:
                    return qs.none()
            qs = qs.filter(
                Q(section_id=section_id) |
                Q(section_assignments__section_id=section_id, section_assignments__is_active=True)
            ).distinct()
        project_id = self.request.query_params.get('project_id')
        if project_id:
            qs = qs.filter(
                Q(section__project_id=project_id) |
                Q(section_assignments__project_id=project_id, section_assignments__is_active=True)
            ).distinct()
        
        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() in ['true', '1'])
            
        approval_status = self.request.query_params.get('approval_status')
        if approval_status:
            if ',' in approval_status:
                status_list = [s.strip() for s in approval_status.split(',') if s.strip()]
                qs = qs.filter(approval_status__in=status_list)
            else:
                qs = qs.filter(approval_status=approval_status)

        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(
                Q(first_name__icontains=search) |
                Q(last_name__icontains=search) |
                Q(national_code__icontains=search) |
                Q(job_title__icontains=search)
            )
        return qs

    def create(self, request, *args, **kwargs):
        data = request.data.copy() if hasattr(request.data, 'copy') else dict(request.data)
        section_id = data.get('section')
        if not data.get('company') and section_id:
            try:
                sec = ProjectSection.objects.filter(id=int(section_id)).select_related('project').first()
                if sec and sec.project and sec.project.company_id:
                    data['company'] = sec.project.company_id
            except (ValueError, TypeError):
                pass
        serializer = self.get_serializer(data=data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        headers = self.get_success_headers(serializer.data)
        return Response(serializer.data, status=status.HTTP_201_CREATED, headers=headers)

    def perform_create(self, serializer):
        user = self.request.user
        is_mgr = user.is_superuser or user.has_perm('accounts.perm_approve_personnel_manager') or user.has_perm('accounts.can_act_as_manager')
        
        section = serializer.validated_data.get('section')
        # گارد انتساب بخش: کاربر عادی باید انتساب فعال به بخش انتخاب‌شده داشته باشد
        if not is_mgr and section and user and user.is_authenticated:
            is_assigned = UserSectionAssignment.objects.filter(user=user, section=section, is_active=True).exists()
            if not is_assigned:
                raise PermissionDenied("شما انتساب فعال در این بخش برای معرفی پرسنل جدید ندارید.")
        # شرکت، بخش و پروژهٔ پرونده باید در قلمرو شرکت‌های مجاز کاربر باشند (حتی برای مدیر؛ عبور هوشمند هم پس از این است)
        ensure_personnel_write_scope(
            user,
            company=serializer.validated_data.get('company'),
            section=section,
            project=serializer.validated_data.get('project'),
        )

        extra_kwargs = {'created_by': user}
        if section and hasattr(section, 'project') and section.project:
            if not serializer.validated_data.get('project'):
                extra_kwargs['project'] = section.project
            if not serializer.validated_data.get('company') and section.project.company_id:
                extra_kwargs['company_id'] = section.project.company_id

        # DEF-02: پرونده همیشه پیش‌نویس ساخته می‌شود و فقط اکشن submit آن را با گیت به سرپرست می‌فرستد.
        # تنها استثنا عبور هوشمند نقش‌های بالاتر است که آن هم با گیت انجام می‌شود (DEF-14).
        req_status = self.request.data.get('approval_status')
        extra_kwargs['approval_status'] = 'draft'
        instance = serializer.save(**extra_kwargs)
        if req_status != 'draft':
            from .lifecycle_services import auto_pass_personnel
            auto_pass_personnel(instance, user, source='create')
        broadcast_personnel_update(instance, action_type='created', message=f'پرونده پرسنل «{instance.full_name}» ثبت شد.', sender_id=user.id if user else None)

    def _ensure_update_scope(self, instance, validated_data):
        """جابه‌جایی شرکت، بخش یا پروژهٔ پرونده فقط درون شرکت‌های مجاز کاربر (پرونده به شرکت بیگانه منتقل نمی‌شود)."""
        def _changed(field):
            return field in validated_data and getattr(validated_data[field], 'pk', None) != getattr(instance, f'{field}_id')

        if not any(_changed(f) for f in ('company', 'section', 'project')):
            return
        ensure_personnel_write_scope(
            self.request.user,
            company=validated_data.get('company', instance.company),
            section=validated_data.get('section', instance.section),
            project=validated_data['project'] if _changed('project') else None,
        )

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        user = request.user
        from .lifecycle_services import (
            FINANCE_STAGE_PERMS,
            change_request_fields_allowed_for,
            compute_profile_diff,
            user_has_any_perm,
        )
        is_fin = user_has_any_perm(user, FINANCE_STAGE_PERMS)
        current_status = instance.approval_status
        audit_excluded = {'id', 'pk', 'created_at', 'created_by', 'updated_at'}

        # اجازه ویرایش مستقیم به حسابدار در pending_accountant (تسک ۳.۷)؛ هر فیلد تغییرکرده با diff لاگ می‌شود
        if current_status == 'pending_accountant':
            if not is_fin:
                return Response(
                    {'error': 'این پرونده در وضعیت «در انتظار حسابداری» قرار دارد و تنها حسابدار مجاز به ویرایش اطلاعات آن است.'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            serializer = self.get_serializer(instance, data=request.data, partial=partial)
            serializer.is_valid(raise_exception=True)
            self._ensure_update_scope(instance, serializer.validated_data)
            diff, old_diff = compute_profile_diff(instance, serializer.validated_data, excluded=audit_excluded)
            with transaction.atomic():
                updated_instance = serializer.save()
                # کش prefetch حساب‌ها و مدارک باطل می‌شود (همان کار UpdateModelMixin در DRF)
                updated_instance._prefetched_objects_cache = {}
                if diff:
                    WorkflowAuditLog.objects.create(
                        content_type='personnel',
                        object_id=updated_instance.id,
                        actor=user,
                        from_status='pending_accountant',
                        to_status='pending_accountant',
                        action='accountant_edit',
                        reason=request.data.get('note', '') or 'ویرایش اطلاعات پرسنلی و مالی توسط حسابدار',
                        metadata={'diff': diff, 'old_diff': old_diff, 'company_id': updated_instance.company_id}
                    )
            broadcast_personnel_update(updated_instance, action_type='updated', message=f'اطلاعات مالی پرسنل «{updated_instance.full_name}» توسط حسابدار به‌روزرسانی شد.', sender_id=user.id if user else None)
            return Response(self.get_serializer(updated_instance).data)

        # DEF-17: پروندهٔ در انتظار مدیر برای همه جز سوپریوزر قفل است؛ مدیر برای اصلاح به حسابدار عودت می‌دهد
        if current_status in ('pending_manager', 'accountant_approved') and not user.is_superuser:
            return Response(
                {'error': f'این پرونده در وضعیت «{instance.get_approval_status_display()}» قرار دارد و اطلاعات آن تا پایان گردش کار قابل تغییر مستقیم نیست. برای اصلاح، پرونده را به حسابدار عودت دهید.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # پروندهٔ ردشده بایگانی است و پیش از هر ویرایش باید با اکشن reopen بازگشایی شود (DEF-15)
        if current_status == 'rejected' and not user.is_superuser:
            return Response(
                {'error': 'این پرونده رد شده است. برای ویرایش و ارسال دوباره، ابتدا آن را بازگشایی کنید.', 'code': 'reopen_required'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # DEF-17: پروندهٔ مصوب برای همه جز سوپریوزر فقط از مسیر درخواست تغییرات ویرایش می‌شود (تسک ۳.۴-الف).
        # DEF-16: ابتدا diff محاسبه می‌شود و فهرست سفید فقط روی کلیدهای واقعاً تغییرکرده اعمال می‌شود.
        if current_status in ('approved', 'manager_approved') and not user.is_superuser:
            # SoD: تلاش صریح نقش غیرمالی برای تغییر دستمزد پروندهٔ مصوب رد می‌شود (نه حذف بی‌صدا)
            if change_request_fields_allowed_for(user) is not None:
                from decimal import Decimal, InvalidOperation
                from .lifecycle_services import WAGE_FIELDS

                def _same_amount(new, old):
                    try:
                        return Decimal(str(new or 0)) == Decimal(str(old or 0))
                    except (InvalidOperation, ValueError):
                        return False

                changed_wages = [
                    f for f in WAGE_FIELDS
                    if f in request.data and hasattr(instance, f) and not _same_amount(request.data.get(f), getattr(instance, f, None))
                ]
                if changed_wages:
                    raise PermissionDenied(f"تعیین دستمزد و حقوق منحصراً با حسابداری است و فیلد «{changed_wages[0]}» برای شما قفل است.")
            serializer = self.get_serializer(instance, data=request.data, partial=partial)
            serializer.is_valid(raise_exception=True)
            diff, old_diff = compute_profile_diff(instance, serializer.validated_data)

            allowed_fields = change_request_fields_allowed_for(user)
            if allowed_fields is not None:
                locked = [k for k in diff if k not in allowed_fields]
                if locked:
                    raise PermissionDenied(f"کارکنان عادی صرفاً مجاز به پیشنهاد تغییر در مشخصات فردی، تماس و اطلاعات بانکی هستند و فیلد «{locked[0]}» قفل است.")

            if diff:
                with transaction.atomic():
                    cr = PersonnelChangeRequest.objects.create(
                        personnel=instance,
                        requested_by=user,
                        proposed_changes=diff,
                        previous_values=old_diff,
                        status='pending_supervisor'
                    )
                    instance.has_pending_changes = True
                    instance.save(update_fields=['has_pending_changes'])
                    WorkflowAuditLog.objects.create(
                        content_type='personnel',
                        object_id=instance.id,
                        actor=user,
                        from_status=instance.approval_status,
                        to_status=instance.approval_status,
                        action='change_request_created',
                        reason=request.data.get('note', '') or 'ثبت درخواست تغییرات پرسنل فعال',
                        metadata={'change_request_id': cr.id, 'diff': diff, 'company_id': instance.company_id}
                    )

                return Response({
                    'message': 'درخواست تغییرات با موفقیت ثبت شد و جهت بررسی در کارتابل سرپرست قرار گرفت. اطلاعات قبلی تا زمان تصویب نهایی مدیر معتبر باقی می‌ماند.',
                    'change_request_id': cr.id,
                    'data': self.get_serializer(instance).data
                }, status=status.HTTP_202_ACCEPTED)
            return Response(self.get_serializer(instance).data)

        # ویرایش رکوردهای پیش‌نویس، عودت‌داده‌شده یا در انتظار سرپرست؛ و ویرایش مستقیم سوپریوزر
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self._ensure_update_scope(instance, serializer.validated_data)

        extra_kwargs = {}
        section = serializer.validated_data.get('section', instance.section)
        if section and hasattr(section, 'project') and section.project:
            if not serializer.validated_data.get('project') and not instance.project_id:
                extra_kwargs['project'] = section.project

        # ویرایش مستقیم سوپریوزر روی پروندهٔ در جریان مدیر، مصوب یا ردشده بی‌ردپا نمی‌ماند
        superuser_diff = None
        if user.is_superuser and current_status not in ('draft', 'revision_required', 'pending_supervisor'):
            superuser_diff = compute_profile_diff(instance, serializer.validated_data, excluded=audit_excluded)

        with transaction.atomic():
            updated_instance = serializer.save(**extra_kwargs)
            updated_instance._prefetched_objects_cache = {}
            if superuser_diff and superuser_diff[0]:
                WorkflowAuditLog.objects.create(
                    content_type='personnel',
                    object_id=updated_instance.id,
                    actor=user,
                    from_status=current_status,
                    to_status=updated_instance.approval_status,
                    action='superuser_edit',
                    reason=request.data.get('note', '') or 'ویرایش مستقیم پرونده توسط مدیر سیستم',
                    metadata={'diff': superuser_diff[0], 'old_diff': superuser_diff[1], 'company_id': updated_instance.company_id}
                )
        broadcast_personnel_update(updated_instance, action_type='updated', message=f'پرونده پرسنل «{updated_instance.full_name}» به‌روزرسانی شد.', sender_id=user.id if user else None)
        return Response(self.get_serializer(updated_instance).data)

    @action(detail=True, methods=['post'], url_path='submit')
    def submit(self, request, pk=None):
        """
        تسک ۲.۷: ارسال پرونده پیش‌نویس یا عودت‌داده‌شده به سرپرست با گیت اعتبارسنجی
        """
        personnel = self.get_object()
        user = request.user

        if personnel.approval_status not in ('draft', 'revision_required'):
            return Response({
                'error': f'امکان ارسال پرونده در وضعیت «{personnel.get_approval_status_display()}» وجود ندارد.'
            }, status=status.HTTP_400_BAD_REQUEST)

        # گیت سروری فیلدهای الزامی مرحله کارمند (تسک ۲.۶)
        from .lifecycle_services import get_missing_fields
        missing = get_missing_fields(personnel, stage='submit')
        if missing:
            labels = [m['label'] for m in missing]
            return Response({
                'error': 'لطفاً فیلدهای الزامی پرونده را پیش از ارسال تکمیل نمایید.',
                'missing_fields': missing,
                'detail': f"فیلدهای ناقص: {', '.join(labels)}"
            }, status=status.HTTP_400_BAD_REQUEST)

        from_status = personnel.approval_status
        to_status = 'pending_supervisor'

        with transaction.atomic():
            personnel.approval_status = to_status
            personnel.rejection_reason = None
            personnel.save(update_fields=['approval_status', 'rejection_reason', 'updated_at'])

            WorkflowAuditLog.objects.create(
                content_type='personnel',
                object_id=personnel.id,
                action='submit',
                from_status=from_status,
                to_status=to_status,
                actor=user,
                reason=request.data.get('note', '') or 'ارسال پرونده توسط ثبت‌کننده به کارتابل سرپرست'
            )

        from .lifecycle_services import dispatch_workflow_notification
        dispatch_workflow_notification(personnel, stage='pending_supervisor', actor=user, note=request.data.get('note', ''))
        broadcast_personnel_update(personnel, action_type='updated', message=f'پرونده پرسنل «{personnel.full_name}» به کارتابل سرپرست ارسال شد.', sender_id=user.id if user else None)
        return Response(self.get_serializer(personnel).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get'], url_path='delete_impact')
    def delete_impact(self, request, pk=None):
        """
        تسک ۲.۵: استعلام اثرات حذف آبشاری و سوابق پرسنل منطبق بر قرارداد DeleteImpactMixin و smart-delete-modal
        """
        instance = self.get_object()
        from django.db.models.deletion import Collector
        from django.db import router
        from .lifecycle_services import check_personnel_history

        using = router.db_for_write(instance.__class__, instance=instance)
        collector = Collector(using=using)
        try:
            collector.collect([instance])
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        impact_summary = []
        for model, instances in collector.data.items():
            if model == instance.__class__:
                continue
            impact_summary.append({
                'model_name': str(model._meta.verbose_name or model.__name__),
                'count': len(instances)
            })
        impact_summary.sort(key=lambda x: x['count'], reverse=True)

        history = check_personnel_history(instance)
        user = request.user
        can_delete = bool(history['can_delete'] and self._user_may_delete(user, instance))

        return Response({
            'impact': impact_summary,
            'total_affected': sum(item['count'] for item in impact_summary),
            'history': history,
            'can_delete': can_delete,
            'has_history': history['has_history'],
            'reasons': history['reasons']
        }, status=status.HTTP_200_OK)

    @staticmethod
    def _user_may_delete(user, instance) -> bool:
        """
        ماتریس حذف تسک ۲.۴ (بدون احتساب سابقه):
        الف) ثبت‌کننده یا هم‌بخشِ غیرمدیر فقط پرونده‌های draft و revision_required را حذف می‌کند؛
        ب) دارندهٔ مجوز اختصاصی perm_delete_personnel_nomovement (تصمیم ۶.۸) هر پروندهٔ بدون سابقه را.
        """
        if not user or not user.is_authenticated:
            return False
        if user.is_superuser or user.has_perm('accounts.perm_delete_personnel_nomovement'):
            return True
        if instance.approval_status not in ('draft', 'revision_required'):
            return False
        return PersonnelProfileViewSet._is_creator_or_co_section(user, instance)

    @staticmethod
    def _is_creator_or_co_section(user, instance) -> bool:
        if instance.created_by_id == user.id:
            return True
        return bool(instance.section_id) and UserSectionAssignment.objects.filter(
            user=user, section_id=instance.section_id, is_active=True
        ).exists()

    def destroy(self, request, *args, **kwargs):
        """
        تسک ۲.۴: حذف پرسنل با ماتریس مجوز و مسدودسازی قطعی رکوردهای دارای سابقه:
        - اگر سابقه کارکرد، فیش، رویداد مالی، تغییرات یا انتساب داشته باشد، حذف فیزیکی مسدود است و قطع همکاری پیشنهاد می‌شود.
        - اگر فاقد سابقه باشد، ثبت‌کننده در draft/revision_required یا کاربر دارای مجوز اختصاصی perm_delete_personnel_nomovement مجاز به حذف فیزیکی است.
        - پیش از حذف، عکس کامل پرونده در WorkflowAuditLog ثبت می‌گردد.
        """
        instance = self.get_object()
        user = request.user
        from .lifecycle_services import build_personnel_snapshot, check_personnel_history

        with transaction.atomic():
            # قفل ردیف پرونده و سنجش سابقه درون همان تراکنش: ثبت هم‌زمان کارکرد یا فیش (کلید خارجی به همین ردیف)
            # تا پایان تراکنش منتظر می‌ماند و سابقه‌ای با CASCADE پاک نمی‌شود (بازبینی از دست رفتن داده)
            instance = PersonnelProfile.objects.select_for_update().get(pk=instance.pk)
            history = check_personnel_history(instance)
            if history['has_history']:
                return Response({
                    'error': 'این پرسنل دارای سابقه کارکرد، فیش حقوقی یا اسناد مالی در سیستم است و حذف فیزیکی آن مجاز نمی‌باشد. لطفاً از طریق ثبت تاریخ پایان کار و غیرفعال‌سازی پرونده اقدام به قطع همکاری نمایید.',
                    'code': 'history_exists',
                    'suggested_action': 'end_collaboration',
                    'history_reasons': history['reasons']
                }, status=status.HTTP_400_BAD_REQUEST)

            if not self._user_may_delete(user, instance):
                if self._is_creator_or_co_section(user, instance):
                    # رفتار فعلی شاخهٔ الف ماتریس ۲.۴: ثبت‌کننده یا هم‌بخش فقط پیش‌نویس و عودت‌شده را حذف می‌کند
                    return Response(
                        {'error': 'فقط پرونده‌های در وضعیت پیش‌نویس یا عودت‌داده‌شده قابل حذف هستند.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
                raise PermissionDenied("شما مجوز حذف این پرونده پرسنلی را ندارید. (نیازمند مجوز اختصاصی حذف پرسنل فاقد سابقه)")

            full_name = instance.full_name
            # DEF-11: عکس کامل پرونده، همهٔ حساب‌ها و فهرست مدارک پیش از حذف (حساب‌ها و مدارک با CASCADE پاک می‌شوند)
            snapshot = build_personnel_snapshot(instance)
            snapshot['deleted_by_user_id'] = user.id if user else None
            snapshot['deleted_by_username'] = user.username if user else None
            WorkflowAuditLog.objects.create(
                content_type='personnel',
                object_id=instance.id,
                action='delete',
                actor=user if (user and user.is_authenticated) else None,
                from_status=instance.approval_status,
                to_status='deleted',
                metadata={'snapshot': snapshot, 'company_id': instance.company_id, 'reason': request.data.get('reason', 'حذف فیزیکی پرونده فاقد سابقه')}
            )
            deleted_id = instance.id
            instance.delete()
            instance.id = deleted_id

        broadcast_personnel_update(instance, action_type='deleted', message=f'پرونده پرسنل «{full_name}» حذف شد.', sender_id=user.id if user else None)
        return Response({'message': f'پرسنل «{full_name}» با موفقیت از سامانه حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get', 'post', 'put', 'patch', 'delete'], url_path='bank-accounts')
    def bank_accounts(self, request, pk=None):
        """
        مدیریت حساب‌های بانکی تودرتوی پرسنل با رعایت تفکیک وظایف (SoD).
        DEF-10: صاحب حساب همیشه همین پرونده است (personnel فقط‌خواندنی)، و تغییر حساب با گارد وضعیت و نقش انجام می‌شود.
        """
        instance = self.get_object()

        if request.method == 'GET':
            accounts = instance.bank_accounts.all()
            serializer = PersonnelBankAccountSerializer(accounts, many=True)
            return Response(serializer.data, status=status.HTTP_200_OK)

        from .lifecycle_services import check_bank_account_edit_allowed
        allowed, deny_msg = check_bank_account_edit_allowed(request.user, instance)
        if not allowed:
            return Response({'error': deny_msg}, status=status.HTTP_400_BAD_REQUEST)

        def _log_bank_change(op, account_id, diff=None):
            # تغییر حساب بیرون از پیش‌نویس (مثلاً ویرایش حسابدار) با diff در لاگ ممیزی ثبت می‌شود؛
            # حذف حساب در هر وضعیتی با عکس کامل حساب ثبت می‌شود تا قابل بازسازی بماند (بازبینی از دست رفتن داده)
            in_draft = instance.approval_status in ('draft', 'revision_required')
            if in_draft and op != 'delete':
                return
            WorkflowAuditLog.objects.create(
                content_type='personnel',
                object_id=instance.id,
                actor=request.user,
                from_status=instance.approval_status,
                to_status=instance.approval_status,
                action='bank_account_deleted' if in_draft else 'accountant_edit',
                reason=request.data.get('note', '') or 'تغییر حساب بانکی پرسنل',
                metadata={'bank_account': {'op': op, 'id': account_id, 'diff': diff or {}}, 'company_id': instance.company_id}
            )

        if request.method == 'POST':
            serializer = PersonnelBankAccountSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            with transaction.atomic():
                acc = serializer.save(personnel=instance)
                _log_bank_change('create', acc.id, {k: str(v) for k, v in serializer.validated_data.items()})
            return Response(PersonnelBankAccountSerializer(acc).data, status=status.HTTP_201_CREATED)

        account_id = request.data.get('id') or request.query_params.get('account_id') or request.query_params.get('id')
        if not account_id:
            return Response({'error': 'شناسه حساب بانکی (id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        acc = instance.bank_accounts.filter(pk=account_id).first()
        if not acc:
            return Response({'error': 'حساب بانکی موردنظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        if request.method in ['PUT', 'PATCH']:
            serializer = PersonnelBankAccountSerializer(acc, data=request.data, partial=True)
            serializer.is_valid(raise_exception=True)
            diff = {
                k: {'old': str(getattr(acc, k)) if getattr(acc, k) is not None else None, 'new': str(v) if v is not None else None}
                for k, v in serializer.validated_data.items() if getattr(acc, k) != v
            }
            with transaction.atomic():
                updated_acc = serializer.save()
                if diff:
                    _log_bank_change('update', updated_acc.id, diff)
            return Response(PersonnelBankAccountSerializer(updated_acc).data, status=status.HTTP_200_OK)

        if request.method == 'DELETE':
            with transaction.atomic():
                from .lifecycle_services import _snapshot_value
                account_snapshot = {
                    f.attname: _snapshot_value(getattr(acc, f.attname))
                    for f in acc._meta.concrete_fields
                }
                _log_bank_change('delete', acc.id, account_snapshot)
                acc.delete()
            return Response({'message': 'حساب بانکی با موفقیت حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get', 'post', 'delete'], url_path='documents')
    def documents(self, request, pk=None):
        """
        مدیریت و بارگذاری مدارک چندگانه پرسنل (تسک ۶.۶ و ۱.۱۱)
        """
        personnel = self.get_object()

        if request.method == 'GET':
            docs = personnel.documents.all().order_by('-created_at')
            return Response(PersonnelDocumentSerializer(docs, many=True).data, status=status.HTTP_200_OK)

        elif request.method == 'POST':
            file_obj = request.FILES.get('file') or request.FILES.get('attachment')
            if not file_obj:
                return Response({'error': 'فایل مدرک الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            doc_type = request.data.get('doc_type') or 'other'
            # DEF-27: سنجش سروری نوع مدرک، پسوند و حجم فایل (هم‌راستا با فرم: تصویر یا PDF تا ۱۰ مگابایت)
            if doc_type not in dict(PersonnelDocument.DOC_TYPE_CHOICES):
                return Response({'error': f'نوع مدرک «{doc_type}» نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
            from .serializers import validate_uploaded_document
            try:
                validate_uploaded_document(
                    file_obj,
                    max_size_bytes=10 * 1024 * 1024,
                    allowed_exts=('.pdf', '.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.heic', '.heif'),
                )
            except serializers.ValidationError as e:
                return Response({'error': ' '.join(str(d) for d in e.detail)}, status=status.HTTP_400_BAD_REQUEST)
            doc = PersonnelDocument.objects.create(
                personnel=personnel,
                file=file_obj,
                doc_type=doc_type,
                uploaded_by=request.user if request.user.is_authenticated else None
            )
            return Response(PersonnelDocumentSerializer(doc).data, status=status.HTTP_201_CREATED)

        elif request.method == 'DELETE':
            doc_id = request.data.get('doc_id') or request.query_params.get('doc_id')
            if not doc_id or not str(doc_id).isdigit():
                return Response({'error': 'شناسه مدرک الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            # حذف مدرک همان گارد وضعیت و نقش حساب‌های بانکی را دارد (پروندهٔ مصوب و در انتظار مدیر قفل است)
            from .lifecycle_services import _snapshot_value, check_bank_account_edit_allowed
            allowed, deny_msg = check_bank_account_edit_allowed(request.user, personnel)
            if not allowed:
                return Response({'error': deny_msg.replace('حساب‌های بانکی', 'مدارک')}, status=status.HTTP_400_BAD_REQUEST)
            doc = personnel.documents.filter(id=int(doc_id)).select_related('uploaded_by').first()
            if not doc:
                return Response({'error': 'مدرک مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
            with transaction.atomic():
                # ردیف حذف می‌شود ولی فایل روی دیسک می‌ماند؛ مسیر و مشخصات آن در لاگ ممیزی ثبت می‌شود تا قابل بازسازی باشد
                WorkflowAuditLog.objects.create(
                    content_type='personnel',
                    object_id=personnel.id,
                    actor=request.user,
                    from_status=personnel.approval_status,
                    to_status=personnel.approval_status,
                    action='document_deleted',
                    reason=request.data.get('note', '') if hasattr(request.data, 'get') else '',
                    metadata={
                        'document': {
                            'id': doc.id,
                            'doc_type': doc.doc_type,
                            'file': doc.file.name if doc.file else None,
                            'uploaded_by_id': doc.uploaded_by_id,
                            'uploaded_by_username': doc.uploaded_by.username if doc.uploaded_by else None,
                            'created_at': _snapshot_value(doc.created_at),
                        },
                        'company_id': personnel.company_id,
                    }
                )
                doc.delete()
            return Response({'message': 'مدرک با موفقیت حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='job-titles')
    def job_titles(self, request):
        """
        ارائه لیست تجمیعی سمت‌های شغلی استاندارد، کدهای ۶ رقمی بیمه و عناوین استفاده‌شده در سامانه
        """
        q = (request.query_params.get('search') or request.query_params.get('q') or '').strip()
        from .lifecycle_services import normalize_job_title as _normalize_title
        q_normalized = _normalize_title(q)

        # استخراج کدهای ۶ رقمی و سمت‌ها از مدل پایگاه داده InsuranceJobTitle
        insurance_qs = InsuranceJobTitle.objects.filter(is_active=True)
        if q:
            insurance_qs = insurance_qs.filter(Q(title__icontains=q) | Q(title__icontains=q_normalized) | Q(job_code__icontains=q))

        title_to_codes = {}
        for item in insurance_qs:
            title_clean = _normalize_title(item.title)
            code_clean = (item.job_code or '').strip()
            if title_clean:
                if code_clean:
                    title_to_codes.setdefault(title_clean, []).append(code_clean)
                else:
                    title_to_codes.setdefault(title_clean, [])

        standard_titles = [
            'کارگر ساده انبار',
            'اپراتور لیفتراک',
            'کمک انباردار',
            'راننده خودرو سبک/سنگین',
            'مسئول انبار',
            'کارشناس اداری و مالی',
            'نگهبان و انتظامات',
            'کارگر فنی و تاسیسات',
            'استادکار تعمیرات و نگهداری',
            'سایر'
        ]
        db_titles = list(
            PersonnelProfile.objects.exclude(job_title__isnull=True)
            .exclude(job_title='')
            .values_list('job_title', flat=True)
            .distinct()
        )

        # DEF-22: تطبیق و حذف تکرار با کلید یکدست («ي/ك» عربی، فاصله و نیم‌فاصله نادیده)؛ عنوان جدول سمت‌ها مقدم است
        from .lifecycle_services import job_title_match_key, normalize_job_title
        q_key = job_title_match_key(q)
        all_candidate_titles = list(title_to_codes.keys()) + standard_titles + db_titles
        merged = []
        seen_keys = set()
        for t in all_candidate_titles:
            clean = normalize_job_title(t)
            key = job_title_match_key(clean)
            if clean and key not in seen_keys:
                t_codes = title_to_codes.get(clean, [])
                match_code = any(q.lower() in str(c).lower() for c in t_codes)
                if not q or (q_key and q_key in key) or match_code:
                    merged.append(clean)
                    seen_keys.add(key)

        items = []
        for t in merged:
            items.append({
                'title': t,
                'job_codes': title_to_codes.get(t, [])
            })

        return Response({
            'job_titles': merged,
            'items': items
        })

    @action(detail=False, methods=['get'], url_path='lookup-by-national-code')
    def lookup_by_national_code(self, request):
        """
        استعلام پرونده پرسنل بر اساس کد ملی با اعتبارسنجی رقم کنترلی، ایزولاسیون شرکتی و پشتیبانی از بازگشایی پرونده‌های ردشده
        """
        raw_code = request.query_params.get('national_code', '').strip()
        if not raw_code:
            return Response({'error': 'کد ملی ارسال نشده است.'}, status=status.HTTP_400_BAD_REQUEST)
        
        from common.validators import validate_iranian_national_code
        try:
            clean_code = validate_iranian_national_code(raw_code)
        except ValueError as e:
            return Response({'error': str(e), 'code': 'invalid_national_code'}, status=status.HTTP_400_BAD_REQUEST)
        
        user = request.user
        cid = validate_user_company_access(user, get_request_company_id(request))
        
        qs = PersonnelProfile.objects.all().select_related('section', 'section__project', 'company')
        if cid:
            qs = qs.filter(Q(company_id=cid) | Q(section__project__company_id=cid))
        elif not user.is_superuser:
            # DEF-15: بدون سربرگ X-Company-ID، جستجو به شرکت‌های مجاز کاربر محدود می‌شود (همان قاعدهٔ قلمرو سراسری)
            allowed_cids = list(get_user_allowed_companies(user).values_list('id', flat=True))
            if allowed_cids:
                qs = qs.filter(Q(company_id__in=allowed_cids) | Q(section__project__company_id__in=allowed_cids))
            else:
                qs = qs.none()
        
        personnel = qs.filter(national_code=clean_code).first()
        if not personnel:
            # وجود کد ملی در شرکت‌های بیرون از قلمرو کاربر افشا نمی‌شود (یکتایی کد ملی درون هر شرکت است)
            return Response({'found': False}, status=status.HTTP_200_OK)
        
        attachment_url = None
        if personnel.attachment:
            try:
                attachment_url = request.build_absolute_uri(personnel.attachment.url)
            except Exception:
                attachment_url = personnel.attachment.url
        
        is_reopenable = (personnel.approval_status == 'rejected')
        from .lifecycle_services import user_can_view_wages
        can_view_wages = user_can_view_wages(user)

        payload = {
            'found': True,
            'is_reopenable': is_reopenable,
            'personnel': {
                'id': personnel.id,
                'first_name': personnel.first_name,
                'last_name': personnel.last_name,
                'national_code': personnel.national_code,
                'father_name': personnel.father_name or '',
                'gender': personnel.gender or 'مرد',
                'id_number': personnel.id_number or '',
                'id_series': personnel.id_series or '',
                'id_serial': personnel.id_serial or '',
                'birth_date': personnel.birth_date or '',
                'birth_place': personnel.birth_place or '',
                'issue_place': personnel.issue_place or '',
                'marital_status': personnel.marital_status or 'single',
                'children_count': personnel.children_count or 0,
                'phone_number': personnel.phone_number or '',
                'address': personnel.address or '',
                'postal_code': personnel.postal_code or '',
                'valid_until_date': personnel.valid_until_date or '',
                'recommender': personnel.recommender or '',
                'job_title': personnel.job_title or '',
                'daily_base_wage': float(personnel.daily_base_wage or 0),
                'bank_name': personnel.bank_name or '',
                'account_number': personnel.account_number or '',
                'sheba_number': personnel.sheba_number or '',
                'card_number': personnel.card_number or '',
                'notes': personnel.notes or '',
                'approval_status': personnel.approval_status,
                'rejection_reason': personnel.rejection_reason or '',
                'current_section_id': personnel.section_id,
                'current_section_name': personnel.section.name if personnel.section else None,
                'current_project_name': personnel.section.project.name if (personnel.section and personnel.section.project) else None,
                'attachment_url': attachment_url,
            }
        }
        if not can_view_wages:
            # DEF-21: دستمزد برای نقش غیرمالی برگردانده نمی‌شود
            payload['personnel'].pop('daily_base_wage', None)
        return Response(payload, status=status.HTTP_200_OK)

    @action(detail=False, methods=['post'], url_path='cost-estimate')
    def cost_estimate(self, request):
        """
        DEF-05: برآورد سروری هزینهٔ ماهانهٔ یک پرسنل با همان موتور حقوق، برای ماشین‌حساب حسابدار و ویجت مدیر.
        بدنه: {personnel_id, year_month?, overrides?}. overrides مقادیر فرم ذخیره‌نشده است و چیزی ذخیره نمی‌شود.
        """
        from .lifecycle_services import estimate_monthly_personnel_cost, user_can_view_wages
        if not user_can_view_wages(request.user):
            raise PermissionDenied("برآورد هزینه و دستمزد فقط برای نقش‌های مالی قابل مشاهده است.")
        personnel_id = request.data.get('personnel_id')
        try:
            personnel = self._scoped_base_queryset().get(pk=int(personnel_id))
        except (TypeError, ValueError, PersonnelProfile.DoesNotExist):
            return Response({'error': 'پرونده پرسنلی یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
        overrides = request.data.get('overrides') or {}
        if not isinstance(overrides, dict):
            return Response({'error': 'overrides باید شیء باشد.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            estimate = estimate_monthly_personnel_cost(personnel, request.data.get('year_month'), overrides=overrides)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response(estimate, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='company-cost-summary')
    def company_cost_summary(self, request):
        """
        DEF-04: میانگین و جمع هزینهٔ ماهانهٔ پرسنل مصوب و فعال در قلمرو کاربر (شرکت فعال)، با همان سرویس برآورد.
        فیلترهای نمایشی (جستجو، بخش) اثر ندارند و فقط عدد تجمیعی برمی‌گردد.
        """
        from .lifecycle_services import estimate_monthly_personnel_cost, user_can_view_wages
        if not user_can_view_wages(request.user):
            raise PermissionDenied("خلاصهٔ هزینهٔ شرکت فقط برای نقش‌های مالی قابل مشاهده است.")
        year_month = request.query_params.get('year_month')
        personnel_qs = self._scoped_base_queryset().filter(approval_status='approved', is_active=True).distinct()
        cache = {}
        count = 0
        total_cost = 0
        total_gross = 0
        try:
            for p in personnel_qs:
                est = estimate_monthly_personnel_cost(p, year_month, cache=cache)
                count += 1
                total_cost += est['employer_monthly_cost']
                total_gross += est['gross_salary']
                year_month = est['year_month']
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        from .lifecycle_services import current_shamsi_year_month
        return Response({
            'year_month': year_month or current_shamsi_year_month(),
            'personnel_count': count,
            'total_employer_monthly_cost': total_cost,
            'average_employer_monthly_cost': round(total_cost / count) if count else 0,
            'average_gross_salary': round(total_gross / count) if count else 0,
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='missing-primary-sheba-report')
    def missing_primary_sheba_report(self, request):
        """
        DEF-23 (گام ۱-ب اصلاحیه): دانلود فایل اکسل گزارش حساب‌های اصلی فاقد شبای معتبر.
        محدود به قلمرو مجاز کاربر (_scoped_base_queryset) و نیازمند دسترسی مالی (user_can_view_wages).
        """
        from django.http import HttpResponse
        from .lifecycle_services import user_can_view_wages, generate_missing_primary_sheba_excel

        if not user_can_view_wages(request.user):
            raise PermissionDenied("مشاهده و دانلود گزارش اطلاعات بانکی و شبا منحصر به نقش مالی و حسابداری است.")

        scoped_qs = self._scoped_base_queryset()
        excel_bytes = generate_missing_primary_sheba_excel(queryset=scoped_qs)

        filename = "Missing_Primary_Sheba_Report.xlsx"
        response = HttpResponse(
            excel_bytes,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    @action(detail=True, methods=['post'], url_path='reopen')
    def reopen(self, request, pk=None):
        """
        DEF-15: بازگشایی پروندهٔ ردشده (rejected ← draft) به‌جای ساخت رکورد تکراری که قید یکتایی
        company + national_code آن را رد می‌کند. فقط در شرکت‌های مجاز کاربر و با ثبت لاگ.
        """
        instance = self.get_object()
        user = request.user
        if instance.approval_status != 'rejected':
            return Response({'error': f'فقط پرونده‌های ردشده قابل بازگشایی هستند (وضعیت فعلی: «{instance.get_approval_status_display()}»).'}, status=status.HTTP_400_BAD_REQUEST)

        if not user.is_superuser:
            company_id = instance.company_id or (instance.section.project.company_id if instance.section_id and instance.section and instance.section.project else None)
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if company_id is None or company_id not in allowed_ids:
                raise PermissionDenied("بازگشایی فقط برای پرونده‌های شرکت‌های مجاز شما ممکن است.")
            is_in_section = bool(instance.section_id) and UserSectionAssignment.objects.filter(user=user, section_id=instance.section_id, is_active=True).exists()
            from .lifecycle_services import APPROVAL_STAGES, user_has_any_perm
            is_approver = any(user_has_any_perm(user, rule['perms']) for rule in APPROVAL_STAGES.values())
            if not (instance.created_by_id == user.id or is_in_section or is_approver):
                raise PermissionDenied("شما مجاز به بازگشایی این پرونده نیستید.")

        note = (request.data.get('note') or request.data.get('reason') or '').strip()
        with transaction.atomic():
            locked = PersonnelProfile.objects.select_for_update().get(pk=instance.pk)
            if locked.approval_status != 'rejected':
                return Response({'error': 'وضعیت پرونده هم‌زمان تغییر کرده است؛ صفحه را تازه کنید.'}, status=status.HTTP_409_CONFLICT)
            locked.approval_status = 'draft'
            locked.save(update_fields=['approval_status', 'updated_at'])
            WorkflowAuditLog.objects.create(
                content_type='personnel',
                object_id=locked.id,
                actor=user,
                from_status='rejected',
                to_status='draft',
                action='reopen',
                reason=note or 'بازگشایی پرونده ردشده جهت اصلاح و ارسال دوباره',
                metadata={'company_id': locked.company_id, 'previous_rejection_reason': locked.rejection_reason}
            )
        broadcast_personnel_update(locked, action_type='updated', message=f'پرونده ردشدهٔ «{locked.full_name}» بازگشایی شد.', sender_id=user.id)
        return Response({
            'message': f'پرونده «{locked.full_name}» بازگشایی شد و در وضعیت پیش‌نویس قرار گرفت.',
            'data': self.get_serializer(locked).data
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='end-collaboration')
    def end_collaboration(self, request, pk=None):
        """
        DEF-18: قطع همکاری پرسنل دارای سابقه (جایگزین حذف فیزیکی). تاریخ پایان شمسی و دلیل الزامی‌اند؛
        پرونده غیرفعال می‌شود و end_date ثبت می‌شود، همراه با لاگ ممیزی و اعلان. فقط با مجوز مدیر.
        """
        from .lifecycle_services import MANAGER_STAGE_PERMS, dispatch_workflow_notification, normalize_shamsi_date, user_has_any_perm
        user = request.user
        if not user_has_any_perm(user, MANAGER_STAGE_PERMS):
            raise PermissionDenied("ثبت قطع همکاری فقط با دسترسی مدیر ممکن است.")

        instance = self.get_object()
        reason = (request.data.get('reason') or '').strip()
        if not reason:
            return Response({'error': 'ثبت دلیل قطع همکاری الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            end_date = normalize_shamsi_date(request.data.get('end_date'))
        except ValueError as e:
            return Response({'error': str(e), 'field': 'end_date'}, status=status.HTTP_400_BAD_REQUEST)
        if not end_date:
            return Response({'error': 'تاریخ پایان همکاری (شمسی) الزامی است.', 'field': 'end_date'}, status=status.HTTP_400_BAD_REQUEST)
        if instance.approval_status not in ('approved', 'manager_approved'):
            return Response({'error': f'قطع همکاری فقط برای پرسنل مصوب و فعال ممکن است (وضعیت فعلی: «{instance.get_approval_status_display()}»).'}, status=status.HTTP_400_BAD_REQUEST)
        if not instance.is_active:
            return Response({'error': 'همکاری این پرسنل پیش‌تر پایان یافته و پرونده غیرفعال است.'}, status=status.HTTP_400_BAD_REQUEST)
        if instance.start_date:
            from .lifecycle_services import parse_shamsi_date
            start = parse_shamsi_date(instance.start_date)
            if start is not None and parse_shamsi_date(end_date) < start:
                return Response({'error': f'تاریخ پایان همکاری نمی‌تواند پیش از تاریخ شروع به کار ({instance.start_date}) باشد.', 'field': 'end_date'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            locked = PersonnelProfile.objects.select_for_update().get(pk=instance.pk)
            previous_end_date = locked.end_date
            locked.end_date = end_date
            locked.is_active = False
            locked.save(update_fields=['end_date', 'is_active', 'updated_at'])
            WorkflowAuditLog.objects.create(
                content_type='personnel',
                object_id=locked.id,
                actor=user,
                from_status=locked.approval_status,
                to_status=locked.approval_status,
                action='end_collaboration',
                reason=reason,
                metadata={'end_date': end_date, 'previous_end_date': previous_end_date, 'company_id': locked.company_id}
            )
            dispatch_workflow_notification(locked, stage='end_collaboration', actor=user, reason=reason)
        broadcast_personnel_update(locked, action_type='updated', message=f'همکاری پرسنل «{locked.full_name}» از تاریخ {end_date} پایان یافت.', sender_id=user.id)
        return Response({
            'message': f'همکاری پرسنل «{locked.full_name}» از تاریخ {end_date} پایان یافت و پرونده غیرفعال شد.',
            'data': self.get_serializer(locked).data
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='assign-section')
    def assign_section(self, request, pk=None):
        """
        انتساب پرسنل موجود به یک بخش دیگر جهت فعالیت همزمان در چند بخش
        """
        personnel = self.get_object()
        section_id = request.data.get('section_id') or request.data.get('section')
        if not section_id:
            return Response({'error': 'شناسه بخش ارسالی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            section_id = int(section_id)
        except (ValueError, TypeError):
            return Response({'error': 'شناسه بخش نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        sec = ProjectSection.objects.filter(id=section_id).select_related('project').first()
        if not sec:
            return Response({'error': 'بخش مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
        # بخش مقصد باید در قلمرو شرکت‌های مجاز کاربر و در همان شرکتِ پرونده باشد
        personnel_company_id = personnel.company_id or (
            personnel.section.project.company_id if personnel.section_id else None
        )
        ensure_personnel_write_scope(request.user, company=personnel_company_id, section=sec)

        job_title = (request.data.get('job_title') or personnel.job_title or '').strip()
        from .lifecycle_services import user_can_view_wages
        # DEF-21: دستمزد انتساب را فقط نقش مالی تعیین می‌کند؛ مقدار ارسالی نقش غیرمالی نادیده گرفته می‌شود
        requested_wage = request.data.get('daily_base_wage') if user_can_view_wages(request.user) else None
        daily_base_wage = requested_wage or personnel.daily_base_wage or 0
        
        assignment, created = PersonnelSectionAssignment.objects.update_or_create(
            personnel=personnel,
            section=sec,
            defaults={
                'project': sec.project,
                'job_title': job_title,
                'daily_base_wage': daily_base_wage,
                'is_active': True
            }
        )
        
        broadcast_personnel_update(personnel, action_type='assigned', message=f'پرسنل «{personnel.full_name}» به بخش «{sec.name}» منتسب شد.', sender_id=request.user.id if request.user else None)
        
        return Response({
            'message': f'پرسنل «{personnel.full_name}» با موفقیت به بخش «{sec.name}» منتسب شد.',
            'assignment_id': assignment.id,
            'personnel': self.get_serializer(personnel).data
        }, status=status.HTTP_200_OK)

    # پیام‌های موفقیت و رویداد وب‌سوکت هر مرحلهٔ تأیید (منطق گذار در lifecycle_services.approve_personnel_stage است)
    APPROVAL_STAGE_MESSAGES = {
        'supervisor': (
            'pending_accountant',
            'پرونده پرسنل «{name}» به تایید سرپرست رسید و به حسابداری ارسال شد.',
            'تایید سرپرست با موفقیت ثبت شد و پرونده به حسابداری ارسال گردید.',
        ),
        'finance': (
            'pending_manager',
            'پرونده پرسنل «{name}» به تایید مالی رسید و به کارتابل مدیر ارسال شد.',
            'تایید مالی با موفقیت ثبت شد و پرونده جهت تصویب نهایی به کارتابل مدیر ارسال گردید.',
        ),
        'manager': (
            'approved',
            'پرونده پرسنل «{name}» به تصویب نهایی مدیر رسید و فعال گردید.',
            'تصویب نهایی مدیر با موفقیت ثبت شد و پرسنل فعال گردید.',
        ),
    }

    def _approve_single(self, request, stage):
        from .lifecycle_services import APPROVAL_STAGES, PersonnelWorkflowError, approve_personnel_stage, user_has_any_perm
        user = request.user
        if not user_has_any_perm(user, APPROVAL_STAGES[stage]['perms']):
            raise PermissionDenied(APPROVAL_STAGES[stage]['perm_error'])
        instance = self.get_object()
        note = (request.data.get('note') or request.data.get('approval_note') or '').strip()
        try:
            instance = approve_personnel_stage(instance.pk, user, stage, note=note)
        except PersonnelWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)
        action_type, broadcast_msg, response_msg = self.APPROVAL_STAGE_MESSAGES[stage]
        broadcast_personnel_update(instance, action_type=action_type, message=broadcast_msg.format(name=instance.full_name), sender_id=user.id)
        return Response({
            'message': response_msg,
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='approve-supervisor')
    def approve_supervisor(self, request, pk=None):
        return self._approve_single(request, 'supervisor')

    @action(detail=True, methods=['post'], url_path='approve-finance')
    def approve_finance(self, request, pk=None):
        return self._approve_single(request, 'finance')

    @action(detail=True, methods=['post'], url_path='approve-manager')
    def approve_manager(self, request, pk=None):
        return self._approve_single(request, 'manager')

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        user = request.user
        can_mgr = (
            user.is_superuser
            or user.has_perm('accounts.perm_approve_personnel_manager')
            or user.has_perm('accounts.can_act_as_manager')
            or user.has_perm('accounts.perm_approve_personnel_supervisor')
            or user.has_perm('accounts.can_act_as_workshop_supervisor')
        )
        can_fin = user.is_superuser or user.has_perm('accounts.perm_approve_personnel_finance')
        if not (can_mgr or can_fin):
            raise PermissionDenied("شما دسترسی لازم برای رد این پرونده را ندارید.")
        
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت دلیل رد الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        
        instance = self.get_object()
        if instance.approval_status in ('approved', 'manager_approved'):
            return Response({'error': 'امکان رد پرونده تصویب‌شده وجود ندارد. در صورت نیاز، تغییرات باید از مسیر درخواست تغییرات یا قطع همکاری پیگیری شود.'}, status=status.HTTP_400_BAD_REQUEST)
        # تسک ۳.۱ و جدول §۳.۱-الف: رد فقط روی پرونده‌های در جریان؛ پیش‌نویسِ ارسال‌نشده و پروندهٔ ردشده از این مسیر رد نمی‌شوند
        in_progress_statuses = ('pending_supervisor', 'pending_accountant', 'pending_manager', 'revision_required', 'supervisor_approved', 'accountant_approved')
        if instance.approval_status not in in_progress_statuses:
            return Response({'error': f'امکان رد پرونده در وضعیت «{instance.get_approval_status_display()}» وجود ندارد؛ رد قطعی فقط برای پرونده‌های در جریان تایید است.'}, status=status.HTTP_400_BAD_REQUEST)

        from .lifecycle_services import build_personnel_snapshot, check_personnel_history, dispatch_workflow_notification
        purge_statuses = ('revision_required', 'pending_supervisor', 'pending_accountant', 'pending_manager')
        with transaction.atomic():
            # قفل ردیف پرونده و سنجش دوبارهٔ وضعیت و سابقه درون همان تراکنش: ثبت هم‌زمان کارکرد یا فیش
            # (کلید خارجی به همین ردیف) تا پایان تراکنش منتظر می‌ماند و سابقه‌ای با CASCADE پاک نمی‌شود
            instance = PersonnelProfile.objects.select_for_update().get(pk=instance.pk)
            if instance.approval_status not in in_progress_statuses:
                return Response({'error': f'امکان رد پرونده در وضعیت «{instance.get_approval_status_display()}» وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)
            history = check_personnel_history(instance)
            was_never_approved = not instance.manager_approved_at
            full_name = instance.full_name

            # شرط Purge (تسک ۳.۱ و بخش ۳.۱-ب): حذف فیزیکی فقط برای پرونده‌های بدون سابقه که هرگز تایید نهایی نشده‌اند
            purged = was_never_approved and not history['has_history'] and instance.approval_status in purge_statuses
            if purged:
                # DEF-11: عکس کامل پرونده، همهٔ حساب‌ها و فهرست مدارک پیش از حذف فیزیکی (شرط ۳ بخش ۳.۱-ب)
                snapshot = build_personnel_snapshot(instance)
                snapshot['reason'] = reason
                snapshot['purged_by'] = user.username
                WorkflowAuditLog.objects.create(
                    content_type='personnel',
                    object_id=instance.id,
                    actor=user,
                    from_status=instance.approval_status,
                    to_status='purged',
                    action='purge',
                    reason=reason,
                    metadata={'snapshot': snapshot, 'purged': True, 'company_id': instance.company_id}
                )
                dispatch_workflow_notification(instance, stage='purged', actor=user, reason=reason)
                deleted_id = instance.id
                instance.delete()
                instance.id = deleted_id
            else:
                old_status = instance.approval_status
                instance.approval_status = 'rejected'
                instance.rejection_reason = reason
                instance.save()

                WorkflowAuditLog.objects.create(
                    content_type='personnel',
                    object_id=instance.id,
                    actor=user,
                    from_status=old_status,
                    to_status='rejected',
                    action='reject',
                    reason=reason,
                    metadata={'actor_role': 'rejecter', 'company_id': instance.company_id}
                )
                dispatch_workflow_notification(instance, stage='rejected', actor=user, reason=reason)

        if purged:
            broadcast_personnel_update(instance, action_type='deleted', message=f'پرونده «{full_name}» به دلیل رد قطعی و عدم وجود سابقه پاک‌سازی (Purge) شد.', sender_id=user.id)
            return Response({'message': f'پرونده پیش‌نویس «{full_name}» به دلیل عدم وجود هرگونه سابقه پاک‌سازی شد.', 'purged': True}, status=status.HTTP_200_OK)
        broadcast_personnel_update(instance, action_type='rejected', message=f'پرونده پرسنل «{instance.full_name}» رد شد: {reason}', sender_id=user.id)
        return Response({'message': 'پرونده پرسنل رد شد.', 'purged': False, 'data': self.get_serializer(instance).data}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='request-revision')
    def request_revision(self, request, pk=None):
        from .lifecycle_services import (
            FINANCE_STAGE_PERMS,
            MANAGER_STAGE_PERMS,
            SUPERVISOR_STAGE_PERMS,
            dispatch_workflow_notification,
            user_has_any_perm,
        )
        user = request.user
        is_sup = user_has_any_perm(user, SUPERVISOR_STAGE_PERMS)
        is_fin = user_has_any_perm(user, FINANCE_STAGE_PERMS)
        is_mgr = user_has_any_perm(user, MANAGER_STAGE_PERMS)
        if not (is_mgr or is_fin or is_sup):
            raise PermissionDenied("شما دسترسی لازم برای ارجاع به بازنگری را ندارید.")

        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت توضیحات جهت بازنگری الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        instance = self.get_object()
        old_status = instance.approval_status
        if old_status in ('approved', 'manager_approved'):
            return Response({'error': 'امکان عودت پرونده تصویب‌شده وجود ندارد. پرونده‌های مصوب در چرخه ثبت کارکرد هستند و تغییرات آن‌ها صرفاً از مسیر درخواست تغییرات مجاز است.'}, status=status.HTTP_400_BAD_REQUEST)

        # DEF-13: جدول گذار بخش ۳.۱-الف طرح. مقصد بر اساس وضعیت فعلی پرونده تعیین می‌شود، نه بالاترین نقش کاربر.
        if old_status == 'pending_supervisor':
            if not is_sup:
                raise PermissionDenied("عودت پرونده در انتظار سرپرست فقط با دسترسی سرپرست ممکن است.")
            target_status = 'revision_required'
        elif old_status == 'pending_accountant':
            if not is_fin:
                raise PermissionDenied("عودت پرونده در انتظار حسابداری فقط با دسترسی حسابدار ممکن است.")
            req_dest = request.data.get('target_status') or request.data.get('destination')
            target_status = req_dest if req_dest in ('pending_supervisor', 'revision_required') else 'revision_required'
        elif old_status == 'pending_manager':
            if not is_mgr:
                raise PermissionDenied("عودت پرونده در انتظار مدیر فقط با دسترسی مدیر ممکن است.")
            target_status = 'pending_accountant'
        else:
            return Response({'error': f'امکان عودت پرونده در وضعیت «{instance.get_approval_status_display()}» وجود ندارد. عودت فقط از مراحل در انتظار سرپرست، حسابدار یا مدیر ممکن است.'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            instance.approval_status = target_status
            instance.rejection_reason = reason
            instance.revision_requested_by = user
            instance.revision_requested_at = timezone.now()
            instance.save()

            WorkflowAuditLog.objects.create(
                content_type='personnel',
                object_id=instance.id,
                actor=user,
                from_status=old_status,
                to_status=target_status,
                action='request_revision',
                reason=reason,
                metadata={'destination': target_status, 'company_id': instance.company_id}
            )
            # DEF-12: گیرندهٔ اعلان عودت همان مقصد است (کارمند، سرپرست یا حسابدار)
            dispatch_workflow_notification(instance, stage=target_status, actor=user, reason=reason, is_return=True)
        dest_title = 'کارمند' if target_status == 'revision_required' else ('سرپرست' if target_status == 'pending_supervisor' else 'حسابدار')
        broadcast_personnel_update(instance, action_type='revision_required', message=f'پرونده پرسنل «{instance.full_name}» به {dest_title} عودت داده شد: {reason}', sender_id=user.id)
        return Response({
            'message': f'پرونده جهت بازنگری به {dest_title} عودت داده شد.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=False, methods=['post'], url_path='batch-approve')
    def batch_approve(self, request):
        """
        تسک ۳.۵: اکشن تایید گروهی پرونده‌های پرسنلی با پردازش مستقل هر رکورد در تراکنش جداگانه
        """
        ids = request.data.get('ids', [])
        if not ids or not isinstance(ids, list):
            return Response({'error': 'لیست شناسه‌های پرسنل (ids) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        stage = request.data.get('stage')  # 'supervisor', 'finance', 'manager'
        note = (request.data.get('note') or '').strip()
        user = request.user

        from .lifecycle_services import APPROVAL_STAGES, PersonnelWorkflowError, approve_personnel_stage

        if stage not in APPROVAL_STAGES:
            return Response({'error': 'مرحله ارسالی نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        # DEF-09: فقط پرونده‌های داخل قلمرو کاربر (همان get_queryset تکی) قابل تأیید گروهی‌اند
        clean_ids = []
        for pid in ids:
            try:
                clean_ids.append(int(pid))
            except (TypeError, ValueError):
                clean_ids.append(None)
        scoped_names = {
            row[0]: f"{row[1]} {row[2]}"
            for row in self.get_queryset().filter(id__in=[i for i in clean_ids if i is not None]).values_list('id', 'first_name', 'last_name')
        }

        action_type, broadcast_msg, _ = self.APPROVAL_STAGE_MESSAGES[stage]
        results = []
        for raw_pid, pid in zip(ids, clean_ids):
            if pid is None or pid not in scoped_names:
                results.append({'id': raw_pid, 'success': False, 'error': 'پرونده پرسنلی یافت نشد.'})
                continue
            try:
                p = approve_personnel_stage(pid, user, stage, note=note, batch=True)
            except PersonnelWorkflowError as e:
                row = {'id': pid, 'full_name': scoped_names[pid], 'success': False, 'error': e.message}
                if e.missing_fields:
                    row['missing_fields'] = e.missing_fields
                results.append(row)
                continue
            except Exception as e:
                import logging
                logging.getLogger(__name__).exception('[batch-approve] personnel %s failed', pid)
                results.append({'id': pid, 'full_name': scoped_names[pid], 'success': False, 'error': str(e)})
                continue
            broadcast_personnel_update(p, action_type=action_type, message=broadcast_msg.format(name=p.full_name), sender_id=user.id)
            results.append({'id': pid, 'full_name': p.full_name, 'success': True})

        successful_count = sum(1 for r in results if r.get('success'))
        return Response({
            'total': len(ids),
            'successful_count': successful_count,
            'failed_count': len(ids) - successful_count,
            'results': results
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        """
        دانلود قالب اکسل خام پرسنل با ساختار ۴۶ ستونی منطبق بر شیت Emp_info و استاندارد ۲ سطری شرکت
        """
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Emp_info'
        ws.sheet_view.rightToLeft = True

        columns_spec = [
            ('status_category', 'وضعیت', 'نفرات شرکتی'),
            ('row_num', 'ردیف', 1),
            ('nationality_code', 'ملیت', '1'),
            ('national_code', 'کد ملی', '0010376488'),
            ('first_name', 'نام', 'علی'),
            ('last_name', 'نام خانوادگی', 'محمدی'),
            ('father_name', 'نام پدر', 'رضا'),
            ('birth_date', 'تاریخ تولد', '1370/01/01'),
            ('id_number', 'شماره شناسنامه', '1234'),
            ('birth_place', 'محل تولد', 'تهران'),
            ('education_level', 'مدرک تحصیلی', '4'),
            ('insurance_type', 'نوع بیمه', '2'),
            ('insurance_number', 'شماره بیمه', '12345678'),
            ('insurance_name', 'نام بیمه', 'تامین اجتماعی'),
            ('exemption_type', 'نوع معافیت', '1'),
            ('citizenship_country_code', 'کشور محل تابعیت', '103'),
            ('residence_country_code', 'کشور محل زندگی', '103'),
            ('postal_code', 'کد پستی', '1234567890'),
            ('address', 'آدرس', 'تهران، خیابان آزادی'),
            ('job_category', 'رسته', '15'),
            ('job_title', 'سمت شغل', 'کارشناس انبار'),
            ('employment_type', 'نوع استخدام', '2'),
            ('start_date', 'تاریخ شروع به کار', '1404/01/01'),
            ('end_date', 'تاریخ پایان کار', ''),
            ('retirement_date', 'تاریخ بازنشستگی', ''),
            ('job_code', 'کد شغل', '027212'),
            ('id_series', 'مسلسل شناسنامه', ''),
            ('id_serial', 'سریال', ''),
            ('issue_place', 'محل صدور', 'تهران'),
            ('issue_date', 'تاریخ صدور', ''),
            ('gender', 'جنسیت', 'مرد'),
            ('phone_number', 'موبایل', '09123456789'),
            ('account_number', 'شماره حساب', '1234567890'),
            ('marital_status', 'وضعیت تاهل', 'مجرد'),
            ('children_count', 'تعداد فرزند', 0),
            ('contract_hours', 'ساعت کار قرارداد شده', 230),
            ('contract_base_salary', 'حقوق قرارداد شده', 0),
            ('job_grade', 'گروه شغلی', '19'),
            ('daily_base_wage', 'مزد روزانه', 3500000),
            ('daily_seniority_bonus', 'پایه سنواتی', 0),
            ('base_years_experience', 'تعداد سال کارکرد', 0),
            ('base_daily_rate', 'مزد مبنا', 3500000),
            ('group_status', 'گروه', 'شاغل'),
            ('include_in_tax', 'maliat', 1),
            ('include_in_insurance', 'بیمه', 1),
            ('include_in_bank', 'bank', 1),
        ]

        f_title = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        f_key = Font(name='Segoe UI', size=9, bold=True, color='CBD5E1')
        fill_title = PatternFill(start_color='4338CA', end_color='4338CA', fill_type='solid')
        fill_key = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 20
        ws.freeze_panes = 'A3'

        for col_idx, (c_en, c_fa, sample_val) in enumerate(columns_spec, 1):
            cell_fa = ws.cell(row=1, column=col_idx, value=c_fa)
            cell_fa.font = f_title
            cell_fa.fill = fill_title
            cell_fa.alignment = align_center
            cell_fa.border = thin_border

            cell_en = ws.cell(row=2, column=col_idx, value=c_en)
            cell_en.font = f_key
            cell_en.fill = fill_key
            cell_en.alignment = align_center
            cell_en.border = thin_border

            c_val = ws.cell(row=3, column=col_idx, value=sample_val)
            c_val.alignment = align_center
            c_val.border = thin_border

        ws.row_dimensions[3].height = 22

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="personnel_template.xlsx"'
        return response

    @action(detail=False, methods=['post', 'POST', 'get', 'GET'], url_path='import-excel')
    def import_excel(self, request):
        """
        درون‌ریزی هوشمند ۴۶ ستونی شیت Emp_info با پشتیبانی از هر دو قالب (خروجی سامانه و فایل مرجع)،
        نرمال‌سازی تاریخ‌ها و شماره‌ها، ایجاد حساب اصلی، و تعیین وضعیت رکورد طبق ماشین حالت
        """
        import openpyxl
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        
        warehouse_id = request.data.get('warehouse_id') or request.query_params.get('warehouse_id')
        section_id = request.data.get('section_id') or request.query_params.get('section_id')
        dry_run = str(request.data.get('dry_run') or request.query_params.get('dry_run', 'false')).lower() in ['true', '1']
        update_existing = str(request.data.get('update_existing') or request.query_params.get('update_existing', 'true')).lower() in ['true', '1']
        
        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
        except Exception as e:
            return Response({
                'success': False,
                'dry_run': dry_run,
                'summary': {'total_rows': 0, 'created': 0, 'updated': 0, 'skipped': 0, 'valid_count': 0, 'error_count': 1},
                'errors': [{'row': 0, 'field': 'فایل اکسل', 'message': f'خطا در باز کردن فایل اکسل: {str(e)}'}]
            }, status=status.HTTP_400_BAD_REQUEST)
        
        target_sheet = None
        for name in wb.sheetnames:
            if name.strip().lower() in ['emp_info', 'emp info', 'empinfo', 'پرسنل', 'پرسنل و کارکنان']:
                target_sheet = wb[name]
                break
        if not target_sheet:
            target_sheet = wb.active

        ws = target_sheet

        FIELD_ALIASES = {
            'national_code': ['کد ملی', 'کدملی', 'national_code', 'national code', 'کد_ملی'],
            'first_name': ['نام', 'first_name', 'firstname', 'نام کوچک'],
            'last_name': ['نام خانوادگی', 'نام  خانوادگی', 'نام‌خانوادگی', 'last_name', 'lastname', 'فامیلی'],
            'father_name': ['نام پدر', 'نام  پدر', 'father_name', 'fathername', 'پدر'],
            'birth_date': ['تاریخ تولد', 'birth_date', 'birthdate', 'تولد'],
            'id_number': ['شماره شناسنامه', 'شماره  شناسنامه', 'id_number', 'شناسنامه'],
            'birth_place': ['محل تولد', 'محل  تولد', 'birth_place', 'صادره'],
            'education_level': ['مدرک تحصیلی', 'مدرک  تحصیلی', 'education_level', 'تحصیلات'],
            'insurance_type': ['نوع بیمه', 'insurance_type'],
            'insurance_number': ['شماره بیمه', 'شماره  بیمه', 'insurance_number', 'کد بیمه'],
            'insurance_name': ['نام بیمه', 'نام  بیمه', 'insurance_name'],
            'exemption_type': ['نوع معافیت', 'نوع  معافیت', 'exemption_type', 'معافیت'],
            'nationality_code': ['ملیت', 'کد ملیت', 'nationality_code', 'تابعیت'],
            'citizenship_country_code': ['کشور محل تابعیت', 'citizenship_country_code', 'کشور تابعیت'],
            'residence_country_code': ['کشور محل زندگی', 'residence_country_code', 'کشور زندگی', 'محل سکونت'],
            'postal_code': ['کد پستی', 'کدپستی', 'postal_code', 'کد_پستی'],
            'address': ['آدرس', 'ادرس', 'آدرس محل سکونت', 'address', 'نشانی'],
            'job_category': ['رسته', 'job_category'],
            'job_title': ['سمت شغل', 'سمت  شغل', 'عنوان شغل', 'سمت شغلی', 'job_title', 'شغل'],
            'employment_type': ['نوع استخدام', 'employment_type'],
            'start_date': ['تاریخ شروع به کار', 'تاریخ شروع', 'start_date', 'شروع به کار'],
            'end_date': ['تاریخ پایان کار', 'تاریخ پایان', 'end_date', 'پایان کار'],
            'retirement_date': ['تاریخ بازنشستگی', 'retirement_date'],
            'job_code': ['کد شغل', 'job_code', 'کد_شغل'],
            'id_series': ['مسلسل شناسنامه', 'مسلسل', 'id_series'],
            'id_serial': ['سریال', 'سریال شناسنامه', 'id_serial'],
            'issue_place': ['محل صدور', 'issue_place'],
            'issue_date': ['تاریخ صدور', 'issue_date'],
            'gender': ['جنسیت', 'gender'],
            'phone_number': ['موبایل', 'شماره همراه', 'شماره تماس', 'phone_number', 'تلفن همراه'],
            'account_number': ['شماره حساب', 'شماره  حساب', 'account_number', 'حساب'],
            'sheba_number': ['شماره شبا', 'شماره  شبا', 'شبا', 'sheba_number', 'sheba'],
            'bank_name': ['نام بانک', 'نام  بانک', 'بانک', 'bank_name'],
            'marital_status': ['وضعیت تاهل', 'وضعیت  تاهل', 'marital_status', 'تاهل'],
            'children_count': ['تعداد فرزند', 'تعداد  فرزند', 'children_count', 'فرزند'],
            'contract_hours': ['ساعت کار قرارداد شده', 'ساعت کار قرارداد', 'contract_hours'],
            'contract_base_salary': ['حقوق قرارداد شده', 'حقوق قرارداد', 'contract_base_salary'],
            'job_grade': ['گروه شغلی', 'job_grade'],
            'daily_base_wage': ['مزد روزانه', 'دستمزد روزانه', 'daily_base_wage', 'مزد روزانه '],
            'daily_seniority_bonus': ['پایه سنواتی', 'سنوات روزانه', 'daily_seniority_bonus'],
            'base_years_experience': ['تعداد سال کارکرد', 'سال کارکرد', 'base_years_experience'],
            'base_daily_rate': ['مزد مبنا', 'دستمزد مبنا', 'base_daily_rate'],
            'status_category': ['وضعیت', 'وضعیت پرسنل', 'status_category'],
            'group_status': ['گروه', 'group_status'],
            'include_in_tax': ['maliat', 'مالیات', 'include_in_tax'],
            'include_in_insurance': ['بیمه', 'include_in_insurance', 'bimeh'],
            'include_in_bank': ['bank', 'بانک', 'include_in_bank'],
        }

        # تشخیص هوشمند ردیف‌های سربرگ و شروع داده‌ها
        header_row_idx = 1
        data_start_row = 2
        col_map = {}

        for r_idx in range(1, min(ws.max_row + 1, 6)):
            row_vals = [str(ws.cell(row=r_idx, column=c).value or '').strip() for c in range(1, ws.max_column + 1)]
            row_str = ' '.join(row_vals)
            if 'کد ملی' in row_str or 'national_code' in row_str or 'نام خانوادگی' in row_str or 'last_name' in row_str:
                header_row_idx = r_idx
                break

        # بررسی اینکه آیا سطر بعدی نیز سطر کلیدهای انگلیسی یا هدر ثانویه است
        next_row_vals = [str(ws.cell(row=header_row_idx + 1, column=c).value or '').strip() for c in range(1, ws.max_column + 1)] if header_row_idx < ws.max_row else []
        next_row_str = ' '.join(next_row_vals).lower()
        if 'national_code' in next_row_str or 'last_name' in next_row_str or 'first_name' in next_row_str or 'job_title' in next_row_str:
            data_start_row = header_row_idx + 2
        else:
            data_start_row = header_row_idx + 1

        for c in range(1, ws.max_column + 1):
            val1 = str(ws.cell(row=header_row_idx, column=c).value or '').strip().replace('  ', ' ')
            val2 = str(ws.cell(row=header_row_idx + 1, column=c).value or '').strip().replace('  ', ' ') if header_row_idx + 1 < data_start_row else ''
            for field, aliases in FIELD_ALIASES.items():
                if field in col_map:
                    continue
                for alias in aliases:
                    if alias == val1 or alias == val2 or alias == val1.lower() or alias == val2.lower():
                        col_map[field] = c
                        break

        if 'national_code' not in col_map and 'last_name' not in col_map:
            return Response({
                'success': False,
                'dry_run': dry_run,
                'summary': {'total_rows': 0, 'created': 0, 'updated': 0, 'skipped': 0, 'valid_count': 0, 'error_count': 1},
                'errors': [{'row': header_row_idx, 'field': 'ستون‌ها', 'message': 'شیت پرسنلی یا ستون‌های الزامی (کد ملی، نام خانوادگی) در فایل یافت نشد.'}]
            }, status=status.HTTP_400_BAD_REQUEST)

        created_count = 0
        updated_count = 0
        skipped_count = 0
        errors = []
        preview_rows = []
        valid_rows_data = []

        persian_arabic_digits = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')

        def get_cell_str(r, field, default=''):
            c_idx = col_map.get(field)
            if c_idx:
                val = ws.cell(row=r, column=c_idx).value
                if val is not None:
                    return str(val).strip()
            return default

        def get_cell_num(r, field, default=0):
            c_idx = col_map.get(field)
            if c_idx:
                val = ws.cell(row=r, column=c_idx).value
                if val is not None:
                    try:
                        clean_str = str(val).translate(persian_arabic_digits).replace(',', '').strip()
                        return float(clean_str)
                    except (ValueError, TypeError):
                        pass
            return default

        def get_cell_bool(r, field, default=True):
            c_idx = col_map.get(field)
            if c_idx:
                val = ws.cell(row=r, column=c_idx).value
                if val is not None:
                    s = str(val).strip().lower()
                    return s in ['1', 'true', 'بله', 'yes']
            return default

        def normalize_date_val(val):
            if not val:
                return ''
            raw_s = str(val).translate(persian_arabic_digits).strip()
            if not raw_s or raw_s.lower() == 'none':
                return ''
            raw_s = raw_s.split('.')[0]
            if len(raw_s) == 8 and raw_s.isdigit():
                return f"{raw_s[:4]}/{raw_s[4:6]}/{raw_s[6:]}"
            raw_s = raw_s.replace('-', '/')
            parts = raw_s.split('/')
            if len(parts) == 3 and len(parts[0]) == 4 and parts[1].isdigit() and parts[2].isdigit():
                return f"{parts[0]}/{int(parts[1]):02d}/{int(parts[2]):02d}"
            return raw_s

        # بخش مقصد یک‌بار و با قلمرو سنجیده می‌شود؛ کاربر غیرسوپریوزر فقط در شرکت‌های مجاز خودش درون‌ریزی می‌کند
        sec_obj = None
        if section_id:
            if not str(section_id).isdigit():
                return Response({'error': 'شناسه بخش نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
            sec_obj = ProjectSection.objects.filter(id=int(section_id)).select_related('project').first()
            if not sec_obj:
                return Response({'error': 'بخش مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
        if requires_company_scope(request.user):
            if sec_obj is not None:
                ensure_personnel_write_scope(request.user, section=sec_obj)
                active_cid = sec_obj.project.company_id
                requested_cid = get_request_company_id(request)
                if requested_cid and str(requested_cid).isdigit() and int(requested_cid) != active_cid:
                    raise PermissionDenied("بخش انتخاب‌شده به شرکت فعال شما تعلق ندارد.")
            else:
                active_cid = resolve_import_company(request)
        else:
            from common.tenant_scope import get_current_tenant_company_id
            active_cid = get_current_tenant_company_id(request=request, user=request.user)

        from common.validators import validate_iranian_national_code
        from .lifecycle_services import normalize_job_title

        for r in range(data_start_row, ws.max_row + 1):
            nat_c = col_map.get('national_code')
            if not nat_c:
                continue
            raw_nc = ws.cell(row=r, column=nat_c).value
            if raw_nc is None or str(raw_nc).strip() == '':
                continue

            nc_raw_str = str(raw_nc).translate(persian_arabic_digits).strip().split('.')[0]
            if nc_raw_str.lower() in ['national_code', 'national code', 'کد ملی', 'کلید']:
                continue

            nc_clean = nc_raw_str.zfill(10)
            try:
                nc_clean = validate_iranian_national_code(nc_clean)
            except ValueError as e:
                errors.append({
                    'row': r,
                    'field': 'کد ملی',
                    'message': f'کد ملی نامعتبر است ({raw_nc}): {str(e)}'
                })
                continue

            first_name = get_cell_str(r, 'first_name')
            last_name = get_cell_str(r, 'last_name')
            if not last_name:
                errors.append({
                    'row': r,
                    'field': 'نام خانوادگی',
                    'message': 'نام خانوادگی در این ردیف خالی است'
                })
                continue

            # استخراج و نرمال‌سازی شماره همراه
            raw_phone = get_cell_str(r, 'phone_number')
            clean_phone = ''
            if raw_phone:
                p_digits = str(raw_phone).translate(persian_arabic_digits).strip().split('.')[0]
                if len(p_digits) == 10 and p_digits.startswith('9'):
                    clean_phone = f"0{p_digits}"
                else:
                    clean_phone = p_digits

            # استخراج و نرمال‌سازی کد شغل ۶ رقمی
            raw_jc = get_cell_str(r, 'job_code')
            clean_jc = ''
            if raw_jc:
                clean_jc = str(raw_jc).translate(persian_arabic_digits).strip().split('.')[0]
                if clean_jc.isdigit() and len(clean_jc) <= 6:
                    clean_jc = clean_jc.zfill(6)

            # ستون‌های حاضر و غیرخالی همین ردیف: رکورد موجود فقط از همین‌ها به‌روز می‌شود، نه از پیش‌فرض ستون‌های غایب
            present_fields = set()
            for field_name, c_idx in col_map.items():
                cell_val = ws.cell(row=r, column=c_idx).value
                if cell_val is not None and str(cell_val).strip() != '':
                    present_fields.add(field_name)

            profile_defaults = {
                'first_name': first_name,
                'last_name': last_name,
                'father_name': get_cell_str(r, 'father_name'),
                'birth_date': normalize_date_val(get_cell_str(r, 'birth_date')),
                'id_number': get_cell_str(r, 'id_number'),
                'birth_place': get_cell_str(r, 'birth_place'),
                'education_level': get_cell_str(r, 'education_level', '4'),
                'insurance_type': get_cell_str(r, 'insurance_type', '2'),
                'insurance_number': get_cell_str(r, 'insurance_number'),
                'insurance_name': get_cell_str(r, 'insurance_name', 'تامین اجتماعی'),
                'exemption_type': get_cell_str(r, 'exemption_type', '1'),
                'nationality_code': get_cell_str(r, 'nationality_code', '1'),
                'citizenship_country_code': get_cell_str(r, 'citizenship_country_code', '103'),
                'residence_country_code': get_cell_str(r, 'residence_country_code', '103'),
                'job_category': get_cell_str(r, 'job_category', '15'),
                # DEF-22: «ي/ك» عربی و فاصله‌های فایل مرجع یکدست می‌شوند (مثل «مديرپروژه»)
                'job_title': normalize_job_title(get_cell_str(r, 'job_title')) or 'پرسنل انبار',
                'employment_type': get_cell_str(r, 'employment_type', '2'),
                'start_date': normalize_date_val(get_cell_str(r, 'start_date')),
                'end_date': normalize_date_val(get_cell_str(r, 'end_date')),
                'retirement_date': normalize_date_val(get_cell_str(r, 'retirement_date')),
                'job_code': clean_jc,
                'id_series': get_cell_str(r, 'id_series'),
                'id_serial': get_cell_str(r, 'id_serial'),
                'issue_place': get_cell_str(r, 'issue_place'),
                'issue_date': normalize_date_val(get_cell_str(r, 'issue_date')),
                'gender': get_cell_str(r, 'gender', 'مرد'),
                'phone_number': clean_phone,
                'account_number': get_cell_str(r, 'account_number'),
                'sheba_number': get_cell_str(r, 'sheba_number'),
                'bank_name': get_cell_str(r, 'bank_name'),
                'address': get_cell_str(r, 'address'),
                'postal_code': get_cell_str(r, 'postal_code'),
                'marital_status': 'married' if 'متاهل' in get_cell_str(r, 'marital_status') or 'متأهل' in get_cell_str(r, 'marital_status') else 'single',
                'children_count': int(get_cell_num(r, 'children_count', 0)),
                'contract_hours': get_cell_num(r, 'contract_hours', 230),
                'contract_base_salary': get_cell_num(r, 'contract_base_salary', 0),
                'job_grade': get_cell_str(r, 'job_grade', '19'),
                'daily_base_wage': get_cell_num(r, 'daily_base_wage', 0),
                'daily_seniority_bonus': get_cell_num(r, 'daily_seniority_bonus', 0),
                'base_daily_rate': get_cell_num(r, 'base_daily_rate', 0),
                'base_years_experience': int(get_cell_num(r, 'base_years_experience', 0)),
                'status_category': get_cell_str(r, 'status_category', 'نفرات شرکتی'),
                'group_status': get_cell_str(r, 'group_status', 'شاغل'),
                'include_in_tax': get_cell_bool(r, 'include_in_tax', True),
                'include_in_insurance': get_cell_bool(r, 'include_in_insurance', True),
                'include_in_bank': get_cell_bool(r, 'include_in_bank', True),
                'contract_type': 'daily',
                'created_by': request.user if request.user and request.user.is_authenticated else None
            }
            if active_cid:
                profile_defaults['company_id'] = active_cid

            if warehouse_id and str(warehouse_id).isdigit():
                profile_defaults['assigned_warehouse_id'] = int(warehouse_id)
            if sec_obj is not None:
                profile_defaults['section_id'] = sec_obj.id
                profile_defaults['project_id'] = sec_obj.project_id
                if not profile_defaults.get('company_id') and sec_obj.project.company_id:
                    profile_defaults['company_id'] = sec_obj.project.company_id

            valid_rows_data.append((r, nc_clean, profile_defaults, present_fields))

        # پروندهٔ موجود با همان قلمرو دید پیدا می‌شود: پروندهٔ بی‌شرکتی که بخشش در همین شرکت است هم «موجود» است
        # (وگرنه درون‌ریزی برای آن پرونده‌ای تکراری در شرکت می‌ساخت)
        existing_qs = PersonnelProfile.objects.filter(national_code__in=[x[1] for x in valid_rows_data])
        if active_cid:
            existing_qs = existing_qs.filter(
                Q(company_id=active_cid) | Q(company__isnull=True, section__project__company_id=active_cid)
            )
        existing_by_code = {}
        ambiguous_codes = set()
        for prof in existing_qs.order_by('company_id', 'id'):  # در Postgres مقدار تهی در ترتیب صعودی آخر می‌آید
            if prof.national_code in existing_by_code:
                if not active_cid or existing_by_code[prof.national_code].company_id == prof.company_id:
                    ambiguous_codes.add(prof.national_code)
                continue
            existing_by_code[prof.national_code] = prof

        for r, nc_clean, p_defaults, _present in valid_rows_data:
            exists = nc_clean in existing_by_code
            full_name = f"{p_defaults['first_name']} {p_defaults['last_name']}".strip()

            if exists:
                if update_existing:
                    updated_count += 1
                else:
                    skipped_count += 1
            else:
                created_count += 1

            if len(preview_rows) < 30:
                preview_rows.append({
                    'row': r,
                    'name': full_name,
                    'username': nc_clean,
                    'national_code': nc_clean,
                    'roles': p_defaults.get('job_title') or 'پرسنل انبار',
                    'warehouses': f"بخش {section_id}" if section_id else "عمومی",
                    'is_update': exists
                })

        import_warnings = []
        if not dry_run:
            user = request.user
            from decimal import Decimal, InvalidOperation
            from .lifecycle_services import (
                FINANCE_STAGE_PERMS,
                MANAGER_STAGE_PERMS,
                WAGE_FIELDS,
                _snapshot_value,
                auto_pass_personnel,
                user_can_view_wages,
                user_has_any_perm,
            )
            is_mgr = user_has_any_perm(user, MANAGER_STAGE_PERMS)
            is_fin = user_has_any_perm(user, FINANCE_STAGE_PERMS)
            can_set_wages = user_can_view_wages(user)
            created_by_manager = []
            scope_keys = ('section_id', 'project_id', 'assigned_warehouse_id')

            def _norm(value):
                if value is None or value == '':
                    return None
                if isinstance(value, bool):
                    return value
                if isinstance(value, (int, float, Decimal)):
                    try:
                        return Decimal(str(value)).normalize()
                    except InvalidOperation:
                        return str(value)
                return str(value).strip()

            with transaction.atomic():
                for r, nc_clean, p_defaults, present_fields in valid_rows_data:
                    existing_prof = existing_by_code.get(nc_clean)
                    if existing_prof is not None and not update_existing:
                        continue

                    if not can_set_wages:
                        # DEF-21: ستون‌های دستمزد را فقط نقش مالی درون‌ریزی می‌کند
                        for wage_field in WAGE_FIELDS:
                            p_defaults.pop(wage_field, None)

                    if existing_prof is None:
                        # DEF-14: ردیف تازه همیشه پیش‌نویس ساخته می‌شود؛ درون‌ریزی مدیر سپس با گیت عبور هوشمند می‌کند
                        p_defaults['approval_status'] = 'draft'
                        prof = PersonnelProfile.objects.create(national_code=nc_clean, **p_defaults)
                        if is_mgr:
                            created_by_manager.append((r, prof))
                        continue

                    full_name = f"{p_defaults.get('first_name', '')} {p_defaults.get('last_name', '')}".strip()
                    if nc_clean in ambiguous_codes:
                        import_warnings.append({
                            'row': r, 'field': 'کد ملی', 'national_code': nc_clean,
                            'message': f'برای «{full_name}» بیش از یک پرونده با این کد ملی هست؛ ابتدا شرکت فعال را انتخاب کنید. این ردیف به‌روز نشد.',
                        })
                        updated_count -= 1
                        skipped_count += 1
                        continue

                    # همان قاعدهٔ ویرایش مستقیم: پروندهٔ مصوب، در انتظار مدیر یا ردشده از درون‌ریزی به‌روز نمی‌شود
                    # (DEF-17؛ تغییر پروندهٔ مصوب فقط از مسیر درخواست تغییرات است)
                    st = existing_prof.approval_status
                    editable = (
                        user.is_superuser
                        or st in ('draft', 'revision_required', 'pending_supervisor')
                        or (st == 'pending_accountant' and is_fin)
                    )
                    if not editable:
                        import_warnings.append({
                            'row': r, 'field': 'وضعیت پرونده', 'national_code': nc_clean,
                            'message': f'پرونده «{existing_prof.full_name}» در وضعیت «{existing_prof.get_approval_status_display()}» است و از درون‌ریزی به‌روز نشد؛ تغییر آن از مسیر درخواست تغییرات انجام می‌شود.',
                        })
                        updated_count -= 1
                        skipped_count += 1
                        continue

                    # فقط ستون‌های حاضر و غیرخالی همین ردیف نوشته می‌شوند؛ ستون غایب یا خالی مقدار فعلی را پاک نمی‌کند
                    updates = {
                        k: v for k, v in p_defaults.items()
                        if (k in present_fields and v not in ('', None)) or k in scope_keys
                    }
                    diff, old_diff = {}, {}
                    for k, v in updates.items():
                        old_v = getattr(existing_prof, k, None)
                        if _norm(old_v) != _norm(v):
                            diff[k] = _snapshot_value(v)
                            old_diff[k] = _snapshot_value(old_v)
                    if not diff:
                        continue
                    for k in diff:
                        setattr(existing_prof, k, updates[k])
                    existing_prof.save()
                    if st not in ('draft', 'revision_required'):
                        WorkflowAuditLog.objects.create(
                            content_type='personnel',
                            object_id=existing_prof.id,
                            actor=user,
                            from_status=st,
                            to_status=existing_prof.approval_status,
                            action='excel_import_update',
                            reason='به‌روزرسانی پرونده از درون‌ریزی اکسل',
                            metadata={'diff': diff, 'old_diff': old_diff, 'row': r, 'company_id': existing_prof.company_id}
                        )

                # DEF-14: گیت برای هر ردیف؛ ردیف ناقص با گزارش، پیش‌نویس می‌ماند و عبور هوشمند لاگ auto_pass دارد
                for r, prof in created_by_manager:
                    result = auto_pass_personnel(prof, user, source='excel_import', draft_on_incomplete=True)
                    if result['status'] == 'draft':
                        import_warnings.append({
                            'row': r,
                            'field': 'گیت تکمیل پرونده',
                            'national_code': prof.national_code,
                            'message': f"پرونده «{prof.full_name}» به‌دلیل نقص اطلاعات پیش‌نویس ماند: "
                                       + '، '.join(m['label'] for m in result['missing_fields']),
                        })

            broadcast_personnel_update(None, action_type='imported', message=f'درون‌ریزی پرسنل انجام شد ({created_count} ایجاد، {updated_count} به‌روزرسانی).', sender_id=request.user.id if request.user else None)

        valid_count = created_count + updated_count
        total_rows = valid_count + len(errors) + skipped_count

        return Response({
            'success': len(errors) == 0 or valid_count > 0,
            'dry_run': dry_run,
            'message': f'درون‌ریزی پرسنل با موفقیت انجام شد ({created_count} پرسنل جدید، {updated_count} به‌روزرسانی).',
            'created_count': created_count,
            'updated_count': updated_count,
            'summary': {
                'total_rows': total_rows,
                'created': created_count,
                'updated': updated_count,
                'skipped': len(errors) + skipped_count,
                'valid_count': valid_count,
                'error_count': len(errors),
            },
            'preview_rows': preview_rows,
            'errors': errors,
            'warnings': import_warnings,
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی رسمی ۴۶ ستونی اطلاعات پرسنل منطبق بر شیت Emp_info و ساختار ۲ سطری
        """
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse

        qs = self.get_queryset()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Emp_info"
        ws.views.sheetView[0].rightToLeft = True

        columns_spec = [
            ('status_category', 'وضعیت'),
            ('row_num', 'ردیف'),
            ('nationality_code', 'ملیت'),
            ('national_code', 'کد ملی'),
            ('first_name', 'نام'),
            ('last_name', 'نام خانوادگی'),
            ('father_name', 'نام پدر'),
            ('birth_date', 'تاریخ تولد'),
            ('id_number', 'شماره شناسنامه'),
            ('birth_place', 'محل تولد'),
            ('education_level', 'مدرک تحصیلی'),
            ('insurance_type', 'نوع بیمه'),
            ('insurance_number', 'شماره بیمه'),
            ('insurance_name', 'نام بیمه'),
            ('exemption_type', 'نوع معافیت'),
            ('citizenship_country_code', 'کشور محل تابعیت'),
            ('residence_country_code', 'کشور محل زندگی'),
            ('postal_code', 'کد پستی'),
            ('address', 'آدرس'),
            ('job_category', 'رسته'),
            ('job_title', 'سمت شغل'),
            ('employment_type', 'نوع استخدام'),
            ('start_date', 'تاریخ شروع به کار'),
            ('end_date', 'تاریخ پایان کار'),
            ('retirement_date', 'تاریخ بازنشستگی'),
            ('job_code', 'کد شغل'),
            ('id_series', 'مسلسل شناسنامه'),
            ('id_serial', 'سریال'),
            ('issue_place', 'محل صدور'),
            ('issue_date', 'تاریخ صدور'),
            ('gender', 'جنسیت'),
            ('phone_number', 'موبایل'),
            ('account_number', 'شماره حساب'),
            ('marital_status', 'وضعیت تاهل'),
            ('children_count', 'تعداد فرزند'),
            ('contract_hours', 'ساعت کار قرارداد شده'),
            ('contract_base_salary', 'حقوق قرارداد شده'),
            ('job_grade', 'گروه شغلی'),
            ('daily_base_wage', 'مزد روزانه'),
            ('daily_seniority_bonus', 'پایه سنواتی'),
            ('base_years_experience', 'تعداد سال کارکرد'),
            ('base_daily_rate', 'مزد مبنا'),
            ('group_status', 'گروه'),
            ('include_in_tax', 'maliat'),
            ('include_in_insurance', 'بیمه'),
            ('include_in_bank', 'bank'),
        ]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        headers_fa = [c[1] for c in columns_spec]
        headers_en = [c[0] for c in columns_spec]
        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(columns_spec) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.freeze_panes = 'A3'

        persian_arabic_digits = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
        def format_date_str(val):
            if not val:
                return ''
            s = str(val).translate(persian_arabic_digits).strip()
            if not s or s.lower() == 'none':
                return ''
            s = s.split('.')[0]
            if len(s) == 8 and s.isdigit():
                return f"{s[:4]}/{s[4:6]}/{s[6:]}"
            s = s.replace('-', '/')
            parts = s.split('/')
            if len(parts) == 3 and len(parts[0]) == 4 and parts[1].isdigit() and parts[2].isdigit():
                return f"{parts[0]}/{int(parts[1]):02d}/{int(parts[2]):02d}"
            return s

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, p in enumerate(qs, start=1):
            curr_row = idx + 2
            row_data = [
                p.status_category or 'نفرات شرکتی',
                idx,
                p.nationality_code or '1',
                p.national_code or '',
                p.first_name or '',
                p.last_name or '',
                p.father_name or '',
                format_date_str(p.birth_date),
                p.id_number or '',
                p.birth_place or '',
                p.education_level or '4',
                p.insurance_type or '2',
                p.insurance_number or '',
                p.insurance_name or 'تامین اجتماعی',
                p.exemption_type or '1',
                p.citizenship_country_code or '103',
                p.residence_country_code or '103',
                p.postal_code or '',
                p.address or '',
                p.job_category or '15',
                p.job_title or '',
                p.employment_type or '2',
                format_date_str(p.start_date),
                format_date_str(p.end_date),
                format_date_str(p.retirement_date),
                str(p.job_code).zfill(6) if p.job_code else '',
                p.id_series or '',
                p.id_serial or '',
                p.issue_place or '',
                format_date_str(p.issue_date),
                p.gender or 'مرد',
                p.phone_number or getattr(p, 'mobile_phone', '') or '',
                p.account_number or '',
                'متاهل' if p.marital_status == 'married' else 'مجرد',
                int(p.children_count or 0),
                float(p.contract_hours or 230),
                float(p.contract_base_salary or 0),
                p.job_grade or '19',
                float(p.daily_base_wage or 0),
                float(p.daily_seniority_bonus or 0),
                int(p.base_years_experience or 0),
                float(p.base_daily_rate or 0),
                p.group_status or 'شاغل',
                1 if p.include_in_tax else 0,
                1 if p.include_in_insurance else 0,
                1 if p.include_in_bank else 0,
            ]
            ws.append(row_data)

            for col_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal='center' if col_idx not in [5, 6, 19] else 'right', vertical='center')
                if col_idx in [2, 4, 8, 9, 23, 24, 25, 26, 30, 32, 33, 35, 36, 37, 39, 40, 41, 42]:
                    c.font = data_font_num
                else:
                    c.font = data_font
                if col_idx in [37, 39, 40, 42]:
                    c.number_format = '#,##0'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"personnel_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response



def log_vehicle_audit(user, action, target_instance, details=None):
    try:
        AuditLog = apps.get_model('accounts', 'AuditLog')
        AuditLog.objects.create(
            user=user if user and user.is_authenticated else None,
            actor_username=user.username if user and user.is_authenticated else 'system',
            actor_name=f"{user.first_name} {user.last_name}".strip() or user.username if user and user.is_authenticated else 'سیستم',
            module='fleet',
            action=action,
            severity='info' if action not in ['REJECT', 'REJECT_CHANGE_REQ', 'DELETE'] else 'warning',
            target_model=target_instance.__class__.__name__,
            target_object_id=str(target_instance.pk),
            target_repr=str(target_instance),
            details=details or {}
        )
    except Exception:
        pass


class VehicleDriverProfileViewSet(viewsets.ModelViewSet):
    queryset = VehicleDriverProfile.objects.all().select_related(
        'user',
        'supervisor_approved_by',
        'accountant_approved_by',
        'manager_approved_by',
        'treasury_paid_by',
        'revision_requested_by',
        'auto_passed_by',
        'created_by',
        'section',
        'project',
    ).prefetch_related('change_requests', 'bank_accounts', 'documents')
    serializer_class = VehicleDriverProfileSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None



    def get_queryset(self):
        qs = super().get_queryset()
        qs = scope_vehicle_queryset(self.request, qs)

        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True))
            
        section_id = self.request.query_params.get('section_id')
        if section_id:
            if not is_vehicle_global_auditor(self.request.user) and self.request.user and self.request.user.is_authenticated:
                user_section_ids = list(
                    UserSectionAssignment.objects.filter(user=self.request.user, is_active=True).values_list('section_id', flat=True)
                )
                if str(section_id).isdigit() and int(section_id) not in user_section_ids:
                    return qs.none()
            qs = qs.filter(section_id=section_id)
        project_id = self.request.query_params.get('project_id')
        if project_id:
            qs = qs.filter(section__project_id=project_id)

        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() in ['true', '1'])

        approval_status = self.request.query_params.get('approval_status')
        if approval_status:
            if ',' in approval_status:
                status_list = [s.strip() for s in approval_status.split(',') if s.strip()]
                qs = qs.filter(approval_status__in=status_list)
            else:
                qs = qs.filter(approval_status=approval_status)

        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(
                Q(driver_name__icontains=search) |
                Q(plate_number__icontains=search) |
                Q(driver_national_code__icontains=search) |
                Q(driver_phone__icontains=search)
            )
        return qs

    def create(self, request, *args, **kwargs):
        data = request.data.copy() if hasattr(request.data, 'copy') else dict(request.data)
        section_id = data.get('section')
        if not data.get('company') and section_id:
            try:
                sec = ProjectSection.objects.filter(id=int(section_id)).select_related('project').first()
                if sec and sec.project and sec.project.company_id:
                    data['company'] = sec.project.company_id
            except (ValueError, TypeError):
                pass
        if not data.get('company'):
            cid = get_request_company_id(request)
            if cid:
                valid_cid = validate_user_company_access(request.user, cid)
                if valid_cid:
                    data['company'] = valid_cid
        serializer = self.get_serializer(data=data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        instance = serializer.instance
        if instance:
            instance.refresh_from_db()
            serializer = self.get_serializer(instance)
        headers = self.get_success_headers(serializer.data)
        return Response(serializer.data, status=status.HTTP_201_CREATED, headers=headers)

    def perform_create(self, serializer):
        user = self.request.user
        is_mgr = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or user.has_perm('accounts.can_act_as_manager')
        
        section = serializer.validated_data.get('section')
        # گارد انتساب بخش: کاربر عادی باید انتساب فعال به بخش انتخاب‌شده داشته باشد
        if not is_mgr and section and user and user.is_authenticated:
            is_assigned = UserSectionAssignment.objects.filter(user=user, section=section, is_active=True).exists()
            if not is_assigned:
                raise PermissionDenied("شما انتساب فعال در این بخش برای معرفی خودرو جدید ندارید.")

        # سنجش قلمرو شرکتی نوشتن (Task 0.3):
        ensure_record_write_scope(
            user,
            company=serializer.validated_data.get('company'),
            section=section,
            project=serializer.validated_data.get('project'),
            label='ناوگان',
        )

        extra_kwargs = {'created_by': user}
        if section and hasattr(section, 'project') and section.project:
            if not serializer.validated_data.get('project'):
                extra_kwargs['project'] = section.project
            if not serializer.validated_data.get('company') and section.project.company_id:
                extra_kwargs['company_id'] = section.project.company_id

        # تسک ۳.۰ و ب-۱۴: ایجاد اولیه همیشه در وضعیت draft، و عبور فقط با auto_pass_vehicle
        extra_kwargs['approval_status'] = 'draft'
        instance = serializer.save(**extra_kwargs)
        from .fleet_lifecycle_services import auto_pass_vehicle
        auto_pass_vehicle(instance, user, source='create')
        broadcast_vehicle_update(instance, action_type='created', message=f'پرونده خودرو «{instance.plate_number}» با راننده «{instance.driver_name}» ثبت شد.', sender_id=user.id if user else None)

    def _ensure_update_scope(self, instance, validated_data):
        """جابه‌جایی شرکت، بخش یا پروژهٔ پرونده فقط درون شرکت‌های مجاز کاربر (پرونده به شرکت بیگانه منتقل نمی‌شود)."""
        def _changed(field):
            return field in validated_data and getattr(validated_data[field], 'pk', None) != getattr(instance, f'{field}_id')

        if not any(_changed(f) for f in ('company', 'section', 'project')):
            return
        ensure_record_write_scope(
            self.request.user,
            company=validated_data.get('company', instance.company),
            section=validated_data.get('section', instance.section),
            project=validated_data['project'] if _changed('project') else None,
            label='ناوگان',
        )

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        # ب-۱ (Task 0.1): عبور از قلمرو و واکشی شیء پیش از قفل ردیف درون تراکنش
        target = self.get_object()
        user = request.user
        is_mgr = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or user.has_perm('accounts.can_act_as_manager')
        is_fin = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_finance')

        from .fleet_lifecycle_services import (
            VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS,
            compute_profile_diff,
            log_vehicle_change_request_event,
            get_vehicle_missing_fields,
            dispatch_vehicle_workflow_notification,
            refresh_vehicle_has_pending_changes,
        )

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)

            # پروندهٔ ردشده بایگانی است و پیش از هر ویرایش باید با اکشن reopen بازگشایی شود (تسک ۱۱.۱، ن-۱، الگوی پرسنل)
            if instance.approval_status == 'rejected' and not user.is_superuser:
                return Response(
                    {'error': 'این پرونده رد شده است. برای ویرایش و ارسال دوباره، ابتدا آن را بازگشایی کنید.', 'code': 'reopen_required'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # ۱. پرونده در انتظار تصویب مدیر: برای همه جز سوپریوزر قفل است (تسک ۳.۷)
            if instance.approval_status in ['pending_manager', 'accountant_approved']:
                if not user.is_superuser:
                    return Response(
                        {'error': f'این پرونده در وضعیت «{instance.get_approval_status_display()}» قرار دارد و اطلاعات آن قفل است؛ در صورت نیاز، مدیر پرونده را به حسابدار عودت دهد.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )

            # ۲. پرونده در انتظار بررسی حسابدار (pending_accountant):
            # ویرایش مستقیم فقط برای حسابدار یا سوپریوزر مجاز است و به صورت accountant_edit لاگ می‌شود (تسک ۳.۷)
            if instance.approval_status in ['pending_accountant', 'supervisor_approved']:
                if not (user.is_superuser or is_fin):
                    return Response(
                        {'error': f'این پرونده در وضعیت «{instance.get_approval_status_display()}» قرار دارد و اطلاعات آن تا پایان بررسی حسابداری قابل تغییر مستقیم توسط کارمند نیست.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
                serializer = self.get_serializer(instance, data=request.data, partial=partial)
                serializer.is_valid(raise_exception=True)
                self._ensure_update_scope(instance, serializer.validated_data)
                diff, old_diff = compute_profile_diff(instance, serializer.validated_data, excluded=VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS)
                updated_instance = serializer.save()
                if diff:
                    WorkflowAuditLog.objects.create(
                        content_type='vehicle',
                        object_id=updated_instance.id,
                        actor=user,
                        from_status=updated_instance.approval_status,
                        to_status=updated_instance.approval_status,
                        action='accountant_edit',
                        reason='ویرایش مستقیم اطلاعات پرونده خودرو توسط حسابدار',
                        metadata={'diff': diff, 'old_diff': old_diff, 'company_id': updated_instance.company_id}
                    )
                broadcast_vehicle_update(updated_instance, action_type='updated', message=f'اطلاعات مالی خودرو «{updated_instance.plate_number}» توسط حسابدار به‌روزرسانی شد.', sender_id=user.id if user else None)
                return Response(self.get_serializer(updated_instance).data)

            # ۳. پرونده مصوب: سوپریوزر ویرایش مستقیم دارد (با لاگ superuser_edit)؛ سایر کاربران به کارتابل درخواست تغییرات هدایت می‌شوند
            if instance.approval_status in ['approved', 'manager_approved']:
                if user.is_superuser:
                    serializer = self.get_serializer(instance, data=request.data, partial=partial)
                    serializer.is_valid(raise_exception=True)
                    self._ensure_update_scope(instance, serializer.validated_data)
                    diff, old_diff = compute_profile_diff(instance, serializer.validated_data, excluded=VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS)
                    updated_instance = serializer.save()
                    if diff:
                        WorkflowAuditLog.objects.create(
                            content_type='vehicle',
                            object_id=updated_instance.id,
                            actor=user,
                            from_status=updated_instance.approval_status,
                            to_status=updated_instance.approval_status,
                            action='superuser_edit',
                            reason='ویرایش مستقیم پرونده مصوب خودرو توسط سوپریوزر',
                            metadata={'diff': diff, 'old_diff': old_diff, 'company_id': updated_instance.company_id}
                        )
                    broadcast_vehicle_update(updated_instance, action_type='updated', message=f'پرونده مصوب خودرو «{updated_instance.plate_number}» توسط سوپریوزر ویرایش شد.', sender_id=user.id if user else None)
                    return Response(self.get_serializer(updated_instance).data)

                # جلوگیری از ایجاد درخواست تغییرات تکراری تا زمان تعیین تکلیف قبلی (CONCUR-01)
                has_pending = VehicleChangeRequest.objects.filter(
                    vehicle=instance,
                    status__in=['pending_supervisor', 'pending_accountant', 'pending_manager', 'supervisor_approved', 'accountant_approved', 'manager_approved']
                ).exists()
                if has_pending:
                    return Response(
                        {'error': 'این خودرو دارای درخواست تغییرات در حال بررسی است. تا زمان تعیین تکلیف درخواست قبلی امکان ثبت تغییرات جدید وجود ندارد.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )

                from decimal import Decimal, InvalidOperation
                from .fleet_lifecycle_services import vehicle_change_request_fields_allowed_for

                allowed_change_fields = vehicle_change_request_fields_allowed_for(user)
                # تسک ۱۱.۳ (ن-۳، ط-۲): تلاش صریح نقش غیرمالی برای تغییر نرخ خودروی مصوب رد می‌شود
                if allowed_change_fields is not None and 'default_service_rate' in request.data:
                    raw_new_rate = request.data.get('default_service_rate')
                    if raw_new_rate not in (None, '', 0, '0'):
                        try:
                            new_rate_dec = Decimal(str(raw_new_rate))
                            old_rate_dec = Decimal(str(instance.default_service_rate or 0))
                            if new_rate_dec > 0 and new_rate_dec != old_rate_dec:
                                return Response(
                                    {'error': 'تغییر نرخ پایه سرویس ناوگان مصوب منحصراً در صلاحیت نقش حسابداری ناوگان است.', 'code': 'forbidden_rate_change'},
                                    status=status.HTTP_403_FORBIDDEN
                                )
                        except (InvalidOperation, ValueError):
                            pass

                serializer = self.get_serializer(instance, data=request.data, partial=partial)
                serializer.is_valid(raise_exception=True)
                self._ensure_update_scope(instance, serializer.validated_data)
                
                diff, old_diff = compute_profile_diff(instance, serializer.validated_data, excluded=VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS)
                if allowed_change_fields is not None:
                    diff = {k: v for k, v in diff.items() if k in allowed_change_fields}
                    old_diff = {k: v for k, v in old_diff.items() if k in diff}

                if diff:
                    cr = VehicleChangeRequest.objects.create(
                        vehicle=instance,
                        requested_by=user,
                        proposed_changes=diff,
                        previous_values=old_diff,
                        status='pending_supervisor'
                    )
                    instance.has_pending_changes = True
                    instance.save(update_fields=['has_pending_changes'])
                    log_vehicle_change_request_event(cr, user, 'change_request_created', reason='ثبت درخواست تغییرات پرونده مصوب خودرو')
                    # تسک ۱۱.۱۰ (ن-۱۰): ارسال اعلان به سرپرست
                    dispatch_vehicle_workflow_notification(instance, stage='change_request_created', actor=user)
                    return Response({
                        'message': 'درخواست تغییرات خودرو با موفقیت ثبت شد و جهت بررسی در کارتابل سرپرست قرار گرفت. اطلاعات قبلی تا زمان تصویب نهایی معتبر باقی می‌ماند.',
                        'change_request_id': cr.id,
                        'data': self.get_serializer(instance).data
                    }, status=status.HTTP_202_ACCEPTED)
                return Response(self.get_serializer(instance).data)

            # ۴. ویرایش رکوردهای پیش‌نویس یا عودت‌داده‌شده:
            serializer = self.get_serializer(instance, data=request.data, partial=partial)
            serializer.is_valid(raise_exception=True)
            self._ensure_update_scope(instance, serializer.validated_data)

            extra_kwargs = {}
            section = serializer.validated_data.get('section', instance.section)
            if section and hasattr(section, 'project') and section.project:
                if not serializer.validated_data.get('project') and not instance.project_id:
                    extra_kwargs['project'] = section.project

            # تسک ۲.۷، ب-۱۳ و تسک ۱۱.۱ (ن-۱، ط-۱): پل سازگاری PATCH برای وضعیت pending_supervisor
            is_submit_attempt = request.data.get('approval_status') == 'pending_supervisor'
            can_bridge_transition = instance.approval_status in ('draft', 'revision_required')
            updated_instance = serializer.save(**extra_kwargs)

            submitted = False
            missing = []
            if is_submit_attempt:
                if can_bridge_transition:
                    missing = get_vehicle_missing_fields(updated_instance, stage='submit')
                    if not missing:
                        old_st = updated_instance.approval_status
                        updated_instance.approval_status = 'pending_supervisor'
                        updated_instance.rejection_reason = None
                        updated_instance.save(update_fields=['approval_status', 'rejection_reason'])
                        submitted = True
                        WorkflowAuditLog.objects.create(
                            content_type='vehicle',
                            object_id=updated_instance.id,
                            actor=user,
                            from_status=old_st,
                            to_status='pending_supervisor',
                            action='submit',
                            reason='ارسال پرونده خودرو به سرپرست جهت بررسی',
                            metadata={'company_id': updated_instance.company_id}
                        )
                        dispatch_vehicle_workflow_notification(updated_instance, stage='pending_supervisor', actor=user)
                else:
                    submitted = False

            broadcast_vehicle_update(updated_instance, action_type='updated', message=f'پرونده خودرو «{updated_instance.plate_number}» به‌روزرسانی شد.', sender_id=user.id if user else None)
            res_data = self.get_serializer(updated_instance).data
            if is_submit_attempt:
                res_data['submitted'] = submitted
                if not submitted:
                    if can_bridge_transition:
                        res_data['missing_fields'] = missing
                        res_data['detail'] = f"تغییرات ذخیره شد، اما به دلیل نقص اطلاعات ارسال به سرپرست انجام نشد: {', '.join(m['label'] for m in missing)}"
                    else:
                        res_data['detail'] = f"تغییرات ذخیره شد، اما گذار وضعیت از «{instance.get_approval_status_display()}» به «در انتظار سرپرست» مجاز نیست."
            return Response(res_data)

    @staticmethod
    def _user_may_delete(user, instance) -> bool:
        if not user or not user.is_authenticated:
            return False
        if user.is_superuser or user.has_perm('accounts.perm_delete_vehicle_nomovement'):
            return True
        if instance.approval_status not in ('draft', 'revision_required'):
            return False
        return VehicleDriverProfileViewSet._is_creator_or_co_section(user, instance)

    @staticmethod
    def _is_creator_or_co_section(user, instance) -> bool:
        if instance.created_by_id == user.id:
            return True
        return bool(instance.section_id) and UserSectionAssignment.objects.filter(
            user=user, section_id=instance.section_id, is_active=True
        ).exists()

    def destroy(self, request, *args, **kwargs):
        """
        تسک ۰.۶، ۱.۲، ۲.۳، ۳.۹، ب-۹ و ب-۱۰: حذف خودرو با قفل ردیف و سنجش سابقه درون تراکنش:
        - اگر خودرو دارای سابقه کارکرد یا تردد، رویداد مالی یا درخواست تغییرات باشد،
          حذف فیزیکی مسدود است و ۴۰۰ با code='history_exists' و پیشنهاد قطع همکاری برمی‌گردد.
        - اگر فاقد سابقه باشد، ثبت‌کننده یا هم‌بخش در draft/revision_required یا دارنده perm_delete_vehicle_nomovement مجاز به حذف است.
        - عکسبرداری کامل پرونده (اسنپ‌شات) در WorkflowAuditLog ثبت می‌گردد.
        """
        from .fleet_lifecycle_services import check_vehicle_history, build_vehicle_snapshot, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            history = check_vehicle_history(instance)
            if history['has_history']:
                return Response({
                    'error': 'این خودرو دارای سابقه کارکرد، رویداد مالی یا درخواست تغییرات در سامانه است و حذف فیزیکی آن مجاز نمی‌باشد. لطفاً از طریق ثبت تاریخ پایان همکاری و غیرفعال‌سازی پرونده اقدام به قطع همکاری نمایید.',
                    'code': 'history_exists',
                    'suggested_action': 'end_collaboration',
                    'history_reasons': history['reasons'],
                    'history_details': history['details']
                }, status=status.HTTP_400_BAD_REQUEST)

            if not self._user_may_delete(user, instance):
                if self._is_creator_or_co_section(user, instance):
                    return Response(
                        {'error': 'فقط پرونده‌های در وضعیت پیش‌نویس یا عودت‌داده‌شده قابل حذف هستند.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
                raise PermissionDenied("شما مجوز حذف این پرونده خودرو را ندارید.")

            snapshot = build_vehicle_snapshot(instance)
            plate_number = instance.plate_number
            cid = instance.company_id
            target_id = instance.id
            approval_status = instance.approval_status
            broadcast_vehicle_update(instance, action_type='deleted', message=f'خودرو «{plate_number}» حذف شد.', sender_id=user.id if user else None)
            dispatch_vehicle_workflow_notification(instance, stage='deleted', actor=user)
            self.perform_destroy(instance)
            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=target_id,
                actor=user,
                from_status=approval_status,
                to_status='deleted',
                action='delete',
                reason='حذف فیزیکی خودرو فاقد سابقه',
                metadata={'snapshot': snapshot, 'company_id': cid, 'plate_number': plate_number}
            )
            return Response({'message': f'خودرو «{plate_number}» با موفقیت از سامانه حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get'], url_path='delete_impact')
    def delete_impact(self, request, pk=None):
        """
        تسک ۵.۵: استعلام اثرات حذف آبشاری و سوابق ناوگان منطبق بر قرارداد smart-delete-modal و check_vehicle_history
        """
        instance = self.get_object()
        from django.db.models.deletion import Collector, ProtectedError, RestrictedError
        from django.db import router
        from .fleet_lifecycle_services import check_vehicle_history

        using = router.db_for_write(instance.__class__, instance=instance)
        collector = Collector(using=using)
        impact_summary = []

        try:
            collector.collect([instance])
            for model, instances in collector.data.items():
                if model == instance.__class__:
                    continue
                impact_summary.append({
                    'model_name': str(model._meta.verbose_name or model.__name__),
                    'count': len(instances)
                })
        except (ProtectedError, RestrictedError) as pe:
            prot_counts = {}
            for obj in getattr(pe, 'protected_objects', []):
                m = obj.__class__
                prot_counts[m] = prot_counts.get(m, 0) + 1
            for m, cnt in prot_counts.items():
                impact_summary.append({
                    'model_name': str(m._meta.verbose_name or m.__name__),
                    'count': cnt
                })
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        impact_summary.sort(key=lambda x: x['count'], reverse=True)

        history = check_vehicle_history(instance)
        user = request.user
        can_delete = bool(history['can_delete'] and self._user_may_delete(user, instance))

        return Response({
            'impact': impact_summary,
            'total_affected': sum(item['count'] for item in impact_summary),
            'history': history,
            'can_delete': can_delete,
            'has_history': history['has_history'],
            'reasons': history['reasons'],
            'blocking_reasons': history['reasons'],
            'history_details': history.get('details', {})
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get', 'post', 'put', 'patch', 'delete'], url_path='bank-accounts')
    def bank_accounts(self, request, pk=None):
        """
        تسک ۵.۱: مدیریت حساب‌های بانکی تودرتوی ناوگان با رعایت تفکیک وظایف (SoD).
        صاحب حساب همیشه همین خودرو است (vehicle فقط‌خواندنی)، و تغییر حساب با گارد وضعیت و نقش انجام می‌شود.
        """
        from .fleet_lifecycle_services import check_vehicle_bank_account_edit_allowed
        from .lifecycle_services import _snapshot_value
        from .serializers import VehicleBankAccountSerializer
        instance = self.get_object()

        if request.method == 'GET':
            accounts = instance.bank_accounts.all()
            serializer = VehicleBankAccountSerializer(accounts, many=True)
            return Response(serializer.data, status=status.HTTP_200_OK)

        allowed, deny_msg = check_vehicle_bank_account_edit_allowed(request.user, instance)
        if not allowed:
            return Response({'error': deny_msg}, status=status.HTTP_400_BAD_REQUEST)

        def _log_bank_change(op, account_id, diff=None):
            in_draft = instance.approval_status in ('draft', 'revision_required')
            if in_draft and op != 'delete':
                return
            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=request.user,
                from_status=instance.approval_status,
                to_status=instance.approval_status,
                action='bank_account_deleted' if in_draft else 'accountant_edit',
                reason=request.data.get('note', '') or 'تغییر حساب بانکی خودرو',
                metadata={'bank_account': {'op': op, 'id': account_id, 'diff': diff or {}}, 'company_id': instance.company_id}
            )

        if request.method == 'POST':
            serializer = VehicleBankAccountSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            with transaction.atomic():
                acc = serializer.save(vehicle=instance)
                _log_bank_change('create', acc.id, {k: str(v) for k, v in serializer.validated_data.items()})
            broadcast_vehicle_update(instance, action_type='updated', message=f'حساب بانکی جدید برای خودرو «{instance.plate_number}» افزوده شد.', sender_id=request.user.id)
            return Response(VehicleBankAccountSerializer(acc).data, status=status.HTTP_201_CREATED)

        account_id = request.data.get('id') or request.query_params.get('account_id') or request.query_params.get('id')
        if not account_id:
            return Response({'error': 'شناسه حساب بانکی (id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        acc = instance.bank_accounts.filter(pk=account_id).first()
        if not acc:
            return Response({'error': 'حساب بانکی موردنظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        if request.method in ['PUT', 'PATCH']:
            serializer = VehicleBankAccountSerializer(acc, data=request.data, partial=True)
            serializer.is_valid(raise_exception=True)
            diff = {
                k: {'old': str(getattr(acc, k)) if getattr(acc, k) is not None else None, 'new': str(v) if v is not None else None}
                for k, v in serializer.validated_data.items() if getattr(acc, k) != v
            }
            with transaction.atomic():
                updated_acc = serializer.save()
                if diff:
                    _log_bank_change('update', updated_acc.id, diff)
            broadcast_vehicle_update(instance, action_type='updated', message=f'حساب بانکی خودرو «{instance.plate_number}» ویرایش شد.', sender_id=request.user.id)
            return Response(VehicleBankAccountSerializer(updated_acc).data, status=status.HTTP_200_OK)

        if request.method == 'DELETE':
            with transaction.atomic():
                account_snapshot = {
                    f.attname: _snapshot_value(getattr(acc, f.attname))
                    for f in acc._meta.concrete_fields
                }
                _log_bank_change('delete', acc.id, account_snapshot)
                acc.delete()
            broadcast_vehicle_update(instance, action_type='updated', message=f'حساب بانکی خودرو «{instance.plate_number}» حذف شد.', sender_id=request.user.id)
            return Response({'message': 'حساب بانکی با موفقیت حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get', 'post', 'delete'], url_path='documents')
    def documents(self, request, pk=None):
        """
        تسک ۵.۲ و ب-۳۰: مدیریت و بارگذاری مدارک چندگانه ناوگان و خودرو با اعتبارسنجی نوع، پسوند و حجم فایل.
        در حذف مدرک، رکورد دیتابیس حذف می‌شود ولی فایل فیزیکی جهت ممیزی حفظ می‌گردد.
        """
        from .models import VehicleDocument
        from .serializers import VehicleDocumentSerializer, validate_uploaded_document
        from .fleet_lifecycle_services import check_vehicle_bank_account_edit_allowed
        instance = self.get_object()

        if request.method == 'GET':
            docs = instance.documents.all().order_by('-created_at')
            return Response(VehicleDocumentSerializer(docs, many=True).data, status=status.HTTP_200_OK)

        elif request.method == 'POST':
            file_obj = request.FILES.get('file') or request.FILES.get('attachment')
            if not file_obj:
                return Response({'error': 'فایل مدرک الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            doc_type = request.data.get('doc_type') or 'other'
            if doc_type not in dict(VehicleDocument.DOC_TYPE_CHOICES):
                return Response({'error': f'نوع مدرک «{doc_type}» نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
            try:
                validate_uploaded_document(
                    file_obj,
                    max_size_bytes=10 * 1024 * 1024,
                    allowed_exts=('.pdf', '.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.heic', '.heif'),
                )
            except serializers.ValidationError as e:
                return Response({'error': ' '.join(str(d) for d in e.detail)}, status=status.HTTP_400_BAD_REQUEST)

            from .fleet_lifecycle_services import normalize_shamsi_date
            title = (request.data.get('title') or '').strip() or None
            raw_issue_date = (request.data.get('issue_date') or '').strip()
            raw_expiry_date = (request.data.get('expiry_date') or '').strip()
            # تاریخ نامعتبر ۴۰۰ روشن می‌گیرد، نه خطای ۵۰۰ (normalize_shamsi_date برای ورودی نامعتبر ValueError می‌دهد)
            try:
                issue_date = normalize_shamsi_date(raw_issue_date) if raw_issue_date else None
                expiry_date = normalize_shamsi_date(raw_expiry_date) if raw_expiry_date else None
            except ValueError as e:
                return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

            with transaction.atomic():
                doc = VehicleDocument.objects.create(
                    vehicle=instance,
                    file=file_obj,
                    doc_type=doc_type,
                    title=title,
                    issue_date=issue_date,
                    expiry_date=expiry_date,
                    uploaded_by=request.user if request.user.is_authenticated else None
                )
                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=request.user if request.user.is_authenticated else None,
                    from_status=instance.approval_status,
                    to_status=instance.approval_status,
                    action='document_uploaded',
                    reason=request.data.get('note', '') or 'بارگذاری مدرک پیوست خودرو',
                    metadata={
                        'document': {
                            'id': doc.id,
                            'doc_type': doc.doc_type,
                            'title': doc.title,
                            'file_name': doc.file.name if doc.file else None,
                            'issue_date': doc.issue_date,
                            'expiry_date': doc.expiry_date,
                            'uploaded_by_id': doc.uploaded_by_id,
                        },
                        'company_id': instance.company_id
                    }
                )
            broadcast_vehicle_update(instance, action_type='updated', message=f'مدرک جدید برای خودرو «{instance.plate_number}» بارگذاری شد.', sender_id=request.user.id)
            return Response(VehicleDocumentSerializer(doc).data, status=status.HTTP_201_CREATED)

        elif request.method == 'DELETE':
            doc_id = request.data.get('doc_id') or request.query_params.get('doc_id') or request.query_params.get('id')
            if not doc_id or not str(doc_id).isdigit():
                return Response({'error': 'شناسه مدرک الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            allowed, deny_msg = check_vehicle_bank_account_edit_allowed(request.user, instance)
            if not allowed:
                return Response({'error': deny_msg.replace('حساب‌های بانکی', 'مدارک')}, status=status.HTTP_400_BAD_REQUEST)
            doc = instance.documents.filter(id=int(doc_id)).select_related('uploaded_by').first()
            if not doc:
                return Response({'error': 'مدرک مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
            with transaction.atomic():
                doc_snapshot = {
                    'id': doc.id,
                    'doc_type': doc.doc_type,
                    'title': doc.title,
                    'file_name': doc.file.name if doc.file else None,
                    'issue_date': doc.issue_date,
                    'expiry_date': doc.expiry_date,
                    'uploaded_by_id': doc.uploaded_by_id,
                }
                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=request.user,
                    from_status=instance.approval_status,
                    to_status=instance.approval_status,
                    action='document_deleted',
                    reason=request.data.get('note', '') or 'حذف مدرک پیوست خودرو',
                    metadata={'document': doc_snapshot, 'company_id': instance.company_id}
                )
                doc.delete()
            broadcast_vehicle_update(instance, action_type='updated', message=f'مدرک خودرو «{instance.plate_number}» حذف شد.', sender_id=request.user.id)
            return Response({'message': 'مدرک با موفقیت حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='submit')
    def submit(self, request, pk=None):
        """
        تسک ۲.۷: اکشن صریح ارسال پرونده خودرو به سرپرست جهت بررسی و تایید
        """
        from .fleet_lifecycle_services import get_vehicle_missing_fields, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user
        is_mgr = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or user.has_perm('accounts.can_act_as_manager')
        if not is_mgr and target.created_by_id != user.id:
            if not (target.section_id and UserSectionAssignment.objects.filter(user=user, section_id=target.section_id, is_active=True).exists()):
                raise PermissionDenied("شما دسترسی ارسال این پرونده به سرپرست را ندارید.")

        if target.approval_status not in ('draft', 'revision_required'):
            return Response(
                {'error': f'ارسال به سرپرست فقط برای پرونده‌های پیش‌نویس یا نیازمند بازنگری مجاز است (وضعیت فعلی: {target.get_approval_status_display()}).'},
                status=status.HTTP_400_BAD_REQUEST
            )

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            missing = get_vehicle_missing_fields(instance, stage='submit')
            if missing:
                return Response({
                    'error': 'فیلدهای الزامی پرونده خودرو جهت ارسال به سرپرست تکمیل نشده است.',
                    'missing_fields': missing,
                    'detail': f"فیلدهای ناقص: {', '.join(m['label'] for m in missing)}"
                }, status=status.HTTP_400_BAD_REQUEST)

            old_status = instance.approval_status
            instance.approval_status = 'pending_supervisor'
            instance.rejection_reason = None
            instance.save(update_fields=['approval_status', 'rejection_reason'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status=old_status,
                to_status='pending_supervisor',
                action='submit',
                reason=(request.data.get('note') or 'ارسال پرونده خودرو به سرپرست').strip(),
                metadata={'company_id': instance.company_id}
            )
            dispatch_vehicle_workflow_notification(instance, stage='pending_supervisor', actor=user)

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» به سرپرست ارسال شد.', sender_id=user.id if user else None)
        return Response({
            'message': 'پرونده خودرو با موفقیت جهت بررسی و تایید به کارتابل سرپرست ارسال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='approve-supervisor')
    def approve_supervisor(self, request, pk=None):
        """
        تسک ۳.۴: تایید مرحله اول (سرپرست) پرونده خودرو با استفاده از ماژول متمرکز گردش کار
        """
        from .fleet_lifecycle_services import approve_vehicle_stage, VehicleWorkflowError
        target = self.get_object()
        note = (request.data.get('note') or request.data.get('approval_note') or '').strip()
        try:
            instance = approve_vehicle_stage(target.id, request.user, stage='supervisor', note=note)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» به تایید سرپرست رسید و به حسابداری ارسال شد.', sender_id=request.user.id if request.user else None)
        return Response({
            'message': 'تایید سرپرست با موفقیت ثبت شد و پرونده خودرو به حسابداری ارسال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='approve-finance')
    def approve_finance(self, request, pk=None):
        """
        تسک ۳.۴: تایید مرحله دوم (مالی) پرونده خودرو همراه با گیت الزامی نرخ و شماره شبا
        """
        from .fleet_lifecycle_services import approve_vehicle_stage, VehicleWorkflowError
        target = self.get_object()
        note = (request.data.get('note') or request.data.get('approval_note') or '').strip()
        try:
            instance = approve_vehicle_stage(target.id, request.user, stage='finance', note=note)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» به تایید مالی رسید و به مدیر ارسال شد.', sender_id=request.user.id if request.user else None)
        return Response({
            'message': 'تایید مالی با موفقیت ثبت شد و پرونده خودرو جهت تصویب نهایی به کارتابل مدیر ارسال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='approve-manager')
    def approve_manager(self, request, pk=None):
        """
        تسک ۳.۴: تصویب نهایی مدیر و فعال‌سازی پرونده خودرو
        """
        from .fleet_lifecycle_services import approve_vehicle_stage, VehicleWorkflowError
        target = self.get_object()
        note = (request.data.get('note') or request.data.get('approval_note') or '').strip()
        try:
            instance = approve_vehicle_stage(target.id, request.user, stage='manager', note=note)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        broadcast_vehicle_update(instance, action_type='updated', message=f'خودرو «{instance.plate_number}» با راننده «{instance.driver_name}» تصویب و فعال شد.', sender_id=request.user.id if request.user else None)
        return Response({
            'message': 'تصویب نهایی مدیر با موفقیت ثبت شد و خودرو فعال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=False, methods=['post'], url_path='batch-approve')
    def batch_approve(self, request):
        """
        تسک ۳.۶ و DEF-09: تایید گروهی پرونده‌های خودرو با تراکنش مستقل برای هر رکورد
        """
        from .fleet_lifecycle_services import approve_vehicle_stage, FLEET_APPROVAL_STAGES, user_has_any_perm, VehicleWorkflowError

        stage = request.data.get('stage')
        ids = request.data.get('ids', [])
        note = (request.data.get('note') or 'تایید گروهی').strip()

        if not stage or stage not in FLEET_APPROVAL_STAGES:
            return Response({'error': 'مرحله ارسالی نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        rule = FLEET_APPROVAL_STAGES[stage]
        if not user_has_any_perm(request.user, rule['perms']):
            raise PermissionDenied(rule['perm_error'])

        if not ids or not isinstance(ids, list):
            return Response({'error': 'لیست شناسه‌های خودرو الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        allowed_qs = self.get_queryset().filter(id__in=ids)
        allowed_ids = set(allowed_qs.values_list('id', flat=True))

        success_ids = []
        errors = []

        for v_id in ids:
            if v_id not in allowed_ids:
                errors.append({'id': v_id, 'error': 'خودرو یافت نشد یا در قلمرو دسترسی شما قرار ندارد.'})
                continue
            try:
                approved_v = approve_vehicle_stage(v_id, request.user, stage=stage, note=note, batch=True)
                success_ids.append(v_id)
                broadcast_vehicle_update(
                    approved_v,
                    action_type='updated',
                    message=f'خودرو «{approved_v.plate_number}» به صورت گروهی در مرحله {stage} تایید شد.',
                    sender_id=request.user.id if request.user else None
                )
            except VehicleWorkflowError as e:
                errors.append({'id': v_id, 'error': e.message, 'missing_fields': e.missing_fields})
            except Exception as e:
                errors.append({'id': v_id, 'error': str(e)})

        return Response({
            'success_count': len(success_ids),
            'failed_count': len(errors),
            'success_ids': success_ids,
            'errors': errors,
        })

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        """
        تسک ۳.۱ و ب-۱۸: رد پرونده با فهرست سفید وضعیت‌های در جریان، قفل ردیف، و Purge مشروط در صورت عدم سابقه
        """
        from .fleet_lifecycle_services import (
            user_has_any_perm,
            FLEET_MANAGER_STAGE_PERMS,
            FLEET_SUPERVISOR_STAGE_PERMS,
            FLEET_FINANCE_STAGE_PERMS,
            check_vehicle_history,
            build_vehicle_snapshot,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        can_mgr = user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS + FLEET_SUPERVISOR_STAGE_PERMS)
        can_fin = user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS)
        if not (can_mgr or can_fin):
            raise PermissionDenied("شما دسترسی لازم برای رد این پرونده را ندارید.")

        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت دلیل رد الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        target = self.get_object()
        valid_statuses = [
            'pending_supervisor', 'pending_accountant', 'pending_manager',
            'supervisor_approved', 'accountant_approved', 'revision_required'
        ]
        if target.approval_status not in valid_statuses:
            return Response(
                {'error': f'امکان رد پرونده در وضعیت «{target.get_approval_status_display()}» وجود ندارد.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            if instance.approval_status not in valid_statuses:
                return Response(
                    {'error': f'وضعیت پرونده تغییر یافته است (وضعیت فعلی: {instance.get_approval_status_display()}).'},
                    status=status.HTTP_409_CONFLICT
                )

            history = check_vehicle_history(instance)
            is_never_approved = instance.manager_approved_at is None
            purge_statuses = ('revision_required', 'pending_supervisor', 'pending_accountant', 'pending_manager')
            can_purge = is_never_approved and history['can_delete'] and instance.approval_status in purge_statuses

            old_status = instance.approval_status
            cid = instance.company_id
            plate = instance.plate_number

            if can_purge:
                snapshot = build_vehicle_snapshot(instance)
                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=user,
                    from_status=old_status,
                    to_status='purged',
                    action='purge',
                    reason=reason,
                    metadata={'snapshot': snapshot, 'company_id': cid}
                )
                instance.delete()
                dispatch_vehicle_workflow_notification(target, stage='purged', actor=user, reason=reason)
                broadcast_vehicle_update(target, action_type='deleted', message=f'پرونده خودرو «{plate}» به دلیل رد قطعی و عدم سابقه پاک‌سازی شد.', sender_id=user.id)
                return Response({
                    'message': f'پرونده خودرو «{plate}» به دلیل رد قطعی و عدم وجود هرگونه سابقه، از سامانه پاک‌سازی شد.',
                    'purged': True
                })
            else:
                instance.approval_status = 'rejected'
                instance.rejection_reason = reason
                instance.save(update_fields=['approval_status', 'rejection_reason'])

                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=user,
                    from_status=old_status,
                    to_status='rejected',
                    action='reject',
                    reason=reason,
                    metadata={'company_id': cid}
                )
                dispatch_vehicle_workflow_notification(instance, stage='rejected', actor=user, reason=reason)
                broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{plate}» رد شد: {reason}', sender_id=user.id)
                return Response({
                    'message': f'پرونده خودرو «{plate}» رد شد.',
                    'data': self.get_serializer(instance).data,
                    'purged': False
                })

    @action(detail=True, methods=['post'], url_path='request-revision')
    def request_revision(self, request, pk=None):
        """
        تسک ۳.۲ و ب-۱۹ و DEF-13: عودت پرونده بر پایه وضعیت فعلی:
        - در pending_supervisor: عودت به کارمند (revision_required)
        - در pending_accountant: عودت به کارمند (revision_required) یا سرپرست (pending_supervisor)
        - در pending_manager: عودت به حسابدار (pending_accountant)
        """
        from .fleet_lifecycle_services import (
            user_has_any_perm,
            FLEET_MANAGER_STAGE_PERMS,
            FLEET_SUPERVISOR_STAGE_PERMS,
            FLEET_FINANCE_STAGE_PERMS,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        target = self.get_object()

        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت توضیحات و دلیل عودت الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            old_status = instance.approval_status

            if old_status == 'pending_supervisor':
                if not user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS):
                    raise PermissionDenied("شما دسترسی سرپرستی برای عودت این پرونده را ندارید.")
                target_status = 'revision_required'

            elif old_status in ('pending_accountant', 'supervisor_approved'):
                if not user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS + FLEET_MANAGER_STAGE_PERMS):
                    raise PermissionDenied("شما دسترسی مالی برای عودت این پرونده را ندارید.")
                target_choice = request.data.get('target_stage', 'employee')
                if target_choice in ('supervisor', 'pending_supervisor'):
                    target_status = 'pending_supervisor'
                else:
                    target_status = 'revision_required'

            elif old_status in ('pending_manager', 'accountant_approved'):
                if not user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS):
                    raise PermissionDenied("شما دسترسی مدیریت برای عودت این پرونده را ندارید.")
                target_status = 'pending_accountant'

            else:
                return Response(
                    {'error': f'امکان عودت پرونده در وضعیت «{instance.get_approval_status_display()}» وجود ندارد.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            instance.approval_status = target_status
            instance.rejection_reason = reason
            instance.save(update_fields=['approval_status', 'rejection_reason'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status=old_status,
                to_status=target_status,
                action='request_revision',
                reason=reason,
                metadata={'company_id': instance.company_id, 'target_stage': target_status}
            )
            dispatch_vehicle_workflow_notification(
                instance, stage=target_status, actor=user, reason=reason, is_return=True
            )

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» عودت داده شد: {reason}', sender_id=user.id)
        return Response({
            'message': f'پرونده خودرو با موفقیت به وضعیت «{instance.get_approval_status_display()}» عودت داده شد.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['get'], url_path='rate-benchmark')
    def rate_benchmark(self, request, pk=None):
        """
        تسک ۹.۳ و ۱۱.۲ (ن-۲): مقایسهٔ نرخ با میانگین هم‌تیپ‌های پروژه (سروری و محدود به قلمرو).
        گارد نقش: فقط برای کاربران مجاز به مشاهده نرخ‌های ناوگان (مالی و مدیریت).
        """
        user = request.user
        from .serializers import user_can_view_fleet_rates
        if not user_can_view_fleet_rates(user):
            raise PermissionDenied("مشاهده مقایسه و شاخص نرخ ناوگان فقط برای نقش‌های مدیریت و مالی مجاز است.")

        from django.db.models import Avg, Min, Max, Count
        target = self.get_object()

        # کوئری هم‌تیپ‌ها در همان پروژه یا بخش با رعایت قلمرو شرکتی
        qs = self.get_queryset().filter(
            vehicle_type=target.vehicle_type,
            default_service_rate__gt=0
        )
        if target.section_id:
            if target.section and target.section.project_id:
                qs = qs.filter(section__project_id=target.section.project_id)
            else:
                qs = qs.filter(section_id=target.section_id)

        stats = qs.aggregate(
            count=Count('id'),
            avg_rate=Avg('default_service_rate'),
            min_rate=Min('default_service_rate'),
            max_rate=Max('default_service_rate')
        )

        return Response({
            'vehicle_id': target.id,
            'vehicle_type': target.vehicle_type,
            'vehicle_type_display': target.get_vehicle_type_display() if hasattr(target, 'get_vehicle_type_display') else target.vehicle_type,
            'project_name': target.section.project.name if (target.section and target.section.project) else (target.section.name if target.section else 'عمومی'),
            'count': stats['count'] or 0,
            'avg_rate': round(stats['avg_rate']) if stats['avg_rate'] else 0,
            'min_rate': stats['min_rate'] or 0,
            'max_rate': stats['max_rate'] or 0,
            'current_rate': target.default_service_rate or 0
        })

    @action(detail=True, methods=['get'], url_path='cost-estimate')
    def cost_estimate(self, request, pk=None):
        """
        تسک ۱۰.۱: ویجت بار مالی ماهانه (سروری، محدود به قلمرو) در کارتابل مدیر.
        برآورد بار مالی ماهانه و سالانه خودرو بر مبنای نرخ روزانه، نوع مالکیت و انحراف از میانگین هم‌تیپ‌ها در پروژه.
        """
        user = request.user
        from .fleet_lifecycle_services import user_has_any_perm, FLEET_MANAGER_STAGE_PERMS, FLEET_FINANCE_STAGE_PERMS
        if not (user.is_superuser or user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS + FLEET_FINANCE_STAGE_PERMS)):
            raise PermissionDenied("مشاهده برآورد مالی و نرخ ناوگان فقط برای نقش‌های مدیریت و مالی مجاز است.")

        target = self.get_object()
        daily_rate = float(target.default_service_rate or 0)
        is_payable = target.ownership_type in ('contract', 'personal')

        # مبنای کاری ماهانه: ۲۶ روز کاری استاندارد (و ۳۰ روز کامل تقویمی)
        standard_days = 26
        full_month_days = 30

        monthly_cost_standard = round(daily_rate * standard_days) if is_payable else 0
        monthly_cost_full = round(daily_rate * full_month_days) if is_payable else 0
        yearly_cost_standard = monthly_cost_standard * 12

        # مقایسه با میانگین هم‌تیپ‌ها در قلمرو پروژه
        from django.db.models import Avg, Count
        qs = self.get_queryset().filter(
            vehicle_type=target.vehicle_type,
            default_service_rate__gt=0
        )
        if target.section_id:
            if target.section and target.section.project_id:
                qs = qs.filter(section__project_id=target.section.project_id)
            else:
                qs = qs.filter(section_id=target.section_id)

        stats = qs.aggregate(
            count=Count('id'),
            avg_rate=Avg('default_service_rate')
        )
        avg_rate = round(stats['avg_rate']) if stats['avg_rate'] else 0
        avg_monthly_cost = round(avg_rate * standard_days) if is_payable else 0
        deviation_percentage = round(((daily_rate - avg_rate) / avg_rate) * 100) if (avg_rate > 0 and is_payable) else 0

        return Response({
            'vehicle_id': target.id,
            'plate_number': target.plate_number,
            'driver_name': target.driver_name,
            'vehicle_type': target.vehicle_type,
            'vehicle_type_display': target.get_vehicle_type_display() if hasattr(target, 'get_vehicle_type_display') else target.vehicle_type,
            'ownership_type': target.ownership_type,
            'ownership_type_display': target.get_ownership_type_display() if hasattr(target, 'get_ownership_type_display') else target.ownership_type,
            'is_payable': is_payable,
            'daily_rate': daily_rate,
            'standard_days': standard_days,
            'full_month_days': full_month_days,
            'monthly_cost_standard': monthly_cost_standard,
            'monthly_cost_full': monthly_cost_full,
            'yearly_cost_standard': yearly_cost_standard,
            'avg_daily_rate': avg_rate,
            'avg_monthly_cost': avg_monthly_cost,
            'deviation_percentage': deviation_percentage,
            'benchmark_count': stats['count'] or 0
        })

    @action(detail=True, methods=['post'], url_path='reopen')
    def reopen(self, request, pk=None):
        """
        تسک ۳.۲-ب: بازگشایی پرونده ردشده (rejected -> draft) با قفل و بررسی وضعیت هم‌زمان
        """
        from .fleet_lifecycle_services import user_has_any_perm, FLEET_SUPERVISOR_STAGE_PERMS, FLEET_MANAGER_STAGE_PERMS, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user

        is_auth = (
            user.is_superuser
            or user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS + FLEET_MANAGER_STAGE_PERMS)
            or target.created_by_id == user.id
            or (target.section_id and UserSectionAssignment.objects.filter(user=user, section_id=target.section_id, is_active=True).exists())
        )
        if not is_auth:
            raise PermissionDenied("شما دسترسی بازگشایی این پرونده را ندارید.")

        reason = (request.data.get('reason') or 'بازگشایی مجدد پرونده ردشده جهت اصلاح').strip()
        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            if instance.approval_status != 'rejected':
                return Response(
                    {'error': f'فقط پرونده‌های در وضعیت ردشده امکان بازگشایی دارند (وضعیت فعلی: {instance.get_approval_status_display()}).'},
                    status=status.HTTP_409_CONFLICT
                )

            instance.approval_status = 'draft'
            instance.rejection_reason = None
            instance.save(update_fields=['approval_status', 'rejection_reason'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status='rejected',
                to_status='draft',
                action='reopen',
                reason=reason,
                metadata={'company_id': instance.company_id}
            )

        dispatch_vehicle_workflow_notification(instance, stage='reopened', actor=user, reason=reason)
        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» بازگشایی شد و به پیش‌نویس بازگشت.', sender_id=user.id)
        return Response({
            'message': 'پرونده خودرو با موفقیت بازگشایی شد و به وضعیت پیش‌نویس بازگشت.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='end-collaboration')
    def end_collaboration(self, request, pk=None):
        """
        تسک ۳.۱-ب: قطع همکاری خودرو؛ فقط برای خودروهای مصوب و فعال با مجوز مدیریت و end_date معتبر
        """
        from .fleet_lifecycle_services import (
            user_has_any_perm,
            FLEET_MANAGER_STAGE_PERMS,
            normalize_shamsi_date,
            parse_shamsi_date,
            _clean_shamsi_date_str,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        if not user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS):
            raise PermissionDenied("فقط مدیر مجاز به ثبت پایان همکاری خودرو است.")

        target = self.get_object()
        if target.approval_status != 'approved':
            return Response(
                {'error': f'قطع همکاری فقط برای خودروهای مصوب امکان‌پذیر است (وضعیت فعلی: {target.get_approval_status_display()}).'},
                status=status.HTTP_400_BAD_REQUEST
            )

        reason = (request.data.get('reason') or '').strip()
        if not reason:
            return Response({'error': 'ثبت علت قطع همکاری الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        raw_end_date = request.data.get('end_date')
        if not raw_end_date or not str(raw_end_date).strip():
            return Response({'error': 'ثبت تاریخ پایان همکاری (end_date) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            end_date = normalize_shamsi_date(raw_end_date)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        end_d_obj = parse_shamsi_date(end_date)

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)

            # تسک ۱۱.۶ (ن-۶): شرط فعال بودن و عدم وجود تاریخ پایان همکاری درون قفل
            if not instance.is_active or instance.end_date:
                return Response(
                    {'error': f'همکاری با این خودرو قبلاً در تاریخ «{instance.end_date}» پایان یافته است و امکان ثبت مجدد وجود ندارد.', 'code': 'already_ended'},
                    status=status.HTTP_409_CONFLICT
                )

            # تسک ۱۱.۶ (ن-۶): یافتن آخرین تردد بر پایه حداکثر شیء تاریخ شمسی (DEF-20)
            trip_dates = list(VehicleTripLog.objects.filter(vehicle=instance, is_deleted=False).values_list('date_shamsi', flat=True).distinct())
            max_trip_d = None
            max_trip_str = None
            for ds in trip_dates:
                if ds:
                    parsed = parse_shamsi_date(_clean_shamsi_date_str(ds))
                    if parsed and (max_trip_d is None or parsed > max_trip_d):
                        max_trip_d = parsed
                        max_trip_str = ds

            if max_trip_d and end_d_obj and end_d_obj < max_trip_d:
                return Response({
                    'error': f'تاریخ پایان همکاری ({end_date}) نمی‌تواند قبل از تاریخ آخرین تردد ثبت‌شده ({max_trip_str}) باشد.'
                }, status=status.HTTP_400_BAD_REQUEST)

            instance.is_active = False
            instance.end_date = end_date
            instance.save(update_fields=['is_active', 'end_date'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status=instance.approval_status,
                to_status=instance.approval_status,
                action='end_collaboration',
                reason=reason,
                metadata={'end_date': end_date, 'company_id': instance.company_id}
            )
            dispatch_vehicle_workflow_notification(instance, stage='end_collaboration', actor=user, reason=reason)

        broadcast_vehicle_update(instance, action_type='updated', message=f'همکاری با خودرو «{instance.plate_number}» در تاریخ {end_date} پایان یافت.', sender_id=user.id)
        return Response({
            'message': f'همکاری با خودرو «{instance.plate_number}» پایان یافت و پرونده غیرفعال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        """
        دانلود قالب اکسل خام ناوگان و رانندگان با ساختار ۲ سطری استاندارد شرکت
        """
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Vehicles'
        ws.sheet_view.rightToLeft = True

        columns_fa = [
            'ردیف', 'شماره پلاک', 'نام راننده', 'کد ملی راننده', 'شماره همراه راننده',
            'نام مالک', 'کد ملی مالک', 'شماره همراه مالک',
            'نوع خودرو', 'نوع مالکیت', 'نرخ پایه سرویس (ریال)', 'شماره شبا', 'نام بانک', 'شماره حساب'
        ]
        columns_en = [
            'row_num', 'plate_number', 'driver_name', 'driver_national_code', 'driver_phone',
            'owner_name', 'owner_national_code', 'owner_phone',
            'vehicle_type', 'ownership_type', 'default_service_rate', 'sheba_number', 'bank_name', 'account_number'
        ]

        f_title = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        f_key = Font(name='Segoe UI', size=9, bold=True, color='CBD5E1')
        fill_title = PatternFill(start_color='0284C7', end_color='0284C7', fill_type='solid') # Sky-600
        fill_key = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 20
        ws.freeze_panes = 'A3'

        for col_idx, (c_fa, c_en) in enumerate(zip(columns_fa, columns_en), 1):
            cell_fa = ws.cell(row=1, column=col_idx, value=c_fa)
            cell_fa.font = f_title
            cell_fa.fill = fill_title
            cell_fa.alignment = align_center
            cell_fa.border = thin_border

            cell_en = ws.cell(row=2, column=col_idx, value=c_en)
            cell_en.font = f_key
            cell_en.fill = fill_key
            cell_en.alignment = align_center
            cell_en.border = thin_border

        # ردیف نمونه راهنما
        sample_row = [
            1, '12 الف 345 ایران 63', 'رضا اکبری', '0010376488', '09123456789',
            'رضا اکبری', '0010376488', '09123456789',
            'وانت نیسان', 'استیجاری', 4500000, 'IR120170000000123456789012', 'بانک ملی ایران', '0101234567001'
        ]
        ws.row_dimensions[3].height = 22
        for col_idx, val in enumerate(sample_row, 1):
            c = ws.cell(row=3, column=col_idx, value=val)
            c.alignment = align_center
            c.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="fleet_vehicles_template.xlsx"'
        return response

    @action(detail=False, methods=['post', 'POST', 'get', 'GET'], url_path='import-excel')
    def import_excel(self, request):
        """
        درون‌ریزی مستقیم شیت ناوگان از فایل اکسل شرکت (Upsert بر مبنای شماره پلاک خودرو)
        پشتیبانی از resolve_import_company، ensure_record_write_scope، اعتبارسنجی کد ملی،
        ایجاد حساب در VehicleBankAccount و عبور هوشمند با auto_pass_vehicle(..., draft_on_incomplete=True).
        """
        import openpyxl
        from decimal import Decimal
        from django.db import transaction
        from common.validators import validate_iranian_national_code
        from .sheba_utils import clean_sheba, validate_sheba, get_bank_from_sheba
        from .fleet_lifecycle_services import normalize_plate, auto_pass_vehicle
        from .serializers import user_can_view_fleet_rates
        from personnel.models import VehicleBankAccount

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        
        section_id = request.data.get('section_id') or request.query_params.get('section_id')
        dry_run = str(request.data.get('dry_run') or request.query_params.get('dry_run', 'false')).lower() in ['true', '1']
        update_existing = str(request.data.get('update_existing') or request.query_params.get('update_existing', 'true')).lower() in ['true', '1']
        
        user = request.user
        is_mgr = bool(user and (user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or user.has_perm('accounts.can_act_as_manager')))
        can_view_rate = user_can_view_fleet_rates(user)

        # تعیین بخش پیش‌فرض
        target_section = None
        if section_id:
            try:
                target_section = ProjectSection.objects.filter(id=int(section_id)).select_related('project').first()
            except (ValueError, TypeError):
                pass
        if not target_section and not is_mgr and user and user.is_authenticated:
            first_assign = UserSectionAssignment.objects.filter(user=user, is_active=True).select_related('section', 'section__project').first()
            if first_assign:
                target_section = first_assign.section

        # تعیین شرکت مقصد با resolve_import_company
        import_company_id = resolve_import_company(request)
        if target_section and target_section.project_id:
            sec_company_id = target_section.project.company_id
            if import_company_id and sec_company_id and import_company_id != sec_company_id:
                return Response({'error': 'بخش انتخاب‌شده با شرکت هدف تطابق ندارد.'}, status=status.HTTP_400_BAD_REQUEST)
            import_company_id = sec_company_id or import_company_id

        if requires_company_scope(user) and not import_company_id:
            return Response({'error': 'تعیین شرکت جهت درون‌ریزی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            ensure_record_write_scope(user, company=import_company_id, section=target_section, label='ناوگان')
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_403_FORBIDDEN)

        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
        except Exception as e:
            return Response({
                'success': False,
                'dry_run': dry_run,
                'summary': {'total_rows': 0, 'created': 0, 'updated': 0, 'skipped': 0, 'valid_count': 0, 'error_count': 1},
                'errors': [{'row': 0, 'field': 'فایل اکسل', 'message': f'خطا در باز کردن فایل اکسل: {str(e)}'}]
            }, status=status.HTTP_400_BAD_REQUEST)
        
        target_sheet = None
        for name in wb.sheetnames:
            if name.strip().lower() in ['vehicles', 'fleet', 'ناوگان', 'خودرو', 'رانندگان']:
                target_sheet = wb[name]
                break
        if not target_sheet:
            target_sheet = wb.active

        # مپینگ نام ستون‌ها
        vtype_map = {
            'وانت نیسان': 'nissan', 'نیسان': 'nissan', 'nissan': 'nissan',
            'خاور': 'khavar', 'کامیونت': 'khavar', 'khavar': 'khavar',
            'تریلی': 'trailer', 'کشنده': 'trailer', 'trailer': 'trailer',
            'وانت بار': 'pickup', 'وانت': 'pickup', 'pickup': 'pickup',
            'کامیون': 'truck', 'تک': 'truck', 'جفت': 'truck', 'truck': 'truck',
            'لیفتراک': 'forklift', 'forklift': 'forklift',
            'سواری': 'sedan', 'sedan': 'sedan',
            'سایر': 'other', 'other': 'other'
        }
        own_map = {
            'استیجاری': 'contract', 'پیمانکاری': 'contract', 'contract': 'contract',
            'شرکتی': 'company', 'ملکی پروژه': 'company', 'company': 'company',
            'ملکی': 'personal', 'ملکی راننده': 'personal', 'personal': 'personal'
        }

        scoped_vehicles_qs = scope_vehicle_queryset(request, VehicleDriverProfile.objects.all())
        created_count = 0
        updated_count = 0
        skipped_count = 0
        errors = []
        warnings = []
        parsed_vehicles = []
        seen_file_plates = {}

        # بررسی سرستون‌های ردیف ۲ جهت نگاشت منعطف ستون‌ها
        col_map = {}
        for c in range(1, min(target_sheet.max_column + 1, 20)):
            val = str(target_sheet.cell(row=2, column=c).value or '').strip().lower()
            if val:
                col_map[val] = c

        has_owner_cols = 'owner_name' in col_map or 'owner_national_code' in col_map

        for row_idx in range(3, target_sheet.max_row + 1):
            plate_raw = target_sheet.cell(row=row_idx, column=col_map.get('plate_number', 2)).value
            if not plate_raw or not str(plate_raw).strip():
                continue

            plate = normalize_plate(str(plate_raw).strip())
            if not plate:
                continue

            if plate in seen_file_plates:
                errors.append({
                    'row': row_idx,
                    'field': 'شماره پلاک',
                    'message': f'شماره پلاک «{plate}» در ردیف {seen_file_plates[plate]} نیز تکرار شده است (پلاک تکراری در فایل اکسل).'
                })
                continue
            seen_file_plates[plate] = row_idx

            driver_name = str(target_sheet.cell(row=row_idx, column=col_map.get('driver_name', 3)).value or '').strip()
            if not driver_name:
                errors.append({'row': row_idx, 'field': 'نام راننده', 'message': f'نام راننده برای پلاک {plate} الزامی است.'})
                continue

            nat_code_raw = target_sheet.cell(row=row_idx, column=col_map.get('driver_national_code', 4)).value
            nat_code = str(nat_code_raw).strip() if nat_code_raw else ''
            if nat_code:
                clean_nat = ''.join(c for c in normalize_digits(nat_code) if c.isdigit())
                clean_nat = clean_nat.zfill(10)
                try:
                    nat_code = validate_iranian_national_code(clean_nat)
                except ValueError as e:
                    errors.append({
                        'row': row_idx,
                        'field': 'کد ملی راننده',
                        'message': f'کد ملی راننده نامعتبر است ({nat_code_raw}): {str(e)}'
                    })
                    continue
            else:
                nat_code = None

            phone_raw = str(target_sheet.cell(row=row_idx, column=col_map.get('driver_phone', 5)).value or '').strip()
            phone = normalize_digits(phone_raw) if phone_raw else None
            
            if has_owner_cols:
                owner_name = str(target_sheet.cell(row=row_idx, column=col_map.get('owner_name', 6)).value or '').strip() or None
                owner_nat_raw = target_sheet.cell(row=row_idx, column=col_map.get('owner_national_code', 7)).value
                owner_nat = str(owner_nat_raw).strip() if owner_nat_raw else ''
                if owner_nat:
                    clean_o_nat = ''.join(c for c in normalize_digits(owner_nat) if c.isdigit())
                    clean_o_nat = clean_o_nat.zfill(10)
                    try:
                        owner_nat = validate_iranian_national_code(clean_o_nat)
                    except ValueError as e:
                        errors.append({
                            'row': row_idx,
                            'field': 'کد ملی مالک',
                            'message': f'کد ملی مالک نامعتبر است ({owner_nat_raw}): {str(e)}'
                        })
                        continue
                else:
                    owner_nat = None
                owner_phone_raw = str(target_sheet.cell(row=row_idx, column=col_map.get('owner_phone', 8)).value or '').strip()
                owner_phone = normalize_digits(owner_phone_raw) if owner_phone_raw else None
                is_driver_owner = bool(not owner_nat or owner_nat == nat_code)

                vtype_col = col_map.get('vehicle_type', 9)
                own_col = col_map.get('ownership_type', 10)
                rate_col = col_map.get('default_service_rate', 11)
                sheba_col = col_map.get('sheba_number', 12)
                bank_col = col_map.get('bank_name', 13)
                acc_col = col_map.get('account_number', 14)
            else:
                owner_name = driver_name
                owner_nat = nat_code
                owner_phone = phone
                is_driver_owner = True

                vtype_col = col_map.get('vehicle_type', 6)
                own_col = col_map.get('ownership_type', 7)
                rate_col = col_map.get('default_service_rate', 8)
                sheba_col = col_map.get('sheba_number', 9)
                bank_col = col_map.get('bank_name', 10)
                acc_col = col_map.get('account_number', 11)

            vtype_raw = str(target_sheet.cell(row=row_idx, column=vtype_col).value or '').strip().lower()
            vehicle_type = vtype_map.get(vtype_raw)
            if not vehicle_type:
                errors.append({
                    'row': row_idx,
                    'field': 'نوع خودرو',
                    'message': f'نوع خودرو «{vtype_raw}» نامعتبر است. انواع معتبر: نیسان، خاور، تریلی، وانت بار، کامیون، لیفتراک، سواری، سایر.'
                })
                continue

            own_raw = str(target_sheet.cell(row=row_idx, column=own_col).value or '').strip().lower()
            ownership_type = own_map.get(own_raw)
            if not ownership_type:
                errors.append({
                    'row': row_idx,
                    'field': 'نوع مالکیت',
                    'message': f'نوع مالکیت «{own_raw}» نامعتبر است. انواع معتبر: استیجاری، شرکتی، ملکی.'
                })
                continue

            # سنجش پلاک درون شرکت جاری و قلمرو مجاز (یکتایی پلاک به‌ازای شرکت)
            existing_in_company = VehicleDriverProfile.objects.filter(company_id=import_company_id, plate_number=plate).first() if import_company_id else None
            if not existing_in_company and plate_raw:
                existing_in_company = VehicleDriverProfile.objects.filter(company_id=import_company_id, plate_number=str(plate_raw).strip()).first() if import_company_id else None

            # اگر خودروی موجود در وضعیت قفل (مصوب، در انتظار و ...) باشد، اکسل مجاز به تغییر آن نیست
            # طبق مصوبه ط-۴: این رکوردها نباید کل ایمپورت را با خطای ۴۰۰ متوقف کنند، بلکه با هشدار کنار گذاشته می‌شوند
            if existing_in_company and existing_in_company.approval_status not in ['draft', 'revision_required']:
                skipped_count += 1
                status_lbl = existing_in_company.get_approval_status_display() if hasattr(existing_in_company, 'get_approval_status_display') else existing_in_company.approval_status
                warnings.append(
                    f'ردیف {row_idx}: خودروی پلاک «{plate_raw}» در وضعیت «{status_lbl}» قرار دارد و تغییر آن از طریق اکسل مجاز نیست (نادیده گرفته شد).'
                )
                continue

            from .fleet_lifecycle_services import is_valid_national_plate, ROAD_VEHICLE_TYPES
            if vehicle_type in ROAD_VEHICLE_TYPES and not is_valid_national_plate(plate):
                errors.append({
                    'row': row_idx,
                    'field': 'شماره پلاک',
                    'message': f'شماره پلاک «{plate_raw}» منطبق بر الگوی استاندارد پلاک ملی ایران نیست (مثال: ۱۲ب۳۴۵ایران۶۷ یا ۱۲-ب-۳۴۵-۶۷).'
                })
                continue

            if existing_in_company and not scoped_vehicles_qs.filter(id=existing_in_company.id).exists():
                errors.append({
                    'row': row_idx,
                    'field': 'شماره پلاک',
                    'message': f'شماره پلاک «{plate}» در این شرکت ثبت شده و شما دسترسی ویرایش آن را ندارید.'
                })
                continue

            if can_view_rate:
                rate_raw = target_sheet.cell(row=row_idx, column=rate_col).value
                try:
                    rate = Decimal(str(rate_raw or 0).replace(',', '').strip())
                    if rate < 0:
                        rate = Decimal(0)
                except Exception:
                    rate = Decimal(0)
            else:
                rate = Decimal(0)

            sheba_raw = str(target_sheet.cell(row=row_idx, column=sheba_col).value or '').strip()
            sheba = clean_sheba(sheba_raw) if sheba_raw else None
            
            bank_name = str(target_sheet.cell(row=row_idx, column=bank_col).value or '').strip() or None
            if sheba and not bank_name:
                b_info = get_bank_from_sheba(sheba)
                if b_info:
                    bank_name = b_info.get('name')

            account_no = str(target_sheet.cell(row=row_idx, column=acc_col).value or '').strip() or None

            parsed_vehicles.append({
                'row_idx': row_idx,
                'plate_number': plate,
                'driver_name': driver_name,
                'driver_national_code': nat_code,
                'driver_phone': phone,
                'is_driver_owner': is_driver_owner,
                'owner_name': owner_name,
                'owner_national_code': owner_nat,
                'owner_phone': owner_phone,
                'vehicle_type': vehicle_type,
                'ownership_type': ownership_type,
                'default_service_rate': rate,
                'sheba_number': sheba,
                'bank_name': bank_name,
                'account_number': account_no,
            })

        if errors:
            return Response({
                'success': False,
                'dry_run': dry_run,
                'summary': {
                    'total_rows': len(parsed_vehicles) + len(errors),
                    'created': 0, 'updated': 0, 'skipped': len(errors),
                    'valid_count': len(parsed_vehicles), 'error_count': len(errors)
                },
                'errors': errors
            }, status=status.HTTP_400_BAD_REQUEST)

        if not dry_run:
            from .fleet_lifecycle_services import build_vehicle_snapshot
            from .lifecycle_services import compute_profile_diff
            with transaction.atomic():
                for item in parsed_vehicles:
                    plate = item['plate_number']
                    q_exist = scoped_vehicles_qs.filter(plate_number=plate)
                    if import_company_id:
                        q_exist = q_exist.filter(company_id=import_company_id)
                    existing = q_exist.first()

                    if existing:
                        if update_existing:
                            # تسک ۱۱.۸ و ط-۳: به‌روزرسانی فقط در draft و revision_required
                            if existing.approval_status not in ['draft', 'revision_required']:
                                skipped_count += 1
                                warnings.append(
                                    f'ردیف {item["row_idx"]}: خودروی پلاک «{plate}» در وضعیت «{existing.get_approval_status_display() or existing.approval_status}» قرار دارد و تغییر آن از طریق اکسل مجاز نیست (نادیده گرفته شد).'
                                )
                                continue

                            old_snap = build_vehicle_snapshot(existing)

                            existing.driver_name = item['driver_name']
                            if item['driver_national_code']:
                                existing.driver_national_code = item['driver_national_code']
                            if item['driver_phone']:
                                existing.driver_phone = item['driver_phone']
                            existing.is_driver_owner = item['is_driver_owner']
                            if item['owner_name']:
                                existing.owner_name = item['owner_name']
                            if item['owner_national_code']:
                                existing.owner_national_code = item['owner_national_code']
                            if item['owner_phone']:
                                existing.owner_phone = item['owner_phone']
                            existing.vehicle_type = item['vehicle_type']
                            existing.ownership_type = item['ownership_type']
                            if can_view_rate and item['default_service_rate'] is not None:
                                existing.default_service_rate = item['default_service_rate']
                            if item['sheba_number']:
                                existing.sheba_number = item['sheba_number']
                            if item['bank_name']:
                                existing.bank_name = item['bank_name']
                            if item['account_number']:
                                existing.account_number = item['account_number']
                            if target_section and not existing.section_id:
                                existing.section = target_section
                                if target_section.project:
                                    existing.project = target_section.project
                            
                            # تسک ۱۱.۸ و ط-۵: اعمال _skip_bank_sync=True جهت جلوگیری از تکرار لاجیک مدل
                            existing._skip_bank_sync = True
                            existing.save()

                            # به‌روزرسانی یا ایجاد حساب اصلی در VehicleBankAccount
                            if item['sheba_number'] or item['account_number']:
                                pri_acc = VehicleBankAccount.objects.filter(vehicle=existing, is_primary=True).first()
                                if pri_acc:
                                    if item['bank_name']:
                                        pri_acc.bank_name = item['bank_name']
                                    if item['account_number']:
                                        pri_acc.account_number = item['account_number']
                                    if item['sheba_number']:
                                        pri_acc.sheba_number = item['sheba_number']
                                    pri_acc.save()
                                else:
                                    VehicleBankAccount.objects.create(
                                        vehicle=existing,
                                        bank_name=item['bank_name'] or '',
                                        account_number=item['account_number'] or '',
                                        sheba_number=item['sheba_number'] or '',
                                        account_owner_name=existing.driver_name if existing.is_driver_owner else (existing.owner_name or existing.driver_name),
                                        account_owner_national_code=existing.driver_national_code if existing.is_driver_owner else (existing.owner_national_code or existing.driver_national_code),
                                        payee_type='driver' if existing.is_driver_owner else 'owner',
                                        is_primary=True,
                                        is_active=True
                                    )

                            new_snap = build_vehicle_snapshot(existing)
                            diff = {}
                            for k, new_val in new_snap.items():
                                old_val = old_snap.get(k)
                                if old_val != new_val:
                                    diff[k] = {'old': old_val, 'new': new_val}

                            WorkflowAuditLog.objects.create(
                                content_type='vehicle',
                                object_id=existing.id,
                                actor=user if user and user.is_authenticated else None,
                                from_status=existing.approval_status,
                                to_status=existing.approval_status,
                                action='import_update',
                                reason=f'به‌روزرسانی مشخصات خودرو از طریق فایل اکسل (ردیف {item["row_idx"]})',
                                metadata={'company_id': existing.company_id, 'row': item['row_idx'], 'diff': diff}
                            )
                            updated_count += 1
                        else:
                            skipped_count += 1
                    else:
                        # تسک ۱۱.۸ و ط-۵: ایجاد اولیه بدون فیلدهای تخت بانک تا حساب اصلی مضاعف ساخته نشود
                        new_v = VehicleDriverProfile.objects.create(
                            company_id=import_company_id,
                            plate_number=plate,
                            driver_name=item['driver_name'],
                            driver_national_code=item['driver_national_code'],
                            driver_phone=item['driver_phone'],
                            is_driver_owner=item['is_driver_owner'],
                            owner_name=item['owner_name'],
                            owner_national_code=item['owner_national_code'],
                            owner_phone=item['owner_phone'],
                            vehicle_type=item['vehicle_type'],
                            ownership_type=item['ownership_type'],
                            default_service_rate=item['default_service_rate'],
                            sheba_number='',
                            bank_name='',
                            account_number='',
                            section=target_section,
                            project=target_section.project if (target_section and target_section.project) else None,
                            created_by=user if user and user.is_authenticated else None,
                            approval_status='draft',
                            is_active=True
                        )
                        # ایجاد صریح رکورد اصلی در VehicleBankAccount و سپس تنظیم فیلدهای تخت با _skip_bank_sync=True
                        if item['sheba_number'] or item['account_number']:
                            VehicleBankAccount.objects.create(
                                vehicle=new_v,
                                bank_name=item['bank_name'] or '',
                                account_number=item['account_number'] or '',
                                sheba_number=item['sheba_number'] or '',
                                account_owner_name=item['driver_name'] if item['is_driver_owner'] else (item['owner_name'] or item['driver_name']),
                                account_owner_national_code=item['driver_national_code'] if item['is_driver_owner'] else (item['owner_national_code'] or item['driver_national_code']),
                                payee_type='driver' if item['is_driver_owner'] else 'owner',
                                is_primary=True,
                                is_active=True
                            )
                            new_v.sheba_number = item['sheba_number'] or ''
                            new_v.bank_name = item['bank_name'] or ''
                            new_v.account_number = item['account_number'] or ''
                            new_v._skip_bank_sync = True
                            new_v.save(update_fields=['sheba_number', 'bank_name', 'account_number'])

                        # اجرای auto_pass_vehicle با draft_on_incomplete=True
                        auto_pass_vehicle(new_v, user, source='import_excel', draft_on_incomplete=True)
                        created_count += 1
        else:
            for item in parsed_vehicles:
                q_chk = scoped_vehicles_qs.filter(plate_number=item['plate_number'])
                if import_company_id:
                    q_chk = q_chk.filter(company_id=import_company_id)
                existing = q_chk.first()
                if existing:
                    if update_existing:
                        if existing.approval_status not in ['draft', 'revision_required']:
                            skipped_count += 1
                            warnings.append(
                                f'ردیف {item["row_idx"]}: خودروی پلاک «{item["plate_number"]}» در وضعیت «{existing.get_approval_status_display() or existing.approval_status}» قرار دارد و از طریق اکسل قابل ویرایش نیست (نادیده گرفته خواهد شد).'
                            )
                        else:
                            updated_count += 1
                    else:
                        skipped_count += 1
                else:
                    created_count += 1

        return Response({
            'success': True,
            'dry_run': dry_run,
            'summary': {
                'total_rows': len(parsed_vehicles),
                'created': created_count,
                'updated': updated_count,
                'skipped': skipped_count,
                'valid_count': len(parsed_vehicles),
                'error_count': 0
            },
            'created_count': created_count,
            'updated_count': updated_count,
            'warnings': warnings,
            'errors': []
        })

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی رسمی ۲ ردیفه اطلاعات ناوگان و رانندگان (عنوان فارسی + کلید لاتین + فریز A3)
        محدود به شرکت‌های مجاز کاربر (readable_company_ids) و تفکیک وظایف مالی (SoD) برای ستون نرخ.
        """
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse
        from django.utils import timezone
        from .serializers import user_can_view_fleet_rates

        qs = self.filter_queryset(self.get_queryset())
        cids = readable_company_ids(request.user)
        if cids is not None:
            qs = qs.filter(company_id__in=cids)

        can_view_rate = user_can_view_fleet_rates(request.user)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ناوگان و رانندگان"
        ws.views.sheetView[0].rightToLeft = True

        # تعریف ستون‌ها به همراه کلید، جهت متن، فونت عددی، و فرمت عددی
        cols_spec = [
            ('ردیف', 'row_num', 'center', True, None),
            ('شماره پلاک', 'plate_number', 'center', True, None),
            ('نوع خودرو', 'vehicle_type_display', 'center', False, None),
            ('نوع مالکیت', 'ownership_type_display', 'center', False, None),
            ('نام راننده', 'driver_name', 'right', False, None),
            ('کد ملی راننده', 'driver_national_code', 'center', True, None),
            ('شماره همراه راننده', 'driver_phone', 'center', True, None),
            ('مالک شخص راننده است', 'is_driver_owner', 'center', False, None),
            ('نام مالک', 'owner_name', 'right', False, None),
            ('کد ملی مالک', 'owner_national_code', 'center', True, None),
            ('شماره همراه مالک', 'owner_phone', 'center', True, None),
        ]
        if can_view_rate:
            cols_spec.append(('نرخ پایه سرویس (ریال)', 'default_service_rate', 'center', True, '#,##0'))
        cols_spec.extend([
            ('شماره حساب', 'account_number', 'center', True, None),
            ('شماره شبا', 'sheba_number', 'center', True, None),
            ('نام بانک', 'bank_name', 'center', False, None),
            ('وضعیت تایید پرونده', 'approval_status_display', 'center', False, None),
            ('وضعیت فعال', 'is_active', 'center', False, None),
        ])

        headers_fa = [col[0] for col in cols_spec]
        headers_en = [col[1] for col in cols_spec]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(headers_fa) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.freeze_panes = 'A3'

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, v in enumerate(qs, start=1):
            curr_row = idx + 2
            is_owner = v.is_driver_owner if v.is_driver_owner is not None else True
            val_map = {
                'row_num': idx,
                'plate_number': v.plate_number,
                'vehicle_type_display': v.get_vehicle_type_display() if hasattr(v, 'get_vehicle_type_display') else v.vehicle_type,
                'ownership_type_display': v.get_ownership_type_display() if hasattr(v, 'get_ownership_type_display') else (v.ownership_type or '-'),
                'driver_name': v.driver_name,
                'driver_national_code': v.driver_national_code or '-',
                'driver_phone': v.driver_phone or '-',
                'is_driver_owner': 'بله' if is_owner else 'خیر',
                'owner_name': v.owner_name or (v.driver_name if is_owner else '-'),
                'owner_national_code': v.owner_national_code or (v.driver_national_code if is_owner else '-'),
                'owner_phone': v.owner_phone or (v.driver_phone if is_owner else '-'),
                'default_service_rate': float(v.default_service_rate or 0),
                'account_number': v.account_number or '-',
                'sheba_number': v.sheba_number or '-',
                'bank_name': v.bank_name or '-',
                'approval_status_display': v.get_approval_status_display() if hasattr(v, 'get_approval_status_display') else str(v.approval_status),
                'is_active': 'فعال' if v.is_active else 'غیرفعال',
            }
            row_data = [val_map[col[1]] for col in cols_spec]
            ws.append(row_data)

            for col_idx, col_spec in enumerate(cols_spec, start=1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal=col_spec[2], vertical='center')
                c.font = data_font_num if col_spec[3] else data_font
                if col_spec[4]:
                    c.number_format = col_spec[4]

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"fleet_vehicles_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class PersonnelChangeRequestViewSet(viewsets.ModelViewSet):
    queryset = PersonnelChangeRequest.objects.all().select_related('personnel', 'requested_by', 'manager_reviewed_by', 'accountant_reviewed_by')
    serializer_class = PersonnelChangeRequestSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        # درخواست تغییرات فقط برای پرونده‌هایی دیده و بررسی می‌شود که کاربر خودِ پرونده را می‌بیند
        # (بخش‌های منتسب و شرکت‌های مجاز، با یا بدون سربرگ شرکت)
        visible_personnel = scope_personnel_queryset(self.request, PersonnelProfile.objects.all())
        qs = super().get_queryset().filter(personnel__in=visible_personnel.values('id'))
        status_filter = self.request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        personnel_id = self.request.query_params.get('personnel_id')
        if personnel_id:
            qs = qs.filter(personnel_id=personnel_id)
        return qs

    def create(self, request, *args, **kwargs):
        return Response(
            {'error': 'ثبت درخواست تغییرات مستقیماً مجاز نیست و باید از طریق ویرایش پرونده پرسنل فعال صورت گیرد.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def update(self, request, *args, **kwargs):
        return Response(
            {'error': 'ویرایش مستقیم درخواست تغییرات مجاز نیست. در صورت لزوم درخواست را لغو کرده و مجدداً ثبت نمایید.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def partial_update(self, request, *args, **kwargs):
        return self.update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if not request.user.is_superuser:
            return Response(
                {'error': 'حذف فیزیکی درخواست تغییرات مجاز نیست. لطفاً از عملیات لغو (Cancel) استفاده کنید.'},
                status=status.HTTP_403_FORBIDDEN
            )
        return super().destroy(request, *args, **kwargs)

    @action(detail=True, methods=['post'], url_path='approve-supervisor')
    def approve_supervisor(self, request, pk=None):
        from .lifecycle_services import SUPERVISOR_STAGE_PERMS, log_change_request_event, user_has_any_perm
        user = request.user
        if not user_has_any_perm(user, SUPERVISOR_STAGE_PERMS):
            raise PermissionDenied("شما دسترسی بررسی مرحله اول (سرپرست) درخواست تغییرات را ندارید.")

        cr = self.get_object()
        if cr.status not in ['draft', 'pending_supervisor', 'revision_required']:
            return Response({'error': f'درخواست در وضعیت «{cr.get_status_display()}» قابل تایید سرپرست نیست.'}, status=status.HTTP_400_BAD_REQUEST)

        note = (request.data.get('note') or '').strip()
        with transaction.atomic():
            cr.status = 'pending_accountant'
            cr.supervisor_reviewed_by = user
            cr.supervisor_reviewed_at = timezone.now()
            cr.save()
            log_change_request_event(cr, user, 'change_request_approve_supervisor', reason=note)
        return Response({'message': 'درخواست تغییرات به تایید سرپرست رسید و به کارتابل حسابداری ارسال شد.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='approve-finance')
    def approve_finance(self, request, pk=None):
        from .lifecycle_services import FINANCE_STAGE_PERMS, log_change_request_event, user_has_any_perm
        user = request.user
        if not user_has_any_perm(user, FINANCE_STAGE_PERMS):
            raise PermissionDenied("شما دسترسی تایید مرحله دوم (مالی) درخواست تغییرات را ندارید.")

        cr = self.get_object()
        valid_statuses = ['pending_accountant', 'supervisor_approved', 'manager_approved']
        if cr.status not in valid_statuses:
            return Response({'error': f'درخواست ابتدا باید به تایید سرپرست برسد (وضعیت فعلی: {cr.get_status_display()}).'}, status=status.HTTP_400_BAD_REQUEST)

        note = (request.data.get('note') or '').strip()
        with transaction.atomic():
            cr.status = 'pending_manager'
            cr.accountant_reviewed_by = user
            cr.accountant_reviewed_at = timezone.now()
            cr.save()
            log_change_request_event(cr, user, 'change_request_approve_finance', reason=note)
        return Response({'message': 'درخواست تغییرات به تایید مالی رسید و به کارتابل مدیر ارسال شد.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='approve-manager')
    def approve_manager(self, request, pk=None):
        from .lifecycle_services import (
            MANAGER_STAGE_PERMS,
            PersonnelWorkflowError,
            apply_personnel_change_request,
            log_change_request_event,
            refresh_has_pending_changes,
            user_has_any_perm,
        )
        user = request.user
        if not user_has_any_perm(user, MANAGER_STAGE_PERMS):
            raise PermissionDenied("شما دسترسی تصویب نهایی درخواست تغییرات را ندارید.")

        cr = self.get_object()
        valid_statuses = ['pending_manager', 'accountant_approved', 'pending_supervisor', 'pending_accountant', 'draft', 'manager_approved']
        if cr.status not in valid_statuses:
            return Response({'error': f'درخواست در وضعیت «{cr.get_status_display()}» قابل تصویب مدیر نیست.'}, status=status.HTTP_400_BAD_REQUEST)

        # DEF-06: اعمال تغییرات فقط از منبع واحد و با سریالایزر؛ اگر نامعتبر باشد ۴۰۰ و درخواست معلق می‌ماند
        try:
            with transaction.atomic():
                personnel = apply_personnel_change_request(cr, request=request)
                previous_status = cr.status
                now = timezone.now()
                cr.status = 'approved'
                cr.manager_reviewed_by = user
                cr.manager_reviewed_at = now
                if not cr.supervisor_reviewed_by:
                    cr.supervisor_reviewed_by = user
                    cr.supervisor_reviewed_at = now
                if not cr.accountant_reviewed_by:
                    cr.accountant_reviewed_by = user
                    cr.accountant_reviewed_at = now
                cr.save()
                refresh_has_pending_changes(personnel)
                log_change_request_event(
                    cr, user, 'change_request_applied',
                    reason=(request.data.get('note') or '').strip() or 'تصویب و اعمال درخواست تغییرات توسط مدیر',
                    extra={'diff': cr.proposed_changes, 'old_diff': cr.previous_values, 'from_change_request_status': previous_status},
                )
        except PersonnelWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        broadcast_personnel_update(personnel, action_type='updated', message=f'درخواست تغییرات پرسنل «{personnel.full_name}» تصویب و اعمال شد.', sender_id=user.id)
        return Response({'message': 'درخواست تغییرات توسط مدیر تصویب و روی پرونده پرسنل اعمال گردید.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        from .lifecycle_services import (
            FINANCE_STAGE_PERMS,
            MANAGER_STAGE_PERMS,
            SUPERVISOR_STAGE_PERMS,
            log_change_request_event,
            refresh_has_pending_changes,
            user_has_any_perm,
        )
        user = request.user
        if not user_has_any_perm(user, SUPERVISOR_STAGE_PERMS + FINANCE_STAGE_PERMS + MANAGER_STAGE_PERMS):
            raise PermissionDenied("دسترسی رد این درخواست را ندارید.")

        reason = request.data.get('reason', '').strip()
        cr = self.get_object()
        if cr.status in ('approved', 'rejected', 'cancelled'):
            return Response({'error': f'درخواست در وضعیت «{cr.get_status_display()}» قابل رد نیست.'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            cr.status = 'rejected'
            cr.rejection_reason = reason
            cr.save()
            refresh_has_pending_changes(cr.personnel)
            log_change_request_event(cr, user, 'change_request_rejected', reason=reason)

        return Response({'message': 'درخواست تغییرات رد شد.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='cancel')
    def cancel(self, request, pk=None):
        from .lifecycle_services import log_change_request_event, refresh_has_pending_changes
        cr = self.get_object()
        user = request.user
        if not (user.is_superuser or cr.requested_by_id == user.id):
            raise PermissionDenied("تنها کاربر ثبت‌کننده درخواست یا مدیر سیستم مجاز به انصراف و لغو درخواست تغییرات است.")

        # DEF-16: ابطال فقط تا پیش از تایید سرپرست
        if cr.status not in ('draft', 'pending_supervisor', 'revision_required'):
            return Response(
                {'error': f'درخواست در وضعیت «{cr.get_status_display()}» قابل لغو نیست؛ ابطال فقط تا پیش از تایید سرپرست ممکن است.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        reason = (request.data.get('reason') or 'لغو شده توسط ثبت‌کننده').strip()
        with transaction.atomic():
            cr.status = 'cancelled'
            cr.rejection_reason = reason
            cr.save(update_fields=['status', 'rejection_reason'])
            refresh_has_pending_changes(cr.personnel)
            # ثبت مستقیم در لاگ ممیزی گردش کار، روی تایملاین همان پرونده
            log_change_request_event(cr, user, 'change_request_cancelled', reason=reason)

        return Response({
            'message': 'درخواست تغییرات با موفقیت لغو شد.',
            'data': self.get_serializer(cr).data
        })



class VehicleChangeRequestViewSet(viewsets.ModelViewSet):
    queryset = VehicleChangeRequest.objects.all().select_related(
        'vehicle', 'requested_by', 'manager_reviewed_by', 'accountant_reviewed_by', 'supervisor_reviewed_by'
    )
    serializer_class = VehicleChangeRequestSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        qs = scope_vehicle_queryset(self.request, qs, prefix='vehicle__')

        status_filter = self.request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        vehicle_id = self.request.query_params.get('vehicle_id')
        if vehicle_id:
            qs = qs.filter(vehicle_id=vehicle_id)
        return qs

    def create(self, request, *args, **kwargs):
        """
        تسک ۳.۵ و ب-۲۳: ثبت مستقیم درخواست تغییرات از این مسیر مسدود است (۴۰۵).
        پیشنهاد تغییرات صرفاً از طریق اکشن اختصاصی پرونده خودرو انجام می‌شود.
        """
        return Response(
            {'error': 'ایجاد مستقیم درخواست تغییرات از این مسیر مجاز نیست. لطفاً از طریق اکشن اختصاصی پرونده خودرو اقدام فرمایید.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def update(self, request, *args, **kwargs):
        """تسک ۳.۵ و ب-۲۳: ویرایش مستقیم درخواست تغییرات مجاز نیست (۴۰۵)."""
        return Response(
            {'error': 'ویرایش مستقیم درخواست تغییرات ناوگان مجاز نیست.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def partial_update(self, request, *args, **kwargs):
        """تسک ۳.۵ و ب-۲۳: ویرایش مستقیم درخواست تغییرات مجاز نیست (۴۰۵)."""
        return Response(
            {'error': 'ویرایش مستقیم درخواست تغییرات ناوگان مجاز نیست.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def destroy(self, request, *args, **kwargs):
        """
        تسک ۳.۵ و ب-۲۳: حذف فیزیکی رکورد درخواست تغییرات منحصراً برای سوپریوزر (۴۰۳ برای بقیه).
        """
        if not request.user.is_superuser:
            raise PermissionDenied("حذف فیزیکی درخواست تغییرات ناوگان منحصراً در اختیار مدیر ارشد سیستم (سوپریوزر) است.")
        from .fleet_lifecycle_services import log_vehicle_change_request_event, refresh_vehicle_has_pending_changes
        cr = self.get_object()
        vehicle = cr.vehicle
        with transaction.atomic():
            log_vehicle_change_request_event(cr, request.user, 'change_request_deleted', reason='حذف توسط سوپریوزر')
            response = super().destroy(request, *args, **kwargs)
            refresh_vehicle_has_pending_changes(vehicle)
            return response

    @action(detail=True, methods=['post'], url_path='cancel')
    def cancel(self, request, pk=None):
        """
        تسک ۳.۵ و ب-۲۳: لغو درخواست تغییرات فقط برای ثبت‌کننده (یا سوپریوزر)
        و فقط تا پیش از تایید سرپرست (draft, pending_supervisor, revision_required).
        """
        from .fleet_lifecycle_services import log_vehicle_change_request_event, refresh_vehicle_has_pending_changes, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user
        if not (user.is_superuser or target.requested_by_id == user.id):
            raise PermissionDenied("تنها کاربر ثبت‌کننده درخواست یا مدیر سیستم مجاز به انصراف و لغو درخواست تغییرات است.")

        if target.status not in ('draft', 'pending_supervisor', 'revision_required'):
            return Response(
                {'error': f'درخواست در وضعیت «{target.get_status_display()}» قابل لغو نیست؛ ابطال فقط تا پیش از تایید سرپرست ممکن است.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        reason = (request.data.get('reason') or 'لغو شده توسط ثبت‌کننده').strip()
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status not in ('draft', 'pending_supervisor', 'revision_required'):
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'cancelled'
            cr.rejection_reason = reason
            cr.save(update_fields=['status', 'rejection_reason'])
            refresh_vehicle_has_pending_changes(cr.vehicle, exclude_cr_id=cr.id)
            log_vehicle_change_request_event(cr, user, 'change_request_cancelled', reason=reason)
            log_vehicle_audit(user, 'CANCEL_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_cancelled', actor=user, reason=reason, recipient_ids=[cr.requested_by_id])
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» لغو شد.', sender_id=user.id if user else None)

        return Response({
            'message': 'درخواست تغییرات با موفقیت لغو شد.',
            'data': self.get_serializer(cr).data
        })

    @action(detail=True, methods=['post'], url_path='approve-supervisor')
    def approve_supervisor(self, request, pk=None):
        from .fleet_lifecycle_services import log_vehicle_change_request_event, FLEET_SUPERVISOR_STAGE_PERMS, user_has_any_perm, dispatch_vehicle_workflow_notification
        user = request.user
        if not (user.is_superuser or user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS)):
            raise PermissionDenied("شما دسترسی بررسی مرحله اول (سرپرست) درخواست تغییرات ناوگان را ندارید.")
        
        target = self.get_object()
        # تسک ۱۱.۵ و ن-۵: فقط وضعیت pending_supervisor مورد تایید سرپرست است و تغییر وضعیت ۴۰۹ می‌گیرد
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status != 'pending_supervisor':
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'pending_accountant'
            cr.supervisor_reviewed_by = user
            cr.supervisor_reviewed_at = timezone.now()
            cr.save(update_fields=['status', 'supervisor_reviewed_by', 'supervisor_reviewed_at'])
            log_vehicle_change_request_event(cr, user, 'change_request_supervisor_approved')
            log_vehicle_audit(user, 'APPROVE_SUPERVISOR_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_supervisor_approved', actor=user)
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» به تایید سرپرست رسید.', sender_id=user.id if user else None)
        return Response({'message': 'درخواست تغییرات خودرو به تایید سرپرست رسید و به کارتابل حسابداری ارسال شد.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='approve-finance')
    def approve_finance(self, request, pk=None):
        from .fleet_lifecycle_services import log_vehicle_change_request_event, FLEET_FINANCE_STAGE_PERMS, user_has_any_perm, dispatch_vehicle_workflow_notification
        user = request.user
        # ب-۸: فقط perm_approve_fleet_finance (بدون can_act_as_accountant)
        if not (user.is_superuser or user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS)):
            raise PermissionDenied("شما دسترسی تایید مرحله دوم (مالی) درخواست تغییرات ناوگان را ندارید.")
        
        target = self.get_object()
        valid_statuses = ['pending_accountant', 'supervisor_approved']
        if target.status not in valid_statuses:
            return Response({'error': f'درخواست ابتدا باید به تایید سرپرست برسد (وضعیت فعلی: {target.get_status_display()}).'}, status=status.HTTP_400_BAD_REQUEST)
        
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status not in valid_statuses:
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'pending_manager'
            cr.accountant_reviewed_by = user
            cr.accountant_reviewed_at = timezone.now()
            cr.save(update_fields=['status', 'accountant_reviewed_by', 'accountant_reviewed_at'])
            log_vehicle_change_request_event(cr, user, 'change_request_finance_approved')
            log_vehicle_audit(user, 'APPROVE_FINANCE_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_finance_approved', actor=user)
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» به تایید مالی رسید.', sender_id=user.id if user else None)
        return Response({'message': 'درخواست تغییرات خودرو به تایید مالی رسید و به کارتابل مدیر ارسال شد.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='approve-manager')
    def approve_manager(self, request, pk=None):
        from .fleet_lifecycle_services import (
            log_vehicle_change_request_event,
            apply_vehicle_change_request,
            refresh_vehicle_has_pending_changes,
            FLEET_MANAGER_STAGE_PERMS,
            user_has_any_perm,
            VehicleWorkflowError,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        if not (user.is_superuser or user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS)):
            raise PermissionDenied("شما دسترسی تصویب نهایی درخواست تغییرات ناوگان را ندارید.")
        
        target = self.get_object()
        valid_statuses = ['pending_manager', 'accountant_approved']
        if user.is_superuser:
            valid_statuses += ['pending_supervisor', 'pending_accountant', 'draft', 'manager_approved']
        
        if target.status not in valid_statuses:
            return Response({'error': f'درخواست در وضعیت «{target.get_status_display()}» قابل تصویب مدیر نیست.'}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            with transaction.atomic():
                cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
                if cr.status not in valid_statuses:
                    return Response(
                        {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                        status=status.HTTP_409_CONFLICT
                    )
                # اعمال امن تغییرات بر روی پرونده خودرو با سریالایزر بدون هیچگونه setattr خام (تسک ۳.۵ و ب-۲۳)
                vehicle = apply_vehicle_change_request(cr, request=request)

                cr.status = 'approved'
                cr.manager_reviewed_by = user
                cr.manager_reviewed_at = timezone.now()
                if not cr.supervisor_reviewed_by:
                    cr.supervisor_reviewed_by = user
                    cr.supervisor_reviewed_at = timezone.now()
                if not cr.accountant_reviewed_by:
                    cr.accountant_reviewed_by = user
                    cr.accountant_reviewed_at = timezone.now()
                cr.save()

                refresh_vehicle_has_pending_changes(vehicle, exclude_cr_id=cr.id)
                log_vehicle_change_request_event(cr, user, 'change_request_applied')
                log_vehicle_audit(user, 'APPROVE_MANAGER_CHANGE_REQ', vehicle, {'change_request_id': cr.id, 'changes': cr.proposed_changes})
                dispatch_vehicle_workflow_notification(vehicle, stage='change_request_manager_approved', actor=user, recipient_ids=[cr.requested_by_id])
                broadcast_vehicle_update(vehicle, action_type='updated', message=f'تغییرات خودرو «{vehicle.plate_number}» توسط مدیر تایید و اعمال گردید.', sender_id=user.id if user else None)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        return Response({'message': 'درخواست تغییرات خودرو توسط مدیر تصویب و اعمال گردید.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        from .fleet_lifecycle_services import (
            log_vehicle_change_request_event,
            refresh_vehicle_has_pending_changes,
            FLEET_SUPERVISOR_STAGE_PERMS,
            FLEET_FINANCE_STAGE_PERMS,
            FLEET_MANAGER_STAGE_PERMS,
            user_has_any_perm,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        can_op = (
            user.is_superuser
            or user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS)
            or user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS)
            or user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS)
        )
        if not can_op:
            raise PermissionDenied("دسترسی رد این درخواست را ندارید.")
        
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت دلیل رد الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        
        target = self.get_object()
        active_statuses = ['draft', 'pending_supervisor', 'pending_accountant', 'pending_manager', 'revision_required', 'supervisor_approved', 'accountant_approved']
        if target.status not in active_statuses:
            return Response({'error': f'درخواست در وضعیت «{target.get_status_display()}» قابل رد نیست.'}, status=status.HTTP_400_BAD_REQUEST)
        
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status not in active_statuses:
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'rejected'
            cr.rejection_reason = reason
            cr.save(update_fields=['status', 'rejection_reason'])
            refresh_vehicle_has_pending_changes(cr.vehicle, exclude_cr_id=cr.id)
            log_vehicle_change_request_event(cr, user, 'change_request_rejected', reason=reason)
            log_vehicle_audit(user, 'REJECT_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id, 'reason': reason})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_rejected', actor=user, reason=reason, recipient_ids=[cr.requested_by_id])
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» رد شد: {reason}', sender_id=user.id if user else None)

        return Response({'message': 'درخواست تغییرات خودرو رد شد.', 'data': self.get_serializer(cr).data})

def normalize_attendance_date(date_str: str) -> str:
    """
    نرمال‌سازی تاریخ شمسی و پشتیبانی کامل از ورودی‌های پیوسته ۸ رقمی (14050607)
    و تبدیل خودکار به استاندارد YYYY/MM/DD بدون بروز خطا
    """
    if not date_str:
        return ""
    cleaned = normalize_digits(str(date_str)).strip()
    digits_only = ''.join(c for c in cleaned if c.isdigit())
    if len(digits_only) == 8 and ('/' not in cleaned and '-' not in cleaned):
        return f"{digits_only[:4]}/{digits_only[4:6]}/{digits_only[6:8]}"
    shamsi_str = format_to_shamsi_str(cleaned)
    if shamsi_str and len(shamsi_str.split('/')) == 3:
        return shamsi_str
    parts = cleaned.replace('-', '/').split('/')
    if len(parts) == 3:
        try:
            return f"{int(parts[0]):04d}/{int(parts[1]):02d}/{int(parts[2]):02d}"
        except Exception:
            pass
    return date_str


def validate_attendance_date_window(date_shamsi: str, user=None, allow_override: bool = False, warehouse_id=None, project_id=None):
    """
    اعتبارسنجی بازه مجاز ثبت و ویرایش کارکرد بر مبنای تنظیمات سالانه حقوق
    -1 به معنای نامحدود، 0 به معنای فقط امروز، اعداد مثبت به معنای سقف روز مجاز
    پشتیبانی کامل از رشته‌های ۸ رقمی پیوسته (مانند 14050607)
    مدیر ارشد سیستم (Superuser) با فلگ allow_override می‌تواند محدودیت زمانی را دور بزند.
    در صورتی که دوره کارکرد در وضعیت REJECTED باشد، جهت امکان اصلاح مغایرت‌ها توسط سرپرست، سقف روزهای گذشته لغو می‌شود.
    """
    if allow_override and user and (getattr(user, 'is_superuser', False) or getattr(user, 'is_staff', False) or user.has_perm('personnel.can_override_attendance_lock')):
        return True

    import jdatetime
    from rest_framework.exceptions import ValidationError as DRFValidationError
    
    if not date_shamsi:
        raise DRFValidationError("تاریخ شمسی معتبر نیست.")
    
    norm_date = normalize_attendance_date(date_shamsi)
    try:
        parts = [int(p) for p in norm_date.split('/')]
        if len(parts) != 3:
            raise ValueError
        target_date = jdatetime.date(parts[0], parts[1], parts[2])
    except Exception:
        raise DRFValidationError(f"فرمت تاریخ شمسی نامعتبر است: {date_shamsi}")

    today = jdatetime.date.today()
    delta_days = (target_date - today).days

    year_month = f"{parts[0]:04d}/{parts[1]:02d}"
    settings = get_effective_payroll_settings(year_month, project_id)

    past_limit = settings.attendance_edit_past_days if settings else 3
    future_limit = settings.attendance_edit_future_days if settings else 0

    # بررسی روزهای گذشته
    if delta_days < 0:
        rej_period_qs = MonthlyWorkPeriod.objects.filter(year_month=year_month, status='REJECTED')
        if warehouse_id:
            rej_period_qs = rej_period_qs.filter(warehouse_id=warehouse_id)
        if rej_period_qs.exists():
            return True

        past_days = abs(delta_days)
        if past_limit is not None and past_limit != -1:
            if past_days > past_limit:
                if past_limit == 0:
                    raise DRFValidationError(f"طبق تنظیمات سیستم، فقط امکان ثبت کارکرد برای تاریخ امروز ({today.strftime('%Y/%m/%d')}) وجود دارد و ویرایش سوابق گذشته مجاز نیست.")
                raise DRFValidationError(f"ویرایش کارکرد برای {past_days} روز قبل مجاز نیست. حداکثر سقف مجاز طبق تنظیمات مدیر {past_limit} روز است.")

    # بررسی روزهای آینده
    elif delta_days > 0:
        future_days = delta_days
        if future_limit is not None and future_limit != -1:
            if future_days > future_limit:
                if future_limit == 0:
                    raise DRFValidationError(f"ثبت کارکرد برای تاریخ‌های آینده ({norm_date}) مجاز نیست.")
                raise DRFValidationError(f"ثبت کارکرد برای {future_days} روز آینده مجاز نیست. حداکثر سقف مجاز طبق تنظیمات مدیر {future_limit} روز است.")

    return True


def is_date_shamsi_friday(date_shamsi_str):
    """
    بررسی جمعه بودن یک تاریخ شمسی
    """
    try:
        import jdatetime
        cleaned = normalize_digits(str(date_shamsi_str)).strip().replace('-', '/')
        parts = [int(p) for p in cleaned.split('/') if p.strip()]
        if len(parts) >= 3:
            return jdatetime.date(parts[0], parts[1], parts[2]).weekday() == 6
    except Exception:
        pass
    return False


def broadcast_attendance_updated(warehouse_id=None, date_shamsi=None, year_month=None, message=None, sender_id=None, client_tab_id=None, extra_data=None, company_id=None):
    """
    ارسال بلادرنگ رویداد ثبت/ویرایش کارکرد به کانال وب‌سوکت شرکتی
    """
    try:
        from channels.layers import get_channel_layer
        from asgiref.sync import async_to_sync
        channel_layer = get_channel_layer()
        if channel_layer is not None:
            cid = company_id
            if not cid and extra_data and isinstance(extra_data, dict):
                cid = extra_data.get('company_id')
            if not cid and warehouse_id:
                try:
                    from platform_core.registry import company_of
                    cid = company_of('warehouse', warehouse_id)
                except Exception:
                    pass


            payload = {
                'type': 'send_notification',
                'type_str': 'attendance_updated',
                'message': message or 'کارکرد پرسنل به‌روزرسانی شد.',
                'date_shamsi': date_shamsi,
                'year_month': year_month,
                'company_id': cid,
                'sender_id': sender_id,
                'client_tab_id': client_tab_id,
            }
            if warehouse_id is not None:
                payload['warehouse_id'] = warehouse_id
            if extra_data and isinstance(extra_data, dict):
                payload.update(extra_data)

            target_group = f'company_{cid}_notifications' if cid else 'global_notifications'
            async_to_sync(channel_layer.group_send)(
                target_group,
                payload
            )
    except Exception as e:
        import logging
        logging.getLogger(__name__).warning(f"[WebSocket] Error broadcasting attendance_updated: {e}")


class DailyAttendanceViewSet(viewsets.ModelViewSet):
    queryset = DailyAttendance.objects.filter(is_deleted=False).select_related('personnel', 'period')
    serializer_class = DailyAttendanceSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(warehouse_id=warehouse_id)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(Q(section_id=section_id) | Q(personnel__section_id=section_id))
        project_id = self.request.query_params.get('project_id')
        if project_id:
            qs = qs.filter(Q(section__project_id=project_id) | Q(personnel__section__project_id=project_id))
        date_shamsi = self.request.query_params.get('date_shamsi')
        if date_shamsi:
            qs = qs.filter(date_shamsi=date_shamsi)
        personnel_id = self.request.query_params.get('personnel_id')
        if personnel_id:
            qs = qs.filter(personnel_id=personnel_id)
        approval_status = self.request.query_params.get('approval_status')
        if approval_status:
            qs = qs.filter(approval_status=approval_status)
        year_month = self.request.query_params.get('year_month')
        if year_month:
            norm_ym = year_month.replace('-', '/')
            qs = qs.filter(date_shamsi__startswith=norm_ym)
        cid = validate_user_company_access(self.request.user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(Q(section__project__company_id=cid) | Q(personnel__company_id=cid) | Q(personnel__section__project__company_id=cid))
        return qs

    @action(detail=False, methods=['get'], url_path='matrix')
    def get_matrix(self, request):
        """
        دریافت داده‌های ماتریسی حضور و غیاب (پشتیبانی از مستقل از انبار یا فیلتر انبار خاص)
        همراه با استخراج هوشمند آخرین وضعیت ثبت‌شده فرد در روز قبل به عنوان پیش‌فرض
        """
        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_date = request.query_params.get('date_shamsi')
        if not raw_date:
            return Response({'error': 'پارامتر date_shamsi الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        date_shamsi = normalize_attendance_date(raw_date)
        year_month = date_shamsi[:7]

        from common.tenant_scope import get_current_tenant_company_id
        company_id = request.query_params.get('company_id') or get_current_tenant_company_id(request=request, user=request.user)

        # بررسی وضعیت قفل دوره بر مبنای انبار یا شرکت
        period = None
        if warehouse_id:
            period = MonthlyWorkPeriod.objects.filter(warehouse_id=warehouse_id, year_month=year_month).first()
        elif company_id:
            period = MonthlyWorkPeriod.objects.filter(company_id=company_id, year_month=year_month).first()

        is_locked = period.status == 'LOCKED' if period else False
        period_status = period.status if period else 'OPEN'

        # لیست پرسنل تاییدشده یا پرسنلی که در این تاریخ کارکرد دارند (مستقل از انبار یا فیلتر انبار خاص)
        if warehouse_id:
            personnel_list = PersonnelProfile.objects.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True),
                Q(is_active=True, approval_status='approved') | Q(daily_attendances__date_shamsi=date_shamsi, daily_attendances__is_deleted=False)
            ).distinct().order_by('last_name', 'first_name')
            if company_id:
                personnel_list = personnel_list.filter(Q(company_id=company_id) | Q(section__project__company_id=company_id))

            attendances_qs = DailyAttendance.objects.filter(
                Q(warehouse_id=warehouse_id) | Q(personnel__in=personnel_list),
                date_shamsi=date_shamsi,
                is_deleted=False
            ).select_related('personnel').order_by('id')
            # اولویت با رکوردهایی است که اختصاصاً برای همین انبار ثبت شده‌اند
            existing_attendances = {}
            for att in attendances_qs:
                if att.personnel_id not in existing_attendances or att.warehouse_id == warehouse_id:
                    existing_attendances[att.personnel_id] = att
        else:
            personnel_list = PersonnelProfile.objects.filter(
                Q(is_active=True, approval_status='approved') | Q(daily_attendances__date_shamsi=date_shamsi, daily_attendances__is_deleted=False)
            ).distinct().order_by('last_name', 'first_name')
            if company_id:
                personnel_list = personnel_list.filter(Q(company_id=company_id) | Q(section__project__company_id=company_id))

            attendances_qs = DailyAttendance.objects.filter(
                date_shamsi=date_shamsi,
                is_deleted=False
            ).select_related('personnel').order_by('id')
            existing_attendances = {att.personnel_id: att for att in attendances_qs}

        raw_sec = request.query_params.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        if section_id:
            personnel_list = personnel_list.filter(
                Q(section_id=section_id) |
                Q(section_assignments__section_id=section_id, section_assignments__is_active=True)
            ).distinct()
            attendances_qs = attendances_qs.filter(
                Q(section_id=section_id) | (Q(section_id__isnull=True) & Q(personnel__section_id=section_id))
            )
            existing_attendances = {att.personnel_id: att for att in attendances_qs}

        # فیلتر کانتکست شرکت فعال در ماتریس کارکرد
        cid = validate_user_company_access(request.user, get_request_company_id(request))
        if cid:
            personnel_list = personnel_list.filter(Q(company_id=cid) | Q(section__project__company_id=cid))

        is_today_friday = is_date_shamsi_friday(date_shamsi)

        matrix_rows = []
        for p in personnel_list:
            att = existing_attendances.get(p.id)
            p_job_title = p.job_title
            if section_id:
                sa = p.section_assignments.filter(section_id=section_id, is_active=True).first()
                if sa and sa.job_title:
                    p_job_title = sa.job_title
            if att:
                matrix_rows.append({
                    'personnel_id': p.id,
                    'full_name': p.full_name,
                    'national_code': p.national_code,
                    'job_title': p_job_title,
                    'is_active': p.is_active,
                    'warehouse_id': att.warehouse_id,
                    'warehouse_name': att.warehouse.name if att.warehouse else (p.assigned_warehouse.name if p.assigned_warehouse else 'شناور'),
                    'attendance_id': att.id,
                    'status': att.status,
                    'effective_hours': float(att.effective_hours),
                    'overtime_hours': float(att.overtime_hours),
                    'is_friday_work': att.is_friday_work,
                    'is_mission': att.is_mission,
                    'advance_payment': float(att.advance_payment),
                    'notes': att.notes or '',
                    'is_existing': True
                })
            else:
                # روزهای ثبت‌نشده یا پاکسازی‌شده با وضعیت خالی و ساعت ۰ بازگردانده می‌شوند
                matrix_rows.append({
                    'personnel_id': p.id,
                    'full_name': p.full_name,
                    'national_code': p.national_code,
                    'job_title': p_job_title,
                    'is_active': p.is_active,
                    'warehouse_id': p.assigned_warehouse_id,
                    'warehouse_name': p.assigned_warehouse.name if p.assigned_warehouse else 'شناور',
                    'attendance_id': None,
                    'status': '',
                    'effective_hours': 0.0,
                    'overtime_hours': 0.0,
                    'is_friday_work': is_today_friday,
                    'is_mission': False,
                    'advance_payment': 0.0,
                    'notes': '',
                    'is_existing': False
                })

        return Response({
            'warehouse_id': warehouse_id,
            'date_shamsi': date_shamsi,
            'year_month': year_month,
            'is_locked': is_locked,
            'period_status': period_status,
            'rows': matrix_rows
        })

    @action(detail=False, methods=['post'], url_path='bulk-save')
    def bulk_save(self, request):
        """
        ذخیره گروهی و سریع ماتریس حضور و غیاب پرسنل در پایان روز
        (پشتیبانی کامل از کارکرد مستقل از انبار و لاگ ممیزی)
        """
        serializer = BulkAttendanceMatrixSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        
        raw_wh = serializer.validated_data.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        project_id = serializer.validated_data.get('project_id')
        section_id = serializer.validated_data.get('section_id')
        if section_id and not project_id:
            sec_ref = ProjectSection.objects.filter(id=section_id).first()
            if sec_ref:
                project_id = sec_ref.project_id

        raw_date = serializer.validated_data['date_shamsi']
        date_shamsi = normalize_attendance_date(raw_date)
        items = serializer.validated_data['items']

        allow_override = bool(request.data.get('is_override', False) or request.headers.get('X-Admin-Override') == 'true')
        is_admin_override = allow_override and (getattr(request.user, 'is_superuser', False) or request.user.has_perm('personnel.can_override_attendance_lock'))

        # اعتبارسنجی بازه زمانی مجاز بر مبنای تنظیمات مدیر
        try:
            validate_attendance_date_window(date_shamsi, user=request.user, allow_override=is_admin_override, warehouse_id=warehouse_id, project_id=project_id)
        except Exception as e:
            err_msg = getattr(e, 'detail', str(e))
            if isinstance(err_msg, list):
                err_msg = err_msg[0]
            return Response({'error': str(err_msg)}, status=status.HTTP_400_BAD_REQUEST)

        year_month = date_shamsi[:7]
        locked_period_statuses = [
            'LOCKED', 'SUBMITTED', 'FINALIZED', 'SUBMITTED_SUPERVISOR',
            'PENDING_ACCOUNTANT', 'SUBMITTED_ACCOUNTANT', 'PENDING_MANAGER',
            'READY_TO_PAY', 'PAID'
        ]
        period = None
        if warehouse_id:
            period, _ = MonthlyWorkPeriod.objects.get_or_create(
                warehouse_id=warehouse_id,
                year_month=year_month,
                defaults={'status': 'OPEN'}
            )
            if period.status in locked_period_statuses and not is_admin_override:
                return Response({'error': f'دوره کارکرد {year_month} در وضعیت «{period.get_status_display()}» بوده و امکان ثبت یا ویرایش وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        saved_count = 0
        updated_count = 0

        with transaction.atomic():
            for item in items:
                personnel_id = item['personnel_id']
                p_obj = PersonnelProfile.objects.filter(id=personnel_id).first()
                if not p_obj:
                    continue

                if p_obj.approval_status != 'approved':
                    return Response({
                        'error': f'امکان ثبت کارکرد برای پرسنل «{p_obj.full_name}» وجود ندارد زیرا پرونده او در وضعیت «{p_obj.get_approval_status_display()}» بوده و هنوز به تایید نهایی مدیر و حسابدار نرسیده است.'
                    }, status=status.HTTP_400_BAD_REQUEST)

                from .lifecycle_services import validate_personnel_validity
                is_valid_date, val_err = validate_personnel_validity(p_obj, date_shamsi)
                if not is_valid_date:
                    return Response({'error': val_err}, status=status.HTTP_400_BAD_REQUEST)

                item_wh = item.get('warehouse_id')
                check_wh = item_wh if item_wh is not None else (warehouse_id if warehouse_id else p_obj.assigned_warehouse_id)
                target_period = period
                if not target_period and check_wh:
                    target_period = MonthlyWorkPeriod.objects.filter(
                        warehouse_id=check_wh,
                        year_month=year_month
                    ).first()
                if target_period and target_period.status in locked_period_statuses and not is_admin_override:
                    return Response({'error': f'دوره کارکرد {year_month} برای انبار پرسنل ({p_obj.full_name}) قفل شده است.'}, status=status.HTTP_400_BAD_REQUEST)

                target_wh = item_wh if item_wh is not None else (warehouse_id if warehouse_id else p_obj.assigned_warehouse_id)

                status_val = item['status']
                eff_h = float(item.get('effective_hours', 0))
                ot_h = float(item.get('overtime_hours', 0))
                is_fri = bool(item.get('is_friday_work', False))
                is_mis = bool(item.get('is_mission', False))
                adv_pay = float(item.get('advance_payment', 0))
                notes_val = item.get('notes', '') or ''

                target_section_id = section_id or p_obj.section_id
                target_project_id = project_id or p_obj.project_id or (p_obj.section.project_id if p_obj.section else None)

                # اگر وضعیت خالی و ساعت صفر باشد (پاکسازی‌شده)، رکورد دیتابیس در این بخش حذف می‌گردد
                if not status_val and eff_h == 0:
                    att_filter = {
                        'personnel_id': personnel_id,
                        'date_shamsi': date_shamsi,
                        'is_deleted': False
                    }
                    if target_section_id:
                        att_obj = DailyAttendance.objects.select_for_update().filter(
                            Q(**att_filter) & (Q(section_id=target_section_id) | (Q(section_id__isnull=True) & Q(personnel__section_id=target_section_id)))
                        ).first()
                    else:
                        att_obj = DailyAttendance.objects.select_for_update().filter(**att_filter).first()
                    if att_obj:
                        att_obj.is_deleted = True
                        att_obj.modified_by = request.user
                        att_obj.save()
                        AttendanceAuditLog.objects.create(
                            attendance=att_obj,
                            personnel_name=p_obj.full_name,
                            date_shamsi=date_shamsi,
                            changed_by=request.user,
                            field_name='is_deleted',
                            old_value='False',
                            new_value='True',
                            reason='پاکسازی کارکرد روزانه'
                        )
                        updated_count += 1
                    continue

                # اعتبارسنجی و ریست منطقی ساعات در صورت غیبت یا مرخصی
                if status_val in ['ABSENT', 'LEAVE']:
                    eff_h = 0.0
                    ot_h = 0.0
                    is_fri = False
                    is_mis = False

                # پیدا کردن رکورد فعال پرسنل در این تاریخ و این بخش
                att_filter = {
                    'personnel_id': personnel_id,
                    'date_shamsi': date_shamsi,
                    'is_deleted': False
                }
                if target_section_id:
                    att_obj = DailyAttendance.objects.select_for_update().filter(
                        Q(**att_filter) & (Q(section_id=target_section_id) | (Q(section_id__isnull=True) & Q(personnel__section_id=target_section_id)))
                    ).first()
                else:
                    att_obj = DailyAttendance.objects.select_for_update().filter(**att_filter).first()

                # اعتبارسنجی عدم تداخل کارکرد در چند بخش مختلف در یک روز
                other_att_qs = DailyAttendance.objects.filter(
                    personnel_id=personnel_id,
                    date_shamsi=date_shamsi,
                    is_deleted=False
                )
                if att_obj:
                    other_att_qs = other_att_qs.exclude(id=att_obj.id)
                elif target_section_id:
                    other_att_qs = other_att_qs.exclude(Q(section_id=target_section_id) | (Q(section_id__isnull=True) & Q(personnel__section_id=target_section_id)))

                if status_val == 'PRESENT_10H':
                    conflicting_full = other_att_qs.filter(status='PRESENT_10H').first()
                    if conflicting_full:
                        sec_title = conflicting_full.section.name if conflicting_full.section else 'بخش دیگری'
                        return Response({
                            'error': f'برای پرسنل «{p_obj.full_name}» در تاریخ {date_shamsi} وضعیت «حاضر کامل» در «{sec_title}» ثبت شده است. در صورت کارکرد همزمان در دو بخش، از وضعیت نیمه‌وقت یا ساعت سفارشی استفاده نمایید.'
                        }, status=status.HTTP_400_BAD_REQUEST)

                other_hours_sum = sum(float(a.effective_hours or 0) for a in other_att_qs)
                if other_hours_sum + eff_h > 24:
                    return Response({
                        'error': f'مجموع کارکرد روزانه پرسنل «{p_obj.full_name}» در تاریخ {date_shamsi} در بخش‌های مختلف ({other_hours_sum + eff_h:.1f} ساعت) از سقف ۲۴ ساعت شبانه‌روز بیشتر است.'
                    }, status=status.HTTP_400_BAD_REQUEST)

                if att_obj:
                    # بررسی تغییرات جهت ثبت در Audit Log
                    changes = []
                    item_sanitized = {
                        'status': status_val,
                        'effective_hours': eff_h,
                        'overtime_hours': ot_h,
                        'is_friday_work': is_fri,
                        'is_mission': is_mis,
                        'advance_payment': adv_pay,
                        'notes': notes_val
                    }
                    fields_to_check = ['status', 'effective_hours', 'overtime_hours', 'is_friday_work', 'is_mission', 'advance_payment', 'notes']
                    for f in fields_to_check:
                        new_val = item_sanitized.get(f)
                        old_val = getattr(att_obj, f)
                        is_diff = False
                        if f in ['effective_hours', 'overtime_hours', 'advance_payment']:
                            if float(new_val or 0) != float(old_val or 0):
                                is_diff = True
                        elif f in ['is_friday_work', 'is_mission']:
                            if bool(new_val) != bool(old_val):
                                is_diff = True
                        else:
                            if str(new_val or '') != str(old_val or ''):
                                is_diff = True
                        if is_diff:
                            changes.append((f, str(old_val), str(new_val)))

                    # اعمال تغییرات
                    att_obj.status = status_val
                    att_obj.effective_hours = eff_h
                    att_obj.overtime_hours = ot_h
                    att_obj.is_friday_work = is_fri
                    att_obj.is_mission = is_mis
                    att_obj.advance_payment = adv_pay
                    att_obj.notes = notes_val
                    att_obj.warehouse_id = target_wh
                    if target_section_id:
                        att_obj.section_id = target_section_id
                    if target_project_id:
                        att_obj.project_id = target_project_id
                    # در صورت ویرایش رکوردی که قبلا تایید شده بود، تاییدیه باطل شده و به وضعیت در انتظار سرپرست بازمی‌گردد
                    if changes and att_obj.approval_status == 'APPROVED':
                        att_obj.approval_status = 'PENDING_SUPERVISOR'
                        att_obj.supervisor_approved_by = None
                        att_obj.supervisor_approved_at = None
                        changes.append(('approval_status', 'APPROVED', 'PENDING_SUPERVISOR'))

                    att_obj.modified_by = request.user
                    if target_period:
                        att_obj.period = target_period
                    att_obj.save()

                    # پاکسازی رکوردهای تکراری قدیمی احتمالی برای این روز در همین بخش
                    dup_qs = DailyAttendance.objects.filter(
                        personnel_id=personnel_id,
                        date_shamsi=date_shamsi,
                        is_deleted=False
                    ).exclude(id=att_obj.id)
                    if target_section_id:
                        dup_qs = dup_qs.filter(Q(section_id=target_section_id) | (Q(section_id__isnull=True) & Q(personnel__section_id=target_section_id)))
                    dup_qs.delete()

                    # ثبت لاگ‌های ممیزی
                    for f_name, o_val, n_val in changes:
                        AttendanceAuditLog.objects.create(
                            attendance=att_obj,
                            personnel_name=p_obj.full_name,
                            date_shamsi=date_shamsi,
                            changed_by=request.user,
                            field_name=f_name,
                            old_value=o_val,
                            new_value=n_val,
                            reason='ویرایش ماتریسی کارکرد روزانه'
                        )
                    updated_count += 1
                else:
                    # ایجاد رکورد جدید با انتساب به بخش و پروژه مشخص
                    att_obj = DailyAttendance.objects.create(
                        personnel_id=personnel_id,
                        warehouse_id=target_wh,
                        project_id=target_project_id,
                        section_id=target_section_id,
                        date_shamsi=date_shamsi,
                        status=status_val,
                        effective_hours=eff_h,
                        overtime_hours=ot_h,
                        is_friday_work=is_fri,
                        is_mission=is_mis,
                        advance_payment=adv_pay,
                        notes=notes_val,
                        period=target_period or period,
                        created_by=request.user,
                        modified_by=request.user
                    )
                    saved_count += 1

        client_tab_id = request.data.get('client_tab_id') or request.headers.get('X-Client-Tab-ID')
        broadcast_attendance_updated(
            warehouse_id=warehouse_id,
            date_shamsi=date_shamsi,
            year_month=year_month,
            sender_id=request.user.id if request.user and request.user.is_authenticated else None,
            client_tab_id=client_tab_id,
            message=f'کارکرد پرسنل برای تاریخ {date_shamsi} به‌روزرسانی شد.'
        )

        return Response({
            'message': f'اطلاعات کارکرد با موفقیت ذخیره شد ({saved_count} جدید، {updated_count} بروزرسانی).',
            'saved_count': saved_count,
            'updated_count': updated_count
        })

    @action(detail=False, methods=['post'], url_path='clear-day')
    def clear_day(self, request):
        """
        پاکسازی کامل و اصولی کارکرد روزانه پرسنل برای یک تاریخ مشخص
        """
        date_shamsi = request.data.get('date_shamsi')
        if not date_shamsi:
            return Response({'error': 'پارامتر date_shamsi الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        raw_wh = request.data.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        project_id = safe_int(request.data.get('project_id'))
        section_id = safe_int(request.data.get('section_id'))
        personnel_ids = request.data.get('personnel_ids')

        # اعتبارسنجی بازه مجاز ویرایش تاریخ
        is_admin_override = request.user.is_superuser or (request.user.role and request.user.role.name in ['Admin', 'Manager']) or request.user.has_perm('personnel.can_override_attendance_lock')
        if not is_admin_override:
            validate_attendance_date_window(date_shamsi, user=request.user, allow_override=is_admin_override, warehouse_id=warehouse_id, project_id=project_id)

        qs = DailyAttendance.objects.filter(date_shamsi=date_shamsi, is_deleted=False)
        if warehouse_id:
            qs = qs.filter(warehouse_id=warehouse_id)
        if section_id:
            qs = qs.filter(section_id=section_id)
        if personnel_ids and isinstance(personnel_ids, list) and len(personnel_ids) > 0:
            qs = qs.filter(personnel_id__in=personnel_ids)

        deleted_count = 0
        with transaction.atomic():
            for att in qs.select_for_update():
                att.is_deleted = True
                att.modified_by = request.user
                att.save()
                AttendanceAuditLog.objects.create(
                    attendance=att,
                    personnel_name=att.personnel.full_name if att.personnel else '',
                    date_shamsi=date_shamsi,
                    changed_by=request.user,
                    field_name='is_deleted',
                    old_value='False',
                    new_value='True',
                    reason='پاکسازی کارکرد روزانه'
                )
                deleted_count += 1

        ym = date_shamsi[:7] if len(date_shamsi) >= 7 else None
        client_tab_id = request.data.get('client_tab_id') or request.headers.get('X-Client-Tab-ID')
        broadcast_attendance_updated(
            warehouse_id=warehouse_id,
            date_shamsi=date_shamsi,
            year_month=ym,
            sender_id=request.user.id if request.user and request.user.is_authenticated else None,
            client_tab_id=client_tab_id,
            message=f'کارکرد روز {date_shamsi} برای {deleted_count} نفر پاکسازی شد.'
        )

        return Response({
            'message': f'کارکرد روز {date_shamsi} برای {deleted_count} نفر با موفقیت پاکسازی شد.',
            'cleared_count': deleted_count
        })

    @action(detail=False, methods=['get'], url_path='monthly-grid')
    def get_monthly_grid(self, request):
        """
        دریافت ماتریس کارکرد ۳۱ روزه کلیه پرسنل در یک ماه مشخص (نمای شیت ماهانه)
        پشتیبانی کامل از کارکرد مستقل از انبار و اصلاح وضعیت روزهای ثبت‌نشده (نباید به عنوان حاضر بزند)
        """
        import jdatetime
        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        project_id = safe_int(request.query_params.get('project_id'))
        section_id = safe_int(request.query_params.get('section_id'))
        raw_year_month = request.query_params.get('year_month')
        if not raw_year_month:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        # نرمال‌سازی سال و ماه (پشتیبانی از 140504 یا 1405/04)
        ym_cleaned = normalize_digits(str(raw_year_month)).strip().replace('-', '/')
        parts = ym_cleaned.split('/')
        if len(parts) >= 2:
            y, m = int(parts[0]), int(parts[1])
            year_month = f"{y:04d}/{m:02d}"
        else:
            digits_only = ''.join(c for c in ym_cleaned if c.isdigit())
            if len(digits_only) >= 6:
                y, m = int(digits_only[:4]), int(digits_only[4:6])
                year_month = f"{y:04d}/{m:02d}"
            else:
                return Response({'error': 'فرمت year_month نامعتبر است (مثال: 1405/04).'}, status=status.HTTP_400_BAD_REQUEST)

        # محاسبه تعداد روزهای ماه خورشیدی
        if 1 <= m <= 6:
            days_in_month = 31
        elif 7 <= m <= 11:
            days_in_month = 30
        elif m == 12:
            days_in_month = 30 if jdatetime.date(y, 1, 1).isleap() else 29
        else:
            return Response({'error': 'شماره ماه باید بین ۱ تا ۱۲ باشد.'}, status=status.HTTP_400_BAD_REQUEST)

        month_names = ["", "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
        month_name = month_names[m] if 1 <= m <= 12 else ""

        is_locked = False
        period_status = 'OPEN'
        period_obj = None
        if warehouse_id:
            period_obj = MonthlyWorkPeriod.objects.filter(warehouse_id=warehouse_id, year_month=year_month).first()
            if period_obj:
                period_status = period_obj.status
                is_locked = period_obj.status in ['LOCKED', 'SUBMITTED', 'FINALIZED']

        # واکشی تنظیمات بازه ویرایش متناسب با سال، ماه و پروژه
        settings = get_effective_payroll_settings(year_month, project_id)
        past_limit = settings.attendance_edit_past_days if settings else 3
        future_limit = settings.attendance_edit_future_days if settings else 0

        today = jdatetime.date.today()

        days_meta = []
        for d in range(1, days_in_month + 1):
            target_date = jdatetime.date(y, m, d)
            weekday = target_date.weekday()  # 0=شنبه ... 6=جمعه
            is_friday = (weekday == 6)
            weekday_names = ["ش", "ی", "د", "س", "چ", "پ", "ج"]
            full_weekday_names = ["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه", "پنج‌شنبه", "جمعه"]
            
            delta_days = (target_date - today).days
            is_editable = not is_locked
            if is_editable:
                if delta_days < 0 and past_limit is not None and past_limit != -1 and abs(delta_days) > past_limit:
                    is_editable = False
                elif delta_days > 0 and future_limit is not None and future_limit != -1 and delta_days > future_limit:
                    is_editable = False

            date_str = f"{year_month}/{d:02d}"
            days_meta.append({
                'day': d,
                'date_shamsi': date_str,
                'weekday_short': weekday_names[weekday],
                'weekday_name': full_weekday_names[weekday],
                'is_friday': is_friday,
                'is_editable': is_editable,
                'is_today': (target_date == today)
            })

        # پرسنل تاییدشده یا پرسنل غیرفعالی که در این ماه کارکرد دارند (مستقل از انبار یا فیلتر انبار خاص)
        if warehouse_id:
            personnel_list = PersonnelProfile.objects.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True),
                Q(is_active=True, approval_status='approved') | Q(daily_attendances__date_shamsi__startswith=year_month, daily_attendances__is_deleted=False)
            ).distinct().order_by('last_name', 'first_name')
            if section_id:
                personnel_list = personnel_list.filter(
                    Q(section_id=section_id) |
                    Q(section_assignments__section_id=section_id, section_assignments__is_active=True)
                ).distinct()
            attendances = DailyAttendance.objects.filter(
                Q(warehouse_id=warehouse_id) | Q(personnel__in=personnel_list),
                date_shamsi__startswith=year_month,
                is_deleted=False
            ).order_by('id')
            if section_id:
                attendances = attendances.filter(section_id=section_id)
        else:
            personnel_list = PersonnelProfile.objects.filter(
                Q(is_active=True, approval_status='approved') | Q(daily_attendances__date_shamsi__startswith=year_month, daily_attendances__is_deleted=False)
            ).distinct().order_by('last_name', 'first_name')
            if section_id:
                personnel_list = personnel_list.filter(
                    Q(section_id=section_id) |
                    Q(section_assignments__section_id=section_id, section_assignments__is_active=True)
                ).distinct()
            attendances = DailyAttendance.objects.filter(
                date_shamsi__startswith=year_month,
                is_deleted=False
            ).order_by('id')
            if section_id:
                attendances = attendances.filter(section_id=section_id)

        att_map = {}
        for a in attendances:
            try:
                d_num = int(a.date_shamsi.split('/')[2])
                att_map[(a.personnel_id, d_num)] = a
            except Exception:
                continue

        status_display_map = {
            'PRESENT_10H': 'حاضر ۱۰س',
            'HALF_5H': 'نیمه‌وقت ۵س',
            'ABSENT': 'غایب',
            'LEAVE': 'مرخصی',
            'MISSION': 'ماموریت',
            'FRIDAY_WORK': 'جمعه‌کاری',
            'CUSTOM': 'سفارشی'
        }

        grid_rows = []
        anomalies = []

        for p in personnel_list:
            p_job_title = p.job_title
            if section_id:
                sa = p.section_assignments.filter(section_id=section_id, is_active=True).first()
                if sa and sa.job_title:
                    p_job_title = sa.job_title
            p_days = []
            total_hours = 0.0
            total_overtime = 0.0
            present_days_count = 0
            recorded_days = []

            for dm in days_meta:
                d = dm['day']
                att = att_map.get((p.id, d))
                if att:
                    eff_h = float(att.effective_hours)
                    ovt_h = float(att.overtime_hours)
                    adv = float(att.advance_payment or 0)
                    st_disp = status_display_map.get(att.status, att.status)

                    p_days.append({
                        'day': d,
                        'date_shamsi': dm['date_shamsi'],
                        'status': att.status,
                        'status_display': st_disp,
                        'effective_hours': eff_h,
                        'overtime_hours': ovt_h,
                        'is_friday_work': att.is_friday_work,
                        'is_mission': att.is_mission,
                        'advance_payment': adv,
                        'notes': att.notes or '',
                        'is_existing': True,
                        'attendance_id': att.id
                    })
                    total_hours += eff_h
                    total_overtime += ovt_h
                    if att.status and att.status != 'ABSENT' and eff_h > 0:
                        present_days_count += 1
                    recorded_days.append(d)

                    # ۱. ناهنجاری کارکرد بیش از ۱۶ ساعت
                    if (eff_h + ovt_h) > 16:
                        anomalies.append({
                            'type': 'EXCESSIVE_HOURS',
                            'severity': 'high',
                            'personnel_id': p.id,
                            'full_name': p.full_name,
                            'day': d,
                            'message': f"کارکرد بیش از سقف مجاز ({eff_h + ovt_h:.1f} ساعت) برای {p.full_name} در روز {d} ثبت شده است."
                        })

                    # ۲. تضاد اضافه‌کار با غیبت یا مرخصی
                    if att.status in ['ABSENT', 'LEAVE'] and ovt_h > 0:
                        anomalies.append({
                            'type': 'CONFLICT_OVERTIME',
                            'severity': 'high',
                            'personnel_id': p.id,
                            'full_name': p.full_name,
                            'day': d,
                            'message': f"ثبت اضافه‌کار ({ovt_h:.1f} ساعت) در روز {d} برای {p.full_name} با وضعیت «{st_disp}» مغایرت دارد."
                        })

                    # ۳. مساعده نقدی در روز غیبت
                    if att.status == 'ABSENT' and adv > 0:
                        anomalies.append({
                            'type': 'ADVANCE_ABSENT',
                            'severity': 'medium',
                            'personnel_id': p.id,
                            'full_name': p.full_name,
                            'day': d,
                            'message': f"ثبت مساعده ({adv:,.0f} ریال) برای {p.full_name} در روز {d} با وضعیت غایب مغایرت دارد."
                        })
                else:
                    p_days.append({
                        'day': d,
                        'date_shamsi': dm['date_shamsi'],
                        'status': '',
                        'status_display': '—',
                        'effective_hours': 0.0,
                        'overtime_hours': 0.0,
                        'is_friday_work': False,
                        'is_mission': False,
                        'advance_payment': 0.0,
                        'notes': '',
                        'is_existing': False,
                        'attendance_id': None
                    })

            # ۴. شناسایی روزهای خالی و ثبت‌نشده در میان ماه
            if len(recorded_days) >= 2:
                min_d, max_d = min(recorded_days), max(recorded_days)
                for d_i in range(min_d + 1, max_d):
                    if d_i not in recorded_days:
                        dm_meta = next((dm for dm in days_meta if dm['day'] == d_i), None)
                        if dm_meta and not dm_meta['is_friday']:
                            anomalies.append({
                                'type': 'GAP_DAY',
                                'severity': 'medium',
                                'personnel_id': p.id,
                                'full_name': p.full_name,
                                'day': d_i,
                                'message': f"روز کاری {d_i} برای {p.full_name} تعیین وضعیت نشده و خالی مانده است."
                            })

            grid_rows.append({
                'personnel_id': p.id,
                'full_name': p.full_name,
                'national_code': p.national_code,
                'job_title': p_job_title,
                'is_active': p.is_active,
                'total_hours': round(total_hours, 2),
                'total_overtime': round(total_overtime, 2),
                'present_days': present_days_count,
                'days': p_days
            })

        period_info = {
            'id': period_obj.id if period_obj else None,
            'status': period_status,
            'status_display': period_obj.get_status_display() if period_obj else 'باز (امکان ثبت و ویرایش)',
            'is_locked': is_locked,
            'submitted_at': period_obj.submitted_at if period_obj else None,
            'submitted_by_name': f"{period_obj.submitted_by.first_name} {period_obj.submitted_by.last_name}".strip() if period_obj and period_obj.submitted_by else None,
            'submission_notes': period_obj.submission_notes if period_obj else '',
            'rejection_reason': period_obj.rejection_reason if period_obj else '',
            'locked_at': period_obj.locked_at if period_obj else None,
            'locked_by_name': f"{period_obj.locked_by.first_name} {period_obj.locked_by.last_name}".strip() if period_obj and period_obj.locked_by else None,
        }

        return Response({
            'warehouse_id': warehouse_id,
            'year_month': year_month,
            'month_name': month_name,
            'days_in_month': days_in_month,
            'is_locked': is_locked,
            'period_status': period_status,
            'period_info': period_info,
            'settings_window': {
                'past_days': past_limit,
                'future_days': future_limit
            },
            'days_meta': days_meta,
            'rows': grid_rows,
            'anomalies': anomalies
        })

    @action(detail=False, methods=['post'], url_path='bulk-save-monthly-grid')
    def bulk_save_monthly_grid(self, request):
        """
        ذخیره گروهی سلول‌های تغییریافته شیت ماهانه
        (پشتیبانی از مستقل از انبار، با اعتبارسنجی بازه تاریخی و ثبت لاگ ممیزی)
        """
        serializer = BulkMonthlyAttendanceGridSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        raw_wh = serializer.validated_data.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_ym = serializer.validated_data['year_month']
        ym_cleaned = normalize_digits(str(raw_ym)).strip().replace('-', '/')
        parts = ym_cleaned.split('/')
        if len(parts) >= 2:
            year_month = f"{int(parts[0]):04d}/{int(parts[1]):02d}"
        else:
            year_month = raw_ym

        items = serializer.validated_data['items']
        section_id = serializer.validated_data.get('section_id')
        project_id = serializer.validated_data.get('project_id')

        allow_override = bool(request.data.get('is_override', False) or request.headers.get('X-Admin-Override') == 'true')
        is_admin_override = allow_override and (getattr(request.user, 'is_superuser', False) or request.user.has_perm('personnel.can_override_attendance_lock'))
        period = None
        if warehouse_id:
            period, _ = MonthlyWorkPeriod.objects.get_or_create(
                warehouse_id=warehouse_id,
                year_month=year_month,
                defaults={'status': 'OPEN'}
            )
            if period.status == 'LOCKED' and not is_admin_override:
                return Response({'error': f'دوره کارکرد {year_month} قفل شده است و امکان ویرایش وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        # واکشی یکجای رکوردهای موجود ماه جاری جهت تشخیص تغییرات (فقط رکوردهای تغییریافته نیاز به ولیدیشن تاریخ دارند)
        existing_qs = DailyAttendance.objects.filter(
            date_shamsi__startswith=year_month,
            is_deleted=False
        )
        if warehouse_id:
            existing_qs = existing_qs.filter(warehouse_id=warehouse_id)
        if section_id:
            existing_qs = existing_qs.filter(section_id=section_id)
        existing_map = {(att.personnel_id, att.date_shamsi): att for att in existing_qs}

        # اعتبارسنجی فقط برای روزهای جدید یا تغییریافته
        for item in items:
            if not item.get('status') and float(item.get('effective_hours', 0)) == 0:
                continue
            day_str = f"{item['day']:02d}"
            date_shamsi = f"{year_month}/{day_str}"
            personnel_id = item['personnel_id']

            existing_att = existing_map.get((personnel_id, date_shamsi))
            is_changed = True
            if existing_att:
                status_val = item.get('status') or ('PRESENT_10H' if float(item.get('effective_hours', 0)) >= 10 else ('HALF_5H' if float(item.get('effective_hours', 0)) >= 5 else 'CUSTOM'))
                if (
                    str(existing_att.status) == str(status_val) and
                    float(existing_att.effective_hours) == float(item.get('effective_hours', 0)) and
                    float(existing_att.overtime_hours) == float(item.get('overtime_hours', 0)) and
                    bool(existing_att.is_friday_work) == bool(item.get('is_friday_work', False)) and
                    bool(existing_att.is_mission) == bool(item.get('is_mission', False)) and
                    float(existing_att.advance_payment) == float(item.get('advance_payment', 0)) and
                    (existing_att.notes or '') == (item.get('notes', '') or '')
                ):
                    is_changed = False

            if is_changed and not is_admin_override:
                try:
                    validate_attendance_date_window(date_shamsi, user=request.user, allow_override=is_admin_override, warehouse_id=warehouse_id, project_id=project_id)
                except Exception as e:
                    err_msg = getattr(e, 'detail', str(e))
                    if isinstance(err_msg, list):
                        err_msg = err_msg[0]
                    return Response({'error': f"خطا در روز {item['day']}: {err_msg}"}, status=status.HTTP_400_BAD_REQUEST)

        saved_count = 0
        updated_count = 0

        with transaction.atomic():
            for item in items:
                status_val = item.get('status') or ''
                eff_h = float(item.get('effective_hours', 0))
                day_str = f"{item['day']:02d}"
                date_shamsi = f"{year_month}/{day_str}"
                personnel_id = item['personnel_id']

                p_obj = PersonnelProfile.objects.filter(id=personnel_id).first()
                if not p_obj:
                    continue

                if p_obj.approval_status != 'approved':
                    return Response({
                        'error': f'امکان ثبت کارکرد برای پرسنل «{p_obj.full_name}» وجود ندارد زیرا پرونده او در وضعیت «{p_obj.get_approval_status_display()}» بوده و هنوز به تایید نهایی مدیر و حسابدار نرسیده است.'
                    }, status=status.HTTP_400_BAD_REQUEST)

                from .lifecycle_services import validate_personnel_validity
                is_valid_date, val_err = validate_personnel_validity(p_obj, date_shamsi)
                if not is_valid_date:
                    return Response({'error': val_err}, status=status.HTTP_400_BAD_REQUEST)

                check_wh = warehouse_id if warehouse_id else p_obj.assigned_warehouse_id
                target_period = period
                if not target_period and check_wh:
                    target_period = MonthlyWorkPeriod.objects.filter(
                        warehouse_id=check_wh,
                        year_month=year_month
                    ).first()
                if target_period and target_period.status == 'LOCKED' and not is_admin_override:
                    return Response({'error': f'دوره کارکرد {year_month} برای انبار پرسنل ({p_obj.full_name}) قفل شده است.'}, status=status.HTTP_400_BAD_REQUEST)

                target_wh = warehouse_id if warehouse_id else p_obj.assigned_warehouse_id

                # پیدا کردن رکورد فعال پرسنل در این تاریخ و این بخش
                att_filter = {
                    'personnel_id': personnel_id,
                    'date_shamsi': date_shamsi,
                    'is_deleted': False
                }
                if section_id:
                    att_filter['section_id'] = section_id
                att_obj = DailyAttendance.objects.select_for_update().filter(**att_filter).first()

                if not att_obj and not status_val and eff_h == 0:
                    continue

                if not status_val:
                    status_val = 'PRESENT_10H' if eff_h >= 10 else ('HALF_5H' if eff_h >= 5 else 'CUSTOM')

                ot_h = float(item.get('overtime_hours', 0))
                is_fri = bool(item.get('is_friday_work', False))
                is_mis = bool(item.get('is_mission', False))
                adv_pay = float(item.get('advance_payment', 0))
                notes_val = item.get('notes', '') or ''

                # اعتبارسنجی و ریست منطقی ساعات در صورت غیبت یا مرخصی
                if status_val in ['ABSENT', 'LEAVE']:
                    eff_h = 0.0
                    ot_h = 0.0
                    is_fri = False
                    is_mis = False

                if att_obj:
                    changes = []
                    item_sanitized = {
                        'status': status_val,
                        'effective_hours': eff_h,
                        'overtime_hours': ot_h,
                        'is_friday_work': is_fri,
                        'is_mission': is_mis,
                        'advance_payment': adv_pay,
                        'notes': notes_val
                    }
                    fields_to_check = ['status', 'effective_hours', 'overtime_hours', 'is_friday_work', 'is_mission', 'advance_payment', 'notes']
                    for f in fields_to_check:
                        new_val = item_sanitized.get(f)
                        old_val = getattr(att_obj, f)
                        is_diff = False
                        if f in ['effective_hours', 'overtime_hours', 'advance_payment']:
                            if float(new_val or 0) != float(old_val or 0):
                                is_diff = True
                        elif f in ['is_friday_work', 'is_mission']:
                            if bool(new_val) != bool(old_val):
                                is_diff = True
                        else:
                            if str(new_val or '') != str(old_val or ''):
                                is_diff = True
                        if is_diff:
                            changes.append((f, str(old_val), str(new_val)))

                    # اگر تغییری نکرده بود از ذخیره و لاگ تکراری صرف‌نظر می‌کنیم
                    if not changes:
                        continue

                    att_obj.status = status_val
                    att_obj.effective_hours = eff_h
                    att_obj.overtime_hours = ot_h
                    att_obj.is_friday_work = is_fri
                    att_obj.is_mission = is_mis
                    att_obj.advance_payment = adv_pay
                    att_obj.notes = notes_val
                    if target_wh:
                        att_obj.warehouse_id = target_wh
                    if section_id:
                        att_obj.section_id = section_id
                    att_obj.modified_by = request.user
                    if target_period:
                        att_obj.period = target_period
                    att_obj.save()

                    # پاکسازی رکوردهای تکراری قدیمی احتمالی برای این روز در همین بخش
                    dup_qs = DailyAttendance.objects.filter(
                        personnel_id=personnel_id,
                        date_shamsi=date_shamsi,
                        is_deleted=False
                    ).exclude(id=att_obj.id)
                    if section_id:
                        dup_qs = dup_qs.filter(section_id=section_id)
                    dup_qs.delete()

                    for f_name, o_val, n_val in changes:
                        AttendanceAuditLog.objects.create(
                            attendance=att_obj,
                            personnel_name=p_obj.full_name,
                            date_shamsi=date_shamsi,
                            changed_by=request.user,
                            field_name=f_name,
                            old_value=o_val,
                            new_value=n_val,
                            reason='ویرایش شیت ماهانه کارکرد'
                        )
                    updated_count += 1
                else:
                    att_obj = DailyAttendance.objects.create(
                        personnel_id=personnel_id,
                        warehouse_id=target_wh,
                        section_id=section_id or p_obj.section_id,
                        project_id=project_id or p_obj.project_id,
                        date_shamsi=date_shamsi,
                        status=status_val,
                        effective_hours=eff_h,
                        overtime_hours=ot_h,
                        is_friday_work=is_fri,
                        is_mission=is_mis,
                        advance_payment=adv_pay,
                        notes=notes_val,
                        period=target_period or period,
                        created_by=request.user,
                        modified_by=request.user
                    )
                    saved_count += 1

        client_tab_id = request.data.get('client_tab_id') or request.headers.get('X-Client-Tab-ID')
        extra_data = {}
        if section_id:
            extra_data['section_id'] = section_id
        if project_id:
            extra_data['project_id'] = project_id

        broadcast_attendance_updated(
            warehouse_id=warehouse_id,
            date_shamsi=f"{year_month}/01",
            year_month=year_month,
            sender_id=request.user.id if request.user and request.user.is_authenticated else None,
            client_tab_id=client_tab_id,
            message=f'شیت کارکرد ماهانه {year_month} به‌روزرسانی شد.',
            extra_data=extra_data
        )

        return Response({
            'message': f'اطلاعات شیت ماهانه با موفقیت ذخیره شد ({saved_count} جدید، {updated_count} بروزرسانی).',
            'saved_count': saved_count,
            'updated_count': updated_count
        })

    @action(detail=False, methods=['get'], url_path='export-monthly-excel')
    def export_monthly_excel(self, request):
        """
        خروجی استاندارد دو سطری اکسل از شیت ماهانه کارکرد پرسنل
        (شامل فرمول‌های جمع، استایل زیبا و فریز ستون‌ها طبق استاندارد)
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_sec = request.query_params.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_proj = request.query_params.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_ym = request.query_params.get('year_month')
        if not raw_ym:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        ym_cleaned = normalize_digits(str(raw_ym)).strip().replace('-', '/')
        parts = ym_cleaned.split('/')
        if len(parts) >= 2:
            y, m = int(parts[0]), int(parts[1])
            year_month = f"{y:04d}/{m:02d}"
        else:
            return Response({'error': 'فرمت year_month نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        if 1 <= m <= 6:
            days_in_month = 31
        elif 7 <= m <= 11:
            days_in_month = 30
        elif m == 12:
            days_in_month = 30 if jdatetime.date(y, 1, 1).isleap() else 29
        else:
            return Response({'error': 'شماره ماه باید بین ۱ تا ۱۲ باشد.'}, status=status.HTTP_400_BAD_REQUEST)

        month_names = ["", "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
        month_name = month_names[m] if 1 <= m <= 12 else ""

        fridays = set()
        weekday_names = ["ش", "ی", "د", "س", "چ", "پ", "ج"]
        day_headers = []
        for d in range(1, days_in_month + 1):
            t_date = jdatetime.date(y, m, d)
            w = t_date.weekday()
            if w == 6:
                fridays.add(d)
            day_headers.append(f"{d} ({weekday_names[w]})")

        personnel_list = PersonnelProfile.objects.filter(is_active=True)
        attendances = DailyAttendance.objects.filter(
            date_shamsi__startswith=year_month,
            is_deleted=False
        )

        if section_id:
            personnel_list = personnel_list.filter(
                Q(section_id=section_id) |
                Q(section_assignments__section_id=section_id, section_assignments__is_active=True)
            ).distinct()
            attendances = attendances.filter(section_id=section_id)
        elif warehouse_id:
            personnel_list = personnel_list.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True)
            )
            attendances = attendances.filter(warehouse_id=warehouse_id)

        personnel_list = personnel_list.order_by('last_name', 'first_name')

        att_map = {(a.personnel_id, int(a.date_shamsi.split('/')[2])): a for a in attendances}

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"کارکرد {month_name} {y}"
        ws.sheet_view.rightToLeft = True

        font_header1 = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        font_header2 = Font(name='Calibri', size=9, bold=False, color='94A3B8')
        font_body = Font(name='B Nazanin', size=10, bold=False, color='0F172A')
        font_bold = Font(name='B Nazanin', size=10, bold=True, color='0F172A')
        font_mono = Font(name='Consolas', size=10, bold=True, color='0F172A')

        fill_header1 = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        fill_header2 = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
        fill_friday = PatternFill(start_color='F3E8FF', end_color='F3E8FF', fill_type='solid')
        fill_total = PatternFill(start_color='EEF2FF', end_color='EEF2FF', fill_type='solid')
        fill_stripe = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        align_right = Alignment(horizontal='right', vertical='center')

        row1_titles = ['ردیف', 'کد ملی', 'نام و نام خانوادگی', 'عنوان شغلی', 'انبار اختصاصی'] + day_headers + ['جمع ساعت کارکرد', 'جمع اضافه کار', 'مجموع کل ساعات (کارکرد + اضافه‌کار)', 'تعداد روز حضور']
        row2_keys = ['row_num', 'national_code', 'full_name', 'job_title', 'warehouse_name'] + [f'day_{d:02d}' for d in range(1, days_in_month + 1)] + ['total_hours', 'total_overtime', 'total_hours_sum', 'present_days']

        for col_idx, (t1, t2) in enumerate(zip(row1_titles, row2_keys), start=1):
            c1 = ws.cell(row=1, column=col_idx, value=t1)
            c1.font = font_header1
            c1.fill = fill_header1
            c1.alignment = align_center
            c1.border = thin_border

            c2 = ws.cell(row=2, column=col_idx, value=t2)
            c2.font = font_header2
            c2.fill = fill_header2
            c2.alignment = align_center
            c2.border = thin_border

        for p_idx, p in enumerate(personnel_list, start=1):
            row_num = p_idx + 2
            is_even = (p_idx % 2 == 0)

            ws.cell(row=row_num, column=1, value=p_idx).alignment = align_center
            ws.cell(row=row_num, column=2, value=p.national_code or '').alignment = align_center
            ws.cell(row=row_num, column=3, value=p.full_name).alignment = align_right
            ws.cell(row=row_num, column=4, value=p.job_title or '').alignment = align_right
            ws.cell(row=row_num, column=5, value=p.assigned_warehouse.name if p.assigned_warehouse else 'تمام انبارها').alignment = align_center

            tot_hours = 0.0
            tot_ot = 0.0
            p_days = 0

            for d in range(1, days_in_month + 1):
                col_idx = 5 + d
                att = att_map.get((p.id, d))
                cell = ws.cell(row=row_num, column=col_idx)
                cell.alignment = align_center

                if att:
                    h = float(att.effective_hours)
                    ot = float(att.overtime_hours)
                    tot_hours += h
                    tot_ot += ot
                    if att.status and att.status != 'ABSENT' and h > 0:
                        p_days += 1

                    if att.status == 'LEAVE':
                        cell.value = 'مرخصی'
                    elif att.status == 'ABSENT':
                        cell.value = 'غایب'
                    elif att.status == 'MISSION':
                        cell.value = f"مامور ({h:g})"
                    elif att.status == 'FRIDAY_WORK':
                        cell.value = f"جمعه ({h:g})"
                    elif h > 0:
                        cell.value = h
                    else:
                        cell.value = ''
                else:
                    cell.value = ''

                if d in fridays:
                    cell.fill = fill_friday
                elif is_even:
                    cell.fill = fill_stripe

                cell.border = thin_border
                cell.font = font_body

            first_col_letter = get_column_letter(6)
            last_col_letter = get_column_letter(5 + days_in_month)
            col_tot_letter = get_column_letter(6 + days_in_month)
            col_ot_letter = get_column_letter(7 + days_in_month)

            c_tot = ws.cell(row=row_num, column=6 + days_in_month, value=f"=SUM({first_col_letter}{row_num}:{last_col_letter}{row_num})")
            c_tot.font = font_bold
            c_tot.fill = fill_total
            c_tot.alignment = align_center
            c_tot.border = thin_border

            c_ot = ws.cell(row=row_num, column=7 + days_in_month, value=tot_ot)
            c_ot.font = font_mono
            c_ot.alignment = align_center
            c_ot.border = thin_border

            c_tot_all = ws.cell(row=row_num, column=8 + days_in_month, value=f"={col_tot_letter}{row_num}+{col_ot_letter}{row_num}")
            c_tot_all.font = font_bold
            c_tot_all.fill = fill_total
            c_tot_all.alignment = align_center
            c_tot_all.border = thin_border

            c_pdays = ws.cell(row=row_num, column=9 + days_in_month, value=f'=COUNTIF({first_col_letter}{row_num}:{last_col_letter}{row_num}, ">0")')
            c_pdays.font = font_bold
            c_pdays.alignment = align_center
            c_pdays.border = thin_border

            for c_i in range(1, 6):
                cell_i = ws.cell(row=row_num, column=c_i)
                cell_i.font = font_bold if c_i == 3 else font_body
                cell_i.border = thin_border
                if is_even and cell_i.fill.fill_type is None:
                    cell_i.fill = fill_stripe

        ws.freeze_panes = 'F3'
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 18

        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 14
        ws.column_dimensions['C'].width = 22
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 16
        for d in range(1, days_in_month + 1):
            col_letter = get_column_letter(5 + d)
            ws.column_dimensions[col_letter].width = 8
        ws.column_dimensions[get_column_letter(6 + days_in_month)].width = 14
        ws.column_dimensions[get_column_letter(7 + days_in_month)].width = 13
        ws.column_dimensions[get_column_letter(8 + days_in_month)].width = 20
        ws.column_dimensions[get_column_letter(9 + days_in_month)].width = 13

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"Timesheet_{year_month.replace('/', '_')}.xlsx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    @action(detail=False, methods=['post'], url_path='import-monthly-excel')
    def import_monthly_excel(self, request):
        """
        بارگذاری مستقیم و هوشمند شیت ماهانه از فایل اکسل
        (تطبیق بر اساس ردیف کلیدها و کد ملی پرسنل همراه با ذخیره اتمیک)
        """
        import openpyxl
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        raw_wh = request.data.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_sec = request.data.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_proj = request.data.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_ym = request.data.get('year_month')
        if not raw_ym:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        ym_cleaned = normalize_digits(str(raw_ym)).strip().replace('-', '/')
        parts = ym_cleaned.split('/')
        if len(parts) >= 2:
            year_month = f"{int(parts[0]):04d}/{int(parts[1]):02d}"
        else:
            return Response({'error': 'فرمت year_month نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        allow_override = bool(request.data.get('is_override', False) or request.headers.get('X-Admin-Override') == 'true')
        is_admin_override = allow_override and (getattr(request.user, 'is_superuser', False) or request.user.has_perm('personnel.can_override_attendance_lock'))

        if warehouse_id:
            period = MonthlyWorkPeriod.objects.filter(warehouse_id=warehouse_id, year_month=year_month).first()
            if period and period.status in ['LOCKED', 'SUBMITTED', 'FINALIZED'] and not is_admin_override:
                return Response({'error': f'دوره کارکرد {year_month} در وضعیت «{period.get_status_display()}» است و امکان بارگذاری وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            ws = wb.active
        except Exception as e:
            return Response({'error': f'خطا در باز کردن فایل اکسل: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        nat_col_idx = None
        day_cols = {}

        for col_idx in range(1, ws.max_column + 1):
            k2 = str(ws.cell(row=2, column=col_idx).value or '').strip()
            k1 = str(ws.cell(row=1, column=col_idx).value or '').strip()

            if k2 == 'national_code' or 'کد ملی' in k1:
                nat_col_idx = col_idx
            elif k2.startswith('day_'):
                try:
                    d_num = int(k2.replace('day_', ''))
                    day_cols[d_num] = col_idx
                except ValueError:
                    pass
            elif k1.startswith('روز '):
                try:
                    d_num = int(k1.replace('روز ', '').split()[0].split('(')[0])
                    day_cols[d_num] = col_idx
                except ValueError:
                    pass

        if not nat_col_idx:
            return Response({'error': 'ستون کد ملی پرسنل در فایل اکسل یافت نشد.'}, status=status.HTTP_400_BAD_REQUEST)
        if not day_cols:
            return Response({'error': 'هیچ ستونی برای روزهای ماه در فایل اکسل شناسایی نشد.'}, status=status.HTTP_400_BAD_REQUEST)

        personnel_qs = PersonnelProfile.objects.filter(is_active=True)
        if section_id:
            personnel_qs = personnel_qs.filter(
                Q(section_id=section_id) |
                Q(section_assignments__section_id=section_id, section_assignments__is_active=True)
            ).distinct()
        elif warehouse_id:
            personnel_qs = personnel_qs.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True)
            )

        personnel_map = {
            normalize_digits(str(p.national_code)).split('.')[0].strip().zfill(10): p
            for p in personnel_qs
            if p.national_code
        }

        updated_cells = 0
        matched_personnel_count = 0

        with transaction.atomic():
            for row_idx in range(3, ws.max_row + 1):
                raw_nat = ws.cell(row=row_idx, column=nat_col_idx).value
                if raw_nat is None or str(raw_nat).strip() == '':
                    continue
                raw_nat_str = str(raw_nat).split('.')[0].strip()
                nat_code = normalize_digits(raw_nat_str).strip().zfill(10)
                p_obj = personnel_map.get(nat_code)
                if not p_obj:
                    continue

                matched_personnel_count += 1
                target_wh = warehouse_id if warehouse_id else p_obj.assigned_warehouse_id

                target_period = None
                if target_wh:
                    target_period, _ = MonthlyWorkPeriod.objects.get_or_create(
                        warehouse_id=target_wh,
                        year_month=year_month,
                        defaults={'status': 'OPEN'}
                    )
                    locked_period_statuses = [
                        'LOCKED', 'SUBMITTED', 'FINALIZED', 'SUBMITTED_SUPERVISOR',
                        'PENDING_ACCOUNTANT', 'SUBMITTED_ACCOUNTANT', 'PENDING_MANAGER',
                        'READY_TO_PAY', 'PAID'
                    ]
                    if target_period and target_period.status in locked_period_statuses and not is_admin_override:
                        return Response({
                            'error': f'امکان بارگذاری اکسل وجود ندارد؛ دوره کارکرد {year_month} برای انبار پرسنل ({p_obj.full_name}) در وضعیت «{target_period.get_status_display()}» است.'
                        }, status=status.HTTP_400_BAD_REQUEST)

                for d_num, col_i in day_cols.items():
                    val = ws.cell(row=row_idx, column=col_i).value
                    if val is None or str(val).strip() == '' or str(val).strip() == '—':
                        continue

                    val_str = str(val).strip()
                    val_str_clean = normalize_digits(val_str)
                    date_shamsi = f"{year_month}/{d_num:02d}"

                    status_val = 'PRESENT_10H'
                    eff_hours = 10.0
                    ot_hours = 0.0
                    is_friday = False
                    is_mission = False

                    try:
                        num = float(val_str_clean)
                        if num == 10:
                            status_val = 'PRESENT_10H'
                            eff_hours = 10.0
                        elif num == 5:
                            status_val = 'HALF_5H'
                            eff_hours = 5.0
                        elif num == 0:
                            status_val = 'ABSENT'
                            eff_hours = 0.0
                        else:
                            status_val = 'CUSTOM'
                            eff_hours = num
                    except ValueError:
                        if 'مرخصی' in val_str:
                            status_val = 'LEAVE'
                            eff_hours = 0.0
                        elif 'غایب' in val_str:
                            status_val = 'ABSENT'
                            eff_hours = 0.0
                        elif 'مامور' in val_str:
                            status_val = 'MISSION'
                            is_mission = True
                            eff_hours = 10.0
                        elif 'جمعه' in val_str:
                            status_val = 'FRIDAY_WORK'
                            is_friday = True
                            eff_hours = 10.0
                        elif 'نیمه' in val_str:
                            status_val = 'HALF_5H'
                            eff_hours = 5.0
                        elif 'حاضر' in val_str:
                            status_val = 'PRESENT_10H'
                            eff_hours = 10.0

                    att_filter = {'personnel_id': p_obj.id, 'date_shamsi': date_shamsi, 'is_deleted': False}
                    if target_wh:
                        att_filter['warehouse_id'] = target_wh

                    att_obj = DailyAttendance.objects.select_for_update().filter(**att_filter).first()
                    if not att_obj and not target_wh:
                        att_obj = DailyAttendance.objects.select_for_update().filter(personnel_id=p_obj.id, date_shamsi=date_shamsi, is_deleted=False).first()

                    # اعتبارسنجی بازه تاریخی در صورت تغییر یا ایجاد
                    is_changed = False
                    if not att_obj:
                        is_changed = True
                    else:
                        if (att_obj.status != status_val or float(att_obj.effective_hours) != eff_hours or
                            att_obj.is_friday_work != is_friday or att_obj.is_mission != is_mission):
                            is_changed = True

                    if is_changed and not is_admin_override:
                        try:
                            target_pid = getattr(p_obj, 'project_id', None)
                            validate_attendance_date_window(
                                date_shamsi=date_shamsi,
                                user=request.user,
                                allow_override=is_admin_override,
                                warehouse_id=target_wh,
                                project_id=target_pid
                            )
                        except Exception as e:
                            err_msg = getattr(e, 'detail', str(e))
                            if isinstance(err_msg, list):
                                err_msg = err_msg[0]
                            return Response({
                                'error': f"خطا در ردیف پرسنل {p_obj.full_name} ({date_shamsi}): {err_msg}"
                            }, status=status.HTTP_400_BAD_REQUEST)

                    if att_obj:
                        changes = []
                        if att_obj.status != status_val:
                            changes.append(('status', str(att_obj.status), str(status_val)))
                        if float(att_obj.effective_hours) != eff_hours:
                            changes.append(('effective_hours', str(att_obj.effective_hours), str(eff_hours)))
                        if att_obj.is_friday_work != is_friday:
                            changes.append(('is_friday_work', str(att_obj.is_friday_work), str(is_friday)))
                        if att_obj.is_mission != is_mission:
                            changes.append(('is_mission', str(att_obj.is_mission), str(is_mission)))

                        if changes:
                            att_obj.status = status_val
                            att_obj.effective_hours = eff_hours
                            att_obj.is_friday_work = is_friday
                            att_obj.is_mission = is_mission
                            if target_wh:
                                att_obj.warehouse_id = target_wh
                            if section_id:
                                att_obj.section_id = section_id
                            if project_id:
                                att_obj.project_id = project_id
                            if target_period:
                                att_obj.period = target_period
                            # در صورت ویرایش رکوردی که قبلا تایید شده بود، تاییدیه باطل شده و به وضعیت در انتظار بازمی‌گردد
                            if changes and att_obj.approval_status == 'APPROVED':
                                att_obj.approval_status = 'PENDING_SUPERVISOR'
                                att_obj.supervisor_approved_by = None
                                att_obj.supervisor_approved_at = None
                                changes.append(('approval_status', 'APPROVED', 'PENDING_SUPERVISOR'))

                            att_obj.modified_by = request.user
                            att_obj.save()

                            for f_name, o_val, n_val in changes:
                                AttendanceAuditLog.objects.create(
                                    attendance=att_obj,
                                    personnel_name=p_obj.full_name,
                                    date_shamsi=date_shamsi,
                                    changed_by=request.user,
                                    field_name=f_name,
                                    old_value=o_val,
                                    new_value=n_val,
                                    reason='ویرایش از طریق بارگذاری اکسل شیت ماهانه'
                                )
                            updated_cells += 1
                    else:
                        att_obj = DailyAttendance.objects.create(
                            personnel=p_obj,
                            warehouse_id=target_wh,
                            section_id=section_id or p_obj.section_id,
                            project_id=project_id or p_obj.project_id,
                            period=target_period,
                            date_shamsi=date_shamsi,
                            status=status_val,
                            effective_hours=eff_hours,
                            overtime_hours=ot_hours,
                            is_friday_work=is_friday,
                            is_mission=is_mission,
                            advance_payment=0,
                            created_by=request.user,
                            modified_by=request.user
                        )
                        AttendanceAuditLog.objects.create(
                            attendance=att_obj,
                            personnel_name=p_obj.full_name,
                            date_shamsi=date_shamsi,
                            changed_by=request.user,
                            field_name='all',
                            old_value=None,
                            new_value=f"{status_val} ({eff_hours}h)",
                            reason='ایجاد از طریق بارگذاری اکسل شیت ماهانه'
                        )
                        updated_cells += 1

        extra_data = {}
        if section_id:
            extra_data['section_id'] = section_id
        if project_id:
            extra_data['project_id'] = project_id

        broadcast_attendance_updated(
            warehouse_id=warehouse_id,
            date_shamsi=f"{year_month}/01",
            year_month=year_month,
            sender_id=request.user.id if request.user and request.user.is_authenticated else None,
            message=f'اطلاعات شیت کارکرد ماهانه {year_month} از فایل اکسل بارگذاری و به‌روزرسانی شد.',
            extra_data=extra_data
        )

        return Response({
            'message': f'اطلاعات شیت کارکرد با موفقیت از اکسل بارگذاری شد ({matched_personnel_count} پرسنل، {updated_cells} سلول کارکرد ثبت گردید).',
            'matched_personnel_count': matched_personnel_count,
            'updated_cells': updated_cells
        })

    @action(detail=False, methods=['get'], url_path='monthly-summary')
    def get_monthly_summary(self, request):
        """
        گزارش تجمیعی ماهانه کارکرد پرسنل با پشتیبانی از فیلتر بخش/پروژه و تفکیک وضعیت‌های تایید
        """
        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '', 'UNDEFINED', 'NULL']) else None
        raw_sec = request.query_params.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '', 'UNDEFINED', 'NULL']) else None
        raw_proj = request.query_params.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '', 'UNDEFINED', 'NULL']) else None

        year_month = request.query_params.get('year_month')
        if not year_month:
            latest_period = MonthlyWorkPeriod.objects.order_by('-year_month').first()
            year_month = latest_period.year_month if latest_period else '1405/06'

        norm_ym = year_month.replace('-', '/')

        # فیلتر پرسنل
        personnel_list = PersonnelProfile.objects.filter(is_active=True)
        if section_id:
            personnel_list = personnel_list.filter(
                Q(section_id=section_id) |
                Q(section_assignments__section_id=section_id, section_assignments__is_active=True)
            ).distinct()
        elif project_id:
            personnel_list = personnel_list.filter(
                Q(section__project_id=project_id) |
                Q(section_assignments__section__project_id=project_id, section_assignments__is_active=True)
            ).distinct()
        elif warehouse_id:
            personnel_list = personnel_list.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True)
            )

        cid = validate_user_company_access(request.user, get_request_company_id(request))
        if cid:
            personnel_list = personnel_list.filter(
                Q(company_id=cid) | Q(section__project__company_id=cid)
            )

        personnel_list = personnel_list.order_by('last_name', 'first_name')

        # فیلتر کارکردها
        att_filter = Q(date_shamsi__startswith=norm_ym, is_deleted=False)
        if section_id:
            att_filter &= (Q(section_id=section_id) | (Q(section_id__isnull=True) & Q(personnel__section_id=section_id)))
        elif project_id:
            att_filter &= (Q(project_id=project_id) | (Q(project_id__isnull=True) & Q(personnel__project_id=project_id)))
        elif warehouse_id:
            att_filter &= Q(warehouse_id=warehouse_id)

        attendances = DailyAttendance.objects.filter(att_filter)

        # جستجوی دوره متناظر
        period = None
        if section_id:
            period = MonthlyWorkPeriod.objects.filter(section_id=section_id, year_month=year_month).first()
        elif project_id:
            period = MonthlyWorkPeriod.objects.filter(project_id=project_id, year_month=year_month).first()
        elif warehouse_id:
            period = MonthlyWorkPeriod.objects.filter(warehouse_id=warehouse_id, year_month=year_month).first()
        if not period:
            period = MonthlyWorkPeriod.objects.filter(year_month=year_month).first()

        # محاسبه کارکردهای تاییدنشده
        pending_approvals_count = attendances.exclude(approval_status='APPROVED').count()

        summary_rows = []
        for p in personnel_list:
            p_atts = [a for a in attendances if a.personnel_id == p.id]
            
            total_hours = sum(float(a.effective_hours) for a in p_atts)
            total_overtime = sum(float(a.overtime_hours) for a in p_atts)
            friday_days = sum(1 for a in p_atts if a.is_friday_work)
            mission_days = sum(1 for a in p_atts if a.is_mission)
            absent_days = sum(1 for a in p_atts if a.status == 'ABSENT')
            leave_days = sum(1 for a in p_atts if a.status == 'LEAVE')
            present_days = sum(1 for a in p_atts if a.status in ['PRESENT_10H', 'HALF_5H', 'FRIDAY_WORK', 'MISSION', 'CUSTOM'] and a.effective_hours > 0)
            total_advances = sum(float(a.advance_payment) for a in p_atts)
            
            # تبدیل به معادل روز (تقسیم بر ۱۰)
            equivalent_days = round(total_hours / 10.0, 2)
            
            # برآورد دستمزد ناخالص کارکرد
            daily_wage = float(p.daily_base_wage)
            hourly_wage = float(p.hourly_rate)
            gross_base_pay = round(equivalent_days * daily_wage, 0)
            overtime_pay = round(total_overtime * (hourly_wage * 1.4), 0)

            summary_rows.append({
                'personnel_id': p.id,
                'full_name': p.full_name,
                'national_code': p.national_code,
                'job_title': p.job_title,
                'contract_type': p.get_contract_type_display(),
                'daily_base_wage': daily_wage,
                'total_hours': round(total_hours, 1),
                'equivalent_days': equivalent_days,
                'present_days': present_days,
                'total_overtime_hours': round(total_overtime, 1),
                'friday_days': friday_days,
                'mission_days': mission_days,
                'absent_days': absent_days,
                'leave_days': leave_days,
                'total_advances': total_advances,
                'gross_base_pay': gross_base_pay,
                'overtime_pay': overtime_pay,
                'estimated_total_pay': gross_base_pay + overtime_pay - total_advances
            })

        locked_statuses = [
            'LOCKED', 'SUBMITTED', 'FINALIZED', 'SUBMITTED_SUPERVISOR',
            'PENDING_ACCOUNTANT', 'SUBMITTED_ACCOUNTANT', 'PENDING_MANAGER',
            'READY_TO_PAY', 'PAID'
        ]
        is_locked = period.status in locked_statuses if period else False

        return Response({
            'warehouse_id': warehouse_id,
            'section_id': section_id,
            'project_id': project_id,
            'year_month': year_month,
            'period_id': period.id if period else None,
            'period_status': period.status if period else 'OPEN',
            'is_locked': is_locked,
            'pending_approvals_count': pending_approvals_count,
            'summary': summary_rows
        })


class VehicleTripViewSet(viewsets.ModelViewSet):
    queryset = VehicleTripLog.objects.filter(is_deleted=False).select_related('vehicle', 'period')
    serializer_class = VehicleTripLogSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def _scoped_vehicles(self, qs=None):
        base_qs = VehicleDriverProfile.objects.all() if qs is None else qs
        return scope_vehicle_queryset(self.request, base_qs)

    def _scoped_trips(self, qs):
        cid = validate_user_company_access(self.request.user, get_request_company_id(self.request))
        return scope_trips_for_user(self.request.user, qs, company_id=cid)

    def _validate_write_target(self, request, section_id, project_id, warehouse_id=None):
        sec_obj = None
        if section_id:
            try:
                sec_id_int = int(section_id)
            except (ValueError, TypeError):
                return Response({'error': 'شناسه بخش نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
            sec_obj = ProjectSection.objects.filter(id=sec_id_int).select_related('project').first()
            if not sec_obj:
                return Response({'error': 'بخش مشخص‌شده یافت نشد.'}, status=status.HTTP_400_BAD_REQUEST)
            if not sec_obj.project or not sec_obj.project.company_id:
                return Response({'error': 'بخش مشخص‌شده فاقد پروژه یا شرکت معتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        proj_obj = None
        if project_id:
            try:
                proj_id_int = int(project_id)
            except (ValueError, TypeError):
                return Response({'error': 'شناسه پروژه نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
            proj_obj = FinancialProject.objects.filter(id=proj_id_int).first()
            if not proj_obj:
                return Response({'error': 'پروژه مشخص‌شده یافت نشد.'}, status=status.HTTP_400_BAD_REQUEST)
            if not proj_obj.company_id:
                return Response({'error': 'پروژه مشخص‌شده فاقد شرکت معتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        if sec_obj and proj_obj:
            if sec_obj.project_id != proj_obj.id:
                return Response({'error': 'بخش انتخابی متعلق به پروژه مشخص‌شده نیست.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            cid = validate_user_company_access(request.user, get_request_company_id(request))
        except PermissionDenied as e:
            return Response({'error': str(e.detail) if hasattr(e, 'detail') else str(e)}, status=status.HTTP_403_FORBIDDEN)

        if cid:
            if sec_obj and sec_obj.project.company_id != cid:
                return Response({'error': 'بخش مشخص‌شده متعلق به شرکت فعال در سربرگ نیست.'}, status=status.HTTP_403_FORBIDDEN)
            if proj_obj and proj_obj.company_id != cid:
                return Response({'error': 'پروژه مشخص‌شده متعلق به شرکت فعال در سربرگ نیست.'}, status=status.HTTP_403_FORBIDDEN)
        elif not request.user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(request.user).values_list('id', flat=True))
            if sec_obj and sec_obj.project.company_id not in allowed_ids:
                return Response({'error': 'شما به بخش یا شرکت مربوطه دسترسی ندارید.'}, status=status.HTTP_403_FORBIDDEN)
            if proj_obj and proj_obj.company_id not in allowed_ids:
                return Response({'error': 'شما به پروژه یا شرکت مربوطه دسترسی ندارید.'}, status=status.HTTP_403_FORBIDDEN)

        if warehouse_id is not None:
            from platform_core.registry import is_module_installed, company_of, contributions
            if is_module_installed('warehouse'):
                try:
                    wh_id_int = int(warehouse_id)
                except (ValueError, TypeError):
                    return Response({'error': 'شناسه انبار نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

                validators = contributions('validate_warehouse')
                wh_exists = any(v(wh_id_int) for v in validators) if validators else (company_of('warehouse', wh_id_int) is not None)
                if not wh_exists:
                    return Response({'error': 'انبار مشخص‌شده یافت نشد.'}, status=status.HTTP_400_BAD_REQUEST)

                wh_cid = company_of('warehouse', wh_id_int)
                if cid:
                    if wh_cid is not None and wh_cid != cid:
                        return Response({'error': 'انبار مشخص‌شده متعلق به شرکت فعال در سربرگ نیست.'}, status=status.HTTP_403_FORBIDDEN)

                if not request.user.is_superuser:
                    if wh_cid is not None:
                        allowed_ids = set(get_user_allowed_companies(request.user).values_list('id', flat=True))
                        if wh_cid not in allowed_ids:
                            return Response({'error': 'شما به انبار مشخص‌شده دسترسی ندارید.'}, status=status.HTTP_403_FORBIDDEN)

                    scoped_wh_ids = set()
                    for ext in contributions('user_extensions'):
                        if hasattr(ext, 'get_all_warehouses'):
                            try:
                                wh_list = ext.get_all_warehouses(user=request.user)
                                for w in wh_list:
                                    w_id = getattr(w, 'id', getattr(w, 'pk', None))
                                    if w_id is not None:
                                        scoped_wh_ids.add(w_id)
                            except Exception:
                                pass

                    # همچنین برای خودروهای یتیم خود کاربر (بی‌شرکت و بی‌بخش)، انبار بی‌شرکت خودرو در دامنه است
                    if wh_cid is None:
                        orphan_wh_ids = set(
                            VehicleDriverProfile.objects.filter(
                                created_by=request.user,
                                company__isnull=True,
                                section__isnull=True,
                                assigned_warehouse_id__isnull=False
                            ).values_list('assigned_warehouse_id', flat=True)
                        )
                        scoped_wh_ids.update(orphan_wh_ids)

                    if wh_id_int not in scoped_wh_ids:
                        return Response({'error': 'شما به انبار مشخص‌شده دسترسی ندارید.'}, status=status.HTTP_403_FORBIDDEN)

        return None

    def get_queryset(self):
        qs = super().get_queryset()
        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(warehouse_id=warehouse_id)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(Q(section_id=section_id) | Q(vehicle__section_id=section_id))
        project_id = self.request.query_params.get('project_id')
        if project_id:
            qs = qs.filter(Q(section__project_id=project_id) | Q(vehicle__section__project_id=project_id))
        date_shamsi = self.request.query_params.get('date_shamsi')
        if date_shamsi:
            qs = qs.filter(date_shamsi=date_shamsi)
        vehicle_id = self.request.query_params.get('vehicle_id')
        if vehicle_id:
            qs = qs.filter(vehicle_id=vehicle_id)
        return self._scoped_trips(qs)

    def perform_create(self, serializer):
        from .fleet_lifecycle_services import assert_vehicle_trip_allowed
        vehicle = serializer.validated_data.get('vehicle')
        date_shamsi = serializer.validated_data.get('date_shamsi')
        trip_count = serializer.validated_data.get('trip_count', 1)
        if (trip_count is None or trip_count > 0) and vehicle and date_shamsi:
            assert_vehicle_trip_allowed(vehicle, date_shamsi)
        serializer.save(created_by=self.request.user if (self.request.user and self.request.user.is_authenticated) else None)

    def perform_update(self, serializer):
        from .fleet_lifecycle_services import assert_vehicle_trip_allowed
        vehicle = serializer.validated_data.get('vehicle') or serializer.instance.vehicle
        date_shamsi = serializer.validated_data.get('date_shamsi') or serializer.instance.date_shamsi
        trip_count = serializer.validated_data.get('trip_count', serializer.instance.trip_count)
        if (trip_count is None or trip_count > 0) and vehicle and date_shamsi:
            assert_vehicle_trip_allowed(vehicle, date_shamsi)
        serializer.save()

    @action(detail=False, methods=['get'], url_path='matrix')
    def get_matrix(self, request):
        """
        دریافت لیست سریع خودروها برای ثبت سرویس روزانه
        (پشتیبانی از فیلتر انبار و همچنین حالت سراسری کلیه خودروها)
        """
        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        date_shamsi = request.query_params.get('date_shamsi')
        if not date_shamsi:
            return Response({'error': 'پارامتر date_shamsi الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        status_filter = request.query_params.get('status', 'active').lower()
        active_filter_q = Q()
        if status_filter == 'active':
            active_filter_q = Q(is_active=True, approval_status='approved')
        elif status_filter == 'inactive':
            active_filter_q = Q(is_active=False)

        vehicles_base = self._scoped_vehicles().filter(active_filter_q)
        if warehouse_id:
            vehicles = vehicles_base.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True)
            ).order_by('driver_name')
            trip_qs = self._scoped_trips(VehicleTripLog.objects.filter(
                warehouse_id=warehouse_id,
                date_shamsi=date_shamsi,
                is_deleted=False
            ))
        else:
            vehicles = vehicles_base.order_by('driver_name')
            trip_qs = self._scoped_trips(VehicleTripLog.objects.filter(
                date_shamsi=date_shamsi,
                is_deleted=False
            ))

        raw_sec = request.query_params.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        if section_id:
            vehicles = vehicles.filter(section_id=section_id)
            trip_qs = trip_qs.filter(Q(section_id=section_id) | Q(vehicle__section_id=section_id))

        existing_trips = {
            t.vehicle_id: t
            for t in trip_qs
        }

        matrix_rows = []
        for v in vehicles:
            t = existing_trips.get(v.id)
            if t:
                matrix_rows.append({
                    'vehicle_id': v.id,
                    'driver_name': v.driver_name,
                    'plate_number': v.plate_number,
                    'vehicle_type_display': v.get_vehicle_type_display(),
                    'default_rate': float(v.default_service_rate),
                    'trip_id': t.id,
                    'trip_count': t.trip_count,
                    'unit_rate': float(t.unit_rate),
                    'total_amount': float(t.total_amount),
                    'dispatch_reference': t.dispatch_reference or '',
                    'origin_destination': t.origin_destination or '',
                    'notes': t.notes or '',
                    'is_existing': True,
                    'is_active': v.is_active
                })
            else:
                matrix_rows.append({
                    'vehicle_id': v.id,
                    'driver_name': v.driver_name,
                    'plate_number': v.plate_number,
                    'vehicle_type_display': v.get_vehicle_type_display(),
                    'default_rate': float(v.default_service_rate),
                    'trip_id': None,
                    'trip_count': 0,
                    'unit_rate': float(v.default_service_rate),
                    'total_amount': 0.0,
                    'dispatch_reference': '',
                    'origin_destination': '',
                    'notes': '',
                    'is_existing': False,
                    'is_active': v.is_active
                })

        return Response({
            'warehouse_id': warehouse_id,
            'date_shamsi': date_shamsi,
            'rows': matrix_rows
        })

    @action(detail=False, methods=['post'], url_path='bulk-save')
    def bulk_save(self, request):
        """
        ثبت سریع سرویس‌های ناوگان در پایان روز
        (پشتیبانی از ذخیره سراسری ناوگان و ارسال برودکست وب‌سوکت)
        """
        serializer = BulkVehicleTripMatrixSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        
        raw_wh = serializer.validated_data.get('warehouse_id') or request.query_params.get('warehouse_id') or request.data.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_sec = serializer.validated_data.get('section_id') or request.query_params.get('section_id') or request.data.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_proj = serializer.validated_data.get('project_id') or request.query_params.get('project_id') or request.data.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '']) else None

        target_err = self._validate_write_target(request, section_id, project_id, warehouse_id=warehouse_id)
        if target_err is not None:
            return target_err

        date_shamsi = serializer.validated_data['date_shamsi']
        items = serializer.validated_data['items']

        year_month = date_shamsi[:7]
        period = None
        if warehouse_id:
            period, _ = MonthlyWorkPeriod.objects.get_or_create(
                warehouse_id=warehouse_id,
                year_month=year_month,
                defaults={'status': 'OPEN'}
            )
            if period.status == 'LOCKED':
                return Response({'error': f'دوره {year_month} قفل شده است.'}, status=status.HTTP_400_BAD_REQUEST)

        saved_count = 0
        with transaction.atomic():
            for item in items:
                vehicle_id = item['vehicle_id']
                trip_count = item.get('trip_count', 0)
                
                v_obj = self._scoped_vehicles().filter(id=vehicle_id).first()
                if not v_obj:
                    continue

                # تسک ۲.۱ و ب-۲۰: گارد جامع ثبت تردد (وضعیت تصویب، پایان همکاری، انقضای بیمه/معاینه/گواهینامه/قرارداد)
                if trip_count > 0:
                    from .fleet_lifecycle_services import validate_vehicle_validity
                    is_valid, reason = validate_vehicle_validity(v_obj, date_shamsi)
                    if not is_valid:
                        return Response({'error': reason}, status=status.HTTP_400_BAD_REQUEST)

                target_wh = warehouse_id if warehouse_id else v_obj.assigned_warehouse_id
                target_period = period
                if not target_period and target_wh:
                    target_period, _ = MonthlyWorkPeriod.objects.get_or_create(
                        warehouse_id=target_wh,
                        year_month=year_month,
                        defaults={'status': 'OPEN'}
                    )
                if target_period and target_period.status == 'LOCKED':
                    return Response({'error': f'دوره {year_month} برای انبار خودرو ({v_obj.driver_name}) قفل شده است.'}, status=status.HTTP_400_BAD_REQUEST)

                trip_obj = VehicleTripLog.objects.filter(
                    vehicle_id=vehicle_id,
                    warehouse_id=target_wh,
                    date_shamsi=date_shamsi,
                    is_deleted=False
                ).first()

                unit_rate = item.get('unit_rate', v_obj.default_service_rate)
                total_amount = unit_rate * trip_count

                if trip_count > 0:
                    effective_section_id = section_id or v_obj.section_id
                    effective_project_id = project_id or (v_obj.section.project_id if v_obj.section else None)
                    if trip_obj:
                        trip_obj.trip_count = trip_count
                        trip_obj.unit_rate = unit_rate
                        trip_obj.total_amount = total_amount
                        trip_obj.dispatch_reference = item.get('dispatch_reference', '')
                        trip_obj.origin_destination = item.get('origin_destination', '')
                        trip_obj.notes = item.get('notes', '')
                        trip_obj.period = target_period
                        if effective_section_id:
                            trip_obj.section_id = effective_section_id
                        if effective_project_id:
                            trip_obj.project_id = effective_project_id
                        trip_obj.save()
                    else:
                        VehicleTripLog.objects.create(
                            vehicle_id=vehicle_id,
                            warehouse_id=target_wh,
                            section_id=effective_section_id,
                            project_id=effective_project_id,
                            date_shamsi=date_shamsi,
                            trip_count=trip_count,
                            unit_rate=unit_rate,
                            total_amount=total_amount,
                            dispatch_reference=item.get('dispatch_reference', ''),
                            origin_destination=item.get('origin_destination', ''),
                            notes=item.get('notes', ''),
                            period=target_period,
                            created_by=request.user if (request.user and request.user.is_authenticated) else None
                        )
                    saved_count += 1
                else:
                    # اگر صفر بود و قبلاً ثبت شده بود، حذف منطقی
                    if trip_obj:
                        trip_obj.is_deleted = True
                        trip_obj.save()

        # ارسال برودکست وب‌سوکت برای ناوگان
        try:
            from channels.layers import get_channel_layer
            from asgiref.sync import async_to_sync
            channel_layer = get_channel_layer()
            client_tab_id = request.data.get('client_tab_id') or request.headers.get('X-Client-Tab-ID')
            if channel_layer:
                from common.tenant_scope import get_current_tenant_company_id
                cid = get_current_tenant_company_id(request=request, user=request.user)
                target_group = f'company_{cid}_notifications' if cid else 'global_notifications'
                async_to_sync(channel_layer.group_send)(
                    target_group,
                    {
                        'type': 'send_notification',
                        'type_str': 'fleet_trips_updated',
                        'message': f'تردد ناوگان برای تاریخ {date_shamsi} ثبت/به‌روزرسانی شد.',
                        'company_id': cid,
                        'warehouse_id': warehouse_id,
                        'section_id': section_id,
                        'date_shamsi': date_shamsi,
                        'year_month': year_month,
                        'sender_id': request.user.id if request.user and request.user.is_authenticated else None,
                        'client_tab_id': client_tab_id
                    }
                )
        except Exception as e:
            import logging
            logging.getLogger(__name__).warning(f"[WebSocket] Error broadcasting fleet_trips_updated: {e}")

        return Response({
            'message': 'سرویس‌های خودروها با موفقیت ذخیره شدند.',
            'saved_count': saved_count
        })

    @action(detail=False, methods=['get'], url_path='monthly-summary')
    def get_monthly_summary(self, request):
        """
        گزارش کاردکس و عملکرد ماهانه خودروها
        """
        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '', 'UNDEFINED', 'NULL']) else None
        year_month = request.query_params.get('year_month')
        if not year_month:
            latest_period = MonthlyWorkPeriod.objects.order_by('-year_month').first()
            year_month = latest_period.year_month if latest_period else '1405/06'

        vehicles_base = self._scoped_vehicles().filter(is_active=True)
        if warehouse_id:
            vehicles = vehicles_base.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True)
            ).order_by('driver_name')
            trips = self._scoped_trips(VehicleTripLog.objects.filter(
                warehouse_id=warehouse_id,
                date_shamsi__startswith=year_month,
                is_deleted=False
            ))
        else:
            vehicles = vehicles_base.order_by('driver_name')
            trips = self._scoped_trips(VehicleTripLog.objects.filter(
                date_shamsi__startswith=year_month,
                is_deleted=False
            ))

        summary_rows = []
        for v in vehicles:
            v_trips = [t for t in trips if t.vehicle_id == v.id]
            total_trips = sum(t.trip_count for t in v_trips)
            total_payable = sum(float(t.total_amount) for t in v_trips)
            active_days = len(set(t.date_shamsi for t in v_trips if t.trip_count > 0))

            summary_rows.append({
                'vehicle_id': v.id,
                'driver_name': v.driver_name,
                'plate_number': v.plate_number,
                'vehicle_type': v.get_vehicle_type_display(),
                'default_rate': float(v.default_service_rate),
                'sheba_number': v.sheba_number or '',
                'active_days': active_days,
                'total_trips': total_trips,
                'total_payable': total_payable
            })

        return Response({
            'warehouse_id': warehouse_id,
            'year_month': year_month,
            'summary': summary_rows
        })

    @action(detail=False, methods=['get'], url_path='monthly-grid')
    def get_monthly_grid(self, request):
        """
        دریافت ماتریس تردد ۳۱ روزه کلیه خودروها در یک ماه مشخص (نمای شیت ماهانه ناوگان)
        """
        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_year_month = request.query_params.get('year_month')
        if not raw_year_month:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        # نرمال‌سازی سال و ماه
        ym_cleaned = normalize_digits(str(raw_year_month)).strip().replace('-', '/')
        parts = ym_cleaned.split('/')
        if len(parts) >= 2:
            y, m = int(parts[0]), int(parts[1])
            year_month = f"{y:04d}/{m:02d}"
        else:
            digits_only = ''.join(c for c in ym_cleaned if c.isdigit())
            if len(digits_only) >= 6:
                y, m = int(digits_only[:4]), int(digits_only[4:6])
                year_month = f"{y:04d}/{m:02d}"
            else:
                return Response({'error': 'فرمت year_month نامعتبر است (مثال: 1405/04).'}, status=status.HTTP_400_BAD_REQUEST)

        if 1 <= m <= 6:
            days_in_month = 31
        elif 7 <= m <= 11:
            days_in_month = 30
        elif m == 12:
            days_in_month = 30 if jdatetime.date(y, 1, 1).isleap() else 29
        else:
            return Response({'error': 'شماره ماه باید بین ۱ تا ۱۲ باشد.'}, status=status.HTTP_400_BAD_REQUEST)

        month_names = ["", "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
        month_name = month_names[m] if 1 <= m <= 12 else ""

        is_locked = False
        period_status = 'OPEN'
        if warehouse_id:
            period_obj = MonthlyWorkPeriod.objects.filter(warehouse_id=warehouse_id, year_month=year_month).first()
            if period_obj:
                period_status = period_obj.status
                is_locked = period_obj.status in ['LOCKED', 'SUBMITTED', 'FINALIZED']

        today = jdatetime.date.today()
        days_meta = []
        weekday_names = ["ش", "ی", "د", "س", "چ", "پ", "ج"]
        full_weekday_names = ["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه", "پنج‌شنبه", "جمعه"]

        for d in range(1, days_in_month + 1):
            target_date = jdatetime.date(y, m, d)
            weekday = target_date.weekday()
            is_friday = (weekday == 6)
            date_str = f"{year_month}/{d:02d}"
            days_meta.append({
                'day': d,
                'date_shamsi': date_str,
                'weekday_short': weekday_names[weekday],
                'weekday_name': full_weekday_names[weekday],
                'is_friday': is_friday,
                'is_editable': not is_locked,
                'is_today': (target_date == today)
            })

        status_filter = request.query_params.get('status', 'active').lower()
        active_filter_q = Q()
        if status_filter == 'active':
            active_filter_q = Q(is_active=True, approval_status='approved')
        elif status_filter == 'inactive':
            active_filter_q = Q(is_active=False)

        vehicles_base = self._scoped_vehicles().filter(active_filter_q)
        if warehouse_id:
            vehicles = vehicles_base.filter(
                Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True)
            ).order_by('driver_name')
        else:
            vehicles = vehicles_base.order_by('driver_name')

        raw_sec = request.query_params.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '', 'UNDEFINED', 'NULL']) else None
        if section_id:
            vehicles = vehicles.filter(section_id=section_id)

        raw_proj = request.query_params.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '', 'UNDEFINED', 'NULL']) else None
        if project_id:
            vehicles = vehicles.filter(section__project_id=project_id)

        trips = self._scoped_trips(VehicleTripLog.objects.filter(
            date_shamsi__startswith=year_month,
            is_deleted=False
        ))
        if warehouse_id:
            trips = trips.filter(warehouse_id=warehouse_id)
        if section_id:
            trips = trips.filter(Q(section_id=section_id) | Q(vehicle__section_id=section_id))
        if project_id:
            trips = trips.filter(Q(section__project_id=project_id) | Q(vehicle__section__project_id=project_id))

        trip_map = {}
        for t in trips:
            try:
                d_num = int(t.date_shamsi.split('/')[2])
                trip_map[(t.vehicle_id, d_num)] = t
            except Exception:
                continue

        grid_rows = []
        for v in vehicles:
            v_days = []
            total_trips = 0
            total_amount = 0.0
            active_days_count = 0

            for dm in days_meta:
                d = dm['day']
                date_str = dm['date_shamsi']
                t = trip_map.get((v.id, d))
                if t:
                    c = t.trip_count
                    rate = float(t.unit_rate)
                    amt = float(t.total_amount)
                    total_trips += c
                    if v.ownership_type in ['contract', 'personal']:
                        total_amount += amt
                    if c > 0:
                        active_days_count += 1

                    v_days.append({
                        'day': d,
                        'date_shamsi': date_str,
                        'trip_id': t.id,
                        'trip_count': c,
                        'unit_rate': rate,
                        'total_amount': amt,
                        'dispatch_reference': t.dispatch_reference or '',
                        'origin_destination': t.origin_destination or '',
                        'notes': t.notes or '',
                        'is_existing': True
                    })
                else:
                    def_rate = float(v.default_service_rate)
                    v_days.append({
                        'day': d,
                        'date_shamsi': date_str,
                        'trip_id': None,
                        'trip_count': 0,
                        'unit_rate': def_rate,
                        'total_amount': 0.0,
                        'dispatch_reference': '',
                        'origin_destination': '',
                        'notes': '',
                        'is_existing': False
                    })

            grid_rows.append({
                'vehicle_id': v.id,
                'driver_name': v.driver_name,
                'plate_number': v.plate_number,
                'vehicle_type': v.vehicle_type,
                'vehicle_type_display': v.get_vehicle_type_display(),
                'ownership_type': v.ownership_type,
                'ownership_type_display': v.get_ownership_type_display(),
                'default_rate': float(v.default_service_rate),
                'sheba_number': v.sheba_number or '',
                'total_trips': total_trips,
                'total_amount': total_amount,
                'active_days': active_days_count,
                'is_active': v.is_active,
                'days': v_days
            })

        return Response({
            'warehouse_id': warehouse_id,
            'year_month': year_month,
            'month_name': month_name,
            'days_in_month': days_in_month,
            'days_meta': days_meta,
            'is_locked': is_locked,
            'period_status': period_status,
            'rows': grid_rows
        })

    @action(detail=False, methods=['post'], url_path='update-day-trip')
    def update_day_trip(self, request):
        """
        ویرایش یا ثبت سریع سرویس یک روز خاص از تقویم ماهانه
        """
        serializer = VehicleDayTripUpdateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        vehicle_id = serializer.validated_data['vehicle_id']
        raw_wh = serializer.validated_data.get('warehouse_id') or request.query_params.get('warehouse_id') or request.data.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        date_shamsi = serializer.validated_data['date_shamsi']
        trip_count = serializer.validated_data.get('trip_count', 0)
        unit_rate = serializer.validated_data.get('unit_rate')
        dispatch_ref = serializer.validated_data.get('dispatch_reference', '') or ''
        orig_dest = serializer.validated_data.get('origin_destination', '') or ''
        notes = serializer.validated_data.get('notes', '') or ''
        client_tab_id = serializer.validated_data.get('client_tab_id')

        raw_sec = serializer.validated_data.get('section_id') or request.query_params.get('section_id') or request.data.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_proj = serializer.validated_data.get('project_id') or request.query_params.get('project_id') or request.data.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '']) else None

        target_err = self._validate_write_target(request, section_id, project_id, warehouse_id=warehouse_id)
        if target_err is not None:
            return target_err

        v_obj = self._scoped_vehicles().filter(id=vehicle_id).first()
        if not v_obj:
            return Response({'error': 'خودرو یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        # تسک ۲.۱ و ب-۲۰: گارد جامع ثبت تردد (وضعیت تصویب، پایان همکاری، انقضای بیمه/معاینه/گواهینامه/قرارداد)
        if trip_count > 0:
            from .fleet_lifecycle_services import validate_vehicle_validity
            is_valid, reason = validate_vehicle_validity(v_obj, date_shamsi)
            if not is_valid:
                return Response({'error': reason}, status=status.HTTP_400_BAD_REQUEST)

        if unit_rate is None:
            unit_rate = v_obj.default_service_rate

        target_wh = warehouse_id if warehouse_id else v_obj.assigned_warehouse_id
        year_month = date_shamsi[:7]

        allow_override = bool(request.data.get('is_override', False) or request.headers.get('X-Admin-Override') == 'true')
        is_admin_override = allow_override and (getattr(request.user, 'is_superuser', False) or request.user.has_perm('personnel.can_override_attendance_lock'))

        # بررسی قفل دوره
        if target_wh and not is_admin_override:
            period = MonthlyWorkPeriod.objects.filter(warehouse_id=target_wh, year_month=year_month).first()
            if period and period.status in ['LOCKED', 'SUBMITTED', 'FINALIZED']:
                return Response({'error': f'دوره کارکرد {year_month} قفل شده است و امکان ویرایش وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            trip_obj = VehicleTripLog.objects.filter(
                vehicle_id=vehicle_id,
                warehouse_id=target_wh,
                date_shamsi=date_shamsi,
                is_deleted=False
            ).first()

            effective_section_id = section_id or v_obj.section_id
            effective_project_id = project_id or (v_obj.section.project_id if v_obj.section else None)

            if trip_count > 0:
                total_amount = unit_rate * trip_count
                if trip_obj:
                    old_trips = trip_obj.trip_count
                    old_rate = trip_obj.unit_rate
                    if old_trips != trip_count:
                        VehicleTripAuditLog.objects.create(
                            trip=trip_obj,
                            driver_name=v_obj.driver_name,
                            plate_number=v_obj.plate_number,
                            date_shamsi=date_shamsi,
                            changed_by=request.user if request.user.is_authenticated else None,
                            field_name='trip_count',
                            old_value=str(old_trips),
                            new_value=str(trip_count),
                            reason=notes or 'ویرایش سریع روزانه'
                        )
                    if old_rate != unit_rate:
                        VehicleTripAuditLog.objects.create(
                            trip=trip_obj,
                            driver_name=v_obj.driver_name,
                            plate_number=v_obj.plate_number,
                            date_shamsi=date_shamsi,
                            changed_by=request.user if request.user.is_authenticated else None,
                            field_name='unit_rate',
                            old_value=str(old_rate),
                            new_value=str(unit_rate),
                            reason=notes or 'تغییر نرخ سرویس'
                        )
                    trip_obj.trip_count = trip_count
                    trip_obj.unit_rate = unit_rate
                    trip_obj.total_amount = total_amount
                    trip_obj.dispatch_reference = dispatch_ref
                    trip_obj.origin_destination = orig_dest
                    trip_obj.notes = notes
                    trip_obj.is_deleted = False
                    if effective_section_id:
                        trip_obj.section_id = effective_section_id
                    if effective_project_id:
                        trip_obj.project_id = effective_project_id
                    trip_obj.save()
                else:
                    trip_obj = VehicleTripLog.objects.create(
                        vehicle_id=vehicle_id,
                        warehouse_id=target_wh,
                        section_id=effective_section_id,
                        project_id=effective_project_id,
                        date_shamsi=date_shamsi,
                        trip_count=trip_count,
                        unit_rate=unit_rate,
                        total_amount=total_amount,
                        dispatch_reference=dispatch_ref,
                        origin_destination=orig_dest,
                        notes=notes,
                        created_by=request.user if request.user.is_authenticated else None
                    )
                    VehicleTripAuditLog.objects.create(
                        trip=trip_obj,
                        driver_name=v_obj.driver_name,
                        plate_number=v_obj.plate_number,
                        date_shamsi=date_shamsi,
                        changed_by=request.user if request.user.is_authenticated else None,
                        field_name='create',
                        old_value='0',
                        new_value=str(trip_count),
                        reason=notes or 'ثبت اولیه تردد'
                    )
            else:
                if trip_obj:
                    VehicleTripAuditLog.objects.create(
                        trip=trip_obj,
                        driver_name=v_obj.driver_name,
                        plate_number=v_obj.plate_number,
                        date_shamsi=date_shamsi,
                        changed_by=request.user if request.user.is_authenticated else None,
                        field_name='delete',
                        old_value=str(trip_obj.trip_count),
                        new_value='0',
                        reason='حذف تردد'
                    )
                    trip_obj.is_deleted = True
                    trip_obj.save()

        # ارسال برودکست وب‌سوکت
        try:
            from channels.layers import get_channel_layer
            from asgiref.sync import async_to_sync
            channel_layer = get_channel_layer()
            if channel_layer:
                cid = getattr(v_obj, 'company_id', None)
                if not cid and hasattr(v_obj, 'section') and v_obj.section and v_obj.section.project:
                    cid = v_obj.section.project.company_id
                target_group = f'company_{cid}_notifications' if cid else 'global_notifications'
                async_to_sync(channel_layer.group_send)(
                    target_group,
                    {
                        'type': 'send_notification',
                        'type_str': 'fleet_trips_updated',
                        'message': f'تردد خودرو {v_obj.driver_name} در تاریخ {date_shamsi} به‌روزرسانی شد.',
                        'company_id': cid,
                        'warehouse_id': warehouse_id,
                        'section_id': section_id,
                        'date_shamsi': date_shamsi,
                        'year_month': year_month,
                        'sender_id': request.user.id if request.user and request.user.is_authenticated else None,
                        'client_tab_id': client_tab_id
                    }
                )
        except Exception as e:
            import logging
            logging.getLogger(__name__).warning(f"[WebSocket] Error broadcasting fleet_trips_updated: {e}")

        return Response({
            'message': f'تردد خودرو {v_obj.driver_name} برای تاریخ {date_shamsi} با موفقیت ثبت شد.',
            'trip_id': trip_obj.id if (trip_count > 0 and trip_obj) else None
        })

    @action(detail=False, methods=['get'], url_path='audit-logs')
    def audit_logs(self, request):
        """
        لیست تاریخچه و ممیزی تغییرات تردد ناوگان
        """
        vehicle_id = request.query_params.get('vehicle_id')
        year_month = request.query_params.get('year_month')
        scoped_trips_qs = self._scoped_trips(VehicleTripLog.objects.all())
        qs = VehicleTripAuditLog.objects.filter(trip__in=scoped_trips_qs).select_related('changed_by')
        if vehicle_id:
            qs = qs.filter(trip__vehicle_id=vehicle_id)
        if year_month:
            qs = qs.filter(date_shamsi__startswith=year_month)

        serializer = VehicleTripAuditLogSerializer(qs[:100], many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['post'], url_path='bulk-save-monthly-grid')
    def bulk_save_monthly_grid(self, request):
        """
        ثبت دسته‌جمعی ماتریس ۳۱ روزه ناوگان
        """
        serializer = BulkVehicleMonthlyGridSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        raw_wh = serializer.validated_data.get('warehouse_id') or request.query_params.get('warehouse_id') or request.data.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_sec = serializer.validated_data.get('section_id') or request.query_params.get('section_id') or request.data.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_proj = serializer.validated_data.get('project_id') or request.query_params.get('project_id') or request.data.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '']) else None

        target_err = self._validate_write_target(request, section_id, project_id, warehouse_id=warehouse_id)
        if target_err is not None:
            return target_err

        year_month = serializer.validated_data['year_month']
        items = serializer.validated_data['items']
        client_tab_id = serializer.validated_data.get('client_tab_id')

        if warehouse_id:
            period = MonthlyWorkPeriod.objects.filter(warehouse_id=warehouse_id, year_month=year_month).first()
            if period and period.status in ['LOCKED', 'SUBMITTED', 'FINALIZED']:
                return Response({'error': f'دوره کارکرد {year_month} قفل شده است.'}, status=status.HTTP_400_BAD_REQUEST)

        saved_count = 0
        with transaction.atomic():
            for item in items:
                v_id = item['vehicle_id']
                day = item['day']
                trip_count = item.get('trip_count', 0)
                v_obj = self._scoped_vehicles().filter(id=v_id).first()
                if not v_obj:
                    continue

                target_wh = warehouse_id if warehouse_id else v_obj.assigned_warehouse_id
                date_shamsi = f"{year_month}/{day:02d}"

                # تسک ۲.۱ و ب-۲۰: گارد جامع ثبت تردد (وضعیت تصویب، پایان همکاری، انقضای بیمه/معاینه/گواهینامه/قرارداد)
                if trip_count > 0:
                    from .fleet_lifecycle_services import validate_vehicle_validity
                    is_valid, reason = validate_vehicle_validity(v_obj, date_shamsi)
                    if not is_valid:
                        return Response({'error': reason}, status=status.HTTP_400_BAD_REQUEST)

                unit_rate = item.get('unit_rate') or v_obj.default_service_rate
                total_amount = unit_rate * trip_count

                trip_obj = VehicleTripLog.objects.filter(
                    vehicle_id=v_id,
                    warehouse_id=target_wh,
                    date_shamsi=date_shamsi,
                    is_deleted=False
                ).first()

                effective_section_id = section_id or v_obj.section_id
                effective_project_id = project_id or (v_obj.section.project_id if v_obj.section else None)

                if trip_count > 0:
                    if trip_obj:
                        trip_obj.trip_count = trip_count
                        trip_obj.unit_rate = unit_rate
                        trip_obj.total_amount = total_amount
                        trip_obj.dispatch_reference = item.get('dispatch_reference', '') or ''
                        trip_obj.origin_destination = item.get('origin_destination', '') or ''
                        trip_obj.notes = item.get('notes', '') or ''
                        trip_obj.is_deleted = False
                        if effective_section_id:
                            trip_obj.section_id = effective_section_id
                        if effective_project_id:
                            trip_obj.project_id = effective_project_id
                        trip_obj.save()
                    else:
                        VehicleTripLog.objects.create(
                            vehicle_id=v_id,
                            warehouse_id=target_wh,
                            section_id=effective_section_id,
                            project_id=effective_project_id,
                            date_shamsi=date_shamsi,
                            trip_count=trip_count,
                            unit_rate=unit_rate,
                            total_amount=total_amount,
                            dispatch_reference=item.get('dispatch_reference', '') or '',
                            origin_destination=item.get('origin_destination', '') or '',
                            notes=item.get('notes', '') or '',
                            created_by=request.user
                        )
                    saved_count += 1
                else:
                    if trip_obj:
                        trip_obj.is_deleted = True
                        trip_obj.save()

        # وب‌سوکت
        try:
            from channels.layers import get_channel_layer
            from asgiref.sync import async_to_sync
            channel_layer = get_channel_layer()
            if channel_layer:
                from common.tenant_scope import get_current_tenant_company_id
                cid = get_current_tenant_company_id(request=request, user=request.user)
                target_group = f'company_{cid}_notifications' if cid else 'global_notifications'
                async_to_sync(channel_layer.group_send)(
                    target_group,
                    {
                        'type': 'send_notification',
                        'type_str': 'fleet_trips_updated',
                        'message': f'ماتریس ۳۱ روزه ناوگان برای دوره {year_month} ذخیره شد.',
                        'company_id': cid,
                        'warehouse_id': warehouse_id,
                        'section_id': section_id,
                        'year_month': year_month,
                        'sender_id': request.user.id if request.user and request.user.is_authenticated else None,
                        'client_tab_id': client_tab_id
                    }
                )
        except Exception as e:
            import logging
            logging.getLogger(__name__).warning(f"[WebSocket] Error broadcasting fleet_trips_updated: {e}")

        return Response({
            'message': f'ماتریس ۳۱ روزه ناوگان برای دوره {year_month} با موفقیت ذخیره شد.',
            'saved_count': saved_count
        })

    @action(detail=False, methods=['get'], url_path='export-monthly-excel')
    def export_monthly_excel(self, request):
        """
        دانلود فایل اکسل ۳۱ روزه استاندارد کارکرد و تردد ناوگان
        """
        raw_wh = request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        year_month = request.query_params.get('year_month')
        if not year_month:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        raw_sec = request.query_params.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None

        excel_buffer = export_fleet_monthly_excel(
            warehouse_id=warehouse_id,
            year_month=year_month,
            section_id=section_id,
            vehicles_qs=self._scoped_vehicles(),
            trips_qs=self._scoped_trips(VehicleTripLog.objects.all())
        )
        filename = f"Fleet_Timesheet_{year_month.replace('/', '_')}.xlsx"
        response = HttpResponse(
            excel_buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    @action(detail=False, methods=['post'], url_path='import-monthly-excel')
    def import_monthly_excel(self, request):
        """
        درون‌ریزی فایل اکسل ۳۱ روزه کارکرد ناوگان
        """
        raw_wh = request.data.get('warehouse_id') or request.query_params.get('warehouse_id')
        warehouse_id = int(raw_wh) if (raw_wh and str(raw_wh).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_sec = request.data.get('section_id') or request.query_params.get('section_id')
        section_id = int(raw_sec) if (raw_sec and str(raw_sec).upper() not in ['ALL', '0', 'NONE', '']) else None
        raw_proj = request.data.get('project_id') or request.query_params.get('project_id')
        project_id = int(raw_proj) if (raw_proj and str(raw_proj).upper() not in ['ALL', '0', 'NONE', '']) else None

        target_err = self._validate_write_target(request, section_id, project_id, warehouse_id=warehouse_id)
        if target_err is not None:
            return target_err

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        year_month = request.data.get('year_month')

        try:
            result = import_fleet_monthly_excel(
                file_obj=file_obj,
                warehouse_id=warehouse_id,
                year_month=year_month,
                user=request.user,
                vehicles_qs=self._scoped_vehicles()
            )
            return Response(result, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)


class MonthlyWorkPeriodViewSet(viewsets.ModelViewSet):
    queryset = MonthlyWorkPeriod.objects.all().select_related('locked_by')
    serializer_class = MonthlyWorkPeriodSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        from common.tenant_scope import scope_tenant_queryset
        qs = scope_tenant_queryset(qs, self.request.user, company_field='company_id', request=self.request)
        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(warehouse_id=warehouse_id)
        year_month = self.request.query_params.get('year_month')
        if year_month:
            qs = qs.filter(year_month=year_month)
        return qs

    @action(detail=True, methods=['post'], url_path='lock')
    def lock_period(self, request, pk=None):
        period = self.get_object()
        period.status = 'LOCKED'
        period.locked_at = timezone.now()
        period.locked_by = request.user
        period.save()
        return Response({'message': f'دوره {period.year_month} با موفقیت قفل شد.'})

    @action(detail=True, methods=['post'], url_path='unlock')
    def unlock_period(self, request, pk=None):
        if not request.user.is_superuser:
            return Response({'error': 'فقط مدیر ارشد سیستم اجازه بازگشایی دوره را دارد.'}, status=status.HTTP_403_FORBIDDEN)
        period = self.get_object()
        period.status = 'OPEN'
        period.locked_at = None
        period.locked_by = None
        period.save()
        return Response({'message': f'دوره {period.year_month} با موفقیت بازگشایی شد.'})

    @action(detail=True, methods=['post'], url_path='submit')
    def submit_period(self, request, pk=None):
        period = self.get_object()
        if period.status == 'LOCKED':
            return Response({'error': 'دوره از قبل قفل شده است.'}, status=status.HTTP_400_BAD_REQUEST)
        period.status = 'SUBMITTED'
        period.submitted_at = timezone.now()
        period.submitted_by = request.user
        period.submission_notes = request.data.get('notes', '')
        period.rejection_reason = None
        period.save()
        return Response({
            'message': f'دوره {period.year_month} جهت بررسی مالی ارسال گردید.',
            'status': period.status,
            'status_display': period.get_status_display()
        })

    @action(detail=True, methods=['post'], url_path='approve')
    def approve_period(self, request, pk=None):
        is_financial = request.user.is_superuser or request.user.groups.filter(name__in=['Accountant', 'Finance', 'مدیر مالی']).exists() or request.user.has_perm('personnel.can_override_attendance_lock')
        if not is_financial:
            return Response({'error': 'فقط حسابدار یا مدیر مالی مجاز به تایید نهایی دوره است.'}, status=status.HTTP_403_FORBIDDEN)
        period = self.get_object()
        period.status = 'LOCKED'
        period.locked_at = timezone.now()
        period.locked_by = request.user
        period.save()
        return Response({
            'message': f'دوره {period.year_month} تایید نهایی و قفل شد.',
            'status': period.status,
            'status_display': period.get_status_display()
        })

    @action(detail=True, methods=['post'], url_path='reject')
    def reject_period(self, request, pk=None):
        is_financial = request.user.is_superuser or request.user.groups.filter(name__in=['Accountant', 'Finance', 'مدیر مالی']).exists() or request.user.has_perm('personnel.can_override_attendance_lock')
        if not is_financial:
            return Response({'error': 'فقط حسابدار یا مدیر مالی مجاز به رد کارکرد است.'}, status=status.HTTP_403_FORBIDDEN)
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'درج علت رد دوره برای اطلاع سرپرست الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        period = self.get_object()
        period.status = 'REJECTED'
        period.rejection_reason = reason
        period.locked_at = None
        period.locked_by = None
        period.save()
        return Response({
            'message': f'دوره {period.year_month} رد شد و جهت بازبینی به سرپرست بازگشت.',
            'status': period.status,
            'status_display': period.get_status_display()
        })

    @action(detail=False, methods=['post'], url_path='workflow-action')
    def workflow_action(self, request):
        """
        اقدام گردش کار کارکرد بر مبنای warehouse_id و year_month (پشتیبانی از فراخوانی سریع بدون دانستن pk)
        """
        warehouse_id = request.data.get('warehouse_id')
        year_month = request.data.get('year_month')
        action_type = request.data.get('action')
        notes = request.data.get('notes', '')
        from common.tenant_scope import get_current_tenant_company_id
        company_id = request.data.get('company_id') or get_current_tenant_company_id(request=request, user=request.user)

        if not year_month or not action_type:
            return Response({'error': 'پارامترهای year_month و action الزامی هستند.'}, status=status.HTTP_400_BAD_REQUEST)
        if not warehouse_id and not company_id:
            return Response({'error': 'شناسه انبار یا شرکت فعال الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        wh_val = int(warehouse_id) if (warehouse_id and str(warehouse_id).isdigit()) else None
        comp_val = int(company_id) if (company_id and str(company_id).isdigit()) else None

        lookup = {'year_month': str(year_month).strip()}
        if wh_val:
            lookup['warehouse_id'] = wh_val
        if comp_val:
            lookup['company_id'] = comp_val

        period, _ = MonthlyWorkPeriod.objects.get_or_create(
            **lookup,
            defaults={'status': 'OPEN', 'company_id': comp_val, 'warehouse_id': wh_val}
        )

        is_financial = request.user.is_superuser or request.user.groups.filter(name__in=['Accountant', 'Finance', 'مدیر مالی']).exists() or request.user.has_perm('personnel.can_override_attendance_lock')

        if action_type == 'submit':
            if period.status == 'LOCKED':
                return Response({'error': 'دوره از قبل قفل شده است.'}, status=status.HTTP_400_BAD_REQUEST)
            period.status = 'SUBMITTED'
            period.submitted_at = timezone.now()
            period.submitted_by = request.user
            period.submission_notes = notes
            period.rejection_reason = None
            period.save()
            return Response({
                'message': f'کارکرد دوره {year_month} جهت بررسی مالی ارسال گردید.',
                'period': MonthlyWorkPeriodSerializer(period).data
            })

        elif action_type == 'approve':
            if not is_financial:
                return Response({'error': 'فقط حسابدار یا مدیر مالی مجاز به تایید نهایی و قفل دوره است.'}, status=status.HTTP_403_FORBIDDEN)
            period.status = 'LOCKED'
            period.locked_at = timezone.now()
            period.locked_by = request.user
            period.save()
            return Response({
                'message': f'کارکرد دوره {year_month} با موفقیت تایید نهایی و قفل شد.',
                'period': MonthlyWorkPeriodSerializer(period).data
            })

        elif action_type == 'reject':
            if not is_financial:
                return Response({'error': 'فقط حسابدار یا مدیر مالی مجاز به رد کارکرد است.'}, status=status.HTTP_403_FORBIDDEN)
            if not notes.strip():
                return Response({'error': 'درج علت رد دوره الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            period.status = 'REJECTED'
            period.rejection_reason = notes.strip()
            period.locked_at = None
            period.locked_by = None
            period.save()
            return Response({
                'message': f'دوره {year_month} رد شد و جهت بازبینی به سرپرست بازگردانده شد.',
                'period': MonthlyWorkPeriodSerializer(period).data
            })

        elif action_type == 'unlock':
            is_unlock_authorized = (
                request.user.is_superuser or
                request.user.groups.filter(name__in=['Accountant', 'Finance', 'مدیر مالی']).exists() or
                request.user.has_perm('personnel.can_override_attendance_lock')
            )
            if not is_unlock_authorized:
                return Response({'error': 'فقط مدیر ارشد سیستم یا مدیر مالی مجاز به بازگشایی دوره است.'}, status=status.HTTP_403_FORBIDDEN)
            period.status = 'OPEN'
            period.locked_at = None
            period.locked_by = None
            period.save()
            return Response({
                'message': f'دوره {year_month} با موفقیت بازگشایی شد.',
                'period': MonthlyWorkPeriodSerializer(period).data
            })

        return Response({'error': 'نوع اقدام ناشناخته است.'}, status=status.HTTP_400_BAD_REQUEST)


# ══════════════════════════════════════════════════════════════════════════════
# ── ۷. ویوست تنظیمات جامع سالانه و ۲۰ گروه شغلی ───────────────────────────────
# ══════════════════════════════════════════════════════════════════════════════

def safe_int(val, default=None):
    if val is None:
        return default
    val_str = str(val).strip().lower()
    if val_str in ('', 'null', 'undefined', 'nan', 'none'):
        return default
    try:
        return int(val_str)
    except (ValueError, TypeError):
        return default


def _copy_payroll_submodels(source, target):
    """کپی ایمن و همه‌جانبه مدل‌های تابعه (کارگاه، مالیات، بانک و ۲۰ گروه شغلی)"""
    if not source:
        return
    if hasattr(source, 'workshop_insurance'):
        wi = source.workshop_insurance
        WorkshopInsuranceSettings.objects.update_or_create(
            yearly_settings=target,
            defaults={
                'workshop_code': wi.workshop_code,
                'workshop_name': wi.workshop_name,
                'employer_name': wi.employer_name,
                'workshop_address': wi.workshop_address,
                'list_type': wi.list_type,
                'list_number': wi.list_number,
                'default_dsk_rate': wi.default_dsk_rate,
                'default_mon_pym': wi.default_mon_pym,
            }
        )
    if hasattr(source, 'tax_settings'):
        ts = source.tax_settings
        TaxRuleSettings.objects.update_or_create(
            yearly_settings=target,
            defaults={
                'payment_type': ts.payment_type,
                'service_location': ts.service_location,
                'exceptions': ts.exceptions,
                'currency_type': ts.currency_type,
                'currency_exchange_rate': ts.currency_exchange_rate,
                'housing_benefit_type': ts.housing_benefit_type,
                'vehicle_benefit_type': ts.vehicle_benefit_type,
                'wh_file_prefix': ts.wh_file_prefix,
                'wp_file_prefix': ts.wp_file_prefix,
            }
        )
    if hasattr(source, 'bank_export_settings'):
        bs = source.bank_export_settings
        BankExportSettings.objects.update_or_create(
            yearly_settings=target,
            defaults={
                'bank_name': bs.bank_name,
                'source_account_number': bs.source_account_number,
                'source_sheba_number': getattr(bs, 'source_sheba_number', None),
                'default_deposit_id': bs.default_deposit_id,
                'deposit_description_template': bs.deposit_description_template,
            }
        )
    for tier in source.job_grades.all():
        JobGradeTier.objects.update_or_create(
            yearly_settings=target,
            grade_number=tier.grade_number,
            defaults={
                'daily_base_wage': tier.daily_base_wage,
                'daily_seniority_bonus': tier.daily_seniority_bonus,
            }
        )


class PayrollYearlySettingsViewSet(viewsets.ModelViewSet):
    queryset = PayrollYearlySettings.objects.all().prefetch_related(
        'job_grades', 'workshop_insurance', 'tax_settings', 'bank_export_settings', 'project'
    )
    serializer_class = PayrollYearlySettingsSerializer
    permission_classes = [IsPayrollSettingsManagerOrReadOnly]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        cid = validate_user_company_access(self.request.user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(Q(company_id=cid) | Q(project__company_id=cid))
        p_id = safe_int(self.request.query_params.get('project_id'))
        if p_id is not None:
            return qs.filter(project_id=p_id)
        is_global = self.request.query_params.get('is_global')
        if is_global == 'true':
            return qs.filter(project__isnull=True)
        return qs

    @action(detail=False, methods=['get', 'post'], url_path='active-or-year')
    def get_active_or_year(self, request):
        if request.method == 'POST':
            pk = request.data.get('id')
            if pk:
                settings_obj = PayrollYearlySettings.objects.filter(id=pk).first()
                if settings_obj:
                    return self.update_all_tabs(request, pk=pk)

        year = request.query_params.get('year', '1405')
        p_id = safe_int(request.query_params.get('project_id'))
        v_id = safe_int(request.query_params.get('version_id'))
        effective_from = request.query_params.get('effective_from')
        if effective_from in ('null', 'undefined', ''):
            effective_from = None

        settings_obj = None
        if v_id:
            candidate = PayrollYearlySettings.objects.filter(id=v_id).first()
            if candidate:
                if p_id is not None:
                    if candidate.project_id == p_id:
                        settings_obj = candidate
                    else:
                        # اگر کاندید برای سراسری است و پروژه تنظیمی ندارد، همان نسخه کاندید را نمایش بده
                        settings_obj = candidate
                else:
                    settings_obj = candidate

        if not settings_obj:
            if effective_from:
                settings_obj = get_effective_payroll_settings(effective_from, p_id)
            else:
                settings_obj = get_effective_payroll_settings(f"{year}/01", p_id)

        if not settings_obj:
            settings_obj = PayrollYearlySettings.objects.filter(fiscal_year=year).first() or PayrollYearlySettings.objects.first()
        if not settings_obj:
            return Response({'error': 'تنظیماتی یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
        return Response(PayrollYearlySettingsSerializer(settings_obj).data)

    @action(detail=False, methods=['get'], url_path='versions')
    def get_versions(self, request):
        """
        دریافت فهرست تمام نسخه‌های احکام در یک سال مالی به تفکیک پروژه
        """
        year = request.query_params.get('year', '1405')
        p_id = safe_int(request.query_params.get('project_id'))
        qs = PayrollYearlySettings.objects.filter(fiscal_year=year)
        if p_id is not None:
            project_qs = qs.filter(project_id=p_id).order_by('effective_from')
            if project_qs.exists():
                return Response(PayrollYearlySettingsSerializer(project_qs, many=True).data)
            # اگر برای پروژه تنظیمات اختصاصی ثبت نشده، نسخه‌های سراسری سازمان بازگردانده می‌شود
            global_qs = qs.filter(project__isnull=True).order_by('effective_from')
            return Response(PayrollYearlySettingsSerializer(global_qs, many=True).data)
        else:
            qs = qs.filter(project__isnull=True).order_by('effective_from')
        return Response(PayrollYearlySettingsSerializer(qs, many=True).data)

    @action(detail=False, methods=['get'], url_path='available-years')
    def get_available_years(self, request):
        """
        دریافت فهرست سال‌های مالی تعریف‌شده در سیستم
        """
        years = list(PayrollYearlySettings.objects.values_list('fiscal_year', flat=True).distinct())
        cleaned_years = sorted(list(set([str(y).strip() for y in years if y and str(y).strip()])), reverse=True)
        if not cleaned_years:
            cleaned_years = ['1405']
        return Response({'years': cleaned_years})

    @action(detail=False, methods=['post'], url_path='create-year')
    def create_fiscal_year(self, request):
        """
        ایجاد سال مالی جدید بر اساس کپی اطلاعات و احکام یک سال مالی مبدأ
        """
        new_year = request.data.get('year')
        source_year = request.data.get('source_year')
        project_id = request.data.get('project_id')
        p_id = safe_int(project_id)

        if not new_year:
            return Response({'error': 'تعیین سال مالی جدید الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        new_year = str(new_year).strip()
        if len(new_year) != 4 or not new_year.isdigit():
            return Response({'error': 'فرمت سال مالی باید یک عدد ۴ رقمی شمسی (مثلاً ۱۴۰۶) باشد.'}, status=status.HTTP_400_BAD_REQUEST)

        # بررسی وجود قبلی در دامنه انتخابی
        existing = PayrollYearlySettings.objects.filter(fiscal_year=new_year, project_id=p_id).first()
        if existing:
            return Response({
                'error': f'سال مالی {new_year} قبلاً در این دامنه تعریف شده است.',
                'settings': PayrollYearlySettingsSerializer(existing).data
            }, status=status.HTTP_400_BAD_REQUEST)

        project = FinancialProject.objects.filter(id=p_id).first() if p_id else None

        # پیدا کردن سال مبدأ
        source = None
        if source_year:
            source = get_effective_payroll_settings(f"{source_year}/01", p_id)
        if not source:
            source = get_effective_payroll_settings("1405/01", p_id) or PayrollYearlySettings.objects.first()

        effective_from = f"{new_year}/01"
        version_title = "احکام مصوب فروردین"

        with transaction.atomic():
            # اگر پروژه انتخاب شده باشد، ابتدا مطمئن می‌شویم رکورد پایه سراسری سازمان برای سال جدید وجود دارد
            if project and not PayrollYearlySettings.objects.filter(fiscal_year=new_year, project__isnull=True).exists():
                global_source = get_effective_payroll_settings(f"{source_year or '1405'}/01", None) or source
                global_settings = PayrollYearlySettings.objects.create(
                    fiscal_year=new_year,
                    effective_from=effective_from,
                    version_title=version_title,
                    project=None,
                    title=f"تنظیمات سال مالی {new_year} (سراسری سازمان)",
                    is_active=True,
                    monthly_food_allowance=global_source.monthly_food_allowance if global_source else 22000000,
                    monthly_housing_allowance=global_source.monthly_housing_allowance if global_source else 30000000,
                    monthly_spouse_allowance=global_source.monthly_spouse_allowance if global_source else 5000000,
                    monthly_child_allowance=global_source.monthly_child_allowance if global_source else 16625549,
                    shift_percent=global_source.shift_percent if global_source else 0,
                    transport_help_percent=global_source.transport_help_percent if global_source else 0,
                    transport_fixed_amount=global_source.transport_fixed_amount if global_source else 2388728,
                    specialist_attraction_percent=global_source.specialist_attraction_percent if global_source else 0,
                    bad_weather_percent=global_source.bad_weather_percent if global_source else 0,
                    remote_hardship_percent=global_source.remote_hardship_percent if global_source else 0,
                    south_pars_percent=global_source.south_pars_percent if global_source else 0,
                    travel_cost_per_day=global_source.travel_cost_per_day if global_source else 8736600,
                    worker_insurance_rate=global_source.worker_insurance_rate if global_source else 7.00,
                    employer_insurance_rate=global_source.employer_insurance_rate if global_source else 20.00,
                    unemployment_insurance_rate=global_source.unemployment_insurance_rate if global_source else 3.00,
                    surplus_overtime_percent=global_source.surplus_overtime_percent if global_source else 50.00,
                    standard_daily_hours=getattr(global_source, 'standard_daily_hours', 10.00) or 10.00,
                    bonus_daily_coefficient=getattr(global_source, 'bonus_daily_coefficient', 5.00) or 5.00,
                    seniority_monthly_coefficient=getattr(global_source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                    overtime_rate_multiplier=getattr(global_source, 'overtime_rate_multiplier', 1.40) or 1.40,
                    friday_work_rate_multiplier=getattr(global_source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                    max_insurable_daily_wage=getattr(global_source, 'max_insurable_daily_wage', None),
                    attendance_edit_past_days=global_source.attendance_edit_past_days if global_source else 3,
                    attendance_edit_future_days=global_source.attendance_edit_future_days if global_source else 0,
                    created_by=request.user,
                    updated_by=request.user
                )
                _copy_payroll_submodels(global_source, global_settings)

            new_settings = PayrollYearlySettings.objects.create(
                fiscal_year=new_year,
                effective_from=effective_from,
                version_title=version_title,
                project=project,
                title=f"تنظیمات سال مالی {new_year} ({project.name if project else 'سراسری سازمان'})",
                is_active=True,
                monthly_food_allowance=source.monthly_food_allowance if source else 22000000,
                monthly_housing_allowance=source.monthly_housing_allowance if source else 30000000,
                monthly_spouse_allowance=source.monthly_spouse_allowance if source else 5000000,
                monthly_child_allowance=source.monthly_child_allowance if source else 16625549,
                shift_percent=source.shift_percent if source else 0,
                transport_help_percent=source.transport_help_percent if source else 0,
                transport_fixed_amount=source.transport_fixed_amount if source else 2388728,
                specialist_attraction_percent=source.specialist_attraction_percent if source else 0,
                bad_weather_percent=source.bad_weather_percent if source else 0,
                remote_hardship_percent=source.remote_hardship_percent if source else 0,
                south_pars_percent=source.south_pars_percent if source else 0,
                travel_cost_per_day=source.travel_cost_per_day if source else 8736600,
                worker_insurance_rate=source.worker_insurance_rate if source else 7.00,
                employer_insurance_rate=source.employer_insurance_rate if source else 20.00,
                unemployment_insurance_rate=source.unemployment_insurance_rate if source else 3.00,
                surplus_overtime_percent=source.surplus_overtime_percent if source else 50.00,
                standard_daily_hours=getattr(source, 'standard_daily_hours', 10.00) or 10.00,
                bonus_daily_coefficient=getattr(source, 'bonus_daily_coefficient', 5.00) or 5.00,
                seniority_monthly_coefficient=getattr(source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                overtime_rate_multiplier=getattr(source, 'overtime_rate_multiplier', 1.40) or 1.40,
                friday_work_rate_multiplier=getattr(source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                max_insurable_daily_wage=getattr(source, 'max_insurable_daily_wage', None),
                attendance_edit_past_days=source.attendance_edit_past_days if source else 3,
                attendance_edit_future_days=source.attendance_edit_future_days if source else 0,
                created_by=request.user,
                updated_by=request.user
            )
            _copy_payroll_submodels(source, new_settings)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='CREATE',
                    severity='warning',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(new_settings.id),
                    target_repr=str(new_settings),
                    details={
                        'action': 'create_fiscal_year',
                        'fiscal_year': new_year,
                        'source_year': source_year,
                        'project_id': p_id
                    }
                )
            except Exception:
                pass

        return Response(PayrollYearlySettingsSerializer(new_settings).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='create-version')
    def create_version(self, request):
        """
        ایجاد نسخه جدید احکام در میانه سال با کپی از آخرین نسخه و تعیین ماه شروع
        """
        import re
        project_id = request.data.get('project_id')
        year = request.data.get('year', '1405')
        effective_from = request.data.get('effective_from')
        version_title = request.data.get('version_title') or f"احکام اصلاحیه {effective_from}"

        if not effective_from:
            return Response({'error': 'ماه شروع اعتبار (effective_from) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        effective_from = str(effective_from).strip()
        if not re.match(r'^\d{4}/(0[1-9]|1[0-2])$', effective_from):
            return Response({'error': 'فرمت ماه شروع اعتبار باید YYYY/MM باشد (مثال: 1405/07).'}, status=status.HTTP_400_BAD_REQUEST)

        p_id = safe_int(project_id)
        project = FinancialProject.objects.filter(id=p_id).first() if p_id else None

        existing = PayrollYearlySettings.objects.filter(
            fiscal_year=year,
            project=project,
            effective_from=effective_from
        ).first()
        if existing:
            return Response({
                'error': f'نسخه احکام با تاریخ شروع {effective_from} قبلاً در این دامنه ایجاد شده است.',
                'settings': PayrollYearlySettingsSerializer(existing).data
            }, status=status.HTTP_400_BAD_REQUEST)

        source = get_effective_payroll_settings(effective_from, p_id)
        if not source:
            source = PayrollYearlySettings.objects.first()

        with transaction.atomic():
            # به‌روزرسانی خودکار ماه پایان نسخه قبلی
            year_str, month_str = effective_from.split('/')
            m_int = int(month_str)
            if m_int > 1:
                prior_effective_to = f"{year_str}/{m_int - 1:02d}"
                PayrollYearlySettings.objects.filter(
                    fiscal_year=year,
                    project=project,
                    effective_from__lt=effective_from
                ).filter(
                    Q(effective_to__isnull=True) | Q(effective_to__gte=effective_from)
                ).update(effective_to=prior_effective_to)

            new_settings = PayrollYearlySettings.objects.create(
                fiscal_year=year,
                effective_from=effective_from,
                version_title=version_title,
                project=project,
                title=f"{version_title} ({project.name if project else 'سراسری سازمان'})",
                is_active=True,
                monthly_food_allowance=source.monthly_food_allowance if source else 22000000,
                monthly_housing_allowance=source.monthly_housing_allowance if source else 30000000,
                monthly_spouse_allowance=source.monthly_spouse_allowance if source else 5000000,
                monthly_child_allowance=source.monthly_child_allowance if source else 16625549,
                shift_percent=source.shift_percent if source else 0,
                transport_help_percent=source.transport_help_percent if source else 0,
                transport_fixed_amount=source.transport_fixed_amount if source else 2388728,
                specialist_attraction_percent=source.specialist_attraction_percent if source else 0,
                bad_weather_percent=source.bad_weather_percent if source else 0,
                remote_hardship_percent=source.remote_hardship_percent if source else 0,
                south_pars_percent=source.south_pars_percent if source else 0,
                travel_cost_per_day=source.travel_cost_per_day if source else 8736600,
                worker_insurance_rate=source.worker_insurance_rate if source else 7.00,
                employer_insurance_rate=source.employer_insurance_rate if source else 20.00,
                unemployment_insurance_rate=source.unemployment_insurance_rate if source else 3.00,
                surplus_overtime_percent=source.surplus_overtime_percent if source else 50.00,
                standard_daily_hours=getattr(source, 'standard_daily_hours', 10.00) or 10.00,
                bonus_daily_coefficient=getattr(source, 'bonus_daily_coefficient', 5.00) or 5.00,
                seniority_monthly_coefficient=getattr(source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                overtime_rate_multiplier=getattr(source, 'overtime_rate_multiplier', 1.40) or 1.40,
                friday_work_rate_multiplier=getattr(source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                max_insurable_daily_wage=getattr(source, 'max_insurable_daily_wage', None),
                attendance_edit_past_days=source.attendance_edit_past_days if source else 3,
                attendance_edit_future_days=source.attendance_edit_future_days if source else 0,
                created_by=request.user,
                updated_by=request.user
            )
            _copy_payroll_submodels(source, new_settings)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='CREATE',
                    severity='warning',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(new_settings.id),
                    target_repr=str(new_settings),
                    details={
                        'action': 'create_version',
                        'fiscal_year': year,
                        'effective_from': effective_from,
                        'version_title': version_title,
                        'project_id': p_id
                    }
                )
            except Exception:
                pass

        return Response(PayrollYearlySettingsSerializer(new_settings).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='clone-for-project')
    def clone_for_project(self, request):
        """
        کپی هوشمند و سریع تنظیمات پایه سراسری برای یک پروژه مشخص
        """
        project_id = request.data.get('project_id')
        year = request.data.get('year', '1405')
        effective_from = request.data.get('effective_from', f"{year}/01")

        p_id = safe_int(project_id)
        if not p_id:
            return Response({'error': 'شناسه پروژه (project_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        project = FinancialProject.objects.filter(id=p_id).first()
        if not project:
            return Response({'error': 'پروژه مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        existing = PayrollYearlySettings.objects.filter(fiscal_year=year, project=project, effective_from=effective_from).first()
        if existing:
            return Response(PayrollYearlySettingsSerializer(existing).data)

        # پیدا کردن تنظیمات مادر (سراسری)
        source = get_effective_payroll_settings(effective_from, None)
        if not source:
            source = PayrollYearlySettings.objects.first()
        if not source:
            return Response({'error': 'تنظیمات مادر سراسری برای کپی یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        with transaction.atomic():
            new_settings = PayrollYearlySettings.objects.create(
                fiscal_year=year,
                effective_from=effective_from,
                version_title=source.version_title or "احکام مصوب فروردین",
                project=project,
                title=f"تنظیمات {project.name} - سال {year}",
                is_active=True,
                monthly_food_allowance=source.monthly_food_allowance,
                monthly_housing_allowance=source.monthly_housing_allowance,
                monthly_spouse_allowance=source.monthly_spouse_allowance,
                monthly_child_allowance=source.monthly_child_allowance,
                shift_percent=source.shift_percent,
                transport_help_percent=source.transport_help_percent,
                transport_fixed_amount=source.transport_fixed_amount,
                specialist_attraction_percent=source.specialist_attraction_percent,
                bad_weather_percent=source.bad_weather_percent,
                remote_hardship_percent=source.remote_hardship_percent,
                south_pars_percent=source.south_pars_percent,
                travel_cost_per_day=source.travel_cost_per_day,
                worker_insurance_rate=source.worker_insurance_rate,
                employer_insurance_rate=source.employer_insurance_rate,
                unemployment_insurance_rate=source.unemployment_insurance_rate,
                surplus_overtime_percent=source.surplus_overtime_percent,
                standard_daily_hours=getattr(source, 'standard_daily_hours', 10.00) or 10.00,
                bonus_daily_coefficient=getattr(source, 'bonus_daily_coefficient', 5.00) or 5.00,
                seniority_monthly_coefficient=getattr(source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                overtime_rate_multiplier=getattr(source, 'overtime_rate_multiplier', 1.40) or 1.40,
                friday_work_rate_multiplier=getattr(source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                max_insurable_daily_wage=getattr(source, 'max_insurable_daily_wage', None),
                attendance_edit_past_days=source.attendance_edit_past_days,
                attendance_edit_future_days=source.attendance_edit_future_days,
                created_by=request.user,
                updated_by=request.user
            )
            _copy_payroll_submodels(source, new_settings)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='CREATE',
                    severity='info',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(new_settings.id),
                    target_repr=str(new_settings),
                    details={
                        'action': 'clone_for_project',
                        'fiscal_year': year,
                        'effective_from': effective_from,
                        'project_id': p_id,
                        'project_name': project.name
                    }
                )
            except Exception:
                pass

        return Response(PayrollYearlySettingsSerializer(new_settings).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='revert-to-global')
    def revert_to_global(self, request, pk=None):
        """
        حذف تنظیمات اختصاصی پروژه و بازگشت به ارث‌بری از تنظیمات سراسری سازمان
        """
        settings_obj = self.get_object()
        if not settings_obj.project:
            return Response({'error': 'تنظیمات سراسری سازمان قابل حذف یا بازگشت به ارث‌بری نیست.'}, status=status.HTTP_400_BAD_REQUEST)

        project_name = settings_obj.project.name
        project_id = settings_obj.project.id
        fiscal_year = settings_obj.fiscal_year

        with transaction.atomic():
            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='DELETE',
                    severity='warning',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(settings_obj.id),
                    target_repr=str(settings_obj),
                    details={
                        'action': 'revert_to_global',
                        'fiscal_year': fiscal_year,
                        'project_id': project_id,
                        'project_name': project_name
                    }
                )
            except Exception:
                pass

            settings_obj.delete()

        return Response({
            'message': f'تنظیمات اختصاصی پروژه «{project_name}» با موفقیت حذف شد و به تنظیمات سراسری سازمان بازگشت.',
            'fiscal_year': fiscal_year
        })

    @action(detail=False, methods=['get'], url_path='job-grade-rate')
    def get_job_grade_rate(self, request):
        """
        دریافت خودکار مزد روزانه و پایه سنواتی بر اساس گروه شغلی (۱ تا ۲۰)
        """
        grade = safe_int(request.query_params.get('grade'), default=19)
        year = request.query_params.get('year', '1405')
        p_id = safe_int(request.query_params.get('project_id'))

        settings_obj = get_effective_payroll_settings(f"{year}/01", p_id)

        tier = None
        if settings_obj:
            tier = JobGradeTier.objects.filter(
                yearly_settings=settings_obj,
                grade_number=grade
            ).first()

        if not tier:
            tier = JobGradeTier.objects.filter(grade_number=grade).first()

        if tier:
            std_h = float(getattr(settings_obj, 'standard_daily_hours', 10.00) or 10.00)
            return Response({
                'grade_number': tier.grade_number,
                'daily_base_wage': float(tier.daily_base_wage),
                'daily_seniority_bonus': float(tier.daily_seniority_bonus),
                'hourly_rate': round(float(tier.daily_base_wage) / (std_h if std_h > 0 else 10.0), 2)
            })
        return Response({'error': 'گروه شغلی در جدول تنظیمات یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

    @action(detail=True, methods=['post'], url_path='update-all')
    def update_all_tabs(self, request, pk=None):
        """
        به‌روزرسانی یکپارچه هر ۵ تب تنظیمات با استفاده از سریالایزر جامع Writable Nested
        """
        settings_obj = self.get_object()
        data = request.data

        serializer = PayrollYearlySettingsSerializer(settings_obj, data=data, partial=True)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            updated_instance = serializer.save(updated_by=request.user)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='UPDATE',
                    severity='info',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(updated_instance.id),
                    target_repr=str(updated_instance),
                    details={
                        'action': 'update_all_tabs',
                        'fiscal_year': updated_instance.fiscal_year,
                        'version_title': updated_instance.version_title,
                        'project_id': updated_instance.project_id
                    }
                )
            except Exception:
                pass

        return Response({
            'message': 'تنظیمات پایه حقوق و دستمزد با موفقیت به‌روزرسانی شد.',
            'settings': PayrollYearlySettingsSerializer(updated_instance).data
        })


# ══════════════════════════════════════════════════════════════════════════════
# ── ۸. ویوست محاسبه ماهانه حقوق و صدور دیسکت‌ها و گزارشات ۵۸ ستونی ────────────
# ══════════════════════════════════════════════════════════════════════════════

class MonthlyPayrollViewSet(viewsets.ModelViewSet):
    queryset = MonthlyPayrollRecord.objects.all().select_related('period', 'personnel')
    serializer_class = MonthlyPayrollRecordSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        from common.tenant_scope import scope_tenant_queryset
        qs = scope_tenant_queryset(qs, user, company_field='personnel__company_id', request=self.request)

        period_id = self.request.query_params.get('period_id')
        if period_id:
            qs = qs.filter(period_id=period_id)
        
        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(period__warehouse_id=warehouse_id)
            
        year_month = self.request.query_params.get('year_month')
        if year_month:
            qs = qs.filter(period__year_month=year_month)

        return qs.order_by('row_number')

    def update(self, request, *args, **kwargs):
        from .lifecycle_services import user_can_view_wages
        if not user_can_view_wages(request.user):
            raise PermissionDenied("ویرایش فیش حقوقی تنها با دسترسی مالی و حسابداری مجاز است.")
        if 'payment_hold' in request.data or 'payment_hold_reason' in request.data:
            return Response(
                {'error': 'تغییر مستقیم وضعیت تعلیق پرداخت مجاز نیست. آزادسازی پرداخت منحصراً از طریق اکشن آزادسازی پس از تمدید اعتبار امکان‌پذیر است.'},
                status=status.HTTP_400_BAD_REQUEST
            )
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        from .lifecycle_services import user_can_view_wages
        if not user_can_view_wages(request.user):
            raise PermissionDenied("ویرایش فیش حقوقی تنها با دسترسی مالی و حسابداری مجاز است.")
        if 'payment_hold' in request.data or 'payment_hold_reason' in request.data:
            return Response(
                {'error': 'تغییر مستقیم وضعیت تعلیق پرداخت مجاز نیست. آزادسازی پرداخت منحصراً از طریق اکشن آزادسازی پس از تمدید اعتبار امکان‌پذیر است.'},
                status=status.HTTP_400_BAD_REQUEST
            )
        return super().partial_update(request, *args, **kwargs)

    @action(detail=False, methods=['post'], url_path='calculate-period')
    def calculate_period(self, request):
        """
        تجمیع کارکرد و محاسبه خودکار ۵۸ ستون حقوق برای دوره انتخابی (با پشتیبانی از سطح شرکت یا انبار)
        """
        warehouse_id = request.data.get('warehouse_id')
        year_month = request.data.get('year_month')
        from common.tenant_scope import get_current_tenant_company_id
        company_id = request.data.get('company_id') or get_current_tenant_company_id(request=request, user=request.user)
        
        if not year_month:
            return Response({'error': 'سال/ماه الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not warehouse_id and not company_id:
            return Response({'error': 'شناسه انبار یا شرکت فعال الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        wh_val = int(warehouse_id) if (warehouse_id and str(warehouse_id).isdigit()) else None
        comp_val = int(company_id) if (company_id and str(company_id).isdigit()) else None

        lookup = {'year_month': str(year_month).strip()}
        if wh_val:
            lookup['warehouse_id'] = wh_val
        if comp_val:
            lookup['company_id'] = comp_val

        # Get or create work period
        period, _ = MonthlyWorkPeriod.objects.get_or_create(
            **lookup,
            defaults={'status': 'OPEN', 'company_id': comp_val, 'warehouse_id': wh_val}
        )

        if period.status == 'LOCKED':
            return Response({'error': f'دوره {period.year_month} قفل شده است و امکان محاسبه مجدد ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        records = calculate_monthly_payroll_for_period(period, user=request.user)
        serialized = MonthlyPayrollRecordSerializer(records, many=True).data

        # Summary Metrics
        total_gross = sum(float(r.gross_salary) for r in records)
        total_payable = sum(float(r.payable_amount) for r in records)
        total_insurance = sum(float(r.total_insurance) for r in records)
        total_tax = sum(float(r.income_tax) for r in records)

        return Response({
            'message': f'محاسبه حقوق دوره {period.year_month} برای {len(records)} نفر پرسنل با موفقیت انجام شد.',
            'period_id': period.id,
            'period_status': period.status,
            'summary': {
                'total_personnel': len(records),
                'total_gross': total_gross,
                'total_payable': total_payable,
                'total_insurance': total_insurance,
                'total_tax': total_tax,
            },
            'records': serialized
        })

    @staticmethod
    def _attach_expired_headers(resp, excluded_expired):
        """
        DEF-20: کنار گذاشتن پرسنل منقضی از دیسکت بی‌صدا نیست. شمار و نام‌ها (JSON با کدگذاری URL،
        چون سربرگ HTTP فقط ASCII می‌پذیرد) همراه فایل برمی‌گردد؛ فرانت پیش از دانلود از اکشن expired-personnel می‌پرسد.
        """
        if not excluded_expired:
            return
        import json
        from urllib.parse import quote
        names = [e['full_name'] for e in excluded_expired][:50]
        resp['X-Expired-Personnel-Excluded'] = str(len(excluded_expired))
        resp['X-Expired-Personnel-Names'] = quote(json.dumps(names, ensure_ascii=False))
        resp['Access-Control-Expose-Headers'] = 'Content-Disposition, X-Expired-Personnel-Excluded, X-Expired-Personnel-Names'

    @action(detail=False, methods=['get'], url_path='expired-personnel')
    def expired_personnel(self, request):
        """
        DEF-20: پیش‌پرسش خروجی دیسکت. فهرست پرسنلی که به‌خاطر پایان تاریخ اعتبار از دیسکت بیمه، مالیات
        و بانک این دوره و پروژه کنار گذاشته می‌شوند تا پیش از دانلود به کاربر هشدار داده شود.
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')
        if not period_id or not project_id:
            return Response({'error': 'شناسه دوره (period_id) و پروژه (project_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not str(period_id).isdigit() or not str(project_id).isdigit():
            return Response({'error': 'شناسه دوره یا پروژه نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        from .lifecycle_services import split_expired_payroll_records
        # فقط فیش‌های درون قلمرو کاربر (get_queryset)؛ دوره مستقیم و بی‌قلمرو خوانده نمی‌شود
        records = [
            r for r in self.get_queryset().filter(period_id=int(period_id)).select_related('personnel', 'period')
            if str(r.personnel.project_id) == str(project_id)
        ]
        year_month = records[0].period.year_month if records else None
        excluded = split_expired_payroll_records(records, year_month)[1] if records else []
        return Response({
            'period_id': int(period_id),
            'year_month': year_month,
            'project_id': int(project_id),
            'count': len(excluded),
            'personnel': excluded,
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='export-bimeh-diskettes')
    def export_bimeh_diskettes(self, request):
        return self.export_dsk_zip(request)

    @action(detail=False, methods=['get'], url_path='export-dsk-zip')
    def export_dsk_zip(self, request):
        """
        صدور پکیج فشرده ZIP شامل دو دیسکت استاندارد DSKWOR00.DBF و DSKKAR00.DBF
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور دیسکت‌های بیمه تامین اجتماعی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = FinancialProject.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        # فیلتر تک‌پروژه‌ای الزامی بر اساس انتساب پرسنل به پروژه
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'هیچ رکوردی برای پروژه «{project.name}» در این دوره محاسبه نشده است.'}, status=status.HTTP_404_NOT_FOUND)

        year_str = period.year_month.split('/')[0] if '/' in period.year_month else '1405'
        month_str = period.year_month.split('/')[1] if '/' in period.year_month else '04'
        fiscal_year_short = int(year_str[-2:]) # 05
        month_num = int(month_str)

        # استخراج تنظیمات کارگاه به ترتیب: تنظیم اختصاصی پروژه -> تنظیم سراسری سال در بازه موثر
        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        ws_settings = getattr(settings_obj, 'workshop_insurance', None) if settings_obj else None
        if not ws_settings:
            ws_settings = WorkshopInsuranceSettings.objects.first()
        if not ws_settings:
            ws_settings = WorkshopInsuranceSettings(
                workshop_code='4894290013',
                workshop_name=f'کارگاه {project.name}',
                employer_name='مدیریت انبارداری',
                workshop_address='منطقه عملیاتی'
            )

        # فیلتر پرسنل معتبر و کنار گذاشتن پرسنل منقضی در دیسکت با گزارش هشدار (تسک ۲.۱ و تصمیم مصوب ۶.۶)
        from .lifecycle_services import split_expired_payroll_records
        records, excluded_expired = split_expired_payroll_records(records, period.year_month)

        insured_recs = [r for r in records if r.include_in_insurance]
        if not insured_recs:
            return Response({'error': f'هیچ پرسنل مشمول بیمه‌ای برای پروژه «{project.name}» در این دوره وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        # Calculate sums for DSKKAR
        sum_days = sum(r.insurance_days for r in insured_recs)
        sum_daily = sum(int(r.daily_wage) for r in insured_recs)
        sum_base = sum(int(r.base_salary) for r in insured_recs)
        sum_maz = sum(int(r.total_taxable_allowances) for r in insured_recs)
        sum_mash = sum(int(r.total_insurable_salary_allowances) for r in insured_recs)
        sum_totl = sum(int(r.gross_salary) for r in insured_recs)
        sum_bime = sum(int(r.worker_insurance) for r in insured_recs)
        sum_koso = sum(int(r.employer_insurance) for r in insured_recs)
        sum_bic = sum(int(r.unemployment_insurance) for r in insured_recs)
        sum_inc = sum(int(r.total_seniority_accumulated) for r in insured_recs)
        sum_spouse = sum(int(r.marital_allowance) for r in insured_recs)

        month_names = {1: 'فروردین', 2: 'اردیبهشت', 3: 'خرداد', 4: 'تیر', 5: 'مرداد', 6: 'شهریور', 7: 'مهر', 8: 'آبان', 9: 'آذر', 10: 'دی', 11: 'بهمن', 12: 'اسفند'}
        m_name = month_names.get(month_num, 'تیر')
        list_desc = f"{project.name} بیمه {m_name} {year_str}"[:50]

        dskkar_bytes = generate_dskkar_bytes(
            workshop_settings=ws_settings,
            fiscal_year_short=fiscal_year_short,
            month_num=month_num,
            list_number='01',
            list_description=list_desc,
            records_count=len(insured_recs),
            sum_insurance_days=sum_days,
            sum_daily_wage=sum_daily,
            sum_base_salary=sum_base,
            sum_insurable_allowances=sum_maz,
            sum_insurable_total=sum_mash,
            sum_gross_salary=sum_totl,
            sum_employee_insurance=sum_bime,
            sum_employer_insurance=sum_koso,
            sum_unemployment_insurance=sum_bic,
            sum_seniority=sum_inc,
            sum_spouse_allowance=sum_spouse
        )

        dskwor_bytes = generate_dskwor_bytes(
            workshop_code=ws_settings.workshop_code,
            fiscal_year_short=fiscal_year_short,
            month_num=month_num,
            list_number='01',
            payroll_records=records
        )

        # Build ZIP in memory
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr("DSKKAR00.DBF", dskkar_bytes)
            z.writestr("DSKWOR00.DBF", dskwor_bytes)

        zip_val = zip_buf.getvalue()
        resp = HttpResponse(zip_val, content_type='application/zip')
        resp['Content-Disposition'] = f'attachment; filename="Bimeh_{project.code}_{year_str}_{month_str}.zip"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=False, methods=['get'], url_path='export-tax-wh')
    def export_tax_wh(self, request):
        """
        صدور فایل متنی WH مالیات حقوق به تفکیک اجباری پروژه
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور فایل متنی WH مالیاتی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = FinancialProject.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'هیچ رکوردی برای پروژه «{project.name}» در این دوره محاسبه نشده است.'}, status=status.HTTP_404_NOT_FOUND)

        # فیلتر پرسنل معتبر و کنار گذاشتن پرسنل منقضی در دیسکت مالیات
        from .lifecycle_services import split_expired_payroll_records
        records, excluded_expired = split_expired_payroll_records(records, period.year_month)

        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        tax_settings = getattr(settings_obj, 'tax_settings', None)

        content = generate_wh_tax_content(records, tax_settings)
        month_code = period.year_month.replace('/', '')[2:] if period else '140504'
        filename = f"WH_{project.code}_{month_code}.txt"

        resp = HttpResponse(content.encode('utf-8'), content_type='text/plain; charset=utf-8')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=False, methods=['get'], url_path='export-tax-wp')
    def export_tax_wp(self, request):
        """
        صدور فایل متنی WP پرسنل مالیاتی به تفکیک اجباری پروژه
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور فایل متنی WP مالیاتی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = FinancialProject.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'هیچ رکوردی برای پروژه «{project.name}» در این دوره محاسبه نشده است.'}, status=status.HTTP_404_NOT_FOUND)

        # فیلتر پرسنل معتبر و کنار گذاشتن پرسنل منقضی در دیسکت مالیات WP
        from .lifecycle_services import split_expired_payroll_records
        records, excluded_expired = split_expired_payroll_records(records, period.year_month)

        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        tax_settings = getattr(settings_obj, 'tax_settings', None)

        content = generate_wp_tax_content(records, tax_settings)
        month_code = period.year_month.replace('/', '')[2:] if period else '140504'
        filename = f"WP_{project.code}_{month_code}.txt"

        resp = HttpResponse(content.encode('utf-8'), content_type='text/plain; charset=utf-8')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=False, methods=['get'], url_path='export-bank-excel')
    def export_bank_excel(self, request):
        """
        صدور فایل اکسل پرداخت گروهی بانک ملی به تفکیک اجباری پروژه.
        DEF-20 (گام ۲ اصلاحیه): بر مبنای علامت payment_hold روی فیش؛
        در صورت ارسال only_released=true، خروجی مجزا فقط شامل پرداخت‌های آزادشده صادر می‌شود.
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')
        only_released = request.query_params.get('only_released') in ['true', '1']

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور فایل پرداخت بانک ملی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = FinancialProject.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'رکوردی برای پروژه «{project.name}» در این دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        if only_released:
            records = [r for r in records if not r.payment_hold and r.payment_hold_released_at is not None]
            if not records:
                return Response({'error': f'هیچ فیش آزادشده‌ای برای پروژه «{project.name}» در این دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
            excluded_expired = []
            filename = f"Bank_Payment_Released_{project.code}_{period.year_month.replace('/', '_')}.xlsx"
        else:
            excluded_expired = [
                {
                    'personnel_id': r.personnel_id,
                    'full_name': r.personnel.full_name,
                    'national_code': r.personnel.national_code,
                    'valid_until_date': getattr(r.personnel, 'valid_until_date', None),
                }
                for r in records if r.payment_hold
            ]
            records = [r for r in records if not r.payment_hold and r.payment_hold_released_at is None]
            if not records:
                return Response({'error': f'هیچ فیش حقوقی عادی برای صدور دیسکت بانک پروژه «{project.name}» در این دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
            filename = f"Bank_Payment_{project.code}_{period.year_month.replace('/', '_')}.xlsx"

        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        bank_settings = getattr(settings_obj, 'bank_export_settings', None)

        excel_bytes = generate_bank_meli_excel(records, bank_settings, year_month_title=f"{project.name} {period.year_month}")

        resp = HttpResponse(excel_bytes, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=True, methods=['post'], url_path='release-payment-hold')
    def release_payment_hold(self, request, pk=None):
        """
        DEF-20 (گام ۲ اصلاحیه): آزادسازی پرداخت معلق پرسنل منقضی پس از تمدید اعتبار پرونده.
        اعتبار را می‌سنجد، لاگ ممیزی ثبت می‌کند و تعلیق پرداخت را آزاد می‌سازد.
        """
        from .lifecycle_services import user_can_view_wages, validate_personnel_validity
        from accounts.audit_utils import log_audit_event
        from django.utils import timezone

        if not user_can_view_wages(request.user):
            raise PermissionDenied("آزادسازی پرداخت معلق حقوق تنها با دسترسی مالی و حسابداری مجاز است.")

        record = self.get_object()
        if not record.payment_hold:
            return Response({'message': 'این فیش حقوقی در وضعیت تعلیق پرداخت قرار ندارد.'}, status=status.HTTP_200_OK)

        # سنجش تمدید اعتبار پرسنل
        is_valid, reason = validate_personnel_validity(record.personnel, record.period.year_month)
        if not is_valid:
            return Response({
                'error': f'اعتبار پرونده هنوز تمدید نشده است ({reason}). لطفاً ابتدا تاریخ اعتبار را از طریق درخواست تغییرات تمدید نمایید.'
            }, status=status.HTTP_400_BAD_REQUEST)

        record.payment_hold = False
        record.payment_hold_released_at = timezone.now()
        record.payment_hold_released_by = request.user
        record.save(update_fields=['payment_hold', 'payment_hold_released_at', 'payment_hold_released_by', 'updated_at'])

        log_audit_event(
            module='personnel',
            action='RELEASE_PAYMENT_HOLD',
            severity='info',
            target_model='MonthlyPayrollRecord',
            target_object_id=str(record.id),
            target_repr=f"آزادسازی پرداخت معلق فیش {record.personnel.full_name} برای دوره {record.period.year_month}",
            user=request.user,
            details={
                'record_id': record.id,
                'personnel_id': record.personnel_id,
                'period_id': record.period_id,
                'payable_amount': str(record.payable_amount),
            }
        )

        return Response({
            'message': f'پرداخت معلق فیش «{record.personnel.full_name}» با موفقیت آزاد شد.',
            'record_id': record.id,
            'payment_hold': False,
            'payment_hold_released_at': record.payment_hold_released_at.isoformat()
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='export-monthly-excel')
    def export_monthly_excel(self, request):
        """
        صدور فایل اکسل ۲ سطری استاندارد ۵۸ ستون شیت ماهانه
        """
        period_id = request.query_params.get('period_id')
        records = list(MonthlyPayrollRecord.objects.filter(period_id=period_id).select_related('personnel'))
        if not records:
            return Response({'error': 'رکوردی یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        period = records[0].period
        excel_bytes = generate_monthly_payroll_excel(records, period_title=period.year_month)
        filename = f"Payroll_Report_{period.year_month.replace('/', '_')}.xlsx"

        resp = HttpResponse(excel_bytes, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        return resp

    @action(detail=False, methods=['post'], url_path='import-tax-excel')
    def import_tax_excel(self, request):
        """
        درون‌ریزی فایل اکسل مالیات محاسبه‌شده توسط دارایی (نمونه مالیات محاسبه شده.xlsx)
        یا لیست داده‌های مالیات ارسالی متنی/JSON و تطبیق خودکار با ستون مالیات دوره بر اساس کد ملی
        """
        excel_file = request.FILES.get('file')
        items = request.data.get('items')
        period_id = request.data.get('period_id') or request.query_params.get('period_id')
        year_month = request.data.get('year_month') or request.query_params.get('year_month')

        if not excel_file and not items:
            return Response({'error': 'لطفاً فایل اکسل مالیات دارایی یا داده‌های متنی را انتخاب کنید.'}, status=status.HTTP_400_BAD_REQUEST)

        # Locate target period
        if period_id:
            try:
                p_id = int(str(period_id).strip())
                period = MonthlyWorkPeriod.objects.filter(id=p_id).first()
            except Exception:
                period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        elif year_month:
            period = MonthlyWorkPeriod.objects.filter(year_month=year_month).first()
        else:
            period = MonthlyWorkPeriod.objects.order_by('-created_at').first()

        if not period:
            return Response({'error': 'دوره ماهانه مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        try:
            matched_count = 0
            unmatched_rows = []
            total_imported_tax = Decimal(0)

            records = MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel')
            record_map = {str(r.national_code).strip().zfill(10): r for r in records}

            if items and isinstance(items, list):
                with transaction.atomic():
                    for idx, item in enumerate(items, 1):
                        raw_code = item.get('national_code')
                        if not raw_code:
                            continue
                        code_str = str(raw_code).split('.')[0].strip().zfill(10)
                        if not code_str or not code_str.isdigit():
                            continue
                        try:
                            tax_val = Decimal(str(item.get('tax_amount', 0)).replace(',', '').strip())
                        except Exception:
                            tax_val = Decimal(0)
                        try:
                            ex_months = int(item.get('exemption_months', 1) or 1)
                        except Exception:
                            ex_months = 1
                        raw_mult = item.get('multiple_employers', 'خیر')
                        has_mult = str(raw_mult).strip() in ['بله', 'true', 'True', '1', 'YES', 'yes']

                        if code_str in record_map:
                            rec = record_map[code_str]
                            rec.income_tax = tax_val
                            rec.tax_source_type = 'IMPORTED_EXCEL'
                            rec.tax_exemption_months = ex_months
                            rec.has_multiple_employers = has_mult
                            rec.is_tax_imported = True

                            rec.net_salary = rec.gross_salary - rec.worker_insurance - rec.income_tax
                            rec.payable_amount = rec.net_salary - rec.advance_payment_deduction
                            rec.save()

                            matched_count += 1
                            total_imported_tax += tax_val
                        else:
                            unmatched_rows.append({
                                'row': idx,
                                'national_code': code_str,
                                'tax_amount': float(tax_val)
                            })

            elif excel_file:
                file_bytes = excel_file.read()
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                ws = wb.active

                # Find header indices
                col_map = {}
                for c in range(1, ws.max_column + 1):
                    header_val = str(ws.cell(1, c).value or '').strip()
                    clean_h = header_val.replace(' ', '').replace('_', '')
                    if 'کدملی' in clean_h or 'کد_ملی' in clean_h:
                        col_map['national_code'] = c
                    elif 'مالیاتمحاسبهشده' in clean_h or 'مالیات' in clean_h:
                        col_map['tax_amount'] = c
                    elif 'تعدادماهاعمالشدهدرمعافیت' in clean_h or 'تعدادماه' in clean_h:
                        col_map['exemption_months'] = c
                    elif 'بیشازیککارفرما' in clean_h or 'چندکارفرما' in clean_h:
                        col_map['multiple_employers'] = c

                if 'national_code' not in col_map or 'tax_amount' not in col_map:
                    col_map.setdefault('national_code', 1)
                    col_map.setdefault('tax_amount', 4)
                    col_map.setdefault('exemption_months', 5)
                    col_map.setdefault('multiple_employers', 6)

                with transaction.atomic():
                    for r_idx in range(2, ws.max_row + 1):
                        raw_code = ws.cell(r_idx, col_map['national_code']).value
                        if raw_code is None:
                            continue
                        
                        code_str = str(raw_code).split('.')[0].strip().zfill(10)
                        if not code_str or not code_str.isdigit():
                            continue

                        raw_tax = ws.cell(r_idx, col_map['tax_amount']).value or 0
                        try:
                            tax_val = Decimal(str(raw_tax).replace(',', '').strip())
                        except Exception:
                            tax_val = Decimal(0)

                        raw_months = ws.cell(r_idx, col_map.get('exemption_months', 5)).value if 'exemption_months' in col_map else 1
                        try:
                            ex_months = int(raw_months or 1)
                        except Exception:
                            ex_months = 1

                        raw_mult = ws.cell(r_idx, col_map.get('multiple_employers', 6)).value if 'multiple_employers' in col_map else 'خیر'
                        has_mult = str(raw_mult).strip() in ['بله', 'true', 'True', '1', 'YES', 'yes']

                        if code_str in record_map:
                            rec = record_map[code_str]
                            rec.income_tax = tax_val
                            rec.tax_source_type = 'IMPORTED_EXCEL'
                            rec.tax_exemption_months = ex_months
                            rec.has_multiple_employers = has_mult
                            rec.is_tax_imported = True

                            rec.net_salary = rec.gross_salary - rec.worker_insurance - rec.income_tax
                            rec.payable_amount = rec.net_salary - rec.advance_payment_deduction
                            rec.save()

                            matched_count += 1
                            total_imported_tax += tax_val
                        else:
                            unmatched_rows.append({
                                'row': r_idx,
                                'national_code': code_str,
                                'tax_amount': float(tax_val)
                            })

            return Response({
                'message': f'اطلاعات مالیاتی با موفقیت بارگذاری شد. {matched_count} پرسنل تطبیق داده شدند.',
                'matched_count': matched_count,
                'unmatched_count': len(unmatched_rows),
                'unmatched_rows': unmatched_rows,
                'total_imported_tax': float(total_imported_tax),
                'period_year_month': period.year_month
            }, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({
                'error': f'خطا در پردازش اطلاعات مالیات: {str(e)}'
            }, status=status.HTTP_400_BAD_REQUEST)


class FleetSettlementViewSet(viewsets.ViewSet):
    """
    ویوست تسویه و محاسبات مالی ناوگان خودرویی
    - محاسبه مبالغ تسویه ماهانه
    - تولید فایل پرداخت گروهی بانک ملی
    """
    permission_classes = [permissions.IsAuthenticated]

    @action(detail=False, methods=['get'], url_path='calculate')
    def calculate(self, request):
        year_month = request.query_params.get('year_month')
        warehouse_id = request.query_params.get('warehouse_id')
        if not year_month:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        wh_id = int(warehouse_id) if (warehouse_id and warehouse_id.isdigit()) else None
        try:
            data = calculate_monthly_fleet_settlement(warehouse_id=wh_id, year_month=year_month)
            return Response(data, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'error': f'خطا در محاسبه تسویه ناوگان: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], url_path='export-bank-excel')
    def export_bank_excel(self, request):
        year_month = request.query_params.get('year_month')
        warehouse_id = request.query_params.get('warehouse_id')
        if not year_month:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        wh_id = int(warehouse_id) if (warehouse_id and warehouse_id.isdigit()) else None
        try:
            data = calculate_monthly_fleet_settlement(warehouse_id=wh_id, year_month=year_month)
            excel_file = generate_fleet_bank_meli_excel(data, year_month)

            clean_ym = year_month.replace('/', '_')
            filename = f"fleet_bank_payment_{clean_ym}.xlsx"
            response = HttpResponse(
                excel_file.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response
        except Exception as e:
            return Response({'error': f'خطا در تولید فایل اکسل پرداخت بانک ملی: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)


# ══════════════════════════════════════════════════════════════════════════════
def broadcast_org_structure_updated(entity_type, action, entity_id=None, name=None, project_id=None, client_tab_id=None, sender_id=None, company_id=None):
    """
    ارسال بلادرنگ رویداد تغییرات ساختار سازمانی به وب‌سوکت شرکتی؛ اگر فقط پروژه داده شده باشد،
    شرکتِ پروژه پیدا می‌شود و ارسال به هستهٔ companies سپرده می‌شود.
    """
    cid = company_id
    if not cid and project_id:
        try:
            cid = FinancialProject.objects.filter(id=project_id).values_list('company_id', flat=True).first()
        except Exception:
            pass
    core_broadcast_org_structure_updated(entity_type, action, entity_id, name, project_id, client_tab_id, sender_id, cid)


# ویوست‌های مدیریت ساختار سازمانی، شرکت، پروژه، بخش، طرف‌حساب و فاکتور
# ══════════════════════════════════════════════════════════════════════════════

class CompanyViewSet(CoreCompanyViewSet):
    """
    مدیریت شرکت‌ها در حسابداری: ویوست هسته به‌علاوهٔ شمار پروژه‌ها، مدارک، حساب‌های بانکی،
    اعضای هیئت‌مدیره و پایش سررسید مدارک.
    """
    queryset = Company.objects.all().annotate(projects_count=Count('projects'))
    serializer_class = CompanySerializer

    def base_queryset(self):
        return Company.objects.all().annotate(
            projects_count=Count('projects')
        ).prefetch_related('documents', 'bank_accounts', 'board_members')

    @action(detail=False, methods=['get'], url_path='expiring-documents')
    def expiring_documents(self, request):
        """
        لیست کلیه مدارک منقضی شده یا در آستانه انقضا (کمتر از ۳۰ روز) برای تمام شرکت‌های مجاز کاربر
        جهت نمایش در ویجت پایش سررسید مدارک داشبورد مرکز عملیات
        """
        from datetime import timedelta
        user = request.user
        qs = CompanyDocument.objects.select_related('company', 'uploaded_by').filter(expiry_date__isnull=False)

        if not user.is_superuser:
            allowed_ids = get_user_allowed_companies(user, required_level='docs_read').values_list('id', flat=True)
            qs = qs.filter(company_id__in=allowed_ids, is_confidential=False)

        today = timezone.now().date()
        threshold = today + timedelta(days=30)
        expiring_qs = qs.filter(expiry_date__lte=threshold).order_by('expiry_date')

        page = self.paginate_queryset(expiring_qs)
        if page is not None:
            serializer = CompanyDocumentSerializer(page, many=True, context={'request': request})
            return self.get_paginated_response(serializer.data)

        serializer = CompanyDocumentSerializer(expiring_qs, many=True, context={'request': request})
        return Response({
            'count': expiring_qs.count(),
            'results': serializer.data
        })


class CompanyDocumentViewSet(viewsets.ModelViewSet):
    """
    مدیریت اسناد، مدارک و مجوزهای رسمی شرکت‌ها با پشتیبانی از آپلود فایل،
    اعتبارسنجی BOLA و حفاظت از اسناد محرمانه
    """
    queryset = CompanyDocument.objects.all().select_related('company', 'uploaded_by')
    serializer_class = CompanyDocumentSerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        user = self.request.user
        qs = CompanyDocument.objects.select_related('company', 'uploaded_by')

        company_id = self.request.query_params.get('company_id') or get_request_company_id(self.request)
        if company_id:
            cid = validate_user_company_access(user, company_id, required_level='docs_read')
            if cid:
                qs = qs.filter(company_id=cid)
        elif not user.is_superuser:
            allowed_ids = get_user_allowed_companies(user, required_level='docs_read').values_list('id', flat=True)
            qs = qs.filter(company_id__in=allowed_ids)

        # عدم نمایش اسناد محرمانه به کاربران غیرسوپریوزر/غیرادمین
        if not (user.is_superuser or user.is_staff):
            qs = qs.filter(is_confidential=False)

        doc_type = self.request.query_params.get('document_type')
        if doc_type:
            qs = qs.filter(document_type=doc_type)

        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(Q(title__icontains=search) | Q(description__icontains=search))

        return qs.order_by('-created_at')

    def perform_create(self, serializer):
        user = self.request.user
        company = serializer.validated_data.get('company')
        if not company:
            raise serializers.ValidationError({"company": "انتخاب شرکت الزامی است."})

        # بررسی دسترسی کاربر به این شرکت (حداقل سطح ثبت اسناد)
        validate_user_company_access(user, company.id, required_level='docs_write')

        # اگر سند محرمانه علامت زده شده، کاربر باید مدیر ارشد باشد
        is_confidential = serializer.validated_data.get('is_confidential', False)
        if is_confidential and not (user.is_superuser or user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به بارگذاری اسناد محرمانه هستند.")

        file_obj = self.request.FILES.get('file')
        file_size = file_obj.size if file_obj else 0

        instance = serializer.save(uploaded_by=user, file_size=file_size)
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_document', 'create', instance.id, instance.title, instance.company_id, tab_id, user.id)

    def perform_update(self, serializer):
        user = self.request.user
        instance = serializer.instance
        validate_user_company_access(user, instance.company_id, required_level='docs_write')

        # عدم دسترسی کاربر عادی به ویرایش اسناد محرمانه
        if instance.is_confidential and not (user.is_superuser or user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به ویرایش اسناد محرمانه هستند.")

        # عدم اجازه به کاربر عادی جهت تغییر وضعیت محرمانگی سند
        new_confidential = serializer.validated_data.get('is_confidential', instance.is_confidential)
        if new_confidential != instance.is_confidential and not (user.is_superuser or user.is_staff):
            raise PermissionDenied("تغییر وضعیت محرمانگی سند تنها توسط مدیران ارشد مجاز است.")

        file_obj = self.request.FILES.get('file')
        kwargs = {}
        if file_obj:
            kwargs['file_size'] = file_obj.size

        updated_instance = serializer.save(**kwargs)
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_document', 'update', updated_instance.id, updated_instance.title, updated_instance.company_id, tab_id, user.id)

    def perform_destroy(self, instance):
        user = self.request.user
        validate_user_company_access(user, instance.company_id, required_level='docs_write')
        if not (user.is_superuser or user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به حذف مدارک رسمی شرکت هستند.")

        instance_id = instance.id
        instance_title = instance.title
        comp_id = instance.company_id
        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_document', 'delete', instance_id, instance_title, comp_id, tab_id, user.id)

    @action(detail=True, methods=['get'], url_path='download')
    def download(self, request, pk=None):
        """
        دانلود امن سند با اعمال گارد امنیتی BOLA، کنترل محرمانگی و بازگرداندن FileResponse
        """
        import mimetypes
        import os
        import re
        from urllib.parse import quote
        from django.http import FileResponse, Http404

        user = request.user
        try:
            document = CompanyDocument.objects.select_related('company').get(pk=pk)
        except CompanyDocument.DoesNotExist:
            raise Http404("سند مورد نظر یافت نشد.")

        # ۱. بررسی دسترسی کاربر به شرکت متبوع (BOLA با حداقل سطح مشاهده)
        validate_user_company_access(user, document.company_id, required_level='docs_read')

        # ۲. بررسی سند محرمانه (صرفاً سوپریوزر و ادمین مجازند)
        if document.is_confidential and not (user.is_superuser or user.is_staff):
            raise PermissionDenied("شما مجوز دسترسی به این سند محرمانه را ندارید.")

        # ۳. بررسی وجود فایل در حافظه ذخیره‌سازی
        if not document.file or not document.file.storage.exists(document.file.name):
            raise Http404("فایل سند در حافظه سرور یافت نشد.")

        file_handle = document.file.open('rb')
        content_type, _ = mimetypes.guess_type(document.file.name)
        content_type = content_type or 'application/octet-stream'

        # پاک‌سازی کاراکترهای غیرمجاز در نام فایل دانلودی و رعایت استاندارد RFC 6266
        ext = os.path.splitext(document.file.name)[1]
        safe_title = re.sub(r'[\\/*?:"<>|]', '_', document.title or 'document').strip()
        download_name = f"{safe_title}{ext}"
        ascii_fallback = f"doc_{document.id}{ext}"
        encoded_name = quote(download_name.encode('utf-8'))

        response = FileResponse(file_handle, content_type=content_type, as_attachment=True, filename=ascii_fallback)
        response['Content-Disposition'] = f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{encoded_name}"

        return response



class CompanyBankAccountViewSet(viewsets.ModelViewSet):
    """
    مدیریت حساب‌های بانکی و شماره‌های شبای شرکت‌ها با پشتیبانی از چند حسابی،
    تعیین حساب اصلی و تفکیک سازمانی
    """
    queryset = CompanyBankAccount.objects.all().select_related('company')
    serializer_class = CompanyBankAccountSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        user = self.request.user
        qs = CompanyBankAccount.objects.select_related('company')

        company_id = self.request.query_params.get('company_id') or get_request_company_id(self.request)
        if company_id:
            cid = validate_user_company_access(user, company_id)
            if cid:
                qs = qs.filter(company_id=cid)
        elif not user.is_superuser:
            allowed_ids = get_user_allowed_companies(user).values_list('id', flat=True)
            qs = qs.filter(company_id__in=allowed_ids)

        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(
                Q(bank_name__icontains=search) |
                Q(account_number__icontains=search) |
                Q(sheba_number__icontains=search) |
                Q(account_title__icontains=search)
            )

        return qs.order_by('-is_primary', '-created_at')

    def perform_create(self, serializer):
        user = self.request.user
        company = serializer.validated_data.get('company')
        if not company:
            raise serializers.ValidationError({"company": "انتخاب شرکت الزامی است."})

        validate_user_company_access(user, company.id)
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_bank_account', 'create', instance.id, instance.sheba_number, instance.company_id, tab_id, user.id)

    def perform_update(self, serializer):
        user = self.request.user
        instance = serializer.instance
        validate_user_company_access(user, instance.company_id)
        updated_instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_bank_account', 'update', updated_instance.id, updated_instance.sheba_number, updated_instance.company_id, tab_id, user.id)

    def perform_destroy(self, instance):
        user = self.request.user
        validate_user_company_access(user, instance.company_id)
        instance_id = instance.id
        sheba = instance.sheba_number
        comp_id = instance.company_id
        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_bank_account', 'delete', instance_id, sheba, comp_id, tab_id, user.id)

    @action(detail=True, methods=['post'], url_path='set-primary')
    def set_primary(self, request, pk=None):
        """
        تنظیم یک حساب بانکی به عنوان حساب پیش‌فرض/اصلی شرکت
        """
        account = self.get_object()
        validate_user_company_access(request.user, account.company_id)
        account.is_primary = True
        account.save()
        tab_id = request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_bank_account', 'update', account.id, account.sheba_number, account.company_id, tab_id, request.user.id)
        return Response(CompanyBankAccountSerializer(account).data, status=status.HTTP_200_OK)


class CompanyBoardMemberViewSet(viewsets.ModelViewSet):
    """
    مدیریت اعضای هیئت‌مدیره، صاحبان امضا و ارکان حاکمیتی شرکت‌ها با پشتیبانی از آپلود مدارک هویتی و احکام
    """
    queryset = CompanyBoardMember.objects.all().select_related('company')
    serializer_class = CompanyBoardMemberSerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        user = self.request.user
        qs = CompanyBoardMember.objects.select_related('company')

        company_id = self.request.query_params.get('company_id') or get_request_company_id(self.request)
        if company_id:
            cid = validate_user_company_access(user, company_id)
            if cid:
                qs = qs.filter(company_id=cid)
        elif not user.is_superuser:
            allowed_ids = get_user_allowed_companies(user).values_list('id', flat=True)
            qs = qs.filter(company_id__in=allowed_ids)

        search = self.request.query_params.get('search')
        if search:
            search_clean = normalize_digits(str(search)).strip()
            qs = qs.filter(
                Q(first_name__icontains=search) |
                Q(last_name__icontains=search) |
                Q(national_code__icontains=search_clean) |
                Q(represented_legal_name__icontains=search)
            )

        return qs.order_by('-has_signature_right', 'role', '-created_at')

    def perform_create(self, serializer):
        user = self.request.user
        company = serializer.validated_data.get('company')
        if not company:
            raise serializers.ValidationError({"company": "انتخاب شرکت الزامی است."})

        validate_user_company_access(user, company.id)
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_board_member', 'create', instance.id, f"{instance.first_name} {instance.last_name}", instance.company_id, tab_id, user.id)

    def perform_update(self, serializer):
        user = self.request.user
        instance = serializer.instance
        validate_user_company_access(user, instance.company_id)
        updated_instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_board_member', 'update', updated_instance.id, f"{updated_instance.first_name} {updated_instance.last_name}", updated_instance.company_id, tab_id, user.id)

    def perform_destroy(self, instance):
        user = self.request.user
        validate_user_company_access(user, instance.company_id)
        if not (user.is_superuser or user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به حذف اعضای هیئت‌مدیره هستند.")
        instance_id = instance.id
        title = f"{instance.first_name} {instance.last_name}"
        comp_id = instance.company_id
        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_board_member', 'delete', instance_id, title, comp_id, tab_id, user.id)


class CompanyFiscalPeriodViewSet(viewsets.ModelViewSet):
    """
    مدیریت دوره‌ها و سال‌های مالی شرکت‌ها (ایجاد، بستن دوره، انجماد کاردکس، بازگشایی)
    """
    queryset = CompanyFiscalPeriod.objects.all().select_related('company', 'closed_by')
    serializer_class = CompanyFiscalPeriodSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        user = self.request.user
        qs = CompanyFiscalPeriod.objects.select_related('company', 'closed_by')

        company_id = self.request.query_params.get('company_id') or get_request_company_id(self.request)
        if company_id:
            cid = validate_user_company_access(user, company_id)
            if cid:
                qs = qs.filter(company_id=cid)
        elif not user.is_superuser:
            allowed_ids = get_user_allowed_companies(user).values_list('id', flat=True)
            qs = qs.filter(company_id__in=allowed_ids)

        year = self.request.query_params.get('fiscal_year')
        if year:
            qs = qs.filter(fiscal_year=year)

        status_param = self.request.query_params.get('status')
        if status_param:
            qs = qs.filter(status=status_param)

        return qs.order_by('-fiscal_year')

    def perform_create(self, serializer):
        user = self.request.user
        company = serializer.validated_data.get('company')
        if not company:
            raise serializers.ValidationError({"company": "انتخاب شرکت الزامی است."})
        validate_user_company_access(user, company.id)
        serializer.save()

    def perform_update(self, serializer):
        user = self.request.user
        instance = serializer.instance
        validate_user_company_access(user, instance.company_id)
        if instance.status == 'closed' and not user.is_superuser:
            raise PermissionDenied("سال مالی بسته شده است و ویرایش آن صرفاً توسط مدیر ارشد امکان‌پذیر است.")
        serializer.save()

    def perform_destroy(self, instance):
        user = self.request.user
        validate_user_company_access(user, instance.company_id)
        if not (user.is_superuser or user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به حذف دوره مالی هستند.")
        if instance.status == 'closed':
            raise serializers.ValidationError("دوره مالی بسته شده را نمی‌توان حذف کرد.")
        instance.delete()

    @action(detail=True, methods=['post'], url_path='close-period')
    def close_period(self, request, pk=None):
        """بستن قطعی دوره و سال مالی"""
        period = self.get_object()
        user = request.user
        validate_user_company_access(user, period.company_id)
        if not (user.is_superuser or user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به بستن سال مالی هستند.")
        period.status = 'closed'
        period.closed_at = timezone.now()
        period.closed_by = user
        period.notes = request.data.get('notes', period.notes)
        period.save(update_fields=['status', 'closed_at', 'closed_by', 'notes', 'updated_at'])
        return Response({'detail': f'سال مالی {period.fiscal_year} با موفقیت بسته و قطعی شد.'})

    @action(detail=True, methods=['post'], url_path='freeze-period')
    def freeze_period(self, request, pk=None):
        """انجماد موقت دوره مالی جهت رسیدگی و انبارگردانی"""
        period = self.get_object()
        user = request.user
        validate_user_company_access(user, period.company_id)
        period.status = 'frozen'
        period.save(update_fields=['status', 'updated_at'])
        return Response({'detail': f'دوره مالی {period.fiscal_year} موقتاً منجمد شد.'})

    @action(detail=True, methods=['post'], url_path='reopen-period')
    def reopen_period(self, request, pk=None):
        """بازگشایی مجدد دوره مالی منجمد یا بسته شده (سوپریوزر)"""
        period = self.get_object()
        user = request.user
        if not user.is_superuser:
            raise PermissionDenied("صرفاً مدیر کل مجاز به بازگشایی دوره مالی بسته شده است.")
        period.status = 'open'
        period.closed_at = None
        period.closed_by = None
        period.save(update_fields=['status', 'closed_at', 'closed_by', 'updated_at'])
        return Response({'detail': f'دوره مالی {period.fiscal_year} با موفقیت بازگشایی شد.'})


class FinancialProjectViewSet(viewsets.ModelViewSet):
    """
    مدیریت پروژه‌های مالی و عملیاتی (مراکز هزینه مستقل)
    """
    queryset = FinancialProject.objects.all().annotate(sections_count=Count('sections'))
    serializer_class = FinancialProjectSerializer
    permission_classes = [permissions.IsAuthenticated, IsProjectStructureManagerOrReadOnly]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user

        # عایق‌سازی شرکت و رفع رخنه امنیتی BOLA:
        if not user.is_authenticated:
            return qs.none()

        company_param = self.request.query_params.get('company_id') or self.request.query_params.get('company') or self.request.headers.get('X-Company-ID')
        if not user.is_superuser:
            allowed_company_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if company_param:
                try:
                    cid = int(company_param)
                    if cid not in allowed_company_ids:
                        raise PermissionDenied("شما به اطلاعات و پروژه‌های این شرکت دسترسی ندارید.")
                    qs = qs.filter(company_id=cid)
                except (ValueError, TypeError):
                    pass
            else:
                qs = qs.filter(Q(company_id__in=allowed_company_ids) | Q(company__isnull=True))
        elif company_param:
            try:
                qs = qs.filter(company_id=int(company_param))
            except (ValueError, TypeError):
                pass

        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() == 'true')
        search = self.request.query_params.get('search')
        if search:
            search_clean = normalize_digits(str(search)).strip()
            qs = qs.filter(Q(name__icontains=search) | Q(code__icontains=search_clean))
        return qs.order_by('code')

    def perform_create(self, serializer):
        user = self.request.user
        company = serializer.validated_data.get('company')
        if requires_company_scope(user):
            ensure_scoped_write(user, company.id if company else None, "پروژه فقط برای شرکت‌های مجاز شما ساخته می‌شود (پروژهٔ سراسری را فقط مدیر ارشد می‌سازد).")
        if company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if company.id not in allowed_ids:
                raise PermissionDenied("شما مجاز به ایجاد پروژه برای این شرکت نیستید.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('project', 'create', instance.id, instance.name, None, tab_id, self.request.user.id)

    def perform_update(self, serializer):
        user = self.request.user
        if requires_company_scope(user):
            ensure_scoped_write(user, serializer.instance.company_id, "فقط پروژه‌های شرکت‌های مجاز خود را می‌توانید ویرایش کنید (پروژهٔ سراسری را فقط مدیر ارشد تغییر می‌دهد).")
            if 'company' in serializer.validated_data:
                new_company = serializer.validated_data.get('company')
                # جابه‌جایی پروژه بین شرکت‌ها اعضای بخش‌هایش را عضو شرکت تازه می‌کند؛ فقط سوپریوزر
                if (new_company.id if new_company else None) != serializer.instance.company_id:
                    raise PermissionDenied("انتقال پروژه به شرکت دیگر یا تبدیل آن به پروژهٔ سراسری را فقط مدیر ارشد انجام می‌دهد.")
        company = serializer.validated_data.get('company') or serializer.instance.company
        if company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if company.id not in allowed_ids:
                raise PermissionDenied("شما مجاز به ویرایش پروژه در این شرکت نیستید.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('project', 'update', instance.id, instance.name, None, tab_id, self.request.user.id)

    def perform_destroy(self, instance):
        user = self.request.user
        if requires_company_scope(user):
            ensure_scoped_write(user, instance.company_id, "فقط پروژه‌های شرکت‌های مجاز خود را می‌توانید حذف کنید (پروژهٔ سراسری را فقط مدیر ارشد تغییر می‌دهد).")
        if instance.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if instance.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به حذف پروژه این شرکت نیستید.")
        instance_id = instance.id
        instance_name = instance.name

        # بررسی محافظت و قوانین عدم حذف زنجیره‌ای (تسک ۱.۲ و ۳۵)
        has_invoices = ExpenseInvoice.objects.filter(section__project=instance).exists()
        has_payroll_settings = PayrollYearlySettings.objects.filter(project=instance).exists()
        has_attendance = DailyAttendance.objects.filter(project=instance).exists()
        has_trips = VehicleTripLog.objects.filter(project=instance).exists()
        has_personnel = PersonnelProfile.objects.filter(project=instance).exists()

        if has_invoices or has_payroll_settings or has_attendance or has_trips or has_personnel:
            # جلوگیری قاطع از حذف فیزیکی جهت حفاظت از اسناد مالی و تبدیل به غیرفعال‌سازی نرم
            instance.is_active = False
            instance.save(update_fields=['is_active', 'updated_at'])
            tab_id = self.request.headers.get('X-Client-Tab-Id')
            broadcast_org_structure_updated('project', 'update', instance_id, instance_name, None, tab_id, self.request.user.id)
            return

        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('project', 'delete', instance_id, instance_name, None, tab_id, self.request.user.id)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        from .org_excel_engine import export_projects_excel
        return export_projects_excel(company_ids=readable_company_ids(request.user))

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        from .org_excel_engine import download_org_template
        return download_org_template('projects')

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request):
        from .org_excel_engine import import_projects_from_excel
        # غیرسوپریوزر فقط درون شرکت فعالِ مجاز خودش درون‌ریزی می‌کند
        company_id = resolve_import_company(request)
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل بارگذاری نشده است.'}, status=status.HTTP_400_BAD_REQUEST)
        dry_run = request.data.get('dry_run') in ['true', 'True', True]
        result = import_projects_from_excel(file_obj, dry_run=dry_run, company_id=company_id)
        if not dry_run:
            tab_id = request.headers.get('X-Client-Tab-Id')
            broadcast_org_structure_updated('project', 'import', None, 'اکسل پروژه‌ها', None, tab_id, request.user.id)
        return Response(result, status=status.HTTP_200_OK if result['success'] else status.HTTP_400_BAD_REQUEST)


class ProjectSectionViewSet(viewsets.ModelViewSet):
    """
    مدیریت بخش‌ها و دپارتمان‌های تابعه هر پروژه
    """
    queryset = ProjectSection.objects.all().select_related('project')
    serializer_class = ProjectSectionSerializer
    permission_classes = [permissions.IsAuthenticated, IsProjectStructureManagerOrReadOnly]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(project__company_id=cid)
        elif not user.is_superuser and user.is_authenticated:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            qs = qs.filter(Q(project__company_id__in=allowed_ids) | Q(project__company__isnull=True))

        project_id = self.request.query_params.get('project_id') or self.request.query_params.get('project')
        if project_id:
            try:
                qs = qs.filter(project_id=int(project_id))
            except (ValueError, TypeError):
                pass
        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() == 'true')
        search = self.request.query_params.get('search')
        if search:
            search_clean = normalize_digits(str(search)).strip()
            qs = qs.filter(Q(name__icontains=search) | Q(code__icontains=search_clean))
        return qs.order_by('project__code', 'code')

    def perform_create(self, serializer):
        user = self.request.user
        project = serializer.validated_data.get('project')
        if requires_company_scope(user):
            ensure_scoped_write(user, project.company_id if project else None, "بخش فقط در پروژه‌های شرکت‌های مجاز شما ساخته می‌شود (پروژهٔ سراسری را فقط مدیر ارشد تغییر می‌دهد).")
        if project and project.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if project.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به ایجاد بخش برای پروژه‌های این شرکت نیستید.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('section', 'create', instance.id, instance.name, instance.project_id, tab_id, self.request.user.id)

    def perform_update(self, serializer):
        user = self.request.user
        if requires_company_scope(user):
            current_project = serializer.instance.project
            ensure_scoped_write(user, current_project.company_id if current_project else None, "فقط بخش‌های شرکت‌های مجاز خود را می‌توانید ویرایش کنید (بخش پروژهٔ سراسری را فقط مدیر ارشد تغییر می‌دهد).")
            if 'project' in serializer.validated_data:
                new_project = serializer.validated_data.get('project')
                ensure_scoped_write(user, new_project.company_id if new_project else None, "انتقال بخش به پروژهٔ شرکت دیگر یا پروژهٔ سراسری مجاز نیست.")
                # جابه‌جایی بخش بین شرکت‌ها اعضایش را عضو شرکت تازه می‌کند؛ فقط سوپریوزر
                if new_project.company_id != (current_project.company_id if current_project else None):
                    raise PermissionDenied("انتقال بخش به پروژهٔ شرکت دیگر را فقط مدیر ارشد انجام می‌دهد.")
        project = serializer.validated_data.get('project') or self.get_object().project
        if project and project.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if project.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به ویرایش بخش در پروژه‌های این شرکت نیستید.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('section', 'update', instance.id, instance.name, instance.project_id, tab_id, self.request.user.id)

    def perform_destroy(self, instance):
        user = self.request.user
        if requires_company_scope(user):
            ensure_scoped_write(user, instance.project.company_id if instance.project else None, "فقط بخش‌های شرکت‌های مجاز خود را می‌توانید حذف کنید (بخش پروژهٔ سراسری را فقط مدیر ارشد تغییر می‌دهد).")
        if instance.project and instance.project.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if instance.project.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به حذف بخش‌های این شرکت نیستید.")
        instance_id = instance.id
        instance_name = instance.name
        project_id = instance.project_id

        # بررسی محافظت و قوانین عدم حذف زنجیره‌ای (تسک ۱.۲ و ۳۵)
        has_invoices = ExpenseInvoice.objects.filter(section=instance).exists()
        has_attendance = DailyAttendance.objects.filter(section=instance).exists()
        has_trips = VehicleTripLog.objects.filter(section=instance).exists()
        has_personnel = PersonnelProfile.objects.filter(section=instance).exists()

        if has_invoices or has_attendance or has_trips or has_personnel:
            instance.is_active = False
            instance.save(update_fields=['is_active', 'updated_at'])
            tab_id = self.request.headers.get('X-Client-Tab-Id')
            broadcast_org_structure_updated('section', 'update', instance_id, instance_name, project_id, tab_id, self.request.user.id)
            return

        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('section', 'delete', instance_id, instance_name, project_id, tab_id, self.request.user.id)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        from .org_excel_engine import export_sections_excel
        project_id = request.query_params.get('project_id')
        return export_sections_excel(project_id, company_ids=readable_company_ids(request.user))

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        from .org_excel_engine import download_org_template
        return download_org_template('sections')

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request):
        from .org_excel_engine import import_sections_from_excel
        # غیرسوپریوزر فقط در پروژه‌های شرکت فعالِ مجاز خودش درون‌ریزی می‌کند
        company_id = resolve_import_company(request)
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل بارگذاری نشده است.'}, status=status.HTTP_400_BAD_REQUEST)
        dry_run = request.data.get('dry_run') in ['true', 'True', True]
        result = import_sections_from_excel(file_obj, dry_run=dry_run, company_id=company_id)
        if not dry_run:
            tab_id = request.headers.get('X-Client-Tab-Id')
            broadcast_org_structure_updated('section', 'import', None, 'اکسل بخش‌ها', None, tab_id, request.user.id)
        return Response(result, status=status.HTTP_200_OK if result['success'] else status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='clone-assignments')
    def clone_assignments(self, request, pk=None):
        """
        کپی دسته‌جمعی پرسنل یک بخش مبدأ به این بخش (مقصد)
        """
        target_section = self.get_object()
        source_section_id = request.data.get('source_section_id')
        if not source_section_id:
            return Response({'error': 'شناسه بخش مبدأ الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            source_sec_id_int = int(source_section_id)
        except (ValueError, TypeError):
            return Response({'error': 'شناسه بخش مبدأ نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        if source_sec_id_int == target_section.id:
            return Response({'error': 'بخش مبدأ و مقصد نمی‌توانند یکسان باشند.'}, status=status.HTTP_400_BAD_REQUEST)

        source_section = ProjectSection.objects.filter(id=source_sec_id_int).first()
        if not source_section:
            return Response({'error': 'بخش مبدأ یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        if requires_company_scope(request.user):
            # مبدأ و مقصد هر دو در شرکت‌های مجاز راهبر و در یک شرکت، تا کپی برای کسی دسترسی شرکت تازه نسازد
            target_company_id = target_section.project.company_id if target_section.project_id else None
            source_company_id = source_section.project.company_id if source_section.project_id else None
            ensure_scoped_write(request.user, target_company_id, "کپی انتساب فقط در بخش‌های شرکت‌های مجاز شما ممکن است.")
            if source_company_id != target_company_id:
                raise PermissionDenied("کپی انتساب فقط بین بخش‌های یک شرکت مجاز است.")

        source_assignments = UserSectionAssignment.objects.filter(section=source_section, is_active=True)
        if not source_assignments.exists():
            return Response({'error': 'هیچ انتساب فعالی در بخش مبدأ جهت کپی وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        # راهبر تشکیلات (بدون مجوز قدیمی نوشتن ساختار) با کپی هم نقش تازه برای خودش نمی‌سازد
        if requires_company_scope(request.user) and not has_structure_write_permission(request.user):
            source_assignments = source_assignments.exclude(user=request.user)

        copied_count = 0
        skipped_count = 0
        with transaction.atomic():
            for sa in source_assignments:
                existing_assignment = UserSectionAssignment.objects.filter(
                    user=sa.user,
                    section=target_section,
                    role=sa.role
                ).first()

                if not existing_assignment:
                    UserSectionAssignment.objects.create(
                        user=sa.user,
                        section=target_section,
                        role=sa.role,
                        is_active=True
                    )
                    copied_count += 1
                else:
                    if not existing_assignment.is_active:
                        existing_assignment.is_active = True
                        existing_assignment.save(update_fields=['is_active', 'updated_at'])
                        copied_count += 1
                    else:
                        skipped_count += 1

        tab_id = request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated(
            'assignment',
            'clone',
            target_section.id,
            f'کپی چارت از {source_section.name} به {target_section.name}',
            target_section.project_id,
            tab_id,
            request.user.id
        )

        return Response({
            'success': True,
            'message': f'ساختار پرسنلی با موفقیت کپی شد. ({copied_count} انتساب جدید/فعال‌شده، {skipped_count} بدون تغییر)',
            'copied_count': copied_count,
            'skipped_count': skipped_count
        }, status=status.HTTP_200_OK)


class UserSectionAssignmentViewSet(viewsets.ModelViewSet):
    """
    مدیریت انتساب کاربران به بخش‌ها و تعیین نقش‌های سازمانی
    """
    queryset = UserSectionAssignment.objects.all().select_related('user', 'section', 'section__project')
    serializer_class = UserSectionAssignmentSerializer
    permission_classes = [permissions.IsAuthenticated, IsProjectStructureManagerOrReadOnly]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if requires_company_scope(user):
            # فهرست اعضای بخش‌های شرکت‌های دیگر نباید دیده شود (بخش‌های سراسری و انتساب‌های خود کاربر می‌مانند)
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            qs = qs.filter(
                Q(section__project__company_id__in=allowed_ids)
                | Q(section__project__company__isnull=True)
                | Q(user=user)
            )
        section_id = self.request.query_params.get('section_id')
        if section_id:
            try:
                qs = qs.filter(section_id=int(section_id))
            except (ValueError, TypeError):
                pass
        project_id = self.request.query_params.get('project_id')
        if project_id:
            try:
                qs = qs.filter(section__project_id=int(project_id))
            except (ValueError, TypeError):
                pass
        user_id = self.request.query_params.get('user_id')
        if user_id:
            try:
                qs = qs.filter(user_id=int(user_id))
            except (ValueError, TypeError):
                pass
        role = self.request.query_params.get('role')
        if role:
            qs = qs.filter(role=role)
        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() == 'true')
        return qs.order_by('-created_at')

    @action(detail=False, methods=['get'], url_path='my-sections')
    def my_sections(self, request):
        """
        دریافت بخش‌های فعال منتسب به کاربر لاگین‌شده جاری
        """
        user = request.user
        if not user.is_authenticated:
            return Response([], status=status.HTTP_200_OK)

        company_param = request.query_params.get('company_id') or request.query_params.get('company') or request.headers.get('X-Company-ID')
        target_company_id = None
        if company_param:
            try:
                target_company_id = int(company_param)
            except (ValueError, TypeError):
                pass

        if not user.is_superuser and target_company_id:
            allowed_cids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if target_company_id not in allowed_cids:
                return Response([], status=status.HTTP_200_OK)

        if user.is_superuser or user.is_staff:
            sections = ProjectSection.objects.filter(is_active=True, project__is_active=True).select_related('project')
            if target_company_id:
                sections = sections.filter(project__company_id=target_company_id)
            return Response(ProjectSectionSerializer(sections, many=True).data)

        assignments = UserSectionAssignment.objects.filter(
            user=user,
            is_active=True,
            section__is_active=True,
            section__project__is_active=True
        ).select_related('section', 'section__project')

        if target_company_id:
            assignments = assignments.filter(section__project__company_id=target_company_id)

        sections = [a.section for a in assignments]
        unique_sections = {s.id: s for s in sections}.values()
        return Response(ProjectSectionSerializer(unique_sections, many=True).data)

    def create(self, request, *args, **kwargs):
        """
        ایجاد انتساب یا فعال‌سازی مجدد رکورد غیرفعال قبلی بدون خطای یکتایی دیتابیس
        """
        user_id = request.data.get('user')
        section_id = request.data.get('section')
        role = request.data.get('role', 'employee')

        try:
            user_id_int = int(user_id) if user_id else None
            sec_id_int = int(section_id) if section_id else None
        except (ValueError, TypeError):
            user_id_int, sec_id_int = None, None

        if user_id_int and sec_id_int and requires_company_scope(request.user):
            # پیش از مسیر فعال‌سازی دوباره هم باید محدودهٔ شرکت راهبر بررسی شود
            section = ProjectSection.objects.select_related('project').filter(id=sec_id_int).first()
            if section is not None:
                ensure_scoped_assignment(request.user, user_id_int, section)

        if user_id_int and sec_id_int:
            existing = UserSectionAssignment.objects.filter(
                user_id=user_id_int,
                section_id=sec_id_int,
                role=role
            ).first()
            if existing:
                if not existing.is_active:
                    existing.is_active = True
                    existing.save(update_fields=['is_active', 'updated_at'])
                    tab_id = request.headers.get('X-Client-Tab-Id') or request.data.get('client_tab_id')
                    proj_id = existing.section.project_id if existing.section else None
                    broadcast_org_structure_updated('assignment', 'create', existing.id, str(existing.user), proj_id, tab_id, request.user.id)
                serializer = self.get_serializer(existing)
                return Response(serializer.data, status=status.HTTP_201_CREATED)

        return super().create(request, *args, **kwargs)

    def perform_create(self, serializer):
        if requires_company_scope(self.request.user):
            target_user = serializer.validated_data.get('user')
            ensure_scoped_assignment(
                self.request.user, target_user.id if target_user else None, serializer.validated_data.get('section'),
                final_active=serializer.validated_data.get('is_active', True)
            )
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        proj_id = instance.section.project_id if instance.section else None
        broadcast_org_structure_updated('assignment', 'create', instance.id, str(instance.user), proj_id, tab_id, self.request.user.id)

    def perform_update(self, serializer):
        if requires_company_scope(self.request.user):
            current = serializer.instance
            current_company_id = current.section.project.company_id if current.section and current.section.project_id else None
            ensure_scoped_write(self.request.user, current_company_id, "فقط انتساب‌های بخش‌های شرکت‌های مجاز خود را می‌توانید ویرایش کنید.")
            final_user = serializer.validated_data.get('user', current.user)
            final_section = serializer.validated_data.get('section', current.section)
            ensure_scoped_assignment(
                self.request.user, final_user.id if final_user else None, final_section,
                final_active=serializer.validated_data.get('is_active', current.is_active)
            )
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        proj_id = instance.section.project_id if instance.section else None
        broadcast_org_structure_updated('assignment', 'update', instance.id, str(instance.user), proj_id, tab_id, self.request.user.id)

    def perform_destroy(self, instance):
        if requires_company_scope(self.request.user):
            company_id = instance.section.project.company_id if instance.section and instance.section.project_id else None
            ensure_scoped_write(self.request.user, company_id, "فقط انتساب‌های بخش‌های شرکت‌های مجاز خود را می‌توانید حذف کنید.")
        instance_id = instance.id
        proj_id = instance.section.project_id if instance.section else None
        # غیرفعال‌سازی نرم جهت حفظ ردپای حسابرسی و انتساب‌های تاریخی
        instance.is_active = False
        instance.save(update_fields=['is_active', 'updated_at'])
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('assignment', 'delete', instance_id, None, proj_id, tab_id, self.request.user.id)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        from .org_excel_engine import export_assignments_excel
        project_id = request.query_params.get('project_id')
        role = request.query_params.get('role')
        return export_assignments_excel(project_id, role, company_ids=readable_company_ids(request.user))

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        from .org_excel_engine import download_org_template
        return download_org_template('assignments')

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request):
        from .org_excel_engine import import_assignments_from_excel
        # غیرسوپریوزر فقط بخش‌های شرکت فعالِ مجاز خودش را و فقط برای اعضای همان شرکت انتساب می‌دهد
        company_id = resolve_import_company(request)
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل بارگذاری نشده است.'}, status=status.HTTP_400_BAD_REQUEST)
        dry_run = request.data.get('dry_run') in ['true', 'True', True]
        forbidden_self = request.user.id if company_id is not None and not has_structure_write_permission(request.user) else None
        result = import_assignments_from_excel(
            file_obj, dry_run=dry_run, company_id=company_id,
            member_check=user_belongs_to_company if requires_company_scope(request.user) else None,
            forbidden_user_id=forbidden_self
        )
        if not dry_run:
            tab_id = request.headers.get('X-Client-Tab-Id')
            broadcast_org_structure_updated('assignment', 'import', None, 'اکسل انتساب‌ها', None, tab_id, request.user.id)
        return Response(result, status=status.HTTP_200_OK if result['success'] else status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], url_path='bulk-assign')
    def bulk_assign(self, request):
        """
        انتساب همزمان چند کاربر به یک بخش و سمت سازمانی با اعتبارسنجی کامل ورودی‌ها
        """
        section_id = request.data.get('section_id') or request.data.get('section')
        user_ids = request.data.get('user_ids') or []
        role = request.data.get('role', 'employee')

        if not section_id:
            return Response({'error': 'شناسه بخش الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            section_id_int = int(section_id)
        except (ValueError, TypeError):
            return Response({'error': 'شناسه بخش باید عددی باشد.'}, status=status.HTTP_400_BAD_REQUEST)

        valid_roles = dict(UserSectionAssignment.ROLE_CHOICES).keys()
        if role not in valid_roles:
            return Response({'error': f'نقش ارسالی نامعتبر است. نقش‌های مجاز: {list(valid_roles)}'}, status=status.HTTP_400_BAD_REQUEST)

        if not user_ids or not isinstance(user_ids, list):
            return Response({'error': 'لیست شناسه‌های کاربران الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        valid_user_ids = []
        for uid in user_ids:
            try:
                valid_user_ids.append(int(uid))
            except (ValueError, TypeError):
                pass

        if not valid_user_ids:
            return Response({'error': 'هیچ شناسه کاربری معتبری ارسال نشده است.'}, status=status.HTTP_400_BAD_REQUEST)

        section = ProjectSection.objects.filter(id=section_id_int).first()
        if not section:
            return Response({'error': 'بخش سازمانی مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        # اعتبارسنجی وجود کاربران در دیتابیس
        existing_users = set(User.objects.filter(id__in=valid_user_ids).values_list('id', flat=True))

        if requires_company_scope(request.user):
            # همه یا هیچ: بخش باید در شرکت‌های مجاز باشد و هر کاربر عضو شرکتِ این بخش.
            # کاربر ناموجود هم مثل غیرعضو گزارش می‌شود تا وجود نام‌های کاربری بیرون از شرکت لو نرود.
            company_id = section.project.company_id if section.project_id else None
            ensure_scoped_write(request.user, company_id, "انتساب فقط در بخش‌های شرکت‌های مجاز شما ممکن است (بخش پروژهٔ سراسری مجاز نیست).")
            rejected = []
            for uid in valid_user_ids:
                if uid not in existing_users:
                    rejected.append(uid)
                    continue
                try:
                    ensure_scoped_assignment(request.user, uid, section)
                except PermissionDenied:
                    rejected.append(uid)
            if rejected:
                # نام کاربری فقط برای کاربرانی که در شرکت‌های مجاز نویسنده دیده می‌شوند (بقیه با شناسه)
                allowed_ids = set(get_user_allowed_companies(request.user).values_list('id', flat=True))
                names = dict(User.objects.filter(id__in=rejected).filter(
                    Q(company_accesses__company_id__in=allowed_ids, company_accesses__access_level='workspace_full')
                    | Q(section_assignments__section__project__company_id__in=allowed_ids)
                ).distinct().values_list('id', 'username'))
                labels = '، '.join(names.get(uid, f'#{uid}') for uid in rejected)
                raise PermissionDenied(f"این کاربران را نمی‌توان در این بخش منتسب کرد (عضو شرکت این بخش نیستند یا انتساب خود است): {labels}. آن‌ها را از انتخاب بردارید.")

        if not existing_users:
            return Response({'error': 'هیچ یک از کاربران ارسالی در سیستم یافت نشدند.'}, status=status.HTTP_404_NOT_FOUND)

        created_assignments = []
        with transaction.atomic():
            for uid in valid_user_ids:
                if uid not in existing_users:
                    continue
                assignment, _ = UserSectionAssignment.objects.get_or_create(
                    user_id=uid,
                    section=section,
                    role=role,
                    defaults={'is_active': True}
                )
                if not assignment.is_active:
                    assignment.is_active = True
                    assignment.save(update_fields=['is_active', 'updated_at'])
                created_assignments.append(assignment)

        tab_id = request.headers.get('X-Client-Tab-Id') or request.data.get('client_tab_id')
        proj_id = section.project_id
        broadcast_org_structure_updated('assignment', 'bulk_create', None, f"{len(created_assignments)} کاربر", proj_id, tab_id, request.user.id)

        assign_ids = [a.id for a in created_assignments]
        qs = UserSectionAssignment.objects.filter(id__in=assign_ids).select_related('user', 'section', 'section__project')
        serializer = self.get_serializer(qs, many=True)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class CounterpartyViewSet(viewsets.ModelViewSet):
    """
    مدیریت طرف‌حساب‌های مالی (رانندگان، تعمیرگاه‌ها، جایگاه سوخت و پیمانکاران)
    """
    queryset = Counterparty.objects.all()
    serializer_class = CounterpartySerializer
    permission_classes = [permissions.IsAuthenticated, IsCounterpartyManagerOrReadOnly]
    pagination_class = OptionalPageNumberPagination

    def get_permissions(self):
        if self.action in ['list', 'retrieve', 'create']:
            return [permissions.IsAuthenticated()]
        return [permissions.IsAuthenticated(), IsCounterpartyManagerOrReadOnly()]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if not user.is_authenticated:
            return qs.none()

        # ایزولاسیون چندمستاجری طرف‌حساب‌ها بر مبنای شرکت فعال
        company_param = get_request_company_id(self.request)
        if not user.is_superuser:
            allowed_cids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if company_param:
                try:
                    cid = int(company_param)
                    if cid not in allowed_cids:
                        raise PermissionDenied("شما به طرف‌حساب‌های این شرکت دسترسی ندارید.")
                    qs = qs.filter(Q(company_id=cid) | Q(company__isnull=True))
                except (ValueError, TypeError):
                    pass
            else:
                qs = qs.filter(Q(company_id__in=allowed_cids) | Q(company__isnull=True))
        elif company_param:
            try:
                qs = qs.filter(Q(company_id=int(company_param)) | Q(company__isnull=True))
            except (ValueError, TypeError):
                pass

        c_type = self.request.query_params.get('counterparty_type')
        if c_type:
            qs = qs.filter(counterparty_type=c_type)
        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() == 'true')
        search = self.request.query_params.get('search')
        if search:
            search_clean = normalize_digits(str(search)).strip()
            qs = qs.filter(Q(name__icontains=search) | Q(national_id__icontains=search_clean) | Q(phone__icontains=search_clean) | Q(account_code__icontains=search_clean))
        return qs.order_by('name')

    def perform_create(self, serializer):
        user = self.request.user
        if requires_company_scope(user):
            # شرکت به ترتیب: بدنه (اگر مجاز) ← هدر شرکت فعال (اگر مجاز) ← تنها شرکت مجاز کاربر.
            # شرکت غیرمجاز نادیده گرفته می‌شود و رد نمی‌شود، چون پاسخ 4xx رکورد صف آفلاین را حذف می‌کند؛
            # پس تزریق به شرکت دیگر ممکن نیست و ثبت مشروعِ صف‌شده هم از دست نمی‌رود.
            allowed_cids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            explicit_company = serializer.validated_data.pop('company', None)
            candidates = [explicit_company.id if explicit_company else None]
            try:
                req_cid = get_request_company_id(self.request)
                candidates.append(int(req_cid) if req_cid else None)
            except (ValueError, TypeError):
                pass
            if len(allowed_cids) == 1:
                candidates.append(next(iter(allowed_cids)))
            company_id = next((c for c in candidates if c is not None and c in allowed_cids), None)
            if company_id is not None:
                serializer.validated_data['company_id'] = company_id
        elif not serializer.validated_data.get('company'):
            # سوپریوزر: شرکت فعال درخواست، اگر در بدنه نیامده باشد
            req_cid = get_request_company_id(self.request)
            if req_cid:
                try:
                    serializer.validated_data['company_id'] = int(req_cid)
                except (ValueError, TypeError):
                    pass
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('counterparty', 'create', instance.id, instance.name, None, tab_id, self.request.user.id)

    def perform_update(self, serializer):
        user = self.request.user
        if requires_company_scope(user):
            ensure_scoped_write(user, serializer.instance.company_id, "فقط طرف‌حساب‌های شرکت‌های مجاز خود را می‌توانید ویرایش کنید (طرف‌حساب سراسری را فقط مدیر ارشد تغییر می‌دهد).")
            if 'company' in serializer.validated_data:
                new_company = serializer.validated_data.get('company')
                ensure_scoped_write(user, new_company.id if new_company else None, "تبدیل طرف‌حساب به طرف‌حساب سراسری یا انتقال آن به شرکت دیگر مجاز نیست.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('counterparty', 'update', instance.id, instance.name, None, tab_id, self.request.user.id)

    def perform_destroy(self, instance):
        if requires_company_scope(self.request.user):
            ensure_scoped_write(self.request.user, instance.company_id, "فقط طرف‌حساب‌های شرکت‌های مجاز خود را می‌توانید حذف کنید (طرف‌حساب سراسری را فقط مدیر ارشد تغییر می‌دهد).")
        instance_id = instance.id
        instance_name = instance.name

        # بررسی وجود فاکتور وابسته جهت جلوگیری از حذف مخرب
        if ExpenseInvoice.objects.filter(counterparty=instance).exists():
            instance.is_active = False
            instance.save(update_fields=['is_active', 'updated_at'])
            tab_id = self.request.headers.get('X-Client-Tab-Id')
            broadcast_org_structure_updated('counterparty', 'update', instance_id, instance_name, None, tab_id, self.request.user.id)
            return

        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('counterparty', 'delete', instance_id, instance_name, None, tab_id, self.request.user.id)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        from .org_excel_engine import export_counterparties_excel
        section_id = request.query_params.get('section_id')
        return export_counterparties_excel(section_id, company_ids=readable_company_ids(request.user))

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        from .org_excel_engine import download_org_template
        return download_org_template('counterparties')

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request):
        from .org_excel_engine import import_counterparties_from_excel
        # غیرسوپریوزر فقط طرف‌حساب‌های شرکت فعالِ مجاز خودش را می‌سازد یا به‌روز می‌کند
        company_id = resolve_import_company(request)
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایلی ارسال نشده است'}, status=status.HTTP_400_BAD_REQUEST)
        dry_run = request.data.get('dry_run') in ['true', 'True', True]
        res = import_counterparties_from_excel(file_obj, dry_run=dry_run, company_id=company_id)
        if not dry_run:
            tab_id = request.headers.get('X-Client-Tab-Id')
            broadcast_org_structure_updated('counterparty', 'import', None, 'Bulk Import', None, tab_id, request.user.id)
        return Response(res, status=status.HTTP_200_OK if res['success'] else status.HTTP_400_BAD_REQUEST)


class ExpenseInvoiceViewSet(viewsets.ModelViewSet):
    """
    مدیریت فاکتورهای هزینه‌ای بخش‌ها
    """
    queryset = ExpenseInvoice.objects.all().select_related('section', 'section__project', 'counterparty', 'created_by')
    serializer_class = ExpenseInvoiceSerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(section__project__company_id=cid)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(section_id=section_id)
        project_id = self.request.query_params.get('project_id')
        if project_id:
            qs = qs.filter(section__project_id=project_id)
        inv_status = self.request.query_params.get('status')
        if inv_status:
            qs = qs.filter(status=inv_status)
        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(Q(invoice_number__icontains=search) | Q(counterparty__name__icontains=search) | Q(description__icontains=search))
        return qs.order_by('-invoice_date_shamsi', '-created_at')

    def perform_create(self, serializer):
        user = self.request.user
        section = serializer.validated_data.get('section')
        if section and section.project and section.project.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if section.project.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به ثبت فاکتور برای این شرکت نیستید.")
        # Guardian G2: default to draft status
        req_status = self.request.data.get('status')
        if not user.is_superuser:
            instance = serializer.save(created_by=user, status='draft')
        else:
            instance = serializer.save(created_by=user, status=req_status or 'draft')
        if req_status and req_status != 'draft':
            from .workflow_engine import process_creation_with_auto_pass
            process_creation_with_auto_pass(instance, user, is_financial=True)
            instance.save()

    def perform_update(self, serializer):
        user = self.request.user
        instance = self.get_object()
        new_status = serializer.validated_data.get('status')
        if new_status and new_status != instance.status and not user.is_superuser:
            from .workflow_engine import determine_user_tier, WorkflowTiers
            tier = determine_user_tier(user)
            if new_status in ['draft', 'pending_supervisor']:
                pass
            elif new_status in ['pending_accountant', 'revision_required', 'rejected'] and tier >= WorkflowTiers.TIER_2_SUPERVISOR:
                pass
            elif new_status in ['pending_manager', 'ready_to_pay', 'revision_required', 'rejected'] and tier >= WorkflowTiers.TIER_3_ACCOUNTANT:
                pass
            elif new_status in ['ready_to_pay', 'approved', 'revision_required', 'rejected'] and tier >= WorkflowTiers.TIER_4_MANAGER:
                pass
            elif new_status == 'paid' and tier >= WorkflowTiers.TIER_5_TREASURY:
                pass
            else:
                raise PermissionDenied("شما سطح دسترسی لازم برای تغییر وضعیت فاکتور به این مرحله را ندارید.")
        serializer.save()

    @action(detail=True, methods=['post'], url_path='submit')
    def submit_invoice(self, request, pk=None):
        instance = self.get_object()
        if instance.status not in ['draft', 'revision_required']:
            return Response({'error': 'فقط فاکتورهای پیش‌نویس یا نیازمند بازنگری قابل ارسال هستند.'}, status=status.HTTP_400_BAD_REQUEST)
        from .workflow_engine import advance_workflow_step
        advance_workflow_step(instance, request.user, target_status='pending_supervisor', is_financial=True)
        return Response({'message': 'فاکتور با موفقیت به سرپرست ارسال شد.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='approve')
    def approve_invoice(self, request, pk=None):
        instance = self.get_object()
        from .workflow_engine import advance_workflow_step
        notes = request.data.get('notes')
        advance_workflow_step(instance, request.user, notes=notes, is_financial=True)
        return Response({'message': 'تایید فاکتور با موفقیت ثبت شد.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='reject')
    def reject_invoice(self, request, pk=None):
        instance = self.get_object()
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'درج دلیل رد فاکتور الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        from .workflow_engine import reject_workflow
        reject_workflow(instance, request.user, reason=reason)
        return Response({'message': 'فاکتور هزینه رد گردید.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='request-revision')
    def request_revision_invoice(self, request, pk=None):
        instance = self.get_object()
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'درج دلیل بازنگری الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        from .workflow_engine import request_workflow_revision
        request_workflow_revision(instance, request.user, reason=reason, target_status='draft')
        return Response({'message': 'فاکتور با ثبت اشکال جهت بازنگری ارجاع شد.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='pay')
    def pay_invoice(self, request, pk=None):
        instance = self.get_object()
        tracking_code = request.data.get('tracking_code') or request.data.get('payment_ref', '').strip()
        if not tracking_code:
            return Response({'error': 'درج شماره پیگیری یا شماره سند بانکی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        from .workflow_engine import execute_treasury_disbursement
        execute_treasury_disbursement(instance, request.user, tracking_code=tracking_code, notes=request.data.get('notes'))
        return Response({'message': 'تسویه فاکتور با موفقیت در خزانه‌داری ثبت گردید.', 'status': instance.status})

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی اکسل استاندارد ۲ ردیفه فاکتورهای هزینه بر اساس معماری سیستم
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        qs = self.get_queryset()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "فاکتورهای هزینه"
        ws.views.sheetView[0].rightToLeft = True

        headers_fa = [
            'ردیف', 'شماره فاکتور', 'تاریخ فاکتور (شمسی)', 'پروژه', 'بخش', 
            'طرف‌حساب', 'سرفصل هزینه', 'مبلغ (ریال)', 'وضعیت', 'ثبت‌کننده', 'شرح هزینه'
        ]
        headers_en = [
            'row_num', 'invoice_number', 'invoice_date_shamsi', 'project_name', 'section_name',
            'counterparty_name', 'category', 'amount', 'status_display', 'created_by_name', 'description'
        ]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(headers_fa) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 16
        ws.freeze_panes = 'A3'

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, inv in enumerate(qs, start=1):
            created_name = f"{inv.created_by.first_name} {inv.created_by.last_name}".strip() or inv.created_by.username if inv.created_by else 'سیستم'
            ws.append([
                idx,
                inv.invoice_number,
                inv.invoice_date_shamsi,
                inv.section.project.name if inv.section and inv.section.project else '',
                inv.section.name if inv.section else '',
                inv.counterparty.name if inv.counterparty else '',
                inv.category,
                inv.amount,
                inv.get_status_display(),
                created_name,
                inv.description or ''
            ])
            curr_row = idx + 2
            ws.row_dimensions[curr_row].height = 20
            for col_idx in range(1, len(headers_fa) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal='center' if col_idx not in [6, 11] else 'right', vertical='center')
                if col_idx in [1, 2, 3, 8]:
                    c.font = data_font_num
                else:
                    c.font = data_font
                if col_idx == 8:
                    c.number_format = '#,##0'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"expense_invoices_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# ==============================================================================
# 0.2 ویوست‌های مدیریت تنخواه‌گردان (Petty Cash Management ViewSets)
# ==============================================================================

class PettyCashAccountViewSet(viewsets.ModelViewSet):
    """
    مدیریت و تعریف حساب‌ها / سقف‌های تنخواه‌گردان کارمندان در بخش‌های پروژه
    """
    queryset = PettyCashAccount.objects.all().select_related('section', 'section__project', 'custodian')
    serializer_class = PettyCashAccountSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(section__project__company_id=cid)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(section_id=section_id)
        custodian_id = self.request.query_params.get('custodian_id')
        if custodian_id:
            qs = qs.filter(custodian_id=custodian_id)
        # کارمندان عادی فقط حساب تنخواه خود را مشاهده می‌کنند
        if not user.is_superuser:
            is_manager = UserSectionAssignment.objects.filter(
                user=user,
                section_id=section_id,
                is_active=True,
                role__in=['supervisor', 'accountant', 'manager', 'treasury']
            ).exists() if section_id else False
            if not is_manager:
                qs = qs.filter(custodian=user)
        return qs.order_by('section', 'custodian')

    def perform_create(self, serializer):
        user = self.request.user
        section = serializer.validated_data.get('section')
        if section and section.project and section.project.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if section.project.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به تعریف حساب تنخواه برای این شرکت نیستید.")
        serializer.save()


class PettyCashTransactionViewSet(viewsets.ModelViewSet):
    """
    مدیریت اسناد، هزینه‌کردها، درخواست‌های شارژ و تسویه تنخواه‌گردان
    """
    queryset = PettyCashTransaction.objects.all().select_related(
        'section', 'section__project', 'custodian', 'counterparty', 'created_by', 'account'
    )
    serializer_class = PettyCashTransactionSerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(section__project__company_id=cid)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(section_id=section_id)
        custodian_id = self.request.query_params.get('custodian_id')
        if custodian_id:
            qs = qs.filter(custodian_id=custodian_id)
        tx_type = self.request.query_params.get('transaction_type')
        if tx_type:
            qs = qs.filter(transaction_type=tx_type)
        status_val = self.request.query_params.get('status')
        if status_val:
            qs = qs.filter(status=status_val)
        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(
                Q(title__icontains=search) |
                Q(receipt_number__icontains=search) |
                Q(category__icontains=search) |
                Q(counterparty__name__icontains=search) |
                Q(description__icontains=search)
            )

        # اگر کاربر دسترسی مدیریتی نداشته باشد، فقط اسناد تنخواه خودش در بخش را می‌بیند
        if not user.is_superuser and not custodian_id:
            is_manager = UserSectionAssignment.objects.filter(
                user=user,
                section_id=section_id,
                is_active=True,
                role__in=['supervisor', 'accountant', 'manager', 'treasury']
            ).exists() if section_id else False
            if not is_manager:
                qs = qs.filter(custodian=user)

        return qs.order_by('-transaction_date_shamsi', '-created_at')

    def perform_create(self, serializer):
        user = self.request.user
        section = serializer.validated_data.get('section')
        if section and section.project and section.project.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if section.project.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به ثبت سند تنخواه برای این شرکت نیستید.")
        custodian = serializer.validated_data.get('custodian') or user

        # اتصال خودکار یا ایجاد حساب تنخواه برای بخش و کاربر
        account = serializer.validated_data.get('account')
        if not account and section:
            account, _ = PettyCashAccount.objects.get_or_create(
                section=section,
                custodian=custodian,
                defaults={'ceiling_amount': Decimal('100000000')}
            )

        instance = serializer.save(
            created_by=user,
            custodian=custodian,
            account=account,
            status='draft'
        )
        from .workflow_engine import process_creation_with_auto_pass
        process_creation_with_auto_pass(instance, user, is_financial=True)
        instance.save()

    def perform_update(self, serializer):
        user = self.request.user
        instance = self.get_object()
        new_status = serializer.validated_data.get('status')
        if new_status and new_status != instance.status and not user.is_superuser:
            from .workflow_engine import determine_user_tier, WorkflowTiers
            tier = determine_user_tier(user)
            if new_status in ['draft', 'pending_supervisor']:
                pass
            elif new_status in ['pending_accountant', 'revision_required', 'rejected'] and tier >= WorkflowTiers.TIER_2_SUPERVISOR:
                pass
            elif new_status in ['approved', 'revision_required', 'rejected'] and tier >= WorkflowTiers.TIER_3_ACCOUNTANT:
                pass
            elif new_status in ['paid'] and tier >= WorkflowTiers.TIER_5_TREASURY:
                pass
            else:
                raise PermissionDenied("شما سطح دسترسی لازم برای تغییر وضعیت سند تنخواه به این مرحله را ندارید.")
        serializer.save()

    @action(detail=True, methods=['post'], url_path='submit')
    def submit_tx(self, request, pk=None):
        instance = self.get_object()
        from .workflow_engine import advance_workflow_step
        advance_workflow_step(instance, request.user, target_status='pending_supervisor', is_financial=True)
        return Response({'message': 'تراکنش تنخواه با موفقیت به سرپرست ارسال شد.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='approve')
    def approve_tx(self, request, pk=None):
        instance = self.get_object()
        from .workflow_engine import advance_workflow_step
        advance_workflow_step(instance, request.user, notes=request.data.get('notes'), is_financial=True)
        return Response({'message': 'تایید تراکنش تنخواه با موفقیت ثبت شد.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='reject')
    def reject_tx(self, request, pk=None):
        instance = self.get_object()
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'درج دلیل رد الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        from .workflow_engine import reject_workflow
        reject_workflow(instance, request.user, reason=reason)
        return Response({'message': 'تراکنش تنخواه رد گردید.', 'status': instance.status})

    @action(detail=False, methods=['get'], url_path='my-balance')
    def my_balance(self, request):
        """
        محاسبه مانده زنده تنخواه، سقف مصوب، هزینه‌های تاییدشده و در جریان (Guardian G1)
        """
        section_id = request.query_params.get('section_id')
        if not section_id:
            return Response(
                {'error': 'شناسه بخش پروژه (section_id) الزامی است.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        target_user_id = request.query_params.get('custodian_id')
        if target_user_id and (request.user.is_superuser or UserSectionAssignment.objects.filter(
            user=request.user, section_id=section_id, is_active=True, role__in=['supervisor', 'accountant', 'manager', 'treasury']
        ).exists()):
            custodian = get_user_model().objects.filter(id=target_user_id).first() or request.user
        else:
            custodian = request.user

        account = PettyCashAccount.objects.filter(
            section_id=section_id, custodian=custodian, is_active=True
        ).first()

        ceiling = account.ceiling_amount if account else Decimal('0')

        # مجموع شارژهای دریافت شده (approved / paid)
        allocations_sum = PettyCashTransaction.objects.filter(
            section_id=section_id,
            custodian=custodian,
            transaction_type='allocation',
            status__in=['approved', 'paid']
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0')

        # مجموع هزینه‌های تایید شده و تسویه شده
        spent_sum = PettyCashTransaction.objects.filter(
            section_id=section_id,
            custodian=custodian,
            transaction_type='expense',
            status='approved'
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0')

        # مجموع هزینه‌های در جریان تایید (draft / pending)
        pending_settlement = PettyCashTransaction.objects.filter(
            section_id=section_id,
            custodian=custodian,
            transaction_type='expense',
            status__in=['draft', 'pending_supervisor', 'pending_accountant']
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0')

        # مانده نقد در دست تنخواه‌دار
        current_balance = allocations_sum - spent_sum

        # درصد مصرف از سقف تنخواه
        utilization_rate = 0.0
        if ceiling > Decimal('0'):
            utilization_rate = float(min(Decimal('100.0'), (spent_sum / ceiling) * Decimal('100.0')))

        custodian_full_name = f"{custodian.first_name} {custodian.last_name}".strip() or custodian.username

        return Response({
            'account_id': account.id if account else None,
            'ceiling_amount': int(ceiling),
            'ceiling_amount_tomans': int(ceiling // 10),
            'total_allocated': int(allocations_sum),
            'total_allocated_tomans': int(allocations_sum // 10),
            'total_spent': int(spent_sum),
            'total_spent_tomans': int(spent_sum // 10),
            'current_balance': int(current_balance),
            'current_balance_tomans': int(current_balance // 10),
            'pending_settlement': int(pending_settlement),
            'pending_settlement_tomans': int(pending_settlement // 10),
            'utilization_rate': round(utilization_rate, 1),
            'card_or_account_number': account.card_or_account_number if account else '',
            'sheba_number': account.sheba_number if account else '',
            'custodian_name': custodian_full_name,
            'custodian_id': custodian.id
        })

    @action(detail=False, methods=['post'], url_path='request-replenishment')
    def request_replenishment(self, request):
        """
        ثبت سریع درخواست شارژ مجدد / واریز به حساب تنخواه
        """
        section_id = request.data.get('section_id')
        amount = request.data.get('amount')
        title = request.data.get('title') or 'درخواست شارژ مجدد تنخواه'
        description = request.data.get('description', '')
        receipt_number = request.data.get('receipt_number', '')

        if not section_id or not amount:
            return Response(
                {'error': 'شناسه بخش و مبلغ الزامی هستند.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            amount_dec = Decimal(str(amount))
            if amount_dec <= 0:
                raise ValueError()
        except Exception:
            return Response({'error': 'مبلغ باید عددی مثبت باشد.'}, status=status.HTTP_400_BAD_REQUEST)

        section = get_object_or_404(ProjectSection, id=section_id)
        custodian = request.user

        account, _ = PettyCashAccount.objects.get_or_create(
            section=section,
            custodian=custodian,
            defaults={'ceiling_amount': Decimal('100000000')}
        )

        from common.date_utils import gregorian_to_jalali
        now = timezone.now()
        shamsi_today = gregorian_to_jalali(now.year, now.month, now.day)

        tx = PettyCashTransaction.objects.create(
            account=account,
            section=section,
            custodian=custodian,
            created_by=request.user,
            transaction_type='allocation',
            amount=amount_dec,
            transaction_date_shamsi=shamsi_today,
            title=title,
            category='شارژ مجدد تنخواه',
            description=description,
            receipt_number=receipt_number,
            status='draft'
        )

        serializer = self.get_serializer(tx)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی اکسل رسمی ۲ ردیفه گردش و تراکنش‌های تنخواه‌گردان مطابق استانداردهای سیستم
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        qs = self.get_queryset()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "گردش تنخواه‌گردان"
        ws.views.sheetView[0].rightToLeft = True

        headers_fa = [
            'ردیف', 'نوع تراکنش', 'شماره رسید/فاکتور', 'تاریخ (شمسی)', 'پروژه', 'بخش',
            'تنخواه‌دار', 'طرف‌حساب/فروشنده', 'سرفصل هزینه', 'عنوان / شرح مختصر',
            'مبلغ (ریال)', 'وضعیت سند', 'ثبت‌کننده', 'توضیحات تفصیلی'
        ]
        headers_en = [
            'row_num', 'transaction_type_display', 'receipt_number', 'transaction_date_shamsi',
            'project_name', 'section_name', 'custodian_name', 'counterparty_name',
            'category', 'title', 'amount', 'status_display', 'created_by_name', 'description'
        ]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(headers_fa) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.freeze_panes = 'A3'

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, tx in enumerate(qs, start=1):
            curr_row = idx + 2
            cust_name = f"{tx.custodian.first_name} {tx.custodian.last_name}".strip() or tx.custodian.username
            cp_name = tx.counterparty.name if tx.counterparty else '-'
            creator_name = f"{tx.created_by.first_name} {tx.created_by.last_name}".strip() if tx.created_by else '-'
            row_data = [
                idx,
                tx.get_transaction_type_display(),
                tx.receipt_number or '-',
                tx.transaction_date_shamsi,
                tx.section.project.name if tx.section else '-',
                tx.section.name if tx.section else '-',
                cust_name,
                cp_name,
                tx.category,
                tx.title,
                float(tx.amount),
                tx.get_status_display(),
                creator_name,
                tx.description or '-'
            ]
            ws.append(row_data)

            for col_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal='center' if col_idx not in [8, 10, 14] else 'right', vertical='center')
                if col_idx in [1, 3, 4, 11]:
                    c.font = data_font_num
                else:
                    c.font = data_font
                if col_idx == 11:
                    c.number_format = '#,##0'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"petty_cash_transactions_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response








class WorkflowAuditLogViewSet(viewsets.ReadOnlyModelViewSet):
    """
    مشاهده تاریخچه ممیزی گردش کار و رویدادهای تایید/رد با ایزولاسیون شرکتی
    """
    queryset = WorkflowAuditLog.objects.all().select_related('actor')
    serializer_class = WorkflowAuditLogSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    # برچسب‌های قدیمی موتور گردش کار ('personnelprofile') هم لاگ پرونده شمرده می‌شوند
    PERSONNEL_LOG_TYPES = ('personnel', 'personnelprofile')
    # مسیر رسیدن به شرکت برای مدل‌های دیگری که موتور گردش کار برایشان لاگ می‌نویسد
    COMPANY_PATH_CANDIDATES = (
        ('company', 'company_id'),
        ('section', 'section__project__company_id'),
        ('project', 'project__company_id'),
        ('personnel', 'personnel__company_id'),
        ('vehicle', 'vehicle__company_id'),
        ('period', 'period__company_id'),
    )

    def _company_scope_q(self, base_qs, allowed_ids):
        """
        DEF-19: محدودیت قطعی لاگ‌ها به شرکت‌های مجاز کاربر، با یا بدون سربرگ شرکت.
        لاگ پروندهٔ حذف‌شده یا Purgeشده با metadata.company_id سنجیده می‌شود.
        """
        personnel_ids = PersonnelProfile.objects.filter(
            Q(company_id__in=allowed_ids) | Q(section__project__company_id__in=allowed_ids)
        ).values('id')
        scope = Q(content_type__in=self.PERSONNEL_LOG_TYPES) & (
            Q(object_id__in=personnel_ids) | Q(metadata__company_id__in=allowed_ids)
        )
        scope |= Q(
            content_type='personnelchangerequest',
            object_id__in=PersonnelChangeRequest.objects.filter(personnel_id__in=personnel_ids).values('id'),
        )

        other_types = (
            base_qs.exclude(content_type__in=self.PERSONNEL_LOG_TYPES + ('personnelchangerequest',))
            .values_list('content_type', flat=True).distinct()
        )
        personnel_models = {m.__name__.lower(): m for m in apps.get_app_config('personnel').get_models()}
        for ct in other_types:
            model = personnel_models.get(ct)
            if model is None:
                continue
            field_names = {f.name for f in model._meta.get_fields()}
            path = next((p for name, p in self.COMPANY_PATH_CANDIDATES if name in field_names), None)
            if path is None:
                continue
            scope |= Q(content_type=ct, object_id__in=model.objects.filter(**{f'{path}__in': allowed_ids}).values('pk'))
        return scope

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))

        content_type = (self.request.query_params.get('content_type') or '').strip().lower()
        object_id = self.request.query_params.get('object_id')
        oid = None
        if object_id:
            try:
                oid = int(object_id)
            except (ValueError, TypeError):
                return qs.none()

        if content_type in self.PERSONNEL_LOG_TYPES:
            qs = qs.filter(content_type__in=self.PERSONNEL_LOG_TYPES)
            if oid is not None:
                qs = qs.filter(object_id=oid)
        elif content_type == 'personnelchangerequest':
            # Diff Viewer: لاگ اکشن‌های درخواست تغییرات روی تایملاین پرونده و با change_request_id ثبت می‌شود
            if oid is not None:
                qs = qs.filter(
                    Q(content_type='personnelchangerequest', object_id=oid) |
                    Q(content_type__in=self.PERSONNEL_LOG_TYPES, metadata__change_request_id=oid)
                )
            else:
                qs = qs.filter(
                    Q(content_type='personnelchangerequest') |
                    Q(content_type__in=self.PERSONNEL_LOG_TYPES, metadata__has_key='change_request_id')
                )
        elif content_type:
            qs = qs.filter(content_type=content_type)
            if oid is not None:
                qs = qs.filter(object_id=oid)
        elif oid is not None:
            qs = qs.filter(object_id=oid)

        # ایزولاسیون شرکتی همیشگی (DEF-19): بدون سربرگ هم به شرکت‌های مجاز کاربر محدود می‌شود
        if not user.is_superuser:
            allowed_ids = [cid] if cid else list(get_user_allowed_companies(user).values_list('id', flat=True))
            if not allowed_ids:
                return qs.none()
            qs = qs.filter(self._company_scope_q(qs, allowed_ids))

        return qs.order_by('-created_at')
