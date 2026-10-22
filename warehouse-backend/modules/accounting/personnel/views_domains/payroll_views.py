# -*- coding: utf-8 -*-
"""
ViewSets for payroll yearly settings, monthly payroll calculation, and fleet settlements.
"""

import io
import zipfile
import openpyxl
from decimal import Decimal
from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied
from django.http import HttpResponse
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from django.apps import apps
from companies.access import (
    get_request_company_id,
    validate_user_company_access,
)
from ..models import (
    PayrollYearlySettings,
    JobGradeTier,
    WorkshopInsuranceSettings,
    MonthlyWorkPeriod,
    MonthlyPayrollRecord,
    Project,
)
from ..serializers import (
    PayrollYearlySettingsSerializer,
    MonthlyPayrollRecordSerializer,
)
from ..permissions import (
    IsPayrollSettingsManagerOrReadOnly,
)
from ..payroll_engine import (
    calculate_monthly_payroll_for_period,
    get_effective_payroll_settings,
)
from ..dbf_generator import generate_dskkar_bytes, generate_dskwor_bytes
from ..tax_bank_exporter import (
    generate_wh_tax_content,
    generate_wp_tax_content,
    generate_bank_meli_excel,
)
from ..payroll_excel_exporter import generate_monthly_payroll_excel
from ..fleet_settlement_engine import (
    calculate_monthly_fleet_settlement,
    generate_fleet_bank_meli_excel,
)
from .scope_helpers import (
    _copy_payroll_submodels,
    safe_int,
)

class PayrollYearlySettingsViewSet(viewsets.ModelViewSet):
    queryset = PayrollYearlySettings.objects.all().prefetch_related(
        'job_grades', 'workshop_insurance', 'tax_settings', 'bank_export_settings', 'project'
    )
    serializer_class = PayrollYearlySettingsSerializer
    permission_classes = [IsPayrollSettingsManagerOrReadOnly]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        cid = validate_user_company_access(self.request.user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(Q(company_id=cid) | Q(project__company_id=cid))
        p_id = safe_int(self.request.query_params.get('project_id'))
        if p_id is not None:
            return qs.filter(project_id=p_id)
        is_global = self.request.query_params.get('is_global')
        if is_global == 'true':
            return qs.filter(project__isnull=True)
        return qs

    @action(detail=False, methods=['get', 'post'], url_path='active-or-year')
    def get_active_or_year(self, request):
        if request.method == 'POST':
            pk = request.data.get('id')
            if pk:
                settings_obj = PayrollYearlySettings.objects.filter(id=pk).first()
                if settings_obj:
                    return self.update_all_tabs(request, pk=pk)

        year = request.query_params.get('year', '1405')
        p_id = safe_int(request.query_params.get('project_id'))
        v_id = safe_int(request.query_params.get('version_id'))
        effective_from = request.query_params.get('effective_from')
        if effective_from in ('null', 'undefined', ''):
            effective_from = None

        settings_obj = None
        if v_id:
            candidate = PayrollYearlySettings.objects.filter(id=v_id).first()
            if candidate:
                if p_id is not None:
                    if candidate.project_id == p_id:
                        settings_obj = candidate
                    else:
                        # اگر کاندید برای سراسری است و پروژه تنظیمی ندارد، همان نسخه کاندید را نمایش بده
                        settings_obj = candidate
                else:
                    settings_obj = candidate

        if not settings_obj:
            if effective_from:
                settings_obj = get_effective_payroll_settings(effective_from, p_id)
            else:
                settings_obj = get_effective_payroll_settings(f"{year}/01", p_id)

        if not settings_obj:
            settings_obj = PayrollYearlySettings.objects.filter(fiscal_year=year).first() or PayrollYearlySettings.objects.first()
        if not settings_obj:
            return Response({'error': 'تنظیماتی یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
        return Response(PayrollYearlySettingsSerializer(settings_obj).data)

    @action(detail=False, methods=['get'], url_path='versions')
    def get_versions(self, request):
        """
        دریافت فهرست تمام نسخه‌های احکام در یک سال مالی به تفکیک پروژه
        """
        year = request.query_params.get('year', '1405')
        p_id = safe_int(request.query_params.get('project_id'))
        qs = PayrollYearlySettings.objects.filter(fiscal_year=year)
        if p_id is not None:
            project_qs = qs.filter(project_id=p_id).order_by('effective_from')
            if project_qs.exists():
                return Response(PayrollYearlySettingsSerializer(project_qs, many=True).data)
            # اگر برای پروژه تنظیمات اختصاصی ثبت نشده، نسخه‌های سراسری سازمان بازگردانده می‌شود
            global_qs = qs.filter(project__isnull=True).order_by('effective_from')
            return Response(PayrollYearlySettingsSerializer(global_qs, many=True).data)
        else:
            qs = qs.filter(project__isnull=True).order_by('effective_from')
        return Response(PayrollYearlySettingsSerializer(qs, many=True).data)

    @action(detail=False, methods=['get'], url_path='available-years')
    def get_available_years(self, request):
        """
        دریافت فهرست سال‌های مالی تعریف‌شده در سیستم
        """
        years = list(PayrollYearlySettings.objects.values_list('fiscal_year', flat=True).distinct())
        cleaned_years = sorted(list(set([str(y).strip() for y in years if y and str(y).strip()])), reverse=True)
        if not cleaned_years:
            cleaned_years = ['1405']
        return Response({'years': cleaned_years})

    @action(detail=False, methods=['post'], url_path='create-year')
    def create_fiscal_year(self, request):
        """
        ایجاد سال مالی جدید بر اساس کپی اطلاعات و احکام یک سال مالی مبدأ
        """
        new_year = request.data.get('year')
        source_year = request.data.get('source_year')
        project_id = request.data.get('project_id')
        p_id = safe_int(project_id)

        if not new_year:
            return Response({'error': 'تعیین سال مالی جدید الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        new_year = str(new_year).strip()
        if len(new_year) != 4 or not new_year.isdigit():
            return Response({'error': 'فرمت سال مالی باید یک عدد ۴ رقمی شمسی (مثلاً ۱۴۰۶) باشد.'}, status=status.HTTP_400_BAD_REQUEST)

        # بررسی وجود قبلی در دامنه انتخابی
        existing = PayrollYearlySettings.objects.filter(fiscal_year=new_year, project_id=p_id).first()
        if existing:
            return Response({
                'error': f'سال مالی {new_year} قبلاً در این دامنه تعریف شده است.',
                'settings': PayrollYearlySettingsSerializer(existing).data
            }, status=status.HTTP_400_BAD_REQUEST)

        project = Project.objects.filter(id=p_id).first() if p_id else None

        # پیدا کردن سال مبدأ
        source = None
        if source_year:
            source = get_effective_payroll_settings(f"{source_year}/01", p_id)
        if not source:
            source = get_effective_payroll_settings("1405/01", p_id) or PayrollYearlySettings.objects.first()

        effective_from = f"{new_year}/01"
        version_title = "احکام مصوب فروردین"

        with transaction.atomic():
            # اگر پروژه انتخاب شده باشد، ابتدا مطمئن می‌شویم رکورد پایه سراسری سازمان برای سال جدید وجود دارد
            if project and not PayrollYearlySettings.objects.filter(fiscal_year=new_year, project__isnull=True).exists():
                global_source = get_effective_payroll_settings(f"{source_year or '1405'}/01", None) or source
                global_settings = PayrollYearlySettings.objects.create(
                    fiscal_year=new_year,
                    effective_from=effective_from,
                    version_title=version_title,
                    project=None,
                    title=f"تنظیمات سال مالی {new_year} (سراسری سازمان)",
                    is_active=True,
                    monthly_food_allowance=global_source.monthly_food_allowance if global_source else 22000000,
                    monthly_housing_allowance=global_source.monthly_housing_allowance if global_source else 30000000,
                    monthly_spouse_allowance=global_source.monthly_spouse_allowance if global_source else 5000000,
                    monthly_child_allowance=global_source.monthly_child_allowance if global_source else 16625549,
                    shift_percent=global_source.shift_percent if global_source else 0,
                    transport_help_percent=global_source.transport_help_percent if global_source else 0,
                    transport_fixed_amount=global_source.transport_fixed_amount if global_source else 2388728,
                    specialist_attraction_percent=global_source.specialist_attraction_percent if global_source else 0,
                    bad_weather_percent=global_source.bad_weather_percent if global_source else 0,
                    remote_hardship_percent=global_source.remote_hardship_percent if global_source else 0,
                    south_pars_percent=global_source.south_pars_percent if global_source else 0,
                    travel_cost_per_day=global_source.travel_cost_per_day if global_source else 8736600,
                    worker_insurance_rate=global_source.worker_insurance_rate if global_source else 7.00,
                    employer_insurance_rate=global_source.employer_insurance_rate if global_source else 20.00,
                    unemployment_insurance_rate=global_source.unemployment_insurance_rate if global_source else 3.00,
                    surplus_overtime_percent=global_source.surplus_overtime_percent if global_source else 50.00,
                    standard_daily_hours=getattr(global_source, 'standard_daily_hours', 10.00) or 10.00,
                    bonus_daily_coefficient=getattr(global_source, 'bonus_daily_coefficient', 5.00) or 5.00,
                    seniority_monthly_coefficient=getattr(global_source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                    overtime_rate_multiplier=getattr(global_source, 'overtime_rate_multiplier', 1.40) or 1.40,
                    friday_work_rate_multiplier=getattr(global_source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                    max_insurable_daily_wage=getattr(global_source, 'max_insurable_daily_wage', None),
                    attendance_edit_past_days=global_source.attendance_edit_past_days if global_source else 3,
                    attendance_edit_future_days=global_source.attendance_edit_future_days if global_source else 0,
                    created_by=request.user,
                    updated_by=request.user
                )
                _copy_payroll_submodels(global_source, global_settings)

            new_settings = PayrollYearlySettings.objects.create(
                fiscal_year=new_year,
                effective_from=effective_from,
                version_title=version_title,
                project=project,
                title=f"تنظیمات سال مالی {new_year} ({project.name if project else 'سراسری سازمان'})",
                is_active=True,
                monthly_food_allowance=source.monthly_food_allowance if source else 22000000,
                monthly_housing_allowance=source.monthly_housing_allowance if source else 30000000,
                monthly_spouse_allowance=source.monthly_spouse_allowance if source else 5000000,
                monthly_child_allowance=source.monthly_child_allowance if source else 16625549,
                shift_percent=source.shift_percent if source else 0,
                transport_help_percent=source.transport_help_percent if source else 0,
                transport_fixed_amount=source.transport_fixed_amount if source else 2388728,
                specialist_attraction_percent=source.specialist_attraction_percent if source else 0,
                bad_weather_percent=source.bad_weather_percent if source else 0,
                remote_hardship_percent=source.remote_hardship_percent if source else 0,
                south_pars_percent=source.south_pars_percent if source else 0,
                travel_cost_per_day=source.travel_cost_per_day if source else 8736600,
                worker_insurance_rate=source.worker_insurance_rate if source else 7.00,
                employer_insurance_rate=source.employer_insurance_rate if source else 20.00,
                unemployment_insurance_rate=source.unemployment_insurance_rate if source else 3.00,
                surplus_overtime_percent=source.surplus_overtime_percent if source else 50.00,
                standard_daily_hours=getattr(source, 'standard_daily_hours', 10.00) or 10.00,
                bonus_daily_coefficient=getattr(source, 'bonus_daily_coefficient', 5.00) or 5.00,
                seniority_monthly_coefficient=getattr(source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                overtime_rate_multiplier=getattr(source, 'overtime_rate_multiplier', 1.40) or 1.40,
                friday_work_rate_multiplier=getattr(source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                max_insurable_daily_wage=getattr(source, 'max_insurable_daily_wage', None),
                attendance_edit_past_days=source.attendance_edit_past_days if source else 3,
                attendance_edit_future_days=source.attendance_edit_future_days if source else 0,
                created_by=request.user,
                updated_by=request.user
            )
            _copy_payroll_submodels(source, new_settings)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='CREATE',
                    severity='warning',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(new_settings.id),
                    target_repr=str(new_settings),
                    details={
                        'action': 'create_fiscal_year',
                        'fiscal_year': new_year,
                        'source_year': source_year,
                        'project_id': p_id
                    }
                )
            except Exception:
                pass

        return Response(PayrollYearlySettingsSerializer(new_settings).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='create-version')
    def create_version(self, request):
        """
        ایجاد نسخه جدید احکام در میانه سال با کپی از آخرین نسخه و تعیین ماه شروع
        """
        import re
        project_id = request.data.get('project_id')
        year = request.data.get('year', '1405')
        effective_from = request.data.get('effective_from')
        version_title = request.data.get('version_title') or f"احکام اصلاحیه {effective_from}"

        if not effective_from:
            return Response({'error': 'ماه شروع اعتبار (effective_from) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        effective_from = str(effective_from).strip()
        if not re.match(r'^\d{4}/(0[1-9]|1[0-2])$', effective_from):
            return Response({'error': 'فرمت ماه شروع اعتبار باید YYYY/MM باشد (مثال: 1405/07).'}, status=status.HTTP_400_BAD_REQUEST)

        p_id = safe_int(project_id)
        project = Project.objects.filter(id=p_id).first() if p_id else None

        existing = PayrollYearlySettings.objects.filter(
            fiscal_year=year,
            project=project,
            effective_from=effective_from
        ).first()
        if existing:
            return Response({
                'error': f'نسخه احکام با تاریخ شروع {effective_from} قبلاً در این دامنه ایجاد شده است.',
                'settings': PayrollYearlySettingsSerializer(existing).data
            }, status=status.HTTP_400_BAD_REQUEST)

        source = get_effective_payroll_settings(effective_from, p_id)
        if not source:
            source = PayrollYearlySettings.objects.first()

        with transaction.atomic():
            # به‌روزرسانی خودکار ماه پایان نسخه قبلی
            year_str, month_str = effective_from.split('/')
            m_int = int(month_str)
            if m_int > 1:
                prior_effective_to = f"{year_str}/{m_int - 1:02d}"
                PayrollYearlySettings.objects.filter(
                    fiscal_year=year,
                    project=project,
                    effective_from__lt=effective_from
                ).filter(
                    Q(effective_to__isnull=True) | Q(effective_to__gte=effective_from)
                ).update(effective_to=prior_effective_to)

            new_settings = PayrollYearlySettings.objects.create(
                fiscal_year=year,
                effective_from=effective_from,
                version_title=version_title,
                project=project,
                title=f"{version_title} ({project.name if project else 'سراسری سازمان'})",
                is_active=True,
                monthly_food_allowance=source.monthly_food_allowance if source else 22000000,
                monthly_housing_allowance=source.monthly_housing_allowance if source else 30000000,
                monthly_spouse_allowance=source.monthly_spouse_allowance if source else 5000000,
                monthly_child_allowance=source.monthly_child_allowance if source else 16625549,
                shift_percent=source.shift_percent if source else 0,
                transport_help_percent=source.transport_help_percent if source else 0,
                transport_fixed_amount=source.transport_fixed_amount if source else 2388728,
                specialist_attraction_percent=source.specialist_attraction_percent if source else 0,
                bad_weather_percent=source.bad_weather_percent if source else 0,
                remote_hardship_percent=source.remote_hardship_percent if source else 0,
                south_pars_percent=source.south_pars_percent if source else 0,
                travel_cost_per_day=source.travel_cost_per_day if source else 8736600,
                worker_insurance_rate=source.worker_insurance_rate if source else 7.00,
                employer_insurance_rate=source.employer_insurance_rate if source else 20.00,
                unemployment_insurance_rate=source.unemployment_insurance_rate if source else 3.00,
                surplus_overtime_percent=source.surplus_overtime_percent if source else 50.00,
                standard_daily_hours=getattr(source, 'standard_daily_hours', 10.00) or 10.00,
                bonus_daily_coefficient=getattr(source, 'bonus_daily_coefficient', 5.00) or 5.00,
                seniority_monthly_coefficient=getattr(source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                overtime_rate_multiplier=getattr(source, 'overtime_rate_multiplier', 1.40) or 1.40,
                friday_work_rate_multiplier=getattr(source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                max_insurable_daily_wage=getattr(source, 'max_insurable_daily_wage', None),
                attendance_edit_past_days=source.attendance_edit_past_days if source else 3,
                attendance_edit_future_days=source.attendance_edit_future_days if source else 0,
                created_by=request.user,
                updated_by=request.user
            )
            _copy_payroll_submodels(source, new_settings)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='CREATE',
                    severity='warning',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(new_settings.id),
                    target_repr=str(new_settings),
                    details={
                        'action': 'create_version',
                        'fiscal_year': year,
                        'effective_from': effective_from,
                        'version_title': version_title,
                        'project_id': p_id
                    }
                )
            except Exception:
                pass

        return Response(PayrollYearlySettingsSerializer(new_settings).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='clone-for-project')
    def clone_for_project(self, request):
        """
        کپی هوشمند و سریع تنظیمات پایه سراسری برای یک پروژه مشخص
        """
        project_id = request.data.get('project_id')
        year = request.data.get('year', '1405')
        effective_from = request.data.get('effective_from', f"{year}/01")

        p_id = safe_int(project_id)
        if not p_id:
            return Response({'error': 'شناسه پروژه (project_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        project = Project.objects.filter(id=p_id).first()
        if not project:
            return Response({'error': 'پروژه مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        existing = PayrollYearlySettings.objects.filter(fiscal_year=year, project=project, effective_from=effective_from).first()
        if existing:
            return Response(PayrollYearlySettingsSerializer(existing).data)

        # پیدا کردن تنظیمات مادر (سراسری)
        source = get_effective_payroll_settings(effective_from, None)
        if not source:
            source = PayrollYearlySettings.objects.first()
        if not source:
            return Response({'error': 'تنظیمات مادر سراسری برای کپی یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        with transaction.atomic():
            new_settings = PayrollYearlySettings.objects.create(
                fiscal_year=year,
                effective_from=effective_from,
                version_title=source.version_title or "احکام مصوب فروردین",
                project=project,
                title=f"تنظیمات {project.name} - سال {year}",
                is_active=True,
                monthly_food_allowance=source.monthly_food_allowance,
                monthly_housing_allowance=source.monthly_housing_allowance,
                monthly_spouse_allowance=source.monthly_spouse_allowance,
                monthly_child_allowance=source.monthly_child_allowance,
                shift_percent=source.shift_percent,
                transport_help_percent=source.transport_help_percent,
                transport_fixed_amount=source.transport_fixed_amount,
                specialist_attraction_percent=source.specialist_attraction_percent,
                bad_weather_percent=source.bad_weather_percent,
                remote_hardship_percent=source.remote_hardship_percent,
                south_pars_percent=source.south_pars_percent,
                travel_cost_per_day=source.travel_cost_per_day,
                worker_insurance_rate=source.worker_insurance_rate,
                employer_insurance_rate=source.employer_insurance_rate,
                unemployment_insurance_rate=source.unemployment_insurance_rate,
                surplus_overtime_percent=source.surplus_overtime_percent,
                standard_daily_hours=getattr(source, 'standard_daily_hours', 10.00) or 10.00,
                bonus_daily_coefficient=getattr(source, 'bonus_daily_coefficient', 5.00) or 5.00,
                seniority_monthly_coefficient=getattr(source, 'seniority_monthly_coefficient', 2.50) or 2.50,
                overtime_rate_multiplier=getattr(source, 'overtime_rate_multiplier', 1.40) or 1.40,
                friday_work_rate_multiplier=getattr(source, 'friday_work_rate_multiplier', 0.40) or 0.40,
                max_insurable_daily_wage=getattr(source, 'max_insurable_daily_wage', None),
                attendance_edit_past_days=source.attendance_edit_past_days,
                attendance_edit_future_days=source.attendance_edit_future_days,
                created_by=request.user,
                updated_by=request.user
            )
            _copy_payroll_submodels(source, new_settings)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='CREATE',
                    severity='info',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(new_settings.id),
                    target_repr=str(new_settings),
                    details={
                        'action': 'clone_for_project',
                        'fiscal_year': year,
                        'effective_from': effective_from,
                        'project_id': p_id,
                        'project_name': project.name
                    }
                )
            except Exception:
                pass

        return Response(PayrollYearlySettingsSerializer(new_settings).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='revert-to-global')
    def revert_to_global(self, request, pk=None):
        """
        حذف تنظیمات اختصاصی پروژه و بازگشت به ارث‌بری از تنظیمات سراسری سازمان
        """
        settings_obj = self.get_object()
        if not settings_obj.project:
            return Response({'error': 'تنظیمات سراسری سازمان قابل حذف یا بازگشت به ارث‌بری نیست.'}, status=status.HTTP_400_BAD_REQUEST)

        project_name = settings_obj.project.name
        project_id = settings_obj.project.id
        fiscal_year = settings_obj.fiscal_year

        with transaction.atomic():
            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='DELETE',
                    severity='warning',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(settings_obj.id),
                    target_repr=str(settings_obj),
                    details={
                        'action': 'revert_to_global',
                        'fiscal_year': fiscal_year,
                        'project_id': project_id,
                        'project_name': project_name
                    }
                )
            except Exception:
                pass

            settings_obj.delete()

        return Response({
            'message': f'تنظیمات اختصاصی پروژه «{project_name}» با موفقیت حذف شد و به تنظیمات سراسری سازمان بازگشت.',
            'fiscal_year': fiscal_year
        })

    @action(detail=False, methods=['get'], url_path='job-grade-rate')
    def get_job_grade_rate(self, request):
        """
        دریافت خودکار مزد روزانه و پایه سنواتی بر اساس گروه شغلی (۱ تا ۲۰)
        """
        grade = safe_int(request.query_params.get('grade'), default=19)
        year = request.query_params.get('year', '1405')
        p_id = safe_int(request.query_params.get('project_id'))

        settings_obj = get_effective_payroll_settings(f"{year}/01", p_id)

        tier = None
        if settings_obj:
            tier = JobGradeTier.objects.filter(
                yearly_settings=settings_obj,
                grade_number=grade
            ).first()

        if not tier:
            tier = JobGradeTier.objects.filter(grade_number=grade).first()

        if tier:
            std_h = float(getattr(settings_obj, 'standard_daily_hours', 10.00) or 10.00)
            return Response({
                'grade_number': tier.grade_number,
                'daily_base_wage': float(tier.daily_base_wage),
                'daily_seniority_bonus': float(tier.daily_seniority_bonus),
                'hourly_rate': round(float(tier.daily_base_wage) / (std_h if std_h > 0 else 10.0), 2)
            })
        return Response({'error': 'گروه شغلی در جدول تنظیمات یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

    @action(detail=True, methods=['post'], url_path='update-all')
    def update_all_tabs(self, request, pk=None):
        """
        به‌روزرسانی یکپارچه هر ۵ تب تنظیمات با استفاده از سریالایزر جامع Writable Nested
        """
        settings_obj = self.get_object()
        data = request.data

        serializer = PayrollYearlySettingsSerializer(settings_obj, data=data, partial=True)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            updated_instance = serializer.save(updated_by=request.user)

            try:
                AuditLog = apps.get_model('accounts', 'AuditLog')
                AuditLog.objects.create(
                    user=request.user,
                    actor_username=request.user.username,
                    actor_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                    module='settings',
                    action='UPDATE',
                    severity='info',
                    target_model='PayrollYearlySettings',
                    target_object_id=str(updated_instance.id),
                    target_repr=str(updated_instance),
                    details={
                        'action': 'update_all_tabs',
                        'fiscal_year': updated_instance.fiscal_year,
                        'version_title': updated_instance.version_title,
                        'project_id': updated_instance.project_id
                    }
                )
            except Exception:
                pass

        return Response({
            'message': 'تنظیمات پایه حقوق و دستمزد با موفقیت به‌روزرسانی شد.',
            'settings': PayrollYearlySettingsSerializer(updated_instance).data
        })

# ══════════════════════════════════════════════════════════════════════════════
# ── ۸. ویوست محاسبه ماهانه حقوق و صدور دیسکت‌ها و گزارشات ۵۸ ستونی ────────────
# ══════════════════════════════════════════════════════════════════════════════

class MonthlyPayrollViewSet(viewsets.ModelViewSet):
    queryset = MonthlyPayrollRecord.objects.all().select_related('period', 'personnel')
    serializer_class = MonthlyPayrollRecordSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        from common.tenant_scope import scope_tenant_queryset
        qs = scope_tenant_queryset(qs, user, company_field='personnel__company_id', request=self.request)

        period_id = self.request.query_params.get('period_id')
        if period_id:
            qs = qs.filter(period_id=period_id)
        
        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(period__warehouse_id=warehouse_id)
            
        year_month = self.request.query_params.get('year_month')
        if year_month:
            qs = qs.filter(period__year_month=year_month)

        return qs.order_by('row_number')

    def update(self, request, *args, **kwargs):
        from ..lifecycle_services import user_can_view_wages
        if not user_can_view_wages(request.user):
            raise PermissionDenied("ویرایش فیش حقوقی تنها با دسترسی مالی و حسابداری مجاز است.")
        if 'payment_hold' in request.data or 'payment_hold_reason' in request.data:
            return Response(
                {'error': 'تغییر مستقیم وضعیت تعلیق پرداخت مجاز نیست. آزادسازی پرداخت منحصراً از طریق اکشن آزادسازی پس از تمدید اعتبار امکان‌پذیر است.'},
                status=status.HTTP_400_BAD_REQUEST
            )
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        from ..lifecycle_services import user_can_view_wages
        if not user_can_view_wages(request.user):
            raise PermissionDenied("ویرایش فیش حقوقی تنها با دسترسی مالی و حسابداری مجاز است.")
        if 'payment_hold' in request.data or 'payment_hold_reason' in request.data:
            return Response(
                {'error': 'تغییر مستقیم وضعیت تعلیق پرداخت مجاز نیست. آزادسازی پرداخت منحصراً از طریق اکشن آزادسازی پس از تمدید اعتبار امکان‌پذیر است.'},
                status=status.HTTP_400_BAD_REQUEST
            )
        return super().partial_update(request, *args, **kwargs)

    @action(detail=False, methods=['post'], url_path='calculate-period')
    def calculate_period(self, request):
        """
        تجمیع کارکرد و محاسبه خودکار ۵۸ ستون حقوق برای دوره انتخابی (با پشتیبانی از سطح شرکت یا انبار)
        """
        warehouse_id = request.data.get('warehouse_id')
        year_month = request.data.get('year_month')
        from common.tenant_scope import get_current_tenant_company_id
        company_id = request.data.get('company_id') or get_current_tenant_company_id(request=request, user=request.user)
        
        if not year_month:
            return Response({'error': 'سال/ماه الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not warehouse_id and not company_id:
            return Response({'error': 'شناسه انبار یا شرکت فعال الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        wh_val = int(warehouse_id) if (warehouse_id and str(warehouse_id).isdigit()) else None
        comp_val = int(company_id) if (company_id and str(company_id).isdigit()) else None

        lookup = {'year_month': str(year_month).strip()}
        if wh_val:
            lookup['warehouse_id'] = wh_val
        if comp_val:
            lookup['company_id'] = comp_val

        # Get or create work period
        period, _ = MonthlyWorkPeriod.objects.get_or_create(
            **lookup,
            defaults={'status': 'OPEN', 'company_id': comp_val, 'warehouse_id': wh_val}
        )

        if period.status == 'LOCKED':
            return Response({'error': f'دوره {period.year_month} قفل شده است و امکان محاسبه مجدد ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        records = calculate_monthly_payroll_for_period(period, user=request.user)
        serialized = MonthlyPayrollRecordSerializer(records, many=True).data

        # Summary Metrics
        total_gross = sum(float(r.gross_salary) for r in records)
        total_payable = sum(float(r.payable_amount) for r in records)
        total_insurance = sum(float(r.total_insurance) for r in records)
        total_tax = sum(float(r.income_tax) for r in records)

        return Response({
            'message': f'محاسبه حقوق دوره {period.year_month} برای {len(records)} نفر پرسنل با موفقیت انجام شد.',
            'period_id': period.id,
            'period_status': period.status,
            'summary': {
                'total_personnel': len(records),
                'total_gross': total_gross,
                'total_payable': total_payable,
                'total_insurance': total_insurance,
                'total_tax': total_tax,
            },
            'records': serialized
        })

    @staticmethod
    def _attach_expired_headers(resp, excluded_expired):
        """
        DEF-20: کنار گذاشتن پرسنل منقضی از دیسکت بی‌صدا نیست. شمار و نام‌ها (JSON با کدگذاری URL،
        چون سربرگ HTTP فقط ASCII می‌پذیرد) همراه فایل برمی‌گردد؛ فرانت پیش از دانلود از اکشن expired-personnel می‌پرسد.
        """
        if not excluded_expired:
            return
        import json
        from urllib.parse import quote
        names = [e['full_name'] for e in excluded_expired][:50]
        resp['X-Expired-Personnel-Excluded'] = str(len(excluded_expired))
        resp['X-Expired-Personnel-Names'] = quote(json.dumps(names, ensure_ascii=False))
        resp['Access-Control-Expose-Headers'] = 'Content-Disposition, X-Expired-Personnel-Excluded, X-Expired-Personnel-Names'

    @action(detail=False, methods=['get'], url_path='expired-personnel')
    def expired_personnel(self, request):
        """
        DEF-20: پیش‌پرسش خروجی دیسکت. فهرست پرسنلی که به‌خاطر پایان تاریخ اعتبار از دیسکت بیمه، مالیات
        و بانک این دوره و پروژه کنار گذاشته می‌شوند تا پیش از دانلود به کاربر هشدار داده شود.
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')
        if not period_id or not project_id:
            return Response({'error': 'شناسه دوره (period_id) و پروژه (project_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not str(period_id).isdigit() or not str(project_id).isdigit():
            return Response({'error': 'شناسه دوره یا پروژه نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        from ..lifecycle_services import split_expired_payroll_records
        # فقط فیش‌های درون قلمرو کاربر (get_queryset)؛ دوره مستقیم و بی‌قلمرو خوانده نمی‌شود
        records = [
            r for r in self.get_queryset().filter(period_id=int(period_id)).select_related('personnel', 'period')
            if str(r.personnel.project_id) == str(project_id)
        ]
        year_month = records[0].period.year_month if records else None
        excluded = split_expired_payroll_records(records, year_month)[1] if records else []
        return Response({
            'period_id': int(period_id),
            'year_month': year_month,
            'project_id': int(project_id),
            'count': len(excluded),
            'personnel': excluded,
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='export-bimeh-diskettes')
    def export_bimeh_diskettes(self, request):
        return self.export_dsk_zip(request)

    @action(detail=False, methods=['get'], url_path='export-dsk-zip')
    def export_dsk_zip(self, request):
        """
        صدور پکیج فشرده ZIP شامل دو دیسکت استاندارد DSKWOR00.DBF و DSKKAR00.DBF
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور دیسکت‌های بیمه تامین اجتماعی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = Project.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        # فیلتر تک‌پروژه‌ای الزامی بر اساس انتساب پرسنل به پروژه
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'هیچ رکوردی برای پروژه «{project.name}» در این دوره محاسبه نشده است.'}, status=status.HTTP_404_NOT_FOUND)

        year_str = period.year_month.split('/')[0] if '/' in period.year_month else '1405'
        month_str = period.year_month.split('/')[1] if '/' in period.year_month else '04'
        fiscal_year_short = int(year_str[-2:]) # 05
        month_num = int(month_str)

        # استخراج تنظیمات کارگاه به ترتیب: تنظیم اختصاصی پروژه -> تنظیم سراسری سال در بازه موثر
        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        ws_settings = getattr(settings_obj, 'workshop_insurance', None) if settings_obj else None
        if not ws_settings:
            ws_settings = WorkshopInsuranceSettings.objects.first()
        if not ws_settings:
            ws_settings = WorkshopInsuranceSettings(
                workshop_code='4894290013',
                workshop_name=f'کارگاه {project.name}',
                employer_name='مدیریت انبارداری',
                workshop_address='منطقه عملیاتی'
            )

        # فیلتر پرسنل معتبر و کنار گذاشتن پرسنل منقضی در دیسکت با گزارش هشدار (تسک ۲.۱ و تصمیم مصوب ۶.۶)
        from ..lifecycle_services import split_expired_payroll_records
        records, excluded_expired = split_expired_payroll_records(records, period.year_month)

        insured_recs = [r for r in records if r.include_in_insurance]
        if not insured_recs:
            return Response({'error': f'هیچ پرسنل مشمول بیمه‌ای برای پروژه «{project.name}» در این دوره وجود ندارد.'}, status=status.HTTP_400_BAD_REQUEST)

        # Calculate sums for DSKKAR
        sum_days = sum(r.insurance_days for r in insured_recs)
        sum_daily = sum(int(r.daily_wage) for r in insured_recs)
        sum_base = sum(int(r.base_salary) for r in insured_recs)
        sum_maz = sum(int(r.total_taxable_allowances) for r in insured_recs)
        sum_mash = sum(int(r.total_insurable_salary_allowances) for r in insured_recs)
        sum_totl = sum(int(r.gross_salary) for r in insured_recs)
        sum_bime = sum(int(r.worker_insurance) for r in insured_recs)
        sum_koso = sum(int(r.employer_insurance) for r in insured_recs)
        sum_bic = sum(int(r.unemployment_insurance) for r in insured_recs)
        sum_inc = sum(int(r.total_seniority_accumulated) for r in insured_recs)
        sum_spouse = sum(int(r.marital_allowance) for r in insured_recs)

        month_names = {1: 'فروردین', 2: 'اردیبهشت', 3: 'خرداد', 4: 'تیر', 5: 'مرداد', 6: 'شهریور', 7: 'مهر', 8: 'آبان', 9: 'آذر', 10: 'دی', 11: 'بهمن', 12: 'اسفند'}
        m_name = month_names.get(month_num, 'تیر')
        list_desc = f"{project.name} بیمه {m_name} {year_str}"[:50]

        dskkar_bytes = generate_dskkar_bytes(
            workshop_settings=ws_settings,
            fiscal_year_short=fiscal_year_short,
            month_num=month_num,
            list_number='01',
            list_description=list_desc,
            records_count=len(insured_recs),
            sum_insurance_days=sum_days,
            sum_daily_wage=sum_daily,
            sum_base_salary=sum_base,
            sum_insurable_allowances=sum_maz,
            sum_insurable_total=sum_mash,
            sum_gross_salary=sum_totl,
            sum_employee_insurance=sum_bime,
            sum_employer_insurance=sum_koso,
            sum_unemployment_insurance=sum_bic,
            sum_seniority=sum_inc,
            sum_spouse_allowance=sum_spouse
        )

        dskwor_bytes = generate_dskwor_bytes(
            workshop_code=ws_settings.workshop_code,
            fiscal_year_short=fiscal_year_short,
            month_num=month_num,
            list_number='01',
            payroll_records=records
        )

        # Build ZIP in memory
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr("DSKKAR00.DBF", dskkar_bytes)
            z.writestr("DSKWOR00.DBF", dskwor_bytes)

        zip_val = zip_buf.getvalue()
        resp = HttpResponse(zip_val, content_type='application/zip')
        resp['Content-Disposition'] = f'attachment; filename="Bimeh_{project.code}_{year_str}_{month_str}.zip"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=False, methods=['get'], url_path='export-tax-wh')
    def export_tax_wh(self, request):
        """
        صدور فایل متنی WH مالیات حقوق به تفکیک اجباری پروژه
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور فایل متنی WH مالیاتی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = Project.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'هیچ رکوردی برای پروژه «{project.name}» در این دوره محاسبه نشده است.'}, status=status.HTTP_404_NOT_FOUND)

        # فیلتر پرسنل معتبر و کنار گذاشتن پرسنل منقضی در دیسکت مالیات
        from ..lifecycle_services import split_expired_payroll_records
        records, excluded_expired = split_expired_payroll_records(records, period.year_month)

        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        tax_settings = getattr(settings_obj, 'tax_settings', None)

        content = generate_wh_tax_content(records, tax_settings)
        month_code = period.year_month.replace('/', '')[2:] if period else '140504'
        filename = f"WH_{project.code}_{month_code}.txt"

        resp = HttpResponse(content.encode('utf-8'), content_type='text/plain; charset=utf-8')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=False, methods=['get'], url_path='export-tax-wp')
    def export_tax_wp(self, request):
        """
        صدور فایل متنی WP پرسنل مالیاتی به تفکیک اجباری پروژه
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور فایل متنی WP مالیاتی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = Project.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'هیچ رکوردی برای پروژه «{project.name}» در این دوره محاسبه نشده است.'}, status=status.HTTP_404_NOT_FOUND)

        # فیلتر پرسنل معتبر و کنار گذاشتن پرسنل منقضی در دیسکت مالیات WP
        from ..lifecycle_services import split_expired_payroll_records
        records, excluded_expired = split_expired_payroll_records(records, period.year_month)

        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        tax_settings = getattr(settings_obj, 'tax_settings', None)

        content = generate_wp_tax_content(records, tax_settings)
        month_code = period.year_month.replace('/', '')[2:] if period else '140504'
        filename = f"WP_{project.code}_{month_code}.txt"

        resp = HttpResponse(content.encode('utf-8'), content_type='text/plain; charset=utf-8')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=False, methods=['get'], url_path='export-bank-excel')
    def export_bank_excel(self, request):
        """
        صدور فایل اکسل پرداخت گروهی بانک ملی به تفکیک اجباری پروژه.
        DEF-20 (گام ۲ اصلاحیه): بر مبنای علامت payment_hold روی فیش؛
        در صورت ارسال only_released=true، خروجی مجزا فقط شامل پرداخت‌های آزادشده صادر می‌شود.
        """
        period_id = request.query_params.get('period_id')
        project_id = request.query_params.get('project_id')
        only_released = request.query_params.get('only_released') in ['true', '1']

        if not period_id:
            return Response({'error': 'شناسه دوره (period_id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if not project_id:
            return Response({'error': 'انتخاب پروژه جهت صدور فایل پرداخت بانک ملی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({'error': 'دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        project = Project.objects.filter(id=project_id).first()
        if not project:
            return Response({'error': 'پروژه انتخاب‌شده یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        all_records = list(MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel', 'personnel__project'))
        records = [r for r in all_records if r.personnel.project_id == project.id]
        if not records:
            return Response({'error': f'رکوردی برای پروژه «{project.name}» در این دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        if only_released:
            records = [r for r in records if not r.payment_hold and r.payment_hold_released_at is not None]
            if not records:
                return Response({'error': f'هیچ فیش آزادشده‌ای برای پروژه «{project.name}» در این دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
            excluded_expired = []
            filename = f"Bank_Payment_Released_{project.code}_{period.year_month.replace('/', '_')}.xlsx"
        else:
            excluded_expired = [
                {
                    'personnel_id': r.personnel_id,
                    'full_name': r.personnel.full_name,
                    'national_code': r.personnel.national_code,
                    'valid_until_date': getattr(r.personnel, 'valid_until_date', None),
                }
                for r in records if r.payment_hold
            ]
            records = [r for r in records if not r.payment_hold and r.payment_hold_released_at is None]
            if not records:
                return Response({'error': f'هیچ فیش حقوقی عادی برای صدور دیسکت بانک پروژه «{project.name}» در این دوره یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
            filename = f"Bank_Payment_{project.code}_{period.year_month.replace('/', '_')}.xlsx"

        settings_obj = get_effective_payroll_settings(period.year_month, project.id)
        bank_settings = getattr(settings_obj, 'bank_export_settings', None)

        excel_bytes = generate_bank_meli_excel(records, bank_settings, year_month_title=f"{project.name} {period.year_month}")

        resp = HttpResponse(excel_bytes, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        self._attach_expired_headers(resp, excluded_expired)
        return resp

    @action(detail=True, methods=['post'], url_path='release-payment-hold')
    def release_payment_hold(self, request, pk=None):
        """
        DEF-20 (گام ۲ اصلاحیه): آزادسازی پرداخت معلق پرسنل منقضی پس از تمدید اعتبار پرونده.
        اعتبار را می‌سنجد، لاگ ممیزی ثبت می‌کند و تعلیق پرداخت را آزاد می‌سازد.
        """
        from ..lifecycle_services import user_can_view_wages, validate_personnel_validity
        from accounts.audit_utils import log_audit_event
        from django.utils import timezone

        if not user_can_view_wages(request.user):
            raise PermissionDenied("آزادسازی پرداخت معلق حقوق تنها با دسترسی مالی و حسابداری مجاز است.")

        record = self.get_object()
        if not record.payment_hold:
            return Response({'message': 'این فیش حقوقی در وضعیت تعلیق پرداخت قرار ندارد.'}, status=status.HTTP_200_OK)

        # سنجش تمدید اعتبار پرسنل
        is_valid, reason = validate_personnel_validity(record.personnel, record.period.year_month)
        if not is_valid:
            return Response({
                'error': f'اعتبار پرونده هنوز تمدید نشده است ({reason}). لطفاً ابتدا تاریخ اعتبار را از طریق درخواست تغییرات تمدید نمایید.'
            }, status=status.HTTP_400_BAD_REQUEST)

        record.payment_hold = False
        record.payment_hold_released_at = timezone.now()
        record.payment_hold_released_by = request.user
        record.save(update_fields=['payment_hold', 'payment_hold_released_at', 'payment_hold_released_by', 'updated_at'])

        log_audit_event(
            module='personnel',
            action='RELEASE_PAYMENT_HOLD',
            severity='info',
            target_model='MonthlyPayrollRecord',
            target_object_id=str(record.id),
            target_repr=f"آزادسازی پرداخت معلق فیش {record.personnel.full_name} برای دوره {record.period.year_month}",
            user=request.user,
            details={
                'record_id': record.id,
                'personnel_id': record.personnel_id,
                'period_id': record.period_id,
                'payable_amount': str(record.payable_amount),
            }
        )

        return Response({
            'message': f'پرداخت معلق فیش «{record.personnel.full_name}» با موفقیت آزاد شد.',
            'record_id': record.id,
            'payment_hold': False,
            'payment_hold_released_at': record.payment_hold_released_at.isoformat()
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='export-monthly-excel')
    def export_monthly_excel(self, request):
        """
        صدور فایل اکسل ۲ سطری استاندارد ۵۸ ستون شیت ماهانه
        """
        period_id = request.query_params.get('period_id')
        records = list(MonthlyPayrollRecord.objects.filter(period_id=period_id).select_related('personnel'))
        if not records:
            return Response({'error': 'رکوردی یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        period = records[0].period
        excel_bytes = generate_monthly_payroll_excel(records, period_title=period.year_month)
        filename = f"Payroll_Report_{period.year_month.replace('/', '_')}.xlsx"

        resp = HttpResponse(excel_bytes, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        return resp

    @action(detail=False, methods=['post'], url_path='import-tax-excel')
    def import_tax_excel(self, request):
        """
        درون‌ریزی فایل اکسل مالیات محاسبه‌شده توسط دارایی (نمونه مالیات محاسبه شده.xlsx)
        یا لیست داده‌های مالیات ارسالی متنی/JSON و تطبیق خودکار با ستون مالیات دوره بر اساس کد ملی
        """
        excel_file = request.FILES.get('file')
        items = request.data.get('items')
        period_id = request.data.get('period_id') or request.query_params.get('period_id')
        year_month = request.data.get('year_month') or request.query_params.get('year_month')

        if not excel_file and not items:
            return Response({'error': 'لطفاً فایل اکسل مالیات دارایی یا داده‌های متنی را انتخاب کنید.'}, status=status.HTTP_400_BAD_REQUEST)

        # Locate target period
        if period_id:
            try:
                p_id = int(str(period_id).strip())
                period = MonthlyWorkPeriod.objects.filter(id=p_id).first()
            except Exception:
                period = MonthlyWorkPeriod.objects.filter(id=period_id).first()
        elif year_month:
            period = MonthlyWorkPeriod.objects.filter(year_month=year_month).first()
        else:
            period = MonthlyWorkPeriod.objects.order_by('-created_at').first()

        if not period:
            return Response({'error': 'دوره ماهانه مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        try:
            matched_count = 0
            unmatched_rows = []
            total_imported_tax = Decimal(0)

            records = MonthlyPayrollRecord.objects.filter(period=period).select_related('personnel')
            record_map = {str(r.national_code).strip().zfill(10): r for r in records}

            if items and isinstance(items, list):
                with transaction.atomic():
                    for idx, item in enumerate(items, 1):
                        raw_code = item.get('national_code')
                        if not raw_code:
                            continue
                        code_str = str(raw_code).split('.')[0].strip().zfill(10)
                        if not code_str or not code_str.isdigit():
                            continue
                        try:
                            tax_val = Decimal(str(item.get('tax_amount', 0)).replace(',', '').strip())
                        except Exception:
                            tax_val = Decimal(0)
                        try:
                            ex_months = int(item.get('exemption_months', 1) or 1)
                        except Exception:
                            ex_months = 1
                        raw_mult = item.get('multiple_employers', 'خیر')
                        has_mult = str(raw_mult).strip() in ['بله', 'true', 'True', '1', 'YES', 'yes']

                        if code_str in record_map:
                            rec = record_map[code_str]
                            rec.income_tax = tax_val
                            rec.tax_source_type = 'IMPORTED_EXCEL'
                            rec.tax_exemption_months = ex_months
                            rec.has_multiple_employers = has_mult
                            rec.is_tax_imported = True

                            rec.net_salary = rec.gross_salary - rec.worker_insurance - rec.income_tax
                            rec.payable_amount = rec.net_salary - rec.advance_payment_deduction
                            rec.save()

                            matched_count += 1
                            total_imported_tax += tax_val
                        else:
                            unmatched_rows.append({
                                'row': idx,
                                'national_code': code_str,
                                'tax_amount': float(tax_val)
                            })

            elif excel_file:
                file_bytes = excel_file.read()
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                ws = wb.active

                # Find header indices
                col_map = {}
                for c in range(1, ws.max_column + 1):
                    header_val = str(ws.cell(1, c).value or '').strip()
                    clean_h = header_val.replace(' ', '').replace('_', '')
                    if 'کدملی' in clean_h or 'کد_ملی' in clean_h:
                        col_map['national_code'] = c
                    elif 'مالیاتمحاسبهشده' in clean_h or 'مالیات' in clean_h:
                        col_map['tax_amount'] = c
                    elif 'تعدادماهاعمالشدهدرمعافیت' in clean_h or 'تعدادماه' in clean_h:
                        col_map['exemption_months'] = c
                    elif 'بیشازیککارفرما' in clean_h or 'چندکارفرما' in clean_h:
                        col_map['multiple_employers'] = c

                if 'national_code' not in col_map or 'tax_amount' not in col_map:
                    col_map.setdefault('national_code', 1)
                    col_map.setdefault('tax_amount', 4)
                    col_map.setdefault('exemption_months', 5)
                    col_map.setdefault('multiple_employers', 6)

                with transaction.atomic():
                    for r_idx in range(2, ws.max_row + 1):
                        raw_code = ws.cell(r_idx, col_map['national_code']).value
                        if raw_code is None:
                            continue
                        
                        code_str = str(raw_code).split('.')[0].strip().zfill(10)
                        if not code_str or not code_str.isdigit():
                            continue

                        raw_tax = ws.cell(r_idx, col_map['tax_amount']).value or 0
                        try:
                            tax_val = Decimal(str(raw_tax).replace(',', '').strip())
                        except Exception:
                            tax_val = Decimal(0)

                        raw_months = ws.cell(r_idx, col_map.get('exemption_months', 5)).value if 'exemption_months' in col_map else 1
                        try:
                            ex_months = int(raw_months or 1)
                        except Exception:
                            ex_months = 1

                        raw_mult = ws.cell(r_idx, col_map.get('multiple_employers', 6)).value if 'multiple_employers' in col_map else 'خیر'
                        has_mult = str(raw_mult).strip() in ['بله', 'true', 'True', '1', 'YES', 'yes']

                        if code_str in record_map:
                            rec = record_map[code_str]
                            rec.income_tax = tax_val
                            rec.tax_source_type = 'IMPORTED_EXCEL'
                            rec.tax_exemption_months = ex_months
                            rec.has_multiple_employers = has_mult
                            rec.is_tax_imported = True

                            rec.net_salary = rec.gross_salary - rec.worker_insurance - rec.income_tax
                            rec.payable_amount = rec.net_salary - rec.advance_payment_deduction
                            rec.save()

                            matched_count += 1
                            total_imported_tax += tax_val
                        else:
                            unmatched_rows.append({
                                'row': r_idx,
                                'national_code': code_str,
                                'tax_amount': float(tax_val)
                            })

            return Response({
                'message': f'اطلاعات مالیاتی با موفقیت بارگذاری شد. {matched_count} پرسنل تطبیق داده شدند.',
                'matched_count': matched_count,
                'unmatched_count': len(unmatched_rows),
                'unmatched_rows': unmatched_rows,
                'total_imported_tax': float(total_imported_tax),
                'period_year_month': period.year_month
            }, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({
                'error': f'خطا در پردازش اطلاعات مالیات: {str(e)}'
            }, status=status.HTTP_400_BAD_REQUEST)

class FleetSettlementViewSet(viewsets.ViewSet):
    """
    ویوست تسویه و محاسبات مالی ناوگان خودرویی
    - محاسبه مبالغ تسویه ماهانه
    - تولید فایل پرداخت گروهی بانک ملی
    """
    permission_classes = [permissions.IsAuthenticated]

    @action(detail=False, methods=['get'], url_path='calculate')
    def calculate(self, request):
        year_month = request.query_params.get('year_month')
        warehouse_id = request.query_params.get('warehouse_id')
        if not year_month:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        wh_id = int(warehouse_id) if (warehouse_id and warehouse_id.isdigit()) else None
        try:
            data = calculate_monthly_fleet_settlement(warehouse_id=wh_id, year_month=year_month)
            return Response(data, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'error': f'خطا در محاسبه تسویه ناوگان: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], url_path='export-bank-excel')
    def export_bank_excel(self, request):
        year_month = request.query_params.get('year_month')
        warehouse_id = request.query_params.get('warehouse_id')
        if not year_month:
            return Response({'error': 'پارامتر year_month الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        wh_id = int(warehouse_id) if (warehouse_id and warehouse_id.isdigit()) else None
        try:
            data = calculate_monthly_fleet_settlement(warehouse_id=wh_id, year_month=year_month)
            excel_file = generate_fleet_bank_meli_excel(data, year_month)

            clean_ym = year_month.replace('/', '_')
            filename = f"fleet_bank_payment_{clean_ym}.xlsx"
            response = HttpResponse(
                excel_file.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response
        except Exception as e:
            return Response({'error': f'خطا در تولید فایل اکسل پرداخت بانک ملی: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

__all__ = [
    'PayrollYearlySettingsViewSet',
    'MonthlyPayrollViewSet',
    'FleetSettlementViewSet',
]
