# -*- coding: utf-8 -*-
"""
ViewSet for inventory items, bulk operations, excel import/export, and photo management.
"""
import os
import re
import time
import json
import uuid
import queue
import asyncio
import logging
import tempfile
import threading
import traceback
from datetime import datetime
from decimal import Decimal

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import PatternFill

from django.core.cache import cache
from django.db import transaction, models
from django.db.models import Q
from django.forms.models import model_to_dict
from django.http import StreamingHttpResponse, HttpResponse
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import viewsets, filters, status
from rest_framework.decorators import action
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from common.mixins import DeleteImpactMixin
from common.warehouse_scope import scope_queryset, can_access_warehouse
from warehouses.models import Warehouse

from ..filters import ItemFilter
from ..models import (
    Item,
    ItemPhoto,
    ImportLog,
    ImportHistory,
    CountTask,
    CountTaskHistory,
    DocTask,
    DocTaskHistory,
    ItemFieldDefinition,
)
from ..serializers import (
    ItemSerializer,
    ItemPhotoSerializer,
    photo_prefetch,
)
from ..signals import (
    broadcast_count_task_update,
    broadcast_doc_task_update,
)
from .helpers import (
    _parse_decimal,
    _parse_date_flexible,
    _parse_positive_int,
    _clean_currency,
    _clean_invoice_type,
    _create_doc_task_snapshot,
    _soft_delete_items_cascade,
    _restore_item_photos_cascade,
    _cleanup_old_import_logs,
    ItemPagination,
    PriorityOrderingFilter,
    logger,
)

BULK_UPDATE_PROTECTED_FIELDS = {
    'id', 'warehouse', 'created_by', 'modified_by', 'created_at', 'updated_at', 'sync_id', 'is_deleted'
}

class ItemViewSet(DeleteImpactMixin, viewsets.ModelViewSet):
    queryset = Item.objects.all()
    serializer_class = ItemSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, PriorityOrderingFilter]
    filterset_class = ItemFilter
    pagination_class = ItemPagination
    search_fields = ['fa_unic_code', 'description', 'po', 'pl', 'pk_number', 'my_tag']
    ordering_fields = '__all__'
    parser_classes = (MultiPartParser, FormParser, *viewsets.ModelViewSet.parser_classes)

    def get_object(self):
        """
        زیرساخت سینک: اگر pk عددی نبود، به‌عنوان sync_id تفسیر می‌شود (تسک ۲۲).
        """
        pk = self.kwargs.get(self.lookup_url_kwarg or self.lookup_field)
        if pk is not None and not str(pk).isdigit():
            from django.shortcuts import get_object_or_404
            queryset = self.filter_queryset(self.get_queryset())
            obj = get_object_or_404(queryset, sync_id=pk)
            self.check_object_permissions(self.request, obj)
            return obj
        return super().get_object()

    def get_queryset(self):
        # photo_prefetch جلوی N+1 عکس‌ها را می‌گیرد: بدون آن هر ردیف یک COUNT و
        # یک SELECT جدا برای بندانگشتی می‌زد (صفحه ۱۰۰ ردیفی = ۲۰۰ کوئری اضافه).
        queryset = super().get_queryset().prefetch_related(photo_prefetch())
        from common.tenant_scope import scope_tenant_queryset
        queryset = scope_tenant_queryset(queryset, self.request.user, company_field='warehouse__company_id', request=self.request)
        return scope_queryset(queryset, self.request.user)

    def create(self, request, *args, **kwargs):
        sync_id = request.data.get('sync_id') if hasattr(request, 'data') and hasattr(request.data, 'get') else None
        if isinstance(sync_id, str):
            sync_id = sync_id.strip()
        if sync_id:
            try:
                existing = Item.all_objects.filter(sync_id=sync_id).first()
            except Exception:
                existing = None
            if existing:
                from common.tenant_scope import scope_tenant_queryset
                scoped_qs = scope_tenant_queryset(
                    Item.all_objects.filter(pk=existing.pk),
                    request.user,
                    company_field='warehouse__company_id',
                    request=request
                )
                scoped_qs = scope_queryset(scoped_qs, request.user)
                if not scoped_qs.filter(pk=existing.pk).exists():
                    return Response(
                        {"error": "دسترسی غیرمجاز: یک یا چند قلم کالا خارج از انبار مجاز شماست."},
                        status=status.HTTP_403_FORBIDDEN
                    )

                return Response(self.get_serializer(existing).data, status=status.HTTP_200_OK)

        return super().create(request, *args, **kwargs)

    def _allowed_warehouse_ids(self):
        cached = getattr(self, '_cached_allowed_warehouse_ids', None)
        if cached is not None:
            return cached
        from common.tenant_scope import scope_tenant_queryset
        from common.warehouse_scope import user_warehouse_ids
        request = getattr(self, 'request', None)
        user = getattr(request, 'user', None)
        if not user or not user.is_authenticated:
            self._cached_allowed_warehouse_ids = set()
            return self._cached_allowed_warehouse_ids

        from companies.models import Company
        has_active_companies = Company.objects.filter(is_active=True).exists()
        if has_active_companies:
            qs = scope_tenant_queryset(
                Warehouse.objects.all(),
                user,
                company_field='company_id',
                request=request
            )
        else:
            qs = Warehouse.objects.all()

        assigned = user_warehouse_ids(user)
        if assigned is not None:
            qs = qs.filter(id__in=assigned)
        self._cached_allowed_warehouse_ids = set(qs.values_list('id', flat=True))
        return self._cached_allowed_warehouse_ids

    def _can_write_warehouse(self, warehouse_id):
        if warehouse_id is None:
            from common.warehouse_scope import can_access_warehouse
            user = getattr(getattr(self, 'request', None), 'user', None)
            return can_access_warehouse(user, None)
        try:
            target_id = getattr(warehouse_id, 'id', warehouse_id)
            return int(target_id) in self._allowed_warehouse_ids()
        except (ValueError, TypeError):
            return False

    def _can_manage_import_log(self, import_log):
        if not import_log:
            return False
        user = getattr(getattr(self, 'request', None), 'user', None)
        if not user or not user.is_authenticated:
            return False
        if user.is_superuser:
            return True
        if import_log.warehouse_id is None:
            return import_log.imported_by_id == user.id
        return self._can_write_warehouse(import_log.warehouse_id)

    def _ensure_target_warehouse_in_scope(self, warehouse):
        from rest_framework.exceptions import PermissionDenied
        wh_id = getattr(warehouse, 'id', warehouse)
        if not self._can_write_warehouse(wh_id):
            raise PermissionDenied("دسترسی غیرمجاز: انبار انتخابشده خارج از محدودهٔ مجاز شماست.")

    def perform_create(self, serializer):
        self._ensure_target_warehouse_in_scope(serializer.validated_data.get('warehouse'))
        from accounts.audit_utils import log_audit_event
        # ۱. بررسی idempotency بر مبنای sync_id (تسک ۱۴)
        sync_id = serializer.validated_data.get('sync_id')
        existing_by_sync = Item.all_objects.filter(sync_id=sync_id).first() if sync_id else None
        if existing_by_sync:
            serializer.instance = existing_by_sync
            instance = serializer.save(is_deleted=False)
            if existing_by_sync.is_deleted:
                _restore_item_photos_cascade([instance.id])
            return

        # ۲. احیای رکورد حذف‌نرم با همان (انبار، کد یکتا) به‌جای INSERT تکراری
        tombstone = Item.all_objects.filter(
            warehouse=serializer.validated_data.get('warehouse'),
            fa_unic_code=serializer.validated_data.get('fa_unic_code'),
            is_deleted=True,
        ).first()
        if tombstone:
            serializer.instance = tombstone
            instance = serializer.save(is_deleted=False)
            # احیای کالا باید عکس‌هایش را هم برگرداند؛ حذف آن‌ها را tombstone کرده بود.
            _restore_item_photos_cascade([instance.id])
        else:
            instance = serializer.save()

        try:
            log_audit_event(
                module='docs',
                action='CREATE',
                target_model='Item',
                target_object_id=instance.id,
                target_repr=f"{instance.fa_unic_code or ''} - {instance.description or ''}"[:255],
                warehouse=instance.warehouse,
                user=self.request.user if self.request.user.is_authenticated else None,
                after_state=model_to_dict(instance, exclude=['photo'])
            )
        except Exception:
            pass

    def perform_update(self, serializer):
        if 'warehouse' in serializer.validated_data:
            self._ensure_target_warehouse_in_scope(serializer.validated_data.get('warehouse'))
        from accounts.audit_utils import log_audit_event, calculate_model_diff
        old_instance = self.get_object()
        before_state = model_to_dict(old_instance, exclude=['photo'])
        instance = serializer.save()
        after_state = model_to_dict(instance, exclude=['photo'])
        diff_b, diff_a = calculate_model_diff(before_state, after_state)
        if diff_b or diff_a:
            try:
                log_audit_event(
                    module='docs',
                    action='UPDATE',
                    target_model='Item',
                    target_object_id=instance.id,
                    target_repr=f"{instance.fa_unic_code or ''} - {instance.description or ''}"[:255],
                    warehouse=instance.warehouse,
                    user=self.request.user if self.request.user.is_authenticated else None,
                    before_state=diff_b,
                    after_state=diff_a
                )
            except Exception:
                pass

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)

        # تشخیص تداخل خوش‌بینانه بر اساس base_updated_at (تسک ۶)
        base_raw = request.data.get('base_updated_at') or request.data.get('baseUpdatedAt')
        if base_raw:
            from django.utils.dateparse import parse_datetime
            base = parse_datetime(str(base_raw))
            if base is not None:
                if timezone.is_naive(base):
                    base = timezone.make_aware(base)
                instance = self.get_object()
                if (instance.updated_at - base).total_seconds() > 0.001:
                    return Response({
                        'detail': 'conflict',
                        'server_record': self.get_serializer(instance).data,
                    }, status=409)
        
        user = request.user
        if not (user.is_superuser or user.has_perm('accounts.perm_wh_edit')):
            instance = self.get_object()
            from warehouses.services import get_setting
            
            if user.has_perm('accounts.view_sys_counter'):
                perms = get_setting('field_permissions_counter', instance.warehouse_id) or {}
            elif user.has_perm('accounts.view_sys_supervisor') or user.has_perm('accounts.view_sys_financial'):
                perms = get_setting('field_permissions_doc', instance.warehouse_id) or {}
            else:
                perms = {}
                
            editable_fields = {k for k, v in perms.items() if isinstance(v, dict) and v.get('editable')}
            
            invalid_fields = []
            for field in request.data.keys():
                if field == 'dynamic_data':
                    dyn_data = request.data.get('dynamic_data')
                    if isinstance(dyn_data, str):
                        try:
                            import json
                            dyn_data = json.loads(dyn_data)
                        except:
                            dyn_data = {}
                    if isinstance(dyn_data, dict):
                        for dyn_k in dyn_data.keys():
                            if dyn_k not in editable_fields:
                                invalid_fields.append(dyn_k)
                elif field not in editable_fields and field not in ['id', 'warehouse']:
                    invalid_fields.append(field)
            
            if invalid_fields:
                from accounts.audit_utils import log_audit_event
                log_audit_event(
                    user=user,
                    module='docs',
                    action='UPDATE',
                    severity='warning',
                    target_model='Item',
                    target_object_id=instance.id,
                    target_repr=f"تلاش غیرمجاز برای ویرایش فیلدها",
                    warehouse=instance.warehouse,
                    details={'invalid_fields': invalid_fields},
                    ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
                )
                return Response({
                    'error': 'شما مجوز ویرایش این فیلدها را ندارید.',
                    'invalid_fields': invalid_fields
                }, status=400)
                
        # Re-inject partial since kwargs is passed to super
        kwargs['partial'] = partial
        return super().update(request, *args, **kwargs)

    def perform_destroy(self, instance):
        from accounts.audit_utils import log_audit_event
        try:
            log_audit_event(
                module='docs',
                action='DELETE',
                severity='critical',
                target_model='Item',
                target_object_id=instance.id,
                target_repr=f"{instance.fa_unic_code or ''} - {instance.description or ''}"[:255],
                warehouse=instance.warehouse,
                user=self.request.user if self.request.user.is_authenticated else None,
                before_state=model_to_dict(instance, exclude=['photo'])
            )
        except Exception:
            pass

        with transaction.atomic():
            _soft_delete_items_cascade(Item.objects.filter(id=instance.id))

    def get_permissions(self):
        from accounts.permissions import HasMenuAccess, CanManageItemPhotos
        from rest_framework.permissions import AllowAny, IsAuthenticated

        if self.action == 'photos':
            # خواندن عکس برای هر کاربر لاگین‌شده، ولی آپلود فقط برای نقش‌هایی که
            # با عکس کالا کار می‌کنند. پیش از این POST هم IsAuthenticated بود.
            if self.request.method in ('GET', 'HEAD', 'OPTIONS'):
                permission_classes = [IsAuthenticated()]
            else:
                permission_classes = [IsAuthenticated(), CanManageItemPhotos()]
        elif self.action in ['list', 'retrieve', 'dashboard_stats', 'export_columns', 'download_template']:
            permission_classes = [IsAuthenticated()]
        elif self.action == 'bulk_assign':
            permission_classes = [HasMenuAccess('perm_rec_dispatch')]
        elif self.action in ['export_excel', 'export_excel_mt']: # I'll just secure export here in case it's added
            permission_classes = [HasMenuAccess('view_sys_export')]
        elif self.action in ['import_excel', 'cancel_import', 'revert_import', 'download_import_log', 'delete_from_excel', 'clear_warehouse_data', 'latest_import', 'parse_headers']:
            permission_classes = [HasMenuAccess('perm_rec_import')]
        elif self.action in ['reject', 'manager_reject']:
            permission_classes = [HasMenuAccess('perm_rec_recount')]
        elif self.action in ['update', 'partial_update', 'bulk_update', 'bulk_tag']:
            permission_classes = [HasMenuAccess('perm_wh_edit') | HasMenuAccess('view_sys_counter') | HasMenuAccess('view_sys_supervisor')]
        else: # create, destroy, etc
            permission_classes = [HasMenuAccess('perm_wh_edit')]

        return permission_classes

    @action(detail=True, methods=['get', 'post'])
    def photos(self, request, pk=None):
        from ..views_photos import handle_item_photos_upload, list_item_photos
        item = self.get_object()
        if request.method == 'GET':
            return list_item_photos(request, item)
        return handle_item_photos_upload(request, item)


    @action(detail=False, methods=['post'])
    def export_excel(self, request):
        import openpyxl
        import io
        import zipfile
        from django.http import HttpResponse
        
        data_scope = request.data.get('data_scope', 'all')
        columns_scope = request.data.get('columns_scope', 'all_db')
        columns_list = request.data.get('columns_list', [])
        warehouse_id = request.data.get('warehouse_id') or request.query_params.get('warehouse_id')
        
        from django.http import QueryDict
        original_query_params = request._request.GET
        try:
            q = QueryDict(mutable=True)
            for k, v in request.data.items():
                if isinstance(v, list):
                    q.setlist(k, [str(x) for x in v])
                else:
                    q[k] = str(v)
            request._request.GET = q
            
            if data_scope == 'selected':
                selected_ids = request.data.get('selected_ids', [])
                queryset = self.get_queryset().filter(id__in=selected_ids)
                queryset = self.filter_queryset(queryset)
            else:
                queryset = self.filter_queryset(self.get_queryset())
        finally:
            request._request.GET = original_query_params
            
        expected_fields_dict = self.get_expected_fields(warehouse_id)
        valid_fields = {f.name: f for f in Item._meta.fields if f.name != 'dynamic_data'}
        
        canonical_headers = list(dict.fromkeys(expected_fields_dict.values()))
        if columns_scope == 'all_db':
            headers = canonical_headers
        elif columns_scope in ['visible', 'custom']:
            headers = list(dict.fromkeys([expected_fields_dict[c] for c in columns_list if c in expected_fields_dict]))
            if not headers:
                headers = canonical_headers
        else:
            headers = canonical_headers
            
        header_labels = [(valid_fields[h].verbose_name if hasattr(valid_fields[h], 'verbose_name') and valid_fields[h].verbose_name else h) if h in valid_fields else f"{h} (داینامیک)" for h in headers]
        header_keys = [h for h in headers]
        
        CHUNK_SIZE = 100_000
        total_count = queryset.count()

        def _format_row(item):
            row = []
            for h in headers:
                if h in valid_fields:
                    val = getattr(item, h, '')
                else:
                    val = item.dynamic_data.get(h, '') if item.dynamic_data else ''
                    
                if val is None:
                    val = ''
                elif hasattr(val, 'username'):
                    val = f"{val.first_name} {val.last_name}".strip() or val.username
                elif val.__class__.__name__ == 'Warehouse':
                    val = str(val)
                elif hasattr(val, 'project_name') and getattr(val, 'project_name'):
                    val = getattr(val, 'project_name')
                else:
                    val = str(val)
                row.append(val)
            return row

        if total_count <= CHUNK_SIZE:
            wb = openpyxl.Workbook(write_only=True)
            ws = wb.create_sheet(title="Export")
            ws.append(header_labels)
            ws.append(header_keys)
            
            for item in queryset.iterator(chunk_size=5000):
                ws.append(_format_row(item))
                
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="export_items.xlsx"'
            response['X-Export-Type'] = 'single'
            response['X-Total-Records'] = str(total_count)
            wb.save(response)
            return response
        else:
            # Multi-part ZIP output for very large datasets
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                part_num = 1
                current_wb = openpyxl.Workbook(write_only=True)
                current_ws = current_wb.create_sheet(title=f"Part_{part_num}")
                current_ws.append(header_labels)
                current_ws.append(header_keys)
                
                rows_in_current_part = 0
                for item in queryset.iterator(chunk_size=5000):
                    current_ws.append(_format_row(item))
                    rows_in_current_part += 1
                    
                    if rows_in_current_part >= CHUNK_SIZE:
                        part_buffer = io.BytesIO()
                        current_wb.save(part_buffer)
                        part_buffer.seek(0)
                        zip_file.writestr(f"export_part_{part_num}.xlsx", part_buffer.getvalue())
                        part_buffer.close()
                        
                        part_num += 1
                        current_wb = openpyxl.Workbook(write_only=True)
                        current_ws = current_wb.create_sheet(title=f"Part_{part_num}")
                        current_ws.append(header_labels)
                        current_ws.append(header_keys)
                        rows_in_current_part = 0
                
                if rows_in_current_part > 0:
                    part_buffer = io.BytesIO()
                    current_wb.save(part_buffer)
                    part_buffer.seek(0)
                    zip_file.writestr(f"export_part_{part_num}.xlsx", part_buffer.getvalue())
                    part_buffer.close()
            
            zip_buffer.seek(0)
            response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
            response['Content-Disposition'] = 'attachment; filename="export_items.zip"'
            response['X-Export-Type'] = 'zip'
            response['X-Total-Records'] = str(total_count)
            response['X-Total-Parts'] = str(part_num)
            return response

    @action(detail=False, methods=['get'])
    def export_columns(self, request):
        warehouse_id = request.query_params.get('warehouse_id')
        if warehouse_id:
            try:
                wh_id_val = int(warehouse_id)
            except (ValueError, TypeError):
                return Response({'error': 'شناسه انبار نامعتبر است.'}, status=400)
            if not self._can_write_warehouse(wh_id_val):
                return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)
            warehouse_id = wh_id_val

        valid_fields = Item._meta.fields
        columns = []
        for f in valid_fields:
            if f.name == 'dynamic_data': continue
            label = getattr(f, 'verbose_name', '') or f.name
            columns.append({"key": f.name, "label": str(label)})
            
        if warehouse_id:
            dynamic_defs = ItemFieldDefinition.objects.filter(warehouse_id=warehouse_id, is_active=True)
            for d in dynamic_defs:
                columns.append({"key": d.name, "label": f"{d.name} (داینامیک)"})
                
        return Response(columns)

    @action(detail=False, methods=['get'])
    def download_template(self, request):
        import openpyxl
        from openpyxl.styles import PatternFill
        from django.http import HttpResponse

        warehouse_id = request.query_params.get('warehouse_id')
        if warehouse_id:
            try:
                wh_id_val = int(warehouse_id)
            except (ValueError, TypeError):
                return Response({'error': 'شناسه انبار نامعتبر است.'}, status=400)
            if not self._can_write_warehouse(wh_id_val):
                return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)
            warehouse_id = wh_id_val

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template"

        expected_fields = self.get_expected_fields(warehouse_id)
        EXCLUDED_TEMPLATE_FIELDS = {'sync_id', 'is_deleted', 'created_at', 'updated_at', 'created_by', 'modified_by'}
        headers = [h for h in dict.fromkeys(expected_fields.values()) if h not in EXCLUDED_TEMPLATE_FIELDS]

        # Add dynamic fields if warehouse_id is provided
        dynamic_fields = []
        if warehouse_id:
            dynamic_defs = ItemFieldDefinition.objects.filter(warehouse_id=warehouse_id, is_active=True)
            for d in dynamic_defs:
                dynamic_fields.append(d)

        # Write headers
        ws.append(headers)

        # Color system fields in the header
        from warehouses.services import get_setting
        restricted_fields = get_setting('SENSITIVE_EXCEL_FIELDS', warehouse_id) or ['doc_status', 'field_status', 'tag_status']
        
        system_fields = ['id', 'created_at', 'updated_at', 'created_by', 'modified_by'] + restricted_fields
        sys_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid") # Warning color (yellow/orange)
        for col_idx, h in enumerate(headers, 1):
            if h in system_fields:
                ws.cell(row=1, column=col_idx).fill = sys_fill

        # Write sample data
        sample_row_1 = []
        sample_row_2 = []
        valid_fields_dict = {f.name: f for f in Item._meta.fields}

        for h in headers:
            field = valid_fields_dict.get(h)
            
            # System fields mock data
            if h in system_fields:
                if h == 'id':
                    sample_row_1.append(101)
                    sample_row_2.append(102)
                elif h in ['created_at', 'updated_at']:
                    sample_row_1.append('2023-10-01 12:00')
                    sample_row_2.append('2023-10-01 12:05')
                elif h in ['created_by', 'modified_by']:
                    sample_row_1.append('سیستم / کاربر ۱')
                    sample_row_2.append('سیستم / کاربر ۲')
                else:
                    sample_row_1.append('مقدار سیستمی')
                    sample_row_2.append('مقدار سیستمی')
                continue
                
            # Dynamic fields mock data
            dyn_field = next((df for df in dynamic_fields if df.name == h), None)
            if dyn_field:
                if dyn_field.field_type == 'number':
                    sample_row_1.append(10)
                    sample_row_2.append(25)
                elif dyn_field.field_type == 'boolean':
                    sample_row_1.append(True)
                    sample_row_2.append(False)
                elif dyn_field.field_type == 'date':
                    sample_row_1.append('1402/05/10')
                    sample_row_2.append('1402/06/15')
                else:
                    sample_row_1.append('مقدار تست ۱')
                    sample_row_2.append('مقدار تست ۲')
                continue

            if h == 'fa_unic_code':
                sample_row_1.append('FA-10001')
                sample_row_2.append('FA-10002')
            elif h == 'warehouse':
                sample_row_1.append('انبار مرکزی')
                sample_row_2.append('انبار مرکزی')
            elif h in ['inventory', 'bal4miv']:
                sample_row_1.append(100.0)
                sample_row_2.append(50.5)
            elif field and field.get_internal_type() == 'BooleanField':
                sample_row_1.append(False)
                sample_row_2.append(True)
            elif field and field.get_internal_type() == 'DateField':
                sample_row_1.append('2023-01-01')
                sample_row_2.append('2023-02-01')
            else:
                sample_row_1.append('نمونه داده')
                sample_row_2.append('نمونه داده ۲')

        ws.append(sample_row_1)
        ws.append(sample_row_2)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Warehouse_Template.xlsx"'
        wb.save(response)

        return response

    @action(detail=False, methods=['post'])
    def bulk_update(self, request):
        try:
            records = request.data
            if not isinstance(records, list):
                return Response({"error": "Data must be a list of records"}, status=400)
            
            record_dict = {str(r['id']): r for r in records if 'id' in r}
            items = self.get_queryset().filter(id__in=record_dict.keys())
            
            # اعتبارسنجی امنیتی: جلوگیری از ویرایش رکوردهای خارج از انبار مجاز کاربر
            if len(items) != len(record_dict):
                missing_ids = set(record_dict.keys()) - set(str(item.id) for item in items)
                if Item.objects.filter(id__in=missing_ids).exists():
                    return Response({"error": "دسترسی غیرمجاز: یک یا چند قلم کالا خارج از انبار مجاز شماست."}, status=403)
            
            user = request.user
            check_field_perms = not (user.is_superuser or user.has_perm('accounts.perm_wh_edit'))
            warehouse_perms_cache = {}
            from warehouses.services import get_setting

            items_to_update = []
            update_fields = set()
            valid_fields = {f.name: f for f in Item._meta.fields}

            for item in items:
                record = record_dict[str(item.id)]

                # بررسی ماتریس دسترسی فیلدها برای کاربران دارای محدودیت نقش (مانند انبارگردان/سرپرست)
                if check_field_perms:
                    wh_id = item.warehouse_id
                    if wh_id not in warehouse_perms_cache:
                        if user.has_perm('accounts.view_sys_counter'):
                            perms = get_setting('field_permissions_counter', wh_id) or {}
                        elif user.has_perm('accounts.view_sys_supervisor') or user.has_perm('accounts.view_sys_financial'):
                            perms = get_setting('field_permissions_doc', wh_id) or {}
                        else:
                            perms = {}
                        warehouse_perms_cache[wh_id] = {k for k, v in perms.items() if isinstance(v, dict) and v.get('editable')}

                    editable_fields = warehouse_perms_cache[wh_id]
                    invalid_fields = []
                    for key in record.keys():
                        if key in ['id', 'warehouse', 'warehouse_id', 'created_at', 'updated_at', 'created_by', 'modified_by']:
                            continue
                        if key == 'dynamic_data':
                            dyn_data = record.get('dynamic_data')
                            if isinstance(dyn_data, str):
                                try:
                                    import json
                                    dyn_data = json.loads(dyn_data)
                                except Exception:
                                    dyn_data = {}
                            if isinstance(dyn_data, dict):
                                for dyn_k in dyn_data.keys():
                                    if dyn_k not in editable_fields:
                                        invalid_fields.append(dyn_k)
                        elif key not in editable_fields:
                            invalid_fields.append(key)

                    if invalid_fields:
                        from accounts.audit_utils import log_audit_event
                        log_audit_event(
                            user=user,
                            module='docs',
                            action='UPDATE',
                            severity='warning',
                            target_model='Item',
                            target_object_id=item.id,
                            target_repr=f"تلاش غیرمجاز برای ویرایش فیلدهای قفل‌شده در ویرایش دسته‌ای کالا {item.fa_unic_code}",
                            warehouse=item.warehouse,
                            details={'invalid_fields': invalid_fields, 'item_id': item.id},
                            ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
                        )
                        return Response({
                            'error': f'شما مجوز ویرایش فیلدهای ({", ".join(invalid_fields)}) را در کالا {item.fa_unic_code} ندارید.',
                            'invalid_fields': invalid_fields
                        }, status=400)

                for key, value in record.items():
                    if key in BULK_UPDATE_PROTECTED_FIELDS:
                        continue
                    if key in valid_fields:
                        field = valid_fields[key]
                        if field.is_relation and field.many_to_one:
                            setattr(item, field.attname, value)
                            update_fields.add(field.name)
                        else:
                            setattr(item, key, value)
                            update_fields.add(key)
                items_to_update.append(item)
            
            if items_to_update and update_fields:
                for item in items_to_update:
                    item.updated_at = timezone.now()
                    item.modified_by = request.user
                update_fields.update(['updated_at', 'modified_by'])
                Item.objects.bulk_update(items_to_update, list(update_fields))

                from accounts.audit_utils import log_audit_event
                first_item = items_to_update[0] if items_to_update else None
                log_audit_event(
                    user=request.user,
                    warehouse=getattr(first_item, 'warehouse', None) if first_item else None,
                    module='docs',
                    action='BULK_UPDATE',
                    severity='info',
                    target_model='Item',
                    target_repr=f"ویرایش دسته‌ای {len(items_to_update)} کالا",
                    details={
                        'updated_count': len(items_to_update),
                        'fields_modified': list(update_fields),
                        'item_ids': [item.id for item in items_to_update[:50]]
                    },
                    ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
                )
            
            return Response({"success": f"Updated {len(items_to_update)} items"})
        except Exception as e:
            logger.exception("خطا در اجرای bulk_update: %s", e)
            return Response({"error": "خطایی در ذخیره‌سازی دسته‌ای کالاها رخ داد. لطفاً داده‌های ارسالی را بررسی فرمایید."}, status=400)

    @action(detail=False, methods=['post'])
    def bulk_assign(self, request):
        ids = request.data.get('item_ids', request.data.get('ids', []))
        
        # Field Counting Assignments
        field_assignee_id = request.data.get('field_assignee')
        supervisor_id = request.data.get('supervisor_assignee')
        manager_id = request.data.get('manager_assignee')
        
        # Document Phase Assignments
        doc_assignee_id = request.data.get('doc_assignee') or request.data.get('doc_assignee_id')
        doc_supervisor_id = request.data.get('doc_supervisor_assignee') or request.data.get('doc_supervisor_id')
        doc_manager_id = request.data.get('doc_manager_assignee') or request.data.get('doc_manager_id')

        
        field_status = request.data.get('field_status')
        doc_status = request.data.get('doc_status')
        force = request.data.get('force', False)
        select_all = request.data.get('select_all', False)
        filters_dict = request.data.get('filters', {})
        
        if select_all:
            from ..filters import ItemFilter
            items = ItemFilter(filters_dict or {}, queryset=self.get_queryset()).qs
            if filters_dict and 'search' in filters_dict and filters_dict['search']:
                s = filters_dict['search']
                items = items.filter(
                    Q(fa_unic_code__icontains=s) |
                    Q(description__icontains=s) |
                    Q(po__icontains=s) |
                    Q(pl__icontains=s) |
                    Q(pk_number__icontains=s) |
                    Q(my_tag__icontains=s)
                )
            if filters_dict and 'warehouse' in filters_dict and filters_dict['warehouse']:
                items = items.filter(warehouse_id=filters_dict['warehouse'])
        elif ids:
            items = self.get_queryset().filter(id__in=ids)
            ids_set = set(str(i) for i in ids)
            if len(items) != len(ids_set):
                missing_ids = ids_set - set(str(item.id) for item in items)
                if Item.objects.filter(id__in=missing_ids).exists():
                    return Response({"error": "دسترسی غیرمجاز: یک یا چند قلم کالا خارج از انبار مجاز شماست."}, status=403)
        else:
            items = self.get_queryset().none()
        
        # بررسی مجوز و اعتبارسنجی در صورت درخواست بازشماری
        if field_status == 'recount':
            if not request.user.is_superuser and not request.user.has_perm('accounts.perm_rec_recount'):
                return Response({'error': 'شما مجوز ثبت درخواست بازشماری (perm_rec_recount) را ندارید.'}, status=403)
                
            invalid_uncounted_count = items.filter(field_status__in=['waiting', 'counting', 'در انتظار شمارش']).count()
            if invalid_uncounted_count > 0:
                return Response({
                    'error': f'تعداد {invalid_uncounted_count} مورد از کالاهای انتخابی هنوز شمارش اولیه نشده‌اند و امکان درخواست بازشماری برای آن‌ها وجود ندارد.'
                }, status=400)
        
        # هشدار برای ارسال مجدد کالایی که CountTask فعال دارد
        if field_status == 'counting' and not force:
            from ..models import CountTask
            active_count_items = CountTask.objects.filter(
                item__in=items
            ).exclude(status='FINAL_APPROVED').values('item_id').distinct().count()
            if active_count_items > 0:
                return Response({
                    'warning': True,
                    'message': f'تعداد {active_count_items} قلم از کالاهای انتخاب شده در حال حاضر دارای فرآیند شمارش فعال هستند. آیا از ارجاع مجدد اطمینان دارید؟'
                }, status=200)
                
        # هشدار برای ارسال مجدد کالایی که DocTask فعال دارد
        if doc_status in ['checking', 'processing'] and not force:
            from ..models import DocTask
            active_doc_items = DocTask.objects.filter(
                item__in=items
            ).exclude(status='DOC_FINAL_APPROVED').values('item_id').distinct().count()
            if active_doc_items > 0:
                return Response({
                    'warning': True,
                    'message': f'تعداد {active_doc_items} قلم از کالاهای انتخاب شده در حال حاضر دارای فرآیند بررسی اسناد فعال هستند. آیا از ارجاع مجدد اطمینان دارید؟'
                }, status=200)
        
        from django.contrib.auth import get_user_model
        User = get_user_model()
        
        def _get_user_or_none(uid):
            if uid:
                try:
                    return User.objects.get(id=uid)
                except (User.DoesNotExist, ValueError):
                    return None
            return None

        counter_user = _get_user_or_none(field_assignee_id)
        manager_user = _get_user_or_none(manager_id)
        
        doc_worker_user = _get_user_or_none(doc_assignee_id)
        doc_manager_user = _get_user_or_none(doc_manager_id)
        
        supervisor_user = None
        skip_supervisor = False
        if supervisor_id:
            if str(supervisor_id) == 'skip':
                skip_supervisor = True
            else:
                supervisor_user = _get_user_or_none(supervisor_id)
                
        doc_supervisor_user = None
        doc_skip_supervisor = False
        if doc_supervisor_id:
            if str(doc_supervisor_id) == 'skip':
                doc_skip_supervisor = True
            else:
                doc_supervisor_user = _get_user_or_none(doc_supervisor_id)

        assignees = [u for u in [counter_user, supervisor_user, manager_user, doc_worker_user, doc_supervisor_user, doc_manager_user] if u is not None]
        if assignees and items.exists():
            from companies.access import get_user_allowed_companies
            item_wh_companies = list(items.values_list('warehouse__company_id', flat=True).distinct())
            has_no_company_items = any(cid is None for cid in item_wh_companies)
            item_company_ids = {cid for cid in item_wh_companies if cid is not None}

            for assignee in assignees:
                if assignee.is_superuser:
                    continue
                if not assignee.is_active:
                    return Response({'error': 'کاربر انتخابشده عضو شرکت این کالاها نیست.'}, status=400)
                if item_company_ids:
                    allowed_company_ids = set(get_user_allowed_companies(assignee).values_list('id', flat=True))
                    if not item_company_ids.issubset(allowed_company_ids):
                        return Response({'error': 'کاربر انتخابشده عضو شرکت این کالاها نیست.'}, status=400)
        
        update_data = {}
        if 'field_assignee' in request.data:
            if counter_user:
                update_data['field_assignee'] = f"{counter_user.first_name} {counter_user.last_name}".strip() or counter_user.username
            else:
                update_data['field_assignee'] = 'استخر عمومی'
                
        if 'doc_assignee' in request.data:
            if doc_worker_user:
                update_data['doc_assignee'] = f"{doc_worker_user.first_name} {doc_worker_user.last_name}".strip() or doc_worker_user.username
            else:
                update_data['doc_assignee'] = 'استخر عمومی'

        if field_status is not None:
            update_data['field_status'] = field_status
            if field_status == 'recount':
                update_data['has_conflict'] = True
        if doc_status is not None:
            update_data['doc_status'] = doc_status
            
        created_count_tasks = False
        created_doc_tasks = False

        with transaction.atomic():
            items_list = list(items.select_for_update())
            if not items_list:
                return Response({'status': 'success', 'updated': 0})

            if update_data:
                update_data['updated_at'] = timezone.now()
                update_data['modified_by'] = request.user
                items.update(**update_data)
                
            from warehouses.services import get_setting
            first_item = items_list[0] if items_list else None
            wh_id = first_item.warehouse_id if first_item else None
            
            # Create CountTasks if it's a field dispatch
            if field_status == 'counting':
                from ..models import CountTask
                from common.sync_models import soft_delete_queryset
                
                # باطل‌سازی تسک‌های فعال قبلی این کالاها برای جلوگیری از ایجاد تسک‌های تکراری در کارتابل
                soft_delete_queryset(CountTask.objects.filter(item__in=items_list).exclude(status='FINAL_APPROVED'))
                
                # بررسی تنظیم تایید سرپرست
                req_supervisor = get_setting('require_supervisor_approval', wh_id)
                
                tasks_to_create = []
                for item in items_list:
                    tasks_to_create.append(CountTask(
                        item=item,
                        counter=counter_user,
                        supervisor=supervisor_user if (req_supervisor and not skip_supervisor) else None,
                        skip_supervisor=skip_supervisor,
                        assigned_manager=manager_user,
                        status='PENDING_COUNT',
                        created_by=request.user,
                        modified_by=request.user
                    ))
                if tasks_to_create:
                    CountTask.objects.bulk_create(tasks_to_create)
                    created_count_tasks = True

            # Create or update CountTasks if it's a recount request
            if field_status == 'recount':
                from ..models import CountTask, CountTaskHistory
                from warehouses.services import get_setting
                
                for item in items_list:
                    latest_task = CountTask.objects.filter(item=item).order_by('-created_at').first()
                    req_supervisor = get_setting('require_supervisor_approval', item.warehouse_id)
                    target_status = 'MANAGER_REJECTED' if (req_supervisor and (latest_task and latest_task.supervisor)) else 'PENDING_COUNT'
                    
                    if latest_task:
                        prev_counted = latest_task.counted_balance
                        latest_task.status = target_status
                        latest_task.manager_note = 'درخواست بازشماری از تخصیص کالا (ثبت مغایرت)'
                        if target_status == 'PENDING_COUNT':
                            latest_task.counted_balance = None
                        latest_task.modified_by = request.user
                        latest_task.updated_at = timezone.now()
                        latest_task.save()
                        
                        CountTaskHistory.objects.create(
                            task=latest_task,
                            action_by=request.user,
                            action_type='MANAGER_REJECTED',
                            counted_balance=prev_counted,
                            note='درخواست بازشماری از صفحه تخصیص'
                        )
                    else:
                        new_task = CountTask.objects.create(
                            item=item,
                            status=target_status,
                            manager_note='درخواست بازشماری از تخصیص کالا (ثبت مغایرت)',
                            created_by=request.user,
                            modified_by=request.user
                        )
                        CountTaskHistory.objects.create(
                            task=new_task,
                            action_by=request.user,
                            action_type='MANAGER_REJECTED',
                            note='ایجاد تسک بازشماری از صفحه تخصیص'
                        )
                created_count_tasks = True

            # Create DocTasks if it's a document dispatch
            if doc_status == 'processing':
                from ..models import DocTask
                
                # حذف تسک‌های فعال قبلی این کالاها برای جلوگیری از ایجاد تسک‌های تکراری در کارتابل اسناد
                DocTask.objects.filter(item__in=items_list).exclude(status='DOC_FINAL_APPROVED').delete()
                
                # بررسی تنظیم تایید سرپرست اسناد
                req_doc_supervisor = get_setting('require_doc_supervisor_approval', wh_id)
                
                doc_tasks_to_create = []
                for item in items_list:
                    # تبدیل امن مقادیر اولیه فیلدهای مالی کالا
                    p_amount = item.price_amount
                    s_price = item.similar_unit_price
                    t_value = item.total_value
                    curr = item.currency if item.currency in ['IRR', 'USD', 'EUR', 'OTHER'] else None
                    inv_type = item.invoice_type if item.invoice_type in ['formal', 'domestic', 'foreign', 'consignment'] else None
                    
                    # صفحه و ردیف فاکتور
                    p_row = int(str(item.page_row).strip()) if item.page_row and str(item.page_row).strip().isdigit() else None
                    inv_page = int(str(item.invoice_page).strip()) if item.invoice_page and str(item.invoice_page).strip().isdigit() else None
                    
                    # مهر و امضا
                    stamp_val = bool(item.stamp) if isinstance(item.stamp, bool) else (str(item.stamp).lower() in ['true', '1', 'بله', 'دارد', 'yes'])
                    sign_val = bool(item.signature) if isinstance(item.signature, bool) else (str(item.signature).lower() in ['true', '1', 'بله', 'دارد', 'yes'])
                    
                    # تاریخ فاکتور
                    inv_date = _parse_date_flexible(item.invoice_date)

                    doc_tasks_to_create.append(DocTask(
                        item=item,
                        doc_worker=doc_worker_user,
                        doc_supervisor=doc_supervisor_user if (req_doc_supervisor and not doc_skip_supervisor) else None,
                        skip_supervisor=doc_skip_supervisor,
                        assigned_manager=doc_manager_user,
                        status='PENDING_DOC',
                        price_amount=p_amount,
                        similar_unit_price=s_price,
                        total_value=t_value,
                        currency=curr,
                        invoice_type=inv_type,
                        invoice_date=inv_date,
                        inv_rti_number=item.inv_rti_number or None,
                        added_rti_no=item.added_rti_no or None,
                        page_row=p_row,
                        invoice_page=inv_page,
                        doc_supplier=item.doc_supplier or None,
                        folder_address=item.folder_address or None,
                        stamp=stamp_val,
                        signature=sign_val,
                        created_by=request.user,
                        modified_by=request.user,
                        sync_id=uuid.uuid4()
                    ))
                if doc_tasks_to_create:
                    DocTask.objects.bulk_create(doc_tasks_to_create)
                    created_doc_tasks = True

        if created_count_tasks:
            broadcast_count_task_update()
        if created_doc_tasks:
            broadcast_doc_task_update()

        from accounts.audit_utils import log_audit_event
        log_audit_event(
            user=request.user,
            warehouse=getattr(first_item, 'warehouse', None) if first_item else None,
            module='dispatch',
            action='UPDATE',
            severity='info',
            target_model='Item',
            target_repr=f"تخصیص و ارجاع گروهی {len(items_list)} کالا",
            details={
                'items_count': len(items_list),
                'field_status': field_status,
                'doc_status': doc_status,
                'field_assignee': str(counter_user) if counter_user else update_data.get('field_assignee'),
                'supervisor_assignee': str(supervisor_user) if supervisor_user else ('عدم نیاز' if skip_supervisor else None),
                'manager_assignee': str(manager_user) if manager_user else None,
                'doc_assignee': str(doc_worker_user) if doc_worker_user else update_data.get('doc_assignee'),
                'doc_supervisor_assignee': str(doc_supervisor_user) if doc_supervisor_user else ('عدم نیاز' if doc_skip_supervisor else None),
                'doc_manager_assignee': str(doc_manager_user) if doc_manager_user else None,
            },
            ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
        )
            
        return Response({'status': 'success', 'updated': len(items_list)})

    @action(detail=False, methods=['post'])
    def bulk_tag(self, request):
        action = request.data.get('action')
        tag_val = (request.data.get('tag') or '').strip()
        item_ids = request.data.get('item_ids', [])
        select_all = request.data.get('select_all', False)
        filters_dict = request.data.get('filters', {})
        
        if action and (item_ids or select_all):
            if select_all:
                from ..filters import ItemFilter
                items = ItemFilter(filters_dict or {}, queryset=self.get_queryset()).qs
                if filters_dict and 'search' in filters_dict and filters_dict['search']:
                    s = filters_dict['search']
                    items = items.filter(
                        Q(fa_unic_code__icontains=s) |
                        Q(description__icontains=s) |
                        Q(po__icontains=s) |
                        Q(pl__icontains=s) |
                        Q(pk_number__icontains=s) |
                        Q(my_tag__icontains=s)
                    )
                if filters_dict and 'warehouse' in filters_dict and filters_dict['warehouse']:
                    items = items.filter(warehouse_id=filters_dict['warehouse'])
            else:
                items = self.get_queryset().filter(id__in=item_ids)
                item_ids_set = set(str(i) for i in item_ids)
                if len(items) != len(item_ids_set):
                    missing_ids = item_ids_set - set(str(item.id) for item in items)
                    if Item.objects.filter(id__in=missing_ids).exists():
                        return Response({"error": "دسترسی غیرمجاز: یک یا چند قلم کالا خارج از انبار مجاز شماست."}, status=403)
            updated_count = 0
            with transaction.atomic():
                for item in items:
                    current_tags = [t.strip() for t in (item.my_tag or '').split('،') if t.strip()]
                    if action == 'add' and tag_val:
                        if tag_val not in current_tags:
                            current_tags.append(tag_val)
                    elif action == 'remove' and tag_val:
                        current_tags = [t for t in current_tags if t != tag_val]
                    elif action == 'clear':
                        current_tags = []
                    
                    new_tag_str = '، '.join(current_tags)
                    if item.my_tag != new_tag_str:
                        item.my_tag = new_tag_str
                        item.updated_at = timezone.now()
                        item.modified_by = request.user
                        item.save(update_fields=['my_tag', 'updated_at', 'modified_by'])
                    updated_count += 1
            return Response({'status': 'success', 'updated': updated_count})

        updates = request.data.get('updates', [])
        updated_count = 0
        if updates:
            with transaction.atomic():
                for up in updates:
                    item_id = up.get('id')
                    tag = up.get('my_tag', '')
                    if item_id:
                        self.get_queryset().filter(id=item_id).update(
                            my_tag=tag,
                            updated_at=timezone.now(), 
                            modified_by=request.user
                        )
                        updated_count += 1
        else:
            # Fallback for old bulk_tag format if any
            ids = request.data.get('ids', [])
            tag = request.data.get('tag')
            items = self.get_queryset().filter(id__in=ids)
            if tag == 'conflict':
                items.update(has_conflict=True, updated_at=timezone.now(), modified_by=request.user)
            updated_count = items.count()
            
        return Response({'status': 'success', 'updated': updated_count})

    def get_expected_fields(self, warehouse_id=None):
        fields = {}
        for f in Item._meta.fields:
            name = f.name
            if name == 'dynamic_data': continue
            if name == 'warehouse':
                fields['warehouse'] = 'warehouse'
            elif name in ['created_by', 'modified_by']:
                fields[name] = name
            else:
                fields[name] = name
                
        if warehouse_id:
            from ..models import ItemFieldDefinition
            dynamic_defs = ItemFieldDefinition.objects.filter(warehouse_id=warehouse_id, is_active=True)
            for d in dynamic_defs:
                fields[d.name] = d.name
                fields[d.name.lower()] = d.name
                
        return fields

    @action(detail=False, methods=['post'])
    def parse_headers(self, request):
        from warehouses.services import get_setting
        file_obj = request.FILES.get('file')
        warehouse_id = request.data.get('warehouse_id')
        
        if not file_obj:
            return Response({'error': 'هیچ فایلی ارسال نشده است.'}, status=400)

        if warehouse_id and not self._can_write_warehouse(warehouse_id):
            return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)
            
        temp_path = None
        try:
            fd, temp_path = tempfile.mkstemp(suffix='.xlsx')
            with os.fdopen(fd, 'wb') as f:
                for chunk in file_obj.chunks():
                    f.write(chunk)
            
            wb = openpyxl.load_workbook(temp_path, read_only=True, data_only=True)
            ws = wb.active
            
            # Read only the first row for headers to maximize speed
            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
            # Convert to lower case to remove case sensitivity
            raw_headers = [str(val).strip().lower() if val is not None else '' for val in first_row]
            wb.close()
            
            expected_fields = self.get_expected_fields(warehouse_id)
            
            found_fields = []
            for h in raw_headers:
                if h in expected_fields:
                    found_fields.append(expected_fields[h])
            
            all_expected = set(expected_fields.values())
            missing_fields = list(all_expected - set(found_fields))
            
            # Check for sensitive fields
            restricted_fields = get_setting('SENSITIVE_EXCEL_FIELDS', warehouse_id) or ['doc_status', 'field_status', 'tag_status']
            has_restricted = any(f in found_fields for f in restricted_fields)
            is_superuser = request.user.is_superuser if request.user and request.user.is_authenticated else False
            
            return Response({
                'status': 'success', 
                'found_fields': list(set(found_fields)), 
                'missing_fields': list(set(missing_fields)),
                'has_restricted_fields': has_restricted,
                'is_superuser': is_superuser
            })
        except Exception as e:
            logger.exception("Error parsing excel headers")
            return Response({'error': 'خطا در خواندن فایل اکسل. لطفاً از معتبر بودن و فرمت استاندارد فایل اطمینان حاصل کنید.'}, status=400)
        finally:
            if temp_path and os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass

    @action(detail=False, methods=['post'])
    def cancel_import(self, request):
        import_id = request.data.get('import_id')
        if not import_id:
            return Response({'error': 'شناسه فرآیند الزامی است.'}, status=400)
            
        import re
        if not re.match(r'^[a-zA-Z0-9_-]{1,64}$', str(import_id)):
            return Response({'error': 'شناسه فرآیند نامعتبر است.'}, status=400)

        import_log = ImportLog.objects.filter(import_id=import_id).first()
        if import_log:
            if not self._can_manage_import_log(import_log):
                return Response({'error': 'شما دسترسی لازم برای لغو این فرآیند را ندارید.'}, status=403)
        else:
            cached_owner = cache.get(f"import_owner_{import_id}")
            if not (request.user.is_superuser or (cached_owner is not None and cached_owner == request.user.id)):
                return Response({'error': 'شما دسترسی لازم برای لغو این فرآیند را ندارید.'}, status=403)

        cache.set(f"cancel_import_{import_id}", True, timeout=3600)
        return Response({'status': 'cancelled'})

    @action(detail=False, methods=['get'])
    def download_import_log(self, request):
        import_id = request.query_params.get('import_id')
        if not import_id:
            return Response({'error': 'شناسه فرآیند الزامی است.'}, status=400)

        import re
        if not re.match(r'^[a-zA-Z0-9_-]{1,64}$', str(import_id)):
            return Response({'error': 'شناسه فرآیند نامعتبر است.'}, status=400)

        import_log = ImportLog.objects.filter(import_id=import_id).first()
        if not import_log:
            return Response({'error': 'فایل لاگ یافت نشد یا منقضی شده است.'}, status=404)

        # کنترل دسترسی کاربر به لاگ (سوپریوزر، ثبت‌کننده لاگ یا هلپر مدیریت لاگ)
        is_owner = (import_log.imported_by_id == request.user.id)
        if not (request.user.is_superuser or is_owner or self._can_manage_import_log(import_log)):
            return Response({'error': 'شما دسترسی لازم برای دریافت این فایل لاگ را ندارید.'}, status=403)

        temp_dir = os.path.abspath(tempfile.gettempdir())
        file_path = os.path.abspath(os.path.join(temp_dir, f"import_log_{import_id}.xlsx"))

        # محافظت قطعی در برابر Path Traversal
        try:
            if os.path.commonpath([temp_dir, file_path]) != temp_dir:
                return Response({'error': 'مسیر فایل نامعتبر است.'}, status=400)
        except ValueError:
            return Response({'error': 'مسیر فایل نامعتبر است.'}, status=400)

        if not os.path.exists(file_path):
            return Response({'error': 'فایل لاگ یافت نشد یا منقضی شده است.'}, status=404)

        with open(file_path, 'rb') as f:
            file_data = f.read()

        response = HttpResponse(
            file_data,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="import_log_{import_id}.xlsx"'
        return response

    @action(detail=False, methods=['post'])
    def import_excel(self, request):
        file_obj = request.FILES.get('file')
        warehouse_id = request.data.get('warehouse_id')

        if warehouse_id is not None and str(warehouse_id).strip() != '':
            try:
                warehouse_id_val = int(warehouse_id)
            except (ValueError, TypeError):
                return Response({'error': 'شناسه انبار نامعتبر است.'}, status=400)

            if not self._can_write_warehouse(warehouse_id_val):
                return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)
            warehouse_id = warehouse_id_val
        else:
            warehouse_id = None
            if not self._can_write_warehouse(None):
                return Response({'error': 'تعیین انبار مجاز برای کاربران با دسترسی محدود الزامی است.'}, status=403)

        from warehouses.services import get_setting
        sys_conflict_strategy = get_setting('default_conflict_strategy', warehouse_id)

        conflict_strategy = request.data.get('conflict_strategy') or sys_conflict_strategy
        import_tag = request.data.get('import_tag', '')
        import_id = request.data.get('import_id', '')
        is_pre_counted_raw = request.data.get('is_pre_counted', False)
        is_pre_counted = str(is_pre_counted_raw).strip().lower() in ['true', '1', 'yes']

        if is_pre_counted:
            can_pre_count = (
                request.user.is_superuser or 
                request.user.has_perm('accounts.perm_inventory_finalize')
            )
            if not can_pre_count:
                return Response({'error': 'شما مجوز ثبت مستقیم اقلام در وضعیت از قبل شمرده‌شده را ندارید.'}, status=403)

        is_doc_pre_approved_raw = request.data.get('is_doc_pre_approved', False)
        is_doc_pre_approved = str(is_doc_pre_approved_raw).strip().lower() in ['true', '1', 'yes']

        if is_doc_pre_approved:
            can_pre_approve_docs = (
                request.user.is_superuser or 
                request.user.has_perm('accounts.perm_doc_approve_action')
            )
            if not can_pre_approve_docs:
                return Response({'error': 'شما مجوز ثبت مستقیم اقلام در وضعیت اسناد تاییدشده را ندارید.'}, status=403)

        if not file_obj:
            return Response({'error': 'هیچ فایلی ارسال نشده است.'}, status=400)
            
        # Save file to temp to allow streaming response generator to read it
        fd, temp_path = tempfile.mkstemp(suffix='.xlsx')
        with os.fdopen(fd, 'wb') as f:
            for chunk in file_obj.chunks():
                f.write(chunk)
                
        user_id = request.user.id if request.user.is_authenticated else None
        original_file_name = file_obj.name
        allowed_warehouse_ids = self._allowed_warehouse_ids()

        if import_id and request.user.is_authenticated:
            cache.set(f"import_owner_{import_id}", request.user.id, timeout=3600)

        def worker(q):
            try:
                q.put(json.dumps({"type": "info", "msg": ">> فایل با موفقیت دریافت شد. در حال خواندن محتوا..."}) + "\n")
                wb = openpyxl.load_workbook(temp_path, data_only=True)
                ws = wb.active
                
                raw_headers = [str(cell.value).strip().lower() if cell.value else '' for cell in ws[1]]
                expected_fields = self.get_expected_fields(warehouse_id)
                
                dynamic_fields_keys = []
                if warehouse_id:
                    dynamic_fields_keys = [d.name for d in ItemFieldDefinition.objects.filter(warehouse_id=warehouse_id, is_active=True)]
                
                found_fields = []
                col_indices = {}
                
                for idx, h in enumerate(raw_headers):
                    if h in expected_fields:
                        found_fields.append(expected_fields[h])
                        col_indices[expected_fields[h]] = idx
                
                all_expected_db_fields = set(expected_fields.values())
                missing_fields = list(all_expected_db_fields - set(found_fields))
                
                if 'fa_unic_code' not in found_fields and 'id' not in found_fields:
                    q.put(json.dumps({"type": "err", "msg": ">> خطای حیاتی: ستون 'fa_unic_code' یا 'id' در فایل اکسل یافت نشد. یکی از این ستون‌ها اجباری است."}) + "\n")
                    q.put(json.dumps({
                        "type": "summary",
                        "status": "failed",
                        "created": 0, "updated": 0, "skipped": 0, "failed": 0,
                        "found_fields": list(set(found_fields)),
                        "missing_fields": list(set(missing_fields)),
                        "error_details": [{"row": 0, "code": "HEADER", "error": "ستون fa_unic_code یا id یافت نشد"}]
                    }) + "\n")
                    return
                
                created = 0
                updated = 0
                skipped = 0
                failed = 0
                error_details = []

                def _mark_item_as_counted(item, counted_qty, note_reason):
                    c_task = CountTask.all_objects.select_for_update().filter(item=item).first()
                    prev_status = c_task.status if c_task else None
                    prev_balance = c_task.counted_balance if c_task else None

                    if c_task:
                        c_task.status = 'FINAL_APPROVED'
                        c_task.counted_balance = counted_qty
                        c_task.manager_note = note_reason
                        c_task.modified_by_id = user_id
                        c_task.is_deleted = False
                        c_task.updated_at = timezone.now()
                        c_task.save(update_fields=['status', 'counted_balance', 'manager_note', 'modified_by', 'updated_at', 'is_deleted'])
                    else:
                        c_task = CountTask.objects.create(
                            item=item,
                            status='FINAL_APPROVED',
                            counted_balance=counted_qty,
                            manager_note=note_reason,
                            created_by_id=user_id,
                            modified_by_id=user_id
                        )

                    audit_note = note_reason
                    if prev_status and prev_status != 'FINAL_APPROVED':
                        audit_note = f"{note_reason} (تغییر وضعیت از: {prev_status}، موجودی قبلی: {prev_balance})"

                    CountTaskHistory.objects.create(
                        task=c_task,
                        action_by_id=user_id,
                        action_type='FINAL_APPROVED',
                        counted_balance=counted_qty,
                        note=f"[IMPORT:{import_id}] {audit_note}" if import_id else audit_note
                    )
                    return c_task

                def _mark_item_docs_as_approved(item, financial_data, note_reason):
                    d_task = DocTask.objects.select_for_update().filter(item=item).first()
                    prev_status = d_task.status if d_task else None
                    prev_snapshot = _create_doc_task_snapshot(d_task) if d_task else None

                    # ثبت اسنپ‌شات وضعیت اولیه در صورتی که تسک از قبل وجود داشته باشد
                    if d_task and prev_snapshot and prev_status != 'DOC_FINAL_APPROVED':
                        if not d_task.history.filter(action_type=prev_status).exists():
                            DocTaskHistory.objects.create(
                                task=d_task,
                                action_by_id=d_task.modified_by_id or d_task.created_by_id or user_id,
                                action_type=prev_status or 'PENDING_DOC',
                                note=f"اسنپ‌شات وضعیت اولیه قبل از به‌روزرسانی با اکسل ({note_reason})",
                                data_snapshot=prev_snapshot
                            )

                    def _get_val(key):
                        # ۱. اولویت اول: مقدار موجود در فایل اکسل جاری
                        if key in financial_data and financial_data[key] is not None and str(financial_data[key]).strip() != '':
                            return financial_data[key]
                        # ۲. اولویت دوم: حفظ مقدار قبلی ثبت‌شده توسط کاربر در پرونده تسک اسناد
                        if d_task and getattr(d_task, key, None) is not None and str(getattr(d_task, key, None)).strip() != '':
                            return getattr(d_task, key, None)
                        # ۳. اولویت سوم: استفاده از مقدار موجود در مدل کالا
                        return getattr(item, key, None)

                    # بررسی هوشمند وضعیت مهر اسناد با حفظ مقدار پیشین تسک
                    raw_stamp = _get_val('stamp')
                    parsed_stamp = None
                    if raw_stamp is not None:
                        if isinstance(raw_stamp, bool):
                            parsed_stamp = raw_stamp
                        elif str(raw_stamp).strip() in ['دارد', '1', 'true', 'True', 'بله']:
                            parsed_stamp = True
                        elif str(raw_stamp).strip() in ['ندارد', '0', 'false', 'False', 'خیر']:
                            parsed_stamp = False

                    final_stamp = parsed_stamp if parsed_stamp is not None else (d_task.stamp if d_task else False)

                    # بررسی هوشمند وضعیت امضای اسناد با حفظ مقدار پیشین تسک
                    raw_sig = _get_val('signature')
                    parsed_sig = None
                    if raw_sig is not None:
                        if isinstance(raw_sig, bool):
                            parsed_sig = raw_sig
                        elif str(raw_sig).strip() in ['دارد', '1', 'true', 'True', 'بله']:
                            parsed_sig = True
                        elif str(raw_sig).strip() in ['ندارد', '0', 'false', 'False', 'خیر']:
                            parsed_sig = False

                    final_sig = parsed_sig if parsed_sig is not None else (d_task.signature if d_task else False)

                    task_fields = {
                        'status': 'DOC_FINAL_APPROVED',
                        'manager_note': note_reason,
                        'modified_by_id': user_id,
                        'updated_at': timezone.now(),
                        'added_rti_no': str(_get_val('added_rti_no'))[:100] if _get_val('added_rti_no') else None,
                        'inv_rti_number': str(_get_val('inv_rti_number'))[:100] if _get_val('inv_rti_number') else None,
                        'invoice_type': _clean_invoice_type(_get_val('invoice_type')),
                        'invoice_date': _parse_date_flexible(_get_val('invoice_date')),
                        'invoice_page': _parse_positive_int(_get_val('invoice_page')),
                        'page_row': _parse_positive_int(_get_val('page_row')),
                        'doc_supplier': str(_get_val('doc_supplier'))[:255] if _get_val('doc_supplier') else None,
                        'total_value': _parse_decimal(_get_val('total_value')),
                        'price_amount': _parse_decimal(_get_val('price_amount')),
                        'similar_unit_price': _parse_decimal(_get_val('similar_unit_price')),
                        'currency': _clean_currency(_get_val('currency')),
                        'folder_address': str(_get_val('folder_address'))[:500] if _get_val('folder_address') else None,
                        'stamp': final_stamp,
                        'signature': final_sig,
                    }

                    if d_task:
                        for k, v in task_fields.items():
                            setattr(d_task, k, v)
                        d_task.save(update_fields=list(task_fields.keys()))
                    else:
                        d_task = DocTask.objects.create(
                            item=item,
                            created_by_id=user_id,
                            **task_fields
                        )

                    audit_doc_note = note_reason
                    if prev_status and prev_status != 'DOC_FINAL_APPROVED':
                        audit_doc_note = f"{note_reason} (تغییر وضعیت از: {prev_status})"

                    DocTaskHistory.objects.create(
                        task=d_task,
                        action_by_id=user_id,
                        action_type='DOC_FINAL_APPROVED',
                        note=f"[IMPORT:{import_id}] {audit_doc_note}" if import_id else audit_doc_note,
                        data_snapshot=_create_doc_task_snapshot(d_task)
                    )
                    return d_task
                
                # Fetch sensitive fields settings and user permissions
                restricted_fields = get_setting('SENSITIVE_EXCEL_FIELDS', warehouse_id) or ['doc_status', 'field_status', 'tag_status']
                is_superuser = request.user.is_superuser if request.user and request.user.is_authenticated else False
                
                # Setup output workbook for log
                out_wb = openpyxl.Workbook(write_only=True)
                out_ws = out_wb.create_sheet('Log')
                log_headers = list(raw_headers) + ['وضعیت پردازش', 'جزئیات پیام']
                out_ws.append(log_headers)
                
                FILL_COLORS = {
                    'created': PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid"),
                    'updated': PatternFill(start_color="BFDBFE", end_color="BFDBFE", fill_type="solid"),
                    'warn': PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid"),
                    'err': PatternFill(start_color="FECDD3", end_color="FECDD3", fill_type="solid")
                }
                
                def append_colored_row(row_data_tuple, status_code, message):
                    row_cells = []
                    for val in list(row_data_tuple) + [status_code, message]:
                        c = WriteOnlyCell(out_ws, value=val)
                        if status_code in FILL_COLORS:
                            c.fill = FILL_COLORS[status_code]
                        row_cells.append(c)
                    out_ws.append(row_cells)

                q.put(json.dumps({"type": "info", "msg": ">> پردازش سطرها شروع شد..."}) + "\n")
                
                try:
                    history_records = []
                    valid_fields_dict = {f.name: f for f in Item._meta.fields}
                    
                    with transaction.atomic():
                        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                            # Check cancel flag
                            if import_id and cache.get(f"cancel_import_{import_id}"):
                                raise Exception("CANCELED_BY_USER")

                            row_data = {}
                            for db_field, col_idx in col_indices.items():
                                row_data[db_field] = row[col_idx]
                            
                            raw_fa = row_data.get('fa_unic_code')
                            fa_unic_code = str(raw_fa).strip() if raw_fa is not None and str(raw_fa).strip() != '' else None
                            
                            raw_id = row_data.pop('id', None)
                            item_id = str(raw_id).strip() if raw_id is not None and str(raw_id).strip() != '' else None
                            
                            if not fa_unic_code and not item_id:
                                failed += 1
                                err_msg = "ستون fa_unic_code و id خالی است"
                                error_details.append({"row": row_idx, "code": "N/A", "error": err_msg})
                                append_colored_row(row, 'err', err_msg)
                                q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا: {err_msg}"}) + "\n")
                                continue
                            
                            # Ignore Excel data for core system fields
                            for ignore_field in ['created_at', 'updated_at', 'created_by', 'modified_by', 'sync_id', 'is_deleted']:
                                row_data.pop(ignore_field, None)
                                
                            # Ignore sensitive fields if not superuser
                            if not is_superuser:
                                for ignore_field in restricted_fields:
                                    row_data.pop(ignore_field, None)
                                
                            if 'hov_date' in row_data:
                                date_val = row_data['hov_date']
                                if isinstance(date_val, datetime):
                                    row_data['hov_date'] = date_val.date()
                                else:
                                    row_data['hov_date'] = None
                                    
                            excel_tag = row_data.pop('my_tag', '')
                            if not excel_tag: excel_tag = ''
                            excel_tag = str(excel_tag).strip().replace(',', '،')
                            
                            final_tags = []
                            if excel_tag:
                                final_tags.extend([t.strip() for t in excel_tag.split('،') if t.strip()])
                            if import_tag:
                                import_tag_clean = import_tag.replace(',', '،')
                                final_tags.extend([t.strip() for t in import_tag_clean.split('،') if t.strip()])
                            
                            unique_tags = list(set(final_tags))
                            if unique_tags:
                                row_data['my_tag'] = '،'.join(unique_tags)
                            else:
                                row_data['my_tag'] = ''

                            warehouse_str = row_data.pop('warehouse', None)
                            target_warehouse_id = int(warehouse_id) if warehouse_id else None
                            
                            if warehouse_str:
                                wh_str = str(warehouse_str).strip()
                                wh = Warehouse.objects.filter(id__in=allowed_warehouse_ids).filter(Q(name__iexact=wh_str) | Q(code__iexact=wh_str)).first()
                                if not wh:
                                    failed += 1
                                    err_msg = f"انبار با نام یا کد '{wh_str}' یافت نشد."
                                    error_details.append({"row": row_idx, "code": fa_unic_code, "error": err_msg})
                                    append_colored_row(row, 'err', err_msg)
                                    q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا در {fa_unic_code}: {err_msg}"}) + "\n")
                                    continue

                                # اگر فرآیند برای انبار مشخصی اجرا شده، ردیف نباید به انبار دیگری اشاره کند
                                if warehouse_id and wh.id != int(warehouse_id):
                                    failed += 1
                                    err_msg = f"مغایرت انبار: فرآیند برای انبار شناسه {warehouse_id} آغاز شده اما این ردیف انبار '{wh.name}' را مشخص کرده است."
                                    error_details.append({"row": row_idx, "code": fa_unic_code, "error": err_msg})
                                    append_colored_row(row, 'err', err_msg)
                                    q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا در {fa_unic_code}: {err_msg}"}) + "\n")
                                    continue

                                target_warehouse_id = wh.id

                            if target_warehouse_id and target_warehouse_id not in allowed_warehouse_ids:
                                failed += 1
                                err_msg = "شما دسترسی لازم به انبار مشخص‌شده در این ردیف را ندارید."
                                error_details.append({"row": row_idx, "code": fa_unic_code, "error": err_msg})
                                append_colored_row(row, 'err', err_msg)
                                q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا در {fa_unic_code}: {err_msg}"}) + "\n")
                                continue

                            # Extract dynamic data
                            dynamic_data_updates = {}
                            if target_warehouse_id:
                                for d_key in dynamic_fields_keys:
                                    if d_key in row_data and row_data[d_key] is not None:
                                        dynamic_data_updates[d_key] = row_data.pop(d_key)

                            if conflict_strategy == 'replace':
                                defaults = {}
                                for k, v in row_data.items():
                                    if k == 'fa_unic_code':
                                        continue
                                    if v is None:
                                        f = valid_fields_dict.get(k)
                                        if f and f.null:
                                            defaults[k] = None
                                        elif f and f.has_default():
                                            defaults[k] = f.default if not callable(f.default) else f.default()
                                        elif f and isinstance(f, (models.CharField, models.TextField)):
                                            defaults[k] = ''
                                    else:
                                        defaults[k] = v
                            else:
                                defaults = {k: v for k, v in row_data.items() if k != 'fa_unic_code' and v is not None}
                            if target_warehouse_id: defaults['warehouse_id'] = target_warehouse_id
                            if user_id: defaults['modified_by_id'] = user_id

                            item_display_code = fa_unic_code or f"ID:{item_id}"
                            existing_item = None
                            if item_id:
                                try:
                                    clean_id = int(float(item_id))
                                except ValueError:
                                    clean_id = item_id

                                existing_item = Item.objects.filter(id=clean_id).first()
                                if existing_item:
                                    # جلوگیری از سرقت یا جابجایی ناخواسته کالا بین انبارها
                                    if target_warehouse_id and existing_item.warehouse_id != target_warehouse_id:
                                        failed += 1
                                        err_msg = f"مغایرت انبار: کالای شناسه {clean_id} متعلق به انبار دیگری است و امکان جابجایی یا ویرایش آن از این انبار وجود ندارد."
                                        error_details.append({"row": row_idx, "code": item_display_code, "error": err_msg})
                                        append_colored_row(row, 'err', err_msg)
                                        q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا: {err_msg}"}) + "\n")
                                        continue

                                    if existing_item.warehouse_id not in allowed_warehouse_ids:
                                        failed += 1
                                        err_msg = f"عدم دسترسی: شما به انبار کالای شناسه {clean_id} دسترسی ندارید."
                                        error_details.append({"row": row_idx, "code": item_display_code, "error": err_msg})
                                        append_colored_row(row, 'err', err_msg)
                                        q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا: {err_msg}"}) + "\n")
                                        continue

                                    # Check if ID points to a different fa_unic_code
                                    if fa_unic_code:
                                        db_code = str(existing_item.fa_unic_code).strip() if existing_item.fa_unic_code else ""
                                        if db_code and db_code != fa_unic_code:
                                            failed += 1
                                            err_msg = f"مغایرت شناسه: کالا با id={clean_id} دارای کد {db_code} است ولی اکسل کد {fa_unic_code} را ارسال کرده است."
                                            error_details.append({"row": row_idx, "code": fa_unic_code, "error": err_msg})
                                            append_colored_row(row, 'err', err_msg)
                                            q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا: {err_msg}"}) + "\n")
                                            continue

                            # هرگز انبار یک کالای موجود را از طریق فایل اکسل تغییر ندهید (جلوگیری از سرقت یا جابجایی انبار)
                            if existing_item:
                                defaults.pop('warehouse_id', None)
                            
                            if not existing_item and fa_unic_code and target_warehouse_id:
                                existing_item = Item.objects.filter(fa_unic_code=fa_unic_code, warehouse_id=target_warehouse_id).first()
                                # Check if fa_unic_code points to a different ID
                                if existing_item and item_id:
                                    if str(existing_item.id) != str(item_id).split('.')[0]: # split to handle '1.0'
                                        failed += 1
                                        err_msg = f"مغایرت شناسه: کد {fa_unic_code} متعلق به id={existing_item.id} است ولی اکسل id={item_id} را ارسال کرده است."
                                        error_details.append({"row": row_idx, "code": fa_unic_code, "error": err_msg})
                                        append_colored_row(row, 'err', err_msg)
                                        q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا: {err_msg}"}) + "\n")
                                        continue
                                
                            # رکورد حذف‌نرم با همین کد؟ (unique_together هنوز فعال است)
                            # → باید احیا شود، وگرنه INSERT خطای دیتابیس می‌دهد.
                            resurrect_tombstone = None
                            if not existing_item and fa_unic_code and target_warehouse_id:
                                resurrect_tombstone = Item.all_objects.filter(
                                    fa_unic_code=fa_unic_code,
                                    warehouse_id=target_warehouse_id,
                                    is_deleted=True,
                                ).first()

                            item_display_code = fa_unic_code or f"ID:{item_id}"
                            
                            # Build merged dynamic_data for this item if needed
                            final_dynamic_data = {}
                            if existing_item and existing_item.dynamic_data:
                                final_dynamic_data = existing_item.dynamic_data.copy()
                            if dynamic_data_updates:
                                final_dynamic_data.update(dynamic_data_updates)
                            
                            if final_dynamic_data:
                                defaults['dynamic_data'] = final_dynamic_data

                            if is_pre_counted:
                                defaults['field_status'] = 'done'
                                ref_item = existing_item or resurrect_tombstone
                                raw_inv = row_data.get('inventory')
                                if (raw_inv is None or str(raw_inv).strip() == '') and ref_item:
                                    inv_val = ref_item.inventory
                                    defaults['inventory'] = ref_item.inventory
                                else:
                                    inv_val = defaults.get('inventory')

                                raw_bal = row_data.get('bal4miv')
                                if (raw_bal is None or str(raw_bal).strip() == '') and ref_item:
                                    bal_val = ref_item.bal4miv
                                    defaults['bal4miv'] = ref_item.bal4miv
                                else:
                                    bal_val = defaults.get('bal4miv')

                                c_dec = _parse_decimal(inv_val)
                                s_dec = _parse_decimal(bal_val)
                                if c_dec is not None and s_dec is not None:
                                    defaults['has_conflict'] = (c_dec != s_dec)
                                elif c_dec is not None:
                                    defaults['has_conflict'] = (c_dec != Decimal('0'))
                                else:
                                    defaults['has_conflict'] = False

                            if is_doc_pre_approved:
                                defaults['doc_status'] = 'done'
                                
                            try:
                                with transaction.atomic():
                                    if resurrect_tombstone:
                                        # احیای رکورد حذف‌نرم: از دید کاربر «رکورد جدید» است،
                                        # اما ردیف قبلی (با همان sync_id) به‌روزرسانی و زنده می‌شود.
                                        from django.core.serializers.json import DjangoJSONEncoder
                                        old_state = model_to_dict(resurrect_tombstone)
                                        old_state_json = json.loads(json.dumps(old_state, cls=DjangoJSONEncoder))

                                        if user_id: defaults['created_by_id'] = user_id
                                        if 'tag_status' not in defaults:
                                            defaults['tag_status'] = 'چاپ نشده'
                                        Item.all_objects.filter(id=resurrect_tombstone.id).update(
                                            fa_unic_code=fa_unic_code,
                                            is_deleted=False,
                                            updated_at=timezone.now(),
                                            **defaults,
                                        )
                                        _restore_item_photos_cascade([resurrect_tombstone.id])
                                        history_records.append(ImportHistory(item=resurrect_tombstone, action='update', previous_state=old_state_json))

                                        if is_pre_counted:
                                            counted_qty = defaults.get('inventory', resurrect_tombstone.inventory)
                                            _mark_item_as_counted(resurrect_tombstone, counted_qty, f"ثبت مستقیم اقلام شمرده‌شده از اکسل ({original_file_name})")

                                        if is_doc_pre_approved:
                                            _mark_item_docs_as_approved(resurrect_tombstone, defaults, f"ثبت مستقیم اسناد تاییدشده از اکسل ({original_file_name})")

                                        created += 1
                                        append_colored_row(row, 'created', 'ثبت رکورد جدید (احیای رکورد حذف‌شده)')
                                        q.put(json.dumps({"type": "created", "msg": f"[ردیف {row_idx}] ثبت رکورد جدید (احیا): {fa_unic_code}"}) + "\n")
                                    elif existing_item:
                                        # Append new tags to existing item's tags
                                        if existing_item.my_tag:
                                            existing_tags = [t.strip() for t in existing_item.my_tag.split('،') if t.strip()]
                                            new_tags = [t.strip() for t in defaults.get('my_tag', '').split('،') if t.strip()]
                                            combined_tags = list(set(existing_tags + new_tags))
                                            defaults['my_tag'] = '،'.join(combined_tags) if combined_tags else ''

                                        if conflict_strategy == 'ignore':
                                            skipped += 1
                                            append_colored_row(row, 'warn', 'نادیده گرفتن رکورد تکراری')
                                            q.put(json.dumps({"type": "warn", "msg": f"[ردیف {row_idx}] نادیده گرفتن رکورد تکراری: {item_display_code}"}) + "\n")
                                            continue
                                        elif conflict_strategy == 'log':
                                            error_details.append({"row": row_idx, "code": item_display_code, "error": "تداخل رکورد (ثبت در لاگ)"})
                                            skipped += 1
                                            append_colored_row(row, 'warn', 'تداخل رکورد (ثبت در لاگ)')
                                            q.put(json.dumps({"type": "warn", "msg": f"[ردیف {row_idx}] تداخل رکورد: {item_display_code}"}) + "\n")
                                            continue
                                        elif conflict_strategy == 'update_empty':
                                            def _is_field_empty(val):
                                                return val is None or (isinstance(val, str) and val.strip() == '')

                                            new_defaults = {
                                                k: v for k, v in defaults.items()
                                                if _is_field_empty(getattr(existing_item, k, None)) and not _is_field_empty(v)
                                            }
                                            # Always update fa_unic_code if empty
                                            if fa_unic_code and _is_field_empty(existing_item.fa_unic_code):
                                                new_defaults['fa_unic_code'] = fa_unic_code
                                            # Always update tag since we append them
                                            if defaults.get('my_tag') and defaults.get('my_tag') != existing_item.my_tag:
                                                new_defaults['my_tag'] = defaults['my_tag']

                                            if is_pre_counted:
                                                new_defaults['field_status'] = 'done'
                                                final_inv = new_defaults.get('inventory', existing_item.inventory)
                                                final_bal = new_defaults.get('bal4miv', existing_item.bal4miv)
                                                c_dec = _parse_decimal(final_inv)
                                                s_dec = _parse_decimal(final_bal)
                                                if c_dec is not None and s_dec is not None:
                                                    new_defaults['has_conflict'] = (c_dec != s_dec)
                                                elif c_dec is not None:
                                                    new_defaults['has_conflict'] = (c_dec != Decimal('0'))
                                                else:
                                                    new_defaults['has_conflict'] = False

                                            if is_doc_pre_approved:
                                                new_defaults['doc_status'] = 'done'

                                            if new_defaults:
                                                from django.core.serializers.json import DjangoJSONEncoder
                                                old_state = model_to_dict(existing_item)
                                                old_state_json = json.loads(json.dumps(old_state, cls=DjangoJSONEncoder))

                                                Item.objects.filter(id=existing_item.id).update(**{**new_defaults, 'updated_at': timezone.now()})
                                                history_records.append(ImportHistory(item=existing_item, action='update', previous_state=old_state_json))

                                                if is_pre_counted:
                                                    counted_qty = new_defaults.get('inventory', getattr(existing_item, 'inventory', None))
                                                    _mark_item_as_counted(existing_item, counted_qty, f"به‌روزرسانی مستقیم اقلام شمرده‌شده از اکسل ({original_file_name})")

                                                if is_doc_pre_approved:
                                                    _mark_item_docs_as_approved(existing_item, new_defaults, f"به‌روزرسانی مستقیم اسناد تاییدشده از اکسل ({original_file_name})")

                                                updated += 1
                                                append_colored_row(row, 'updated', 'تکمیل نواقص رکورد')
                                                q.put(json.dumps({"type": "updated", "msg": f"[ردیف {row_idx}] تکمیل نواقص رکورد: {item_display_code}"}) + "\n")
                                            else:
                                                skipped += 1
                                                append_colored_row(row, 'warn', 'بدون نقص، نادیده گرفته شد')
                                                q.put(json.dumps({"type": "warn", "msg": f"[ردیف {row_idx}] بدون نقص، نادیده گرفته شد: {item_display_code}"}) + "\n")
                                        elif conflict_strategy == 'replace':
                                            if fa_unic_code: defaults['fa_unic_code'] = fa_unic_code

                                            from django.core.serializers.json import DjangoJSONEncoder
                                            old_state = model_to_dict(existing_item)
                                            old_state_json = json.loads(json.dumps(old_state, cls=DjangoJSONEncoder))

                                            Item.objects.filter(id=existing_item.id).update(**{**defaults, 'updated_at': timezone.now()})
                                            history_records.append(ImportHistory(item=existing_item, action='update', previous_state=old_state_json))

                                            if is_pre_counted:
                                                counted_qty = defaults.get('inventory', getattr(existing_item, 'inventory', None))
                                                _mark_item_as_counted(existing_item, counted_qty, f"به‌روزرسانی مستقیم اقلام شمرده‌شده از اکسل ({original_file_name})")

                                            if is_doc_pre_approved:
                                                _mark_item_docs_as_approved(existing_item, defaults, f"به‌روزرسانی مستقیم اسناد تاییدشده از اکسل ({original_file_name})")

                                            updated += 1
                                            append_colored_row(row, 'updated', 'بروزرسانی کامل رکورد')
                                            q.put(json.dumps({"type": "updated", "msg": f"[ردیف {row_idx}] بروزرسانی رکورد: {item_display_code}"}) + "\n")
                                    else:
                                        if not target_warehouse_id:
                                            failed += 1
                                            err_msg = "بدون انبار امکان ساخت رکورد جدید نیست"
                                            error_details.append({"row": row_idx, "code": item_display_code, "error": err_msg})
                                            append_colored_row(row, 'err', err_msg)
                                            q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا در {item_display_code}: {err_msg}"}) + "\n")
                                            continue

                                        if not fa_unic_code:
                                            failed += 1
                                            err_msg = "برای ساخت رکورد جدید، کد یکتا (fa_unic_code) الزامی است"
                                            error_details.append({"row": row_idx, "code": "N/A", "error": err_msg})
                                            append_colored_row(row, 'err', err_msg)
                                            q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا: {err_msg}"}) + "\n")
                                            continue

                                        if user_id: defaults['created_by_id'] = user_id
                                        if 'tag_status' not in defaults:
                                            defaults['tag_status'] = 'چاپ نشده'
                                        new_item = Item.objects.create(fa_unic_code=fa_unic_code, **defaults)
                                        history_records.append(ImportHistory(item=new_item, action='create'))

                                        if is_pre_counted:
                                            counted_qty = new_item.inventory
                                            _mark_item_as_counted(new_item, counted_qty, f"ثبت مستقیم اقلام شمرده‌شده از فایل اکسل: {original_file_name}")

                                        if is_doc_pre_approved:
                                            _mark_item_docs_as_approved(new_item, defaults, f"ثبت مستقیم اسناد تاییدشده از فایل اکسل: {original_file_name}")

                                        created += 1
                                        append_colored_row(row, 'created', 'ثبت رکورد جدید')
                                        q.put(json.dumps({"type": "created", "msg": f"[ردیف {row_idx}] ثبت رکورد جدید: {fa_unic_code}"}) + "\n")
                            except Exception as row_err:
                                if str(row_err) == "CANCELED_BY_USER":
                                    raise
                                failed += 1
                                from django.db import DatabaseError
                                if isinstance(row_err, DatabaseError):
                                    logger.warning(f"[Import Row {row_idx}] Database error for {item_display_code}: {row_err}")
                                    err_msg = "خطا در قالب داده یا محدودیت‌های پایگاه داده (طول فیلد یا نوع نامعتبر)"
                                else:
                                    err_msg = str(row_err)
                                error_details.append({"row": row_idx, "code": item_display_code, "error": err_msg})
                                append_colored_row(row, 'err', err_msg)
                                q.put(json.dumps({"type": "err", "msg": f"[ردیف {row_idx}] خطا در {item_display_code}: {err_msg}"}) + "\n")

                        log_record = ImportLog.objects.create(
                            import_id=import_id,
                            warehouse_id=target_warehouse_id if 'target_warehouse_id' in locals() and target_warehouse_id else None,
                            imported_by_id=user_id,
                            file_name=original_file_name,
                            records_created=created,
                            records_updated=updated,
                            records_skipped=skipped,
                            records_failed=failed,
                            conflict_strategy=conflict_strategy,
                            error_details=error_details
                        )
                        
                        for hr in history_records:
                            hr.import_log = log_record
                        ImportHistory.objects.bulk_create(history_records)

                    req_wh_id = int(warehouse_id) if (warehouse_id and str(warehouse_id).isdigit()) else None
                    effective_wh_id = req_wh_id or (target_warehouse_id if 'target_warehouse_id' in locals() and target_warehouse_id else None)

                    if is_pre_counted:
                        try:
                            if effective_wh_id:
                                broadcast_count_task_update(warehouse_id=effective_wh_id)
                            else:
                                broadcast_count_task_update()
                        except Exception as bc_err:
                            logger.warning(f"Failed to broadcast count task update after pre-counted import: {bc_err}")

                    if is_doc_pre_approved:
                        try:
                            if effective_wh_id:
                                broadcast_doc_task_update(warehouse_id=effective_wh_id)
                            else:
                                broadcast_doc_task_update()
                        except Exception as bc_err:
                            logger.warning(f"Failed to broadcast doc task update after pre-approved docs import: {bc_err}")

                    # ثبت لاگ کلان ممیزی برای بارگذاری اکسل (خارج از تراکنش جهت جلوگیری از رول‌بک ایمپورت در صورت خطای جانبی)
                    try:
                        from accounts.audit_utils import log_audit_event
                        from django.contrib.auth import get_user_model
                        actor = get_user_model().objects.filter(id=user_id).first() if user_id else None
                        target_wh = Warehouse.objects.filter(id=effective_wh_id).first() if effective_wh_id else None
                        log_audit_event(
                            user=actor,
                            warehouse=target_wh,
                            module='docs',
                            action='IMPORT',
                            severity='info',
                            target_model='Item',
                            target_repr=f"بارگذاری اکسل کالاها ({original_file_name})",
                            details={
                                'import_id': import_id,
                                'file_name': original_file_name,
                                'is_pre_counted': is_pre_counted,
                                'is_doc_pre_approved': is_doc_pre_approved,
                                'records_created': created,
                                'records_updated': updated,
                                'records_skipped': skipped,
                                'records_failed': failed,
                                'conflict_strategy': conflict_strategy
                            }
                        )
                    except Exception as log_err:
                        logger.warning(f"Failed to log audit event for excel import: {log_err}")

                    # Save the colored log workbook
                    if import_id:
                        try:
                            _cleanup_old_import_logs(max_age_hours=24)
                            out_file_path = os.path.join(tempfile.gettempdir(), f"import_log_{import_id}.xlsx")
                            out_wb.save(out_file_path)
                        except Exception as wb_err:
                            logger.warning(f"Failed to save import log workbook: {wb_err}")

                except Exception as ex:
                    if str(ex) == "CANCELED_BY_USER":
                        q.put(json.dumps({
                            "type": "summary",
                            "status": "cancelled",
                            "created": 0, "updated": 0, "skipped": 0, "failed": 0,
                            "found_fields": [], "missing_fields": [],
                            "error_details": []
                        }) + "\n")
                        return
                    else:
                        traceback.print_exc()
                        q.put(json.dumps({"type": "err", "msg": f">> خطای بحرانی در فرآیند: {str(ex)}"}) + "\n")
                        q.put(json.dumps({
                            "type": "summary",
                            "status": "failed",
                            "created": 0,
                            "updated": 0,
                            "skipped": skipped if 'skipped' in locals() else 0,
                            "failed": (failed if 'failed' in locals() else 0) + 1,
                            "found_fields": list(set(found_fields)) if 'found_fields' in locals() else [],
                            "missing_fields": list(set(missing_fields)) if 'missing_fields' in locals() else [],
                            "error_details": (error_details if 'error_details' in locals() else []) + [{"row": 0, "code": "CRASH", "error": f"خطای سیستمی: {str(ex)}"}]
                        }) + "\n")
                        return

                q.put(json.dumps({
                    "type": "summary",
                    "status": "success",
                    "created": created,
                    "updated": updated,
                    "skipped": skipped,
                    "failed": failed,
                    "found_fields": list(set(found_fields)),
                    "missing_fields": list(set(missing_fields)),
                    "error_details": error_details
                }) + "\n")
                
            except Exception as e:
                traceback.print_exc()
                q.put(json.dumps({"type": "err", "msg": f">> خطای سیستمی: {str(e)}"}) + "\n")
                q.put(json.dumps({
                    "type": "summary",
                    "status": "failed",
                    "created": 0,
                    "updated": 0,
                    "skipped": 0,
                    "failed": 1,
                    "found_fields": [],
                    "missing_fields": [],
                    "error_details": [{"row": 0, "code": "FATAL", "error": str(e)}]
                }) + "\n")
            finally:
                q.put(None)
                if import_id:
                    cache.delete(f"cancel_import_{import_id}")
                if os.path.exists(temp_path):
                    try:
                        os.remove(temp_path)
                    except:
                        pass
                try:
                    from django.db import connection
                    connection.close()
                except Exception:
                    pass

        async def stream_logs_async():
            q = queue.Queue()
            thread = threading.Thread(target=worker, args=(q,))
            thread.start()
            
            while True:
                chunk = await asyncio.to_thread(q.get)
                if chunk is None:
                    break
                yield chunk

        response = StreamingHttpResponse(
            stream_logs_async(), 
            content_type='application/x-ndjson'
        )
        response['X-Accel-Buffering'] = 'no'
        response['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response['X-Content-Type-Options'] = 'nosniff'
        return response

    @action(detail=False, methods=['post'])
    def revert_import(self, request):
        import_id = request.data.get('import_id')
        if not import_id:
            return Response({'error': 'شناسه فرآیند الزامی است.'}, status=400)
            
        try:
            import_log = ImportLog.objects.get(import_id=import_id)
        except ImportLog.DoesNotExist:
            return Response({'error': 'فرآیندی با این شناسه یافت نشد.'}, status=404)

        if not self._can_manage_import_log(import_log):
            return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)

        if import_log.is_reverted:
            return Response({'error': 'این فرآیند قبلاً بازگردانی شده است.'}, status=400)
            
        time_elapsed = (timezone.now() - import_log.imported_at).total_seconds()
        if time_elapsed > 900: # 15 minutes
            return Response({'error': 'مهلت ۱۵ دقیقه‌ای برای بازگردانی این فرآیند به پایان رسیده است.'}, status=400)
            
        try:
            with transaction.atomic():
                histories = import_log.histories.all()
                items_to_create = []
                for history in histories:
                    if history.action == 'create':
                        if history.item:
                            # بازگردانی ایجاد = حذف (نرم) رکورد ساخته‌شده
                            # توجه: تسک اسناد وابسته حذف سخت نمی‌شود تا تاریخچه اسناد و تقارن با تسک شمارش حفظ شود
                            _soft_delete_items_cascade(Item.objects.filter(id=history.item_id))
                    elif history.action == 'update' or history.action == 'delete':
                        if history.previous_state:
                            state = history.previous_state.copy()

                            # Handle foreign keys correctly for both update and create
                            fk_fields = ['warehouse', 'created_by', 'modified_by']
                            for fk in fk_fields:
                                if fk in state and isinstance(state[fk], int):
                                    state[f'{fk}_id'] = state.pop(fk)

                            if history.action == 'update' and history.item:
                                # بازگردانی مشخصات کالا
                                Item.all_objects.filter(id=history.item.id).update(**{**state, 'updated_at': timezone.now()})

                                # بازگردانی وضعیت تسک شمارش بر اساس تاریخچه قبل از زمان ایمپورت
                                prev_field_status = state.get('field_status')
                                c_task = CountTask.all_objects.filter(item=history.item).first()
                                if c_task:
                                    if prev_field_status and prev_field_status != 'done':
                                        hist_qs = CountTaskHistory.all_objects.filter(task=c_task)
                                        if import_log.import_id:
                                            hist_qs = hist_qs.exclude(note__startswith=f"[IMPORT:{import_log.import_id}]")
                                        last_hist = hist_qs.order_by('-created_at').first()
                                        if last_hist:
                                            c_task.status = last_hist.action_type
                                            c_task.counted_balance = last_hist.counted_balance
                                            c_task.save(update_fields=['status', 'counted_balance', 'updated_at'])
                                        else:
                                            c_task.status = 'PENDING_COUNT'
                                            c_task.counted_balance = None
                                            c_task.save(update_fields=['status', 'counted_balance', 'updated_at'])
                                    elif prev_field_status == 'done':
                                        prev_inv = state.get('inventory')
                                        if prev_inv is not None:
                                            c_task.counted_balance = prev_inv
                                            c_task.save(update_fields=['counted_balance', 'updated_at'])

                                # بازگردانی وضعیت و مقادیر مالی تسک اسناد بر اساس تاریخچه قبل از ایمپورت
                                prev_doc_status = state.get('doc_status')
                                d_task = DocTask.objects.filter(item=history.item).first()
                                if d_task:
                                    if prev_doc_status and prev_doc_status != 'done':
                                        d_hist_qs = DocTaskHistory.objects.filter(task=d_task)
                                        if import_log.import_id:
                                            d_hist_qs = d_hist_qs.exclude(note__startswith=f"[IMPORT:{import_log.import_id}]")
                                        last_d_hist = d_hist_qs.order_by('-created_at').first()
                                        if last_d_hist:
                                            valid_statuses = [choice[0] for choice in DocTask.STATUS_CHOICES]
                                            target_status = None
                                            if last_d_hist.data_snapshot and isinstance(last_d_hist.data_snapshot, dict):
                                                snap_status = last_d_hist.data_snapshot.get('status')
                                                if snap_status in valid_statuses:
                                                    target_status = snap_status

                                            if not target_status:
                                                if last_d_hist.action_type in valid_statuses:
                                                    target_status = last_d_hist.action_type
                                                else:
                                                    target_status = 'PENDING_DOC'

                                            d_task.status = target_status
                                            if last_d_hist.data_snapshot and isinstance(last_d_hist.data_snapshot, dict):
                                                snap = last_d_hist.data_snapshot
                                                financial_keys = [
                                                    'added_rti_no', 'inv_rti_number', 'invoice_type', 'invoice_date',
                                                    'invoice_page', 'page_row', 'doc_supplier', 'total_value',
                                                    'price_amount', 'similar_unit_price', 'currency', 'folder_address',
                                                    'stamp', 'signature'
                                                ]
                                                for f_key in financial_keys:
                                                    if f_key in snap:
                                                        val = snap[f_key]
                                                        if f_key == 'invoice_date' and val:
                                                            val = _parse_date_flexible(val)
                                                        elif f_key in ['total_value', 'price_amount', 'similar_unit_price'] and val is not None:
                                                            val = _parse_decimal(val)
                                                        elif f_key in ['added_rti_no', 'inv_rti_number'] and val:
                                                            val = str(val)[:100]
                                                        elif f_key == 'invoice_type' and val:
                                                            val = _clean_invoice_type(val)
                                                        elif f_key == 'currency' and val:
                                                            val = _clean_currency(val)
                                                        elif f_key == 'doc_supplier' and val:
                                                            val = str(val)[:255]
                                                        elif f_key == 'folder_address' and val:
                                                            val = str(val)[:500]
                                                        elif f_key in ['invoice_page', 'page_row'] and val is not None:
                                                            val = _parse_positive_int(val)
                                                        setattr(d_task, f_key, val)
                                                d_task.save()
                                            else:
                                                d_task.save(update_fields=['status', 'updated_at'])
                                        else:
                                            d_task.status = 'PENDING_DOC'
                                            d_task.save(update_fields=['status', 'updated_at'])
                                    elif prev_doc_status == 'done':
                                        d_task.status = 'DOC_FINAL_APPROVED'
                                        if 'added_rti_no' in state:
                                            d_task.added_rti_no = str(state['added_rti_no'])[:100] if state['added_rti_no'] else None
                                        if 'inv_rti_number' in state:
                                            d_task.inv_rti_number = str(state['inv_rti_number'])[:100] if state['inv_rti_number'] else None
                                        if 'invoice_type' in state:
                                            d_task.invoice_type = _clean_invoice_type(state['invoice_type'])
                                        if 'currency' in state:
                                            d_task.currency = _clean_currency(state['currency'])
                                        if 'total_value' in state and state['total_value'] is not None:
                                            d_task.total_value = _parse_decimal(state['total_value'])
                                        if 'price_amount' in state and state['price_amount'] is not None:
                                            d_task.price_amount = _parse_decimal(state['price_amount'])
                                        if 'similar_unit_price' in state and state['similar_unit_price'] is not None:
                                            d_task.similar_unit_price = _parse_decimal(state['similar_unit_price'])
                                        if 'doc_supplier' in state:
                                            d_task.doc_supplier = str(state['doc_supplier'])[:255] if state['doc_supplier'] else None
                                        if 'folder_address' in state:
                                            d_task.folder_address = str(state['folder_address'])[:500] if state['folder_address'] else None
                                        if 'invoice_date' in state and state['invoice_date']:
                                            d_task.invoice_date = _parse_date_flexible(state['invoice_date'])
                                        if 'invoice_page' in state and state['invoice_page'] is not None:
                                            d_task.invoice_page = _parse_positive_int(state['invoice_page'])
                                        if 'page_row' in state and state['page_row'] is not None:
                                            d_task.page_row = _parse_positive_int(state['page_row'])
                                        d_task.save()
                            elif history.action == 'delete':
                                # حذف حالا نرم است؛ ردیف tombstone هنوز با همه داده‌ها موجود است
                                # → فقط احیا می‌شود. اگر (به هر دلیل) واقعاً حذف شده بود، بازسازی.
                                tombstone_ids = list(
                                    Item.all_objects.filter(
                                        warehouse_id=state.get('warehouse_id'),
                                        fa_unic_code=state.get('fa_unic_code'),
                                        is_deleted=True,
                                    ).values_list('id', flat=True)
                                )
                                if tombstone_ids:
                                    Item.all_objects.filter(id__in=tombstone_ids).update(
                                        is_deleted=False, updated_at=timezone.now()
                                    )
                                    # عکس‌ها هم با حذف کالا tombstone شده بودند
                                    _restore_item_photos_cascade(tombstone_ids)
                                else:
                                    items_to_create.append(Item(**state))
                
                if items_to_create:
                    Item.objects.bulk_create(items_to_create, ignore_conflicts=True)
                            
                import_log.is_reverted = True
                import_log.save()

                try:
                    wh_id = import_log.warehouse_id
                    if wh_id:
                        broadcast_count_task_update(warehouse_id=wh_id)
                        broadcast_doc_task_update(warehouse_id=wh_id)
                    else:
                        broadcast_count_task_update()
                        broadcast_doc_task_update()
                except Exception as bc_err:
                    logger.warning(f"Failed to broadcast task updates after revert_import: {bc_err}")

                from accounts.audit_utils import log_audit_event
                log_audit_event(
                    user=request.user,
                    warehouse=import_log.warehouse,
                    module='docs',
                    action='ROLLBACK',
                    severity='warning',
                    target_model='Item',
                    target_repr=f"بازگردانی بارگذاری اکسل ({import_log.file_name})",
                    details={
                        'import_id': import_id,
                        'file_name': import_log.file_name,
                        'affected_records': len(histories)
                    },
                    ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
                )
                
            return Response({'status': 'success', 'msg': 'فرآیند با موفقیت بازگردانی شد.', 'affected_records': len(histories)})
        except Exception as e:
            logger.exception(f"Error reverting import {import_id}")
            return Response({'error': 'خطا در بازگردانی فرآیند در دیتابیس.'}, status=500)

    @action(detail=False, methods=['post'])
    def clear_warehouse_data(self, request):
        warehouse_id = request.data.get('warehouse_id')
        if not warehouse_id:
            return Response({'error': 'شناسه انبار الزامی است.'}, status=400)

        try:
            warehouse_id_val = int(warehouse_id)
        except (ValueError, TypeError):
            return Response({'error': 'شناسه انبار نامعتبر است.'}, status=400)

        if not self._can_write_warehouse(warehouse_id_val):
            return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)
        warehouse_id = warehouse_id_val

        try:
            items_deleted = 0
            user_id = request.user.id if request.user.is_authenticated else None
            import_id = f"clear_{datetime.now().strftime('%Y%m%d%H%M%S')}_{warehouse_id}"
            
            with transaction.atomic():
                import_log = ImportLog.objects.create(
                    import_id=import_id,
                    warehouse_id=warehouse_id,
                    imported_by_id=user_id,
                    file_name="حذف تمامی داده‌های انبار"
                )
                
                items_qs = Item.objects.filter(warehouse_id=warehouse_id)
                
                # We need to bulk read the items to save them
                from django.core.serializers.json import DjangoJSONEncoder
                histories = []
                for item in items_qs.iterator(chunk_size=1000):
                    state = model_to_dict(item)
                    state = json.loads(json.dumps(state, cls=DjangoJSONEncoder))
                    
                    histories.append(ImportHistory(
                        import_log=import_log,
                        item=None, 
                        action='delete',
                        previous_state=state
                    ))
                    if len(histories) >= 500:
                        ImportHistory.objects.bulk_create(histories, batch_size=500)
                        histories.clear()
                
                if histories:
                    ImportHistory.objects.bulk_create(histories, batch_size=500)
                    histories.clear()
                    
                deleted = _soft_delete_items_cascade(items_qs)
                items_deleted = deleted
                
                import_log.records_created = items_deleted # Store count here
                import_log.save()

                from accounts.audit_utils import log_audit_event
                log_audit_event(
                    user=request.user,
                    warehouse_id=warehouse_id,
                    module='docs',
                    action='DELETE',
                    severity='critical',
                    target_model='Warehouse',
                    target_object_id=warehouse_id,
                    target_repr=f"پاکسازی کامل داده‌های انبار #{warehouse_id}",
                    details={
                        'operation': 'CLEAR_WAREHOUSE_DATA',
                        'items_deleted': items_deleted,
                        'import_id': import_id
                    },
                    ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
                )
                
            return Response({'status': 'success', 'msg': f'{items_deleted} رکورد با موفقیت حذف شدند.', 'import_id': import_id})
        except Exception as e:
            return Response({'error': f'خطا در حذف داده‌ها: {str(e)}'}, status=500)

    @action(detail=False, methods=['get'])
    def latest_import(self, request):
        warehouse_id = request.query_params.get('warehouse_id')
        if not warehouse_id:
            return Response({'error': 'شناسه انبار الزامی است.'}, status=400)

        try:
            warehouse_id_val = int(warehouse_id)
        except (ValueError, TypeError):
            return Response({'error': 'شناسه انبار نامعتبر است.'}, status=400)

        if not self._can_write_warehouse(warehouse_id_val):
            return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)
        warehouse_id = warehouse_id_val

        try:
            logs = ImportLog.objects.filter(warehouse_id=warehouse_id, is_reverted=False).order_by('-imported_at')[:10]
            
            recent_imports = []
            for log in logs:
                time_elapsed = (timezone.now() - log.imported_at).total_seconds()
                if time_elapsed <= 900:
                    recent_imports.append({
                        'import_id': log.import_id,
                        'file_name': log.file_name,
                        'imported_at': log.imported_at.isoformat(),
                        'time_remaining_seconds': 900 - int(time_elapsed),
                        'records_affected': log.records_created + log.records_updated
                    })
                
            return Response({
                'latest_import': recent_imports[0] if recent_imports else None,
                'recent_imports': recent_imports
            })
        except Exception as e:
            return Response({'error': str(e)}, status=500)

    @action(detail=False, methods=['post'], parser_classes=[MultiPartParser, FormParser])
    def delete_from_excel(self, request):
        warehouse_id = request.data.get('warehouse_id')
        file_obj = request.FILES.get('file')

        if not warehouse_id or not file_obj:
            return Response({'error': 'پارامترهای warehouse_id و file الزامی است.'}, status=400)

        try:
            warehouse_id_val = int(warehouse_id)
        except (ValueError, TypeError):
            return Response({'error': 'شناسه انبار نامعتبر است.'}, status=400)

        if not self._can_write_warehouse(warehouse_id_val):
            return Response({'error': 'شما به این انبار دسترسی ندارید.'}, status=403)
        warehouse_id = warehouse_id_val

        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            sheet = wb.active
            headers = [str(cell.value).strip() if cell.value else '' for cell in sheet[1]]
            
            id_idx = headers.index('id') if 'id' in headers else None
            fa_unic_idx = headers.index('fa_unic_code') if 'fa_unic_code' in headers else None
            
            if id_idx is None and fa_unic_idx is None:
                return Response({'error': 'ستون fa_unic_code یا id در فایل اکسل یافت نشد.'}, status=400)
                
            ids_to_delete = []
            fa_unics_to_delete = []
            
            for row in sheet.iter_rows(min_row=2, values_only=True):
                id_val = row[id_idx] if id_idx is not None else None
                fa_val = row[fa_unic_idx] if fa_unic_idx is not None else None
                
                if id_val is not None and str(id_val).strip():
                    raw_str = str(id_val).strip()
                    try:
                        clean_id = int(float(raw_str))
                        ids_to_delete.append(clean_id)
                    except (ValueError, TypeError):
                        ids_to_delete.append(raw_str)
                elif fa_val is not None and str(fa_val).strip():
                    fa_unics_to_delete.append(str(fa_val).strip())
                    
            if not ids_to_delete and not fa_unics_to_delete:
                return Response({'error': 'هیچ شناسه یا کدی در فایل یافت نشد.'}, status=400)
                
            items_deleted = 0
            user_id = request.user.id if request.user.is_authenticated else None
            import_id = f"del_{datetime.now().strftime('%Y%m%d%H%M%S')}_{warehouse_id}"
            
            with transaction.atomic():
                import_log = ImportLog.objects.create(
                    import_id=import_id,
                    warehouse_id=warehouse_id,
                    imported_by_id=user_id,
                    file_name=f"حذف گروهی (اکسل): {file_obj.name}"
                )
                
                query = Q()
                if ids_to_delete:
                    query |= Q(id__in=ids_to_delete)
                if fa_unics_to_delete:
                    query |= Q(fa_unic_code__in=fa_unics_to_delete)
                    
                items_qs = Item.objects.filter(Q(warehouse_id=warehouse_id) & query)
                
                # We need to bulk read the items to save them
                from django.core.serializers.json import DjangoJSONEncoder
                histories = []
                for item in items_qs.iterator(chunk_size=1000):
                    # For delete action, item foreign key might be nullified when item is deleted. 
                    # We store it anyway.
                    state = model_to_dict(item)
                    state = json.loads(json.dumps(state, cls=DjangoJSONEncoder))
                    
                    histories.append(ImportHistory(
                        import_log=import_log,
                        item=None, # item will be deleted, don't set foreign key
                        action='delete',
                        previous_state=state
                    ))
                    if len(histories) >= 500:
                        ImportHistory.objects.bulk_create(histories, batch_size=500)
                        histories.clear()
                
                if histories:
                    ImportHistory.objects.bulk_create(histories, batch_size=500)
                    histories.clear()
                    
                deleted = _soft_delete_items_cascade(items_qs)
                items_deleted = deleted
                
                import_log.records_created = items_deleted # Store count in this field for history
                import_log.save()

                from accounts.audit_utils import log_audit_event
                log_audit_event(
                    user=request.user,
                    warehouse_id=warehouse_id,
                    module='docs',
                    action='DELETE',
                    severity='critical',
                    target_model='Item',
                    target_repr=f"حذف گروهی {items_deleted} کالا از فایل اکسل ({file_obj.name})",
                    details={
                        'file_name': file_obj.name,
                        'items_deleted': items_deleted,
                        'import_id': import_id
                    },
                    ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
                )
                
            return Response({'status': 'success', 'msg': f'{items_deleted} رکورد با موفقیت از انبار حذف شدند.', 'import_id': import_id})
        except Exception as e:
            return Response({'error': f'خطا در پردازش فایل: {str(e)}'}, status=500)

    @action(detail=False, methods=['get'])
    def dashboard_stats(self, request):
        import jdatetime
        from datetime import timedelta
        from django.utils import timezone
        from django.db.models import Q
        
        company_id = request.headers.get('X-Company-ID') or request.query_params.get('company_id')
        project_id = request.query_params.get('project_id')
        
        items = Item.objects.all()
        
        # ۱. انزوای چندشرکتی انبار: اعمال فیلتر بر مبنای شرکت فعال
        if company_id:
            try:
                cid = int(company_id)
                items = items.filter(warehouse__company_id=cid)
            except (ValueError, TypeError):
                pass
        elif not request.user.is_superuser:
            try:
                from companies.access import get_user_allowed_companies
                allowed_comp_ids = list(get_user_allowed_companies(request.user, required_level='workspace_full').values_list('id', flat=True))
                items = items.filter(warehouse__company_id__in=allowed_comp_ids)
            except Exception:
                pass

        # ۲. اعمال اسکوپ کاربر به انبارهای مجاز
        items = scope_queryset(items, request.user)

        if project_id and project_id != 'ALL':
            items = items.filter(warehouse_id=project_id)
            
        now = timezone.localtime(timezone.now()) if timezone.is_aware(timezone.now()) else timezone.now()
        today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        
        # Overall
        total_quantity = items.count()
        total_counted = items.exclude(field_status__in=['waiting', 'counting', 'recount', 'در انتظار شمارش', 'بازشماری']).count()
        printed_tags = items.filter(tag_status__in=['printed', 'reprint', 'چاپ شده', 'چاپ مجدد']).count()
        docs_approved = items.filter(doc_status__in=['done', 'approved']).count()
        conflicts = items.filter(Q(has_conflict=True) | Q(field_status='recount')).count()
        done = items.filter(field_status='done', doc_status__in=['done', 'approved']).count()
        
        # Days stats (last 7 days)
        weekly_data = []
        
        for i in range(6, -1, -1):
            d_start = today_start - timedelta(days=i)
            d_end = d_start + timedelta(days=1)
            
            day_items = items.filter(updated_at__gte=d_start, updated_at__lt=d_end)
            
            c_count = day_items.exclude(field_status__in=['waiting', 'counting', 'recount', 'در انتظار شمارش', 'بازشماری']).count()
            c_docs = day_items.filter(doc_status__in=['done', 'approved']).count()
            c_feed = day_items.filter(field_status='done', doc_status__in=['done', 'approved']).count()
            
            jdt = jdatetime.datetime.fromgregorian(datetime=d_start)
            day_label = 'امروز' if i == 0 else ('دیروز' if i == 1 else jdt.strftime('%A'))
            
            weekly_data.append({
                'day': day_label,
                'count': c_count,
                'docs': c_docs,
                'feed': c_feed
            })
            
        today_stats = weekly_data[-1]
        yesterday_stats = weekly_data[-2]
        
        last_7_days_count = sum(d['count'] for d in weekly_data)
        last_7_days_docs = sum(d['docs'] for d in weekly_data)
        last_7_days_feed = sum(d['feed'] for d in weekly_data)
        
        max_val = max([max(d['count'], d['docs'], d['feed']) for d in weekly_data] + [10])
        
        return Response({
            'overall': {
                'total': total_quantity,
                'counted': total_counted,
                'printed': printed_tags,
                'docs_approved': docs_approved,
                'conflicts': conflicts,
                'done': done,
                'ready_to_feed': max(0, total_counted - done)
            },
            'today': today_stats,
            'yesterday': yesterday_stats,
            'last_week_totals': {
                'count': last_7_days_count,
                'docs': last_7_days_docs,
                'feed': last_7_days_feed
            },
            'weekly_data': weekly_data,
            'overallMax': max_val
        })

__all__ = [
    'ItemViewSet',
]
