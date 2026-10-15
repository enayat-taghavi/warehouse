# -*- coding: utf-8 -*-
"""
ViewSet for inventory counting tasks (انبارگردانی).
"""
from decimal import Decimal
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from common.sync_models import soft_delete_queryset
from common.warehouse_scope import scope_queryset, can_access_warehouse
from ..models import CountTask, CountTaskHistory, Item
from ..serializers import CountTaskSerializer, photo_prefetch
from ..signals import broadcast_count_task_update
from .helpers import _parse_decimal

class CountTaskViewSet(viewsets.ModelViewSet):
    serializer_class = CountTaskSerializer

    def get_serializer_context(self):
        context = super().get_serializer_context()
        as_role = self.request.query_params.get('as_role')
        user = self.request.user
        
        # اگر نقش یکی از نقش‌های مدیریتی، نظارتی یا پیگیری باشد، یا اکشن خروجی اکسل باشد
        # شمارش نباید کور باشد و اطلاعات موجودی کامل باید ارسال گردد.
        if as_role in ['manager', 'supervisor', 'count_tracking', 'doc_supervisor', 'doc_worker']:
            context['is_blind'] = False
        elif self.action in ['export_excel', 'get_export_columns', 'bulk_manager_approve', 'bulk_manager_reject', 'bulk_approve', 'bulk_reject']:
            context['is_blind'] = False
        elif as_role == 'counter':
            context['is_blind'] = True
        elif user.is_authenticated and (user.is_superuser or user.has_perm('accounts.view_sys_manager_review') or user.has_perm('accounts.view_sys_supervisor') or user.has_perm('accounts.view_wh_stocktaking')):
            if not user.has_perm('accounts.can_act_as_counter') and not user.has_perm('accounts.view_sys_counter'):
                context['is_blind'] = False
        return context

    def get_permissions(self):
        from accounts.permissions import HasMenuAccess
        from rest_framework.permissions import IsAuthenticated
        
        if self.action in ['list', 'retrieve', 'pool_tasks', 'claim_tasks', 'get_export_columns', 'export_excel']:
            permission_classes = [HasMenuAccess('view_sys_counter') | HasMenuAccess('view_sys_supervisor') | HasMenuAccess('view_sys_manager_review') | HasMenuAccess('view_sys_recounts') | HasMenuAccess('view_wh_stocktaking')]
        elif self.action == 'bulk_submit':
            permission_classes = [HasMenuAccess('view_sys_counter') | HasMenuAccess('can_act_as_counter')]
        elif self.action == 'bulk_approve':
            permission_classes = [HasMenuAccess('view_sys_supervisor') | HasMenuAccess('can_act_as_supervisor') | HasMenuAccess('perm_feed_approve_action')]
        elif self.action in ['reject', 'bulk_reject']:
            permission_classes = [HasMenuAccess('perm_rec_recount') | HasMenuAccess('view_sys_supervisor') | HasMenuAccess('view_sys_manager_review') | HasMenuAccess('can_act_as_supervisor')]
        elif self.action in ['manager_reject', 'bulk_manager_reject']:
            permission_classes = [HasMenuAccess('perm_rec_recount') | HasMenuAccess('view_sys_manager_review') | HasMenuAccess('can_act_as_manager') | HasMenuAccess('perm_inventory_finalize')]
        elif self.action in ['bulk_manager_approve', 'bulk_cancel']:
            permission_classes = [HasMenuAccess('view_sys_manager_review') | HasMenuAccess('can_act_as_manager') | HasMenuAccess('perm_inventory_finalize')]
        else: # create, update, etc
            permission_classes = [IsAuthenticated()]
            
        return permission_classes

    def perform_destroy(self, instance):
        # حذف نرم + tombstone تاریخچه، تا کلاینت‌های آفلاین باخبر شوند
        with transaction.atomic():
            soft_delete_queryset(CountTaskHistory.objects.filter(task_id=instance.id))
            instance.soft_delete()

    def get_object(self):
        """
        زیرساخت سینک: اگر pk عددی نبود، به‌عنوان sync_id تفسیر می‌شود
        (کلاینت آفلاین ممکن است فقط sync_id پایدار را داشته باشد).
        """
        pk = self.kwargs.get(self.lookup_url_kwarg or self.lookup_field)
        if pk is not None and not str(pk).isdigit():
            from django.shortcuts import get_object_or_404
            queryset = self.filter_queryset(self.get_queryset())
            obj = get_object_or_404(queryset, sync_id=pk)
            self.check_object_permissions(self.request, obj)
            return obj
        return super().get_object()

    def update(self, request, *args, **kwargs):
        """
        تشخیص تداخل خوش‌بینانه: کلاینت می‌تواند base_updated_at (نسخه‌ای که
        رویش تغییر داده) بفرستد؛ اگر سرور جدیدتر باشد → 409 + رکورد سروری
        برای reconciliation سمت کلاینت. بدون این پارامتر، رفتار قبلی حفظ است.
        """
        base_raw = request.data.get('base_updated_at')
        if base_raw:
            from django.utils.dateparse import parse_datetime
            base = parse_datetime(str(base_raw))
            if base is None:
                return Response({'detail': 'base_updated_at نامعتبر است.'}, status=400)
            if timezone.is_naive(base):
                base = timezone.make_aware(base)
            instance = self.get_object()
            # تلورانس ۱ms برای جلوگیری از تداخل کاذب ناشی از گرد شدن میکروثانیه
            if (instance.updated_at - base).total_seconds() > 0.001:
                return Response({
                    'detail': 'conflict',
                    'server_record': self.get_serializer(instance).data,
                }, status=409)
        return super().update(request, *args, **kwargs)

    def perform_update(self, serializer):
        from rest_framework.exceptions import ValidationError
        instance = serializer.instance
        old_status = instance.status
        old_counted = str(instance.counted_balance)
        new_status = serializer.validated_data.get('status', old_status)

        # جلوگیری از جهش غیرمجاز وضعیت با متد مستقیم PATCH
        if new_status != old_status:
            allowed_patch_transitions = {
                'PENDING_COUNT': ['INITIAL_COUNT', 'PENDING_COUNT'],
                'INITIAL_COUNT': ['INITIAL_COUNT', 'PENDING_COUNT'],
                'SUPERVISOR_REJECTED': ['INITIAL_COUNT', 'SUPERVISOR_REJECTED'],
                'MANAGER_REJECTED': ['INITIAL_COUNT', 'MANAGER_REJECTED'],
            }
            allowed = allowed_patch_transitions.get(old_status, [])
            if new_status not in allowed:
                raise ValidationError({
                    'status': f'تغییر مستقیم وضعیت از {old_status} به {new_status} با متد PATCH مجاز نیست. لطفاً از اکشن‌های گردش‌کار استفاده کنید.'
                })

        updated_instance = serializer.save(modified_by=self.request.user)

        from accounts.audit_utils import log_audit_event
        from ..models import CountTaskHistory

        note = ''
        if new_status in ['MANAGER_REJECTED', 'FINAL_APPROVED']:
            note = updated_instance.manager_note or ''
        elif new_status == 'SUPERVISOR_REJECTED':
            note = updated_instance.supervisor_note or ''
        elif new_status in ['COUNTED', 'INITIAL_COUNT']:
            note = updated_instance.counter_note or ''

        if old_status != new_status or old_counted != str(updated_instance.counted_balance):
            CountTaskHistory.objects.create(
                task=updated_instance,
                action_by=self.request.user if self.request.user.is_authenticated else None,
                action_type=new_status,
                counted_balance=updated_instance.counted_balance,
                note=note
            )

            try:
                item_label = updated_instance.item.fa_unic_code if updated_instance.item else f"#{updated_instance.id}"
                wh = updated_instance.item.warehouse if updated_instance.item else None
                log_audit_event(
                    module='counter' if new_status in ['INITIAL_COUNT', 'COUNTED', 'PENDING_COUNT'] else 'supervisor',
                    action='UPDATE' if old_status == new_status else 'STATUS_CHANGE',
                    target_model='CountTask',
                    target_object_id=updated_instance.id,
                    target_repr=f"شمارش کالا {item_label} ({new_status})",
                    warehouse=wh,
                    user=self.request.user if self.request.user.is_authenticated else None,
                    before_state={'status': old_status, 'counted_balance': old_counted},
                    after_state={'status': new_status, 'counted_balance': str(updated_instance.counted_balance), 'note': note}
                )
            except Exception:
                pass

    def get_queryset(self):
        user = self.request.user
        from django.db.models import Prefetch
        from ..models import CountTaskHistory
        queryset = CountTask.objects.all().select_related(
            'item', 'counter', 'supervisor', 'assigned_manager', 'created_by', 'modified_by'
        ).prefetch_related(
            Prefetch('history', queryset=CountTaskHistory.objects.order_by('created_at').select_related('action_by')),
            # item_details کل ItemSerializer را صدا می‌زند؛ بدون این، هر تسک دو
            # کوئری عکس اضافه می‌زد.
            photo_prefetch('item__photos'),
        )
        
        as_role = self.request.query_params.get('as_role')
        warehouse_id = self.request.query_params.get('warehouse_id')
        status_filter = self.request.query_params.get('status')
        date_filter = self.request.query_params.get('date')
        q_filter = self.request.query_params.get('q')
        
        from common.tenant_scope import scope_tenant_queryset
        queryset = scope_tenant_queryset(queryset, user, company_field='item__warehouse__company_id', request=self.request)

        if warehouse_id and str(warehouse_id) not in ['ALL', '-1']:
            try:
                requested_wh = int(warehouse_id)
            except (TypeError, ValueError):
                return CountTask.objects.none()
            if not can_access_warehouse(user, requested_wh):
                return CountTask.objects.none()
            queryset = queryset.filter(item__warehouse_id=requested_wh)
        else:
            queryset = scope_queryset(queryset, user, field='item__warehouse_id')
        
        # 1. فیلتر بر اساس نقش (Role Filter)
        if as_role == 'counter':
            queryset = queryset.filter(counter=user)
        elif as_role == 'supervisor':
            queryset = queryset.filter(supervisor=user)
        elif as_role == 'manager':
            from django.db.models import Q
            queryset = queryset.filter(Q(assigned_manager=user) | Q(assigned_manager__isnull=True))
        elif as_role == 'tracking':
            # بازگرداندن همه موارد در حال شمارش برای پیگیری
            show_completed = self.request.query_params.get('show_completed', 'false').lower() == 'true'
            if not show_completed:
                queryset = queryset.exclude(status='FINAL_APPROVED')
        else:
            # Fallback to permission checking if as_role is not provided
            if user.is_superuser or user.has_perm('accounts.view_sys_manager_review') or user.has_perm('accounts.can_act_as_manager') or user.has_perm('inventory.can_act_as_manager'):
                from django.db.models import Q
                queryset = queryset.filter(Q(assigned_manager=user) | Q(assigned_manager__isnull=True))
            elif user.has_perm('accounts.view_sys_supervisor') or user.has_perm('accounts.can_act_as_supervisor') or user.has_perm('inventory.can_act_as_supervisor'):
                queryset = queryset.filter(supervisor=user)
            elif user.has_perm('accounts.view_sys_counter') or user.has_perm('accounts.can_act_as_counter') or user.has_perm('inventory.can_act_as_counter'):
                queryset = queryset.filter(counter=user)
            else:
                return CountTask.objects.none()

        # 2. فیلتر جستجوی متنی (Text Search Q)
        if q_filter:
            from django.db.models import Q
            q_clean = q_filter.strip()
            queryset = queryset.filter(
                Q(item__fa_unic_code__icontains=q_clean) |
                Q(item__description__icontains=q_clean) |
                Q(item__tag__icontains=q_clean) |
                Q(item__po__icontains=q_clean) |
                Q(item__pk_number__icontains=q_clean) |
                Q(item__new_location__icontains=q_clean)
            )

        # 3. فیلتر تاریخ (Date Filter)
        if date_filter and date_filter != 'all':
            from django.utils import timezone
            from datetime import timedelta
            now = timezone.localtime(timezone.now())
            today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
            if date_filter == 'today':
                queryset = queryset.filter(updated_at__gte=today_start)
            elif date_filter == 'yesterday':
                yesterday_start = today_start - timedelta(days=1)
                queryset = queryset.filter(updated_at__gte=yesterday_start, updated_at__lt=today_start)
            elif date_filter == 'week':
                week_start = today_start - timedelta(days=7)
                queryset = queryset.filter(updated_at__gte=week_start)

        # 4. فیلتر وضعیت (Status Filter)
        if status_filter and status_filter != 'all':
            from django.db.models import Q
            if as_role == 'counter':
                if status_filter == 'pending':
                    queryset = queryset.filter(status='PENDING_COUNT', counted_balance__isnull=True)
                elif status_filter == 'initial':
                    queryset = queryset.filter(Q(status='INITIAL_COUNT') | Q(status='PENDING_COUNT', counted_balance__isnull=False))
                elif status_filter == 'recount':
                    queryset = queryset.filter(status__in=['SUPERVISOR_REJECTED', 'MANAGER_REJECTED'])
                elif status_filter == 'completed':
                    queryset = queryset.filter(status__in=['COUNTED', 'SUPERVISOR_APPROVED', 'MANAGER_REVIEW', 'FINAL_APPROVED'])
            elif as_role in ['supervisor', 'doc_supervisor']:
                if status_filter == 'counted':
                    queryset = queryset.filter(status='COUNTED')
                elif status_filter == 'recount':
                    queryset = queryset.filter(status__in=['SUPERVISOR_REJECTED', 'MANAGER_REJECTED'])
                else:
                    queryset = queryset.filter(status=status_filter)
            else:
                queryset = queryset.filter(status=status_filter)
            
        return queryset

    @action(detail=False, methods=['get'])
    def pool_tasks(self, request):
        as_role = request.query_params.get('as_role')
        warehouse_id = request.query_params.get('warehouse_id')
        queryset = CountTask.objects.all().select_related('item', 'counter', 'supervisor', 'created_by', 'modified_by').prefetch_related(photo_prefetch('item__photos'))
        
        if warehouse_id:
            queryset = queryset.filter(item__warehouse_id=warehouse_id)
            
        if as_role == 'counter':
            queryset = queryset.filter(counter__isnull=True, status='PENDING_COUNT')
        elif as_role == 'supervisor':
            queryset = queryset.filter(supervisor__isnull=True, status='COUNTED')
        elif as_role == 'manager':
            queryset = queryset.filter(assigned_manager__isnull=True, status='MANAGER_REVIEW')
        else:
            return Response({'error': 'نقش نامعتبر است.'}, status=400)

        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['post'])
    def claim_tasks(self, request):
        task_ids = request.data.get('task_ids', [])
        as_role = request.data.get('as_role')

        if not task_ids or not as_role:
            return Response({'error': 'شناسه تسک‌ها یا نقش ارسال نشده است.'}, status=400)

        with transaction.atomic():
            if as_role == 'counter':
                tasks = CountTask.objects.select_for_update().filter(
                    id__in=task_ids,
                    counter__isnull=True,
                    status='PENDING_COUNT'
                )
                valid_task_ids = list(tasks.values_list('id', flat=True))
                if not valid_task_ids:
                    broadcast_count_task_update()
                    return Response({'success': False, 'claimed_count': 0, 'message': 'این کالا(ها) قبلاً توسط انبارگردان دیگری بر عهده گرفته شده است.'})

                from ..models import Item
                item_ids = list(CountTask.objects.filter(id__in=valid_task_ids).values_list('item_id', flat=True))

                # Update CountTasks
                updated = CountTask.objects.filter(id__in=valid_task_ids).update(counter=request.user, updated_at=timezone.now())

                # Update Item field_assignee
                assignee_name = f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username
                Item.objects.filter(id__in=item_ids).update(field_assignee=assignee_name, updated_at=timezone.now())

            elif as_role == 'supervisor':
                tasks = CountTask.objects.select_for_update().filter(
                    id__in=task_ids,
                    supervisor__isnull=True,
                    status='COUNTED'
                )
                updated = tasks.update(supervisor=request.user, updated_at=timezone.now())
            elif as_role == 'manager':
                tasks = CountTask.objects.select_for_update().filter(
                    id__in=task_ids,
                    assigned_manager__isnull=True,
                    status='MANAGER_REVIEW'
                )
                updated = tasks.update(assigned_manager=request.user, updated_at=timezone.now())
            else:
                return Response({'error': 'نقش نامعتبر است.'}, status=400)

        broadcast_count_task_update()
        return Response({'success': True, 'claimed_count': updated})

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user, modified_by=self.request.user)

    @action(detail=False, methods=['post'])
    def bulk_submit(self, request):
        user = request.user
        task_ids = request.data.get('task_ids', [])
        sync_ids = request.data.get('sync_ids', [])
        warehouse_id = request.data.get('warehouse_id') or request.query_params.get('warehouse_id')

        # کلاینت آفلاین ممکن است فقط sync_id پایدار را داشته باشد
        if sync_ids:
            task_ids = list(set(task_ids) | set(
                CountTask.objects.filter(sync_id__in=sync_ids).values_list('id', flat=True)
            ))

        with transaction.atomic():
            tasks_qs = CountTask.objects.select_for_update().filter(
                counter=user,
                status='INITIAL_COUNT',
                counted_balance__isnull=False
            )
            if task_ids:
                tasks_qs = tasks_qs.filter(id__in=task_ids)

            if warehouse_id and str(warehouse_id) not in ['ALL', '-1']:
                tasks_qs = tasks_qs.filter(item__warehouse_id=warehouse_id)

            tasks_list = list(tasks_qs.select_related('item'))

            if not tasks_list:
                # Idempotency: اگر تسک‌های درخواستی قبلاً ارسال شده‌اند (مثلاً retry صف
                # آفلاین پس از گم شدن پاسخ در تونل)، موفقیت no-op برگردان نه ابهام.
                if task_ids:
                    already = CountTask.objects.filter(
                        id__in=task_ids, counter=user,
                        status__in=['COUNTED', 'MANAGER_REVIEW', 'FINAL_APPROVED'],
                    ).count()
                    if already > 0:
                        return Response({
                            'message': f'{already} مورد قبلاً ارسال شده بود.',
                            'already_submitted': already,
                        })
                return Response({'message': 'هیچ موردی برای ارسال یافت نشد.'})
                
            from warehouses.services import get_setting
            wh_setting_cache = {}

            def get_warehouse_req_sup(wh_id):
                if wh_id not in wh_setting_cache:
                    wh_setting_cache[wh_id] = get_setting('require_supervisor_approval', wh_id)
                return wh_setting_cache[wh_id]
            
            from ..models import CountTaskHistory
            histories = []
            counted_count = 0
            manager_count = 0
            for task in tasks_list:
                wh_id = task.item.warehouse_id if task.item else None
                req_sup_app = get_warehouse_req_sup(wh_id)
                task_req_sup = req_sup_app and not task.skip_supervisor
                target_status = 'COUNTED' if task_req_sup else 'MANAGER_REVIEW'
                if target_status == 'COUNTED':
                    counted_count += 1
                else:
                    manager_count += 1
                histories.append(CountTaskHistory(
                    task=task,
                    action_by=user,
                    action_type=target_status,
                    counted_balance=task.counted_balance,
                    note=task.counter_note
                ))
                task.status = target_status
                task.modified_by = user
                task.updated_at = timezone.now()  # bulk_update سیگنال auto_now را رد می‌کند

            count = len(tasks_list)
            if count > 0:
                CountTask.objects.bulk_update(tasks_list, ['status', 'modified_by', 'updated_at'])
            if histories:
                CountTaskHistory.objects.bulk_create(histories)
                
        if counted_count > 0 and manager_count > 0:
            msg = f'{count} مورد ارسال شد ({counted_count} مورد به سرپرست و {manager_count} مورد مستقیم به مدیر).'
        elif counted_count > 0:
            msg = f'{count} مورد به سرپرست ارسال شد.'
        else:
            msg = f'{count} مورد مستقیماً به مدیر ارسال شد.'
        broadcast_count_task_update()
        return Response({'message': msg})

    @action(detail=False, methods=['post'])
    def bulk_approve(self, request):
        user = request.user
        task_ids = request.data.get('task_ids', [])
        if not task_ids:
            return Response({'error': 'هیچ موردی انتخاب نشده است.'}, status=400)
            
        note = request.data.get('note', '')
        from ..models import CountTaskHistory
        from django.db import transaction
        
        with transaction.atomic():
            tasks = CountTask.objects.select_for_update(of=('self',)).filter(
                id__in=task_ids, supervisor=user, status__in=['COUNTED', 'MANAGER_REJECTED']
            )
            histories = []
            tasks_list = list(tasks)
            for task in tasks_list:
                task.supervisor_note = note
                task.status = 'MANAGER_REVIEW'
                task.modified_by = user
                task.updated_at = timezone.now()
                histories.append(CountTaskHistory(
                    task=task,
                    action_by=user,
                    action_type='MANAGER_REVIEW',
                    counted_balance=task.counted_balance,
                    note=note
                ))
                
            if tasks_list:
                CountTask.objects.bulk_update(tasks_list, ['status', 'supervisor_note', 'modified_by', 'updated_at'])
            if histories:
                CountTaskHistory.objects.bulk_create(histories)
                
        broadcast_count_task_update()
        return Response({'message': f'{len(tasks_list)} مورد تایید و برای مدیر ارسال شد.'})

    @action(detail=False, methods=['post'])
    def bulk_manager_approve(self, request):
        user = request.user
        task_ids = request.data.get('task_ids', [])
        note = request.data.get('note', '')
        
        if not task_ids:
            return Response({'error': 'هیچ موردی انتخاب نشده است.'}, status=400)
            
        from ..models import CountTaskHistory
        from django.db import transaction
        from decimal import Decimal, InvalidOperation
        from accounts.audit_utils import log_audit_event

        with transaction.atomic():
            tasks_qs = CountTask.objects.select_for_update(of=('self',)).filter(
                id__in=task_ids, status='MANAGER_REVIEW'
            ).select_related('item', 'item__warehouse')
            tasks = list(tasks_qs)
            if not tasks:
                return Response({'message': 'هیچ موردی در وضعیت بررسی مدیر یافت نشد.'})

            # بررسی دسترسی انبار برای کاربر (IDOR Check)
            if not user.is_superuser:
                from common.warehouse_scope import can_access_warehouse
                for task in tasks:
                    wh_id = task.item.warehouse_id if task.item else None
                    if not can_access_warehouse(user, wh_id):
                        return Response({'error': f'شما به انبار کالا #{task.id} دسترسی ندارید.'}, status=403)

            histories = []
            items_to_update = []
            wh_ids = set()
            
            for task in tasks:
                old_status = task.status
                old_counted = str(task.counted_balance)
                task.manager_note = note
                task.status = 'FINAL_APPROVED'
                task.modified_by = user
                task.updated_at = timezone.now()

                histories.append(CountTaskHistory(
                    task=task,
                    action_by=user,
                    action_type='FINAL_APPROVED',
                    counted_balance=task.counted_balance,
                    note=note
                ))
                
                # بروزرسانی کالای اصلی پس از تایید نهایی
                item = task.item
                if item:
                    if item.warehouse_id:
                        wh_ids.add(item.warehouse_id)
                    item.field_status = 'done'
                    if task.counted_balance is not None:
                        c_dec = _parse_decimal(task.counted_balance)
                        s_dec = _parse_decimal(item.bal4miv)
                        if c_dec is not None and s_dec is not None:
                            item.has_conflict = (c_dec != s_dec)
                        elif c_dec is not None:
                            item.has_conflict = (c_dec != Decimal('0'))
                        else:
                            item.has_conflict = False
                        item.inventory = task.counted_balance
                    item.modified_by = user
                    item.updated_at = timezone.now()
                    items_to_update.append(item)

                try:
                    item_label = item.fa_unic_code if item else f"#{task.id}"
                    wh = item.warehouse if item else None
                    log_audit_event(
                        module='manager',
                        action='APPROVE',
                        target_model='CountTask',
                        target_object_id=task.id,
                        target_repr=f"تایید نهایی شمارش کالا {item_label}",
                        warehouse=wh,
                        user=user if user.is_authenticated else None,
                        before_state={'status': old_status, 'counted_balance': old_counted},
                        after_state={'status': 'FINAL_APPROVED', 'counted_balance': str(task.counted_balance), 'note': note}
                    )
                except Exception:
                    pass
            
            CountTask.objects.bulk_update(tasks, ['status', 'manager_note', 'modified_by', 'updated_at'])
            count = len(tasks)
            
            if items_to_update:
                Item.objects.bulk_update(items_to_update, ['field_status', 'inventory', 'has_conflict', 'modified_by', 'updated_at'])
            
            if histories:
                CountTaskHistory.objects.bulk_create(histories)
            
        for wid in wh_ids:
            broadcast_count_task_update(warehouse_id=wid)
        if not wh_ids:
            broadcast_count_task_update()

        return Response({'message': f'{count} مورد به صورت گروهی تایید نهایی شد.'})

    @action(detail=True, methods=['post'])
    def reject(self, request, pk=None):
        task = self.get_object()
        note = request.data.get('note') or request.data.get('reason', '')
        
        if task.status not in ['COUNTED', 'MANAGER_REJECTED']:
            return Response({'error': 'فقط موارد شمرده شده قابل رد هستند.'}, status=400)
            
        old_status = task.status
        task.status = 'SUPERVISOR_REJECTED'
        task.supervisor_note = note
        task.modified_by = request.user
        task.updated_at = timezone.now()
        task.save()
        
        from ..models import CountTaskHistory
        from accounts.audit_utils import log_audit_event

        CountTaskHistory.objects.create(
            task=task,
            action_by=request.user,
            action_type='SUPERVISOR_REJECTED',
            counted_balance=task.counted_balance,
            note=note
        )

        try:
            item_label = task.item.fa_unic_code if task.item else f"#{task.id}"
            wh = task.item.warehouse if task.item else None
            log_audit_event(
                module='supervisor',
                action='REJECT',
                target_model='CountTask',
                target_object_id=task.id,
                target_repr=f"رد سرپرست برای بازشماری کالا {item_label}",
                warehouse=wh,
                user=request.user if request.user.is_authenticated else None,
                before_state={'status': old_status, 'counted_balance': str(task.counted_balance)},
                after_state={'status': 'SUPERVISOR_REJECTED', 'counted_balance': str(task.counted_balance), 'note': note}
            )
        except Exception:
            pass
        
        wh_id = task.item.warehouse_id if task.item else None
        broadcast_count_task_update(warehouse_id=wh_id)
        return Response({'message': 'مورد با موفقیت رد شد و به شمارشگر ارجاع داده شد.'})

    @action(detail=True, methods=['post'])
    def manager_reject(self, request, pk=None):
        """رد توسط مدیر — بسته به تنظیم سرپرست، به سرپرست یا انبارگردان برمی‌گردد"""
        task = self.get_object()
        note = request.data.get('note', '')
        
        if not note.strip():
            return Response({'error': 'لطفاً علت بازشماری را بنویسید.'}, status=400)
        
        if task.status != 'MANAGER_REVIEW':
            return Response({'error': 'فقط موارد در انتظار تایید مدیر قابل رد هستند.'}, status=400)
        
        from warehouses.services import get_setting
        from accounts.audit_utils import log_audit_event
        
        req_supervisor = get_setting('require_supervisor_approval', task.item.warehouse_id) if task.item else False
        old_status = task.status
        prev_counted = task.counted_balance

        if req_supervisor and task.supervisor:
            # ارسال به سرپرست (مقدار شمارش شده حفظ می‌شود تا سرپرست آن را بازبینی کند)
            target_status = 'MANAGER_REJECTED'
            target_msg = 'مورد برای بازشماری به سرپرست ارجاع شد.'
        else:
            # ارسال مستقیم به انبارگردان
            target_status = 'PENDING_COUNT'
            target_msg = 'مورد برای بازشماری مستقیماً به انبارگردان ارجاع شد.'
            task.counted_balance = None  # پاک کردن مقدار قبلی برای شمارش مجدد
        
        task.status = target_status
        task.manager_note = note
        task.modified_by = request.user
        task.updated_at = timezone.now()
        task.save()
        
        from ..models import CountTaskHistory
        CountTaskHistory.objects.create(
            task=task,
            action_by=request.user,
            action_type='MANAGER_REJECTED',
            counted_balance=prev_counted,
            note=note
        )

        try:
            item_label = task.item.fa_unic_code if task.item else f"#{task.id}"
            wh = task.item.warehouse if task.item else None
            log_audit_event(
                module='manager',
                action='REJECT',
                target_model='CountTask',
                target_object_id=task.id,
                target_repr=f"رد مدیر برای بازشماری کالا {item_label}",
                warehouse=wh,
                user=request.user if request.user.is_authenticated else None,
                before_state={'status': old_status, 'counted_balance': str(prev_counted)},
                after_state={'status': target_status, 'counted_balance': str(task.counted_balance), 'note': note}
            )
        except Exception:
            pass
        
        wh_id = task.item.warehouse_id if task.item else None
        broadcast_count_task_update(warehouse_id=wh_id)
        return Response({'message': target_msg})

    @action(detail=False, methods=['post'])
    def bulk_reject(self, request):
        """رد گروهی توسط سرپرست و ارجاع به شمارشگر"""
        user = request.user
        task_ids = request.data.get('task_ids', [])
        note = request.data.get('note', '')
        
        if not task_ids:
            return Response({'error': 'هیچ موردی انتخاب نشده است.'}, status=400)
            
        from ..models import CountTaskHistory
        from django.db import transaction
        from accounts.audit_utils import log_audit_event
        
        with transaction.atomic():
            tasks = CountTask.objects.select_for_update(of=('self',)).filter(
                id__in=task_ids, supervisor=user, status__in=['COUNTED', 'MANAGER_REJECTED']
            ).select_related('item', 'item__warehouse')
            histories = []
            tasks_list = list(tasks)
            wh_ids = set()

            for task in tasks_list:
                old_status = task.status
                if task.item and task.item.warehouse_id:
                    wh_ids.add(task.item.warehouse_id)
                task.supervisor_note = note
                task.status = 'SUPERVISOR_REJECTED'
                task.modified_by = user
                task.updated_at = timezone.now()
                histories.append(CountTaskHistory(
                    task=task,
                    action_by=user,
                    action_type='SUPERVISOR_REJECTED',
                    counted_balance=task.counted_balance,
                    note=note
                ))

                try:
                    item_label = task.item.fa_unic_code if task.item else f"#{task.id}"
                    wh = task.item.warehouse if task.item else None
                    log_audit_event(
                        module='supervisor',
                        action='REJECT',
                        target_model='CountTask',
                        target_object_id=task.id,
                        target_repr=f"رد سرپرست برای بازشماری کالا {item_label}",
                        warehouse=wh,
                        user=user if user.is_authenticated else None,
                        before_state={'status': old_status, 'counted_balance': str(task.counted_balance)},
                        after_state={'status': 'SUPERVISOR_REJECTED', 'counted_balance': str(task.counted_balance), 'note': note}
                    )
                except Exception:
                    pass
                
            if tasks_list:
                CountTask.objects.bulk_update(tasks_list, ['status', 'supervisor_note', 'modified_by', 'updated_at'])
            if histories:
                CountTaskHistory.objects.bulk_create(histories)
            
        for wid in wh_ids:
            broadcast_count_task_update(warehouse_id=wid)
        if not wh_ids:
            broadcast_count_task_update()

        return Response({'message': f'{len(tasks_list)} مورد با موفقیت رد شد و به شمارشگر ارجاع داده شد.'})

    @action(detail=False, methods=['post'])
    def bulk_manager_reject(self, request):
        """رد گروهی توسط مدیر — ارجاع به سرپرست یا شمارشگر بر اساس تنظیمات انبار"""
        user = request.user
        task_ids = request.data.get('task_ids', [])
        note = request.data.get('note', '')
        
        if not note.strip():
            return Response({'error': 'لطفاً علت بازشماری را بنویسید.'}, status=400)
            
        if not task_ids:
            return Response({'error': 'هیچ موردی انتخاب نشده است.'}, status=400)
            
        from warehouses.services import get_setting
        from ..models import CountTaskHistory
        from django.db import transaction
        from accounts.audit_utils import log_audit_event

        with transaction.atomic():
            tasks = CountTask.objects.select_for_update(of=('self',)).filter(
                id__in=task_ids, status='MANAGER_REVIEW'
            ).select_related('item', 'item__warehouse')
            tasks_list = list(tasks)
            if not tasks_list:
                return Response({'message': 'هیچ موردی برای رد یافت نشد.'})

            # بررسی دسترسی انبار برای کاربر (IDOR Check)
            if not user.is_superuser:
                from common.warehouse_scope import can_access_warehouse
                for task in tasks_list:
                    wh_id = task.item.warehouse_id if task.item else None
                    if not can_access_warehouse(user, wh_id):
                        return Response({'error': f'شما به انبار کالا #{task.id} دسترسی ندارید.'}, status=403)

            histories = []
            wh_ids = set()

            for task in tasks_list:
                old_status = task.status
                prev_counted = task.counted_balance
                wh_id = task.item.warehouse_id if task.item else None
                if wh_id:
                    wh_ids.add(wh_id)
                req_supervisor = get_setting('require_supervisor_approval', wh_id) if wh_id else False
                if req_supervisor and task.supervisor:
                    target_status = 'MANAGER_REJECTED'
                else:
                    target_status = 'PENDING_COUNT'
                    task.counted_balance = None
                    
                task.status = target_status
                task.manager_note = note
                task.modified_by = user
                task.updated_at = timezone.now()
                
                histories.append(CountTaskHistory(
                    task=task,
                    action_by=user,
                    action_type='MANAGER_REJECTED',
                    counted_balance=prev_counted,
                    note=note
                ))

                try:
                    item_label = task.item.fa_unic_code if task.item else f"#{task.id}"
                    wh = task.item.warehouse if task.item else None
                    log_audit_event(
                        module='manager',
                        action='REJECT',
                        target_model='CountTask',
                        target_object_id=task.id,
                        target_repr=f"رد مدیر برای بازشماری کالا {item_label}",
                        warehouse=wh,
                        user=user if user.is_authenticated else None,
                        before_state={'status': old_status, 'counted_balance': str(prev_counted)},
                        after_state={'status': target_status, 'counted_balance': str(task.counted_balance), 'note': note}
                    )
                except Exception:
                    pass
                
            if tasks_list:
                CountTask.objects.bulk_update(tasks_list, ['status', 'manager_note', 'counted_balance', 'modified_by', 'updated_at'])
            if histories:
                CountTaskHistory.objects.bulk_create(histories)
            
        for wid in wh_ids:
            broadcast_count_task_update(warehouse_id=wid)
        if not wh_ids:
            broadcast_count_task_update()

        return Response({'message': f'{len(tasks_list)} مورد با موفقیت رد شد و جهت بازشماری ارجاع داده شد.'})

    @action(detail=False, methods=['post'])
    def bulk_cancel(self, request):
        """لغو تخصیص گروهی — رکوردهای PENDING_COUNT و INITIAL_COUNT مجاز هستند."""
        user = request.user
        task_ids = request.data.get('task_ids', [])
        if not task_ids:
            return Response({'error': 'هیچ رکوردی انتخاب نشده است.'}, status=400)

        from accounts.audit_utils import log_audit_event
        from ..models import CountTaskHistory
        from django.db import transaction

        with transaction.atomic():
            all_tasks = CountTask.objects.select_for_update(of=('self',)).filter(id__in=task_ids).select_related('item', 'item__warehouse', 'counter')
            eligible_tasks = all_tasks.filter(status__in=['PENDING_COUNT', 'INITIAL_COUNT'])
            eligible_list = list(eligible_tasks)
            ineligible_count = all_tasks.count() - len(eligible_list)

            if not eligible_list:
                return Response(
                    {'error': 'هیچ‌یک از رکوردهای انتخاب شده قابل لغو تخصیص نیستند. فقط رکوردهای «در انتظار شمارش» و «شمارش اولیه» مجاز هستند.'},
                    status=400
                )

            item_ids = [t.item_id for t in eligible_list if t.item_id]
            task_ids_to_del = [t.id for t in eligible_list]
            wh_ids = set()

            for t in eligible_list:
                if t.item and t.item.warehouse_id:
                    wh_ids.add(t.item.warehouse_id)
                try:
                    item_label = t.item.fa_unic_code if t.item else f"#{t.id}"
                    wh = t.item.warehouse if t.item else None
                    log_audit_event(
                        module='dispatch',
                        action='CANCEL_ALLOCATION',
                        target_model='CountTask',
                        target_object_id=t.id,
                        target_repr=f"لغو تخصیص کالا {item_label}",
                        warehouse=wh,
                        user=user if user.is_authenticated else None,
                        before_state={'status': t.status, 'counter': t.counter.username if t.counter else None},
                        after_state={'status': 'CANCELLED_AND_DELETED'}
                    )
                except Exception:
                    pass

            # پاک کردن تاریخچه مربوطه و خود تسک‌ها به صورت نرم جهت حفظ سینک آفلاین
            soft_delete_queryset(CountTaskHistory.objects.filter(task_id__in=task_ids_to_del))
            deleted_count = len(eligible_list)
            soft_delete_queryset(CountTask.objects.filter(id__in=task_ids_to_del))

            # بازگرداندن وضعیت آیتم‌ها به استخر
            Item.objects.filter(id__in=item_ids).update(
                field_status='checking',
                field_assignee=None,
                updated_at=timezone.now()
            )

        msg = f'{deleted_count} وظیفه شمارش با موفقیت لغو تخصیص شد.'
        if ineligible_count > 0:
            msg += f' ({ineligible_count} رکورد به دلیل وضعیت نامعتبر نادیده گرفته شد.)'

        for wid in wh_ids:
            broadcast_count_task_update(warehouse_id=wid)
        if not wh_ids:
            broadcast_count_task_update()

        return Response({'message': msg})

    @action(detail=False, methods=['get'], url_path='get_export_columns')
    def get_export_columns(self, request):
        """لیست ستون‌های مجاز برای خروجی اکسل پیگیری شمارش"""
        columns = [
            {'key': 'warehouse_name',    'label': 'نام انبار'},
            {'key': 'fa_unic_code',      'label': 'کد یکتا'},
            {'key': 'description',       'label': 'شرح کالا'},
            {'key': 'counter_name',      'label': 'شمارنده'},
            {'key': 'supervisor_name',   'label': 'سرپرست'},
            {'key': 'manager_name',      'label': 'مدیر'},
            {'key': 'status',            'label': 'وضعیت'},
            {'key': 'counted_balance',   'label': 'موجودی شمارش شده'},
            {'key': 'inventory',         'label': 'موجودی سیستم'},
            {'key': 'difference',        'label': 'اختلاف'},
            {'key': 'created_at',        'label': 'تاریخ ایجاد'},
            {'key': 'updated_at',        'label': 'تاریخ بروزرسانی'},
            {'key': 'counter_note',      'label': 'یادداشت شمارنده'},
            {'key': 'supervisor_note',   'label': 'یادداشت سرپرست'},
            {'key': 'manager_note',      'label': 'یادداشت مدیر'},
        ]
        return Response(columns)

    @action(detail=False, methods=['post'], url_path='export_excel')
    def export_excel(self, request):
        """خروجی اکسل صفحه پیگیری شمارش با تاریخ شمسی و وضعیت فارسی"""
        import openpyxl
        from django.http import HttpResponse

        STATUS_FA = {
            'PENDING_COUNT':      'در انتظار شمارش',
            'INITIAL_COUNT':      'آماده ارسال (پیش‌نویس)',
            'COUNTED':            'شمارش شده (ارسال به سرپرست)',
            'SUPERVISOR_REVIEW':  'در بررسی سرپرست',
            'SUPERVISOR_APPROVED':'تایید سرپرست',
            'SUPERVISOR_REJECTED':'رد شده توسط سرپرست (بازشماری)',
            'MANAGER_REVIEW':     'در بررسی مدیر',
            'MANAGER_REJECTED':   'رد شده توسط مدیر (بازشماری)',
            'FINAL_APPROVED':     'تأیید نهایی',
        }

        def to_jalali(dt):
            if not dt:
                return ''
            try:
                import jdatetime
                from django.utils import timezone
                if timezone.is_aware(dt):
                    dt = timezone.localtime(dt)
                jdt = jdatetime.datetime.fromgregorian(datetime=dt)
                return jdt.strftime('%Y/%m/%d %H:%M:%S')
            except Exception as e:
                # اگر خطایی رخ داد، لاگ بگیریم تا اشکال‌زدایی راحت‌تر باشد
                import logging
                logging.getLogger(__name__).error(f"Date conversion error: {e}")
                return str(dt)

        data_scope    = request.data.get('data_scope', 'all')
        columns_scope = request.data.get('columns_scope', 'all_db')
        columns_list  = request.data.get('columns_list', [])

        # اعمال همان فیلترهای get_queryset
        from django.http import QueryDict
        original_query_params = request._request.GET
        try:
            q = QueryDict(mutable=True)
            for k, v in request.data.items():
                if isinstance(v, list):
                    q.setlist(k, [str(x) for x in v])
                elif v is not None:
                    q[k] = str(v)
            request._request.GET = q

            if data_scope == 'selected':
                selected_ids = request.data.get('selected_ids', [])
                queryset = self.get_queryset().filter(id__in=selected_ids)
            else:
                queryset = self.get_queryset()
        finally:
            request._request.GET = original_query_params

        # ستون‌ها
        ALL_COLUMNS = [
            {'key': 'warehouse_name',    'label': 'نام انبار'},
            {'key': 'fa_unic_code',      'label': 'کد یکتا'},
            {'key': 'description',       'label': 'شرح کالا'},
            {'key': 'counter_name',      'label': 'شمارنده'},
            {'key': 'supervisor_name',   'label': 'سرپرست'},
            {'key': 'manager_name',      'label': 'مدیر'},
            {'key': 'status',            'label': 'وضعیت'},
            {'key': 'counted_balance',   'label': 'موجودی شمارش شده'},
            {'key': 'inventory',         'label': 'موجودی سیستم'},
            {'key': 'difference',        'label': 'اختلاف'},
            {'key': 'created_at',        'label': 'تاریخ ایجاد'},
            {'key': 'updated_at',        'label': 'تاریخ بروزرسانی'},
            {'key': 'counter_note',      'label': 'یادداشت شمارنده'},
            {'key': 'supervisor_note',   'label': 'یادداشت سرپرست'},
            {'key': 'manager_note',      'label': 'یادداشت مدیر'},
        ]
        all_keys = [c['key'] for c in ALL_COLUMNS]
        key_to_label = {c['key']: c['label'] for c in ALL_COLUMNS}

        if columns_scope in ('visible', 'custom') and columns_list:
            selected_keys = [k for k in columns_list if k in all_keys]
            if not selected_keys:
                selected_keys = all_keys
        else:
            selected_keys = all_keys

        # بررسی وضعیت شمارش کور برای هر انبار جهت حفظ محرمانگی داده‌ها
        from warehouses.services import get_setting
        warehouse_blind_cache = {}

        def check_is_blind(task):
            wh_id = task.item.warehouse_id if task.item else None
            if wh_id not in warehouse_blind_cache:
                warehouse_blind_cache[wh_id] = (get_setting('blind_counting', wh_id) == 'blind')
            return warehouse_blind_cache[wh_id]

        def get_cell(task, key):
            if key == 'warehouse_name':
                if task.item and task.item.warehouse:
                    return getattr(task.item.warehouse, 'project_name', None) or task.item.warehouse.name or ''
                return ''
            elif key == 'fa_unic_code':
                return getattr(task.item, 'fa_unic_code', '') if task.item else ''
            elif key == 'description':
                return getattr(task.item, 'description', '') if task.item else ''
            elif key == 'counter_name':
                if task.counter:
                    return f"{task.counter.first_name} {task.counter.last_name}".strip() or task.counter.username
                return ''
            elif key == 'supervisor_name':
                if task.supervisor:
                    return f"{task.supervisor.first_name} {task.supervisor.last_name}".strip() or task.supervisor.username
                return ''
            elif key == 'manager_name':
                if task.assigned_manager:
                    return f"{task.assigned_manager.first_name} {task.assigned_manager.last_name}".strip() or task.assigned_manager.username
                return ''
            elif key == 'status':
                return STATUS_FA.get(task.status, task.status)
            elif key == 'counted_balance':
                return task.counted_balance if task.counted_balance is not None else ''
            elif key == 'inventory':
                if check_is_blind(task):
                    return ''
                base_bal = getattr(task.item, 'bal4miv', None) if task.item else None
                if base_bal is None and task.item:
                    base_bal = getattr(task.item, 'inventory', '')
                return base_bal if base_bal is not None else ''
            elif key == 'difference':
                if check_is_blind(task):
                    return ''
                base_bal = getattr(task.item, 'bal4miv', None) if task.item else None
                if base_bal is None and task.item:
                    base_bal = getattr(task.item, 'inventory', None)
                cnt = task.counted_balance
                if base_bal is not None and cnt is not None:
                    try:
                        return float(cnt) - float(base_bal)
                    except Exception:
                        return ''
                return ''
            elif key == 'created_at':
                return to_jalali(task.created_at)
            elif key == 'updated_at':
                return to_jalali(task.updated_at)
            elif key == 'counter_note':
                return task.counter_note or ''
            elif key == 'supervisor_note':
                return task.supervisor_note or ''
            elif key == 'manager_note':
                return task.manager_note or ''
            return ''

        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'پیگیری شمارش'
        
        # راست‌چین کردن شیت
        ws.sheet_view.rightToLeft = True

        # استایل‌های هدر ۲ سطری
        header_fa_font = Font(name='Tahoma', bold=True, color='FFFFFF', size=11)
        header_fa_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # سرمه‌ای تیره
        header_en_font = Font(name='Consolas', size=9, color='64748B', italic=True) # فیلد دیتابیس
        header_en_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid') # طوسی روشن
        
        cell_font = Font(name='Tahoma', size=10)
        
        # رنگ‌های یکی‌درمیان ردیف‌ها
        fill_even = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
        )

        # ساخت هدر ۲ سطری (سطر ۱: فارسی، سطر ۲: کلید فنی دیتابیس)
        headers_fa = [key_to_label[k] for k in selected_keys]
        headers_en = [k for k in selected_keys]
        ws.append(headers_fa)
        ws.append(headers_en)
        
        # استایل سطر اول
        for col_idx, cell in enumerate(ws[1], 1):
            cell.font = header_fa_font
            cell.fill = header_fa_fill
            cell.alignment = center_alignment
            cell.border = thin_border
            ws.row_dimensions[1].height = 28

        # استایل سطر دوم
        for col_idx, cell in enumerate(ws[2], 1):
            cell.font = header_en_font
            cell.fill = header_en_fill
            cell.alignment = center_alignment
            cell.border = thin_border
            ws.row_dimensions[2].height = 18

        # فریز کردن ۲ سطر اول (سطر ۳ به بعد اسکرول شود)
        ws.freeze_panes = 'A3'

        # محاسبه عرض اولیه بر اساس هدرها
        col_widths = {i: max(len(headers_fa[i]), len(headers_en[i])) + 6 for i in range(len(headers_fa))}

        # درج داده‌ها از سطر ۳
        row_idx = 3
        for task in queryset.select_related('item', 'item__warehouse', 'counter', 'supervisor', 'assigned_manager').iterator(chunk_size=500):
            row_data = [str(get_cell(task, k)) for k in selected_keys]
            ws.append(row_data)
            
            # تعیین رنگ پس‌زمینه ردیف
            current_fill = fill_even if row_idx % 2 == 0 else fill_odd
            
            # استایل‌دهی به سلول‌ها و محاسبه عرض
            for col_idx, val in enumerate(row_data):
                cell = ws.cell(row=row_idx, column=col_idx + 1)
                cell.font = cell_font
                cell.fill = current_fill
                cell.alignment = center_alignment
                cell.border = thin_border
                
                # بروزرسانی عرض ستون (محدود کردن به حداکثر 60)
                lines = val.split('\n')
                max_line_len = max([len(line) for line in lines]) if lines else 0
                if max_line_len > col_widths.get(col_idx, 10):
                    col_widths[col_idx] = min(max_line_len + 4, 60)
            
            ws.row_dimensions[row_idx].height = 22
            row_idx += 1

        # اعمال عرض ستون‌ها
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx + 1)].width = max(width, 14)

        # افزودن قابلیت فیلتر روی هدرها
        if row_idx > 3:
            last_col_letter = get_column_letter(len(selected_keys))
            ws.auto_filter.ref = f"A1:{last_col_letter}{row_idx - 1}"

        from io import BytesIO
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="count_tracking_export.xlsx"'
        return response

__all__ = [
    'CountTaskViewSet',
]
