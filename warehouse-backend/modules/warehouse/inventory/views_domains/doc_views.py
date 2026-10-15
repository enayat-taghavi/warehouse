# -*- coding: utf-8 -*-
"""
ViewSet for inventory document tasks (اسناد انبار - ورود و خروج کالا).
"""
from decimal import Decimal
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from common.warehouse_scope import scope_queryset, can_access_warehouse
from ..models import DocTask, DocTaskHistory, Item
from ..serializers import DocTaskSerializer, photo_prefetch
from ..signals import broadcast_doc_task_update
from .helpers import _create_doc_task_snapshot

class DocTaskViewSet(viewsets.ModelViewSet):
    serializer_class = DocTaskSerializer
    permission_classes = [IsAuthenticated]

    def get_object(self):
        """
        زیرساخت سینک: اگر pk عددی نبود، به‌عنوان sync_id تفسیر می‌شود (تسک ۲۲).
        """
        pk = self.kwargs.get(self.lookup_url_kwarg or self.lookup_field)
        if pk is not None and not str(pk).isdigit():
            from django.shortcuts import get_object_or_404
            queryset = self.filter_queryset(self.get_queryset())
            obj = get_object_or_404(queryset, sync_id=pk)
            self.check_object_permissions(self.request, obj)
            return obj
        return super().get_object()

    def perform_create(self, serializer):
        # بررسی idempotency بر مبنای sync_id (تسک ۱۴)
        sync_id = serializer.validated_data.get('sync_id')
        existing = DocTask.all_objects.filter(sync_id=sync_id).first() if sync_id else None
        if existing:
            serializer.instance = existing
            serializer.save(is_deleted=False)
            return
        serializer.save(created_by=self.request.user)

    def update(self, request, *args, **kwargs):
        # تشخیص تداخل خوش‌بینانه بر اساس base_updated_at (تسک ۶)
        base_raw = request.data.get('base_updated_at') or request.data.get('baseUpdatedAt')
        if base_raw:
            from django.utils.dateparse import parse_datetime
            base = parse_datetime(str(base_raw))
            if base is not None:
                if timezone.is_naive(base):
                    base = timezone.make_aware(base)
                instance = self.get_object()
                if (instance.updated_at - base).total_seconds() > 0.001:
                    return Response({
                        'detail': 'conflict',
                        'server_record': self.get_serializer(instance).data,
                    }, status=409)
        return super().update(request, *args, **kwargs)

    def get_permissions(self):
        from accounts.permissions import HasMenuAccess
        from rest_framework.permissions import IsAuthenticated
        
        if self.action in ['bulk_approve', 'reject']:
            return [HasMenuAccess('perm_doc_approve_action') | HasMenuAccess('view_wh_doc_approvals') | HasMenuAccess('can_act_as_doc_supervisor') | HasMenuAccess('can_act_as_manager') | HasMenuAccess('view_sys_manager_review')]
        elif self.action in ['bulk_manager_approve', 'manager_reject']:
            return [HasMenuAccess('perm_doc_approve_action') | HasMenuAccess('view_wh_doc_approvals') | HasMenuAccess('can_act_as_manager') | HasMenuAccess('view_sys_manager_review')]
        return [IsAuthenticated()]

    def get_queryset(self):
        user = self.request.user
        queryset = DocTask.objects.filter(item__is_deleted=False).select_related('item', 'doc_worker', 'doc_supervisor', 'created_by', 'modified_by').prefetch_related(photo_prefetch('item__photos'))
        
        from common.tenant_scope import scope_tenant_queryset
        queryset = scope_tenant_queryset(queryset, user, company_field='item__warehouse__company_id', request=self.request)

        as_role = self.request.query_params.get('as_role')
        warehouse_id = self.request.query_params.get('warehouse_id')
        
        if warehouse_id and str(warehouse_id) not in ['ALL', '-1']:
            queryset = queryset.filter(item__warehouse_id=warehouse_id)
        else:
            queryset = scope_queryset(queryset, user, field='item__warehouse_id')
        
        if as_role == 'doc_worker':
            queryset = queryset.filter(doc_worker=user)
        elif as_role == 'doc_supervisor':
            queryset = queryset.filter(doc_supervisor=user)
        elif as_role == 'manager':
            queryset = queryset.filter(assigned_manager=user, status='DOC_MANAGER_REVIEW')
        elif as_role == 'tracking':
            show_completed = self.request.query_params.get('show_completed', 'false').lower() == 'true'
            if not show_completed:
                queryset = queryset.exclude(status='DOC_FINAL_APPROVED')
        else:
            # Fallback
            from django.db.models import Q
            if user.is_superuser or user.groups.filter(name__in=['admin', 'manager']).exists() or user.has_perm('accounts.view_sys_manager_review'):
                queryset = queryset.filter(Q(assigned_manager=user) | Q(assigned_manager__isnull=True) | Q(doc_worker=user) | Q(doc_supervisor=user))
            else:
                queryset = queryset.filter(Q(doc_worker=user) | Q(doc_supervisor=user) | Q(assigned_manager=user))

        # اعمال فیلترهای تکمیلی جستجو، وضعیت و تاریخ
        q = self.request.query_params.get('q') or self.request.data.get('q')
        if q:
            q_clean = str(q).strip()
            from django.db.models import Q
            queryset = queryset.filter(
                Q(item__fa_unic_code__icontains=q_clean) |
                Q(item__description__icontains=q_clean) |
                Q(item__tag__icontains=q_clean) |
                Q(item__po__icontains=q_clean) |
                Q(item__pk_number__icontains=q_clean) |
                Q(item__new_location__icontains=q_clean) |
                Q(inv_rti_number__icontains=q_clean) |
                Q(added_rti_no__icontains=q_clean) |
                Q(doc_supplier__icontains=q_clean)
            )

        status_param = self.request.query_params.get('status') or self.request.data.get('status')
        if status_param and status_param != 'all':
            if status_param == 'untouched':
                queryset = queryset.filter(status='PENDING_DOC', inv_rti_number__isnull=True, added_rti_no__isnull=True, price_amount__isnull=True, doc_supplier__isnull=True)
            elif status_param == 'rejected':
                queryset = queryset.filter(status__in=['DOC_SUPERVISOR_REJECTED', 'DOC_MANAGER_REJECTED'])
            elif status_param == 'ready':
                queryset = queryset.filter(status='PENDING_DOC').exclude(inv_rti_number__isnull=True, added_rti_no__isnull=True, price_amount__isnull=True, doc_supplier__isnull=True)
            elif status_param == 'completed':
                queryset = queryset.filter(status='DOC_FINAL_APPROVED')
            else:
                queryset = queryset.filter(status=status_param)

        date_param = self.request.query_params.get('date') or self.request.data.get('date')
        if date_param and date_param != 'all':
            from django.utils import timezone
            import datetime
            now = timezone.now()
            today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
            if date_param == 'today':
                queryset = queryset.filter(updated_at__gte=today_start)
            elif date_param == 'yesterday':
                yesterday_start = today_start - datetime.timedelta(days=1)
                queryset = queryset.filter(updated_at__gte=yesterday_start, updated_at__lt=today_start)
            elif date_param == 'week':
                week_start = today_start - datetime.timedelta(days=7)
                queryset = queryset.filter(updated_at__gte=week_start)
            
        return queryset

    @action(detail=False, methods=['get'])
    def pool_tasks(self, request):
        as_role = request.query_params.get('as_role')
        warehouse_id = request.query_params.get('warehouse_id')
        queryset = DocTask.objects.filter(item__is_deleted=False).select_related('item', 'doc_worker', 'doc_supervisor', 'created_by', 'modified_by').prefetch_related(photo_prefetch('item__photos'))
        
        if warehouse_id and str(warehouse_id) not in ['ALL', '-1']:
            queryset = queryset.filter(item__warehouse_id=warehouse_id)
            
        if as_role == 'doc_worker':
            queryset = queryset.filter(doc_worker__isnull=True, status='PENDING_DOC')
        elif as_role == 'doc_supervisor':
            queryset = queryset.filter(doc_supervisor__isnull=True, status='DOC_PROCESSED')
        elif as_role == 'manager':
            queryset = queryset.filter(assigned_manager__isnull=True, status='DOC_MANAGER_REVIEW')
        else:
            return Response({'error': 'نقش نامعتبر است.'}, status=400)

        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['post'])
    def claim_tasks(self, request):
        task_ids = request.data.get('task_ids', [])
        sync_ids = request.data.get('sync_ids', [])
        as_role = request.data.get('as_role')

        if (not task_ids and not sync_ids) or not as_role:
            return Response({'error': 'لیست شناسه‌ها یا نقش ارسال نشده است.'}, status=400)

        from django.db.models import Q
        q_filter = Q()
        if task_ids:
            q_filter |= Q(id__in=task_ids)
        if sync_ids:
            q_filter |= Q(sync_id__in=sync_ids)

        with transaction.atomic():
            tasks = DocTask.objects.select_for_update().filter(q_filter).select_related('item')
            from ..models import DocTaskHistory

            if as_role == 'doc_worker':
                tasks = tasks.filter(doc_worker__isnull=True, status='PENDING_DOC')
                valid_tasks = list(tasks)
                if not valid_tasks:
                    broadcast_doc_task_update()
                    return Response({'success': True, 'claimed_count': 0})

                valid_task_ids = [t.id for t in valid_tasks]
                item_ids = [t.item_id for t in valid_tasks if t.item_id]

                # Update DocTasks
                updated = DocTask.objects.filter(id__in=valid_task_ids).update(doc_worker=request.user, updated_at=timezone.now())

                # Update Item doc_assignee
                assignee_name = f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username
                from ..models import Item
                Item.objects.filter(id__in=item_ids).update(doc_assignee=assignee_name, updated_at=timezone.now())

                # Log History
                histories = [
                    DocTaskHistory(
                        task=t,
                        action_by=request.user,
                        action_type='CLAIMED',
                        note=f'بر عهده گرفته شد توسط کارشناس مالی ({assignee_name})',
                        data_snapshot=_create_doc_task_snapshot(t)
                    ) for t in valid_tasks
                ]
                DocTaskHistory.objects.bulk_create(histories)

            elif as_role == 'doc_supervisor':
                tasks = tasks.filter(doc_supervisor__isnull=True, status='DOC_PROCESSED')
                valid_tasks = list(tasks)
                if not valid_tasks:
                    broadcast_doc_task_update()
                    return Response({'success': True, 'claimed_count': 0})
                valid_task_ids = [t.id for t in valid_tasks]
                updated = DocTask.objects.filter(id__in=valid_task_ids).update(doc_supervisor=request.user, updated_at=timezone.now())
                supervisor_name = f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username
                histories = [
                    DocTaskHistory(
                        task=t,
                        action_by=request.user,
                        action_type='CLAIMED',
                        note=f'بر عهده گرفته شد توسط سرپرست اسناد ({supervisor_name})',
                        data_snapshot=_create_doc_task_snapshot(t)
                    ) for t in valid_tasks
                ]
                DocTaskHistory.objects.bulk_create(histories)

            elif as_role == 'manager':
                tasks = tasks.filter(assigned_manager__isnull=True, status='DOC_MANAGER_REVIEW')
                valid_tasks = list(tasks)
                if not valid_tasks:
                    broadcast_doc_task_update()
                    return Response({'success': True, 'claimed_count': 0})
                valid_task_ids = [t.id for t in valid_tasks]
                updated = DocTask.objects.filter(id__in=valid_task_ids).update(assigned_manager=request.user, updated_at=timezone.now())
                manager_name = f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username
                histories = [
                    DocTaskHistory(
                        task=t,
                        action_by=request.user,
                        action_type='CLAIMED',
                        note=f'بر عهده گرفته شد توسط مدیر ({manager_name})',
                        data_snapshot=_create_doc_task_snapshot(t)
                    ) for t in valid_tasks
                ]
                DocTaskHistory.objects.bulk_create(histories)
            else:
                return Response({'error': 'نقش نامعتبر است.'}, status=400)

        broadcast_doc_task_update()
        return Response({'success': True, 'claimed_count': updated})

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user, modified_by=self.request.user)

    def perform_update(self, serializer):
        instance = serializer.instance
        old_status = instance.status
        new_status = serializer.validated_data.get('status', old_status)
        
        updated_instance = serializer.save(modified_by=self.request.user)
        
        if old_status != new_status:
            from ..models import DocTaskHistory
            
            note = ''
            if new_status in ['DOC_MANAGER_REJECTED', 'DOC_FINAL_APPROVED']:
                note = updated_instance.manager_note
            elif new_status == 'DOC_SUPERVISOR_REJECTED':
                note = updated_instance.supervisor_note
            elif new_status == 'DOC_PROCESSED':
                note = updated_instance.worker_note
                
            DocTaskHistory.objects.create(
                task=updated_instance,
                action_by=self.request.user,
                action_type=new_status,
                note=note,
                data_snapshot=_create_doc_task_snapshot(updated_instance)
            )

    @action(detail=False, methods=['post'])
    def bulk_submit(self, request):
        user = request.user
        task_ids = request.data.get('task_ids', [])
        sync_ids = request.data.get('sync_ids', [])
        warehouse_id = request.data.get('warehouse_id') or request.query_params.get('warehouse_id')
        
        from django.db.models import Q
        q_filter = Q()
        if task_ids:
            q_filter |= Q(id__in=task_ids)
        if sync_ids:
            q_filter |= Q(sync_id__in=sync_ids)
            
        with transaction.atomic():
            if task_ids or sync_ids:
                tasks = DocTask.objects.select_for_update().filter(q_filter, doc_worker=user, status__in=['PENDING_DOC', 'DOC_SUPERVISOR_REJECTED', 'DOC_MANAGER_REJECTED']).select_related('item')
            else:
                tasks = DocTask.objects.select_for_update().filter(doc_worker=user, status__in=['PENDING_DOC', 'DOC_SUPERVISOR_REJECTED', 'DOC_MANAGER_REJECTED']).select_related('item')
            
            if warehouse_id and str(warehouse_id) not in ['ALL', '-1']:
                tasks = tasks.filter(item__warehouse_id=warehouse_id)
            
            def has_doc_info(t):
                return bool(
                    t.status in ['DOC_SUPERVISOR_REJECTED', 'DOC_MANAGER_REJECTED'] or
                    t.inv_rti_number or t.added_rti_no or t.doc_supplier or
                    t.price_amount or t.total_value or t.similar_unit_price or
                    t.invoice_type or t.invoice_date or t.folder_address or
                    t.stamp or t.signature or (t.worker_note and str(t.worker_note).strip())
                )

            tasks_list = [t for t in tasks if has_doc_info(t)]
            if not tasks_list:
                if task_ids or sync_ids:
                    already = DocTask.objects.filter(
                        Q(id__in=task_ids) | Q(sync_id__in=sync_ids),
                        doc_worker=user,
                        status__in=['DOC_PROCESSED', 'DOC_MANAGER_REVIEW', 'DOC_FINAL_APPROVED']
                    ).count()
                    if already > 0:
                        return Response({
                            'message': f'{already} مورد قبلاً ارسال شده بود.',
                            'already_submitted': already
                        })
                return Response({'message': 'هیچ سند مالی تکمیل‌شده‌ای برای ارجاع یافت نشد.'})
                
            from warehouses.services import get_setting
            from ..models import DocTaskHistory
            histories = []
            
            # بررسی تنظیم تایید سرپرست به تفکیک انبار هر تسک
            wh_settings_cache = {}
            target_status_counts = {'DOC_PROCESSED': 0, 'DOC_MANAGER_REVIEW': 0}
            
            now_time = timezone.now()
            for task in tasks_list:
                wh_id = task.item.warehouse_id if task.item else None
                if wh_id not in wh_settings_cache:
                    wh_settings_cache[wh_id] = get_setting('require_doc_supervisor_approval', wh_id)
                req_sup_app = wh_settings_cache[wh_id]
                
                task_req_sup = req_sup_app and not task.skip_supervisor
                target_status = 'DOC_PROCESSED' if task_req_sup else 'DOC_MANAGER_REVIEW'
                task.status = target_status
                task.modified_by = user
                task.updated_at = now_time
                target_status_counts[target_status] += 1
                
                histories.append(DocTaskHistory(
                    task=task,
                    action_by=user,
                    action_type=target_status,
                    note=task.worker_note,
                    data_snapshot=_create_doc_task_snapshot(task)
                ))
                
            count = len(tasks_list)
            if count > 0:
                DocTask.objects.bulk_update(tasks_list, ['status', 'modified_by', 'updated_at'])
                broadcast_doc_task_update()
            if histories:
                DocTaskHistory.objects.bulk_create(histories)
                
            if target_status_counts['DOC_PROCESSED'] > 0 and target_status_counts['DOC_MANAGER_REVIEW'] > 0:
                msg = f'{count} کالا ارسال شد ({target_status_counts["DOC_PROCESSED"]} مورد جهت بررسی سرپرست و {target_status_counts["DOC_MANAGER_REVIEW"]} مورد مستقیماً به مدیر).'
            elif target_status_counts['DOC_PROCESSED'] > 0:
                msg = f'{count} کالا جهت بررسی سرپرست ارسال شد.'
            else:
                msg = f'{count} کالا مستقیماً جهت بررسی مدیر ارسال شد.'
                
            return Response({'message': msg})

    @action(detail=False, methods=['post'])
    def bulk_approve(self, request):
        user = request.user
        task_ids = request.data.get('task_ids', [])
        if not task_ids:
            return Response({'error': 'هیچ کالایی انتخاب نشده است.'}, status=400)
            
        from django.db.models import Q
        from django.db import transaction
        from ..models import DocTaskHistory
        note = request.data.get('note', '')
        
        with transaction.atomic():
            tasks = DocTask.objects.select_for_update().filter(id__in=task_ids, status__in=['DOC_PROCESSED', 'DOC_MANAGER_REJECTED'])
            if not (user.is_superuser or user.groups.filter(name__in=['admin', 'manager']).exists()):
                tasks = tasks.filter(Q(doc_supervisor=user) | Q(doc_supervisor__isnull=True))
            
            histories = []
            tasks_list = list(tasks)
            for task in tasks_list:
                task.supervisor_note = note
                task.status = 'DOC_MANAGER_REVIEW'
                task.modified_by = user
                task.updated_at = timezone.now()
                if task.doc_supervisor_id is None:
                    task.doc_supervisor = user
                histories.append(DocTaskHistory(
                    task=task,
                    action_by=user,
                    action_type='DOC_MANAGER_REVIEW',
                    note=note,
                    data_snapshot=_create_doc_task_snapshot(task)
                ))
                
            if tasks_list:
                DocTask.objects.bulk_update(tasks_list, ['status', 'modified_by', 'supervisor_note', 'doc_supervisor', 'updated_at'])
            if histories:
                DocTaskHistory.objects.bulk_create(histories)
                
        broadcast_doc_task_update()
        return Response({'message': f'{len(histories)} رکورد جهت تایید نهایی مدیر ارسال شد.'})

    @action(detail=False, methods=['post'])
    def reject(self, request):
        task_ids = request.data.get('task_ids', [])
        note = request.data.get('note', '')
        
        if not task_ids:
            return Response({'error': 'هیچ رکوردی انتخاب نشده است.'}, status=400)
            
        from django.db import transaction
        from ..models import DocTaskHistory
        
        with transaction.atomic():
            tasks = DocTask.objects.select_for_update().filter(id__in=task_ids, status__in=['DOC_PROCESSED', 'DOC_MANAGER_REJECTED'])
            histories = []
            tasks_list = list(tasks)
            for task in tasks_list:
                task.supervisor_note = note
                task.status = 'DOC_SUPERVISOR_REJECTED'
                task.modified_by = request.user
                task.updated_at = timezone.now()
                histories.append(DocTaskHistory(
                    task=task,
                    action_by=request.user,
                    action_type='DOC_SUPERVISOR_REJECTED',
                    note=note,
                    data_snapshot=_create_doc_task_snapshot(task)
                ))
                
            if tasks_list:
                DocTask.objects.bulk_update(tasks_list, ['status', 'modified_by', 'supervisor_note', 'updated_at'])
            if histories:
                DocTaskHistory.objects.bulk_create(histories)
                
        broadcast_doc_task_update()
        return Response({'message': f'{len(histories)} رکورد به بررسی‌کننده اسناد برگشت داده شد.'})

    @action(detail=False, methods=['post'])
    def manager_reject(self, request):
        task_ids = request.data.get('task_ids', [])
        note = request.data.get('note', '')
        
        if not task_ids:
            return Response({'error': 'هیچ رکوردی انتخاب نشده است.'}, status=400)
            
        tasks = DocTask.objects.filter(id__in=task_ids, status='DOC_MANAGER_REVIEW')
        
        from ..models import DocTaskHistory
        histories = []
        for task in tasks:
            task.manager_note = note
            task.status = 'DOC_MANAGER_REJECTED'
            task.modified_by = request.user
            histories.append(DocTaskHistory(
                task=task,
                action_by=request.user,
                action_type='DOC_MANAGER_REJECTED',
                note=note,
                data_snapshot=_create_doc_task_snapshot(task)
            ))
            
        if tasks:
            DocTask.objects.bulk_update(tasks, ['status', 'modified_by', 'manager_note'])
            DocTaskHistory.objects.bulk_create(histories)
            broadcast_doc_task_update()
            
        return Response({'message': f'{len(histories)} رکورد به سرپرست اسناد (یا بررسی‌کننده) برگشت داده شد.'})

    @action(detail=False, methods=['post'])
    def bulk_manager_approve(self, request):
        task_ids = request.data.get('task_ids', [])
        note = request.data.get('note', '')
        if not task_ids:
            return Response({'error': 'هیچ رکوردی انتخاب نشده است.'}, status=400)
            
        tasks_list = list(DocTask.objects.filter(id__in=task_ids, status='DOC_MANAGER_REVIEW').select_related('item'))
        if not tasks_list:
            return Response({'message': 'هیچ رکوردی در انتظار تایید مدیر یافت نشد.'})
        
        from ..models import DocTaskHistory, Item
        from django.db import transaction
        
        histories = []
        items_to_update = []
        now = timezone.now()
        
        with transaction.atomic():
            for task in tasks_list:
                if note:
                    task.manager_note = note
                task.status = 'DOC_FINAL_APPROVED'
                task.modified_by = request.user
                histories.append(DocTaskHistory(
                    task=task,
                    action_by=request.user,
                    action_type='DOC_FINAL_APPROVED',
                    note=task.manager_note,
                    data_snapshot=_create_doc_task_snapshot(task)
                ))
                
                # انتقال و همگام‌سازی فیلدهای مالی به رکورد کالا (Item)
                item = task.item
                if item:
                    item.doc_status = 'done'
                    item.price_amount = task.price_amount
                    item.total_value = task.total_value
                    item.similar_unit_price = task.similar_unit_price
                    item.currency = task.currency
                    item.invoice_type = task.invoice_type
                    item.invoice_date = str(task.invoice_date) if task.invoice_date else None
                    item.inv_rti_number = task.inv_rti_number
                    item.added_rti_no = task.added_rti_no
                    item.page_row = str(task.page_row) if task.page_row is not None else None
                    item.invoice_page = str(task.invoice_page) if task.invoice_page is not None else None
                    item.doc_supplier = task.doc_supplier
                    item.folder_address = task.folder_address
                    if task.stamp is not None:
                        item.stamp = 'دارد' if task.stamp else 'ندارد'
                    if task.signature is not None:
                        item.signature = 'دارد' if task.signature else 'ندارد'
                    item.modified_by = request.user
                    item.updated_at = now
                    items_to_update.append(item)
                
            if tasks_list:
                DocTask.objects.bulk_update(tasks_list, ['status', 'modified_by', 'manager_note'])
                DocTaskHistory.objects.bulk_create(histories)
                if items_to_update:
                    Item.objects.bulk_update(items_to_update, [
                        'doc_status', 'price_amount', 'total_value', 'similar_unit_price',
                        'currency', 'invoice_type', 'invoice_date', 'inv_rti_number',
                        'added_rti_no', 'page_row', 'invoice_page', 'doc_supplier',
                        'folder_address', 'stamp', 'signature', 'modified_by', 'updated_at'
                    ])
                broadcast_doc_task_update()
            
        return Response({'message': f'{len(histories)} رکورد نهایی شد.'})
        
    @action(detail=False, methods=['post'])
    def bulk_cancel(self, request):
        task_ids = request.data.get('task_ids', [])
        if not task_ids:
            return Response({'error': 'هیچ رکوردی انتخاب نشده است.'}, status=400)
            
        tasks = DocTask.objects.filter(id__in=task_ids)
        item_ids = list(tasks.values_list('item_id', flat=True))
        
        from ..models import Item
        deleted_count, _ = tasks.delete()
        Item.objects.filter(id__in=item_ids).update(doc_status='checking', doc_assignee=None, updated_at=timezone.now())
        
        broadcast_doc_task_update()
        return Response({'message': f'{deleted_count} وظیفه ارجاع اسناد با موفقیت لغو شد.'})

    @action(detail=False, methods=['get'], url_path='download_template')
    def download_template(self, request):
        """تولید فایل اکسل نمونه آزمایشی برای کارتابل مالی متناسب با فیلدهای قابل ویرایش"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        from django.http import HttpResponse
        from warehouses.services import get_setting

        warehouse_id = request.query_params.get('warehouse_id')
        if warehouse_id and str(warehouse_id) in ['ALL', '-1', 'null', 'undefined']:
            warehouse_id = None
        num_wh_id = int(warehouse_id) if warehouse_id and str(warehouse_id).isdigit() else None

        # دریافت تنظیمات فیلدهای کارتابل مالی
        saved_perms = get_setting('field_permissions_doc', num_wh_id) or {}
        if isinstance(saved_perms, str):
            import json
            try:
                saved_perms = json.loads(saved_perms)
            except Exception:
                saved_perms = {}

        # لیست فیلدهای استاندارد مالی
        STANDARD_DOC_FIELDS = [
            {'key': 'price_amount',       'default_label': 'قیمت واحد',        'type': 'number',   'sample1': '2500000',               'sample2': '1850000'},
            {'key': 'similar_unit_price', 'default_label': 'قیمت کالای مشابه', 'type': 'number',   'sample1': '2400000',               'sample2': '1800000'},
            {'key': 'total_value',        'default_label': 'ارزش کل',          'type': 'number',   'sample1': '25000000',              'sample2': '18500000'},
            {'key': 'currency',           'default_label': 'ارز',              'type': 'currency', 'sample1': 'ریال',                  'sample2': 'دلار'},
            {'key': 'invoice_type',       'default_label': 'نوع فاکتور',        'type': 'inv_type', 'sample1': 'رسمی/مالیاتی',          'sample2': 'خریدهای داخلی'},
            {'key': 'invoice_date',       'default_label': 'تاریخ فاکتور',     'type': 'date',     'sample1': '1405/05/25 07:10',      'sample2': '1405/06/01 10:30'},
            {'key': 'inv_rti_number',     'default_label': 'شماره RTI فاکتور', 'type': 'text',     'sample1': 'RTI-1405-001',          'sample2': 'RTI-1405-002'},
            {'key': 'added_rti_no',       'default_label': 'شماره RTI افزوده‌شده', 'type': 'text', 'sample1': 'RTI-ADD-99',           'sample2': 'RTI-ADD-100'},
            {'key': 'invoice_page',       'default_label': 'صفحه فاکتور',      'type': 'number',   'sample1': '1',                     'sample2': '2'},
            {'key': 'page_row',           'default_label': 'ردیف فاکتور',      'type': 'number',   'sample1': '1',                     'sample2': '3'},
            {'key': 'doc_supplier',       'default_label': 'تأمین‌کننده',      'type': 'text',     'sample1': 'شرکت تأمین تجهیز پارس', 'sample2': 'شرکت مهندسی پویا'},
            {'key': 'folder_address',     'default_label': 'مسیر پوشه اسناد',  'type': 'text',     'sample1': 'Z:/Docs/Archive/1405',  'sample2': 'Z:/Docs/Archive/1405'},
            {'key': 'stamp',              'default_label': 'مهر',              'type': 'boolean',  'sample1': 'بله',                   'sample2': 'خیر'},
            {'key': 'signature',          'default_label': 'امضا',             'type': 'boolean',  'sample1': 'بله',                   'sample2': 'بله'},
            {'key': 'worker_note',        'default_label': 'یادداشت کارشناس',  'type': 'text',     'sample1': 'مدارک مالی مطابقت دارد', 'sample2': 'فاکتور ضمیمه شد'},
        ]

        # همیشه ستون شناساگر کد یکتا در ابتدا قرار می‌گیرد
        columns = [
            {'key': 'fa_unic_code', 'label': 'کد یکتا', 'type': 'text', 'sample1': 'FA-10001', 'sample2': 'FA-10002'}
        ]

        for f in STANDARD_DOC_FIELDS:
            cfg = saved_perms.get(f['key'], {}) if isinstance(saved_perms, dict) else {}
            is_editable = cfg.get('editable', True)
            is_visible = cfg.get('visible', True)
            if is_visible and is_editable:
                label = (cfg.get('custom_label') or '').strip() or f['default_label']
                columns.append({
                    'key': f['key'],
                    'label': label,
                    'type': f['type'],
                    'sample1': f['sample1'],
                    'sample2': f['sample2'],
                })

        # افزودن فیلدهای داینامیک فعال و قابل ویرایش انبار
        if num_wh_id:
            from ..models import ItemFieldDefinition
            dyn_defs = ItemFieldDefinition.objects.filter(warehouse_id=num_wh_id, is_active=True)
            for d in dyn_defs:
                dyn_key = f"dyn_{d.name}"
                cfg = saved_perms.get(dyn_key, {}) if isinstance(saved_perms, dict) else {}
                is_editable = cfg.get('editable', True)
                is_visible = cfg.get('visible', True)
                if is_visible and is_editable:
                    label = (cfg.get('custom_label') or '').strip() or d.label or d.name
                    dtype = d.field_type
                    if dtype == 'number':
                        s1, s2 = '100', '250'
                    elif dtype == 'boolean':
                        s1, s2 = 'بله', 'خیر'
                    elif dtype == 'date':
                        s1, s2 = '1405/05/25', '1405/06/01'
                    else:
                        s1, s2 = 'مقدار نمونه ۱', 'مقدار نمونه ۲'
                    columns.append({
                        'key': dyn_key,
                        'label': label,
                        'type': 'boolean' if dtype == 'boolean' else ('number' if dtype == 'number' else ('date' if dtype == 'date' else 'text')),
                        'sample1': s1,
                        'sample2': s2,
                    })

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'قالب اسناد مالی'
        ws.sheet_view.rightToLeft = True

        header_font = Font(name='Tahoma', bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        sample_font = Font(name='Tahoma', size=10)
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'), right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'), bottom=Side(style='thin', color='E2E8F0')
        )
        fill_row1 = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        fill_row2 = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

        headers = [c['label'] for c in columns]
        ws.append(headers)

        for col_idx, cell in enumerate(ws[1], 1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = thin_border

        row1_data = [c['sample1'] for c in columns]
        row2_data = [c['sample2'] for c in columns]
        ws.append(row1_data)
        ws.append(row2_data)

        for r_idx, (r_data, r_fill) in enumerate([(row1_data, fill_row1), (row2_data, fill_row2)], start=2):
            for c_idx, val in enumerate(r_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = sample_font
                cell.fill = r_fill
                cell.alignment = center_alignment
                cell.border = thin_border

        # اعمال کشوها با DataValidation
        bool_dv = DataValidation(type="list", formula1='"بله,خیر"', allow_blank=True)
        inv_dv = DataValidation(type="list", formula1='"رسمی/مالیاتی,خریدهای داخلی,خریدهای خارجی,امانی"', allow_blank=True)
        cur_dv = DataValidation(type="list", formula1='"ریال,دلار,یورو,سایر"', allow_blank=True)

        ws.add_data_validation(bool_dv)
        ws.add_data_validation(inv_dv)
        ws.add_data_validation(cur_dv)

        for col_idx, col in enumerate(columns, 1):
            col_letter = get_column_letter(col_idx)
            if col['type'] == 'boolean':
                bool_dv.add(f"{col_letter}2:{col_letter}500")
            elif col['type'] == 'inv_type':
                inv_dv.add(f"{col_letter}2:{col_letter}500")
            elif col['type'] == 'currency':
                cur_dv.add(f"{col_letter}2:{col_letter}500")

            max_len = max(len(str(col['label'])), len(str(col.get('sample1', ''))), len(str(col.get('sample2', ''))))
            ws.column_dimensions[col_letter].width = max(max_len + 6, 14)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Customs_Doc_Template.xlsx"'
        wb.save(response)
        return response

    @action(detail=False, methods=['get'], url_path='get_export_columns')
    def get_export_columns(self, request):
        """لیست ستون‌های مجاز برای خروجی اکسل کارتابل مالی و اسناد"""
        columns = [
            {'key': 'warehouse_name',        'label': 'نام انبار'},
            {'key': 'fa_unic_code',          'label': 'کد یکتا'},
            {'key': 'description',           'label': 'شرح کالا'},
            {'key': 'po',                    'label': 'شماره PO'},
            {'key': 'new_location',          'label': 'لوکیشن'},
            {'key': 'status',                'label': 'وضعیت'},
            {'key': 'doc_worker_name',       'label': 'کارشناس مالی'},
            {'key': 'doc_supervisor_name',   'label': 'سرپرست'},
            {'key': 'assigned_manager_name', 'label': 'مدیر'},
            {'key': 'inv_rti_number',        'label': 'شماره RTI فاکتور'},
            {'key': 'added_rti_no',          'label': 'شماره RTI افزوده‌شده'},
            {'key': 'invoice_type',          'label': 'نوع فاکتور'},
            {'key': 'invoice_date',          'label': 'تاریخ فاکتور'},
            {'key': 'invoice_page',          'label': 'صفحه فاکتور'},
            {'key': 'page_row',              'label': 'ردیف فاکتور'},
            {'key': 'doc_supplier',          'label': 'تأمین‌کننده'},
            {'key': 'price_amount',          'label': 'قیمت واحد'},
            {'key': 'similar_unit_price',    'label': 'قیمت کالای مشابه'},
            {'key': 'total_value',           'label': 'ارزش کل'},
            {'key': 'currency',              'label': 'ارز'},
            {'key': 'folder_address',        'label': 'مسیر پوشه اسناد'},
            {'key': 'stamp',                 'label': 'مهر'},
            {'key': 'signature',             'label': 'امضا'},
            {'key': 'worker_note',           'label': 'یادداشت کارشناس'},
            {'key': 'supervisor_note',       'label': 'یادداشت سرپرست'},
            {'key': 'manager_note',          'label': 'یادداشت مدیر'},
            {'key': 'created_at',            'label': 'تاریخ ایجاد'},
            {'key': 'updated_at',            'label': 'تاریخ بروزرسانی'},
        ]
        return Response(columns)

    @action(detail=False, methods=['post'], url_path='export_excel')
    def export_excel(self, request):
        """خروجی اکسل صفحه کارتابل مالی با تاریخ شمسی و وضعیت فارسی"""
        import openpyxl
        from django.http import HttpResponse

        STATUS_FA = {
            'PENDING_DOC':             'در انتظار بررسی',
            'DOC_PROCESSED':           'ارسال‌شده به سرپرست',
            'DOC_SUPERVISOR_REJECTED': 'رد سرپرست',
            'DOC_MANAGER_REVIEW':      'در بررسی مدیر',
            'DOC_MANAGER_REJECTED':    'رد مدیر',
            'DOC_FINAL_APPROVED':      'تأیید نهایی',
        }

        INVOICE_TYPE_FA = {
            'formal':      'رسمی/مالیاتی',
            'domestic':    'خریدهای داخلی',
            'foreign':     'خریدهای خارجی',
            'consignment': 'امانی',
        }

        CURRENCY_FA = {
            'IRR':   'ریال',
            'USD':   'دلار',
            'EUR':   'یورو',
            'OTHER': 'سایر',
        }

        def to_jalali(dt):
            if not dt:
                return ''
            try:
                import jdatetime
                from django.utils import timezone
                if timezone.is_aware(dt):
                    dt = timezone.localtime(dt)
                jdt = jdatetime.datetime.fromgregorian(datetime=dt)
                return jdt.strftime('%Y/%m/%d %H:%M')
            except Exception:
                return str(dt)

        data_scope    = request.data.get('data_scope', 'all')
        columns_scope = request.data.get('columns_scope', 'all_db')
        columns_list  = request.data.get('columns_list', [])

        from django.http import QueryDict
        original_query_params = request._request.GET
        try:
            q = QueryDict(mutable=True)
            for k, v in request.data.items():
                if isinstance(v, list):
                    q.setlist(k, [str(x) for x in v])
                elif v is not None:
                    q[k] = str(v)
            request._request.GET = q

            if data_scope == 'selected':
                selected_ids = request.data.get('selected_ids', [])
                queryset = self.get_queryset().filter(id__in=selected_ids)
            else:
                queryset = self.get_queryset()
        finally:
            request._request.GET = original_query_params

        ALL_COLUMNS = [
            {'key': 'warehouse_name',        'label': 'نام انبار'},
            {'key': 'fa_unic_code',          'label': 'کد یکتا'},
            {'key': 'description',           'label': 'شرح کالا'},
            {'key': 'po',                    'label': 'شماره PO'},
            {'key': 'new_location',          'label': 'لوکیشن'},
            {'key': 'status',                'label': 'وضعیت'},
            {'key': 'doc_worker_name',       'label': 'کارشناس مالی'},
            {'key': 'doc_supervisor_name',   'label': 'سرپرست'},
            {'key': 'assigned_manager_name', 'label': 'مدیر'},
            {'key': 'inv_rti_number',        'label': 'شماره RTI فاکتور'},
            {'key': 'added_rti_no',          'label': 'شماره RTI افزوده‌شده'},
            {'key': 'invoice_type',          'label': 'نوع فاکتور'},
            {'key': 'invoice_date',          'label': 'تاریخ فاکتور'},
            {'key': 'invoice_page',          'label': 'صفحه فاکتور'},
            {'key': 'page_row',              'label': 'ردیف فاکتور'},
            {'key': 'doc_supplier',          'label': 'تأمین‌کننده'},
            {'key': 'price_amount',          'label': 'قیمت واحد'},
            {'key': 'similar_unit_price',    'label': 'قیمت کالای مشابه'},
            {'key': 'total_value',           'label': 'ارزش کل'},
            {'key': 'currency',              'label': 'ارز'},
            {'key': 'folder_address',        'label': 'مسیر پوشه اسناد'},
            {'key': 'stamp',                 'label': 'مهر'},
            {'key': 'signature',             'label': 'امضا'},
            {'key': 'worker_note',           'label': 'یادداشت کارشناس'},
            {'key': 'supervisor_note',       'label': 'یادداشت سرپرست'},
            {'key': 'manager_note',          'label': 'یادداشت مدیر'},
            {'key': 'created_at',            'label': 'تاریخ ایجاد'},
            {'key': 'updated_at',            'label': 'تاریخ بروزرسانی'},
        ]
        all_keys = [c['key'] for c in ALL_COLUMNS]
        key_to_label = {c['key']: c['label'] for c in ALL_COLUMNS}

        if columns_scope in ('visible', 'custom') and columns_list:
            selected_keys = [k for k in columns_list if k in all_keys]
            if not selected_keys:
                selected_keys = all_keys
        else:
            selected_keys = all_keys

        def get_cell(task, key):
            if key == 'warehouse_name':
                if task.item and task.item.warehouse:
                    return getattr(task.item.warehouse, 'project_name', None) or task.item.warehouse.name or ''
                return ''
            elif key == 'fa_unic_code':
                return getattr(task.item, 'fa_unic_code', '') if task.item else ''
            elif key == 'description':
                return getattr(task.item, 'description', '') if task.item else ''
            elif key == 'po':
                return getattr(task.item, 'po', '') if task.item else ''
            elif key == 'new_location':
                return getattr(task.item, 'new_location', '') if task.item else ''
            elif key == 'status':
                return STATUS_FA.get(task.status, task.status)
            elif key == 'doc_worker_name':
                if task.doc_worker:
                    return f"{task.doc_worker.first_name} {task.doc_worker.last_name}".strip() or task.doc_worker.username
                return ''
            elif key == 'doc_supervisor_name':
                if task.doc_supervisor:
                    return f"{task.doc_supervisor.first_name} {task.doc_supervisor.last_name}".strip() or task.doc_supervisor.username
                return ''
            elif key == 'assigned_manager_name':
                if task.assigned_manager:
                    return f"{task.assigned_manager.first_name} {task.assigned_manager.last_name}".strip() or task.assigned_manager.username
                return ''
            elif key == 'inv_rti_number':
                return task.inv_rti_number or ''
            elif key == 'added_rti_no':
                return task.added_rti_no or ''
            elif key == 'invoice_type':
                return INVOICE_TYPE_FA.get(task.invoice_type, task.invoice_type or '')
            elif key == 'invoice_date':
                return task.invoice_date or ''
            elif key == 'invoice_page':
                return task.invoice_page if task.invoice_page is not None else ''
            elif key == 'page_row':
                return task.page_row if task.page_row is not None else ''
            elif key == 'doc_supplier':
                return task.doc_supplier or ''
            elif key == 'price_amount':
                return task.price_amount or ''
            elif key == 'similar_unit_price':
                return task.similar_unit_price or ''
            elif key == 'total_value':
                return task.total_value or ''
            elif key == 'currency':
                return CURRENCY_FA.get(task.currency, task.currency or '')
            elif key == 'folder_address':
                return task.folder_address or ''
            elif key == 'stamp':
                return 'بله' if task.stamp else 'خیر'
            elif key == 'signature':
                return 'بله' if task.signature else 'خیر'
            elif key == 'worker_note':
                return task.worker_note or ''
            elif key == 'supervisor_note':
                return task.supervisor_note or ''
            elif key == 'manager_note':
                return task.manager_note or ''
            elif key == 'created_at':
                return to_jalali(task.created_at)
            elif key == 'updated_at':
                return to_jalali(task.updated_at)
            return ''

        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'کارتابل مالی'
        ws.sheet_view.rightToLeft = True

        header_font = Font(name='Tahoma', bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # سرمه‌ای دودی مدرن
        cell_font = Font(name='Tahoma', size=10)

        fill_even = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'), right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'), bottom=Side(style='thin', color='E2E8F0')
        )

        headers = [key_to_label[k] for k in selected_keys]
        ws.append(headers)

        for col_idx, cell in enumerate(ws[1], 1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = thin_border

        ws.freeze_panes = 'A2'
        col_widths = {i: len(headers[i]) + 6 for i in range(len(headers))}

        row_idx = 2
        for task in queryset.select_related('item', 'item__warehouse', 'doc_worker', 'doc_supervisor', 'assigned_manager').iterator():
            row_data = [str(get_cell(task, k)) for k in selected_keys]
            ws.append(row_data)

            current_fill = fill_even if row_idx % 2 == 0 else fill_odd

            for col_idx, val in enumerate(row_data):
                cell = ws.cell(row=row_idx, column=col_idx + 1)
                cell.font = cell_font
                cell.fill = current_fill
                cell.alignment = center_alignment
                cell.border = thin_border

                lines = val.split('\n')
                max_line_len = max([len(line) for line in lines]) if lines else 0
                if max_line_len > col_widths.get(col_idx, 10):
                    col_widths[col_idx] = min(max_line_len + 4, 60)

            row_idx += 1

        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx + 1)].width = width

        if row_idx > 2:
            last_col_letter = get_column_letter(len(selected_keys))
            ws.auto_filter.ref = f"A1:{last_col_letter}{row_idx - 1}"

        from io import BytesIO
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="customs_cartable_export.xlsx"'
        return response

__all__ = [
    'DocTaskViewSet',
]
