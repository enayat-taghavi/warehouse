"""
Characterization tests for AuditLogViewSet domain app-scope filters.
Covers list, stats, and export_excel actions on the current codebase.

Session 6 of 15 - Phase 4 backend modularization.
"""
import io
import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from rest_framework.test import APIClient

from accounts.models import AuditLog
from warehouses.models import Warehouse

User = get_user_model()


class AuditAppScopeCharacterizationTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Clean slate for AuditLog fixtures
        AuditLog.objects.all().delete()

        # 1. Warehouses: wh1 is assigned to non-superuser; wh2 is not assigned
        self.wh1 = Warehouse.objects.create(name='انبار مرکزی تست', code='WH-CHAR-1')
        self.wh2 = Warehouse.objects.create(name='انبار فرعی تست', code='WH-CHAR-2')

        # 2. Superuser
        self.superuser = User.objects.create_superuser(
            username='audit_char_admin',
            password='Password123!',
            email='admin_char@test.com'
        )

        # 3. Non-superuser: has view_wh_audit permission only, assigned to wh1 only.
        # Has no accounts.perm_sys_logs permission.
        self.non_superuser = User.objects.create_user(
            username='audit_char_user',
            password='Password123!',
            email='user_char@test.com'
        )
        perm = Permission.objects.get(codename='view_wh_audit', content_type__app_label='accounts')
        self.non_superuser.user_permissions.add(perm)
        if hasattr(self.non_superuser, '_perm_cache'):
            del self.non_superuser._perm_cache
        if hasattr(self.non_superuser, '_user_perm_cache'):
            del self.non_superuser._user_perm_cache
        self.non_superuser.assigned_warehouses.add(self.wh1)

        # 4. Fixed set of AuditLog rows covering all 8 required cases:
        # Row 1: warehouse-module key with a warehouse id
        self.wh_log_wh1 = AuditLog.objects.create(
            module='docs',
            warehouse_id=self.wh1.id,
            action='CREATE',
            details={}
        )

        # Row 2: accounting-module key with no warehouse id
        self.acct_log_nowh = AuditLog.objects.create(
            module='payroll',
            warehouse_id=None,
            action='UPDATE',
            details={}
        )

        # Row 3: accounting key with a warehouse id
        self.acct_log_wh1 = AuditLog.objects.create(
            module='finance',
            warehouse_id=self.wh1.id,
            action='UPDATE',
            details={}
        )

        # Row 4: target_model='PersonnelProfile' and a non-accounting module
        self.target_model_log = AuditLog.objects.create(
            module='system',
            target_model='PersonnelProfile',
            warehouse_id=None,
            action='UPDATE',
            details={}
        )

        # Row 5: details={'target_module': 'finance'}
        self.details_finance_log = AuditLog.objects.create(
            module='system',
            warehouse_id=None,
            action='UPDATE',
            details={'target_module': 'finance'}
        )

        # Row 6: CROSS_APP_DENIED row
        self.cross_app_denied_log = AuditLog.objects.create(
            module='security',
            warehouse_id=self.wh1.id,
            action='UPDATE',
            details={'event': 'CROSS_APP_DENIED'}
        )

        # Row 7: APP_SCOPE_SWITCH row
        self.app_scope_switch_log = AuditLog.objects.create(
            module='security',
            warehouse_id=self.wh1.id,
            action='UPDATE',
            details={'event': 'APP_SCOPE_SWITCH'}
        )

        # Row 8: warehouse the user is not assigned to (wh2)
        self.unassigned_wh_log = AuditLog.objects.create(
            module='docs',
            warehouse_id=self.wh2.id,
            action='CREATE',
            details={}
        )

        self.all_log_ids = {
            self.wh_log_wh1.id,
            self.acct_log_nowh.id,
            self.acct_log_wh1.id,
            self.target_model_log.id,
            self.details_finance_log.id,
            self.cross_app_denied_log.id,
            self.app_scope_switch_log.id,
            self.unassigned_wh_log.id,
        }

    def _extract_ids(self, response):
        data = response.data
        if isinstance(data, dict) and 'results' in data:
            return {item['id'] for item in data['results']}
        elif isinstance(data, list):
            return {item['id'] for item in data}
        raise ValueError(f"Unexpected response format: {data}")

    def test_superuser_list_warehouse_scope(self):
        """
        Superuser list with app_scope=warehouse.
        Current behaviour: returns empty set because exclude(details__event=...)
        on PostgreSQL evaluates to SQL NULL for rows lacking 'event' in details.
        """
        self.client.force_authenticate(self.superuser)
        res = self.client.get('/api/accounts/audit-logs/?app_scope=warehouse')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._extract_ids(res), set())

    def test_superuser_list_finance_scope(self):
        """
        Superuser list with app_scope=finance.
        Current behaviour: returns empty set because exclude(details__event=...)
        on PostgreSQL evaluates to SQL NULL for rows lacking 'event' in details.
        """
        self.client.force_authenticate(self.superuser)
        res = self.client.get('/api/accounts/audit-logs/?app_scope=finance')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._extract_ids(res), set())

    def test_superuser_list_personnel_scope(self):
        """
        Superuser list with app_scope=personnel.
        Current behaviour: returns empty set because exclude(details__event=...)
        on PostgreSQL evaluates to SQL NULL for rows lacking 'event' in details.
        """
        self.client.force_authenticate(self.superuser)
        res = self.client.get('/api/accounts/audit-logs/?app_scope=personnel')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._extract_ids(res), set())

    def test_superuser_list_all_scope(self):
        """Superuser list with app_scope=all: returns all 8 rows."""
        self.client.force_authenticate(self.superuser)
        res = self.client.get('/api/accounts/audit-logs/?app_scope=all')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._extract_ids(res), self.all_log_ids)

    def test_superuser_list_absent_scope(self):
        """Superuser list with app_scope absent: returns all 8 rows."""
        self.client.force_authenticate(self.superuser)
        res = self.client.get('/api/accounts/audit-logs/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._extract_ids(res), self.all_log_ids)

    def test_nonsuperuser_list_warehouse_scope(self):
        """
        Non-superuser (assigned to wh1 only) list with app_scope=warehouse.
        Current behaviour: returns empty set due to SQL NULL on exclude(details__event=...).
        """
        self.client.force_authenticate(self.non_superuser)
        res = self.client.get('/api/accounts/audit-logs/?app_scope=warehouse')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._extract_ids(res), set())

    def test_nonsuperuser_list_finance_scope(self):
        """
        Non-superuser (assigned to wh1 only) list with app_scope=finance.
        Current behaviour: returns empty set due to SQL NULL on exclude(details__event=...).
        """
        self.client.force_authenticate(self.non_superuser)
        res = self.client.get('/api/accounts/audit-logs/?app_scope=finance')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._extract_ids(res), set())

    def test_nonsuperuser_list_absent_scope(self):
        """
        Non-superuser (assigned to wh1 only) list with app_scope absent.
        Current behaviour: returns rows in assigned warehouse (wh1) plus accounting
        rows with null warehouse. Note: target_model_log and details_finance_log
        are excluded because warehouse_id is null and module is 'system' (not in
        audit_scope_keys('accounting')), while CROSS_APP_DENIED and APP_SCOPE_SWITCH
        are included because warehouse restriction keeps wh1 and no downstream exclude runs.
        """
        self.client.force_authenticate(self.non_superuser)
        res = self.client.get('/api/accounts/audit-logs/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            self._extract_ids(res),
            {
                self.wh_log_wh1.id,
                self.acct_log_nowh.id,
                self.acct_log_wh1.id,
                self.cross_app_denied_log.id,
                self.app_scope_switch_log.id,
            }
        )

    def test_superuser_stats_counts(self):
        """
        Stats action for superuser: warehouse_logs_count and finance_logs_count.
        Current behaviour: returns 0 and 0 due to SQL NULL on exclude(details__event=...).
        """
        self.client.force_authenticate(self.superuser)
        res = self.client.get('/api/accounts/audit-logs/stats/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data.get('warehouse_logs_count'), 0)
        self.assertEqual(res.data.get('finance_logs_count'), 0)

    def test_superuser_export_excel_warehouse_scope(self):
        """
        Excel export for superuser with app_scope=warehouse in request data.
        Current behaviour: 0 data rows due to SQL NULL on exclude(details__event=...).
        """
        self.client.force_authenticate(self.superuser)
        res = self.client.post('/api/accounts/audit-logs/export_excel/', {'app_scope': 'warehouse'}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertIn('spreadsheetml', res['Content-Type'])
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        data_rows = [row for row in ws.iter_rows(min_row=3, values_only=True) if any(cell is not None for cell in row)]
        self.assertEqual(len(data_rows), 0)

    def test_superuser_export_excel_finance_scope(self):
        """
        Excel export for superuser with app_scope=finance in request data.
        Current behaviour: 0 data rows due to SQL NULL on exclude(details__event=...).
        """
        self.client.force_authenticate(self.superuser)
        res = self.client.post('/api/accounts/audit-logs/export_excel/', {'app_scope': 'finance'}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertIn('spreadsheetml', res['Content-Type'])
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        data_rows = [row for row in ws.iter_rows(min_row=3, values_only=True) if any(cell is not None for cell in row)]
        self.assertEqual(len(data_rows), 0)

    def test_superuser_export_excel_query_param(self):
        """
        Excel export with app_scope=warehouse via GET query parameter.
        Current behaviour: 0 data rows due to SQL NULL on exclude(details__event=...).
        """
        self.client.force_authenticate(self.superuser)
        res = self.client.get('/api/accounts/audit-logs/export_excel/?app_scope=warehouse')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        data_rows = [row for row in ws.iter_rows(min_row=3, values_only=True) if any(cell is not None for cell in row)]
        self.assertEqual(len(data_rows), 0)
