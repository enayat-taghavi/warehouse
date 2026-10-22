"""
آزمون‌های یکپارچگی تفکیک چندشرکتی در موتور گزارش‌ساز و ورود کاربران از Excel.
طرح پیگیری‌های ناظر — نشست ۱.
"""
import io
from companies.testing import grant_membership
from openpyxl import Workbook, load_workbook

from django.apps import apps
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from accounts.excel_utils import USERS_COLUMNS
from companies.models import Company, UserCompanyAccess
from platform_core.query_engine import get_all_entities, ReportEngine
from reports.registry import register_warehouse_reports

User = get_user_model()


def _make_users_excel(rows_data):
    """تولید فایل اکسل نمونه با ساختار استاندارد ۲ سطری هدر برای آزمون‌های ورود کاربران."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'قالب کاربران'
    # سطر ۱: عناوین فارسی
    ws.append([c['label'] for c in USERS_COLUMNS])
    # سطر ۲: کلیدهای دیتابیسی
    ws.append([c['key'] for c in USERS_COLUMNS])
    # داده‌ها
    for r in rows_data:
        row_cells = []
        for c in USERS_COLUMNS:
            row_cells.append(r.get(c['key'], ''))
        ws.append(row_cells)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


class CompanyScopeReportsAndImportIntegrationTests(TestCase):
    def setUp(self):
        register_warehouse_reports()
        self.client = APIClient()

        self.company_a = Company.objects.create(code='COMP-A', name='شرکت الف', is_active=True)
        self.company_b = Company.objects.create(code='COMP-B', name='شرکت ب', is_active=True)

        self.superuser = User.objects.create_superuser(
            username='scope_super',
            password='Password123!',
            national_code='0011111111',
            phone_number='09121111111'
        )
        self.user_a = User.objects.create_user(
            username='user_comp_a',
            password='Password123!',
            national_code='0022222222',
            phone_number='09122222222',
            first_name='کاربر',
            last_name='شرکت الف'
        )
        self.user_b = User.objects.create_user(
            username='user_comp_b',
            password='Password123!',
            national_code='0033333333',
            phone_number='09123333333',
            first_name='کاربر',
            last_name='شرکت ب'
        )
        self.user_no_company = User.objects.create_user(
            username='user_no_comp',
            password='Password123!',
            national_code='0044444444',
            phone_number='09124444444',
            first_name='کاربر',
            last_name='بدون شرکت'
        )

        UserCompanyAccess.objects.create(
            user=self.user_a,
            company=self.company_a,
            is_default=True
        )
        UserCompanyAccess.objects.create(
            user=self.user_b,
            company=self.company_b,
            is_default=True
        )
        grant_membership(self.user_a, self.company_a)
        grant_membership(self.user_b, self.company_b)

        perms = Permission.objects.filter(
            codename__in=[
                'view_sys_reports',
                'view_sys_users',
                'view_sys_personnel',
                'view_sys_payroll',
                'view_sys_fleet_settlement',
                'perm_usr_add',
            ]
        )
        for u in (self.user_a, self.user_b, self.user_no_company):
            u.user_permissions.set(perms)

        if apps.is_installed('warehouses'):
            from warehouses.models import Warehouse
            self.wh_a = Warehouse.objects.create(name='انبار شرکت الف', code='WH-A-1', company=self.company_a)
            self.wh_b = Warehouse.objects.create(name='انبار شرکت ب', code='WH-B-1', company=self.company_b)

        if apps.is_installed('personnel'):
            from personnel.models import (
                FinancialProject, ProjectSection, PersonnelProfile,
                MonthlyPayrollRecord, MonthlyWorkPeriod,
                VehicleDriverProfile, VehicleTripLog
            )
            self.proj_a = FinancialProject.objects.create(company=self.company_a, code='PA', name='پروژه الف')
            self.sec_a = ProjectSection.objects.create(project=self.proj_a, code='SA', name='بخش الف')
            self.proj_b = FinancialProject.objects.create(company=self.company_b, code='PB', name='پروژه ب')
            self.sec_b = ProjectSection.objects.create(project=self.proj_b, code='SB', name='بخش ب')

            self.period_a = MonthlyWorkPeriod.objects.create(
                company=self.company_a,
                project=self.proj_a,
                section=self.sec_a,
                year_month='1404/05',
                status='OPEN',
            )
            self.period_b = MonthlyWorkPeriod.objects.create(
                company=self.company_b,
                project=self.proj_b,
                section=self.sec_b,
                year_month='1404/05',
                status='OPEN',
            )

            self.profile_a = PersonnelProfile.objects.create(
                company=self.company_a,
                section=self.sec_a,
                first_name='پرسنل الف',
                last_name='احمدی',
                national_code='0055555555',
                personnel_id_code='EMP-A',
            )
            self.profile_b = PersonnelProfile.objects.create(
                company=self.company_b,
                section=self.sec_b,
                first_name='پرسنل ب',
                last_name='رضایی',
                national_code='0066666666',
                personnel_id_code='EMP-B',
            )

            self.payroll_a = MonthlyPayrollRecord.objects.create(
                period=self.period_a,
                personnel=self.profile_a,
                full_name=f"{self.profile_a.first_name} {self.profile_a.last_name}",
                national_code=self.profile_a.national_code,
                gross_salary=15000000,
                net_salary=13000000,
                payable_amount=13000000,
                payment_status='PENDING',
            )
            self.payroll_b = MonthlyPayrollRecord.objects.create(
                period=self.period_b,
                personnel=self.profile_b,
                full_name=f"{self.profile_b.first_name} {self.profile_b.last_name}",
                national_code=self.profile_b.national_code,
                gross_salary=18000000,
                net_salary=16000000,
                payable_amount=16000000,
                payment_status='PENDING',
            )

            self.veh_a = VehicleDriverProfile.objects.create(
                company=self.company_a,
                driver_name='راننده الف',
                plate_number='11-111-11',
            )
            self.veh_b = VehicleDriverProfile.objects.create(
                company=self.company_b,
                driver_name='راننده ب',
                plate_number='22-222-22',
            )

            self.trip_a = VehicleTripLog.objects.create(
                vehicle=self.veh_a,
                section=self.sec_a,
                project=self.proj_a,
                date_shamsi='1405/01/10',
                trip_count=1,
                total_amount=1000000,
            )
            self.trip_b = VehicleTripLog.objects.create(
                vehicle=self.veh_b,
                section=self.sec_b,
                project=self.proj_b,
                date_shamsi='1405/01/11',
                trip_count=2,
                total_amount=2000000,
            )

            # کاربر غیرممیز پرسنل فقط بخش‌های منتسب فعال خود را می‌بیند (همان قاعدهٔ صفحهٔ پرسنل)
            from personnel.models import UserSectionAssignment
            UserSectionAssignment.objects.create(user=self.user_a, section=self.sec_a, is_active=True)
            UserSectionAssignment.objects.create(user=self.user_b, section=self.sec_b, is_active=True)

    # -------------------------------------------------------------------------
    # ۱. run برای personnel_profiles: کاربر شرکت الف فقط الف، سوپریوزر هر دو، کاربر بی‌شرکت هیچ
    # -------------------------------------------------------------------------
    def test_run_report_personnel_profiles_scope(self):
        if not apps.is_installed('personnel'):
            self.skipTest('ماژول پرسنل و حسابداری نصب نیست.')

        spec = {'entity': 'personnel_profiles', 'fields': ['id', 'first_name']}

        # کاربر شرکت الف
        self.client.force_authenticate(self.user_a)
        res_a = self.client.post('/api/reports/run/', spec, format='json').json()
        ids_a = [r['id'] for r in res_a['rows']]
        self.assertIn(self.profile_a.id, ids_a)
        self.assertNotIn(self.profile_b.id, ids_a)
        self.assertEqual(res_a['count'], 1)

        # سوپریوزر
        self.client.force_authenticate(self.superuser)
        res_s = self.client.post('/api/reports/run/', spec, format='json').json()
        ids_s = [r['id'] for r in res_s['rows']]
        self.assertIn(self.profile_a.id, ids_s)
        self.assertIn(self.profile_b.id, ids_s)
        self.assertGreaterEqual(res_s['count'], 2)

        # کاربر بی‌شرکت
        self.client.force_authenticate(self.user_no_company)
        res_none = self.client.post('/api/reports/run/', spec, format='json').json()
        self.assertEqual(res_none['count'], 0)
        self.assertEqual(len(res_none['rows']), 0)

    # -------------------------------------------------------------------------
    # ۲. run برای monthly_payrolls: کاربر شرکت الف فقط الف، سوپریوزر هر دو، کاربر بی‌شرکت هیچ
    # -------------------------------------------------------------------------
    def test_run_report_monthly_payrolls_scope(self):
        if not apps.is_installed('personnel'):
            self.skipTest('ماژول پرسنل و حسابداری نصب نیست.')

        spec = {'entity': 'monthly_payrolls', 'fields': ['id', 'full_name']}

        # کاربر شرکت الف
        self.client.force_authenticate(self.user_a)
        res_a = self.client.post('/api/reports/run/', spec, format='json').json()
        ids_a = [r['id'] for r in res_a['rows']]
        self.assertIn(self.payroll_a.id, ids_a)
        self.assertNotIn(self.payroll_b.id, ids_a)
        self.assertEqual(res_a['count'], 1)

        # سوپریوزر
        self.client.force_authenticate(self.superuser)
        res_s = self.client.post('/api/reports/run/', spec, format='json').json()
        ids_s = [r['id'] for r in res_s['rows']]
        self.assertIn(self.payroll_a.id, ids_s)
        self.assertIn(self.payroll_b.id, ids_s)
        self.assertGreaterEqual(res_s['count'], 2)

        # کاربر بی‌شرکت
        self.client.force_authenticate(self.user_no_company)
        res_none = self.client.post('/api/reports/run/', spec, format='json').json()
        self.assertEqual(res_none['count'], 0)
        self.assertEqual(len(res_none['rows']), 0)

    # -------------------------------------------------------------------------
    # ۳. run برای fleet_trips: کاربر شرکت الف فقط الف، سوپریوزر هر دو، کاربر بی‌شرکت هیچ
    # -------------------------------------------------------------------------
    def test_run_report_fleet_trips_scope(self):
        if not apps.is_installed('personnel'):
            self.skipTest('ماژول پرسنل و حسابداری نصب نیست.')

        spec = {'entity': 'fleet_trips', 'fields': ['id', 'total_amount']}

        # کاربر شرکت الف
        self.client.force_authenticate(self.user_a)
        res_a = self.client.post('/api/reports/run/', spec, format='json').json()
        ids_a = [r['id'] for r in res_a['rows']]
        self.assertIn(self.trip_a.id, ids_a)
        self.assertNotIn(self.trip_b.id, ids_a)
        self.assertEqual(res_a['count'], 1)

        # سوپریوزر
        self.client.force_authenticate(self.superuser)
        res_s = self.client.post('/api/reports/run/', spec, format='json').json()
        ids_s = [r['id'] for r in res_s['rows']]
        self.assertIn(self.trip_a.id, ids_s)
        self.assertIn(self.trip_b.id, ids_s)
        self.assertGreaterEqual(res_s['count'], 2)

        # کاربر بی‌شرکت
        self.client.force_authenticate(self.user_no_company)
        res_none = self.client.post('/api/reports/run/', spec, format='json').json()
        self.assertEqual(res_none['count'], 0)
        self.assertEqual(len(res_none['rows']), 0)

    # -------------------------------------------------------------------------
    # ۴. run برای users، به‌علاوهٔ پنهان بودن سوپریوزر برای غیرسوپریوزر
    # -------------------------------------------------------------------------
    def test_run_report_users_scope_and_superuser_hidden(self):
        spec = {'entity': 'users', 'fields': ['id', 'username']}

        # کاربر شرکت الف
        self.client.force_authenticate(self.user_a)
        res_a = self.client.post('/api/reports/run/', spec, format='json').json()
        unames_a = [r['username'] for r in res_a['rows']]
        self.assertIn(self.user_a.username, unames_a)
        self.assertNotIn(self.user_b.username, unames_a)
        self.assertNotIn(self.superuser.username, unames_a)

        # سوپریوزر
        self.client.force_authenticate(self.superuser)
        res_s = self.client.post('/api/reports/run/', spec, format='json').json()
        unames_s = [r['username'] for r in res_s['rows']]
        self.assertIn(self.user_a.username, unames_s)
        self.assertIn(self.user_b.username, unames_s)
        self.assertIn(self.superuser.username, unames_s)
        self.assertIn(self.user_no_company.username, unames_s)

        # کاربر بی‌شرکت: کاربران شرکت الف و ب و سوپریوزر را نمی‌بیند
        self.client.force_authenticate(self.user_no_company)
        res_none = self.client.post('/api/reports/run/', spec, format='json').json()
        unames_none = [r['username'] for r in res_none['rows']]
        self.assertNotIn(self.user_a.username, unames_none)
        self.assertNotIn(self.user_b.username, unames_none)
        self.assertNotIn(self.superuser.username, unames_none)

    # -------------------------------------------------------------------------
    # ۵. مسیر export همان دامنه را دارد
    # -------------------------------------------------------------------------
    def test_export_report_company_scope(self):
        if not apps.is_installed('personnel'):
            self.skipTest('ماژول پرسنل و حسابداری نصب نیست.')

        spec = {'entity': 'personnel_profiles', 'fields': ['id', 'first_name']}

        # کاربر شرکت الف
        eng_a = ReportEngine(self.user_a, spec)
        qs_a, _, total_a = eng_a.export_queryset()
        self.assertEqual(total_a, 1)
        self.assertEqual(list(qs_a.values_list('id', flat=True)), [self.profile_a.id])

        # سوپریوزر
        eng_s = ReportEngine(self.superuser, spec)
        _, _, total_s = eng_s.export_queryset()
        self.assertGreaterEqual(total_s, 2)

        # کاربر بی‌شرکت
        eng_none = ReportEngine(self.user_no_company, spec)
        _, _, total_none = eng_none.export_queryset()
        self.assertEqual(total_none, 0)

        # اعتبارسنجی از مسیر HTTP
        self.client.force_authenticate(self.user_a)
        r_exp = self.client.post('/api/reports/export/', spec, format='json')
        self.assertEqual(r_exp.status_code, 200)

    # -------------------------------------------------------------------------
    # ۶. پیمایش همهٔ موجودیت‌های ثبت‌شده: scope_path یا scope_resolver یا در فهرست سفید
    # -------------------------------------------------------------------------
    def test_all_registered_entities_have_scope_or_whitelisted(self):
        """
        فهرست سفید موجودیت‌های سراسری بدون قلمرو (GLOBAL_WHITELIST):
        در حال حاضر کلیهٔ موجودیت‌های سیستم باید دارای قلمرو مشخص (انبار یا شرکت یا قلاب) باشند.
        اگر موجودیتی کاملاً عمومی در پلتفرم اضافه شود (مانند تقویم تعطیلات سراسری)،
        باید با ذکر دلیل موجه در این دیکشنری ثبت شود.
        """
        GLOBAL_WHITELIST = {
            # مثال: 'holidays': 'تقویم رسمی تعطیلات کشوری عمومی بوده و مستقل از شرکت است'
        }

        registry = get_all_entities()
        self.assertGreater(len(registry), 0, 'حداقل یک موجودیت باید ثبت شده باشد.')

        for key, cfg in registry.items():
            has_scope_path = bool(cfg.scope_path)
            has_resolver = callable(getattr(cfg, 'scope_resolver', None))
            is_whitelisted = key in GLOBAL_WHITELIST
            self.assertTrue(
                has_scope_path or has_resolver or is_whitelisted,
                f"موجودیت «{key}» فاقد scope_path یا scope_resolver است و در فهرست سفید سراسری هم نیست!"
            )

    # -------------------------------------------------------------------------
    # ۷. ورود با update_existing=1 و ردیفی برای کاربر شرکت ب: خطای ردیف و کاربر هدف دست‌نخورده
    # -------------------------------------------------------------------------
    def test_excel_import_update_existing_target_in_other_company_rejected(self):
        orig_first_name = self.user_b.first_name
        orig_phone = self.user_b.phone_number

        excel_data = _make_users_excel([
            {
                'first_name': 'مهاجم',
                'last_name': 'تست',
                'username': self.user_b.username,
                'national_code': self.user_b.national_code,
                'phone_number': '09129999999',
                'is_active': 'خیر',
            }
        ])

        file = SimpleUploadedFile('import.xlsx', excel_data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.force_authenticate(self.user_a)
        res = self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file, 'update_existing': 'true'},
            format='multipart',
            HTTP_X_COMPANY_ID=str(self.company_a.id)
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        errors = data.get('errors', [])
        self.assertTrue(any(e.get('field') == 'username' for e in errors), f"خطای فیلد username دریافت نشد: {errors}")

        # اطمینان از دست‌نخورده ماندن کاربر شرکت ب در دیتابیس
        self.user_b.refresh_from_db()
        self.assertEqual(self.user_b.first_name, orig_first_name)
        self.assertEqual(self.user_b.phone_number, orig_phone)
        self.assertTrue(self.user_b.is_active)

    # -------------------------------------------------------------------------
    # ۸. ورود با update_existing=1 و ردیفی برای سوپریوزر: خطای ردیف و سوپریوزر دست‌نخورده
    # -------------------------------------------------------------------------
    def test_excel_import_update_existing_target_superuser_rejected(self):
        orig_phone = self.superuser.phone_number

        excel_data = _make_users_excel([
            {
                'first_name': 'هک سوپریوزر',
                'last_name': 'ناموفق',
                'username': self.superuser.username,
                'national_code': self.superuser.national_code,
                'phone_number': '09120000000',
                'is_active': 'خیر',
            }
        ])

        file = SimpleUploadedFile('import.xlsx', excel_data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.force_authenticate(self.user_a)
        res = self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file, 'update_existing': 'true'},
            format='multipart',
            HTTP_X_COMPANY_ID=str(self.company_a.id)
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        errors = data.get('errors', [])
        self.assertTrue(any(e.get('field') == 'username' for e in errors), f"خطای شناسه ورود سوپریوزر یافت نشد: {errors}")

        self.superuser.refresh_from_db()
        self.assertEqual(self.superuser.phone_number, orig_phone)
        self.assertTrue(self.superuser.is_active)

    # -------------------------------------------------------------------------
    # ۹. ورود با انبار شرکت ب: خطای ردیف
    # -------------------------------------------------------------------------
    def test_excel_import_warehouse_in_other_company_rejected(self):
        if not apps.is_installed('warehouses'):
            self.skipTest('ماژول انبار نصب نیست.')

        excel_data = _make_users_excel([
            {
                'first_name': 'کاربر',
                'last_name': 'تست انبار',
                'username': 'wh_attacker',
                'national_code': '0098765432',
                'warehouses': self.wh_b.name,
                'is_active': 'بله',
            }
        ])

        file = SimpleUploadedFile('import.xlsx', excel_data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.force_authenticate(self.user_a)
        res = self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file, 'update_existing': 'false'},
            format='multipart',
            HTTP_X_COMPANY_ID=str(self.company_a.id)
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        errors = data.get('errors', [])
        self.assertTrue(any(e.get('field') == 'warehouses' for e in errors), f"خطای انبار دریافت نشد: {errors}")
        self.assertFalse(User.objects.filter(username='wh_attacker').exists())

    # -------------------------------------------------------------------------
    # ۱۰. قالب دانلودی نام انبار شرکت ب را ندارد
    # -------------------------------------------------------------------------
    def test_download_template_excludes_other_company_warehouses(self):
        if not apps.is_installed('warehouses'):
            self.skipTest('ماژول انبار نصب نیست.')

        self.client.force_authenticate(self.user_a)
        res = self.client.get('/api/auth/users/download_template/', HTTP_X_COMPANY_ID=str(self.company_a.id))
        self.assertEqual(res.status_code, 200)

        wb = load_workbook(io.BytesIO(res.content), data_only=True)
        all_values = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for val in row:
                    if val is not None:
                        all_values.append(str(val))

        self.assertIn(self.wh_a.name, all_values, 'نام انبار شرکت الف باید در قالب باشد.')
        self.assertNotIn(self.wh_b.name, all_values, 'نام انبار شرکت ب نباید در قالب کاربر شرکت الف باشد.')

    # -------------------------------------------------------------------------
    # ۱۱. ورود موفق کاربر مجاز در دامنهٔ خودش و ایجاد UserCompanyAccess با workspace_full
    # -------------------------------------------------------------------------
    def test_excel_import_success_in_own_scope_creates_user_company_access(self):
        new_username = 'legit_user_a'
        excel_data = _make_users_excel([
            {
                'first_name': 'رضا',
                'last_name': 'صادقی',
                'username': new_username,
                'national_code': '0087654321',
                'phone_number': '09128887766',
                'is_active': 'بله',
            }
        ])

        file = SimpleUploadedFile('import.xlsx', excel_data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.force_authenticate(self.user_a)
        res = self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file, 'update_existing': 'false'},
            format='multipart',
            HTTP_X_COMPANY_ID=str(self.company_a.id)
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data.get('success'))
        self.assertEqual(data.get('summary', {}).get('created'), 1)

        new_user = User.objects.filter(username=new_username).first()
        self.assertIsNotNone(new_user)
        access = UserCompanyAccess.objects.filter(user=new_user, company=self.company_a).first()
        self.assertIsNotNone(access, 'رکورد UserCompanyAccess برای کاربر جدید ساخته نشده است.')
        self.assertEqual(access.section_levels, {})

    # -------------------------------------------------------------------------
    # ۱۲. غیرسوپریوزر بی‌شرکت فعال: خطای روشن فارسی و هیچ کاربری ساخته نمی‌شود
    # -------------------------------------------------------------------------
    def test_excel_import_non_superuser_without_active_company_rejected(self):
        new_username = 'ghost_user'
        excel_data = _make_users_excel([
            {
                'first_name': 'شبح',
                'last_name': 'ناشناس',
                'username': new_username,
                'national_code': '0076543210',
                'is_active': 'بله',
            }
        ])

        file = SimpleUploadedFile('import.xlsx', excel_data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.force_authenticate(self.user_a)
        # فراخوانی بدون هدر X-Company-ID
        res = self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file, 'update_existing': 'false'},
            format='multipart'
        )
        self.assertEqual(res.status_code, 400)
        data = res.json()
        errors = data.get('errors', [])
        self.assertTrue(any(e.get('field') == 'company' for e in errors), f"خطای الزام شرکت فعال دریافت نشد: {errors}")
        self.assertFalse(User.objects.filter(username=new_username).exists())

    # -------------------------------------------------------------------------
    # ۱۳. تکرار نام کاربری یا کد ملی کاربرِ بیرون از دامنه: خطای ردیف، بی IntegrityError و ۵۰۰
    # -------------------------------------------------------------------------
    def test_excel_import_duplicate_username_or_nid_outside_scope_gives_row_error_no_500(self):
        # تکرار نام کاربری متعلق به شرکت دیگر
        excel_data_uname = _make_users_excel([
            {
                'first_name': 'تکراری',
                'last_name': 'کاربری',
                'username': self.user_b.username,
                'national_code': '0091122334',
                'is_active': 'بله',
            }
        ])
        file_uname = SimpleUploadedFile('uname.xlsx', excel_data_uname, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.force_authenticate(self.user_a)
        res_u = self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file_uname, 'update_existing': 'false'},
            format='multipart',
            HTTP_X_COMPANY_ID=str(self.company_a.id)
        )
        self.assertNotEqual(res_u.status_code, 500)
        data_u = res_u.json()
        errors_u = data_u.get('errors', [])
        self.assertTrue(any(e.get('field') == 'username' for e in errors_u), f"خطای ردیف برای نام کاربری بیرون از دامنه صادر نشد: {errors_u}")

        # تکرار کد ملی متعلق به شرکت دیگر
        excel_data_nid = _make_users_excel([
            {
                'first_name': 'تکراری',
                'last_name': 'کدملی',
                'username': 'brand_new_user_xyz',
                'national_code': self.user_b.national_code,
                'is_active': 'بله',
            }
        ])
        file_nid = SimpleUploadedFile('nid.xlsx', excel_data_nid, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        res_n = self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file_nid, 'update_existing': 'false'},
            format='multipart',
            HTTP_X_COMPANY_ID=str(self.company_a.id)
        )
        self.assertNotEqual(res_n.status_code, 500)
        data_n = res_n.json()
        errors_n = data_n.get('errors', [])
        self.assertTrue(any(e.get('field') == 'national_code' for e in errors_n), f"خطای ردیف برای کد ملی بیرون از دامنه صادر نشد: {errors_n}")
        self.assertFalse(User.objects.filter(username='brand_new_user_xyz').exists())

    # -------------------------------------------------------------------------
    # بازبینی ناظر: یافته‌های tenant-scope-reviewer روی نشست ۱
    # -------------------------------------------------------------------------
    def _make_user_ab(self, username, national_code):
        """کاربر غیرسوپریوزر با دسترسی به هر دو شرکت."""
        u = User.objects.create_user(
            username=username, password='Password123!', national_code=national_code,
            first_name='کاربر', last_name='دوشرکتی',
        )
        UserCompanyAccess.objects.create(user=u, company=self.company_a, is_default=True)
        grant_membership(u, self.company_a)
        UserCompanyAccess.objects.create(user=u, company=self.company_b)
        grant_membership(u, self.company_b)
        return u

    def _import(self, actor, rows, update_existing, company):
        excel_data = _make_users_excel(rows)
        file = SimpleUploadedFile('import.xlsx', excel_data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.client.force_authenticate(actor)
        return self.client.post(
            '/api/auth/users/import_excel/',
            {'file': file, 'update_existing': 'true' if update_existing else 'false'},
            format='multipart',
            HTTP_X_COMPANY_ID=str(company.id)
        )

    def test_report_personnel_profiles_limited_to_assigned_sections(self):
        """کاربر غیرممیز شرکت الف که فقط به بخش الف منتسب است، پرسنل بخش دیگر همان شرکت را نمی‌بیند."""
        if not apps.is_installed('personnel'):
            self.skipTest('ماژول پرسنل و حسابداری نصب نیست.')
        from personnel.models import ProjectSection, PersonnelProfile
        sec_a2 = ProjectSection.objects.create(project=self.proj_a, code='SA2', name='بخش دوم الف')
        profile_a2 = PersonnelProfile.objects.create(
            company=self.company_a, section=sec_a2, first_name='پرسنل', last_name='بخش دوم',
            national_code='0077777777', personnel_id_code='EMP-A2',
        )
        spec = {'entity': 'personnel_profiles', 'fields': ['id', 'first_name']}

        self.client.force_authenticate(self.user_a)
        ids_a = [r['id'] for r in self.client.post('/api/reports/run/', spec, format='json').json()['rows']]
        self.assertIn(self.profile_a.id, ids_a)
        self.assertNotIn(profile_a2.id, ids_a)

        # ممیز پرسنل (حسابدار کل) به بخش محدود نیست ولی همچنان به شرکت‌های مجاز محدود است
        auditor = User.objects.create_user(username='pers_auditor_a', password='Password123!', national_code='0088888888')
        UserCompanyAccess.objects.create(user=auditor, company=self.company_a, is_default=True)
        grant_membership(auditor, self.company_a)
        auditor.user_permissions.set(Permission.objects.filter(
            codename__in=['view_sys_reports', 'view_sys_personnel', 'perm_approve_personnel_finance']
        ))
        self.client.force_authenticate(auditor)
        ids_aud = [r['id'] for r in self.client.post('/api/reports/run/', spec, format='json').json()['rows']]
        self.assertIn(self.profile_a.id, ids_aud)
        self.assertIn(profile_a2.id, ids_aud)
        self.assertNotIn(self.profile_b.id, ids_aud)

    def test_trips_endpoint_without_company_header_is_company_scoped(self):
        """GET /api/personnel/trips/ بی سربرگ شرکت: کاربر الف فقط سفر الف، سوپریوزر هر دو، کاربر بی‌شرکت هیچ."""
        if not apps.is_installed('personnel'):
            self.skipTest('ماژول پرسنل و حسابداری نصب نیست.')

        def trip_ids(user):
            self.client.force_authenticate(user)
            res = self.client.get('/api/personnel/trips/')
            self.assertEqual(res.status_code, 200, res.content)
            data = res.json()
            return {r['id'] for r in (data if isinstance(data, list) else data.get('results', []))}

        self.assertEqual(trip_ids(self.user_a), {self.trip_a.id})
        self.assertTrue({self.trip_a.id, self.trip_b.id} <= trip_ids(self.superuser))
        self.assertEqual(trip_ids(self.user_no_company), set())

        # دسترسی مستقیم با شناسه به سفر شرکت دیگر هم بسته است
        self.client.force_authenticate(self.user_a)
        self.assertEqual(self.client.get(f'/api/personnel/trips/{self.trip_b.id}/').status_code, 404)

    def test_report_users_grouped_count_has_no_duplicates(self):
        """گزارش گروه‌بندی‌شدهٔ کاربران، کاربری را که به چند شرکت مجاز دسترسی دارد یک بار می‌شمارد."""
        actor = self._make_user_ab('actor_ab_count', '0012121212')
        actor.user_permissions.set(Permission.objects.filter(codename__in=['view_sys_reports', 'view_sys_users']))
        self._make_user_ab('target_ab_count', '0013131313')

        self.client.force_authenticate(actor)
        flat = self.client.post('/api/reports/run/', {'entity': 'users', 'fields': ['id']}, format='json').json()
        grouped = self.client.post('/api/reports/run/', {
            'entity': 'users',
            'group_by': ['is_active'],
            'aggregations': [{'field': 'id', 'fn': 'count', 'alias': 'n'}],
        }, format='json').json()
        self.assertEqual(sum(int(r['n']) for r in grouped['rows']), flat['count'])

    def test_excel_import_update_keeps_warehouses_outside_actor_scope(self):
        """به‌روزرسانی با سازندهٔ فقط‌الف، انبار شرکت ب کاربر دوشرکتی را پاک نمی‌کند."""
        if not apps.is_installed('warehouses'):
            self.skipTest('ماژول انبار نصب نیست.')
        target = self._make_user_ab('target_ab_wh', '0014141414')
        target.assigned_warehouses.set([self.wh_a, self.wh_b])

        res = self._import(self.user_a, [{
            'first_name': 'کاربر', 'last_name': 'دوشرکتی', 'username': target.username,
            'national_code': target.national_code, 'warehouses': self.wh_a.name, 'is_active': 'بله',
        }], update_existing=True, company=self.company_a)
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json().get('summary', {}).get('updated'), 1, res.json())

        target.refresh_from_db()
        self.assertEqual(set(target.assigned_warehouses.all()), {self.wh_a, self.wh_b})

    def test_excel_import_warehouse_must_belong_to_new_user_company(self):
        """سازندهٔ دوشرکتی با شرکت فعال الف: انبار شرکت ب برای کاربر تازه خطای ردیف است و «همه انبارها» فقط انبارهای الف را می‌دهد."""
        if not apps.is_installed('warehouses'):
            self.skipTest('ماژول انبار نصب نیست.')
        actor = self._make_user_ab('actor_ab_wh', '0015151515')
        actor.user_permissions.set(Permission.objects.filter(codename__in=['view_sys_users', 'perm_usr_add']))

        res = self._import(actor, [
            {'first_name': 'تازه', 'last_name': 'یک', 'username': 'new_wh_b_user', 'national_code': '0016161616',
             'warehouses': self.wh_b.name, 'is_active': 'بله'},
            {'first_name': 'تازه', 'last_name': 'دو', 'username': 'new_all_wh_user', 'national_code': '0017171717',
             'warehouses': 'all_warehouses', 'is_active': 'بله'},
        ], update_existing=False, company=self.company_a)
        self.assertEqual(res.status_code, 200, res.content)
        errors = res.json().get('errors', [])
        self.assertTrue(any(e.get('field') == 'warehouses' for e in errors), errors)
        self.assertFalse(User.objects.filter(username='new_wh_b_user').exists())

        all_user = User.objects.get(username='new_all_wh_user')
        self.assertIn(self.wh_a, set(all_user.assigned_warehouses.all()))
        self.assertNotIn(self.wh_b, set(all_user.assigned_warehouses.all()))
