import io

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets, permissions, parsers, serializers
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response

from accounts.views import OptionalPageNumberPagination
from common.date_utils import normalize_digits

from platform_core.registry import contributions

from .access import (
    derived_company_ids,
    ensure_scoped_write,
    get_user_allowed_companies,
)
from .models import Company, UserCompanyAccess
from .permissions import IsOrgStructureManagerOrReadOnly, requires_company_scope
from .realtime import broadcast_org_structure_updated
from .serializers import CompanySerializer, UserCompanyAccessSerializer


# وابستگی‌های شرکت که در صورت وجود، حذف شرکت را مسدود می‌کنند.
# ماژول‌ها این وابستگی‌ها را در نقطهٔ توسعهٔ company_delete_blockers ثبت می‌کنند (گام ۴).


class CompanyViewSet(viewsets.ModelViewSet):
    """
    مدیریت شرکت‌ها (هلدینگ و شرکت‌های تابعه) در هسته. حسابداری با زیرکلاس خود
    (personnel.views.CompanyViewSet) شمارش پروژه‌ها، مدارک و پایش سررسید مدارک را اضافه می‌کند.
    """
    queryset = Company.objects.all()
    serializer_class = CompanySerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    pagination_class = OptionalPageNumberPagination

    def base_queryset(self):
        """نقطهٔ گسترش برای ماژول‌ها (annotate و prefetch)."""
        return Company.objects.all()

    def get_queryset(self):
        user = self.request.user
        qs = self.base_queryset()

        if not user.is_superuser and self.action != 'user_available':
            allowed_ids = get_user_allowed_companies(user).values_list('id', flat=True)
            qs = qs.filter(id__in=allowed_ids)

        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() == 'true')

        search = self.request.query_params.get('search')
        if search:
            search_clean = normalize_digits(str(search)).strip()
            qs = qs.filter(
                Q(name__icontains=search) |
                Q(code__icontains=search_clean) |
                Q(national_id__icontains=search_clean) |
                Q(ceo_name__icontains=search)
            )
        return qs.order_by('code')

    def perform_create(self, serializer):
        if not (self.request.user.is_superuser or self.request.user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به ایجاد شرکت جدید هستند.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('company', 'create', instance.id, instance.name, None, tab_id, self.request.user.id)

    def perform_update(self, serializer):
        if not (self.request.user.is_superuser or self.request.user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به ویرایش اطلاعات شرکت هستند.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('company', 'update', instance.id, instance.name, None, tab_id, self.request.user.id)

    def perform_destroy(self, instance):
        if not (self.request.user.is_superuser or self.request.user.is_staff):
            raise PermissionDenied("تنها مدیران ارشد مجاز به حذف شرکت هستند.")
        blockers = sorted(contributions('company_delete_blockers'), key=lambda b: b[0])
        for _order, related_name, message in blockers:
            if hasattr(instance, related_name) and getattr(instance, related_name).exists():
                raise serializers.ValidationError(message)
        instance_id = instance.id
        instance_name = instance.name
        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company', 'delete', instance_id, instance_name, None, tab_id, self.request.user.id)

    @action(detail=False, methods=['get'], url_path='user-available')
    def user_available(self, request):
        """
        لیست شرکت‌های در دسترس کاربر جاری برای نمایش در مودال ورود، هدر سوئیچر و کارتابل مدارک
        اگر scope='documents' یا 'docs' باشد: شرکت‌هایی که کاربر حداقل دسترسی docs_read دارد.
        اگر scope='all' باشد: تمامی شرکت‌های منتسب به کاربر.
        پیش‌فرض: شرکت‌هایی که کاربر عضویت فضای کاری (workspace_full) دارد.
        """
        user = request.user
        scope = request.query_params.get('scope', 'workspace')
        if scope in ('documents', 'docs'):
            companies = get_user_allowed_companies(user, required_level='docs_read')
        elif scope == 'all':
            companies = get_user_allowed_companies(user, required_level=None)
        else:
            companies = get_user_allowed_companies(user, required_level='workspace_full')

        serializer = self.get_serializer(companies, many=True)
        data = list(serializer.data)

        # الحاق سطح دسترسی کاربر در هر شرکت برای هدایت هوشمند در فرانت‌اند
        user_accesses = {
            a.company_id: a.access_level
            for a in UserCompanyAccess.objects.filter(user=user)
        }
        derived_ids = set() if user.is_superuser else derived_company_ids(user)
        for c in data:
            if user.is_superuser:
                c['user_access_level'] = 'workspace_full'
            elif c['id'] in user_accesses:
                c['user_access_level'] = user_accesses[c['id']]
            elif c['id'] in derived_ids:
                c['user_access_level'] = 'workspace_full'
            else:
                c['user_access_level'] = 'docs_read'

        return Response({
            'companies': data,
            'is_superuser': user.is_superuser,
            'count': len(data)
        })

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی اکسل استاندارد ۲ ردیفه اطلاعات شرکت‌ها منطبق بر استانداردهای مهندسی
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        qs = self.get_queryset()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "شرکت‌های هلدینگ"
        ws.views.sheetView[0].rightToLeft = True

        headers_fa = [
            'ردیف', 'کد یکتا', 'نام کامل شرکت', 'نوع شرکت', 'شناسه ملی', 'کد اقتصادی',
            'شماره ثبت', 'تلفن تماس', 'مدیرعامل', 'رئیس هیئت‌مدیره', 'کد کارگاه بیمه',
            'شناسه مودیان', 'شماره شبا رسمی', 'تعداد پروژه‌ها', 'تعداد اسناد', 'وضعیت', 'نشانی دفتر مرکزی'
        ]
        headers_en = [
            'row_num', 'code', 'name', 'company_type', 'national_id', 'economic_code',
            'registration_number', 'phone', 'ceo_name', 'board_chairman', 'workshop_code',
            'tax_memory_id', 'primary_iban', 'projects_count', 'documents_count', 'is_active', 'address'
        ]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(headers_fa) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 16
        ws.freeze_panes = 'A3'

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, comp in enumerate(qs, start=1):
            # شمار پروژه‌ها و اسناد فقط وقتی حسابداری نصب باشد معنا دارد
            if hasattr(comp, 'projects_count'):
                p_count = comp.projects_count
            else:
                p_count = comp.projects.count() if hasattr(comp, 'projects') else 0
            d_count = comp.documents.count() if hasattr(comp, 'documents') else 0
            ws.append([
                idx,
                comp.code,
                comp.name,
                comp.get_company_type_display(),
                comp.national_id or '',
                comp.economic_code or '',
                comp.registration_number or '',
                comp.phone or '',
                comp.ceo_name or '',
                comp.board_chairman or '',
                comp.workshop_code or '',
                comp.tax_memory_id or '',
                comp.primary_iban or '',
                p_count,
                d_count,
                'فعال' if comp.is_active else 'غیرفعال',
                comp.address or ''
            ])
            curr_row = idx + 2
            ws.row_dimensions[curr_row].height = 20
            for col_idx in range(1, len(headers_fa) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal='center' if col_idx not in [3, 17] else 'right', vertical='center')
                if col_idx in [1, 2, 5, 6, 7, 8, 11, 12, 13, 14, 15]:
                    c.font = data_font_num
                else:
                    c.font = data_font

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"companies_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class UserCompanyAccessViewSet(viewsets.ModelViewSet):
    """
    مدیریت انتساب دسترسی کاربران به شرکت‌ها
    """
    queryset = UserCompanyAccess.objects.all().select_related('user', 'company')
    serializer_class = UserCompanyAccessSerializer
    permission_classes = [permissions.IsAuthenticated, IsOrgStructureManagerOrReadOnly]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if requires_company_scope(user):
            # کاربر غیرسوپریوزر فقط دسترسی‌های شرکت‌های مجاز خودش و ردیف‌های خودش را می‌بیند
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            qs = qs.filter(Q(company_id__in=allowed_ids) | Q(user=user))
        user_id = self.request.query_params.get('user_id')
        company_id = self.request.query_params.get('company_id')
        if user_id:
            qs = qs.filter(user_id=user_id)
        if company_id:
            qs = qs.filter(company_id=company_id)
        return qs.order_by('-created_at')

    def perform_create(self, serializer):
        if requires_company_scope(self.request.user):
            company = serializer.validated_data.get('company')
            ensure_scoped_write(self.request.user, company.id if company else None, "دسترسی فقط برای شرکت‌های مجاز شما قابل اعطاست.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_access', 'create', instance.id, str(instance.user), client_tab_id=tab_id, sender_id=self.request.user.id, company_id=instance.company_id)

    def perform_update(self, serializer):
        if requires_company_scope(self.request.user):
            ensure_scoped_write(self.request.user, serializer.instance.company_id, "فقط دسترسی‌های شرکت‌های مجاز خود را می‌توانید تغییر دهید.")
            if 'company' in serializer.validated_data:
                company = serializer.validated_data.get('company')
                ensure_scoped_write(self.request.user, company.id if company else None, "انتقال دسترسی به شرکتی که مجاز نیستید ممکن نیست.")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_access', 'update', instance.id, str(instance.user), client_tab_id=tab_id, sender_id=self.request.user.id, company_id=instance.company_id)

    def perform_destroy(self, instance):
        if requires_company_scope(self.request.user):
            ensure_scoped_write(self.request.user, instance.company_id, "فقط دسترسی‌های شرکت‌های مجاز خود را می‌توانید حذف کنید.")
        instance_id = instance.id
        user_title = str(instance.user)
        comp_id = instance.company_id
        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_access', 'delete', instance_id, user_title, client_tab_id=tab_id, sender_id=self.request.user.id, company_id=comp_id)
