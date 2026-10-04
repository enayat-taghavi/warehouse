from django.test import TestCase
from rest_framework.test import APIClient
from django.contrib.auth.models import Permission
from decimal import Decimal
from accounts.models import CustomUser
from personnel.models import (
    FinancialProject,
    ProjectSection,
    VehicleDriverProfile,
    VehicleChangeRequest,
    VehicleTripLog,
    UserSectionAssignment,
)
from personnel.sheba_utils import generate_sheba_from_account


class VehicleLifecycleTests(TestCase):
    """
    تست‌های جامع چرخه حیات ناوگان، امنیت قلمرو، جلوگیری از تنزل رتبه در تغییرات و حذف امن
    """

    def setUp(self):
        from personnel.models import Company
        self.company, _ = Company.objects.get_or_create(
            code="DALAN-CO",
            defaults={"name": "شرکت دالان", "is_active": True}
        )
        self.project = FinancialProject.objects.create(code="1", name="پروژه دالان", company=self.company, is_active=True)
        self.section_a = ProjectSection.objects.create(project=self.project, code="11", name="کارگاه دالان", is_active=True)
        self.section_b = ProjectSection.objects.create(project=self.project, code="12", name="انبار دالان", is_active=True)

        self.operator = CustomUser.objects.create(username="veh_operator", first_name="رضا", last_name="اکبری")
        self.supervisor = CustomUser.objects.create(username="veh_supervisor", first_name="حمید", last_name="صالحی")
        self.accountant = CustomUser.objects.create(username="veh_accountant", first_name="مهری", last_name="باقری")
        self.manager = CustomUser.objects.create(username="veh_manager", first_name="سید", last_name="موسوی")
        self.outsider = CustomUser.objects.create(username="veh_outsider", first_name="کاربر", last_name="بی‌دسترسی")

        # انتساب مجوزها
        for codename in ["perm_approve_fleet_supervisor", "can_act_as_workshop_supervisor"]:
            p = Permission.objects.filter(codename=codename).first()
            if p:
                self.supervisor.user_permissions.add(p)

        p_fin = Permission.objects.filter(codename="perm_approve_fleet_finance").first()
        if p_fin:
            self.accountant.user_permissions.add(p_fin)

        for codename in ["perm_approve_fleet_manager", "can_act_as_manager"]:
            p = Permission.objects.filter(codename=codename).first()
            if p:
                self.manager.user_permissions.add(p)

        # انتساب بخش A به کاربران مجاز
        UserSectionAssignment.objects.create(user=self.operator, section=self.section_a, role="employee", is_active=True)
        UserSectionAssignment.objects.create(user=self.supervisor, section=self.section_a, role="supervisor", is_active=True)
        UserSectionAssignment.objects.create(user=self.accountant, section=self.section_a, role="accountant", is_active=True)
        UserSectionAssignment.objects.create(user=self.manager, section=self.section_a, role="manager", is_active=True)

        self.client = APIClient()

    def test_complete_vehicle_approval_lifecycle(self):
        """
        تست چرخه کامل: ثبت پیش‌نویس -> تایید سرپرست -> تایید مالی -> تصویب نهایی مدیر
        """
        valid_sheba = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        self.client.force_authenticate(user=self.operator)

        # ۱. ثبت با وضعیت پیش‌نویس
        payload = {
            "plate_number": "12الف345ایران63",
            "vehicle_type": "nissan",
            "ownership_type": "contract",
            "driver_name": "احمد قاسمی",
            "driver_national_code": "0010376488",
            "driver_phone": "09121234567",
            "default_service_rate": 4500000,
            "insurance_expiry_date": "1406/01/01",
            "bank_name": "بانک ملی ایران",
            "sheba_number": valid_sheba,
            "section": self.section_a.id,
            "approval_status": "draft",
        }
        res_create = self.client.post("/api/personnel/vehicles/", payload, format="json")
        self.assertEqual(res_create.status_code, 201)
        veh_id = res_create.data["id"]
        self.assertEqual(res_create.data["approval_status"], "draft")
        self.assertEqual(len(res_create.data["bank_accounts"]), 1)
        self.assertTrue(res_create.data["bank_accounts"][0]["is_primary"])
        self.assertEqual(res_create.data["bank_accounts"][0]["sheba_number"], valid_sheba)

        # ۲. اپراتور ارسال به سرپرست می‌کند
        res_submit = self.client.patch(f"/api/personnel/vehicles/{veh_id}/", {"approval_status": "pending_supervisor"}, format="json")
        self.assertEqual(res_submit.status_code, 200)
        self.assertEqual(res_submit.data["approval_status"], "pending_supervisor")

        # ۳. تایید سرپرست
        self.client.force_authenticate(user=self.supervisor)
        res_sup = self.client.post(f"/api/personnel/vehicles/{veh_id}/approve-supervisor/")
        self.assertEqual(res_sup.status_code, 200)
        self.assertEqual(res_sup.data["data"]["approval_status"], "pending_accountant")

        # ۴. تایید مالی (حسابدار نرخ مصوب را تعیین و تایید مالی می‌کند)
        self.client.force_authenticate(user=self.accountant)
        res_rate = self.client.patch(f"/api/personnel/vehicles/{veh_id}/", {"default_service_rate": 4500000}, format="json")
        self.assertEqual(res_rate.status_code, 200)
        res_fin = self.client.post(f"/api/personnel/vehicles/{veh_id}/approve-finance/")
        self.assertEqual(res_fin.status_code, 200)
        self.assertEqual(res_fin.data["data"]["approval_status"], "pending_manager")

        # ۵. تصویب نهایی مدیر
        self.client.force_authenticate(user=self.manager)
        res_mgr = self.client.post(f"/api/personnel/vehicles/{veh_id}/approve-manager/")
        self.assertEqual(res_mgr.status_code, 200)
        self.assertEqual(res_mgr.data["data"]["approval_status"], "approved")
        self.assertTrue(res_mgr.data["data"]["is_active"])

    def test_prevent_demotion_bug_on_change_request_approval(self):
        """
        راستی‌آزمایی رفع باگ تنزل رتبه:
        تایید درخواست تغییرات یک خودروی مصوب، نباید وضعیت خودرو را به pending یا draft بازگرداند!
        """
        # ایجاد یک خودروی تاییدشده اولیه
        vehicle = VehicleDriverProfile.objects.create(
            plate_number="34ب567ایران63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="محسن رضایی",
            driver_national_code="0010376488",
            default_service_rate=Decimal("6000000"),
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )

        # حسابدار درخواست تغییر نرخ و نام راننده ارسال می‌کند (SoD: تغییر نرخ فقط با دسترسی مالی پذیرفته می‌شود)
        self.client.force_authenticate(user=self.accountant)
        edit_payload = {
            "driver_name": "محسن رضایی اصل",
            "default_service_rate": 6500000,
            "approval_status": "pending_supervisor",  # حتی اگر فرانت‌اند وضعیت را بفرستد
        }
        res_edit = self.client.patch(f"/api/personnel/vehicles/{vehicle.id}/", edit_payload, format="json")
        self.assertEqual(res_edit.status_code, 202)  # درخواست به کارتابل تغییرات رفت
        cr_id = res_edit.data["change_request_id"]

        cr = VehicleChangeRequest.objects.get(id=cr_id)
        self.assertEqual(cr.status, "pending_supervisor")

        # تایید مرحله‌ای درخواست تغییرات طبق SoD (سرپرست -> مالی -> مدیر)
        self.client.force_authenticate(user=self.supervisor)
        res_sup = self.client.post(f"/api/personnel/vehicle-change-requests/{cr_id}/approve-supervisor/")
        self.assertEqual(res_sup.status_code, 200)

        self.client.force_authenticate(user=self.accountant)
        res_fin = self.client.post(f"/api/personnel/vehicle-change-requests/{cr_id}/approve-finance/")
        self.assertEqual(res_fin.status_code, 200)

        self.client.force_authenticate(user=self.manager)
        res_approve_cr = self.client.post(f"/api/personnel/vehicle-change-requests/{cr_id}/approve-manager/")
        self.assertEqual(res_approve_cr.status_code, 200)

        # بررسی پرونده اصلی خودرو در دیتابیس
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.driver_name, "محسن رضایی اصل")
        self.assertEqual(vehicle.default_service_rate, Decimal("6500000"))
        # مهم‌ترین اعتبارسنجی: وضعیت خودرو همچنان باید مصوب باقی بماند و به pending تنزل نیابد!
        self.assertEqual(vehicle.approval_status, "approved")
        self.assertFalse(vehicle.has_pending_changes)

    def test_section_scoping_and_bola_protection(self):
        """
        بررسی ایزولاسیون قلمرو (Guardian G1): کاربر بدون انتساب نباید به خودروهای بخش دیگر دسترسی داشته باشد
        """
        VehicleDriverProfile.objects.create(
            plate_number="99ج999ایران99",
            vehicle_type="nissan",
            driver_name="راننده بخش B",
            section=self.section_b,
            approval_status="approved",
            is_active=True
        )

        # کاربر outsider دسترسی به بخش B ندارد
        self.client.force_authenticate(user=self.outsider)
        res_list = self.client.get(f"/api/personnel/vehicles/?section_id={self.section_b.id}")
        self.assertEqual(len(res_list.data), 0)

        # کاربر outsider نمی‌تواند در بخش B خودرو ثبت کند
        create_payload = {
            "plate_number": "88د888ایران88",
            "vehicle_type": "nissan",
            "driver_name": "تلاش غیرمجاز",
            "section": self.section_b.id,
        }
        res_create = self.client.post("/api/personnel/vehicles/", create_payload, format="json")
        self.assertEqual(res_create.status_code, 403)

    def test_safe_destroy_prevents_hard_cascade_delete(self):
        """
        بررسی حذف امن: خودرویی که دارای سابقه کارکرد و تردد است نباید حذف فیزیکی شود، بلکه نرم غیرفعال می‌گردد
        """
        vehicle = VehicleDriverProfile.objects.create(
            plate_number="55س555ایران55",
            vehicle_type="nissan",
            driver_name="راننده دارای سابقه",
            section=self.section_a,
            approval_status="draft",
            is_active=True
        )
        # ثبت یک تردد برای این خودرو
        VehicleTripLog.objects.create(
            vehicle=vehicle,
            date_shamsi="1403/04/01",
            trip_count=2,
            unit_rate=4000000,
            total_amount=8000000
        )

        self.client.force_authenticate(user=self.operator)
        res_del = self.client.delete(f"/api/personnel/vehicles/{vehicle.id}/")
        self.assertEqual(res_del.status_code, 400)
        self.assertEqual(res_del.data.get("code"), "history_exists")
        self.assertEqual(res_del.data.get("suggested_action"), "end_collaboration")

        # رکورد نباید از دیتابیس پاک شده باشد (محافظت در برابر حذف آبشاری)
        vehicle.refresh_from_db()
        self.assertTrue(VehicleDriverProfile.objects.filter(id=vehicle.id).exists())

    def test_destroy_nomovement_vehicle_with_permission(self):
        """
        تست حذف قطعی خودروی فاقد سابقه با مجوز perm_delete_vehicle_nomovement یا توسط ثبت‌کننده در پیش‌نویس
        """
        # خودروی ۱: ثبت توسط اپراتور در وضعیت پیش‌نویس
        v1 = VehicleDriverProfile.objects.create(
            plate_number="11ع111-11",
            vehicle_type="pickup",
            driver_name="راننده بدون سابقه ۱",
            section=self.section_a,
            approval_status="draft",
            created_by=self.operator,
            is_active=True
        )
        # کاربر بدون دسترسی نمی‌تواند حذف کند
        self.client.force_authenticate(user=self.outsider)
        res_outsider = self.client.delete(f"/api/personnel/vehicles/{v1.id}/")
        self.assertIn(res_outsider.status_code, [403, 404])
        self.assertTrue(VehicleDriverProfile.objects.filter(id=v1.id).exists())

        # ثبت‌کننده در وضعیت draft می‌تواند خودروی بدون سابقه را حذف کند
        self.client.force_authenticate(user=self.operator)
        res_del1 = self.client.delete(f"/api/personnel/vehicles/{v1.id}/")
        self.assertEqual(res_del1.status_code, 200)
        self.assertFalse(VehicleDriverProfile.objects.filter(id=v1.id).exists())

        # خودروی ۲: در وضعیت عودت‌داده‌شده با مجوز perm_delete_vehicle_nomovement
        user_deleter = CustomUser.objects.create(username="veh_deleter", first_name="مسئول", last_name="حذف")
        p_del = Permission.objects.filter(codename="perm_delete_vehicle_nomovement").first()
        if p_del:
            user_deleter.user_permissions.add(p_del)
        UserSectionAssignment.objects.create(user=user_deleter, section=self.section_a, role="supervisor", is_active=True)

        v2 = VehicleDriverProfile.objects.create(
            plate_number="22ع222-22",
            vehicle_type="nissan",
            driver_name="راننده بدون سابقه ۲",
            section=self.section_a,
            approval_status="revision_required",
            created_by=self.operator,
            is_active=True
        )
        self.client.force_authenticate(user=user_deleter)
        res_del2 = self.client.delete(f"/api/personnel/vehicles/{v2.id}/")
        self.assertEqual(res_del2.status_code, 200)
        self.assertFalse(VehicleDriverProfile.objects.filter(id=v2.id).exists())

    def test_vehicle_bank_account_and_documents(self):
        """
        تست جامع حساب‌های بانکی ناوگان (VehicleBankAccount): قید یکتایی حساب اصلی،
        همگام‌سازی خودکار ستون‌های آینه، جابه‌جایی حساب اصلی و ثبت مدارک (VehicleDocument)
        """
        from personnel.models import VehicleBankAccount, VehicleDocument
        v = VehicleDriverProfile.objects.create(
            plate_number="77ع777-77",
            vehicle_type="nissan",
            driver_name="راننده چند حسابی",
            section=self.section_a,
            approval_status="draft",
            created_by=self.operator,
            is_active=True
        )
        # حساب اول: خودکار اصلی می‌شود
        acc1 = VehicleBankAccount.objects.create(
            vehicle=v,
            bank_name="بانک تجارت",
            account_number="12345678",
            sheba_number="IR120180000000001234567890",
            payee_type="driver",
            is_primary=True
        )
        v.refresh_from_db()
        self.assertEqual(v.bank_name, "بانک تجارت")
        self.assertEqual(v.account_number, "12345678")
        self.assertEqual(v.sheba_number, "IR120180000000001234567890")

        # حساب دوم: با is_primary=True ذخیره می‌شود و حساب اول خودکار غیر اصلی می‌شود
        acc2 = VehicleBankAccount.objects.create(
            vehicle=v,
            bank_name="بانک ملت",
            account_number="87654321",
            sheba_number="IR160120000000008765432100",
            payee_type="owner",
            is_primary=True
        )
        acc1.refresh_from_db()
        self.assertFalse(acc1.is_primary)
        v.refresh_from_db()
        self.assertEqual(v.bank_name, "بانک ملت")
        self.assertEqual(v.sheba_number, "IR160120000000008765432100")

        # حذف حساب اصلی: حساب فعال بعدی خودکار اصلی می‌شود
        acc2.delete()
        acc1.refresh_from_db()
        self.assertTrue(acc1.is_primary)
        v.refresh_from_db()
        self.assertEqual(v.bank_name, "بانک تجارت")

        # مدارک خودرو
        doc = VehicleDocument.objects.create(
            vehicle=v,
            doc_type="insurance",
            file="vehicle_docs/test.pdf",
            title="بیمه‌نامه ثالث",
            issue_date="1403/01/01",
            expiry_date="1404/01/01",
            uploaded_by=self.operator
        )
        self.assertEqual(v.documents.count(), 1)
        self.assertEqual(v.documents.first().doc_type, "insurance")

    def test_excel_download_template(self):
        """
        دانلود قالب اکسل ناوگان
        """
        self.client.force_authenticate(user=self.operator)
        res = self.client.get("/api/personnel/vehicles/download-template/")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_excel_import_lifecycle_vehicles(self):
        """
        تست چرخه کامل درون‌ریزی اکسل ناوگان: دانلود قالب، ارسال در حالت پیش‌نمایش و اعمال قطعی
        """
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_authenticate(user=self.operator)

        # ۱. دریافت قالب اکسل
        tpl_res = self.client.get("/api/personnel/vehicles/download-template/")
        self.assertEqual(tpl_res.status_code, 200)

        # ۲. پیش‌نمایش (Dry-Run)
        excel_file = SimpleUploadedFile("template.xlsx", tpl_res.content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        res_dry = self.client.post(
            "/api/personnel/vehicles/import-excel/",
            {"file": excel_file, "section_id": self.section_a.id, "dry_run": "true"},
            format="multipart"
        )
        self.assertEqual(res_dry.status_code, 200)
        self.assertTrue(res_dry.data["success"])
        self.assertTrue(res_dry.data["dry_run"])
        self.assertEqual(res_dry.data["summary"]["created"], 1)
        self.assertFalse(VehicleDriverProfile.objects.filter(plate_number="12 الف 345 ایران 63").exists())

        # ۳. اعمال قطعی
        excel_file_real = SimpleUploadedFile("template.xlsx", tpl_res.content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        res_commit = self.client.post(
            "/api/personnel/vehicles/import-excel/",
            {"file": excel_file_real, "section_id": self.section_a.id, "dry_run": "false"},
            format="multipart"
        )
        self.assertEqual(res_commit.status_code, 200)
        self.assertTrue(res_commit.data["success"])
        self.assertFalse(res_commit.data["dry_run"])
        self.assertEqual(res_commit.data["created_count"], 1)

        # بررسی در دیتابیس
        v = VehicleDriverProfile.objects.filter(plate_number="12 الف 345 ایران 63").first()
        self.assertIsNotNone(v)
        self.assertEqual(v.driver_name, "رضا اکبری")
        self.assertEqual(v.approval_status, "draft")
        self.assertEqual(v.section_id, self.section_a.id)

    def test_vehicle_owner_fields_and_validation(self):
        """
        تست تفکیک راننده و مالک خودرو و اعتبارسنجی کدملی مالک
        """
        self.client.force_authenticate(user=self.operator)

        # ۱. تست ثبت موفق با مالک مجزا
        payload = {
            "plate_number": "99ط888ایران63",
            "vehicle_type": "truck",
            "ownership_type": "contract",
            "driver_name": "سعید رضایی",
            "driver_national_code": "0010376488",
            "is_driver_owner": False,
            "owner_name": "حسین حسینی",
            "owner_national_code": "0010376488",
            "owner_phone": "09123334455",
            "section": self.section_a.id,
            "approval_status": "draft"
        }
        res = self.client.post("/api/personnel/vehicles/", payload, format="json")
        self.assertEqual(res.status_code, 201)
        v = VehicleDriverProfile.objects.get(id=res.data["id"])
        self.assertFalse(v.is_driver_owner)
        self.assertEqual(v.owner_name, "حسین حسینی")
        self.assertEqual(v.owner_national_code, "0010376488")
        self.assertEqual(v.owner_phone, "09123334455")

        # ۲. تست رد شدن کدملی نامعتبر مالک
        bad_payload = {
            "plate_number": "88ق777ایران63",
            "vehicle_type": "nissan",
            "ownership_type": "contract",
            "driver_name": "سعید رضایی",
            "is_driver_owner": False,
            "owner_name": "حسین حسینی",
            "owner_national_code": "1111111111",
            "section": self.section_a.id,
            "approval_status": "draft"
        }
        res_bad = self.client.post("/api/personnel/vehicles/", bad_payload, format="json")
        self.assertEqual(res_bad.status_code, 400)
        self.assertIn("owner_national_code", res_bad.data)

    def test_export_excel_includes_owner_columns(self):
        """
        تست وجود ستون‌های مالک و نوع مالکیت در خروجی رسمی اکسل ناوگان
        """
        import io
        import openpyxl

        self.client.force_authenticate(user=self.operator)
        VehicleDriverProfile.objects.create(
            plate_number="55 ع 333 ایران 63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="محمود کاظمی",
            driver_national_code="0010376488",
            driver_phone="09121112233",
            is_driver_owner=False,
            owner_name="جواد احمدی",
            owner_national_code="0078901235",
            owner_phone="09124445566",
            default_service_rate=Decimal("2500000"),
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )

        res = self.client.get(f"/api/personnel/vehicles/export-excel/?section_id={self.section_a.id}")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, "ناوگان و رانندگان")

        # بررسی سطرهای هدر
        headers_fa = [cell.value for cell in ws[1]]
        headers_en = [cell.value for cell in ws[2]]

        self.assertIn("مالک شخص راننده است", headers_fa)
        self.assertIn("نام مالک", headers_fa)
        self.assertIn("کد ملی مالک", headers_fa)
        self.assertIn("شماره همراه مالک", headers_fa)
        self.assertIn("نوع مالکیت", headers_fa)

        self.assertIn("is_driver_owner", headers_en)
        self.assertIn("owner_name", headers_en)
        self.assertIn("owner_national_code", headers_en)
        self.assertIn("owner_phone", headers_en)
        self.assertIn("ownership_type_display", headers_en)

        # بررسی مقادیر ردیف داده
        data_row = [cell.value for cell in ws[3]]
        self.assertIn("55 ع 333 ایران 63", data_row)
        self.assertIn("خیر", data_row)
        self.assertIn("جواد احمدی", data_row)
        self.assertIn("0078901235", data_row)
        self.assertIn("09124445566", data_row)

    def test_submit_action_and_patch_compatibility_bridge(self):
        """
        تسک ۲.۷، ب-۱۳، ۲.۵ و ب-۱۷:
        ۱. اکشن صریح POST /submit/ در صورت نقص فیلدها ۴۰۰ با missing_fields برمی‌گرداند.
        ۲. پل سازگاری PATCH: کلاینت با PATCH وضعیت pending_supervisor می‌فرستد؛
           اگر ناقص باشد، داده ذخیره می‌شود اما وضعیت draft می‌ماند و submitted=False است.
           با رفع نواقص، به pending_supervisor می‌رود.
        """
        self.client.force_authenticate(user=self.operator)
        v = VehicleDriverProfile.objects.create(
            plate_number="15الف234ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="محسن رضایی",
            section=self.section_a,
            approval_status="draft",
            created_by=self.operator,
            is_active=True
        )

        # ۱. اکشن صریح submit بدون کدملی و موبایل معتبر و بیمه
        res_submit = self.client.post(f"/api/personnel/vehicles/{v.id}/submit/")
        self.assertEqual(res_submit.status_code, 400)
        self.assertIn("missing_fields", res_submit.data)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "draft")

        # ۲. پل سازگاری PATCH در حالت ناقص
        res_patch_incomplete = self.client.patch(
            f"/api/personnel/vehicles/{v.id}/",
            {"approval_status": "pending_supervisor", "driver_phone": "09121112233"},
            format="json"
        )
        self.assertEqual(res_patch_incomplete.status_code, 200)
        self.assertFalse(res_patch_incomplete.data.get("submitted", False))
        self.assertEqual(res_patch_incomplete.data["approval_status"], "draft")
        v.refresh_from_db()
        self.assertEqual(v.driver_phone, "09121112233")
        self.assertEqual(v.approval_status, "draft")

        # ۳. تکمیل فیلدهای الزامی مرحله submit (کدملی معتبر، تاریخ بیمه)
        v.driver_national_code = "0010376488"
        v.insurance_expiry_date = "1406/01/01"
        v.save()

        # اکنون فراخوانی submit باید موفق باشد
        res_submit_ok = self.client.post(f"/api/personnel/vehicles/{v.id}/submit/")
        self.assertEqual(res_submit_ok.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "pending_supervisor")

    def test_accountant_stage_gate_and_direct_edit(self):
        """
        تسک ۲.۵، ب-۱۷، ۳.۴ و ۳.۷:
        ۱. در مرحله pending_accountant، حسابدار نرخ و حساب بانکی را تعیین می‌کند.
        ۲. تایید حسابدار بدون نرخ یا بدون حساب شبا رد می‌شود (گیت تجمعی).
        ۳. ویرایش مستقیم حسابدار لاگ accountant_edit ثبت می‌کند.
        ۴. در مرحله pending_manager اطلاعات برای کاربر عادی قفل است.
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="25ب345ایران63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="جواد کریمی",
            driver_national_code="0010376488",
            driver_phone="09125556677",
            insurance_expiry_date="1406/05/01",
            section=self.section_a,
            approval_status="pending_accountant",
            is_active=True
        )

        self.client.force_authenticate(user=self.accountant)
        # تلاش برای تایید مالی بدون تعیین نرخ و حساب بانکی
        res_fail = self.client.post(f"/api/personnel/vehicles/{v.id}/approve-finance/")
        self.assertEqual(res_fail.status_code, 400)
        self.assertIn("missing_fields", res_fail.data)

        # ویرایش مستقیم حسابدار: ثبت نرخ و شماره شبا
        valid_sheba = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        res_edit = self.client.patch(
            f"/api/personnel/vehicles/{v.id}/",
            {"default_service_rate": 5500000, "sheba_number": valid_sheba, "bank_name": "بانک ملی ایران"},
            format="json"
        )
        self.assertEqual(res_edit.status_code, 200)

        # بررسی لاگ accountant_edit
        from personnel.models import WorkflowAuditLog
        log_entry = WorkflowAuditLog.objects.filter(object_id=v.id, action="accountant_edit").first()
        self.assertIsNotNone(log_entry)
        self.assertEqual(log_entry.actor_id, self.accountant.id)

        # اکنون تایید مالی موفقیت‌آمیز است
        res_fin_ok = self.client.post(f"/api/personnel/vehicles/{v.id}/approve-finance/")
        self.assertEqual(res_fin_ok.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "pending_manager")

        # در pending_manager، ویرایش توسط کاربر غیرسوپریوزر قفل است (۴۰۰)
        res_locked = self.client.patch(
            f"/api/personnel/vehicles/{v.id}/",
            {"driver_name": "تغییر نام در قفل"},
            format="json"
        )
        self.assertEqual(res_locked.status_code, 400)

    def test_workflow_state_machine_transitions(self):
        """
        تسک ۳.۱ تا ۳.۶: تست کامل عودت (request-revision)، بازگشایی (reopen)،
        قطع همکاری (end-collaboration)، و رد با Purge مشروط.
        """
        # ۱. عودت توسط سرپرست
        v = VehicleDriverProfile.objects.create(
            plate_number="35ج456ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="علی اکبری",
            section=self.section_a,
            approval_status="pending_supervisor",
            is_active=True
        )
        self.client.force_authenticate(user=self.supervisor)
        res_rev = self.client.post(f"/api/personnel/vehicles/{v.id}/request-revision/", {"reason": "نقص مدارک فنی"}, format="json")
        self.assertEqual(res_rev.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "revision_required")
        self.assertEqual(v.rejection_reason, "نقص مدارک فنی")

        # ۲. رد پرونده هرگز تصویب‌نشده و فاقد سابقه -> Purge مشروط (حذف قطعی از دیتابیس)
        v_purge = VehicleDriverProfile.objects.create(
            plate_number="45د567ایران63",
            vehicle_type="pickup",
            ownership_type="contract",
            driver_name="فرهاد صادقی",
            section=self.section_a,
            approval_status="pending_supervisor",
            is_active=True
        )
        purge_id = v_purge.id
        res_purge = self.client.post(f"/api/personnel/vehicles/{purge_id}/reject/", {"reason": "خودرو شرایط ندارد"}, format="json")
        self.assertEqual(res_purge.status_code, 200)
        self.assertTrue(res_purge.data.get("purged", False))
        self.assertFalse(VehicleDriverProfile.objects.filter(id=purge_id).exists())

        # ۳. رد پرونده دارای سابقه -> rejected بدون حذف فیزیکی
        v_with_hist = VehicleDriverProfile.objects.create(
            plate_number="55س678ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="حسین علوی",
            section=self.section_a,
            approval_status="pending_supervisor",
            is_active=True
        )
        VehicleTripLog.objects.create(
            vehicle=v_with_hist,
            date_shamsi="1404/01/01",
            trip_count=1,
            unit_rate=3000000,
            total_amount=3000000
        )
        res_rej = self.client.post(f"/api/personnel/vehicles/{v_with_hist.id}/reject/", {"reason": "رد به دلیل تخلف"}, format="json")
        self.assertEqual(res_rej.status_code, 200)
        self.assertFalse(res_rej.data.get("purged", False))
        v_with_hist.refresh_from_db()
        self.assertEqual(v_with_hist.approval_status, "rejected")

        # ۴. بازگشایی پرونده ردشده (reopen) -> وضعیت draft
        res_reopen = self.client.post(f"/api/personnel/vehicles/{v_with_hist.id}/reopen/")
        self.assertEqual(res_reopen.status_code, 200)
        v_with_hist.refresh_from_db()
        self.assertEqual(v_with_hist.approval_status, "draft")

        # ۵. قطع همکاری (end-collaboration) روی خودروی مصوب
        v_active = VehicleDriverProfile.objects.create(
            plate_number="65ص789ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="قاسم نوری",
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )
        VehicleTripLog.objects.create(
            vehicle=v_active,
            date_shamsi="1405/06/15",
            trip_count=2,
            unit_rate=4000000,
            total_amount=8000000
        )
        self.client.force_authenticate(user=self.manager)
        # تاریخ پایان قبل از آخرین تردد -> خطای ۴۰۰
        res_end_bad = self.client.post(
            f"/api/personnel/vehicles/{v_active.id}/end-collaboration/",
            {"end_date": "1405/06/10", "reason": "اتمام قرارداد"},
            format="json"
        )
        self.assertEqual(res_end_bad.status_code, 400)

        # تاریخ پایان بعد از آخرین تردد -> موفق
        res_end_ok = self.client.post(
            f"/api/personnel/vehicles/{v_active.id}/end-collaboration/",
            {"end_date": "1405/06/20", "reason": "اتمام قرارداد"},
            format="json"
        )
        self.assertEqual(res_end_ok.status_code, 200)
        v_active.refresh_from_db()
        self.assertFalse(v_active.is_active)
        self.assertEqual(v_active.end_date, "1405/06/20")

    def test_trip_guard_in_attendance_paths(self):
        """
        تسک ۲.۱ و ب-۲۰: تست جامع گارد ثبت تردد در مسیرهای ثبت تردد:
        - خودروی غیرفعال یا تصویب‌نشده
        - پایان همکاری (تردد بعد از end_date ممنوع، تا خود end_date مجاز)
        - انقضای بیمه‌نامه
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="75ط890ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="محمود تهرانی",
            section=self.section_a,
            approval_status="approved",
            is_active=True,
            end_date="1405/06/10",
            insurance_expiry_date="1405/07/01",
            default_service_rate=Decimal("4000000")
        )

        self.client.force_authenticate(user=self.operator)

        # ۱. ثبت در روز پایان همکاری (1405/06/10) -> مجاز است
        payload_ok = {
            "warehouse_id": None,
            "section_id": self.section_a.id,
            "date_shamsi": "1405/06/10",
            "items": [{"vehicle_id": v.id, "trip_count": 2, "unit_rate": 4000000}]
        }
        res_ok = self.client.post("/api/personnel/trips/bulk-save/", payload_ok, format="json")
        self.assertEqual(res_ok.status_code, 200)

        # ۲. ثبت بعد از تاریخ پایان همکاری (1405/06/11) -> مسدود (۴۰۰)
        payload_after_end = {
            "warehouse_id": None,
            "section_id": self.section_a.id,
            "date_shamsi": "1405/06/11",
            "items": [{"vehicle_id": v.id, "trip_count": 1, "unit_rate": 4000000}]
        }
        res_after_end = self.client.post("/api/personnel/trips/bulk-save/", payload_after_end, format="json")
        self.assertEqual(res_after_end.status_code, 400)
        self.assertIn("پایان یافته", res_after_end.data.get("error", ""))

        # ۳. خودروی با بیمه منقضی‌شده
        v_expired_ins = VehicleDriverProfile.objects.create(
            plate_number="85ظ901ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="کریم حسنی",
            section=self.section_a,
            approval_status="approved",
            is_active=True,
            insurance_expiry_date="1405/05/01",
            default_service_rate=Decimal("4000000")
        )
        payload_expired_ins = {
            "warehouse_id": None,
            "section_id": self.section_a.id,
            "date_shamsi": "1405/06/01",
            "items": [{"vehicle_id": v_expired_ins.id, "trip_count": 1, "unit_rate": 4000000}]
        }
        res_exp = self.client.post("/api/personnel/trips/bulk-save/", payload_expired_ins, format="json")
        self.assertEqual(res_exp.status_code, 400)
        self.assertIn("بیمه‌نامه شخص ثالث", res_exp.data.get("error", ""))

        # ۴. تست مسیر update_day_trip
        res_day_trip = self.client.post(
            "/api/personnel/trips/update-day-trip/",
            {
                "vehicle_id": v_expired_ins.id,
                "date_shamsi": "1405/06/01",
                "trip_count": 1,
                "section_id": self.section_a.id
            },
            format="json"
        )
        self.assertEqual(res_day_trip.status_code, 400)

    def test_vehicle_change_request_security_and_operations(self):
        """
        تسک ۳.۵ و ب-۲۳: تست ایمن‌سازی کامل VehicleChangeRequestViewSet:
        - create, update, partial_update بازگرداننده ۴۰۵ باشند.
        - destroy برای غیرسوپریوزر ۴۰۳ باشد.
        - اکشن cancel فقط توسط ثبت‌کننده در وضعیت‌های اولیه مجاز باشد.
        - SoD در تصویب مدیر بدون بازنویسی فیلدهای سیستمی.
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="95ع012ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="سعید احمدی",
            driver_national_code="0010376488",
            driver_phone="09121113355",
            section=self.section_a,
            approval_status="approved",
            is_active=True,
            default_service_rate=Decimal("5000000")
        )

        cr = VehicleChangeRequest.objects.create(
            vehicle=v,
            requested_by=self.operator,
            status="pending_supervisor",
            proposed_changes={"driver_phone": "09129998877"}
        )
        v.has_pending_changes = True
        v.save(update_fields=["has_pending_changes"])

        # ۱. متدهای مسدودشده (۴۰۵)
        self.client.force_authenticate(user=self.operator)
        res_post = self.client.post("/api/personnel/vehicle-change-requests/", {"vehicle": v.id})
        self.assertEqual(res_post.status_code, 405)

        res_put = self.client.put(f"/api/personnel/vehicle-change-requests/{cr.id}/", {"status": "approved"})
        self.assertEqual(res_put.status_code, 405)

        res_patch = self.client.patch(f"/api/personnel/vehicle-change-requests/{cr.id}/", {"status": "approved"})
        self.assertEqual(res_patch.status_code, 405)

        # ۲. تلاش برای حذف توسط کاربر غیرسوپریوزر -> ۴۰۳
        res_del_forbidden = self.client.delete(f"/api/personnel/vehicle-change-requests/{cr.id}/")
        self.assertEqual(res_del_forbidden.status_code, 403)

        # ۳. لغو (cancel) توسط کاربر غیر ثبت‌کننده در همان بخش -> ۴۰۳
        self.client.force_authenticate(user=self.supervisor)
        res_cancel_bad = self.client.post(f"/api/personnel/vehicle-change-requests/{cr.id}/cancel/")
        self.assertEqual(res_cancel_bad.status_code, 403)

        # ۴. لغو (cancel) توسط ثبت‌کننده در وضعیت pending_supervisor -> موفق
        self.client.force_authenticate(user=self.operator)
        res_cancel_ok = self.client.post(f"/api/personnel/vehicle-change-requests/{cr.id}/cancel/")
        self.assertEqual(res_cancel_ok.status_code, 200)
        cr.refresh_from_db()
        self.assertEqual(cr.status, "cancelled")
        v.refresh_from_db()
        self.assertFalse(v.has_pending_changes)

    def test_nested_bank_accounts_api_and_sod(self):
        """
        تسک ۵.۱ و ب-۱۵ و DEF-10: تست اکشن تودرتوی حساب‌های بانکی ناوگان (bank-accounts):
        - لیست‌گیری GET
        - ایجاد حساب POST و همگام‌سازی خودکار فیلدهای اصلی خودرو
        - انحصار حساب اصلی و جابجایی خودکار is_primary
        - ویرایش PUT/PATCH حساب
        - حذف DELETE و جانشینی خودکار حساب اصلی
        - بررسی گارد SoD با check_vehicle_bank_account_edit_allowed در وضعیت تصویب‌شده
        """
        from personnel.models import VehicleBankAccount
        valid_sheba_1 = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        valid_sheba_2 = generate_sheba_from_account("بانک ملت", "1234567890")

        v = VehicleDriverProfile.objects.create(
            plate_number="18الف999ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="قاسم سلیمانی",
            driver_national_code="0010376488",
            driver_phone="09121112233",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="draft",
            is_active=True
        )

        self.client.force_authenticate(user=self.operator)

        # ۱. لیست‌گیری اولیه (خالی)
        res_list = self.client.get(f"/api/personnel/vehicles/{v.id}/bank-accounts/")
        self.assertEqual(res_list.status_code, 200)
        self.assertEqual(len(res_list.data), 0)

        # ۲. ایجاد حساب اول با is_primary=True
        payload_acc1 = {
            "bank_name": "بانک ملی ایران",
            "account_number": "0105678901004",
            "sheba_number": valid_sheba_1,
            "account_owner_name": "قاسم سلیمانی",
            "account_owner_national_code": "0010376488",
            "payee_type": "driver",
            "is_primary": True,
            "is_active": True
        }
        res_create1 = self.client.post(f"/api/personnel/vehicles/{v.id}/bank-accounts/", payload_acc1, format="json")
        self.assertEqual(res_create1.status_code, 201)
        acc1_id = res_create1.data["id"]
        v.refresh_from_db()
        self.assertEqual(v.sheba_number, valid_sheba_1)
        self.assertEqual(v.bank_name, "بانک ملی ایران")

        # ۳. ایجاد حساب دوم با is_primary=True -> حساب اول باید is_primary=False شود
        payload_acc2 = {
            "bank_name": "بانک ملت",
            "account_number": "1234567890",
            "sheba_number": valid_sheba_2,
            "account_owner_name": "قاسم سلیمانی",
            "account_owner_national_code": "0010376488",
            "payee_type": "driver",
            "is_primary": True,
            "is_active": True
        }
        res_create2 = self.client.post(f"/api/personnel/vehicles/{v.id}/bank-accounts/", payload_acc2, format="json")
        self.assertEqual(res_create2.status_code, 201)
        acc2_id = res_create2.data["id"]

        acc1 = VehicleBankAccount.objects.get(id=acc1_id)
        acc2 = VehicleBankAccount.objects.get(id=acc2_id)
        self.assertFalse(acc1.is_primary)
        self.assertTrue(acc2.is_primary)
        v.refresh_from_db()
        self.assertEqual(v.sheba_number, valid_sheba_2)

        # ۴. حذف حساب دوم -> حساب اول باید دوباره حساب اصلی شود
        res_del = self.client.delete(f"/api/personnel/vehicles/{v.id}/bank-accounts/?account_id={acc2_id}")
        self.assertEqual(res_del.status_code, 200)
        self.assertFalse(VehicleBankAccount.objects.filter(id=acc2_id).exists())
        acc1.refresh_from_db()
        self.assertTrue(acc1.is_primary)
        v.refresh_from_db()
        self.assertEqual(v.sheba_number, valid_sheba_1)

        # ۵. تست گارد SoD: در وضعیت pending_accountant اپراتور عادی بدون دسترسی مالی نباید بتواند حساب بانکی اضافه کند
        v.approval_status = "pending_accountant"
        v.save(update_fields=["approval_status"])
        res_forbidden = self.client.post(f"/api/personnel/vehicles/{v.id}/bank-accounts/", payload_acc2, format="json")
        self.assertEqual(res_forbidden.status_code, 400)
        self.assertIn("مجاز به تغییر حساب‌های بانکی", res_forbidden.data.get("error", ""))

        # اما با دسترسی حسابدار در وضعیت pending_accountant باید مجاز باشد
        self.client.force_authenticate(user=self.accountant)
        res_accountant_ok = self.client.post(f"/api/personnel/vehicles/{v.id}/bank-accounts/", payload_acc2, format="json")
        self.assertEqual(res_accountant_ok.status_code, 201)

        # و در وضعیت approved حساب‌های بانکی قفل است و نیازمند درخواست تغییرات است
        v.approval_status = "approved"
        v.save(update_fields=["approval_status"])
        res_approved_locked = self.client.post(f"/api/personnel/vehicles/{v.id}/bank-accounts/", payload_acc1, format="json")
        self.assertEqual(res_approved_locked.status_code, 400)
        self.assertIn("ویرایش یا تغییر حساب‌های بانکی برای پرونده مصوب مجاز نیست", res_approved_locked.data.get("error", ""))

    def test_nested_documents_api_and_physical_file_preservation(self):
        """
        تسک ۵.۲ و ب-۳۰: تست اکشن مدارک ناوگان (documents):
        - اعتبارسنجی و آپلود فایل مدرک
        - لیست‌گیری مدارک
        - حذف ردیف مدرک از دیتابیس با حفظ فایل فیزیکی روی دیسک (audit safety)
        """
        import os
        from django.core.files.uploadedfile import SimpleUploadedFile
        from personnel.models import VehicleDocument

        v = VehicleDriverProfile.objects.create(
            plate_number="21ج777ایران63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="جواد رضویان",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="draft",
            is_active=True
        )

        self.client.force_authenticate(user=self.operator)

        # ۱. آپلود مدرک
        fake_file = SimpleUploadedFile("insurance_doc.pdf", b"%PDF-1.4 test document content for fleet", content_type="application/pdf")
        res_upload = self.client.post(
            f"/api/personnel/vehicles/{v.id}/documents/",
            {
                "file": fake_file,
                "doc_type": "insurance",
                "title": "بیمه‌نامه ثالث خاور",
                "issue_date": "1404/01/01",
                "expiry_date": "1405/01/01"
            },
            format="multipart"
        )
        self.assertEqual(res_upload.status_code, 201)
        doc_id = res_upload.data["id"]
        doc = VehicleDocument.objects.get(id=doc_id)
        self.assertEqual(doc.vehicle_id, v.id)
        self.assertEqual(doc.title, "بیمه‌نامه ثالث خاور")
        file_disk_path = doc.file.path
        self.assertTrue(os.path.exists(file_disk_path))

        # ۲. لیست‌گیری مدارک
        res_list = self.client.get(f"/api/personnel/vehicles/{v.id}/documents/")
        self.assertEqual(res_list.status_code, 200)
        self.assertEqual(len(res_list.data), 1)

        try:
            # ۳. حذف مدرک: رکورد از دیتابیس پاک شود اما فایل فیزیکی روی دیسک باقی بماند (ب-۳۰)
            res_delete = self.client.delete(f"/api/personnel/vehicles/{v.id}/documents/?doc_id={doc_id}")
            self.assertEqual(res_delete.status_code, 200)
            self.assertFalse(VehicleDocument.objects.filter(id=doc_id).exists())
            self.assertTrue(os.path.exists(file_disk_path), "فایل فیزیکی مدرک نباید از دیسک حذف شود (Audit Preservation)")
        finally:
            # تمیزکاری دیسک تست
            if os.path.exists(file_disk_path):
                os.remove(file_disk_path)

    def test_delete_impact_api_and_contract(self):
        """
        تسک ۵.۵ و ب-۱۰: تست استعلام اثرات حذف ناوگان (delete_impact):
        - مطابقت با قرارداد smart-delete-modal
        - مجاز بودن حذف پیش‌نویس بدون تردد
        - مسدودسازی در صورت وجود تردد با ارائه گزارش تاریخچه
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="33ب444ایران63",
            vehicle_type="sedan",
            ownership_type="contract",
            driver_name="فرهاد مجیدی",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="draft",
            is_active=True
        )

        self.client.force_authenticate(user=self.operator)

        # ۱. وضعیت بدون تردد -> can_delete=True
        res_impact = self.client.get(f"/api/personnel/vehicles/{v.id}/delete_impact/")
        self.assertEqual(res_impact.status_code, 200)
        self.assertTrue(res_impact.data["can_delete"])
        self.assertEqual(len(res_impact.data["blocking_reasons"]), 0)
        self.assertEqual(res_impact.data["history"]["details"]["total_trips_count"], 0)

        # ۲. ثبت یک تردد برای خودرو
        VehicleTripLog.objects.create(
            vehicle=v,
            section=self.section_a,
            date_shamsi="1405/01/10",
            trip_count=1,
            unit_rate=Decimal("1000000"),
            total_amount=Decimal("1000000")
        )

        # ۳. مجدداً استعلام -> can_delete=False و هشدار مسدودکننده
        res_blocked = self.client.get(f"/api/personnel/vehicles/{v.id}/delete_impact/")
        self.assertEqual(res_blocked.status_code, 200)
        self.assertFalse(res_blocked.data["can_delete"])
        self.assertGreater(len(res_blocked.data["blocking_reasons"]), 0)
        self.assertEqual(res_blocked.data["history"]["details"]["total_trips_count"], 1)

    def test_export_excel_company_isolation_and_rate_sod(self):
        """
        تسک ۵.۳ و ب-۱۲: تست خروجی اکسل ۲ سطری ناوگان:
        - حذف ستون نرخ برای کاربر غیرمالی (SoD)
        - وجود ستون نرخ برای نقش مالی
        - ایزولاسیون شرکت‌ها (محدود به readable_company_ids)
        """
        import io
        import openpyxl
        from personnel.models import Company

        # خودروی شرکت دالان (مجاز برای کاربر ما)
        VehicleDriverProfile.objects.create(
            plate_number="77ص111ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده دالان",
            section=self.section_a,
            company=self.company,
            default_service_rate=Decimal("3500000"),
            approval_status="approved",
            is_active=True
        )

        # شرکت و خودروی دیگر (نامربوط)
        other_co = Company.objects.create(code="OTHER-CO", name="شرکت دیگر")
        other_proj = FinancialProject.objects.create(code="2", name="پروژه دیگر", company=other_co, is_active=True)
        other_sec = ProjectSection.objects.create(project=other_proj, code="21", name="بخش دیگر", is_active=True)
        VehicleDriverProfile.objects.create(
            plate_number="88ص222ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده شرکت دیگر",
            section=other_sec,
            company=other_co,
            default_service_rate=Decimal("9000000"),
            approval_status="approved",
            is_active=True
        )

        # ۱. کاربر اپراتور (غیرمالی): نباید نرخ را ببیند
        self.client.force_authenticate(user=self.operator)
        res_op = self.client.get("/api/personnel/vehicles/export-excel/")
        self.assertEqual(res_op.status_code, 200)
        wb_op = openpyxl.load_workbook(io.BytesIO(res_op.content))
        ws_op = wb_op.active
        headers_fa_op = [ws_op.cell(row=1, column=c).value for c in range(1, ws_op.max_column + 1)]
        headers_en_op = [ws_op.cell(row=2, column=c).value for c in range(1, ws_op.max_column + 1)]
        self.assertNotIn("نرخ پایه سرویس (ریال)", headers_fa_op)
        self.assertNotIn("default_service_rate", headers_en_op)

        # ۲. کاربر حسابدار (مالی): باید ستون نرخ را ببیند
        self.client.force_authenticate(user=self.accountant)
        res_acc = self.client.get("/api/personnel/vehicles/export-excel/")
        self.assertEqual(res_acc.status_code, 200)
        wb_acc = openpyxl.load_workbook(io.BytesIO(res_acc.content))
        ws_acc = wb_acc.active
        headers_fa_acc = [ws_acc.cell(row=1, column=c).value for c in range(1, ws_acc.max_column + 1)]
        headers_en_acc = [ws_acc.cell(row=2, column=c).value for c in range(1, ws_acc.max_column + 1)]
        self.assertIn("نرخ پایه سرویس (ریال)", headers_fa_acc)
        self.assertIn("default_service_rate", headers_en_acc)

    def test_import_excel_national_codes_and_auto_pass(self):
        """
        تسک ۵.۴ و ب-۵: تست درون‌ریزی اکسل ناوگان با اعتبارسنجی کد ملی،
        ایجاد حساب در VehicleBankAccount و عدم استفاده از manager_approved منسوخ.
        """
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        from personnel.models import VehicleBankAccount

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ناوگان"
        ws.append(["ردیف", "شماره پلاک", "نام راننده", "کد ملی راننده", "تلفن", "نوع خودرو", "مالکیت", "نرخ", "شبا", "بانک", "حساب"])
        ws.append(["row_num", "plate_number", "driver_name", "driver_national_code", "driver_phone", "vehicle_type", "ownership_type", "default_service_rate", "sheba_number", "bank_name", "account_number"])
        valid_sheba = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        # ردیف معتبر
        ws.append([1, "14الف123ایران63", "سعید اکبری", "0010376488", "09121112233", "وانت نیسان", "استیجاری", 3000000, valid_sheba, "بانک ملی ایران", "0105678901004"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        excel_file = SimpleUploadedFile("fleet_import.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        self.client.force_authenticate(user=self.operator)
        res_import = self.client.post(
            "/api/personnel/vehicles/import-excel/",
            {"file": excel_file, "section_id": self.section_a.id, "dry_run": "false"},
            format="multipart"
        )
        self.assertEqual(res_import.status_code, 200)
        self.assertEqual(res_import.data["created_count"], 1)

        v = VehicleDriverProfile.objects.filter(plate_number="14 الف 123 ایران 63").first()
        self.assertIsNotNone(v)
        self.assertNotEqual(v.approval_status, "manager_approved")
        self.assertEqual(v.company_id, self.company.id)
        # حساب اصلی در VehicleBankAccount ایجاد شده باشد
        pri_acc = VehicleBankAccount.objects.filter(vehicle=v, is_primary=True).first()
        self.assertIsNotNone(pri_acc)
        self.assertEqual(pri_acc.sheba_number, valid_sheba)

    def test_workflow_notifications_reopen_and_deleted(self):
        """
        تسک ۴.۱ و ۴.۲: تست ارسال اعلان‌های جریان کار در گذارهای reopen و destroy
        """
        from notifications.models import Notification

        v = VehicleDriverProfile.objects.create(
            plate_number="41ق555ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="محسن رضایی",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="rejected",
            rejection_reason="نقص مدارک",
            is_active=False
        )

        self.client.force_authenticate(user=self.operator)

        # ۱. بازگشایی پرونده
        with self.captureOnCommitCallbacks(execute=True):
            res_reopen = self.client.post(f"/api/personnel/vehicles/{v.id}/reopen/", {"reason": "مدارک تکمیل شد"}, format="json")
        self.assertEqual(res_reopen.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "draft")

        # بررسی ایجاد اعلان برای ناظران/سرپرستان با مرحله reopened
        notif_reopened = Notification.objects.filter(target_model="VehicleDriverProfile", target_id=str(v.id)).first()
        self.assertIsNotNone(notif_reopened)
        self.assertIn("بازگشایی", notif_reopened.title)

        # ۲. حذف خودروی پیش‌نویس بدون تردد توسط مدیر
        self.client.force_authenticate(user=self.manager)
        notif_count_before = Notification.objects.count()
        with self.captureOnCommitCallbacks(execute=True):
            res_del = self.client.delete(f"/api/personnel/vehicles/{v.id}/")
        self.assertIn(res_del.status_code, [200, 204])
        notif_count_after = Notification.objects.count()
        self.assertGreater(notif_count_after, notif_count_before)

    def test_rate_benchmark_calculation(self):
        """
        تسک ۹.۳: تست محاسبه سروری شاخص و میانگین نرخ هم‌تیپ‌ها در پروژه
        """
        v1 = VehicleDriverProfile.objects.create(
            plate_number="11ع111ایران63",
            vehicle_type="truck",
            ownership_type="contract",
            driver_name="راننده یک",
            default_service_rate=Decimal("10000000"),
            section=self.section_a,
            company=self.company,
            approval_status="approved",
            is_active=True
        )
        v2 = VehicleDriverProfile.objects.create(
            plate_number="22ع222ایران63",
            vehicle_type="truck",
            ownership_type="contract",
            driver_name="راننده دو",
            default_service_rate=Decimal("20000000"),
            section=self.section_a,
            company=self.company,
            approval_status="approved",
            is_active=True
        )
        target = VehicleDriverProfile.objects.create(
            plate_number="33ع333ایران63",
            vehicle_type="truck",
            ownership_type="contract",
            driver_name="راننده هدف",
            default_service_rate=Decimal("15000000"),
            section=self.section_a,
            company=self.company,
            approval_status="pending_accountant",
            is_active=True
        )

        self.client.force_authenticate(user=self.accountant)
        res = self.client.get(f"/api/personnel/vehicles/{target.id}/rate-benchmark/")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["count"], 3)
        self.assertEqual(res.data["avg_rate"], 15000000)
        self.assertEqual(res.data["min_rate"], Decimal("10000000"))
        self.assertEqual(res.data["max_rate"], Decimal("20000000"))
        self.assertEqual(res.data["current_rate"], Decimal("15000000"))

        # نقش غیرمالی (اپراتور یا سرپرست) باید خطای ۴۰۳ دریافت کند
        self.client.force_authenticate(user=self.operator)
        res_forbidden_op = self.client.get(f"/api/personnel/vehicles/{target.id}/rate-benchmark/")
        self.assertEqual(res_forbidden_op.status_code, 403)

        self.client.force_authenticate(user=self.supervisor)
        res_forbidden_sup = self.client.get(f"/api/personnel/vehicles/{target.id}/rate-benchmark/")
        self.assertEqual(res_forbidden_sup.status_code, 403)

    def test_dual_destination_revision_for_accountant(self):
        """
        تسک ۹.۵: تست عودت دومقصده حسابدار (به سرپرست یا کارمند)
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="77ط777ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده نیسان",
            section=self.section_a,
            company=self.company,
            approval_status="pending_accountant",
            is_active=True
        )

        self.client.force_authenticate(user=self.accountant)

        # ۱. عودت به سرپرست
        res_to_sup = self.client.post(
            f"/api/personnel/vehicles/{v.id}/request-revision/",
            {"reason": "بررسی فنی بیشتر توسط سرپرست", "target_stage": "supervisor"},
            format="json"
        )
        self.assertEqual(res_to_sup.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "pending_supervisor")

        # بازگرداندن به pending_accountant توسط سرپرست
        self.client.force_authenticate(user=self.supervisor)
        self.client.post(f"/api/personnel/vehicles/{v.id}/approve-supervisor/")
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "pending_accountant")

        # ۲. عودت به کارمند
        self.client.force_authenticate(user=self.accountant)
        res_to_emp = self.client.post(
            f"/api/personnel/vehicles/{v.id}/request-revision/",
            {"reason": "اصلاح مدارک هویتی توسط کارمند", "target_stage": "employee"},
            format="json"
        )
        self.assertEqual(res_to_emp.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "revision_required")

    def test_batch_approve_supervisor_and_finance(self):
        """
        تسک ۸.۴ و ۹.۶: تست جامع تایید گروهی سرپرست و حسابدار همراه با گیت فیلدهای مالی
        """
        valid_sheba = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        v1 = VehicleDriverProfile.objects.create(
            plate_number="55ع555ایران63",
            vehicle_type="truck",
            ownership_type="contract",
            driver_name="راننده دسته یک",
            driver_phone="09171112233",
            driver_national_code="0010376488",
            insurance_expiry_date="1406/01/01",
            default_service_rate=Decimal("15000000"),
            sheba_number=valid_sheba,
            bank_name="بانک ملی ایران",
            section=self.section_a,
            company=self.company,
            approval_status="pending_supervisor",
            is_active=True
        )
        v2 = VehicleDriverProfile.objects.create(
            plate_number="66ع666ایران63",
            vehicle_type="truck",
            ownership_type="contract",
            driver_name="راننده دسته دو",
            driver_phone="09174445566",
            driver_national_code="0010376488",
            insurance_expiry_date="1406/01/01",
            default_service_rate=Decimal("0"), # نرخ ناقص!
            sheba_number=valid_sheba,
            bank_name="بانک ملی ایران",
            section=self.section_a,
            company=self.company,
            approval_status="pending_supervisor",
            is_active=True
        )

        # ۱. تایید گروهی سرپرست (هر دو باید تایید و به مرحله حسابدار بروند)
        self.client.force_authenticate(user=self.supervisor)
        res_sup = self.client.post(
            "/api/personnel/vehicles/batch-approve/",
            {"stage": "supervisor", "ids": [v1.id, v2.id], "note": "تایید گروهی سرپرست"},
            format="json"
        )
        self.assertEqual(res_sup.status_code, 200)
        self.assertEqual(res_sup.data["success_count"], 2)
        self.assertEqual(res_sup.data["failed_count"], 0)
        v1.refresh_from_db()
        v2.refresh_from_db()
        self.assertEqual(v1.approval_status, "pending_accountant")
        self.assertEqual(v2.approval_status, "pending_accountant")

        # ۲. تایید گروهی مالی حسابدار (v1 به دلیل تکمیل بودن تایید می‌شود و v2 به دلیل نداشتن نرخ رد می‌شود)
        self.client.force_authenticate(user=self.accountant)
        res_fin = self.client.post(
            "/api/personnel/vehicles/batch-approve/",
            {"stage": "finance", "ids": [v1.id, v2.id], "note": "تایید گروهی مالی"},
            format="json"
        )
        self.assertEqual(res_fin.status_code, 200)
        self.assertEqual(res_fin.data["success_count"], 1)
        self.assertEqual(res_fin.data["failed_count"], 1)
        self.assertIn(v1.id, res_fin.data["success_ids"])

        v1.refresh_from_db()
        v2.refresh_from_db()
        self.assertEqual(v1.approval_status, "pending_manager")
        self.assertEqual(v2.approval_status, "pending_accountant")

    def test_vehicle_cost_estimate_endpoint_and_sod(self):
        """
        تسک ۱۰.۱ و ۱۰.۴: تست برآورد بار مالی سروری ناوگان و کنترل دسترسی تفکیک وظایف (SoD)
        """
        # خودروی استیجاری با نرخ روزانه ۱۵ میلیون ریال
        v_contract = VehicleDriverProfile.objects.create(
            plate_number="88ع888ایران63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="راننده استیجاری",
            default_service_rate=Decimal("15000000"),
            section=self.section_a,
            company=self.company,
            approval_status="pending_manager",
            is_active=True
        )

        # خودروی شرکتی (فاقد مبلغ کرایه/تسویه)
        v_company = VehicleDriverProfile.objects.create(
            plate_number="99ع999ایران63",
            vehicle_type="khavar",
            ownership_type="company",
            driver_name="راننده سازمانی",
            default_service_rate=Decimal("0"),
            section=self.section_a,
            company=self.company,
            approval_status="pending_manager",
            is_active=True
        )

        # ۱. دسترسی مدیر مجاز است
        self.client.force_authenticate(user=self.manager)
        res_mgr = self.client.get(f"/api/personnel/vehicles/{v_contract.id}/cost-estimate/")
        self.assertEqual(res_mgr.status_code, 200)
        self.assertTrue(res_mgr.data["is_payable"])
        self.assertEqual(res_mgr.data["daily_rate"], 15000000.0)
        self.assertEqual(res_mgr.data["standard_days"], 26)
        self.assertEqual(res_mgr.data["monthly_cost_standard"], 390000000) # 15,000,000 * 26
        self.assertEqual(res_mgr.data["yearly_cost_standard"], 4680000000)  # 390,000,000 * 12

        # خودروی شرکتی باید هزینه صفر برگرداند
        res_comp = self.client.get(f"/api/personnel/vehicles/{v_company.id}/cost-estimate/")
        self.assertEqual(res_comp.status_code, 200)
        self.assertFalse(res_comp.data["is_payable"])
        self.assertEqual(res_comp.data["monthly_cost_standard"], 0)
        self.assertEqual(res_comp.data["yearly_cost_standard"], 0)

        # ۲. دسترسی حسابدار نیز مجاز است
        self.client.force_authenticate(user=self.accountant)
        res_acc = self.client.get(f"/api/personnel/vehicles/{v_contract.id}/cost-estimate/")
        self.assertEqual(res_acc.status_code, 200)

        # ۳. کاربر فاقد دسترسی مالی/مدیریتی خطای ۴۰۳ می‌گیرد
        self.client.force_authenticate(user=self.operator)
        res_denied = self.client.get(f"/api/personnel/vehicles/{v_contract.id}/cost-estimate/")
        self.assertEqual(res_denied.status_code, 403)

    def test_manager_return_to_accountant_and_final_approval(self):
        """
        تسک ۱۰.۲ و ۱۰.۴: تست چرخه عودت پرونده به حسابدار توسط مدیر و تصویب نهایی
        """
        from personnel.models import WorkflowAuditLog
        valid_sheba = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        v = VehicleDriverProfile.objects.create(
            plate_number="34ب123ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده نیسان تست مدیر",
            driver_phone="09171234567",
            driver_national_code="0010376488",
            insurance_expiry_date="1406/05/01",
            default_service_rate=Decimal("12000000"),
            sheba_number=valid_sheba,
            bank_name="بانک ملی ایران",
            section=self.section_a,
            company=self.company,
            approval_status="pending_manager",
            is_active=True
        )

        # ۱. مدیر پرونده را جهت اصلاح نرخ به حسابدار عودت می‌دهد
        self.client.force_authenticate(user=self.manager)
        res_rev = self.client.post(
            f"/api/personnel/vehicles/{v.id}/request-revision/",
            {"reason": "نرخ نیازمند بازنگری و بررسی مجدد حسابداری است"},
            format="json"
        )
        self.assertEqual(res_rev.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "pending_accountant")

        # بررسی ثبت لاگ ممیزی گردش کار
        log_rev = WorkflowAuditLog.objects.filter(object_id=v.id, action="request_revision").first()
        self.assertIsNotNone(log_rev)
        self.assertEqual(log_rev.to_status, "pending_accountant")

        # ۲. حسابدار پس از بررسی مجدد، پرونده را به مدیر تایید مالی می‌کند
        self.client.force_authenticate(user=self.accountant)
        res_fin = self.client.post(f"/api/personnel/vehicles/{v.id}/approve-finance/")
        self.assertEqual(res_fin.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "pending_manager")

        # ۳. مدیر پرونده را با ثبت دستور مدیریتی تصویب نهایی می‌کند
        self.client.force_authenticate(user=self.manager)
        res_appr = self.client.post(
            f"/api/personnel/vehicles/{v.id}/approve-manager/",
            {"note": "تصویب شد؛ هماهنگی جهت شروع سرویس انجام شود."},
            format="json"
        )
        self.assertEqual(res_appr.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "approved")
        self.assertTrue(v.is_active)

    def test_manager_end_collaboration_and_trip_guard_enforcement(self):
        """
        تسک ۱۰.۳ و ۱۰.۴: تست قطع همکاری با ثبت end_date و مسدودسازی ثبت تردد پس از پایان همکاری
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="78ج456ایران63",
            vehicle_type="pickup",
            ownership_type="contract",
            driver_name="راننده وانت تست قطع همکاری",
            default_service_rate=Decimal("8000000"),
            section=self.section_a,
            company=self.company,
            approval_status="approved",
            is_active=True
        )

        # ۱. کاربر عادی نمی‌تواند قطع همکاری ثبت کند
        self.client.force_authenticate(user=self.operator)
        res_op = self.client.post(
            f"/api/personnel/vehicles/{v.id}/end-collaboration/",
            {"end_date": "1405/07/15", "reason": "درخواست راننده"},
            format="json"
        )
        self.assertEqual(res_op.status_code, 403)

        # ۲. مدیر قطع همکاری را برای تاریخ ۱۴۰۵/۰۷/۱۵ ثبت می‌کند
        self.client.force_authenticate(user=self.manager)
        res_mgr = self.client.post(
            f"/api/personnel/vehicles/{v.id}/end-collaboration/",
            {"end_date": "1405/07/15", "reason": "اتمام قرارداد و پایان پروژه"},
            format="json"
        )
        self.assertEqual(res_mgr.status_code, 200)
        v.refresh_from_db()
        self.assertFalse(v.is_active)
        self.assertEqual(v.end_date, "1405/07/15")

        # ۳. ثبت تردد تا تاریخ پایان همکاری مجاز است
        from personnel.fleet_lifecycle_services import assert_vehicle_trip_allowed, VehicleWorkflowError
        assert_vehicle_trip_allowed(v, "1405/07/15") # بدون خطا عبور می‌کند

        # ۴. تلاش برای ثبت تردد پس از تاریخ پایان همکاری مسدود است
        with self.assertRaises(VehicleWorkflowError) as ctx:
            assert_vehicle_trip_allowed(v, "1405/07/16")
        self.assertEqual(ctx.exception.code, 'trip_not_allowed')
        self.assertIn("پایان یافته", str(ctx.exception.message if hasattr(ctx.exception, 'message') else ctx.exception))

    def test_idor_and_company_scoping_in_vehicle_update(self):
        """
        تسک ۰.۱ و ۱۰.۴: تست پیشگیری از IDOR در ویرایش خودرو بین شرکت‌های مختلف
        """
        from personnel.models import Company
        other_company = Company.objects.create(code="OTHER-CO", name="شرکت دیگر", is_active=True)
        other_project = FinancialProject.objects.create(code="99", name="پروژه دیگر", company=other_company, is_active=True)
        other_section = ProjectSection.objects.create(project=other_project, code="991", name="کارگاه دیگر", is_active=True)

        v_other = VehicleDriverProfile.objects.create(
            plate_number="19د789ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده شرکت دیگر",
            section=other_section,
            company=other_company,
            approval_status="draft",
            is_active=True
        )

        # کاربر شرکت اول تلاش می‌کند خودروی شرکت دوم را ویرایش کند
        self.client.force_authenticate(user=self.operator)
        res = self.client.patch(
            f"/api/personnel/vehicles/{v_other.id}/",
            {"driver_name": "تلاش برای نفوذ IDOR"},
            format="json"
        )
        # به دلیل قلمرو شرکت، شیء یافت نمی‌شود (404 Not Found)
        self.assertEqual(res.status_code, 404)

    def test_vehicle_trip_log_clean_defense_in_depth(self):
        """
        تسک ۱۰.۷ و طرح اصلاحی بند ۶: آزمون متد clean مدل VehicleTripLog
        """
        from django.core.exceptions import ValidationError
        v = VehicleDriverProfile.objects.create(
            plate_number="55ب999ایران77",
            vehicle_type="pickup",
            driver_name="راننده منقضی",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=False,
            end_date="1405/06/01"
        )
        trip = VehicleTripLog(
            vehicle=v,
            warehouse_id=1,
            date_shamsi="1405/06/10",
            trip_count=2,
            unit_rate=5000000
        )
        with self.assertRaises(ValidationError) as ctx:
            trip.clean()
        self.assertIn('vehicle', ctx.exception.message_dict)

    def test_workflow_audit_log_serializer_sod_masking(self):
        """
        تسک ۱۰.۸ و طرح اصلاحی بند ۵: فیلتر نرخ خودرو در لاگ ممیزی برای کاربران غیرمالی
        """
        from personnel.models import WorkflowAuditLog
        from personnel.serializers import WorkflowAuditLogSerializer
        from rest_framework.test import APIRequestFactory

        log = WorkflowAuditLog.objects.create(
            content_type="vehicle",
            object_id="999",
            action="update",
            actor=self.accountant,
            metadata={
                "diff": {"default_service_rate": 6000000, "plate_number": "11ب222ایران33"},
                "old_diff": {"default_service_rate": 5000000},
                "snapshot": {
                    "id": 999,
                    "driver_name": "احمد رضایی",
                    "plate_number": "11ب222ایران33",
                    "vehicle_type": "nissan",
                    "default_service_rate": 6000000,
                    "bank_name": "بانک ملت"
                }
            }
        )

        factory = APIRequestFactory()

        # ۱. کاربر سرپرست/اپراتور (غیرمالی)
        req_op = factory.get('/')
        req_op.user = self.operator
        serializer_op = WorkflowAuditLogSerializer(log, context={'request': req_op})
        rep_op = serializer_op.data
        meta_op = rep_op['metadata']

        self.assertNotIn('default_service_rate', meta_op['diff'])
        self.assertIn('plate_number', meta_op['diff'])
        self.assertNotIn('default_service_rate', meta_op['old_diff'])
        # در اسنپ‌شات فقط کلیدهای خلاصه مجاز وجود دارند
        self.assertIn('driver_name', meta_op['snapshot'])
        self.assertIn('plate_number', meta_op['snapshot'])
        self.assertNotIn('default_service_rate', meta_op['snapshot'])
        self.assertNotIn('bank_name', meta_op['snapshot'])

        # ۲. کاربر حسابدار (مالی)
        req_acc = factory.get('/')
        req_acc.user = self.accountant
        serializer_acc = WorkflowAuditLogSerializer(log, context={'request': req_acc})
        rep_acc = serializer_acc.data
        meta_acc = rep_acc['metadata']

        self.assertIn('default_service_rate', meta_acc['diff'])
        self.assertIn('default_service_rate', meta_acc['old_diff'])
        self.assertIn('default_service_rate', meta_acc['snapshot'])

    def test_fleet_notification_stage_urls_remediation(self):
        """
        تسک ۱۰.۹ و طرح اصلاحی بند ۴: هدایت اعلان‌های بازگشایی و رد به employee-new-vehicle
        """
        from personnel.fleet_lifecycle_services import FLEET_STAGE_URLS
        self.assertEqual(FLEET_STAGE_URLS['revision_required'], '/app/finance/employee-new-vehicle')
        self.assertEqual(FLEET_STAGE_URLS['rejected'], '/app/finance/employee-new-vehicle')
        self.assertEqual(FLEET_STAGE_URLS['reopened'], '/app/finance/employee-new-vehicle')
        self.assertEqual(FLEET_STAGE_URLS['purged'], '/app/finance/employee-new-vehicle')
        self.assertEqual(FLEET_STAGE_URLS['deleted'], '/app/finance/employee-new-vehicle')
        self.assertEqual(FLEET_STAGE_URLS['approved'], '/app/finance/employee-fleet')

    def test_scope_vehicle_queryset_without_company_header(self):
        """
        تسک ۱۱.۱۷ (ن-۱۷) و ب-۲: اعتبارسنجی قلمرو شرکت در صورت عدم ارسال سربرگ X-Company-ID
        کاربر غیرسوپریوزر هرگز نباید خودروهای شرکت‌های دیگر را ببیند حتی اگر هدر شرکت ارسال نشود.
        """
        from personnel.models import Company
        company_other = Company.objects.create(code="CO-OTHER-SCOPE", name="شرکت بیگانه", is_active=True)
        proj_other = FinancialProject.objects.create(code="998", name="پروژه بیگانه", company=company_other, is_active=True)
        sec_other = ProjectSection.objects.create(project=proj_other, code="9981", name="کارگاه بیگانه", is_active=True)

        v_other = VehicleDriverProfile.objects.create(
            plate_number="88ج999ایران55",
            vehicle_type="nissan",
            driver_name="راننده شرکت دیگر",
            company=company_other,
            section=sec_other,
            approval_status="approved",
            is_active=True
        )

        # کاربر اپراتور متعلق به شرکت اول (DALAN-CO) است و بدون هدر درخواست می‌دهد
        self.client.force_authenticate(user=self.operator)
        res = self.client.get("/api/personnel/vehicles/")
        self.assertEqual(res.status_code, 200)
        plates = [item["plate_number"] for item in res.data]
        self.assertNotIn("88ج999ایران55", plates)

    def test_vehicle_trip_log_perform_create_and_update_allowed(self):
        """
        تسک ۱۱.۱۷ (ن-۱۷) و ب-۲۰: اعتبارسنجی گارد assert_vehicle_trip_allowed در perform_create و perform_update
        """
        # ۱. خودروی پیش‌نویس (تایید نشده)
        v_unapproved = VehicleDriverProfile.objects.create(
            plate_number="33ب111ایران44",
            vehicle_type="nissan",
            driver_name="راننده پیش‌نویس",
            company=self.company,
            section=self.section_a,
            approval_status="draft",
            is_active=True
        )
        self.client.force_authenticate(user=self.operator)
        trip_payload = {
            "vehicle": v_unapproved.id,
            "section": self.section_a.id,
            "warehouse_id": 1,
            "date_shamsi": "1405/07/15",
            "trip_count": 2,
            "unit_rate": 3000000
        }
        res_fail = self.client.post("/api/personnel/trips/", trip_payload, format="json")
        self.assertEqual(res_fail.status_code, 400)
        self.assertIn("تصویب", str(res_fail.data))

        # ۲. خودروی مصوب با تاریخ معتبر
        v_approved = VehicleDriverProfile.objects.create(
            plate_number="33ب222ایران44",
            vehicle_type="nissan",
            driver_name="راننده مصوب",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=True,
            default_service_rate=4000000
        )
        trip_payload["vehicle"] = v_approved.id
        res_ok = self.client.post("/api/personnel/trips/", trip_payload, format="json")
        self.assertEqual(res_ok.status_code, 201)
        trip_id = res_ok.data["id"]

        # ۳. ویرایش با تاریخ نامعتبر (انقضا یا پس از قطع همکاری)
        v_approved.is_active = False
        v_approved.end_date = "1405/07/10"
        v_approved.save()

        res_update_fail = self.client.patch(f"/api/personnel/trips/{trip_id}/", {"date_shamsi": "1405/07/15"}, format="json")
        self.assertEqual(res_update_fail.status_code, 400)

    def test_fleet_excel_trip_allowed_gate(self):
        """
        تسک ۱۱.۱۷ (ن-۱۷) و ب-۲۰: اعتبارسنجی گارد assert_vehicle_trip_allowed در درون‌ریزی تردد اکسل
        """
        import openpyxl
        import io
        from personnel.fleet_excel_engine import import_fleet_monthly_excel

        v_unapproved = VehicleDriverProfile.objects.create(
            plate_number="77ج333ایران88",
            vehicle_type="nissan",
            driver_name="راننده اکسل تایید نشده",
            company=self.company,
            section=self.section_a,
            approval_status="pending_supervisor",
            is_active=True
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Fleet_Monthly"
        # Row 1: Persian display titles
        ws.cell(row=1, column=1, value="ردیف")
        ws.cell(row=1, column=2, value="نام راننده")
        ws.cell(row=1, column=3, value="شماره پلاک")
        ws.cell(row=1, column=7, value="روز ۰۱")

        # Row 2: Field keys (standard 2-row architecture)
        ws.cell(row=2, column=1, value="row_num")
        ws.cell(row=2, column=2, value="driver_name")
        ws.cell(row=2, column=3, value="plate_number")
        ws.cell(row=2, column=7, value="day_01")

        # Row 3: Data
        ws.cell(row=3, column=1, value=1)
        ws.cell(row=3, column=2, value="راننده اکسل تایید نشده")
        ws.cell(row=3, column=3, value="77ج333ایران88")
        ws.cell(row=3, column=7, value=2) # ۲ تردد در روز اول

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        with self.assertRaises(ValueError) as ctx:
            import_fleet_monthly_excel(
                file_obj=buf,
                warehouse_id=1,
                year_month="1405/04",
                user=self.operator
            )
        self.assertIn("77ج333ایران88", str(ctx.exception))
        self.assertIn("تصویب", str(ctx.exception))

    def test_vehicle_trip_log_protect_on_delete(self):
        """
        تسک ۱۱.۱۷ (ن-۱۷) و ب-۹: اعتبارسنجی قید PROTECT مدل VehicleTripLog
        حذف مستقیم خودروی دارای سابقه تردد از سطح مدل باید ProtectedError پرتاب کند.
        """
        from django.db.models.deletion import ProtectedError
        v = VehicleDriverProfile.objects.create(
            plate_number="99الف111ایران22",
            vehicle_type="nissan",
            driver_name="راننده دارای تردد",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )
        VehicleTripLog.objects.create(
            vehicle=v,
            warehouse_id=1,
            date_shamsi="1405/07/01",
            trip_count=1,
            unit_rate=2000000
        )
        with self.assertRaises(ProtectedError):
            v.delete()

    def test_sod_rate_silent_drop_and_forbidden_on_approved(self):
        """
        تسک ۱۱.۱۷ (ن-۱۷)، ن-۳ و ط-۲: تفکیک وظایف نرخ؛ حذف بی‌صدا در برابر ۴۰۳ روی پرونده مصوب
        """
        # ۱. پرونده پیش‌نویس: ارسال نرخ توسط کاربر غیرمالی بی‌صدا نادیده گرفته می‌شود (بدون ۴۰۰ یا ۴۰۳)
        v_draft = VehicleDriverProfile.objects.create(
            plate_number="44د555ایران66",
            vehicle_type="nissan",
            driver_name="راننده پیش‌نویس نرخ",
            company=self.company,
            section=self.section_a,
            approval_status="draft",
            is_active=True,
            default_service_rate=0
        )
        self.client.force_authenticate(user=self.operator)
        res_draft = self.client.patch(f"/api/personnel/vehicles/{v_draft.id}/", {"default_service_rate": 8000000}, format="json")
        self.assertEqual(res_draft.status_code, 200)
        v_draft.refresh_from_db()
        self.assertEqual(v_draft.default_service_rate or 0, 0) # دست‌نخورده باقی می‌ماند

        # ۲. پرونده مصوب: تلاش صریح کاربر غیرمالی برای تغییر نرخ به مقدار متفاوت ناتهی باید ۴۰۳ بگیرد
        v_approved = VehicleDriverProfile.objects.create(
            plate_number="44د777ایران66",
            vehicle_type="nissan",
            driver_name="راننده مصوب نرخ",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=True,
            default_service_rate=5000000
        )
        res_app_forbidden = self.client.patch(f"/api/personnel/vehicles/{v_approved.id}/", {"default_service_rate": 7500000}, format="json")
        self.assertEqual(res_app_forbidden.status_code, 403)
        self.assertEqual(res_app_forbidden.data.get("code"), "forbidden_rate_change")

        # ۳. پرونده مصوب: اگر نرخ صفر یا تهی ارسال شود (فرم‌های قدیمی یا صف آفلاین)، نباید ۴۰۳ بگیرد و باید بی‌صدا نادیده گرفته شود
        res_app_silent = self.client.patch(f"/api/personnel/vehicles/{v_approved.id}/", {"default_service_rate": 0, "driver_phone": "09129998877"}, format="json")
        self.assertIn(res_app_silent.status_code, (200, 202))
        v_approved.refresh_from_db()
        self.assertEqual(v_approved.default_service_rate, Decimal('5000000')) # نرخ دست‌نخورده

    def test_excel_import_update_draft_vehicle_audit_log_and_tenant_isolation(self):
        """
        تست ن-۸، ن-۱۷ و ن-۵: درون‌ریزی اکسل برای به‌روزرسانی پیش‌نویس، ثبت diff در لاگ ممیزی و تفکیک قلمرو شرکت
        """
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        from personnel.models import WorkflowAuditLog, Company
        from personnel.fleet_lifecycle_services import normalize_plate

        norm_p1 = normalize_plate("55ج666ایران77")
        norm_p2 = normalize_plate("88س999ایران11")

        # خودروی پیش‌نویس در شرکت فعلی با شماره پلاک نرمال‌شده
        v = VehicleDriverProfile.objects.create(
            company=self.company,
            section=self.section_a,
            plate_number=norm_p1,
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده اولیه",
            driver_national_code="0010376488",
            driver_phone="09121112233",
            approval_status="draft"
        )

        # خودرویی با پلاک دیگر در شرکت مجزا
        other_co = Company.objects.create(code="CO-OTHER-TEST", name="شرکت دیگر")
        VehicleDriverProfile.objects.create(
            company=other_co,
            plate_number=norm_p2,
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده شرکت دیگر",
            driver_national_code="0010376488",
            approval_status="approved"
        )

        # ساخت شیت اکسل
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Fleet"
        ws.append(["ردیف", "شماره پلاک", "نام راننده", "کد ملی", "موبایل", "نوع خودرو", "مالکیت", "نرخ", "شبا", "بانک", "حساب"])
        ws.append(["row", "plate_number", "driver_name", "driver_national_code", "driver_phone", "vehicle_type", "ownership_type", "default_service_rate", "sheba_number", "bank_name", "account_number"])
        ws.append([1, "55ج666ایران77", "راننده به‌روز شده", "0010376488", "09129998877", "نیسان", "استیجاری", 4000000, "", "", ""])
        ws.append([2, "88س999ایران11", "راننده جدید شرکت خودمان", "0010376488", "09123334455", "نیسان", "استیجاری", 5000000, "", "", ""])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        excel_file = SimpleUploadedFile("vehicles.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        self.client.force_authenticate(user=self.manager)
        res = self.client.post(
            "/api/personnel/vehicles/import-excel/",
            {
                "file": excel_file,
                "section_id": self.section_a.id,
                "company_id": self.company.id,
                "update_existing": "true",
            },
            format="multipart"
        )
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.data.get("success"))
        self.assertEqual(res.data["summary"]["updated"], 1)
        self.assertEqual(res.data["summary"]["created"], 1)

        # ۱. بررسی به‌روزرسانی موفق خودروی پیش‌نویس بدون خطای ۵۰۰
        v.refresh_from_db()
        self.assertEqual(v.driver_name, "راننده به‌روز شده")
        self.assertEqual(v.driver_phone, "09129998877")

        # ۲. بررسی ایجاد لاگ ممیزی و وجود diff در metadata
        log = WorkflowAuditLog.objects.filter(content_type="vehicle", object_id=v.id, action="import_update").last()
        self.assertIsNotNone(log)
        self.assertIn("diff", log.metadata)
        self.assertIn("driver_name", log.metadata["diff"])
        self.assertEqual(log.metadata["diff"]["driver_name"]["old"], "راننده اولیه")
        self.assertEqual(log.metadata["diff"]["driver_name"]["new"], "راننده به‌روز شده")

        # ۳. بررسی تفکیک قلمرو شرکت (ثبت موفق پلاک مشابه شرکت دیگر در شرکت خودمان)
        new_v = VehicleDriverProfile.objects.filter(company=self.company, plate_number=norm_p2).first()
        self.assertIsNotNone(new_v)
        self.assertEqual(new_v.driver_name, "راننده جدید شرکت خودمان")

    def test_change_request_workflow_notifications(self):
        """
        تست ن-۱۰ و ن-۱۷: ارسال اعلان‌های مراحل میانی درخواست تغییرات به سرپرست، حسابدار و مدیر
        """
        from notifications.models import Notification
        from personnel.fleet_lifecycle_services import dispatch_vehicle_workflow_notification

        v_app = VehicleDriverProfile.objects.create(
            plate_number="33ب444ایران55",
            vehicle_type="nissan",
            driver_name="راننده تست اعلان",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )

        # ۱. ثبت درخواست تغییرات -> ارسال اعلان به سرپرست کارگاه
        Notification.objects.all().delete()
        with self.captureOnCommitCallbacks(execute=True):
            dispatch_vehicle_workflow_notification(v_app, stage="change_request_created", actor=self.operator)
        sup_notifs = Notification.objects.filter(recipient=self.supervisor, notification_type="fleet_change_request")
        self.assertTrue(sup_notifs.exists())
        # کاربر ایجادکننده نباید اعلان دریافت کند
        self.assertFalse(Notification.objects.filter(recipient=self.operator).exists())

        # ۲. تایید سرپرست کارگاه -> ارسال اعلان به حسابدار مالی
        Notification.objects.all().delete()
        with self.captureOnCommitCallbacks(execute=True):
            dispatch_vehicle_workflow_notification(v_app, stage="change_request_supervisor_approved", actor=self.supervisor)
        acc_notifs = Notification.objects.filter(recipient=self.accountant, notification_type="fleet_change_request")
        self.assertTrue(acc_notifs.exists())

        # ۳. تایید حسابدار مالی -> ارسال اعلان به مدیر
        Notification.objects.all().delete()
        with self.captureOnCommitCallbacks(execute=True):
            dispatch_vehicle_workflow_notification(v_app, stage="change_request_finance_approved", actor=self.accountant)
        mgr_notifs = Notification.objects.filter(recipient=self.manager, notification_type="fleet_change_request")
        self.assertTrue(mgr_notifs.exists())

    def test_n1_rejected_status_locked_against_patch_and_reopen(self):
        """
        تست ن-۱ و تسک ۱۱.۱: قفل پرونده rejected در برابر PATCH و بازگشایی انحصاری با reopen
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="11الف111ایران11",
            vehicle_type="nissan",
            driver_name="راننده رد شده",
            company=self.company,
            section=self.section_a,
            approval_status="rejected",
            rejection_reason="عدم انطباق مدارک",
            is_active=True
        )
        self.client.force_authenticate(user=self.operator)

        # ۱. تلاش برای ویرایش مستقیم با PATCH باید با ۴۰۰ و reopen_required مسدود شود
        res_patch = self.client.patch(
            f"/api/personnel/vehicles/{v.id}/",
            {"driver_name": "نام تغییر یافته"},
            format="json"
        )
        self.assertEqual(res_patch.status_code, 400)
        self.assertEqual(res_patch.data.get("code"), "reopen_required")

        # ۲. تلاش برای ارسال مستقیم به سرپرست نیز باید با ۴۰۰ مسدود شود
        res_submit = self.client.patch(
            f"/api/personnel/vehicles/{v.id}/",
            {"approval_status": "pending_supervisor"},
            format="json"
        )
        self.assertEqual(res_submit.status_code, 400)

        # ۳. بازگشایی پرونده با اکشن اختصاصی reopen
        res_reopen = self.client.post(f"/api/personnel/vehicles/{v.id}/reopen/", {"reason": "اصلاح مدارک"}, format="json")
        self.assertEqual(res_reopen.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "draft")

        # ۴. اکنون ویرایش با PATCH مجاز است
        res_patch_ok = self.client.patch(
            f"/api/personnel/vehicles/{v.id}/",
            {"driver_name": "نام اصلاح شده"},
            format="json"
        )
        self.assertEqual(res_patch_ok.status_code, 200)
        v.refresh_from_db()
        self.assertEqual(v.driver_name, "نام اصلاح شده")

    def test_n4_national_plate_validation_and_normalization(self):
        """
        تست ن-۴ و تسک ۱۱.۴: اعتبارسنجی رگکس پلاک ملی (پذیرش فرمت با دش و رد کلمات کامل به جای حرف)
        """
        from personnel.fleet_lifecycle_services import is_valid_national_plate, normalize_plate

        # ۱. فرمت‌های معتبر با خط فاصله
        self.assertTrue(is_valid_national_plate("12-ب-345-67"))
        self.assertTrue(is_valid_national_plate("12-ب-345-ایران-67"))
        self.assertTrue(is_valid_national_plate("12ب345ایران67"))
        self.assertTrue(is_valid_national_plate("12 الف 345 - 67"))

        # بررسی نرمال‌سازی
        self.assertEqual(normalize_plate("12-ب-345-67"), "12 ب 345 ایران 67")
        self.assertEqual(normalize_plate("12-الف-345-ایران-67"), "12 الف 345 ایران 67")

        # ۲. فرمت‌های نامعتبر: کلمات کامل به جای یک حرف (مانند «ایران» یا «تهران» در بخش حرف)
        self.assertFalse(is_valid_national_plate("12 ایران 345 67"))
        self.assertFalse(is_valid_national_plate("12تهران34567"))
        self.assertFalse(is_valid_national_plate("1234567"))

        # ۳. اعتبارسنجی در سطح API سریالایزر
        self.client.force_authenticate(user=self.operator)
        valid_sheba = generate_sheba_from_account("بانک صادرات ایران", "0201234567001")
        # تلاش با پلاک نامعتبر (کلمه به جای حرف)
        res_invalid = self.client.post("/api/personnel/vehicles/", {
            "plate_number": "12 ایران 345 67",
            "vehicle_type": "nissan",
            "driver_name": "تست پلاک",
            "section": self.section_a.id,
            "bank_name": "بانک صادرات ایران",
            "sheba_number": valid_sheba,
        }, format="json")
        self.assertEqual(res_invalid.status_code, 400)
        self.assertIn("plate_number", res_invalid.data)

        # ارسال با فرمت دش معتبر
        res_valid = self.client.post("/api/personnel/vehicles/", {
            "plate_number": "78-ج-901-23",
            "vehicle_type": "nissan",
            "driver_name": "تست پلاک معتبر",
            "section": self.section_a.id,
            "bank_name": "بانک صادرات ایران",
            "sheba_number": valid_sheba,
        }, format="json")
        self.assertEqual(res_valid.status_code, 201)
        self.assertEqual(res_valid.data["plate_number"], "78 ج 901 ایران 23")

    def test_n5_change_request_concurrent_conflict_409_and_supervisor_whitelist(self):
        """
        تست ن-۵ و تسک ۱۱.۵: خطای ۴۰۹ در صورت تغییر هم‌زمان وضعیت درخواست تغییرات و کنترل تایید سرپرست
        """
        v_app = VehicleDriverProfile.objects.create(
            plate_number="55ج666ایران77",
            vehicle_type="nissan",
            driver_name="راننده مصوب",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )
        cr = VehicleChangeRequest.objects.create(
            vehicle=v_app,
            requested_by=self.operator,
            proposed_changes={"driver_phone": "09121112233"},
            previous_values={"driver_phone": ""},
            status="pending_supervisor"
        )

        self.client.force_authenticate(user=self.supervisor)

        # ۱. اگر وضعیت قبل از تایید به دلیل کنش موازی تغییر کند (مثلاً لغو توسط کاربر)
        cr_concurrent = VehicleChangeRequest.objects.get(pk=cr.pk)
        cr_concurrent.status = "cancelled"
        cr_concurrent.save(update_fields=["status"])

        res_conflict = self.client.post(f"/api/personnel/vehicle-change-requests/{cr.id}/approve-supervisor/")
        self.assertEqual(res_conflict.status_code, 409)
        self.assertEqual(res_conflict.data.get("code"), "conflict")

        # ۲. تلاش برای تایید درخواست در وضعیت draft باید با ۴۰۹ مسدود شود
        cr_concurrent.status = "draft"
        cr_concurrent.save(update_fields=["status"])
        res_draft = self.client.post(f"/api/personnel/vehicle-change-requests/{cr.id}/approve-supervisor/")
        self.assertEqual(res_draft.status_code, 409)
        self.assertEqual(res_draft.data.get("code"), "conflict")

        # ۳. تلاش برای تایید درخواست در وضعیت revision_required باید با ۴۰۹ مسدود شود
        cr_concurrent.status = "revision_required"
        cr_concurrent.save(update_fields=["status"])
        res_rev = self.client.post(f"/api/personnel/vehicle-change-requests/{cr.id}/approve-supervisor/")
        self.assertEqual(res_rev.status_code, 409)
        self.assertEqual(res_rev.data.get("code"), "conflict")

        # ۴. اکنون با بازگرداندن وضعیت به pending_supervisor، تایید سرپرست با موفقیت انجام می‌شود
        cr_concurrent.status = "pending_supervisor"
        cr_concurrent.save(update_fields=["status"])
        res_ok = self.client.post(f"/api/personnel/vehicle-change-requests/{cr.id}/approve-supervisor/")
        self.assertEqual(res_ok.status_code, 200)
        cr.refresh_from_db()
        self.assertEqual(cr.status, "pending_accountant")
        self.assertEqual(cr.supervisor_reviewed_by, self.supervisor)

    def test_n6_end_collaboration_invariants(self):
        """
        تست ن-۶ و تسک ۱۱.۶: قطع همکاری خودرو، الزامات تاریخ، اعتبارسنجی نسبت به آخرین تردد و ممانعت از تکرار (۴۰۹)
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="88د999ایران00",
            vehicle_type="nissan",
            driver_name="راننده فعال",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )
        # ثبت یک تردد در تاریخ 1404/05/15
        VehicleTripLog.objects.create(
            vehicle=v,
            date_shamsi="1404/05/15",
            trip_count=2,
            is_deleted=False
        )

        self.client.force_authenticate(user=self.manager)

        # ۱. بدون تاریخ پایان همکاری -> ۴۰۰
        res_no_date = self.client.post(f"/api/personnel/vehicles/{v.id}/end-collaboration/", {"reason": "تسویه"}, format="json")
        self.assertEqual(res_no_date.status_code, 400)

        # ۲. تاریخ پایان قبل از آخرین تردد (1404/05/10 < 1404/05/15) -> ۴۰۰
        res_before_trip = self.client.post(f"/api/personnel/vehicles/{v.id}/end-collaboration/", {
            "end_date": "1404/05/10",
            "reason": "تسویه حساب"
        }, format="json")
        self.assertEqual(res_before_trip.status_code, 400)
        self.assertIn("قبل از تاریخ آخرین تردد", res_before_trip.data.get("error", ""))

        # ۳. تاریخ پایان معتبر (1404/05/20) -> موفقیت ۲۰۰
        res_success = self.client.post(f"/api/personnel/vehicles/{v.id}/end-collaboration/", {
            "end_date": "1404/05/20",
            "reason": "اتمام قرارداد پروژه"
        }, format="json")
        self.assertEqual(res_success.status_code, 200)
        v.refresh_from_db()
        self.assertFalse(v.is_active)
        self.assertEqual(v.end_date, "1404/05/20")

        # ۴. تلاش مجدد برای قطع همکاری خودرویی که همکاری‌اش پایان یافته -> ۴۰۹ conflict
        res_repeat = self.client.post(f"/api/personnel/vehicles/{v.id}/end-collaboration/", {
            "end_date": "1404/06/01",
            "reason": "تکرار"
        }, format="json")
        self.assertEqual(res_repeat.status_code, 409)
        self.assertEqual(res_repeat.data.get("code"), "already_ended")

    def test_n7_reject_purge_vs_archived_rejection(self):
        """
        تست ن-۷ و تسک ۱۱.۷: رفتار reject و تفکیک purge (حذف سخت بدون سابقه) از بایگانی با وضعیت rejected
        """
        self.client.force_authenticate(user=self.supervisor)

        # ۱. پرونده هرگز مصوب نشده و بدون سابقه در وضعیت pending_supervisor -> پاک‌سازی کامل (purge)
        v_clean = VehicleDriverProfile.objects.create(
            plate_number="21ط321ایران45",
            vehicle_type="nissan",
            driver_name="راننده بدون سابقه",
            company=self.company,
            section=self.section_a,
            approval_status="pending_supervisor",
            is_active=True
        )
        clean_id = v_clean.id
        res_purge = self.client.post(f"/api/personnel/vehicles/{clean_id}/reject/", {"reason": "عدم نیاز کارگاه"}, format="json")
        self.assertEqual(res_purge.status_code, 200)
        self.assertTrue(res_purge.data.get("purged"))
        self.assertFalse(VehicleDriverProfile.objects.filter(id=clean_id).exists())

        # ۲. پرونده دارای سابقه (تردد) -> حفظ رکورد و تبدیل وضعیت به rejected
        v_history = VehicleDriverProfile.objects.create(
            plate_number="22ط322ایران45",
            vehicle_type="nissan",
            driver_name="راننده دارای تردد",
            company=self.company,
            section=self.section_a,
            approval_status="pending_supervisor",
            is_active=True
        )
        VehicleTripLog.objects.create(
            vehicle=v_history,
            date_shamsi="1404/06/01",
            trip_count=1,
            is_deleted=False
        )
        hist_id = v_history.id
        res_reject = self.client.post(f"/api/personnel/vehicles/{hist_id}/reject/", {"reason": "تخلف راننده"}, format="json")
        self.assertEqual(res_reject.status_code, 200)
        self.assertFalse(res_reject.data.get("purged"))
        v_history.refresh_from_db()
        self.assertEqual(v_history.approval_status, "rejected")
        self.assertEqual(v_history.rejection_reason, "تخلف راننده")

        # ۳. پرونده بدون سابقه کارکرد اما در وضعیت supervisor_approved (خارج از purge_statuses)
        # طبق قانون، پاک‌سازی کامل (purge) فقط در وضعیت‌های در جریان مجاز است و در اینجا باید بایگانی rejected شود
        v_outside_purge = VehicleDriverProfile.objects.create(
            plate_number="23ط323ایران45",
            vehicle_type="nissan",
            driver_name="راننده تایید سرپرست",
            company=self.company,
            section=self.section_a,
            approval_status="supervisor_approved",
            is_active=True
        )
        self.client.force_authenticate(user=self.manager)
        res_reject_sup = self.client.post(f"/api/personnel/vehicles/{v_outside_purge.id}/reject/", {"reason": "رد پس از بررسی سرپرست"}, format="json")
        self.assertEqual(res_reject_sup.status_code, 200)
        self.assertFalse(res_reject_sup.data.get("purged"))
        v_outside_purge.refresh_from_db()
        self.assertEqual(v_outside_purge.approval_status, "rejected")
        self.assertEqual(v_outside_purge.rejection_reason, "رد پس از بررسی سرپرست")

    def test_n9_document_upload_normalization_and_audit_without_status_guard(self):
        """
        تست ن-۹ و تسک ۱۱.۹: بارگذاری مدارک بدون گارد وضعیت، نرمال‌سازی تاریخ‌های نامنفصل و ثبت لاگ ممیزی
        """
        from django.core.files.uploadedfile import SimpleUploadedFile
        from personnel.models import WorkflowAuditLog

        v = VehicleDriverProfile.objects.create(
            plate_number="33ی444ایران55",
            vehicle_type="nissan",
            driver_name="راننده تست مدارک",
            company=self.company,
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )
        self.client.force_authenticate(user=self.operator)

        dummy_file = SimpleUploadedFile("insurance.pdf", b"%PDF-1.4 test document content", content_type="application/pdf")
        # ارسال تاریخ ۸ رقمی بدون خط مورب جهت سنجش نرمال‌سازی
        res_doc = self.client.post(
            f"/api/personnel/vehicles/{v.id}/documents/",
            {
                "file": dummy_file,
                "doc_type": "insurance",
                "title": "بیمه‌نامه شخص ثالث",
                "issue_date": "14040101",
                "expiry_date": "14050101",
                "note": "پیوست بیمه‌نامه جدید"
            },
            format="multipart"
        )
        self.assertEqual(res_doc.status_code, 201)
        self.assertEqual(res_doc.data["issue_date"], "1404/01/01")
        self.assertEqual(res_doc.data["expiry_date"], "1405/01/01")

        # بررسی لاگ ممیزی
        audit = WorkflowAuditLog.objects.filter(content_type="vehicle", object_id=v.id, action="document_uploaded").last()
        self.assertIsNotNone(audit)
        self.assertIn("document", audit.metadata)
        self.assertEqual(audit.metadata["document"]["doc_type"], "insurance")
        self.assertEqual(audit.metadata["document"]["expiry_date"], "1405/01/01")

    def test_create_vehicle_auto_pass_response_status(self):
        """
        تست ن-۱۲: ایجاد خودرو توسط کاربر با نقش بالا (مدیر)، بازگشت وضعیت تصویب‌شده خودکار (approved) در پاسخ HTTP 201
        """
        valid_sheba = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        self.client.force_authenticate(user=self.manager)
        payload = {
            "plate_number": "99ج888ایران77",
            "vehicle_type": "nissan",
            "ownership_type": "contract",
            "driver_name": "احمد مدیر",
            "driver_national_code": "0010376488",
            "driver_phone": "09121234567",
            "default_service_rate": 5000000,
            "insurance_expiry_date": "1406/01/01",
            "bank_name": "بانک ملی ایران",
            "sheba_number": valid_sheba,
            "section": self.section_a.id,
        }
        res = self.client.post("/api/personnel/vehicles/", payload, format="json")
        self.assertEqual(res.status_code, 201)
        self.assertEqual(res.data["approval_status"], "approved")
        self.assertTrue(res.data["is_active"])


    def test_rejected_record_requires_reopen_and_bridge_limited_to_draft_states(self):
        """
        تسک ۱۱.۱ (ن-۱، ط-۱): پروندهٔ ردشده بدون بازگشایی ویرایش نمی‌شود؛
        پل PATCH فقط از draft و revision_required گذار می‌دهد و pending_supervisor قفل نیست
        """
        from personnel.models import WorkflowAuditLog

        # ۱. پروندهٔ ردشده: ویرایش کاربر غیرسوپریوزر ۴۰۰ با کد reopen_required می‌گیرد و چیزی ذخیره نمی‌شود
        v_rejected = VehicleDriverProfile.objects.create(
            plate_number="23ب456ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده ردشده",
            driver_phone="09121112233",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="rejected",
            rejection_reason="نقص مدارک",
            is_active=True
        )
        self.client.force_authenticate(user=self.operator)
        res_rejected = self.client.patch(
            f"/api/personnel/vehicles/{v_rejected.id}/",
            {"driver_phone": "09124445566", "approval_status": "pending_supervisor"},
            format="json"
        )
        self.assertEqual(res_rejected.status_code, 400)
        self.assertEqual(res_rejected.data.get("code"), "reopen_required")
        v_rejected.refresh_from_db()
        self.assertEqual(v_rejected.approval_status, "rejected")
        self.assertEqual(v_rejected.driver_phone, "09121112233")

        # ۲. پروندهٔ در انتظار سرپرست: ویرایش ذخیره می‌شود ولی پل PATCH گذار تازه‌ای انجام نمی‌دهد
        v_pending = VehicleDriverProfile.objects.create(
            plate_number="23ب457ایران63",
            vehicle_type="nissan",
            ownership_type="company",
            driver_name="راننده در انتظار سرپرست",
            driver_phone="09121112233",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="pending_supervisor",
            is_active=True
        )
        res_pending = self.client.patch(
            f"/api/personnel/vehicles/{v_pending.id}/",
            {"driver_phone": "09127778899", "approval_status": "pending_supervisor"},
            format="json"
        )
        self.assertEqual(res_pending.status_code, 200, res_pending.data)
        self.assertFalse(res_pending.data.get("submitted"))
        self.assertIn("مجاز نیست", res_pending.data.get("detail", ""))
        v_pending.refresh_from_db()
        self.assertEqual(v_pending.approval_status, "pending_supervisor")
        self.assertEqual(v_pending.driver_phone, "09127778899")
        self.assertFalse(
            WorkflowAuditLog.objects.filter(content_type="vehicle", object_id=v_pending.id, action="submit").exists()
        )

    def test_rate_benchmark_requires_fleet_rate_role(self):
        """
        تسک ۱۱.۲ (ن-۲): مقایسهٔ نرخ فقط برای نقش‌های مالی و مدیریت؛ کارمند و سرپرست ۴۰۳ می‌گیرند
        """
        target = VehicleDriverProfile.objects.create(
            plate_number="24ب456ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده شاخص نرخ",
            default_service_rate=Decimal("6000000"),
            section=self.section_a,
            company=self.company,
            approval_status="pending_accountant",
            is_active=True
        )
        for user in (self.operator, self.supervisor):
            self.client.force_authenticate(user=user)
            res = self.client.get(f"/api/personnel/vehicles/{target.id}/rate-benchmark/")
            self.assertEqual(res.status_code, 403, user.username)
            self.assertNotIn("avg_rate", res.data)

        self.client.force_authenticate(user=self.accountant)
        res_fin = self.client.get(f"/api/personnel/vehicles/{target.id}/rate-benchmark/")
        self.assertEqual(res_fin.status_code, 200)

    def test_national_plate_validation_for_road_vehicles(self):
        """
        تسک ۱۱.۴ (ن-۴، ب-۲۵): قالب سخت پلاک ملی فقط برای خودروهای جاده‌ای (گیت، سریالایزر و درون‌ریزی)؛
        شناسهٔ آزاد برای لیفتراک؛ مقدار ذخیره‌شدهٔ دست‌نخورده دوباره سنجیده نمی‌شود
        """
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        from personnel.fleet_lifecycle_services import is_valid_national_plate, get_vehicle_missing_fields

        # ۱. تابع اعتبارسنجی
        for plate in ("12ب345ایران63", "12 الف 345 ایران 63", "۱۲ب۳۴۵ایران۶۳"):
            self.assertTrue(is_valid_national_plate(plate), plate)
        for plate in ("", None, "ABC-123", "12AB345ایران63", "00ب345ایران63", "12ب000ایران63", "12ب345ایران00", "1ب345ایران63"):
            self.assertFalse(is_valid_national_plate(plate), plate)

        # ۲. سریالایزر: خودروی جاده‌ای با پلاک نامنطبق رد می‌شود، لیفتراک شناسهٔ آزاد می‌پذیرد
        self.client.force_authenticate(user=self.operator)
        base_payload = {"ownership_type": "company", "driver_name": "راننده تست پلاک", "section": self.section_a.id}
        res_bad = self.client.post("/api/personnel/vehicles/", {**base_payload, "plate_number": "ABC-123", "vehicle_type": "nissan"}, format="json")
        self.assertEqual(res_bad.status_code, 400)
        self.assertIn("plate_number", res_bad.data)
        res_lift = self.client.post("/api/personnel/vehicles/", {**base_payload, "plate_number": "LIFT-07", "vehicle_type": "forklift"}, format="json")
        self.assertEqual(res_lift.status_code, 201, res_lift.data)
        self.assertEqual(res_lift.data["plate_number"], "LIFT-07")

        # ۳. گیت: پلاک قدیمی نامنطبق برای خودروی جاده‌ای نقص است، برای لیفتراک نه
        legacy_road = VehicleDriverProfile.objects.create(
            plate_number="ABC-123",
            vehicle_type="nissan",
            ownership_type="company",
            driver_name="راننده پلاک قدیمی",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="draft",
            is_active=True
        )
        lift = VehicleDriverProfile.objects.get(pk=res_lift.data["id"])
        self.assertIn("plate_number", [m["field"] for m in get_vehicle_missing_fields(legacy_road, "draft")])
        self.assertNotIn("plate_number", [m["field"] for m in get_vehicle_missing_fields(lift, "draft")])

        # ۴. مقدار ذخیره‌شدهٔ دست‌نخورده دوباره سنجیده نمی‌شود
        res_same = self.client.patch(
            f"/api/personnel/vehicles/{legacy_road.id}/",
            {"plate_number": "ABC-123", "driver_name": "راننده پلاک قدیمی ویرایش‌شده"},
            format="json"
        )
        self.assertEqual(res_same.status_code, 200, res_same.data)
        legacy_road.refresh_from_db()
        self.assertEqual(legacy_road.driver_name, "راننده پلاک قدیمی ویرایش‌شده")

        # ۵. درون‌ریزی: پلاک نامنطبق و نوع ناشناخته خطای ردیف می‌گیرند و هیچ رکوردی ساخته نمی‌شود
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Fleet"
        ws.append(["ردیف", "شماره پلاک", "نام راننده", "کد ملی", "موبایل", "نوع خودرو", "مالکیت", "نرخ", "شبا", "بانک", "حساب"])
        ws.append(["row", "plate_number", "driver_name", "driver_national_code", "driver_phone", "vehicle_type", "ownership_type", "default_service_rate", "sheba_number", "bank_name", "account_number"])
        ws.append([1, "XYZ-999", "راننده پلاک نامعتبر اکسل", "0010376488", "09121234567", "نیسان", "استیجاری", 4000000, "", "", ""])
        ws.append([2, "45ب678ایران63", "راننده نوع ناشناخته اکسل", "0010376488", "09121234567", "هواپیما", "استیجاری", 4000000, "", "", ""])
        buf = io.BytesIO()
        wb.save(buf)
        excel_file = SimpleUploadedFile("vehicles.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        self.client.force_authenticate(user=self.manager)
        res_imp = self.client.post(
            "/api/personnel/vehicles/import-excel/",
            {"file": excel_file, "section_id": self.section_a.id, "company_id": self.company.id},
            format="multipart"
        )
        self.assertEqual(res_imp.status_code, 400)
        error_fields = [e["field"] for e in res_imp.data["errors"]]
        self.assertIn("شماره پلاک", error_fields)
        self.assertIn("نوع خودرو", error_fields)
        self.assertFalse(VehicleDriverProfile.objects.filter(driver_name__endswith="اکسل").exists())

    def test_change_request_supervisor_approval_whitelist_and_conflict(self):
        """
        تسک ۱۱.۵ (ن-۵): تأیید سرپرست درخواست تغییرات فقط در pending_supervisor؛
        وضعیت درون قفل دوباره سنجیده می‌شود و تغییر هم‌زمان ۴۰۹ می‌گیرد
        """
        from unittest import mock
        from personnel.views import VehicleChangeRequestViewSet

        v = VehicleDriverProfile.objects.create(
            plate_number="25ب456ایران63",
            vehicle_type="nissan",
            driver_name="راننده درخواست تغییرات",
            section=self.section_a,
            company=self.company,
            approval_status="approved",
            is_active=True
        )
        self.client.force_authenticate(user=self.supervisor)

        # ۱. وضعیت‌های draft و revision_required دیگر به تأیید سرپرست نمی‌رسند
        for st in ("draft", "revision_required"):
            cr = VehicleChangeRequest.objects.create(
                vehicle=v,
                requested_by=self.operator,
                status=st,
                proposed_changes={"driver_phone": "09129998877"}
            )
            res = self.client.post(f"/api/personnel/vehicle-change-requests/{cr.id}/approve-supervisor/")
            self.assertEqual(res.status_code, 409, st)
            self.assertEqual(res.data.get("code"), "conflict")
            cr.refresh_from_db()
            self.assertEqual(cr.status, st)
            self.assertIsNone(cr.supervisor_reviewed_by_id)

        # ۲. تغییر هم‌زمان: شیء خوانده‌شده هنوز pending_supervisor است، ولی ردیف دیتابیس پیش از قفل لغو شده است
        cr_live = VehicleChangeRequest.objects.create(
            vehicle=v,
            requested_by=self.operator,
            status="pending_supervisor",
            proposed_changes={"driver_phone": "09129998877"}
        )
        stale = VehicleChangeRequest.objects.get(pk=cr_live.pk)
        VehicleChangeRequest.objects.filter(pk=cr_live.pk).update(status="cancelled")
        with mock.patch.object(VehicleChangeRequestViewSet, "get_object", return_value=stale):
            res_conflict = self.client.post(f"/api/personnel/vehicle-change-requests/{cr_live.id}/approve-supervisor/")
        self.assertEqual(res_conflict.status_code, 409)
        self.assertEqual(res_conflict.data.get("code"), "conflict")
        cr_live.refresh_from_db()
        self.assertEqual(cr_live.status, "cancelled")
        self.assertIsNone(cr_live.supervisor_reviewed_by_id)

    def test_end_collaboration_reason_repeat_guard_and_date_based_last_trip(self):
        """
        تسک ۱۱.۶ (ن-۶): دلیل قطع همکاری الزامی است؛ آخرین تردد با شیء تاریخ پیدا می‌شود (نه مرتب‌سازی رشته‌ای)؛
        ثبت دوبارهٔ قطع همکاری ۴۰۹ می‌گیرد و end_date جابه‌جا نمی‌شود
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="26ب456ایران63",
            vehicle_type="pickup",
            ownership_type="contract",
            driver_name="راننده تست تاریخ تردد",
            default_service_rate=Decimal("8000000"),
            section=self.section_a,
            company=self.company,
            approval_status="approved",
            is_active=True
        )
        # در مرتب‌سازی رشته‌ای «1405/9/30» از «1405/10/01» بزرگ‌تر است؛ آخرین تردد واقعی 1405/10/01 است
        for ds in ("1405/9/30", "1405/10/01"):
            VehicleTripLog.objects.create(vehicle=v, date_shamsi=ds, trip_count=1, unit_rate=3000000, total_amount=3000000)

        self.client.force_authenticate(user=self.manager)
        url = f"/api/personnel/vehicles/{v.id}/end-collaboration/"

        # ۱. بدون دلیل
        res_no_reason = self.client.post(url, {"end_date": "1405/10/01"}, format="json")
        self.assertEqual(res_no_reason.status_code, 400)

        # ۲. تاریخ پایان پیش از آخرین تردد واقعی
        res_before_last = self.client.post(url, {"end_date": "1405/09/30", "reason": "اتمام قرارداد"}, format="json")
        self.assertEqual(res_before_last.status_code, 400)
        self.assertIn("1405/10/01", res_before_last.data.get("error", ""))
        v.refresh_from_db()
        self.assertTrue(v.is_active)
        self.assertIsNone(v.end_date)

        # ۳. ثبت معتبر
        res_ok = self.client.post(url, {"end_date": "1405/10/01", "reason": "اتمام قرارداد"}, format="json")
        self.assertEqual(res_ok.status_code, 200, res_ok.data)

        # ۴. تکرار قطع همکاری ۴۰۹ می‌گیرد و تاریخ پایان قبلی حفظ می‌شود
        res_repeat = self.client.post(url, {"end_date": "1405/10/05", "reason": "ثبت تکراری"}, format="json")
        self.assertEqual(res_repeat.status_code, 409)
        self.assertEqual(res_repeat.data.get("code"), "already_ended")
        v.refresh_from_db()
        self.assertFalse(v.is_active)
        self.assertEqual(v.end_date, "1405/10/01")

    def test_reject_purges_only_in_purge_statuses(self):
        """
        تسک ۱۱.۷ (ن-۷): پروندهٔ هرگز تصویب‌نشده و بی‌سابقه در وضعیت قدیمی supervisor_approved
        بیرون از purge_statuses است؛ پس رد می‌شود و Purge نمی‌شود
        """
        v = VehicleDriverProfile.objects.create(
            plate_number="27ب456ایران63",
            vehicle_type="nissan",
            ownership_type="contract",
            driver_name="راننده وضعیت قدیمی",
            section=self.section_a,
            company=self.company,
            approval_status="supervisor_approved",
            is_active=True
        )
        self.assertIsNone(v.manager_approved_at)

        self.client.force_authenticate(user=self.accountant)
        res = self.client.post(f"/api/personnel/vehicles/{v.id}/reject/", {"reason": "رد پروندهٔ قدیمی"}, format="json")
        self.assertEqual(res.status_code, 200, res.data)
        self.assertFalse(res.data.get("purged"))
        v.refresh_from_db()
        self.assertEqual(v.approval_status, "rejected")

    def test_document_upload_audit_log_and_date_normalization(self):
        """
        تسک ۱۱.۹ (ن-۹، ط-۶): بارگذاری مدرک لاگ document_uploaded با company_id دارد و تاریخ‌ها نرمال می‌شوند؛
        تاریخ نامعتبر ۴۰۰ می‌گیرد (نه ۵۰۰) و مدرکی ساخته نمی‌شود
        """
        import shutil
        import tempfile
        from django.core.files.uploadedfile import SimpleUploadedFile
        from personnel.models import VehicleDocument, WorkflowAuditLog

        media_root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media_root, ignore_errors=True)

        v = VehicleDriverProfile.objects.create(
            plate_number="28ب456ایران63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="راننده تست مدرک",
            section=self.section_a,
            company=self.company,
            created_by=self.operator,
            approval_status="draft",
            is_active=True
        )
        self.client.force_authenticate(user=self.operator)
        url = f"/api/personnel/vehicles/{v.id}/documents/"

        with self.settings(MEDIA_ROOT=media_root):
            res = self.client.post(
                url,
                {
                    "file": SimpleUploadedFile("insurance.pdf", b"%PDF-1.4 fleet document", content_type="application/pdf"),
                    "doc_type": "insurance",
                    "issue_date": "۱۴۰۵/۷/۱",
                    "expiry_date": "1406/7/1",
                },
                format="multipart"
            )
            self.assertEqual(res.status_code, 201, res.data)
            self.assertEqual(res.data["issue_date"], "1405/07/01")
            self.assertEqual(res.data["expiry_date"], "1406/07/01")

            log = WorkflowAuditLog.objects.filter(content_type="vehicle", object_id=v.id, action="document_uploaded").last()
            self.assertIsNotNone(log)
            self.assertEqual(log.metadata.get("company_id"), self.company.id)
            self.assertEqual(log.metadata["document"]["id"], res.data["id"])
            self.assertEqual(log.metadata["document"]["issue_date"], "1405/07/01")

            res_bad = self.client.post(
                url,
                {
                    "file": SimpleUploadedFile("insurance2.pdf", b"%PDF-1.4 fleet document", content_type="application/pdf"),
                    "doc_type": "insurance",
                    "expiry_date": "1405/13/40",
                },
                format="multipart"
            )
            self.assertEqual(res_bad.status_code, 400)
            self.assertIn("error", res_bad.data)
            self.assertEqual(VehicleDocument.objects.filter(vehicle=v).count(), 1)


class FleetDataMigrationReverseTests(TestCase):
    """
    تسک ۱۱.۱۱ (ن-۱۱): برگشت مایگریشن 0035 شرکت را فقط وقتی تهی می‌کند که هنوز همان مقدار forward باشد
    """

    def test_reverse_clears_only_forward_assigned_company(self):
        import importlib
        from django.apps import apps as global_apps
        from personnel.models import Company

        # شناسه‌ها ثابت‌اند، چون مایگریشن شناسه‌های عملیاتی را صریح نام می‌برد
        company_1, _ = Company.objects.get_or_create(id=1, defaults={"code": "MIG-C1", "name": "شرکت آزمون ۱"})
        company_2, _ = Company.objects.get_or_create(id=2, defaults={"code": "MIG-C2", "name": "شرکت آزمون ۲"})
        company_other = Company.objects.create(id=990001, code="MIG-C3", name="شرکت آزمون ۳")

        def make(pk, company, plate):
            return VehicleDriverProfile.objects.create(id=pk, company=company, plate_number=plate, vehicle_type="nissan", driver_name=f"راننده {pk}")

        v144 = make(144, company_1, "11ب111ایران11")            # مقدار forward → تهی می‌شود
        v2 = make(2, company_2, "22ب222ایران22")                # مقدار forward → تهی می‌شود
        v3 = make(3, company_other, "33ب333ایران33")            # پس از forward عوض شده → دست‌نخورده
        v4 = make(4, company_1, "44ب444ایران44")                # شرکت ۱ ولی forward آن شرکت ۲ بود → دست‌نخورده
        v145 = make(145, company_1, "55ب555ایران55")            # بیرون از فهرست مایگریشن → دست‌نخورده

        migration = importlib.import_module("personnel.migrations.0035_fleet_data_migration")
        migration.reverse_fleet_data(global_apps, None)

        for v in (v144, v2, v3, v4, v145):
            v.refresh_from_db()
        self.assertIsNone(v144.company_id)
        self.assertIsNone(v2.company_id)
        self.assertEqual(v3.company_id, company_other.id)
        self.assertEqual(v4.company_id, company_1.id)
        self.assertEqual(v145.company_id, company_1.id)



