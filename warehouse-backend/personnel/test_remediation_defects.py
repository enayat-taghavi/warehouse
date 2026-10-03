"""
آزمون‌های رگرسیون برنامهٔ اصلاح نقص‌های چرخهٔ ثبت تا تصویب پرسنل
(Documents/Personnel_Registration_To_Approval_Workflow/remediation_plan_defects.md).
هر آزمون یکی از نقص‌های DEF-xx را با payload واقعی فرانت می‌سنجد؛ آزمون‌های کاوشی پیوست ۸ با انتظار وارونه.
"""
import shutil
import tempfile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from personnel.models import (
    Company,
    FinancialProject,
    PersonnelBankAccount,
    PersonnelChangeRequest,
    PersonnelDocument,
    PersonnelProfile,
    ProjectSection,
    UserCompanyAccess,
    UserSectionAssignment,
    WorkflowAuditLog,
)

User = get_user_model()
SHEBA1 = 'IR062960000000100324200001'
SHEBA2 = 'IR820540102680020817909002'


_TEST_MEDIA_ROOT = tempfile.mkdtemp(prefix='rm_media_')


# فایل‌های آزمون نباید در MEDIA_ROOT واقعی (مدیای عملیاتی همین ماشین) نوشته شوند
@override_settings(MEDIA_ROOT=_TEST_MEDIA_ROOT)
class RemediationTestBase(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(_TEST_MEDIA_ROOT, ignore_errors=True)

    def setUp(self):
        self.client = APIClient()
        self.c1 = Company.objects.create(name='شرکت الف', code='RM-A')
        self.c2 = Company.objects.create(name='شرکت ب', code='RM-B')
        self.p1 = FinancialProject.objects.create(company=self.c1, name='پروژه الف', code='RPA')
        self.p2 = FinancialProject.objects.create(company=self.c2, name='پروژه ب', code='RPB')
        self.s1 = ProjectSection.objects.create(project=self.p1, name='بخش الف', code='RSA')
        self.s2 = ProjectSection.objects.create(project=self.p2, name='بخش ب', code='RSB')
        self.emp = User.objects.create_user(username='rm_emp', password='p')
        self.sup = User.objects.create_user(username='rm_sup', password='p')
        self.acc = User.objects.create_user(username='rm_acc', password='p')
        self.mgr = User.objects.create_user(username='rm_mgr', password='p')
        ct, _ = ContentType.objects.get_or_create(app_label='accounts', model='customuser')

        def grant(u, code):
            perm = Permission.objects.filter(codename=code).first() or Permission.objects.create(codename=code, name=code, content_type=ct)
            u.user_permissions.add(perm)

        grant(self.sup, 'perm_approve_personnel_supervisor')
        grant(self.acc, 'perm_approve_personnel_finance')
        grant(self.mgr, 'perm_approve_personnel_manager')
        grant(self.mgr, 'perm_delete_personnel_nomovement')
        for u in (self.emp, self.sup, self.acc, self.mgr):
            UserSectionAssignment.objects.create(user=u, section=self.s1, is_active=True)
            UserCompanyAccess.objects.create(user=u, company=self.c1, access_level='workspace_full')

    def as_user(self, user):
        self.client.force_authenticate(user)

    def complete_profile(self, nc, status='draft', section=None, company=None, **extra):
        """پروندهٔ کامل (گیت ارسال و گیت حسابدار هر دو پاس می‌شوند)."""
        data = dict(
            first_name='علی', last_name='رضایی', national_code=nc,
            section=section or self.s1, company=company or self.c1,
            father_name='حسن', birth_date='1370/01/01', id_number='12', gender='مرد', job_title='انباردار',
            phone_number='09121112233', address='تهران، خیابان آزادی', postal_code='1998877665',
            sheba_number=SHEBA1, bank_name='ملت', job_code='000255', start_date='1405/01/01',
            daily_base_wage=7000000, approval_status=status, created_by=self.emp,
        )
        data.update(extra)
        return PersonnelProfile.objects.create(**data)

    def employee_form_payload(self, nc, target_status='draft', **overrides):
        """همان بدنه‌ای که employee-new-personnel.ts:savePersonnel در حالت JSON می‌فرستد."""
        payload = {
            'first_name': 'سارا', 'last_name': 'احمدی', 'national_code': nc,
            'father_name': 'محمد', 'gender': 'زن', 'id_number': '345', 'birth_date': '1372/02/03',
            'birth_place': 'تهران', 'issue_place': 'تهران', 'job_title': 'کارگر انبار',
            'contract_type': 'daily', 'marital_status': 'single', 'children_count': 0,
            'phone_number': '09121234567', 'address': 'تهران، خیابان انقلاب', 'postal_code': '1234567890',
            'bank_name': 'ملت', 'account_number': '', 'sheba_number': SHEBA1,
            'bank_accounts': [{
                'bank_name': 'ملت', 'account_number': '', 'sheba_number': SHEBA1,
                'card_number': '', 'is_primary': True, 'is_active': True,
            }],
            'daily_base_wage': 0, 'section': self.s1.id, 'company': self.c1.id,
            'approval_status': target_status, 'is_active': True,
        }
        payload.update(overrides)
        return payload


class Step1CriticalBackendTests(RemediationTestBase):
    # ── DEF-01: ثبت با حساب بانکی و تعیین حساب اصلی ──
    def test_def01_create_with_flat_and_nested_bank_creates_single_primary(self):
        self.as_user(self.emp)
        r = self.client.post('/api/personnel/profiles/', self.employee_form_payload('0076543218', 'pending_supervisor'), format='json')
        self.assertEqual(r.status_code, 201, r.data)
        p = PersonnelProfile.objects.get(pk=r.data['id'])
        accounts = list(PersonnelBankAccount.objects.filter(personnel=p))
        self.assertEqual(len(accounts), 1, 'فقط یک حساب باید ساخته شود (نویسندهٔ واحد)')
        self.assertTrue(accounts[0].is_primary)
        self.assertEqual(p.sheba_number, SHEBA1)
        self.assertEqual(r.data['sheba_number'], SHEBA1, 'پاسخ ایجاد باید آینهٔ حساب اصلی را داشته باشد')

    def test_def01_create_with_two_nested_accounts_keeps_requested_primary(self):
        self.as_user(self.emp)
        payload = self.employee_form_payload('0076543218', bank_accounts=[
            {'bank_name': 'ملت', 'sheba_number': SHEBA1, 'is_primary': False, 'is_active': True},
            {'bank_name': 'صادرات', 'sheba_number': SHEBA2, 'is_primary': True, 'is_active': True},
        ])
        r = self.client.post('/api/personnel/profiles/', payload, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        p = PersonnelProfile.objects.get(pk=r.data['id'])
        self.assertEqual(PersonnelBankAccount.objects.filter(personnel=p).count(), 2)
        self.assertEqual(PersonnelBankAccount.objects.get(personnel=p, is_primary=True).sheba_number, SHEBA2)
        self.assertEqual(p.sheba_number, SHEBA2)

    def test_def01_second_primary_and_set_primary_switch_cleanly(self):
        p = self.complete_profile('0098765434')
        self.as_user(self.emp)
        r = self.client.post(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'bank_name': 'صادرات', 'sheba_number': SHEBA2, 'is_primary': True}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(PersonnelBankAccount.objects.filter(personnel=p, is_primary=True).count(), 1)
        p.refresh_from_db()
        self.assertEqual(p.sheba_number, SHEBA2)
        old = PersonnelBankAccount.objects.get(personnel=p, sheba_number=SHEBA1)
        r2 = self.client.put(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'id': old.id, 'is_primary': True}, format='json')
        self.assertEqual(r2.status_code, 200, r2.data)
        self.assertEqual(PersonnelBankAccount.objects.filter(personnel=p, is_primary=True).count(), 1)
        p.refresh_from_db()
        self.assertEqual(p.sheba_number, SHEBA1)

    def test_def01_unsetting_primary_directly_is_rejected(self):
        p = self.complete_profile('0098765434')
        acc = PersonnelBankAccount.objects.get(personnel=p, is_primary=True)
        self.as_user(self.emp)
        r = self.client.patch(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'id': acc.id, 'is_primary': False}, format='json')
        self.assertEqual(r.status_code, 400)
        acc.refresh_from_db()
        self.assertTrue(acc.is_primary)

    def test_def01_delete_by_query_id_promotes_remaining_account(self):
        p = self.complete_profile('0098765434')
        self.as_user(self.emp)
        self.client.post(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'bank_name': 'صادرات', 'sheba_number': SHEBA2, 'is_primary': False}, format='json')
        primary = PersonnelBankAccount.objects.get(personnel=p, is_primary=True)
        # personnel-api.service.ts شناسه را با ?id= می‌فرستد
        r = self.client.delete(f'/api/personnel/profiles/{p.id}/bank-accounts/?id={primary.id}')
        self.assertEqual(r.status_code, 200, r.data)
        remaining = PersonnelBankAccount.objects.get(personnel=p)
        self.assertTrue(remaining.is_primary)
        p.refresh_from_db()
        self.assertEqual(p.sheba_number, SHEBA2)

    # ── DEF-10: جابه‌جایی حساب و گارد وضعیت/نقش ──
    def test_def10_account_cannot_be_moved_to_other_personnel(self):
        p = self.complete_profile('0098765434')
        victim = self.complete_profile('0012345679', status='approved', section=self.s2, company=self.c2)
        acc = PersonnelBankAccount.objects.get(personnel=p)
        self.as_user(self.emp)
        r = self.client.put(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'id': acc.id, 'personnel': victim.id}, format='json')
        self.assertIn(r.status_code, (200, 400))
        acc.refresh_from_db()
        self.assertEqual(acc.personnel_id, p.id)
        self.assertFalse(PersonnelBankAccount.objects.filter(personnel=victim).exclude(sheba_number=SHEBA1).exists())

    def test_def10_bank_account_status_and_role_guard(self):
        p = self.complete_profile('0098765434', status='pending_accountant')
        self.as_user(self.emp)
        r = self.client.post(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'bank_name': 'صادرات', 'sheba_number': SHEBA2}, format='json')
        self.assertEqual(r.status_code, 400, 'کارمند در pending_accountant نباید حساب را تغییر دهد')
        self.as_user(self.acc)
        r = self.client.post(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'bank_name': 'صادرات', 'sheba_number': SHEBA2}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertTrue(WorkflowAuditLog.objects.filter(object_id=p.id, action='accountant_edit').exists())
        for locked in ('pending_manager', 'manager_approved', 'approved'):
            p.approval_status = locked
            p.save(update_fields=['approval_status'])
            for user in (self.acc, self.mgr):
                self.as_user(user)
                r = self.client.post(f'/api/personnel/profiles/{p.id}/bank-accounts/', {'bank_name': 'تجارت', 'sheba_number': SHEBA2}, format='json')
                self.assertEqual(r.status_code, 400, f'{locked} باید برای {user.username} قفل باشد')

    # ── DEF-08: ویرایش با بدنهٔ کامل شامل bank_accounts ──
    def test_def08_full_body_patch_by_employee_and_accountant(self):
        p = self.complete_profile('0061122335')
        self.as_user(self.emp)
        body = dict(self.client.get(f'/api/personnel/profiles/{p.id}/').data)
        self.assertEqual(len(body['bank_accounts']), 1)
        body['phone_number'] = '09125554433'
        r = self.client.patch(f'/api/personnel/profiles/{p.id}/', body, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        p.refresh_from_db()
        self.assertEqual(p.phone_number, '09125554433')
        self.assertEqual(PersonnelBankAccount.objects.filter(personnel=p).count(), 1)

        p.approval_status = 'pending_accountant'
        p.save(update_fields=['approval_status'])
        self.as_user(self.acc)
        body2 = dict(self.client.get(f'/api/personnel/profiles/{p.id}/').data)
        body2['note'] = 'اصلاح گروه شغلی'
        body2['job_grade'] = '18'
        r2 = self.client.patch(f'/api/personnel/profiles/{p.id}/', body2, format='json')
        self.assertEqual(r2.status_code, 200, r2.data)
        p.refresh_from_db()
        self.assertEqual(p.job_grade, '18')

    # ── DEF-02: دور زدن گیت در ایجاد ──
    def test_def02_create_always_draft_then_submit_gate(self):
        self.as_user(self.emp)
        r = self.client.post('/api/personnel/profiles/', {
            'first_name': 'x', 'last_name': 'y', 'national_code': '0012345679', 'section': self.s1.id,
            'approval_status': 'pending_supervisor',
        }, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data['approval_status'], 'draft')
        self.assertTrue(r.data['missing_fields'])
        r2 = self.client.post(f"/api/personnel/profiles/{r.data['id']}/submit/", {}, format='json')
        self.assertEqual(r2.status_code, 400)
        self.assertTrue(r2.data['missing_fields'])

    def test_def02_complete_employee_form_draft_then_submit(self):
        self.as_user(self.emp)
        r = self.client.post('/api/personnel/profiles/', self.employee_form_payload('0076543218', 'pending_supervisor'), format='json')
        self.assertEqual(r.data['approval_status'], 'draft')
        self.assertEqual(r.data['missing_fields'], [])
        r2 = self.client.post(f"/api/personnel/profiles/{r.data['id']}/submit/", {}, format='json')
        self.assertEqual(r2.status_code, 200, r2.data)
        self.assertEqual(r2.data['approval_status'], 'pending_supervisor')

    def test_def02_legacy_patch_cannot_change_status(self):
        p = self.complete_profile('0010376488')
        self.as_user(self.emp)
        self.client.patch(f'/api/personnel/profiles/{p.id}/', {'approval_status': 'approved'}, format='json')
        p.refresh_from_db()
        self.assertEqual(p.approval_status, 'draft')

    # ── DEF-14 (بخش ایجاد): عبور هوشمند با گیت ──
    def test_def14_manager_minimal_create_stays_draft(self):
        self.as_user(self.mgr)
        r = self.client.post('/api/personnel/profiles/', {'first_name': 'x', 'last_name': 'y', 'national_code': '0012345679', 'section': self.s1.id}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data['approval_status'], 'draft')
        self.assertFalse(WorkflowAuditLog.objects.filter(object_id=r.data['id'], action='auto_pass').exists())

    def test_def14_manager_complete_create_auto_passes_with_log(self):
        self.as_user(self.mgr)
        payload = self.employee_form_payload('0076543218', 'pending_supervisor', job_code='000255', start_date='1405/01/01', daily_base_wage=7000000)
        r = self.client.post('/api/personnel/profiles/', payload, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data['approval_status'], 'approved')
        log = WorkflowAuditLog.objects.get(object_id=r.data['id'], action='auto_pass')
        self.assertEqual(log.content_type, 'personnel')
        self.assertEqual(log.metadata['passed_stages'], ['supervisor', 'accountant', 'manager'])

    def test_def14_supervisor_create_passes_only_submit_gate(self):
        self.as_user(self.sup)
        r = self.client.post('/api/personnel/profiles/', self.employee_form_payload('0076543218', 'pending_supervisor'), format='json')
        self.assertEqual(r.data['approval_status'], 'pending_accountant')
        self.assertTrue(r.data['missing_fields'], 'نواقص مرحلهٔ حسابدار باید برگردد')

    # ── DEF-09: تأیید گروهی ──
    def test_def09_batch_approve_is_scoped_to_user_companies(self):
        other = self.complete_profile('0088776654', status='pending_supervisor', section=self.s2, company=self.c2)
        self.as_user(self.sup)
        r = self.client.post('/api/personnel/profiles/batch-approve/', {'ids': [other.id], 'stage': 'supervisor'}, format='json', HTTP_X_COMPANY_ID=str(self.c1.id))
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.data['results'][0]['success'])
        other.refresh_from_db()
        self.assertEqual(other.approval_status, 'pending_supervisor')

    def test_def09_batch_manager_rejects_rejected_and_incomplete(self):
        rejected = PersonnelProfile.objects.create(first_name='a', last_name='b', national_code='0061122335', section=self.s1, company=self.c1, approval_status='rejected')
        incomplete = PersonnelProfile.objects.create(first_name='a', last_name='b', national_code='0010376488', section=self.s1, company=self.c1, approval_status='pending_accountant')
        self.as_user(self.mgr)
        r = self.client.post('/api/personnel/profiles/batch-approve/', {'ids': [rejected.id, incomplete.id], 'stage': 'manager'}, format='json')
        self.assertEqual(r.data['successful_count'], 0)
        self.assertTrue(r.data['results'][1].get('missing_fields'))
        rejected.refresh_from_db()
        incomplete.refresh_from_db()
        self.assertEqual(rejected.approval_status, 'rejected')
        self.assertEqual(incomplete.approval_status, 'pending_accountant')

    def test_def09_batch_and_single_share_rules_and_log(self):
        p = self.complete_profile('0098765434', status='pending_supervisor')
        self.as_user(self.sup)
        r = self.client.post('/api/personnel/profiles/batch-approve/', {'ids': [p.id], 'stage': 'supervisor', 'note': 'ok'}, format='json')
        self.assertEqual(r.data['successful_count'], 1, r.data)
        log = WorkflowAuditLog.objects.get(object_id=p.id, action='approve_supervisor')
        self.assertTrue(log.metadata['batch'])
        self.assertEqual(log.metadata['company_id'], self.c1.id)
        self.as_user(self.acc)
        r2 = self.client.post(f'/api/personnel/profiles/{p.id}/approve-finance/', {'note': 'ok'}, format='json')
        self.assertEqual(r2.status_code, 200, r2.data)
        self.as_user(self.mgr)
        r3 = self.client.post(f'/api/personnel/profiles/{p.id}/approve-manager/', {}, format='json')
        self.assertEqual(r3.status_code, 200, r3.data)
        p.refresh_from_db()
        self.assertEqual(p.approval_status, 'approved')
        self.as_user(self.sup)
        r4 = self.client.post(f'/api/personnel/profiles/{p.id}/approve-supervisor/', {}, format='json')
        self.assertEqual(r4.status_code, 400)

    # ── DEF-11: عکس کامل پیش از Purge و حذف ──
    def _profile_with_accounts_and_docs(self, status):
        p = self.complete_profile('0010376488', status=status)
        PersonnelBankAccount.objects.create(personnel=p, bank_name='صادرات', sheba_number=SHEBA2, is_primary=False)
        PersonnelDocument.objects.create(personnel=p, doc_type='national_card', file=SimpleUploadedFile('card.png', b'x'), uploaded_by=self.emp)
        p.refresh_from_db()
        return p

    def _assert_snapshot_rebuilds(self, snapshot, original_values):
        rebuilt = PersonnelProfile(**snapshot['profile'])
        for attname, value in original_values.items():
            if attname in ('created_at', 'updated_at'):
                continue
            rebuilt_value = getattr(rebuilt, attname)
            if hasattr(rebuilt_value, 'field') and hasattr(rebuilt_value, 'name'):  # FieldFile
                rebuilt_value = rebuilt_value.name or None
            self.assertEqual(str(rebuilt_value) if rebuilt_value is not None else None,
                             str(value) if value is not None else None, f'فیلد {attname} از عکس بازسازی نمی‌شود')
        self.assertEqual(len(snapshot['bank_accounts']), 2)
        self.assertEqual({a['sheba_number'] for a in snapshot['bank_accounts']}, {SHEBA1, SHEBA2})
        self.assertEqual(len(snapshot['documents']), 1)
        self.assertTrue(snapshot['documents'][0]['file'])
        self.assertEqual(snapshot['documents'][0]['uploaded_by_username'], 'rm_emp')

    def test_def11_purge_snapshot_is_complete_and_restorable(self):
        p = self._profile_with_accounts_and_docs('pending_supervisor')
        pid = p.id
        original = {f.attname: getattr(p, f.attname) for f in p._meta.concrete_fields}
        original['attachment'] = p.attachment.name or None
        self.as_user(self.sup)
        r = self.client.post(f'/api/personnel/profiles/{pid}/reject/', {'reason': 'ناقص'}, format='json')
        self.assertTrue(r.data['purged'])
        self.assertFalse(PersonnelProfile.objects.filter(pk=pid).exists())
        log = WorkflowAuditLog.objects.get(object_id=pid, action='purge')
        self.assertEqual(log.metadata['company_id'], self.c1.id)
        self._assert_snapshot_rebuilds(log.metadata['snapshot'], original)

    def test_def11_destroy_snapshot_is_complete(self):
        p = self._profile_with_accounts_and_docs('draft')
        pid = p.id
        original = {f.attname: getattr(p, f.attname) for f in p._meta.concrete_fields}
        original['attachment'] = p.attachment.name or None
        self.as_user(self.mgr)
        r = self.client.delete(f'/api/personnel/profiles/{pid}/')
        self.assertEqual(r.status_code, 200, r.data)
        log = WorkflowAuditLog.objects.get(object_id=pid, action='delete')
        self._assert_snapshot_rebuilds(log.metadata['snapshot'], original)


class Step2HighSeverityTests(RemediationTestBase):
    def setUp(self):
        super().setUp()
        from notifications.models import Notification
        self.Notification = Notification

    # ── DEF-12: اعلان‌ها ──
    def test_def12_next_stage_and_return_recipients(self):
        p = self.complete_profile('0098765434', status='pending_supervisor')
        self.as_user(self.sup)
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(f'/api/personnel/profiles/{p.id}/approve-supervisor/', {'note': 'ok'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(self.Notification.objects.filter(recipient=self.acc).count(), 1)
        n = self.Notification.objects.get(recipient=self.acc)
        self.assertEqual(n.target_id, str(p.id))

        p.refresh_from_db()
        p.approval_status = 'pending_manager'
        p.save(update_fields=['approval_status'])
        self.as_user(self.mgr)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f'/api/personnel/profiles/{p.id}/request-revision/', {'reason': 'اصلاح'}, format='json')
        self.assertEqual(self.Notification.objects.filter(recipient=self.acc).count(), 2, 'عودت مدیر باید به حسابدار اعلان دهد')
        self.assertEqual(self.Notification.objects.filter(recipient=self.emp).count(), 0, 'ثبت‌کننده نباید اعلان عودت به حسابدار را بگیرد')

        self.as_user(self.acc)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f'/api/personnel/profiles/{p.id}/request-revision/', {'reason': 'بررسی میدانی', 'target_status': 'pending_supervisor'}, format='json')
        self.assertEqual(self.Notification.objects.filter(recipient=self.sup).count(), 1)

    def test_def12_other_company_users_are_not_notified(self):
        outsider = User.objects.create_user(username='rm_acc_b', password='p')
        outsider.user_permissions.add(*self.acc.user_permissions.all())
        UserCompanyAccess.objects.create(user=outsider, company=self.c2, access_level='workspace_full')
        p = self.complete_profile('0098765434', status='pending_supervisor')
        self.as_user(self.sup)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f'/api/personnel/profiles/{p.id}/approve-supervisor/', {}, format='json')
        self.assertFalse(self.Notification.objects.filter(recipient=outsider).exists())

    def test_def12_purge_notification_keeps_target_id(self):
        p = self.complete_profile('0010376488', status='pending_supervisor')
        pid = p.id
        self.as_user(self.sup)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f'/api/personnel/profiles/{pid}/reject/', {'reason': 'ناقص'}, format='json')
        n = self.Notification.objects.get(recipient=self.emp)
        self.assertEqual(n.target_id, str(pid))

    # ── DEF-13: جدول گذار عودت ──
    def test_def13_revision_transition_table(self):
        draft = self.complete_profile('0010376488')
        self.as_user(self.mgr)
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{draft.id}/request-revision/', {'reason': 'x'}, format='json').status_code, 400)
        rejected = self.complete_profile('0076543218', status='rejected')
        self.as_user(self.acc)
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{rejected.id}/request-revision/', {'reason': 'x', 'target_status': 'pending_supervisor'}, format='json').status_code, 400)
        pa = self.complete_profile('0061122335', status='pending_accountant')
        self.as_user(self.sup)
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{pa.id}/request-revision/', {'reason': 'x'}, format='json').status_code, 403)
        self.as_user(self.acc)
        r = self.client.post(f'/api/personnel/profiles/{pa.id}/request-revision/', {'reason': 'x', 'target_status': 'pending_supervisor'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        pa.refresh_from_db()
        self.assertEqual(pa.approval_status, 'pending_supervisor')
        pm = self.complete_profile('0098765434', status='pending_manager')
        self.as_user(self.mgr)
        self.client.post(f'/api/personnel/profiles/{pm.id}/request-revision/', {'reason': 'x'}, format='json')
        pm.refresh_from_db()
        self.assertEqual(pm.approval_status, 'pending_accountant')
        log = WorkflowAuditLog.objects.filter(object_id=pm.id, action='request_revision').latest('id')
        self.assertEqual((log.from_status, log.to_status), ('pending_manager', 'pending_accountant'))

    def test_def13_reject_only_in_progress(self):
        draft = self.complete_profile('0010376488')
        self.as_user(self.sup)
        r = self.client.post(f'/api/personnel/profiles/{draft.id}/reject/', {'reason': 'x'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertTrue(PersonnelProfile.objects.filter(pk=draft.id).exists(), 'پیش‌نویس کارمند نباید با رد سرپرست پاک شود')

    # ── DEF-15: استعلام و بازگشایی ──
    def test_def15_lookup_without_header_is_scoped(self):
        self.complete_profile('0061122335', status='approved')
        self.complete_profile('0012345679', status='approved', section=self.s2, company=self.c2)
        self.as_user(self.emp)
        r = self.client.get('/api/personnel/profiles/lookup-by-national-code/?national_code=0061122335')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertTrue(r.data['found'])
        self.assertNotIn('daily_base_wage', r.data['personnel'])
        r2 = self.client.get('/api/personnel/profiles/lookup-by-national-code/?national_code=0012345679')
        self.assertFalse(r2.data['found'])
        # وجود کد ملی در شرکت بیرون از قلمرو هم افشا نمی‌شود (بازبینی قلمرو)
        self.assertNotIn('exists_in_other_company', r2.data)
        self.assertNotIn('personnel', r2.data)

    def test_def15_reopen_rejected_record(self):
        p = self.complete_profile('0010376488', status='rejected', rejection_reason='ناقص')
        self.as_user(self.emp)
        self.assertEqual(self.client.patch(f'/api/personnel/profiles/{p.id}/', {'phone_number': '09120000000'}, format='json').status_code, 400)
        r = self.client.post(f'/api/personnel/profiles/{p.id}/reopen/', {}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        p.refresh_from_db()
        self.assertEqual(p.approval_status, 'draft')
        self.assertTrue(WorkflowAuditLog.objects.filter(object_id=p.id, action='reopen', from_status='rejected', to_status='draft').exists())
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{p.id}/submit/', {}, format='json').status_code, 200)
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{p.id}/reopen/', {}, format='json').status_code, 400)

    def test_def15_reopen_other_company_is_blocked(self):
        other = self.complete_profile('0012345679', status='rejected', section=self.s2, company=self.c2)
        self.as_user(self.mgr)
        r = self.client.post(f'/api/personnel/profiles/{other.id}/reopen/', {}, format='json')
        self.assertIn(r.status_code, (403, 404))
        other.refresh_from_db()
        self.assertEqual(other.approval_status, 'rejected')

    # ── DEF-16 و DEF-17: درخواست تغییرات و ویرایش پروندهٔ مصوب ──
    def test_def16_employee_change_request_with_full_form_body(self):
        p = self.complete_profile('0098765434', status='approved', issue_date='1380/01/01')
        self.as_user(self.emp)
        body = {k: v for k, v in dict(self.client.get(f'/api/personnel/profiles/{p.id}/').data).items() if k != 'bank_accounts'}
        body['phone_number'] = '09125554433'
        r = self.client.patch(f'/api/personnel/profiles/{p.id}/', body, format='json')
        self.assertEqual(r.status_code, 202, r.data)
        cr = PersonnelChangeRequest.objects.get(pk=r.data['change_request_id'])
        self.assertEqual(cr.proposed_changes, {'phone_number': '09125554433'})
        p.refresh_from_db()
        self.assertEqual(p.phone_number, '09121112233', 'تا تصویب مدیر مقدار قبلی معتبر است')

    def test_def16_employee_cannot_propose_locked_field(self):
        p = self.complete_profile('0098765434', status='approved')
        self.as_user(self.emp)
        r = self.client.patch(f'/api/personnel/profiles/{p.id}/', {'job_code': '000111'}, format='json')
        self.assertEqual(r.status_code, 403)

    def test_def16_cancel_route_status_window_and_log(self):
        p = self.complete_profile('0098765434', status='approved')
        cr = PersonnelChangeRequest.objects.create(personnel=p, requested_by=self.emp, proposed_changes={'phone_number': '09120000000'}, previous_values={}, status='pending_supervisor')
        p.has_pending_changes = True
        p.save(update_fields=['has_pending_changes'])
        self.as_user(self.emp)
        r = self.client.post(f'/api/personnel/personnel-change-requests/{cr.id}/cancel/', {'reason': 'انصراف'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        p.refresh_from_db()
        self.assertFalse(p.has_pending_changes)
        log = WorkflowAuditLog.objects.get(object_id=p.id, action='change_request_cancelled')
        self.assertEqual(log.content_type, 'personnel')
        self.assertEqual(log.metadata['change_request_id'], cr.id)
        cr2 = PersonnelChangeRequest.objects.create(personnel=p, requested_by=self.emp, proposed_changes={'phone_number': '09120000000'}, previous_values={}, status='pending_accountant')
        self.assertEqual(self.client.post(f'/api/personnel/personnel-change-requests/{cr2.id}/cancel/', {}, format='json').status_code, 400)

    def test_def17_manager_cannot_directly_edit_approved_or_pending_manager(self):
        p = self.complete_profile('0061122335', status='approved')
        self.as_user(self.mgr)
        r = self.client.patch(f'/api/personnel/profiles/{p.id}/', {'sheba_number': SHEBA2, 'daily_base_wage': 9999999}, format='json')
        self.assertEqual(r.status_code, 202, r.data)
        p.refresh_from_db()
        self.assertEqual(p.sheba_number, SHEBA1)
        self.assertEqual(int(p.daily_base_wage), 7000000)
        self.assertEqual(PersonnelBankAccount.objects.get(personnel=p, is_primary=True).sheba_number, SHEBA1)
        p.approval_status = 'pending_manager'
        p.save(update_fields=['approval_status'])
        self.assertEqual(self.client.patch(f'/api/personnel/profiles/{p.id}/', {'phone_number': '09120000000'}, format='json').status_code, 400)

    def test_def17_superuser_direct_edit_is_logged(self):
        admin = User.objects.create_superuser(username='rm_admin', password='p')
        p = self.complete_profile('0061122335', status='approved')
        self.as_user(admin)
        r = self.client.patch(f'/api/personnel/profiles/{p.id}/', {'phone_number': '09120000000'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        log = WorkflowAuditLog.objects.get(object_id=p.id, action='superuser_edit')
        self.assertEqual(log.metadata['diff'], {'phone_number': '09120000000'})

    # ── DEF-06: منبع واحد اعمال درخواست تغییرات ──
    def test_def06_change_request_with_fk_applies_and_logs(self):
        p = self.complete_profile('0098765434', status='approved')
        cr = PersonnelChangeRequest.objects.create(personnel=p, requested_by=self.acc, proposed_changes={'project': self.p1.id, 'phone_number': '09127777777'}, previous_values={'project': None}, status='pending_manager')
        self.as_user(self.mgr)
        r = self.client.post(f'/api/personnel/personnel-change-requests/{cr.id}/approve-manager/', {}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        p.refresh_from_db()
        self.assertEqual(p.project_id, self.p1.id)
        self.assertEqual(p.phone_number, '09127777777')
        self.assertTrue(WorkflowAuditLog.objects.filter(object_id=p.id, action='change_request_applied').exists())

    def test_def06_invalid_change_request_is_not_applied(self):
        p = self.complete_profile('0098765434', status='approved')
        cr = PersonnelChangeRequest.objects.create(personnel=p, requested_by=self.acc, proposed_changes={'phone_number': '123'}, previous_values={}, status='pending_manager')
        self.as_user(self.mgr)
        r = self.client.post(f'/api/personnel/personnel-change-requests/{cr.id}/approve-manager/', {}, format='json')
        self.assertEqual(r.status_code, 400)
        cr.refresh_from_db()
        p.refresh_from_db()
        self.assertEqual(cr.status, 'pending_manager')
        self.assertEqual(p.phone_number, '09121112233')

    def test_def06_requester_whitelist_rechecked_on_apply(self):
        p = self.complete_profile('0098765434', status='approved')
        cr = PersonnelChangeRequest.objects.create(personnel=p, requested_by=self.emp, proposed_changes={'daily_base_wage': '9999999'}, previous_values={}, status='pending_manager')
        self.as_user(self.mgr)
        r = self.client.post(f'/api/personnel/personnel-change-requests/{cr.id}/approve-manager/', {}, format='json')
        self.assertEqual(r.status_code, 400)
        p.refresh_from_db()
        self.assertEqual(int(p.daily_base_wage), 7000000)

    # ── DEF-18: قطع همکاری ──
    def test_def18_end_collaboration(self):
        p = self.complete_profile('0098765434', status='approved')
        self.as_user(self.emp)
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{p.id}/end-collaboration/', {'end_date': '1405/07/01', 'reason': 'x'}, format='json').status_code, 403)
        self.as_user(self.mgr)
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{p.id}/end-collaboration/', {'end_date': '1405/07/01'}, format='json').status_code, 400)
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{p.id}/end-collaboration/', {'end_date': 'abc', 'reason': 'x'}, format='json').status_code, 400)
        r = self.client.post(f'/api/personnel/profiles/{p.id}/end-collaboration/', {'end_date': '۱۴۰۵/۷/۱', 'reason': 'پایان قرارداد'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        p.refresh_from_db()
        self.assertFalse(p.is_active)
        self.assertEqual(p.end_date, '1405/07/01')
        self.assertTrue(WorkflowAuditLog.objects.filter(object_id=p.id, action='end_collaboration').exists())
        self.assertEqual(self.client.post(f'/api/personnel/profiles/{p.id}/end-collaboration/', {'end_date': '1405/07/02', 'reason': 'x'}, format='json').status_code, 400)

    # ── DEF-19: قلمرو و یکدستی لاگ ──
    def test_def19_audit_log_scope_without_header(self):
        other = self.complete_profile('0088776654', status='pending_supervisor', section=self.s2, company=self.c2)
        WorkflowAuditLog.objects.create(content_type='personnel', object_id=other.id, action='submit', from_status='draft', to_status='pending_supervisor')
        WorkflowAuditLog.objects.create(content_type='personnel', object_id=999999, action='purge', metadata={'company_id': self.c2.id, 'snapshot': {'id': 999999}})
        mine = self.complete_profile('0098765434', status='pending_supervisor')
        WorkflowAuditLog.objects.create(content_type='personnel', object_id=mine.id, action='submit', from_status='draft', to_status='pending_supervisor')
        WorkflowAuditLog.objects.create(content_type='personnel', object_id=888888, action='delete', metadata={'company_id': self.c1.id, 'snapshot': {'id': 888888, 'full_name': 'x', 'sheba_number': SHEBA1}})
        self.as_user(self.emp)
        r = self.client.get('/api/personnel/workflow-audit-logs/')
        ids = {row['object_id'] for row in r.data}
        self.assertEqual(ids, {mine.id, 888888})
        deleted = next(row for row in r.data if row['object_id'] == 888888)
        self.assertNotIn('sheba_number', deleted['metadata']['snapshot'], 'عکس کامل حذف برای نقش غیرمالی خلاصه می‌شود')

    def test_def19_change_request_logs_found_by_diff_viewer_filter(self):
        p = self.complete_profile('0098765434', status='approved')
        cr = PersonnelChangeRequest.objects.create(personnel=p, requested_by=self.emp, proposed_changes={'phone_number': '09120000000'}, previous_values={}, status='pending_supervisor')
        self.as_user(self.sup)
        self.client.post(f'/api/personnel/personnel-change-requests/{cr.id}/approve-supervisor/', {}, format='json')
        r = self.client.get(f'/api/personnel/workflow-audit-logs/?content_type=personnelchangerequest&object_id={cr.id}')
        self.assertEqual([row['action'] for row in r.data], ['change_request_approve_supervisor'])

    def test_def19_engine_logs_use_unified_content_type(self):
        from personnel.workflow_engine import _record_audit_log
        p = self.complete_profile('0098765434', status='pending_supervisor')
        _record_audit_log(p, self.sup, 'pending_supervisor', 'pending_accountant', 'approve')
        self.assertTrue(WorkflowAuditLog.objects.filter(content_type='personnel', object_id=p.id, action='approve').exists())
        self.assertFalse(WorkflowAuditLog.objects.filter(content_type='personnelprofile').exists())

    # ── DEF-20: تاریخ اعتبار و دیسکت ──
    def test_def20_valid_until_date_validation_and_comparison(self):
        from personnel.lifecycle_services import validate_personnel_validity
        p = self.complete_profile('0098765434')
        self.as_user(self.emp)
        self.assertEqual(self.client.patch(f'/api/personnel/profiles/{p.id}/', {'valid_until_date': 'abc'}, format='json').status_code, 400)
        self.assertEqual(self.client.patch(f'/api/personnel/profiles/{p.id}/', {'valid_until_date': '1405/13/01'}, format='json').status_code, 400)
        r = self.client.patch(f'/api/personnel/profiles/{p.id}/', {'valid_until_date': '۱۴۰۵/۴/۵'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        p.refresh_from_db()
        self.assertEqual(p.valid_until_date, '1405/04/05')
        p.valid_until_date = '1405/4/5'
        self.assertFalse(validate_personnel_validity(p, '1405/04/15')[0])
        self.assertTrue(validate_personnel_validity(p, '1405/04/05')[0])
        self.assertFalse(validate_personnel_validity(p, '1405/04')[0])
        p.valid_until_date = '1405/04/31'
        self.assertTrue(validate_personnel_validity(p, '1405/04')[0])

    def test_def20_diskette_preflight_lists_expired_personnel(self):
        from personnel.models import MonthlyPayrollRecord, MonthlyWorkPeriod
        expired = self.complete_profile('0098765434', status='approved', valid_until_date='1405/03/31', project=self.p1)
        valid = self.complete_profile('0061122335', status='approved', project=self.p1)
        period = MonthlyWorkPeriod.objects.create(year_month='1405/04', company=self.c1)
        for prof in (expired, valid):
            MonthlyPayrollRecord.objects.create(period=period, personnel=prof, national_code=prof.national_code, full_name=prof.full_name)
        self.as_user(self.acc)
        r = self.client.get(f'/api/personnel/monthly-payroll/expired-personnel/?period_id={period.id}&project_id={self.p1.id}', HTTP_X_COMPANY_ID=str(self.c1.id))
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['count'], 1)
        self.assertEqual(r.data['personnel'][0]['personnel_id'], expired.id)

    # ── DEF-21: دستمزد برای نقش غیرمالی ──
    def test_def21_wages_hidden_and_not_writable_for_employee(self):
        self.as_user(self.emp)
        r = self.client.post('/api/personnel/profiles/', self.employee_form_payload('0076543218', daily_base_wage=99999999), format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertNotIn('daily_base_wage', r.data)
        p = PersonnelProfile.objects.get(pk=r.data['id'])
        self.assertEqual(int(p.daily_base_wage), 0)
        listing = self.client.get(f'/api/personnel/profiles/?section_id={self.s1.id}').data
        self.assertTrue(all('daily_base_wage' not in row and 'contract_base_salary' not in row for row in listing))
        self.as_user(self.acc)
        self.assertIn('daily_base_wage', self.client.get(f'/api/personnel/profiles/{p.id}/').data)

    # ── DEF-23: گیت شبای معتبر حساب اصلی ──
    def test_def23_gate_requires_valid_primary_sheba(self):
        from personnel.lifecycle_services import get_missing_fields
        p = self.complete_profile('0098765434')
        PersonnelBankAccount.objects.filter(personnel=p).update(sheba_number='')
        p.refresh_from_db()
        self.assertIn('primary_bank_account', [m['field'] for m in get_missing_fields(p, 'submit')])
        PersonnelBankAccount.objects.filter(personnel=p).update(sheba_number='IR000000000000000000000000')
        self.assertIn('primary_bank_account', [m['field'] for m in get_missing_fields(p, 'submit')])
        PersonnelBankAccount.objects.filter(personnel=p).update(sheba_number=SHEBA1)
        self.assertNotIn('primary_bank_account', [m['field'] for m in get_missing_fields(p, 'submit')])

    # ── DEF-05 و DEF-04: برآورد سروری ──
    def test_def05_cost_estimate_follows_payroll_engine_rules(self):
        p = self.complete_profile('0098765434', status='pending_accountant')
        self.as_user(self.acc)
        single = self.client.post('/api/personnel/profiles/cost-estimate/', {'personnel_id': p.id, 'year_month': '1405/04', 'overrides': {'marital_status': 'single', 'children_count': 0}}, format='json')
        married = self.client.post('/api/personnel/profiles/cost-estimate/', {'personnel_id': p.id, 'year_month': '1405/04', 'overrides': {'marital_status': 'married', 'children_count': 0}}, format='json')
        kids = self.client.post('/api/personnel/profiles/cost-estimate/', {'personnel_id': p.id, 'year_month': '1405/04', 'overrides': {'marital_status': 'single', 'children_count': 2}}, format='json')
        self.assertEqual(single.status_code, 200, single.data)
        self.assertEqual(single.data['marital_allowance'], 0)
        self.assertGreater(married.data['marital_allowance'], 0)
        self.assertEqual(kids.data['marital_allowance'], 0, 'حق اولاد مشروط به تأهل نیست و حق تأهل به مجرد داده نمی‌شود')
        self.assertGreater(kids.data['child_allowance'], 0)
        p.refresh_from_db()
        self.assertEqual(p.marital_status, 'single', 'overrides نباید ذخیره شود')
        self.as_user(self.emp)
        self.assertEqual(self.client.post('/api/personnel/profiles/cost-estimate/', {'personnel_id': p.id}, format='json').status_code, 403)

    def test_def04_company_cost_summary_counts_only_approved_active(self):
        self.complete_profile('0098765434', status='approved')
        self.complete_profile('0061122335', status='approved', is_active=False)
        self.complete_profile('0010376488', status='draft')
        self.complete_profile('0012345679', status='approved', section=self.s2, company=self.c2)
        self.as_user(self.mgr)
        r = self.client.get('/api/personnel/profiles/company-cost-summary/?year_month=1405/04&search=zzz', HTTP_X_COMPANY_ID=str(self.c1.id))
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['personnel_count'], 1)
        self.assertGreater(r.data['average_employer_monthly_cost'], 0)

    # ── DEF-14: درون‌ریزی مدیر با گیت ──
    def test_def14_manager_import_keeps_incomplete_rows_as_draft(self):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Emp_info'
        ws.append(['کد ملی', 'نام', 'نام خانوادگی', 'نام پدر'])
        ws.append(['0076543218', 'سارا', 'احمدی', 'محمد'])
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile('emp.xlsx', buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.as_user(self.mgr)
        r = self.client.post('/api/personnel/profiles/import-excel/', {'file': upload, 'section_id': self.s1.id}, format='multipart', HTTP_X_COMPANY_ID=str(self.c1.id))
        self.assertEqual(r.status_code, 200, r.data)
        p = PersonnelProfile.objects.get(national_code='0076543218')
        self.assertEqual(p.approval_status, 'draft')
        self.assertEqual(len(r.data['warnings']), 1)

    # ── DEF-27: انحراف‌های کوچک ──
    def test_def27_co_section_member_can_delete_draft(self):
        colleague = User.objects.create_user(username='rm_colleague', password='p')
        UserSectionAssignment.objects.create(user=colleague, section=self.s1, is_active=True)
        p = self.complete_profile('0098765434')
        self.as_user(colleague)
        r = self.client.delete(f'/api/personnel/profiles/{p.id}/')
        self.assertEqual(r.status_code, 200, r.data)

    def test_def27_document_upload_is_validated(self):
        p = self.complete_profile('0098765434')
        self.as_user(self.emp)
        bad_type = self.client.post(f'/api/personnel/profiles/{p.id}/documents/', {'file': SimpleUploadedFile('a.exe', b'x'), 'doc_type': 'other'}, format='multipart')
        self.assertEqual(bad_type.status_code, 400)
        bad_doc_type = self.client.post(f'/api/personnel/profiles/{p.id}/documents/', {'file': SimpleUploadedFile('a.pdf', b'x'), 'doc_type': 'passport'}, format='multipart')
        self.assertEqual(bad_doc_type.status_code, 400)
        ok = self.client.post(f'/api/personnel/profiles/{p.id}/documents/', {'file': SimpleUploadedFile('a.pdf', b'x'), 'doc_type': 'contract'}, format='multipart')
        self.assertEqual(ok.status_code, 201, ok.data)

    def test_def27_cartable_does_not_bypass_personnel_workflow(self):
        draft = self.complete_profile('0098765434')
        pending = self.complete_profile('0061122335', status='pending_supervisor')
        self.as_user(self.sup)
        r = self.client.post('/api/personnel/cartable/supervisor/', {'action': 'approve', 'model': 'personnel', 'id': draft.id}, format='json')
        self.assertEqual(r.status_code, 400)
        draft.refresh_from_db()
        self.assertEqual(draft.approval_status, 'draft')
        listing = self.client.get('/api/personnel/cartable/supervisor/', HTTP_X_COMPANY_ID=str(self.c1.id)).data
        self.assertEqual([row['id'] for row in listing['personnel']], [pending.id])

    # ── DEF-23: گزارش حساب‌های اصلی فاقد شبای معتبر (گام ۱-الف و ۱-ب اصلاحیه) ──
    def test_def23_missing_primary_sheba_service_and_excel_readonly(self):
        from personnel.lifecycle_services import (
            get_personnel_missing_primary_sheba,
            generate_missing_primary_sheba_excel,
        )
        from personnel.models import PersonnelBankAccount
        import openpyxl
        import io

        # پرونده ۱: بدون حساب بانکی و بدون شبای آینه (باید در گزارش بیاید)
        p1 = self.complete_profile('0098765434', sheba_number='', bank_name='', account_number='')

        # پرونده ۲: با حساب اصلی دارای شبای معتبر (نباید در گزارش بیاید)
        p2 = self.complete_profile('0061122335', sheba_number=SHEBA1, bank_name='بانک ملی')

        # پرونده ۳: با حساب اصلی فاقد شبا (باید در گزارش بیاید)
        p3 = self.complete_profile('0076543218', sheba_number='', bank_name='', account_number='')
        PersonnelBankAccount.objects.create(
            personnel=p3,
            bank_name='بانک صادرات',
            account_number='0201111111002',
            sheba_number='',
            is_primary=True,
            is_active=True
        )

        before_profiles_count = PersonnelProfile.objects.count()
        before_accounts_count = PersonnelBankAccount.objects.count()

        results = get_personnel_missing_primary_sheba()
        result_ids = [r['personnel_id'] for r in results]

        self.assertIn(p1.id, result_ids)
        self.assertNotIn(p2.id, result_ids)
        self.assertIn(p3.id, result_ids)

        # بررسی ساختار فیلدها و راهنمای تکمیل
        p3_row = next(r for r in results if r['personnel_id'] == p3.id)
        self.assertEqual(p3_row['account_number'], '0201111111002')
        self.assertEqual(p3_row['sheba_number'], '')
        self.assertIn('درخواست تغییرات', p3_row['completion_guide'])

        # تولید اکسل و اعتبارسنجی معماری ۲ سطری و freeze_panes
        excel_bytes = generate_missing_primary_sheba_excel()
        self.assertTrue(len(excel_bytes) > 0)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        self.assertEqual(ws.freeze_panes, 'A3')
        self.assertEqual(ws.cell(row=1, column=1).value, 'ردیف')
        self.assertEqual(ws.cell(row=2, column=1).value, 'row_number')

        # تضمین اصل عدم تغییر داده (Read-only Invariant)
        self.assertEqual(PersonnelProfile.objects.count(), before_profiles_count)
        self.assertEqual(PersonnelBankAccount.objects.count(), before_accounts_count)

    def test_def23_missing_primary_sheba_endpoint_financial_role_and_tenant_scope(self):
        # ساخت یک پرونده فاقد شبا در شرکت ۱ (تحت دسترسی self.acc)
        p_c1 = self.complete_profile('0098765434', company=self.c1, sheba_number='', bank_name='', account_number='')
        # ساخت یک پرونده فاقد شبا در شرکت ۲ (خارج از دسترسی self.acc)
        p_c2 = self.complete_profile('0076543218', company=self.c2, section=self.s2, sheba_number='', bank_name='', account_number='')

        # ۱. کاربر عادی غیرمالی باید 403 Forbidden دریافت کند
        self.as_user(self.emp)
        r_forbidden = self.client.get('/api/personnel/profiles/missing-primary-sheba-report/')
        self.assertEqual(r_forbidden.status_code, 403)

        # ۲. حسابدار دسترسی دارد و خروجی اکسل ۲ سطری استاندارد دانلود می‌شود
        self.as_user(self.acc)
        r_ok = self.client.get('/api/personnel/profiles/missing-primary-sheba-report/', HTTP_X_COMPANY_ID=str(self.c1.id))
        self.assertEqual(r_ok.status_code, 200)
        self.assertEqual(r_ok['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('Missing_Primary_Sheba_Report.xlsx', r_ok['Content-Disposition'])

        # ۳. اعتبارسنجی قلمرو شرکت (Tenant Scope Reviewer): فقط رکوردهای شرکت الف در خروجی باشد
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(r_ok.content))
        ws = wb.active
        national_codes = [ws.cell(row=i, column=5).value for i in range(3, ws.max_row + 1)]
        self.assertIn('0098765434', national_codes)
        self.assertNotIn('0076543218', national_codes)

    # ── DEF-20: سیاست تعلیق پرداخت فیش‌های منقضی و دیسکت بانک (گام ۲ اصلاحیه) ──
    def test_def20_payroll_engine_sets_payment_hold_for_expired_personnel_with_work(self):
        from personnel.models import MonthlyWorkPeriod, MonthlyPayrollRecord, DailyAttendance
        from personnel.payroll_engine import calculate_monthly_payroll_for_period

        # پرسنل با پایان تاریخ اعتبار قبل از دوره ۱۴۰۵/۰۷
        p_expired = self.complete_profile(
            '0011223344',
            status='approved',
            company=self.c1,
            project=self.p1,
            valid_until_date='1405/06/31',
            daily_base_wage=5000000
        )
        period = MonthlyWorkPeriod.objects.create(
            company=self.c1,
            project=self.p1,
            year_month='1405/07',
            status='OPEN'
        )
        # ثبت حضور و کارکرد برای پرسنل منقضی
        DailyAttendance.objects.create(
            personnel=p_expired,
            project=self.p1,
            section=self.s1,
            date_shamsi='1405/07/05',
            status='PRESENT_10H',
            approval_status='APPROVED'
        )

        # محاسبه مکانیزه دوره حقوق
        recs = calculate_monthly_payroll_for_period(period, user=self.acc)
        self.assertGreater(len(recs), 0)

        record = MonthlyPayrollRecord.objects.get(period=period, personnel=p_expired)
        self.assertGreater(record.gross_salary, 0)
        self.assertTrue(record.payment_hold, 'فیش حقوقی پرسنل منقضیِ دارای کارکرد باید با payment_hold=True علامت بخورد')
        self.assertIn('منقضی', record.payment_hold_reason)

    def test_def20_bank_export_excludes_held_and_supports_only_released(self):
        from personnel.models import MonthlyWorkPeriod, MonthlyPayrollRecord
        from django.utils import timezone

        p_normal = self.complete_profile('0033445566', status='approved', company=self.c1, project=self.p1, valid_until_date='1406/01/01')
        p_held = self.complete_profile('0044556677', status='approved', company=self.c1, project=self.p1, valid_until_date='1405/01/01')

        period = MonthlyWorkPeriod.objects.create(
            company=self.c1,
            project=self.p1,
            year_month='1405/07',
            status='OPEN'
        )
        r_normal = MonthlyPayrollRecord.objects.create(
            period=period,
            personnel=p_normal,
            gross_salary=10000000,
            payable_amount=9000000,
            payment_hold=False
        )
        r_held = MonthlyPayrollRecord.objects.create(
            period=period,
            personnel=p_held,
            gross_salary=8000000,
            payable_amount=7000000,
            payment_hold=True,
            payment_hold_reason='پرسنل منقضی'
        )

        self.as_user(self.acc)

        # ۱. صدور دیسکت پیش‌فرض: رکوردهای payment_hold مستثنی می‌شوند
        res_default = self.client.get(f'/api/personnel/monthly-payroll/export-bank-excel/?period_id={period.id}&project_id={self.p1.id}')
        self.assertEqual(res_default.status_code, 200)
        self.assertIn('Bank_Payment_', res_default['Content-Disposition'])
        self.assertNotIn('Released', res_default['Content-Disposition'])

        # ۲. صدور با only_released=true پیش از آزادسازی: باید 404 بدهد زیرا هیچ رکورد آزادشده‌ای نیست
        res_rel_empty = self.client.get(f'/api/personnel/monthly-payroll/export-bank-excel/?period_id={period.id}&project_id={self.p1.id}&only_released=true')
        self.assertEqual(res_rel_empty.status_code, 404)

        # آزادسازی رکورد معلق
        r_held.payment_hold = False
        r_held.payment_hold_released_at = timezone.now()
        r_held.payment_hold_released_by = self.acc
        r_held.save()

        # ۳. صدور با only_released=true پس از آزادسازی: با موفقیت صادر شده و عنوان فایل Released دارد
        res_rel = self.client.get(f'/api/personnel/monthly-payroll/export-bank-excel/?period_id={period.id}&project_id={self.p1.id}&only_released=true')
        self.assertEqual(res_rel.status_code, 200)
        self.assertIn('Bank_Payment_Released_', res_rel['Content-Disposition'])

    def test_def20_release_payment_hold_action_checks_extension_and_audit(self):
        from personnel.models import MonthlyWorkPeriod, MonthlyPayrollRecord
        from accounts.models import AuditLog

        p_expired = self.complete_profile('0055667788', status='approved', company=self.c1, project=self.p1, valid_until_date='1405/01/01')
        period = MonthlyWorkPeriod.objects.create(
            company=self.c1,
            project=self.p1,
            year_month='1405/07',
            status='OPEN'
        )
        rec = MonthlyPayrollRecord.objects.create(
            period=period,
            personnel=p_expired,
            gross_salary=5000000,
            payable_amount=4500000,
            payment_hold=True,
            payment_hold_reason='پرسنل منقضی'
        )

        # ۱. کاربر عادی غیرمالی ۴۰۳ می‌گیرد
        self.as_user(self.emp)
        r_forbidden = self.client.post(f'/api/personnel/monthly-payroll/{rec.id}/release-payment-hold/')
        self.assertEqual(r_forbidden.status_code, 403)

        # ۲. حسابدار تلاش می‌کند اما پرونده هنوز منقضی است -> خطا ۴۰۰
        self.as_user(self.acc)
        r_bad = self.client.post(f'/api/personnel/monthly-payroll/{rec.id}/release-payment-hold/')
        self.assertEqual(r_bad.status_code, 400)
        self.assertIn('اعتبار پرونده هنوز تمدید نشده است', r_bad.data['error'])

        # ۳. تمدید اعتبار پرسنل (مثلاً تا پایان سال ۱۴۰۵)
        p_expired.valid_until_date = '1405/12/29'
        p_expired.save(update_fields=['valid_until_date'])

        # ۴. حسابدار دوباره فراخوانی می‌کند -> موفقیت ۲۰۰، ثبت زمان و کاربر و لاگ ممیزی
        r_ok = self.client.post(f'/api/personnel/monthly-payroll/{rec.id}/release-payment-hold/')
        self.assertEqual(r_ok.status_code, 200)
        rec.refresh_from_db()
        self.assertFalse(rec.payment_hold)
        self.assertIsNotNone(rec.payment_hold_released_at)
        self.assertEqual(rec.payment_hold_released_by, self.acc)

        # بررسی لاگ ممیزی در AuditLog
        log = AuditLog.objects.filter(
            action='RELEASE_PAYMENT_HOLD',
            target_object_id=str(rec.id)
        ).first()
        self.assertIsNotNone(log)
        self.assertEqual(log.user, self.acc)

    def test_def20_no_duplicate_payment_between_default_and_released_diskettes(self):
        """
        تست تضمین عدم پرداخت دوباره:
        اشتراک رکوردهای خروجی دیسکت عادی و دیسکت آزادشده همواره باید تهی باشد (set() == set()).
        صدور مجدد دیسکت عادی پس از آزادسازی هرگز نباید رکورد آزادشده را شامل شود.
        """
        import io
        import openpyxl
        from personnel.models import MonthlyWorkPeriod, MonthlyPayrollRecord

        p_valid = self.complete_profile('0011223344', status='approved', company=self.c1, project=self.p1, valid_until_date='1405/12/29')
        p_expired = self.complete_profile('0022334455', status='approved', company=self.c1, project=self.p1, valid_until_date='1405/01/01')

        period = MonthlyWorkPeriod.objects.create(
            company=self.c1,
            project=self.p1,
            year_month='1405/07',
            status='OPEN'
        )

        r_valid = MonthlyPayrollRecord.objects.create(
            period=period,
            personnel=p_valid,
            gross_salary=6000000,
            payable_amount=5500000,
            bank_account_number='11111111',
            payment_hold=False
        )

        r_expired = MonthlyPayrollRecord.objects.create(
            period=period,
            personnel=p_expired,
            gross_salary=7000000,
            payable_amount=6500000,
            bank_account_number='22222222',
            payment_hold=True,
            payment_hold_reason='انقضای اعتبار تا تمدید'
        )

        def extract_national_codes(content):
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            return {str(row[6]).strip() for row in ws.iter_rows(values_only=True) if row[6] and str(row[6]).strip() != 'کد ملی (اختیاری)'}

        self.as_user(self.acc)

        # مرحله ۱: پیش از آزادسازی - دیسکت عادی فقط رکورد معتبر را دارد
        res1 = self.client.get(f'/api/personnel/monthly-payroll/export-bank-excel/?period_id={period.id}&project_id={self.p1.id}')
        self.assertEqual(res1.status_code, 200)
        nc1 = extract_national_codes(res1.content)
        self.assertIn('0011223344', nc1)
        self.assertNotIn('0022334455', nc1)

        # مرحله ۲: تمدید اعتبار و آزادسازی رسمی رکورد معلق
        p_expired.valid_until_date = '1405/12/29'
        p_expired.save(update_fields=['valid_until_date'])
        rel_resp = self.client.post(f'/api/personnel/monthly-payroll/{r_expired.id}/release-payment-hold/')
        self.assertEqual(rel_resp.status_code, 200)

        # مرحله ۳: صدور مجدد دیسکت عادی پس از آزادسازی - باید کماکان رکورد آزادشده را استثنا کند
        res_default_after_release = self.client.get(f'/api/personnel/monthly-payroll/export-bank-excel/?period_id={period.id}&project_id={self.p1.id}')
        self.assertEqual(res_default_after_release.status_code, 200)
        nc_default = extract_national_codes(res_default_after_release.content)

        # مرحله ۴: صدور دیسکت رکوردهای آزادشده
        res_released = self.client.get(f'/api/personnel/monthly-payroll/export-bank-excel/?period_id={period.id}&project_id={self.p1.id}&only_released=true')
        self.assertEqual(res_released.status_code, 200)
        nc_released = extract_national_codes(res_released.content)

        # بررسی و تضمین ریاضی:
        self.assertIn('0011223344', nc_default)
        self.assertNotIn('0022334455', nc_default, "خطای اساسی: رکورد آزادشده در دیسکت عادی تکرار شده است (خطر پرداخت دوباره)!")
        self.assertIn('0022334455', nc_released)
        self.assertNotIn('0011223344', nc_released)

        # تقاطع دو مجموعه حتماً باید تهی باشد
        intersection = nc_default.intersection(nc_released)
        self.assertEqual(intersection, set(), f"تداخل و خطر پرداخت مضاعف کشف شد: {intersection}")

    def test_def20_recalculate_does_not_bypass_payment_hold(self):
        """
        تست جلوگیری از دور زدن تعلیق:
        اگر پرسنل منقضی بوده و فیش او معلق شده، با تمدید تاریخ اعتبار و اجرای دوباره محاسبه حقوق،
        قفل تعلیق نباید بی‌صدا برداشته شود؛ تنها راه آزادسازی اکشن صریح حسابدار است.
        """
        from personnel.models import MonthlyWorkPeriod, MonthlyPayrollRecord, DailyAttendance
        from personnel.payroll_engine import calculate_monthly_payroll_for_period

        p = self.complete_profile('0033445566', status='approved', company=self.c1, project=self.p1, valid_until_date='1405/01/01')
        period = MonthlyWorkPeriod.objects.create(
            company=self.c1,
            project=self.p1,
            year_month='1405/07',
            status='OPEN'
        )
        DailyAttendance.objects.create(
            personnel=p,
            project=self.p1,
            section=self.s1,
            date_shamsi='1405/07/15',
            status='PRESENT_10H',
            approval_status='APPROVED'
        )

        # محاسبه اول: چون تاریخ اعتبار گذشته است، فیش باید معلق شود
        calculate_monthly_payroll_for_period(period, user=self.acc)
        rec = MonthlyPayrollRecord.objects.get(period=period, personnel=p)
        self.assertTrue(rec.payment_hold)
        self.assertIn('اعتبار', rec.payment_hold_reason)

        # تمدید تاریخ اعتبار پرسنل در دیتابیس
        p.valid_until_date = '1405/12/29'
        p.save(update_fields=['valid_until_date'])

        # محاسبه مجدد توسط موتور حقوق:
        # قفل تعلیق نباید خودکار باز شود تا فرآیند ممیزی دور زده نشود
        calculate_monthly_payroll_for_period(period, user=self.acc)
        rec.refresh_from_db()
        self.assertTrue(rec.payment_hold, "خطا: محاسبه مجدد قفل تعلیق را بدون ممیزی برداشت!")
        self.assertIsNone(rec.payment_hold_released_at)

        # حسابدار رسماً دکمه آزادسازی را می‌زند
        self.as_user(self.acc)
        rel_resp = self.client.post(f'/api/personnel/monthly-payroll/{rec.id}/release-payment-hold/')
        self.assertEqual(rel_resp.status_code, 200)
        rec.refresh_from_db()
        self.assertFalse(rec.payment_hold)
        self.assertIsNotNone(rec.payment_hold_released_at)

        # اگر بعد از آزادسازی باز هم محاسبه مجدد اجرا شود، وضعیت آزادسازی و متادیتا پاک نمی‌شود
        calculate_monthly_payroll_for_period(period, user=self.acc)
        rec.refresh_from_db()
        self.assertFalse(rec.payment_hold)
        self.assertIsNotNone(rec.payment_hold_released_at)
        self.assertEqual(rec.payment_hold_released_by, self.acc)

    def test_def20_cannot_bypass_payment_hold_via_patch_or_put(self):
        """
        تست امنیتی جلوگیری از دور زدن تعلیق پرداخت با PATCH یا PUT:
        هیچ کاربری حتی حسابدار نباید بتواند با ارسال مستقیم فیلد payment_hold در درخواست update،
        تعلیق را بدون ممیزی و بدون تمدید اعتبار آزاد کند.
        """
        from personnel.models import MonthlyWorkPeriod, MonthlyPayrollRecord

        p_expired = self.complete_profile('0099887766', status='approved', company=self.c1, project=self.p1, valid_until_date='1405/01/01')
        period = MonthlyWorkPeriod.objects.create(
            company=self.c1,
            project=self.p1,
            year_month='1405/07',
            status='OPEN'
        )
        rec = MonthlyPayrollRecord.objects.create(
            period=period,
            personnel=p_expired,
            gross_salary=8000000,
            payable_amount=7200000,
            payment_hold=True,
            payment_hold_reason='انقضای اعتبار تا تمدید'
        )

        # ۱. کاربر عادی غیرمالی تلاش می‌کند با PATCH فیش را ویرایش کند -> ۴۰۳
        self.as_user(self.emp)
        patch_emp = self.client.patch(f'/api/personnel/monthly-payroll/{rec.id}/', {'payment_hold': False}, format='json')
        self.assertEqual(patch_emp.status_code, 403)
        rec.refresh_from_db()
        self.assertTrue(rec.payment_hold)

        # ۲. حسابدار تلاش می‌کند با PATCH تعلیق را دور بزند -> خطای ۴۰۰ صریح
        self.as_user(self.acc)
        patch_acc = self.client.patch(f'/api/personnel/monthly-payroll/{rec.id}/', {'payment_hold': False}, format='json')
        self.assertEqual(patch_acc.status_code, 400)
        self.assertIn('تغییر مستقیم وضعیت تعلیق پرداخت مجاز نیست', patch_acc.data['error'])
        rec.refresh_from_db()
        self.assertTrue(rec.payment_hold, "خطای بحرانی: قفل تعلیق پرداخت با PATCH دور زده شد!")

        # ۳. ارسال PUT کامل نیز مسدود است
        put_acc = self.client.put(f'/api/personnel/monthly-payroll/{rec.id}/', {'payment_hold': False}, format='json')
        self.assertEqual(put_acc.status_code, 400)
        rec.refresh_from_db()
        self.assertTrue(rec.payment_hold)
