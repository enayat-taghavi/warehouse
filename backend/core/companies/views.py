import io

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets, permissions, parsers, serializers, status
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response

from accounts.views import OptionalPageNumberPagination
from common.date_utils import normalize_digits
from common.mixins import DeleteImpactMixin

from platform_core.registry import contributions

from .access import (
    companies_with_any_section,
    companies_with_section,
    company_section_level,
    ensure_section,
    user_section_levels,
    derived_company_access_details,
    derived_company_ids,
    ensure_scoped_write,
    get_user_allowed_companies,
)
from .models import Company, UserCompanyAccess
from .permissions import IsOrgStructureManagerOrReadOnly, requires_company_scope
from .realtime import broadcast_org_structure_updated
from .serializers import CompanySerializer, UserCompanyAccessSerializer


# وابستگی‌های شرکت که در صورت وجود، حذف شرکت را مسدود می‌کنند.
# ماژول‌ها این وابستگی‌ها را در نقطهٔ توسعهٔ company_delete_blockers ثبت می‌کنند (گام ۴).


class CompanyViewSet(DeleteImpactMixin, viewsets.ModelViewSet):
    """
    مدیریت شرکت‌ها (هلدینگ و شرکت‌های تابعه) در هسته. حسابداری با زیرکلاس خود
    (personnel.views.CompanyViewSet) شمارش پروژه‌ها، مدارک و پایش سررسید مدارک را اضافه می‌کند.
    """
    queryset = Company.objects.all()
    serializer_class = CompanySerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    pagination_class = OptionalPageNumberPagination

    def base_queryset(self):
        """نقطهٔ گسترش برای ماژول‌ها (annotate و prefetch)."""
        return Company.objects.all()

    def get_queryset(self):
        user = self.request.user
        qs = self.base_queryset()

        if not user.is_superuser and self.action != 'user_available':
            # عضو (دادهٔ ماژول‌ها) یا دارندهٔ دست‌کم یک بخش با سطح خواندن (حالت محدود استودیو)
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            allowed_ids |= set(companies_with_any_section(user).values_list('id', flat=True))
            qs = qs.filter(id__in=allowed_ids)

        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() == 'true')

        search = self.request.query_params.get('search')
        if search:
            search_clean = normalize_digits(str(search)).strip()
            qs = qs.filter(
                Q(name__icontains=search) |
                Q(code__icontains=search_clean) |
                Q(national_id__icontains=search_clean) |
                Q(ceo_name__icontains=search)
            )
        return qs.order_by('code')

    def create(self, request, *args, **kwargs):
        # پیش از اعتبارسنجی، تا خطای یکتایی کد یا شناسهٔ ملی وجود شرکت دیگر را لو ندهد
        if not request.user.is_superuser:
            raise PermissionDenied("تنها مدیر کل سامانه مجاز به ایجاد شرکت جدید است.")
        return super().create(request, *args, **kwargs)

    def perform_create(self, serializer):
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('company', 'create', instance.id, instance.name, company_id=instance.id,
                                        client_tab_id=tab_id, sender_id=self.request.user.id, notify_superusers=True)

    def perform_update(self, serializer):
        user = self.request.user
        if not user.is_superuser:
            # هر فیلد باید به بخشی نگاشت شده باشد که کاربر «نوشتن» آن را در همین شرکت دارد؛
            # فیلد بی‌نگاشت (مثل is_active و disabled_modules) فقط برای superuser است.
            from .permission_sections import field_sections
            mapping = field_sections()
            company_id = serializer.instance.id
            denied = []
            for field in serializer.validated_data:
                section = mapping.get(field)
                if section is None or company_section_level(user, company_id, section) != 'write':
                    denied.append(field)
            if denied:
                raise PermissionDenied(f"مجوز ویرایش این فیلدها را ندارید: {', '.join(sorted(denied))}")
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id') or self.request.data.get('client_tab_id')
        broadcast_org_structure_updated('company', 'update', instance.id, instance.name, company_id=instance.id,
                                        client_tab_id=tab_id, sender_id=self.request.user.id)

    def _ensure_can_delete(self, user):
        if not user.is_superuser:
            raise PermissionDenied("تنها مدیر کل سامانه مجاز به حذف شرکت است.")

    def check_delete_impact_permission(self, request, instance):
        # اثر حذف (شمار رکوردهای وابسته) فقط برای کسی است که اجازهٔ حذف دارد؛
        # get_object() پیش از این، شرکت بیرون از دامنهٔ کاربر را با ۴۰۴ رد کرده است.
        self._ensure_can_delete(request.user)

    def get_delete_blockers(self, instance):
        blockers = sorted(contributions('company_delete_blockers'), key=lambda b: b[0])

        def has_related(name):
            # رابطهٔ چندتایی exists() دارد؛ رابطهٔ یک‌به‌یکِ ناموجود در hasattr همان False است
            related = getattr(instance, name, None)
            return related.exists() if hasattr(related, 'exists') else related is not None

        return [message for _order, related_name, message in blockers if has_related(related_name)]

    def perform_destroy(self, instance):
        self._ensure_can_delete(self.request.user)
        blockers = self.get_delete_blockers(instance)
        if blockers:
            raise serializers.ValidationError(blockers[0])
        instance_id = instance.id
        instance_name = instance.name
        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company', 'delete', instance_id, instance_name, company_id=instance_id,
                                        client_tab_id=tab_id, sender_id=self.request.user.id)

    @action(detail=False, methods=['get'], url_path='user-available')
    def user_available(self, request):
        """
        لیست شرکت‌های در دسترس کاربر جاری برای مودال ورود، هدر سوئیچر و حالت محدود استودیو.
        scope='workspace' (پیش‌فرض): شرکت‌هایی که کاربر عضو آن‌هاست.
        scope='all': شرکت‌های پیوندخورده با کاربر.
        scope='manage': شرکت‌هایی که کاربر دست‌کم یک بخش آن‌ها را می‌خواند (بایگانی مدارک هم از همین راه).
        """
        user = request.user
        scope = request.query_params.get('scope', 'workspace')
        if scope == 'all':
            companies = get_user_allowed_companies(user, required_level=None)
        elif scope == 'manage':
            companies = companies_with_any_section(user)
        elif scope == 'workspace':
            companies = get_user_allowed_companies(user)
        else:
            # scope='documents' کلاینت کهنه با فهرست اعضا اشتباه گرفته نشود
            return Response({'detail': 'دامنهٔ فهرست شرکت‌ها نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        serializer = self.get_serializer(companies, many=True)
        data = list(serializer.data)

        # عضویت (با منبع) و سطح بخش‌ها برای هر شرکت
        from .permission_sections import section_map
        all_write = {k: 'write' for k in section_map()}
        levels = {} if user.is_superuser else user_section_levels(user)
        sources = {}
        if not user.is_superuser:
            for d in derived_company_access_details(user):
                sources.setdefault(d['company_id'], []).append(d.get('source_label') or d.get('company_name') or '')
        member_ids = None if user.is_superuser else set(get_user_allowed_companies(user).values_list('id', flat=True))
        for c in data:
            if user.is_superuser:
                c['is_member'] = True
                c['membership_sources'] = ['مدیر کل سامانه']
                c['section_levels'] = dict(all_write)
            else:
                c['is_member'] = c['id'] in member_ids
                c['membership_sources'] = sources.get(c['id'], [])
                c['section_levels'] = levels.get(c['id'], {})

        return Response({
            'companies': data,
            'is_superuser': user.is_superuser,
            'count': len(data)
        })

    IMPORT_FIELDS = (
        'code', 'name', 'company_type', 'national_id', 'economic_code', 'registration_number',
        'phone', 'ceo_name', 'board_chairman', 'address', 'is_active',
    )

    @action(detail=False, methods=['post'], url_path='import-excel',
            parser_classes=[parsers.MultiPartParser, parsers.FormParser])
    def import_excel(self, request):
        """
        ورود شرکت‌ها از اکسل دوسطری (ردیف ۱ عنوان فارسی، ردیف ۲ کلید فیلد؛ هم‌شکل export-excel).
        فقط ساخت شرکت تازه و فقط برای superuser (ت۵). کد تکراری رد و گزارش می‌شود؛ ردیف‌های سالم ثبت می‌شوند.
        """
        import openpyxl
        from django.db import transaction

        if not request.user.is_superuser:
            raise PermissionDenied("تنها مدیر کل سامانه مجاز به ورود شرکت از اکسل است.")
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایلی ارسال نشده است'}, status=400)
        dry_run = str(request.data.get('dry_run')).lower() == 'true'
        try:
            ws = openpyxl.load_workbook(file_obj, read_only=True, data_only=True).active
            rows = list(ws.iter_rows(values_only=True))
        except Exception:
            return Response({'error': 'فایل اکسل خوانا نیست.'}, status=400)
        if len(rows) < 2:
            return Response({'error': 'سربرگ دوسطری (عنوان فارسی و کلید فیلد) در فایل نیست.'}, status=400)
        keys = [str(k).strip() if k is not None else '' for k in rows[1]]
        if 'code' not in keys or 'name' not in keys:
            return Response({'error': 'ستون‌های code و name در ردیف دوم الزامی‌اند.'}, status=400)

        errors, preview, created, skipped = [], [], 0, 0
        data_rows = [(i + 3, r) for i, r in enumerate(rows[2:]) if any(v not in (None, '') for v in r)]
        for row_num, values in data_rows:
            payload = {}
            for key, value in zip(keys, values):
                if key in self.IMPORT_FIELDS and value not in (None, ''):
                    payload[key] = normalize_digits(str(value).strip()) if key != 'is_active' else value
            if 'is_active' in payload:
                payload['is_active'] = str(payload['is_active']).strip().lower() not in ('0', 'false', 'غیرفعال', 'no')
            serializer = self.get_serializer(data=payload)
            if not serializer.is_valid():
                skipped += 1
                for field, msgs in serializer.errors.items():
                    errors.append({'row': row_num, 'field': field, 'message': '؛ '.join(str(m) for m in msgs)})
                continue
            preview.append({'row': row_num, 'name': payload.get('name', ''), 'username': payload.get('code', '')})
            if not dry_run:
                with transaction.atomic():
                    instance = serializer.save()
                broadcast_org_structure_updated('company', 'create', instance.id, instance.name, company_id=instance.id,
                                                client_tab_id=request.headers.get('X-Client-Tab-Id'),
                                                sender_id=request.user.id, notify_superusers=True)
            created += 1
        return Response({
            'success': not errors,
            'dry_run': dry_run,
            'summary': {'total_rows': len(data_rows), 'created': created, 'updated': 0, 'skipped': skipped,
                        'valid_count': created, 'error_count': len(errors)},
            'preview_rows': preview,
            'errors': errors,
        })

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی اکسل استاندارد ۲ ردیفه اطلاعات شرکت‌ها منطبق بر استانداردهای مهندسی
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        # خروجی همهٔ شناسه‌ها و شبا فقط برای superuser (یافتهٔ tenant-scope-reviewer گام ۸)
        if not request.user.is_superuser:
            raise PermissionDenied("خروجی اکسل شرکت‌ها فقط برای مدیر کل سامانه است.")
        qs = self.get_queryset()
        if request.query_params.get('template') in ('1', 'true'):
            # الگوی خالی ورود از اکسل: همان سربرگ دوسطری بدون ردیف داده
            qs = qs.none()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "شرکت‌های هلدینگ"
        ws.views.sheetView[0].rightToLeft = True

        headers_fa = [
            'ردیف', 'کد یکتا', 'نام کامل شرکت', 'نوع شرکت', 'شناسه ملی', 'کد اقتصادی',
            'شماره ثبت', 'تلفن تماس', 'مدیرعامل', 'رئیس هیئت‌مدیره',
            'تعداد پروژه‌ها', 'تعداد اسناد', 'وضعیت', 'نشانی دفتر مرکزی'
        ]
        headers_en = [
            'row_num', 'code', 'name', 'company_type', 'national_id', 'economic_code',
            'registration_number', 'phone', 'ceo_name', 'board_chairman',
            'projects_count', 'documents_count', 'is_active', 'address'
        ]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(headers_fa) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 16
        ws.freeze_panes = 'A3'

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, comp in enumerate(qs, start=1):
            # شمار پروژه‌ها و اسناد فقط وقتی حسابداری نصب باشد معنا دارد
            if hasattr(comp, 'projects_count'):
                p_count = comp.projects_count
            else:
                p_count = comp.projects.count() if hasattr(comp, 'projects') else 0
            d_count = comp.documents.count() if hasattr(comp, 'documents') else 0
            ws.append([
                idx,
                comp.code,
                comp.name,
                comp.get_company_type_display(),
                comp.national_id or '',
                comp.economic_code or '',
                comp.registration_number or '',
                comp.phone or '',
                comp.ceo_name or '',
                comp.board_chairman or '',
                p_count,
                d_count,
                'فعال' if comp.is_active else 'غیرفعال',
                comp.address or ''
            ])
            curr_row = idx + 2
            ws.row_dimensions[curr_row].height = 20
            for col_idx in range(1, len(headers_fa) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal='center' if col_idx not in [3, 14] else 'right', vertical='center')
                if col_idx in [1, 2, 5, 6, 7, 8, 11, 12]:
                    c.font = data_font_num
                else:
                    c.font = data_font

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"companies_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class UserCompanyAccessViewSet(viewsets.ModelViewSet):
    """
    مدیریت انتساب دسترسی کاربران به شرکت‌ها
    """
    queryset = UserCompanyAccess.objects.all().select_related('user', 'company')
    serializer_class = UserCompanyAccessSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if not user.is_superuser:
            # ردیف‌های شرکت‌هایی که کاربر «دسترسی کاربران» را دست‌کم می‌خواند، و ردیف‌های خودش
            readable = set(companies_with_section(user, 'user_access', 'read').values_list('id', flat=True))
            qs = qs.filter(Q(company_id__in=readable) | Q(user=user))
        user_id = self.request.query_params.get('user_id')
        company_id = self.request.query_params.get('company_id')
        if user_id:
            qs = qs.filter(user_id=user_id)
        if company_id:
            qs = qs.filter(company_id=company_id)
        return qs.order_by('-created_at')

    @action(detail=False, methods=['get'], url_path='members')
    def members(self, request):
        """اعضای خودکار یک شرکت با منبع عضویت (فقط نمایش؛ ق۲). نیازمند «دسترسی کاربران» خواندن."""
        from django.contrib.auth import get_user_model
        try:
            company_id = int(request.query_params.get('company_id') or '')
        except ValueError:
            return Response({'error': 'company_id الزامی است.'}, status=400)
        ensure_section(request.user, company_id, 'user_access', 'read')
        sources = {}
        for provider in contributions('company_member_sources'):
            for uid, labels in provider(company_id).items():
                bucket = sources.setdefault(uid, [])
                bucket.extend(lb for lb in labels if lb not in bucket)
        users = get_user_model().objects.filter(id__in=sources.keys(), is_active=True).order_by('username')
        return Response([
            {'user_id': u.id, 'username': u.username,
             'user_full_name': f"{u.first_name} {u.last_name}".strip() or u.username,
             'sources': sources.get(u.id, [])}
            for u in users
        ])

    def _ensure_can_grant(self, company_id, target_user, old_levels, new_levels):
        """
        نوشتنِ «دسترسی کاربران» فقط تا سطح خودِ نویسنده (§۵-۱): نه ردیف خود، نه superuser،
        نه ردیف کسی که در بخشی از نویسنده بالاتر است؛ کارهای حساس فقط با superuser (ت۹).
        """
        from .permission_sections import LEVEL_RANK, section_map
        actor = self.request.user
        if actor.is_superuser:
            return
        ensure_section(actor, company_id, 'user_access', 'write', "مجوز نوشتن «دسترسی کاربران» این شرکت را ندارید.")
        if target_user is not None and (target_user.id == actor.id or target_user.is_superuser):
            raise PermissionDenied("دسترسی خودتان یا مدیر کل سامانه را نمی‌توانید تغییر دهید.")
        sections = section_map()
        mine = user_section_levels(actor).get(int(company_id), {})
        for levels in (old_levels or {}, new_levels or {}):
            for key, level in levels.items():
                spec = sections.get(key, {})
                if spec.get('sensitive') and level != 'none':
                    raise PermissionDenied("اعطا یا تغییر کارهای حساس فقط با مدیر کل سامانه است.")
                if LEVEL_RANK.get(level, 0) > LEVEL_RANK.get(mine.get(key, 'none'), 0):
                    raise PermissionDenied(f"سطح «{spec.get('label_fa', key)}» از سطح خود شما بالاتر است.")

    def perform_create(self, serializer):
        company = serializer.validated_data.get('company')
        target = serializer.validated_data.get('user')
        actor = self.request.user
        if not actor.is_superuser and target is not None:
            # طرح ۲.۹ (S1): «کشیدن» کاربر نادیده به شرکت خود فقط کار مدیر کل است (لایهٔ دوم پس از queryset سریالایزر)
            from accounts.scope import get_visible_users_for_user
            if not get_visible_users_for_user(actor).filter(pk=target.pk).exists():
                raise PermissionDenied("پیوند دادن کاربری که در قلمرو شما نیست فقط با مدیر کل سامانه است.")
        self._ensure_can_grant(company.id if company else None, target,
                               {}, serializer.validated_data.get('section_levels') or {})
        instance = serializer.save()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_access', 'create', instance.id, str(instance.user), client_tab_id=tab_id, sender_id=self.request.user.id, company_id=instance.company_id)

    def perform_update(self, serializer):
        inst = serializer.instance
        if 'company' in serializer.validated_data and serializer.validated_data['company'].id != inst.company_id:
            raise PermissionDenied("انتقال ردیف دسترسی به شرکت دیگر ممکن نیست؛ ردیف تازه بسازید.")
        if 'user' in serializer.validated_data and serializer.validated_data['user'].id != inst.user_id:
            raise PermissionDenied("تغییر کاربرِ ردیف دسترسی ممکن نیست.")
        from .access import _row_levels
        from .permission_sections import section_map
        self._ensure_can_grant(inst.company_id, inst.user, _row_levels(inst),
                               serializer.validated_data.get('section_levels') or {})
        extra = {}
        if 'section_levels' in serializer.validated_data:
            # کلید بخشِ ماژولِ نصب‌نشده به‌صورت داده می‌ماند و با ویرایش پاک نمی‌شود (§۹)
            installed = section_map()
            kept = {k: v for k, v in (inst.section_levels or {}).items() if k not in installed}
            extra['section_levels'] = {**kept, **serializer.validated_data['section_levels']}
        instance = serializer.save(**extra)
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_access', 'update', instance.id, str(instance.user), client_tab_id=tab_id, sender_id=self.request.user.id, company_id=instance.company_id)

    def perform_destroy(self, instance):
        from .access import _row_levels
        self._ensure_can_grant(instance.company_id, instance.user, _row_levels(instance), {})
        instance_id = instance.id
        user_title = str(instance.user)
        comp_id = instance.company_id
        instance.delete()
        tab_id = self.request.headers.get('X-Client-Tab-Id')
        broadcast_org_structure_updated('company_access', 'delete', instance_id, user_title, client_tab_id=tab_id, sender_id=self.request.user.id, company_id=comp_id)


@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def company_permission_sections(request):
    """بخش‌های مجوز شرکتِ نصب‌شده، سطح‌ها و قالب‌ها (برای جدول تب «دسترسی»)."""
    from .permission_sections import public_spec
    return Response(public_spec())
