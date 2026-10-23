"""
آزمون‌های API هستهٔ پروژه، بخش و انتساب (طرح Project_Section_To_Core، گام ۲):
سطح بخش projects (خواندن و نوشتن)، جداسازی شرکت‌ها، منبع عضویتِ انتساب در هسته و نبود نقش اعلام‌شده.
بی‌نیاز از هر ماژول؛ در پروفایل‌های WH_MODULES=warehouse و WH_MODULES=accounting هم اجرا می‌شود.
"""
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APIClient

from companies.access import derived_company_ids, get_user_allowed_companies
from companies.models import Company, Project, ProjectSection, SectionAssignment, UserCompanyAccess
from companies.org_membership import build_section_member_q, section_member_sources
from companies.permission_sections import template_levels
from platform_core import registry

User = get_user_model()

ROLES = (('employee', 'کارمند'), ('manager', 'مدیر'))


def with_roles(roles=ROLES):
    """نقش‌های انتساب را مستقل از ماژول‌های نصب‌شده تعیین می‌کند."""
    return mock.patch.dict(registry._CONTRIBUTIONS, {'section_roles': [roles] if roles else []})


def rows(res):
    body = res.json()
    return body.get('results', body) if isinstance(body, dict) else body


class OrgApiBase(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.a = Company.objects.create(code='ORG-A', name='شرکت الف')
        self.b = Company.objects.create(code='ORG-B', name='شرکت ب')
        self.pa = Project.objects.create(company=self.a, code='PA', name='پروژه الف')
        self.sa = ProjectSection.objects.create(project=self.pa, code='SA', name='بخش الف')
        self.pb = Project.objects.create(company=self.b, code='PB', name='پروژه ب')
        self.sb = ProjectSection.objects.create(project=self.pb, code='SB', name='بخش ب')
        self.writer = self._user('org_writer', self.a, {'projects': 'write'})
        self.reader = self._user('org_reader', self.a, {'projects': 'read'})
        self.nolevel = self._user('org_nolevel', self.a, {'identity': 'read'})
        self.linked = self._user('org_linked', self.a, {})
        self.outsider = self._user('org_outsider', self.b, {})

    def _user(self, username, company, levels):
        u = User.objects.create_user(username=username, password='x-pass-123')
        UserCompanyAccess.objects.create(user=u, company=company, section_levels=levels)
        return u

    def _as(self, user):
        self.client.force_authenticate(user=user)


class ProjectsSectionLevelTests(OrgApiBase):
    def test_templates_carry_projects_section(self):
        self.assertEqual(template_levels('manager')['projects'], 'write')
        self.assertEqual(template_levels('viewer')['projects'], 'read')

    def test_read_level_lists_but_cannot_write(self):
        self._as(self.reader)
        res = self.client.get('/api/org/projects/', {'company_id': self.a.id})
        self.assertEqual(res.status_code, 200)
        self.assertEqual({r['code'] for r in rows(res)}, {'PA'})
        self.assertEqual({r['code'] for r in rows(self.client.get('/api/org/project-sections/', {'company_id': self.a.id}))}, {'SA'})
        self.assertEqual(self.client.post('/api/org/projects/', {'company': self.a.id, 'code': 'PX', 'name': 'x'}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.pa.id}/', {'name': 'تغییر'}, format='json').status_code, 403)
        self.assertEqual(self.client.delete(f'/api/org/projects/{self.pa.id}/').status_code, 403)
        self.assertEqual(self.client.post('/api/org/project-sections/', {'project': self.pa.id, 'code': 'SX', 'name': 'x'}, format='json').status_code, 403)
        with with_roles():
            res = self.client.post('/api/org/section-assignments/', {'user': self.linked.id, 'section': self.sa.id, 'role': 'employee'}, format='json')
        self.assertEqual(res.status_code, 403)
        self.pa.refresh_from_db()
        self.assertEqual(self.pa.name, 'پروژه الف')
        self.assertTrue(Project.objects.filter(pk=self.pa.pk).exists())

    def test_user_without_level_gets_403_or_empty(self):
        self._as(self.nolevel)
        self.assertEqual(self.client.get('/api/org/projects/', {'company_id': self.a.id}).status_code, 403)
        self.assertEqual(self.client.get('/api/org/project-sections/', {'company_id': self.a.id}).status_code, 403)
        self.assertEqual(rows(self.client.get('/api/org/projects/')), [])
        self.assertEqual(rows(self.client.get('/api/org/project-sections/')), [])
        self.assertEqual(rows(self.client.get('/api/org/section-assignments/')), [])
        self.assertEqual(self.client.post('/api/org/projects/', {'company': self.a.id, 'code': 'PX', 'name': 'x'}, format='json').status_code, 403)

    def test_write_level_creates_updates_and_assigns(self):
        self._as(self.writer)
        res = self.client.post('/api/org/projects/', {'company': self.a.id, 'code': 'PN', 'name': 'پروژهٔ تازه'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        for key in ('id', 'company', 'company_name', 'company_code', 'sections_count', 'code', 'name', 'is_active'):
            self.assertIn(key, res.json())
        res = self.client.post('/api/org/project-sections/', {'project': res.json()['id'], 'code': 'SN', 'name': 'بخش تازه'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.json()['company_id'], self.a.id)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.pa.id}/', {'name': 'نام تازه'}, format='json').status_code, 200)
        with with_roles():
            res = self.client.post('/api/org/section-assignments/', {'user': self.linked.id, 'section': self.sa.id, 'role': 'manager'}, format='json')
            self.assertEqual(res.status_code, 201, res.content)
            self.assertEqual(res.json()['role_display'], 'مدیر')
            # کاربرِ بی‌پیوند با شرکت الف منتسب نمی‌شود (انتساب خودش عضویت می‌دهد)
            res = self.client.post('/api/org/section-assignments/', {'user': self.outsider.id, 'section': self.sa.id, 'role': 'employee'}, format='json')
            self.assertEqual(res.status_code, 403)
            # نویسنده خودش را منتسب نمی‌کند
            res = self.client.post('/api/org/section-assignments/', {'user': self.writer.id, 'section': self.sa.id, 'role': 'manager'}, format='json')
            self.assertEqual(res.status_code, 403)
        self.assertFalse(SectionAssignment.objects.filter(user__in=[self.outsider, self.writer]).exists())

    def test_global_project_is_superuser_only(self):
        self._as(self.writer)
        self.assertEqual(self.client.post('/api/org/projects/', {'code': 'PG', 'name': 'سراسری'}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.pa.id}/', {'company': None}, format='json').status_code, 403)
        admin = User.objects.create_superuser(username='org_admin_su', password='x-pass-123')
        self._as(admin)
        self.assertEqual(self.client.post('/api/org/projects/', {'code': 'PG', 'name': 'سراسری'}, format='json').status_code, 201)
        self.assertEqual({r['code'] for r in rows(self.client.get('/api/org/projects/'))}, {'PA', 'PB', 'PG'})

    def test_no_contributed_roles_rejects_assignment(self):
        self._as(self.writer)
        with with_roles(None):
            res = self.client.post('/api/org/section-assignments/', {'user': self.linked.id, 'section': self.sa.id}, format='json')
            self.assertEqual(res.status_code, 400)
            self.assertIn('نقش انتساب', str(res.json()))
            res = self.client.post('/api/org/section-assignments/bulk-assign/', {'section_id': self.sa.id, 'user_ids': [self.linked.id]}, format='json')
            self.assertEqual(res.status_code, 400)
        self.assertFalse(SectionAssignment.objects.exists())

    def test_unused_project_is_deleted_and_sections_cascade(self):
        self._as(self.writer)
        self.assertEqual(self.client.delete(f'/api/org/projects/{self.pa.id}/').status_code, 204)
        self.assertFalse(Project.objects.filter(pk=self.pa.pk).exists())
        self.assertFalse(ProjectSection.objects.filter(pk=self.sa.pk).exists())


class ReviewFindingTests(OrgApiBase):
    """یافته‌های tenant-scope-reviewer گام ۲."""

    def test_inactive_assignment_cannot_link_unlinked_user(self):
        # انتساب غیرفعال کاربر را «پیوندخورده» می‌کرد و فعال‌سازی بعدی عضویت می‌داد
        self._as(self.writer)
        base = '/api/org/section-assignments/'
        with with_roles():
            res = self.client.post(base, {'user': self.outsider.id, 'section': self.sa.id, 'role': 'employee', 'is_active': False}, format='json')
            self.assertEqual(res.status_code, 403)
            self.assertEqual(self.client.post(base, {'user': self.outsider.id, 'section': self.sa.id, 'role': 'employee'}, format='json').status_code, 403)
            # ردیف موجودِ کاربر پیوندخورده را می‌توان بی‌شرط غیرفعال کرد، ولی نه به کاربر بی‌پیوند منتقل کرد
            own = SectionAssignment.objects.create(user=self.linked, section=self.sa, role='employee', is_active=True)
            res = self.client.patch(f'{base}{own.id}/', {'user': self.outsider.id, 'is_active': False}, format='json')
            self.assertEqual(res.status_code, 403)
            self.assertEqual(self.client.patch(f'{base}{own.id}/', {'is_active': False}, format='json').status_code, 200)
        self.assertFalse(SectionAssignment.objects.filter(user=self.outsider).exists())
        self.assertNotIn(self.a.id, set(get_user_allowed_companies(self.outsider).values_list('id', flat=True)))

    def test_staff_my_sections_is_limited_to_own_assignments(self):
        staff = User.objects.create_user(username='org_staff', password='x-pass-123', is_staff=True)
        self._as(staff)
        self.assertEqual(self.client.get('/api/org/section-assignments/my-sections/').json(), [])
        SectionAssignment.objects.create(user=staff, section=self.sa, role='employee', is_active=True)
        self.assertEqual([s['code'] for s in self.client.get('/api/org/section-assignments/my-sections/').json()], ['SA'])

    def test_inactive_company_gives_no_section_level(self):
        from companies.access import company_section_level
        self.a.is_active = False
        self.a.save(update_fields=['is_active'])
        self.assertEqual(company_section_level(self.writer, self.a.id, 'projects'), 'none')
        self._as(self.writer)
        self.assertEqual(self.client.post('/api/org/projects/', {'company': self.a.id, 'code': 'PX', 'name': 'x'}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.pa.id}/', {'name': 'x'}, format='json').status_code, 404)
        self.pa.refresh_from_db()
        self.assertEqual(self.pa.name, 'پروژه الف')


class CrossCompanyIsolationTests(OrgApiBase):
    def test_writer_of_a_cannot_read_b(self):
        SectionAssignment.objects.create(user=self.outsider, section=self.sb, role='employee')
        self._as(self.writer)
        self.assertEqual(self.client.get('/api/org/projects/', {'company_id': self.b.id}).status_code, 403)
        self.assertEqual(self.client.get('/api/org/project-sections/', HTTP_X_COMPANY_ID=str(self.b.id)).status_code, 403)
        self.assertEqual(self.client.get('/api/org/section-assignments/', {'company_id': self.b.id}).status_code, 403)
        self.assertNotIn('PB', {r['code'] for r in rows(self.client.get('/api/org/projects/'))})
        self.assertNotIn('SB', {r['code'] for r in rows(self.client.get('/api/org/project-sections/'))})
        self.assertEqual(rows(self.client.get('/api/org/section-assignments/')), [])
        self.assertEqual(self.client.get(f'/api/org/projects/{self.pb.id}/').status_code, 404)

    def test_writer_of_a_cannot_write_b(self):
        foreign = SectionAssignment.objects.create(user=self.outsider, section=self.sb, role='employee')
        self._as(self.writer)
        self.assertEqual(self.client.post('/api/org/projects/', {'company': self.b.id, 'code': 'PX', 'name': 'x'}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.pb.id}/', {'name': 'x'}, format='json').status_code, 404)
        self.assertEqual(self.client.delete(f'/api/org/project-sections/{self.sb.id}/').status_code, 404)
        self.assertEqual(self.client.post('/api/org/project-sections/', {'project': self.pb.id, 'code': 'SX', 'name': 'x'}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.pa.id}/', {'company': self.b.id}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org/project-sections/{self.sa.id}/', {'project': self.pb.id}, format='json').status_code, 403)
        self.assertEqual(self.client.delete(f'/api/org/section-assignments/{foreign.id}/').status_code, 404)
        with with_roles():
            res = self.client.post('/api/org/section-assignments/', {'user': self.outsider.id, 'section': self.sb.id, 'role': 'employee'}, format='json')
            self.assertEqual(res.status_code, 403)
            res = self.client.post('/api/org/section-assignments/bulk-assign/', {'section_id': self.sb.id, 'user_ids': [self.linked.id], 'role': 'employee'}, format='json')
            self.assertEqual(res.status_code, 403)
            res = self.client.post(f'/api/org/project-sections/{self.sa.id}/clone-assignments/', {'source_section_id': self.sb.id}, format='json')
            self.assertEqual(res.status_code, 403)
        foreign.refresh_from_db()
        self.assertTrue(foreign.is_active)
        self.pb.refresh_from_db()
        self.assertEqual(self.pb.name, 'پروژه ب')
        self.assertEqual(SectionAssignment.objects.count(), 1)

    def test_export_contains_only_readable_companies(self):
        import io
        import openpyxl
        self._as(self.reader)
        res = self.client.get('/api/org/projects/export-excel/')
        self.assertEqual(res.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        self.assertEqual({ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)}, {'PA'})
        self.assertEqual(self.client.get('/api/org/projects/export-excel/', {'company_id': self.b.id}).status_code, 403)


class SectionMembershipInCoreTests(OrgApiBase):
    def test_active_assignment_grants_membership(self):
        self.assertNotIn(self.a.id, derived_company_ids(self.linked))
        sa = SectionAssignment.objects.create(user=self.linked, section=self.sa, role='employee', is_active=True)
        self.assertIn(self.a.id, set(get_user_allowed_companies(self.linked).values_list('id', flat=True)))
        self.assertIn(self.linked.id, section_member_sources(self.a.id))
        self.assertTrue(User.objects.filter(build_section_member_q([self.a.id]), id=self.linked.id).exists())
        sa.is_active = False
        sa.save(update_fields=['is_active'])
        self.assertNotIn(self.a.id, derived_company_ids(self.linked))

    def test_member_reads_projects_without_level(self):
        SectionAssignment.objects.create(user=self.linked, section=self.sa, role='employee', is_active=True)
        self._as(self.linked)
        self.assertEqual({r['code'] for r in rows(self.client.get('/api/org/projects/'))}, {'PA'})
        res = self.client.get('/api/org/section-assignments/my-sections/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual([s['code'] for s in res.json()], ['SA'])
        # عضویت بی سطح «نوشتن» نوشتن نمی‌دهد
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.pa.id}/', {'name': 'x'}, format='json').status_code, 403)

    def test_company_with_project_cannot_be_deleted(self):
        admin = User.objects.create_superuser(username='org_del_su', password='x-pass-123')
        self._as(admin)
        self.assertEqual(self.client.delete(f'/api/org/companies/{self.a.id}/').status_code, 400)
        self.assertTrue(Company.objects.filter(pk=self.a.pk).exists())
