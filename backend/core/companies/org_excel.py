"""
موتور استاندارد ورود و خروج اکسل پروژه‌ها، بخش‌ها و انتساب کاربران به بخش در هسته
(طرح Project_Section_To_Core، گام ۲؛ طرف‌حساب‌ها در حسابداری می‌مانند و helperهای سربرگ همین فایل را به کار می‌برند):
- ساختار ۲ سطری هدر (سطر ۱: عناوین فارسی، سطر ۲: کلیدهای دیتابیسی با فونت طوسی)
- فریز پنل در A3 (freeze_panes = 'A3')
- راست‌به‌چپ (rightToLeft = True)
- تبدیل زمان‌ها به تاریخ و زمان شمسی
- پشتیبانی کامل از پیش‌نمایش واقعی (dry_run) بدون ایجاد رکورد در دیتابیس
"""

import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from django.http import HttpResponse
from django.db import transaction
from django.db.models import Count, Q
import jdatetime

from django.contrib.auth import get_user_model
from common.date_utils import normalize_digits
from .models import Company, Project, ProjectSection, SectionAssignment, section_role_choices
from .org_roles import assignments_in_companies_q

User = get_user_model()


def _format_datetime_shamsi(dt):
    if not dt:
        return '-'
    try:
        jdt = jdatetime.datetime.fromgregorian(datetime=dt)
        return jdt.strftime('%Y/%m/%d %H:%M:%S')
    except Exception:
        return str(dt)


def _apply_2row_header(ws, columns, header_bg='059669'):
    """
    اعمال هدر دو سطری استاندارد شرکت:
    سطر ۱: عناوین فارسی
    سطر ۲: کلیدهای سیستمی دیتابیسی
    """
    f_title = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
    f_key = Font(name='Segoe UI', size=9, bold=True, color='CBD5E1')
    fill_title = PatternFill(start_color=header_bg, end_color=header_bg, fill_type='solid')
    fill_key = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 20
    ws.freeze_panes = 'A3'
    ws.sheet_view.rightToLeft = True

    for col_idx, col in enumerate(columns, 1):
        # سطر ۱
        c1 = ws.cell(row=1, column=col_idx, value=col['title'])
        c1.font = f_title
        c1.fill = fill_title
        c1.alignment = align_center
        c1.border = thin_border

        # سطر ۲
        c2 = ws.cell(row=2, column=col_idx, value=col['key'])
        c2.font = f_key
        c2.fill = fill_key
        c2.alignment = align_center
        c2.border = thin_border

        ws.column_dimensions[get_column_letter(col_idx)].width = col.get('width', 18)


# -------------------------------------------------------------
# ۱. پروژه‌های مالی (Financial Projects)
# -------------------------------------------------------------
PROJECT_COLUMNS = [
    {'title': 'کد پروژه', 'key': 'code', 'width': 15},
    {'title': 'نام پروژه مالی / عملیاتی', 'key': 'name', 'width': 28},
    {'title': 'توضیحات و دامنه فعالیت', 'key': 'description', 'width': 35},
    {'title': 'وضعیت فعال (بله/خیر)', 'key': 'is_active', 'width': 16},
    {'title': 'تعداد بخش‌ها', 'key': 'sections_count', 'width': 14},
    {'title': 'تاریخ ایجاد', 'key': 'created_at', 'width': 22},
]

def export_projects_excel(company_ids=None):
    """company_ids تهی یعنی همهٔ پروژه‌ها (سوپریوزر)؛ وگرنه فقط پروژه‌های همان شرکت‌ها."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'پروژه‌های مالی'
    _apply_2row_header(ws, PROJECT_COLUMNS, header_bg='4338CA')  # رنگ ایندیگو

    font_data = Font(name='B Nazanin', size=11)
    align_data = Alignment(horizontal='center', vertical='center')
    align_text = Alignment(horizontal='right', vertical='center')

    # بهینه‌سازی N+1 کوئری با Count('sections')
    projects = Project.objects.annotate(sections_count=Count('sections'))
    if company_ids is not None:
        projects = projects.filter(company_id__in=company_ids)
    for row_idx, p in enumerate(projects.order_by('code'), 3):
        ws.cell(row=row_idx, column=1, value=p.code).alignment = align_data
        ws.cell(row=row_idx, column=2, value=p.name).alignment = align_text
        ws.cell(row=row_idx, column=3, value=p.description or '-').alignment = align_text
        ws.cell(row=row_idx, column=4, value='بله' if p.is_active else 'خیر').alignment = align_data
        ws.cell(row=row_idx, column=5, value=p.sections_count).alignment = align_data
        ws.cell(row=row_idx, column=6, value=_format_datetime_shamsi(p.created_at)).alignment = align_data

        for c in range(1, 7):
            ws.cell(row=row_idx, column=c).font = font_data

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="financial_projects.xlsx"'
    return response


# -------------------------------------------------------------
# ۲. بخش‌ها و دپارتمان‌ها (Project Sections)
# -------------------------------------------------------------
SECTION_COLUMNS = [
    {'title': 'کد پروژه مادر', 'key': 'project_code', 'width': 16},
    {'title': 'نام پروژه مادر', 'key': 'project_name', 'width': 24},
    {'title': 'کد بخش', 'key': 'code', 'width': 15},
    {'title': 'نام بخش / دپارتمان', 'key': 'name', 'width': 28},
    {'title': 'وضعیت فعال (بله/خیر)', 'key': 'is_active', 'width': 16},
    {'title': 'تاریخ ایجاد', 'key': 'created_at', 'width': 22},
]

def export_sections_excel(project_id=None, company_ids=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'بخش‌های پروژه'
    _apply_2row_header(ws, SECTION_COLUMNS, header_bg='0D9488')  # سبز کله‌غازی

    qs = ProjectSection.objects.select_related('project').all()
    if company_ids is not None:
        qs = qs.filter(project__company_id__in=company_ids)
    if project_id:
        try:
            qs = qs.filter(project_id=int(project_id))
        except (ValueError, TypeError):
            pass
    qs = qs.order_by('project__code', 'code')

    font_data = Font(name='B Nazanin', size=11)
    align_data = Alignment(horizontal='center', vertical='center')
    align_text = Alignment(horizontal='right', vertical='center')

    for row_idx, s in enumerate(qs, 3):
        ws.cell(row=row_idx, column=1, value=s.project.code if s.project else '-').alignment = align_data
        ws.cell(row=row_idx, column=2, value=s.project.name if s.project else '-').alignment = align_text
        ws.cell(row=row_idx, column=3, value=s.code).alignment = align_data
        ws.cell(row=row_idx, column=4, value=s.name).alignment = align_text
        ws.cell(row=row_idx, column=5, value='بله' if s.is_active else 'خیر').alignment = align_data
        ws.cell(row=row_idx, column=6, value=_format_datetime_shamsi(s.created_at)).alignment = align_data

        for c in range(1, 7):
            ws.cell(row=row_idx, column=c).font = font_data

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="project_sections.xlsx"'
    return response



# -------------------------------------------------------------
# ۴. انتساب‌های سازمانی (User Section Assignments)
# -------------------------------------------------------------
ASSIGNMENT_COLUMNS = [
    {'title': 'نام کاربری', 'key': 'username', 'width': 18},
    {'title': 'نام و نام‌خانوادگی', 'key': 'full_name', 'width': 24},
    {'title': 'کد پروژه', 'key': 'project_code', 'width': 16},
    {'title': 'کد بخش', 'key': 'section_code', 'width': 16},
    {'title': 'نقش سازمانی', 'key': 'role', 'width': 18},
    {'title': 'عنوان نقش', 'key': 'role_display', 'width': 22},
    {'title': 'وضعیت فعال', 'key': 'is_active', 'width': 14},
    {'title': 'تاریخ انتساب', 'key': 'created_at', 'width': 22},
    {'title': 'کد شرکت', 'key': 'company_code', 'width': 16},
    {'title': 'سطح انتساب', 'key': 'level', 'width': 14},
]

# سطح انتساب (طرح ۲.۱، ت۳): کد پروژه و کد بخش خالی یعنی سطح شرکت؛ فقط کد پروژه یعنی سطح پروژه
LEVEL_LABELS = {'company': 'شرکت', 'project': 'پروژه', 'section': 'بخش'}


def export_assignments_excel(project_id=None, role=None, company_ids=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ماتریس انتساب‌های سازمانی'
    _apply_2row_header(ws, ASSIGNMENT_COLUMNS, header_bg='6366F1')

    qs = SectionAssignment.objects.select_related(
        'user', 'company', 'project__company', 'section__project__company'
    ).all()
    if company_ids is not None:
        qs = qs.filter(assignments_in_companies_q(company_ids))
    if project_id:
        try:
            pid = int(project_id)
            qs = qs.filter(Q(project_id=pid) | Q(section__project_id=pid))
        except (ValueError, TypeError):
            pass
    if role and role != 'all':
        qs = qs.filter(role=role)
    qs = qs.order_by('company__code', 'project__code', 'section__project__code', 'section__code', 'role')

    font_data = Font(name='B Nazanin', size=11)
    align_data = Alignment(horizontal='center', vertical='center')
    align_text = Alignment(horizontal='right', vertical='center')

    role_map = dict(section_role_choices())

    for row_idx, a in enumerate(qs, 3):
        u = a.user
        full_name = f"{u.first_name} {u.last_name}".strip() if u else '-'
        ws.cell(row=row_idx, column=1, value=u.username if u else '-').alignment = align_data
        ws.cell(row=row_idx, column=2, value=full_name).alignment = align_text
        project = a.target_project
        company = a.company if a.company_id else project.company
        # کد خالی (نه «-») تا خروجی بی‌تغییر دوباره درون‌ریزی شود
        ws.cell(row=row_idx, column=3, value=project.code if project else '').alignment = align_data
        ws.cell(row=row_idx, column=4, value=a.section.code if a.section_id else '').alignment = align_data
        ws.cell(row=row_idx, column=5, value=a.role).alignment = align_data
        ws.cell(row=row_idx, column=6, value=role_map.get(a.role, a.role)).alignment = align_data
        ws.cell(row=row_idx, column=7, value='بله' if a.is_active else 'خیر').alignment = align_data
        ws.cell(row=row_idx, column=8, value=_format_datetime_shamsi(a.created_at)).alignment = align_data
        ws.cell(row=row_idx, column=9, value=company.code if company else '').alignment = align_data
        ws.cell(row=row_idx, column=10, value=LEVEL_LABELS[a.level]).alignment = align_data

        for c in range(1, 11):
            ws.cell(row=row_idx, column=c).font = font_data

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="assignments_matrix.xlsx"'
    return response


# -------------------------------------------------------------
# دانلود قالب‌های آماده برای اکسل (Templates)
# -------------------------------------------------------------
def _yes_no_list(wb, column=1):
    """شیت Lists با «بله/خیر» در ستون داده‌شده (برای منوی کشویی وضعیت فعال)."""
    ws_lists = wb['Lists'] if 'Lists' in wb.sheetnames else wb.create_sheet(title='Lists')
    ws_lists.sheet_view.rightToLeft = True
    ws_lists.cell(row=1, column=column, value='بله')
    ws_lists.cell(row=2, column=column, value='خیر')
    return ws_lists


def download_org_template(entity_type):
    """قالب خالی ورود از اکسل برای 'projects'، 'sections' یا 'assignments'."""
    wb = openpyxl.Workbook()
    ws = wb.active

    if entity_type == 'projects':
        ws.title = 'قالب پروژه‌ها'
        _apply_2row_header(ws, PROJECT_COLUMNS[:4], header_bg='4338CA')
        # ردیف نمونه
        ws.cell(row=3, column=1, value='PRJ-01')
        ws.cell(row=3, column=2, value='پروژه عملیات مرکزی')
        ws.cell(row=3, column=3, value='مدیریت کل عملیات و انبارها')
        ws.cell(row=3, column=4, value='بله')
        _yes_no_list(wb)
        dv_active = DataValidation(type="list", formula1="Lists!$A$1:$A$2", allow_blank=True)
        ws.add_data_validation(dv_active)
        dv_active.add("D3:D500")
        filename = 'projects_template.xlsx'

    elif entity_type == 'sections':
        ws.title = 'قالب بخش‌ها'
        _apply_2row_header(ws, SECTION_COLUMNS[:5], header_bg='0D9488')
        ws.cell(row=3, column=1, value='PRJ-01')
        ws.cell(row=3, column=2, value='پروژه عملیات مرکزی')
        ws.cell(row=3, column=3, value='SEC-101')
        ws.cell(row=3, column=4, value='دپارتمان لجستیک و ترابری')
        ws.cell(row=3, column=5, value='بله')
        _yes_no_list(wb)
        dv_active = DataValidation(type="list", formula1="Lists!$A$1:$A$2", allow_blank=True)
        ws.add_data_validation(dv_active)
        dv_active.add("E3:E500")
        filename = 'sections_template.xlsx'

    else:
        ws.title = 'قالب انتساب پرسنل'
        _apply_2row_header(ws, ASSIGNMENT_COLUMNS[:7], header_bg='6366F1')
        # نقش‌ها از ماژول‌های نصب‌شده (section_role_choices)؛ ستون ۵ کد نقش و ستون ۶ عنوان آن است
        roles = list(section_role_choices())
        sample_code, sample_label = roles[0] if roles else ('', '')
        ws.cell(row=3, column=1, value='admin')
        ws.cell(row=3, column=2, value='مدیر سیستم')
        ws.cell(row=3, column=3, value='PRJ-01')
        ws.cell(row=3, column=4, value='SEC-101')
        ws.cell(row=3, column=5, value=sample_code)
        ws.cell(row=3, column=6, value=sample_label)
        ws.cell(row=3, column=7, value='بله')
        # نمونهٔ سطح پروژه (کد بخش خالی) و سطح شرکت (کد پروژه و بخش خالی)
        for row, proj in ((4, 'PRJ-01'), (5, '')):
            ws.cell(row=row, column=1, value='admin')
            ws.cell(row=row, column=2, value='مدیر سیستم')
            ws.cell(row=row, column=3, value=proj)
            ws.cell(row=row, column=5, value=sample_code)
            ws.cell(row=row, column=6, value=sample_label)
            ws.cell(row=row, column=7, value='بله')
        ws_lists = _yes_no_list(wb, column=3)
        for idx, (code, label) in enumerate(roles, 1):
            ws_lists.cell(row=idx, column=1, value=code)
            ws_lists.cell(row=idx, column=2, value=label)
        if roles:
            dv_code = DataValidation(type="list", formula1=f"Lists!$A$1:$A${len(roles)}", allow_blank=True)
            ws.add_data_validation(dv_code)
            dv_code.add("E3:E500")
            dv_role = DataValidation(type="list", formula1=f"Lists!$B$1:$B${len(roles)}", allow_blank=True)
            ws.add_data_validation(dv_role)
            dv_role.add("F3:F500")
        dv_active = DataValidation(type="list", formula1="Lists!$C$1:$C$2", allow_blank=True)
        ws.add_data_validation(dv_active)
        dv_active.add("G3:G500")
        filename = 'assignments_template.xlsx'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def import_projects_from_excel(file_obj, dry_run=False, company_id=None):
    """
    خواندن و ثبت دسته‌جمعی پروژه‌های مالی از اکسل با پشتیبانی از dry_run.
    تطبیق و ساخت فقط در شرکت company_id (کد پروژه در هر شرکت یکتاست). پروژهٔ بی‌شرکت ساخته نمی‌شود (طرح ۲.۱، ت۲).
    """
    if company_id is None:
        return {
            'success': False,
            'dry_run': dry_run,
            'summary': {'total_rows': 0, 'valid_count': 0, 'error_count': 1, 'created': 0, 'updated': 0, 'skipped': 1},
            'errors': [{'row': 0, 'message': 'شرکت مقصد پروژه‌ها مشخص نیست؛ ابتدا شرکت را در سربرگ انتخاب کنید.'}]
        }
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        ws = wb.active
    except Exception as e:
        return {
            'success': False,
            'dry_run': dry_run,
            'summary': {'total_rows': 0, 'created': 0, 'updated': 0, 'skipped': 0},
            'errors': [{'row': 0, 'message': f'فایل اکسل نامعتبر یا آسیب‌دیده است: {str(e)}'}]
        }

    created = 0
    updated = 0
    errors = []

    for row_idx in range(3, ws.max_row + 1):
        code = ws.cell(row=row_idx, column=1).value
        if not code or str(code).strip() == '':
            continue
        code = normalize_digits(str(code)).strip().upper()
        name = str(ws.cell(row=row_idx, column=2).value or '').strip()
        if not name:
            errors.append(f"ردیف {row_idx}: نام پروژه الزامی است.")
            continue
        desc = str(ws.cell(row=row_idx, column=3).value or '').strip()
        active_val = str(ws.cell(row=row_idx, column=4).value or 'بله').strip().lower()
        is_active = active_val not in ['خیر', 'false', '0', 'no']

        lookup = {'code': code, 'company_id': company_id}
        try:
            exists = Project.objects.filter(**lookup).exists()
            if dry_run:
                if exists:
                    updated += 1
                else:
                    created += 1
            else:
                with transaction.atomic():
                    p, is_new = Project.objects.update_or_create(
                        **lookup,
                        defaults={'name': name, 'description': desc or None, 'is_active': is_active}
                    )
                    if is_new:
                        created += 1
                    else:
                        updated += 1
        except Exception as e:
            errors.append(f"ردیف {row_idx} ({code}): {str(e)}")

    return {
        'success': len(errors) == 0,
        'dry_run': dry_run,
        'summary': {
            'total_rows': created + updated + len(errors),
            'valid_count': created + updated,
            'error_count': len(errors),
            'created': created,
            'updated': updated,
            'skipped': len(errors)
        },
        'errors': [{'row': idx, 'message': err} for idx, err in enumerate(errors, 1)]
    }


def import_sections_from_excel(file_obj, dry_run=False, company_id=None):
    """
    خواندن و ثبت دسته‌جمعی بخش‌های پروژه از اکسل با پشتیبانی از dry_run.
    company_id داده‌شده: پروژهٔ مادر فقط از پروژه‌های همان شرکت پیدا می‌شود.
    """
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        ws = wb.active
    except Exception as e:
        return {
            'success': False,
            'dry_run': dry_run,
            'summary': {'total_rows': 0, 'valid_count': 0, 'error_count': 1, 'created': 0, 'updated': 0, 'skipped': 1},
            'errors': [{'row': 0, 'message': f'فایل اکسل نامعتبر یا آسیب‌دیده است: {str(e)}'}]
        }

    created = 0
    updated = 0
    errors = []

    for row_idx in range(3, ws.max_row + 1):
        proj_code = ws.cell(row=row_idx, column=1).value
        if not proj_code or str(proj_code).strip() == '':
            continue
        proj_code = normalize_digits(str(proj_code)).strip().upper()
        proj_qs = Project.objects.filter(code=proj_code)
        if company_id is not None:
            proj_qs = proj_qs.filter(company_id=company_id)
        proj = proj_qs.first()
        if not proj:
            scope_hint = "در شرکت فعال " if company_id is not None else ""
            errors.append(f"ردیف {row_idx}: پروژه با کد «{proj_code}» {scope_hint}یافت نشد.")
            continue

        code = ws.cell(row=row_idx, column=3).value
        if not code or str(code).strip() == '':
            continue
        code = normalize_digits(str(code)).strip().upper()
        name = str(ws.cell(row=row_idx, column=4).value or '').strip()
        if not name:
            errors.append(f"ردیف {row_idx}: نام بخش الزامی است.")
            continue

        active_val = str(ws.cell(row=row_idx, column=5).value or 'بله').strip().lower()
        is_active = active_val not in ['خیر', 'false', '0', 'no']

        try:
            exists = ProjectSection.objects.filter(project=proj, code=code).exists()
            if dry_run:
                if exists:
                    updated += 1
                else:
                    created += 1
            else:
                with transaction.atomic():
                    s, is_new = ProjectSection.objects.update_or_create(
                        project=proj,
                        code=code,
                        defaults={'name': name, 'is_active': is_active}
                    )
                    if is_new:
                        created += 1
                    else:
                        updated += 1
        except Exception as e:
            errors.append(f"ردیف {row_idx} ({code}): {str(e)}")

    return {
        'success': len(errors) == 0,
        'dry_run': dry_run,
        'summary': {
            'total_rows': created + updated + len(errors),
            'valid_count': created + updated,
            'error_count': len(errors),
            'created': created,
            'updated': updated,
            'skipped': len(errors)
        },
        'errors': [{'row': idx, 'message': err} for idx, err in enumerate(errors, 1)]
    }


def _assignment_row_target(proj_code, sec_code, company_code, company_id):
    """
    هدف یک ردیف اکسل انتساب (طرح ۲.۱، ت۳): (فیلد هدف، شیء هدف، پیام خطا).
    کد بخش: سطح بخش؛ فقط کد پروژه: سطح پروژه؛ هر دو خالی: سطح شرکت. شرکت: شرکت فعال درون‌ریزی (company_id)،
    یا برای superuser بی شرکت فعال، ستون «کد شرکت». کد شرکتِ دیگر در ردیف رد می‌شود.
    """
    if company_id is not None:
        company = Company.objects.filter(id=company_id).first()
        if company is None or (company_code and company_code != company.code.upper()):
            return None, None, "این ردیف به شرکت دیگری تعلق دارد؛ فقط انتساب‌های شرکت فعال درون‌ریزی می‌شوند."
    elif company_code:
        company = Company.objects.filter(code__iexact=company_code).first()
        if company is None:
            return None, None, f"شرکت با کد «{company_code}» یافت نشد."
    else:
        company = None

    if sec_code:
        sec_qs = ProjectSection.objects.filter(code=sec_code).select_related('project')
        if company is not None:
            sec_qs = sec_qs.filter(project__company=company)
        if proj_code:
            sec_qs = sec_qs.filter(project__code=proj_code)
        if sec_qs.count() > 1:
            return None, None, f"کد بخش «{sec_code}» بین چند پروژه یا شرکت مشترک است؛ کد پروژه و کد شرکت را درج کنید."
        sec = sec_qs.first()
        if not sec:
            scope_hint = "در شرکت فعال " if company_id is not None else ""
            return None, None, f"بخش با کد «{sec_code}» {scope_hint}یافت نشد."
        return 'section', sec, None

    if proj_code:
        proj_qs = Project.objects.filter(code=proj_code)
        if company is not None:
            proj_qs = proj_qs.filter(company=company)
        if proj_qs.count() > 1:
            return None, None, f"کد پروژه «{proj_code}» در چند شرکت هست؛ کد شرکت را درج کنید."
        proj = proj_qs.first()
        if not proj:
            scope_hint = "در شرکت فعال " if company_id is not None else ""
            return None, None, f"پروژه با کد «{proj_code}» {scope_hint}یافت نشد."
        return 'project', proj, None

    if company is None:
        return None, None, "برای انتساب در سطح شرکت، شرکت فعال را انتخاب کنید یا ستون «کد شرکت» را پر کنید."
    return 'company', company, None


def import_assignments_from_excel(file_obj, dry_run=False, company_id=None, member_check=None, forbidden_user_id=None):
    """
    خواندن و ثبت دسته‌جمعی انتساب کاربران در سطح شرکت، پروژه یا بخش از اکسل با پشتیبانی از dry_run
    (هدف هر ردیف: `_assignment_row_target`).
    company_id داده‌شده: هدف فقط در همان شرکت پیدا می‌شود و کاربر باید از قبل عضو همان شرکت باشد
    (member_check)، چون انتساب خودش دسترسی شرکت می‌دهد.
    forbidden_user_id: کاربری که نباید خودش را منتسب کند (نویسندهٔ «پروژه‌ها و بخش‌ها» بی اجازهٔ انتساب خود).
    """
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        ws = wb.active
    except Exception as e:
        return {
            'success': False,
            'dry_run': dry_run,
            'summary': {'total_rows': 0, 'valid_count': 0, 'error_count': 1, 'created': 0, 'updated': 0, 'skipped': 1},
            'errors': [{'row': 0, 'message': f'فایل اکسل نامعتبر یا آسیب‌دیده است: {str(e)}'}]
        }

    created = 0
    updated = 0
    errors = []

    # نقش‌ها از ماژول‌های نصب‌شده: کد نقش یا عنوان آن پذیرفته می‌شود؛ خانهٔ خالی یعنی نخستین نقش اعلام‌شده
    roles = list(section_role_choices())
    if not roles:
        from .org_serializers import NO_SECTION_ROLES_MESSAGE
        return {
            'success': False,
            'dry_run': dry_run,
            'summary': {'total_rows': 0, 'valid_count': 0, 'error_count': 1, 'created': 0, 'updated': 0, 'skipped': 1},
            'errors': [{'row': 0, 'message': NO_SECTION_ROLES_MESSAGE}]
        }
    role_lookup = {}
    for code, label in roles:
        role_lookup[str(code).strip().lower()] = code
        role_lookup.setdefault(str(label).strip().lower(), code)
    default_role = roles[0][0]

    for row_idx in range(3, ws.max_row + 1):
        raw_username = ws.cell(row=row_idx, column=1).value
        if not raw_username or str(raw_username).strip() == '':
            continue
        username = str(raw_username).strip()
        user = User.objects.filter(username=username).first()
        if not user:
            if company_id is not None:
                # همان پیام غیرعضو، تا وجود نام‌های کاربری بیرون از شرکت لو نرود
                errors.append(f"ردیف {row_idx}: کاربر «{username}» عضو شرکت فعال نیست یا یافت نشد.")
            else:
                errors.append(f"ردیف {row_idx}: کاربر با نام کاربری «{username}» یافت نشد.")
            continue
        if forbidden_user_id is not None and user.id == forbidden_user_id:
            errors.append(f"ردیف {row_idx}: انتساب خودتان را مدیر کل سامانه انجام می‌دهد.")
            continue

        proj_code = normalize_digits(str(ws.cell(row=row_idx, column=3).value or '')).strip().upper()
        sec_code = normalize_digits(str(ws.cell(row=row_idx, column=4).value or '')).strip().upper()
        company_code = normalize_digits(str(ws.cell(row=row_idx, column=9).value or '')).strip().upper()
        field, target, error = _assignment_row_target(proj_code, sec_code, company_code, company_id)
        if error:
            errors.append(f"ردیف {row_idx}: {error}")
            continue

        if company_id is not None and member_check is not None and not member_check(user.id, company_id):
            errors.append(f"ردیف {row_idx}: کاربر «{username}» عضو شرکت فعال نیست یا یافت نشد.")
            continue

        raw_role = str(ws.cell(row=row_idx, column=5).value or '').strip()
        role = role_lookup.get(raw_role.lower()) if raw_role else default_role
        if role is None:
            errors.append(f"ردیف {row_idx}: نقش «{raw_role}» معتبر نیست. نقش‌های مجاز: {', '.join(c for c, _ in roles)}")
            continue

        active_val = str(ws.cell(row=row_idx, column=7).value or 'بله').strip().lower()
        is_active = active_val not in ['خیر', 'false', '0', 'no']

        try:
            exists = SectionAssignment.objects.filter(user=user, role=role, **{field: target}).exists()
            if dry_run:
                if exists:
                    updated += 1
                else:
                    created += 1
            else:
                with transaction.atomic():
                    a, is_new = SectionAssignment.objects.update_or_create(
                        user=user,
                        role=role,
                        **{field: target},
                        defaults={'is_active': is_active}
                    )
                    if is_new:
                        created += 1
                    else:
                        updated += 1
        except Exception as e:
            errors.append(f"ردیف {row_idx} ({username}): {str(e)}")

    return {
        'success': len(errors) == 0,
        'dry_run': dry_run,
        'summary': {
            'total_rows': created + updated + len(errors),
            'valid_count': created + updated,
            'error_count': len(errors),
            'created': created,
            'updated': updated,
            'skipped': len(errors)
        },
        'errors': [{'row': idx, 'message': err} for idx, err in enumerate(errors, 1)]
    }
