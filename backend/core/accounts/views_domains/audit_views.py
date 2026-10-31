# -*- coding: utf-8 -*-
"""
Audit logging, login logging, database storage metrics, and export ViewSets.
"""
import csv
from datetime import datetime
from io import StringIO

from django.db.models import Q, Count
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime, parse_date
from django.utils.timezone import now, timedelta
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import viewsets, filters, permissions, status
from rest_framework.decorators import action
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response

from common.excel_utils import sanitize_excel_row
from platform_core.registry import contributions, entity_view_decision
from ..models import UserLoginLog, AuditLog
from ..serializers import UserLoginLogSerializer, AuditLogSerializer, AuditLogListSerializer
from ..audit_utils import log_login_event
from ..audit_scope import apply_audit_app_scope, restrict_audit_for_user

def _parse_query_date(date_str, is_end_of_day=False):
    if not date_str:
        return None
    try:
        s = str(date_str).strip()
        if len(s) == 10 and '-' in s:
            d = parse_date(s)
            if d:
                time_val = datetime.max.time() if is_end_of_day else datetime.min.time()
                dt = datetime.combine(d, time_val)
                return timezone.make_aware(dt) if timezone.is_naive(dt) else dt
        dt = parse_datetime(s)
        if not dt and 'Z' in s:
            try:
                dt = datetime.fromisoformat(s.replace('Z', '+00:00'))
            except Exception:
                dt = None
        if not dt:
            d = parse_date(s[:10])
            if d:
                time_val = datetime.max.time() if is_end_of_day else datetime.min.time()
                dt = datetime.combine(d, time_val)
        if dt:
            if timezone.is_naive(dt):
                dt = timezone.make_aware(dt)
            return dt
    except Exception:
        pass
    return None

class StandardLogPagination(PageNumberPagination):
    page_size = 20
    page_size_query_param = 'page_size'
    max_page_size = 500

def get_db_total_storage_info():
    """
    استعلام حجم فیزیکی کل دیتابیس فعلی بر روی دیسک با PostgreSQL
    """
    from django.db import connection
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT pg_database_size(current_database())")
            row = cursor.fetchone()
            if row and row[0] is not None:
                return int(row[0])
    except Exception:
        pass
    return None

def get_table_storage_info(table_name):
    """
    استعلام حجم فیزیکی دیسک برای یک جدول با پشتیبانی از PostgreSQL
    """
    from django.db import connection
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT pg_total_relation_size(%s)", [table_name])
            row = cursor.fetchone()
            if row and row[0] is not None:
                return int(row[0])
    except Exception:
        pass
    return None

def format_bytes_to_human(size_bytes):
    """
    تبدیل بایت به فرمت خوانا (B, KB, MB, GB)
    """
    if size_bytes is None or size_bytes < 0:
        return '—'
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    elif size_bytes < 1024 * 1024 * 1024:
        return f"{size_bytes / (1024 * 1024):.2f} MB"
    else:
        return f"{size_bytes / (1024 * 1024 * 1024):.2f} GB"

def build_log_storage_payload(audit_count=0, login_count=0):
    """
    تولید ساختار متادیتای حجم کل پایگاه داده و تفکیک درصدی لاگ‌ها
    """
    db_total_bytes = get_db_total_storage_info()
    audit_table_size = get_table_storage_info('accounts_auditlog')
    login_table_size = get_table_storage_info('accounts_userloginlog')

    # Fallback در صورت دسترسی نداشتن به توابع سیستمی
    if not audit_table_size and audit_count > 0:
        audit_table_size = audit_count * 2560
    elif not audit_table_size:
        audit_table_size = 0

    if not login_table_size and login_count > 0:
        login_table_size = login_count * 1536
    elif not login_table_size:
        login_table_size = 0

    total_logs_size = audit_table_size + login_table_size

    if not db_total_bytes or db_total_bytes <= 0:
        db_total_bytes = max(total_logs_size * 5, 15 * 1024 * 1024)

    audit_percent = round((audit_table_size / db_total_bytes) * 100, 2) if db_total_bytes > 0 else 0.0
    login_percent = round((login_table_size / db_total_bytes) * 100, 2) if db_total_bytes > 0 else 0.0
    total_logs_percent = round((total_logs_size / db_total_bytes) * 100, 2) if db_total_bytes > 0 else 0.0

    return {
        'db_total_bytes': db_total_bytes,
        'db_total_formatted': format_bytes_to_human(db_total_bytes),
        'audit_bytes': audit_table_size,
        'audit_formatted': format_bytes_to_human(audit_table_size),
        'audit_percent': audit_percent,
        'login_bytes': login_table_size,
        'login_formatted': format_bytes_to_human(login_table_size),
        'login_percent': login_percent,
        'total_logs_bytes': total_logs_size,
        'total_logs_formatted': format_bytes_to_human(total_logs_size),
        'total_logs_percent': total_logs_percent,
        'avg_row_size_kb': round((audit_table_size / (audit_count * 1024)), 2) if (audit_table_size and audit_count > 0) else 2.5
    }

def format_shamsi_datetime(dt):
    """
    تبدیل شیء تاریخ/زمان به رشته فرمت‌بندی‌شده شمسی (YYYY/MM/DD HH:MM:SS)
    """
    if not dt:
        return "—"
    try:
        import jdatetime
        from django.utils import timezone
        if timezone.is_aware(dt):
            dt = timezone.localtime(dt)
        jdt = jdatetime.datetime.fromgregorian(datetime=dt)
        return jdt.strftime('%Y/%m/%d %H:%M:%S')
    except Exception:
        return dt.strftime('%Y-%m-%d %H:%M:%S') if hasattr(dt, 'strftime') else str(dt)

class UserLoginLogViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = UserLoginLog.objects.all().select_related('user').prefetch_related('user__groups__customrole').order_by('-created_at')
    serializer_class = UserLoginLogSerializer
    pagination_class = StandardLogPagination

    def get_permissions(self):
        from ..permissions import HasMenuAccess
        # پالس حضور روزانه را هر کاربر واردشده برای خودش می‌فرستد؛ بقیهٔ گزارش ورودها فقط برای دارندگان مجوز گزارش
        if getattr(self, 'action', None) == 'heartbeat':
            return [permissions.IsAuthenticated()]
        return [HasMenuAccess('perm_sys_logs') | HasMenuAccess('perm_sys_audit_export') | HasMenuAccess('perm_sys_purge_logs')]

    def get_queryset(self):
        req = getattr(self, 'request', None)
        qs = super().get_queryset()
        params = getattr(req, 'query_params', getattr(req, 'GET', {})) if req else {}


        username = params.get('username')
        if username:
            qs = qs.filter(username_attempted__icontains=username)

        status = params.get('status')
        if status:
            s_clean = str(status).strip().upper()
            if s_clean == 'FAILED':
                qs = qs.filter(status__startswith='FAILED')
            else:
                qs = qs.filter(status=s_clean)

        ip_address = params.get('ip_address')
        if ip_address:
            qs = qs.filter(ip_address__icontains=ip_address)

        user_id = params.get('user')
        if user_id:
            qs = qs.filter(user_id=user_id)

        search = params.get('search')
        if search:
            from django.db.models import Q
            qs = qs.filter(
                Q(username_attempted__icontains=search) |
                Q(user__first_name__icontains=search) |
                Q(user__last_name__icontains=search) |
                Q(ip_address__icontains=search) |
                Q(failure_reason__icontains=search)
            )

        from_date = _parse_query_date(params.get('from_date'), is_end_of_day=False)
        if from_date:
            qs = qs.filter(created_at__gte=from_date)

        to_date = _parse_query_date(params.get('to_date'), is_end_of_day=True)
        if to_date:
            qs = qs.filter(created_at__lte=to_date)

        return qs

    @action(detail=False, methods=['get'])
    def stats(self, request):
        from django.db.models import Count
        now_time = now()
        today_start = now_time.replace(hour=0, minute=0, second=0, microsecond=0)
        last_24h = now_time - timedelta(hours=24)

        base_qs = self.filter_queryset(self.get_queryset())
        total_logins = base_qs.count()
        logins_24h = base_qs.filter(created_at__gte=last_24h).count()
        failed_24h = base_qs.filter(created_at__gte=last_24h).exclude(status__in=['SUCCESS', 'LOGOUT']).count()
        success_24h = base_qs.filter(created_at__gte=last_24h, status='SUCCESS').count()

        status_breakdown = dict(
            base_qs.filter(created_at__gte=last_24h)
            .values_list('status')
            .annotate(c=Count('id'))
        )

        audit_total = AuditLog.objects.count()
        storage_info = build_log_storage_payload(audit_count=audit_total, login_count=total_logins)

        return Response({
            'total_all_time': total_logins,
            'logins_24h': logins_24h,
            'success_24h': success_24h,
            'failed_24h': failed_24h,
            'status_breakdown': status_breakdown,
            'storage': storage_info
        })

    @action(detail=False, methods=['post'], permission_classes=[permissions.IsAuthenticated])
    def heartbeat(self, request):
        """
        ثبت حضور فعال روزانه کاربر (یک‌بار در هر روز تقویمی)
        """
        user = request.user
        if not user or not user.is_authenticated:
            return Response({'detail': 'احراز هویت نشده'}, status=status.HTTP_401_UNAUTHORIZED)
        
        now_time = now()
        today_start = now_time.replace(hour=0, minute=0, second=0, microsecond=0)
        
        # بررسی اینکه آیا امروز قبلاً رکوردی برای کاربر ثبت شده است یا خیر
        already_logged_today = UserLoginLog.objects.filter(
            user=user,
            created_at__gte=today_start
        ).exists()
        
        if not already_logged_today:
            from ..middleware import get_client_ip
            device_model = request.data.get('device_model') or request.headers.get('Sec-Ch-Ua-Model')
            log_login_event(
                username=user.username,
                status='DAILY_ACTIVE',
                user=user,
                ip_address=get_client_ip(request),
                user_agent=request.META.get('HTTP_USER_AGENT', ''),
                device_model=device_model,
                failure_reason=None
            )
            return Response({'status': 'logged', 'message': 'لاگ حضور روزانه با موفقیت ثبت شد.'})
        
        return Response({'status': 'already_recorded', 'message': 'حضور کاربر در تاریخ امروز قبلاً ثبت شده است.'})

    def _check_export_permission(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_sys_audit_export') or
            user.has_perm('accounts.perm_sys_purge_logs') or
            user.has_perm('accounts.perm_sys_logs')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها کاربران دارای مجوز خروجی ممیزی (perm_sys_audit_export) یا مدیر سیستم مجاز به دریافت گزارش ورودها هستند.")

    @action(detail=False, methods=['get', 'post'])
    def export_excel(self, request):
        self._check_export_permission(request)
        data = request.data if request.method == 'POST' and request.data else request.query_params

        qs = self.get_queryset()
        if request.method == 'POST' and request.data:
            username = data.get('username')
            if username:
                qs = qs.filter(username_attempted__icontains=username)
            status_val = data.get('status')
            if status_val:
                s_clean = str(status_val).strip().upper()
                if s_clean == 'FAILED':
                    qs = qs.filter(status__startswith='FAILED')
                else:
                    qs = qs.filter(status=s_clean)
            ip_address = data.get('ip_address')
            if ip_address:
                qs = qs.filter(ip_address__icontains=ip_address)
            search = data.get('search')
            if search:
                from django.db.models import Q
                qs = qs.filter(
                    Q(username_attempted__icontains=search) |
                    Q(user__first_name__icontains=search) |
                    Q(user__last_name__icontains=search) |
                    Q(ip_address__icontains=search) |
                    Q(failure_reason__icontains=search)
                )
            from_date = _parse_query_date(data.get('from_date'), is_end_of_day=False)
            if from_date:
                qs = qs.filter(created_at__gte=from_date)
            to_date = _parse_query_date(data.get('to_date'), is_end_of_day=True)
            if to_date:
                qs = qs.filter(created_at__lte=to_date)

        export_limit = 10000
        qs = qs[:export_limit]

        all_cols_def = [
            ('index', 'ردیف', 8),
            ('id', 'شناسه', 10),
            ('username_attempted', 'نام کاربری ورودی', 20),
            ('user_name', 'نام و نام خانوادگی', 22),
            ('status', 'وضعیت ورود', 20),
            ('ip_address', 'آدرس آی‌پی', 18),
            ('device_model', 'مدل دستگاه', 20),
            ('user_agent', 'مرورگر / دستگاه', 30),
            ('failure_reason', 'علت شکست', 25),
            ('created_at', 'تاریخ و زمان', 22),
        ]

        req_cols = data.get('columns')
        if req_cols:
            if isinstance(req_cols, str):
                req_cols = [c.strip() for c in req_cols.split(',') if c.strip()]
            selected_cols = [c for c in all_cols_def if c[0] in req_cols]
            if not selected_cols:
                selected_cols = all_cols_def
        else:
            selected_cols = all_cols_def

        format_type = str(data.get('file_format') or data.get('export_format') or data.get('format', 'xlsx')).lower().strip()

        if format_type == 'csv':
            output = StringIO()
            writer = csv.writer(output)
            writer.writerow([c[1] for c in selected_cols])  # ردیف ۱: عنوان فارسی
            writer.writerow([c[0] for c in selected_cols])  # ردیف ۲: نام فیلد دیتابیس
            for idx, item in enumerate(qs, 1):
                user_name = f"{item.user.first_name} {item.user.last_name}".strip() if item.user else "—"
                row_vals = []
                for col_key, _, _ in selected_cols:
                    if col_key == 'index':
                        row_vals.append(idx)
                    elif col_key == 'id':
                        row_vals.append(item.id)
                    elif col_key == 'username_attempted':
                        row_vals.append(item.username_attempted)
                    elif col_key == 'user_name':
                        row_vals.append(user_name)
                    elif col_key == 'status':
                        row_vals.append(item.get_status_display())
                    elif col_key == 'ip_address':
                        row_vals.append(item.ip_address or "—")
                    elif col_key == 'device_model':
                        row_vals.append(item.device_model or "—")
                    elif col_key == 'user_agent':
                        row_vals.append(item.user_agent or "—")
                    elif col_key == 'failure_reason':
                        row_vals.append(item.failure_reason or "—")
                    elif col_key == 'created_at':
                        row_vals.append(format_shamsi_datetime(item.created_at))
                writer.writerow(sanitize_excel_row(row_vals))

            response = HttpResponse(output.getvalue().encode('utf-8-sig'), content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = 'attachment; filename="login_history_report.csv"'
            return response

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "تاریخچه ورود کاربران"
        ws.views.sheetView[0].rightToLeft = True

        header_font = Font(name='Tahoma', size=10, bold=True, color='FFFFFF')
        header_fill = PatternFill('solid', fgColor='4F46E5')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        key_font = Font(name='Consolas', size=9, bold=True, color='475569')
        key_fill = PatternFill('solid', fgColor='F1F5F9')
        key_align = Alignment(horizontal='center', vertical='center')

        data_font = Font(name='Tahoma', size=9)
        data_align_center = Alignment(horizontal='center', vertical='center')
        data_align_right = Alignment(horizontal='right', vertical='center')
        zebra_fill = PatternFill('solid', fgColor='F8FAFC')
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # ردیف ۱: عناوین فارسی
        header_titles = [c[1] for c in selected_cols]
        ws.append(header_titles)
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(selected_cols) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # ردیف ۲: نام‌های دیتابیسی ستون‌ها
        db_keys = [c[0] for c in selected_cols]
        ws.append(db_keys)
        ws.row_dimensions[2].height = 20
        for col_idx in range(1, len(selected_cols) + 1):
            cell = ws.cell(row=2, column=col_idx)
            cell.font = key_font
            cell.fill = key_fill
            cell.alignment = key_align
            cell.border = thin_border

        # سطرهای داده از ردیف ۳
        for idx, item in enumerate(qs, 1):
            user_name = f"{item.user.first_name} {item.user.last_name}".strip() if item.user else "—"
            row_vals = []
            for col_key, _, _ in selected_cols:
                if col_key == 'index':
                    row_vals.append(idx)
                elif col_key == 'id':
                    row_vals.append(item.id)
                elif col_key == 'username_attempted':
                    row_vals.append(item.username_attempted)
                elif col_key == 'user_name':
                    row_vals.append(user_name)
                elif col_key == 'status':
                    row_vals.append(item.get_status_display())
                elif col_key == 'ip_address':
                    row_vals.append(item.ip_address or "—")
                elif col_key == 'device_model':
                    row_vals.append(item.device_model or "—")
                elif col_key == 'user_agent':
                    row_vals.append((item.user_agent or "—")[:120])
                elif col_key == 'failure_reason':
                    row_vals.append(item.failure_reason or "—")
                elif col_key == 'created_at':
                    row_vals.append(format_shamsi_datetime(item.created_at))

            ws.append(sanitize_excel_row(row_vals))
            row_num = idx + 2
            ws.row_dimensions[row_num].height = 22
            is_zebra = (idx % 2 == 0)

            for col_idx, (col_key, _, _) in enumerate(selected_cols, 1):
                c = ws.cell(row=row_num, column=col_idx)
                c.font = data_font
                c.border = thin_border
                if is_zebra:
                    c.fill = zebra_fill
                if col_key in ('user_name', 'user_agent', 'failure_reason'):
                    c.alignment = data_align_right
                else:
                    c.alignment = data_align_center

        for col_idx, (_, _, width_val) in enumerate(selected_cols, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width_val

        ws.freeze_panes = 'A3'

        import io
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="login_history_report.xlsx"'
        return response

    @action(detail=False, methods=['get'])
    def export_csv(self, request):
        return self.export_excel(request)

    @action(detail=False, methods=['post'])
    def purge(self, request):
        """
        پاکسازی تاریخچه ورود کاربران توسط مدیر سیستم بر اساس فیلترهای زمانی و وضعیت
        """
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_sys_purge_logs')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها مدیر ارشد سامانه یا کاربران دارای مجوز اختصاصی پاکسازی لاگ‌ها (perm_sys_purge_logs) مجاز به حذف لاگ‌های ورود هستند.")

        data = request.data or {}
        dry_run = data.get('dry_run') in (True, 'true', '1', 1)
        
        qs = UserLoginLog.objects.all()

        from_date = _parse_query_date(data.get('from_date'), is_end_of_day=False)
        if from_date:
            qs = qs.filter(created_at__gte=from_date)

        to_date = _parse_query_date(data.get('to_date'), is_end_of_day=True)
        if to_date:
            qs = qs.filter(created_at__lte=to_date)

        status_val = data.get('status')
        if status_val:
            qs = qs.filter(status=status_val)

        username = data.get('username')
        if username:
            qs = qs.filter(username_attempted__icontains=username)

        days = data.get('days')
        if days and str(days).isdigit():
            cutoff = now() - timedelta(days=int(days))
            qs = qs.filter(created_at__lt=cutoff)

        count = qs.count()

        if dry_run:
            return Response({
                'success': True,
                'dry_run': True,
                'count': count,
                'message': f"{count} رکورد لاگ ورود کاربران منطبق با شرایط انتخابی برای پاکسازی است."
            })

        confirm_text = str(data.get('confirm_text', '')).strip()
        if confirm_text != 'PURGE_LOGIN_LOGS_CONFIRM':
            return Response({
                'success': False,
                'error': 'جهت پاکسازی قطعی تاریخچه ورود، عبارت تاییدیه امنیتی PURGE_LOGIN_LOGS_CONFIRM باید به صورت دقیق وارد شود.'
            }, status=status.HTTP_400_BAD_REQUEST)

        if count == 0:
            return Response({
                'success': True,
                'purged_count': 0,
                'message': 'هیچ رکوردی منطبق با فیلترهای انتخابی یافت نشد.'
            })

        batch_size = 5000
        while True:
            batch_ids = list(qs.values_list('id', flat=True)[:batch_size])
            if not batch_ids:
                break
            UserLoginLog.objects.filter(id__in=batch_ids).delete()
            if len(batch_ids) < batch_size:
                break

        from ..audit_utils import log_audit_event
        desc_parts = []
        if from_date:
            desc_parts.append(f"از: {from_date.strftime('%Y-%m-%d')}")
        if to_date:
            desc_parts.append(f"تا: {to_date.strftime('%Y-%m-%d')}")
        if status_val:
            desc_parts.append(f"وضعیت: {status_val}")
        if username:
            desc_parts.append(f"کاربر: {username}")
        if days:
            desc_parts.append(f"قدیمی‌تر از {days} روز")
        criteria_str = ' | '.join(desc_parts) if desc_parts else 'کلیه لاگ‌های ورود'

        log_audit_event(
            user=user,
            module='system',
            action='DELETE',
            severity='critical',
            details={
                'description': f"پاکسازی قطعی {count} رکورد از لاگ‌های ورود کاربران ({criteria_str})",
                'purged_count': count,
                'filters': {
                    'from_date': str(from_date) if from_date else None,
                    'to_date': str(to_date) if to_date else None,
                    'status': status_val,
                    'username': username,
                    'days': days
                }
            },
            ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
        )

        return Response({
            'success': True,
            'purged_count': count,
            'message': f"{count} رکورد تاریخچه ورود با موفقیت پاکسازی شد."
        })

    @action(detail=False, methods=['get'])
    def locked_users(self, request):
        return self._get_locked_users_response(request)

    @action(detail=False, methods=['get'], url_path='locked_status')
    def locked_status(self, request):
        return self._get_locked_users_response(request)

    def _get_locked_users_response(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_sys_logs') or
            any(user.has_perm(f'accounts.{p}') for p in contributions('audit_view_perms')) or
            user.has_perm('accounts.perm_usr_edit')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("عدم دسترسی به مشاهده وضعیت کاربران مسدود شده.")

        from axes.models import AccessAttempt
        from django.conf import settings
        failure_limit = getattr(settings, 'AXES_FAILURE_LIMIT', 5)

        attempts = AccessAttempt.objects.all().order_by('-attempt_time')
        # طرح ۲.۹ (S6): غیرمدیر کل فقط قفل کاربران دامنهٔ خودش (هرگز مدیر کل) را می‌بیند، بی IP و مرورگر دیگران
        full_view = user.is_superuser
        if not full_view:
            from ..scope import get_visible_users_for_user
            attempts = attempts.filter(username__in=get_visible_users_for_user(user).values('username'))
        results = []
        for att in attempts:
            show_client = full_view or att.username == user.username
            results.append({
                'username': att.username,
                'ip_address': att.ip_address if show_client else '',
                'failures': att.failures_since_start,
                'is_locked': att.failures_since_start >= failure_limit,
                'attempt_time': att.attempt_time.isoformat() if att.attempt_time else None,
                'user_agent': (att.user_agent or '') if show_client else '',
            })

        return Response({
            'locked_users': results,
            'total_locked': len([r for r in results if r['is_locked']])
        })

    @action(detail=False, methods=['post'])
    def reset_lockout(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_usr_edit') or
            user.has_perm('accounts.perm_sys_logs')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("شما دسترسی لازم برای رفع مسدودیت کاربران را ندارید.")

        data = request.data or {}
        username = str(data.get('username', '')).strip()
        ip_address = str(data.get('ip_address', '')).strip()

        if not username and not ip_address:
            return Response({'error': 'نام کاربری یا آدرس آی‌پی جهت رفع مسدودیت الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        # طرح ۲.۹ (S6): غیرمدیر کل فقط قفل کاربران دامنهٔ خودش را باز می‌کند؛ باز کردن قفل IP فقط کار مدیر کل
        if not user.is_superuser:
            from rest_framework.exceptions import PermissionDenied
            if ip_address:
                raise PermissionDenied("باز کردن قفل آدرس آی‌پی فقط در اختیار مدیر کل سامانه است.")
            from ..scope import get_visible_users_for_user
            if not get_visible_users_for_user(user).filter(username=username).exists():
                raise PermissionDenied("این کاربر در دامنهٔ دسترسی شما نیست.")

        from axes.utils import reset
        from axes.models import AccessAttempt

        if username:
            try:
                reset(username=username)
            except Exception:
                pass
            AccessAttempt.objects.filter(username=username).delete()

        if ip_address:
            try:
                reset(ip=ip_address)
            except Exception:
                pass
            AccessAttempt.objects.filter(ip_address=ip_address).delete()

        # Log security audit event
        from ..audit_utils import log_audit_event
        from ..middleware import get_client_ip
        target_name = f"کاربر «{username}»" if username else f"آی‌پی «{ip_address}»"
        log_audit_event(
            user=user,
            module='users',
            action='UPDATE',
            target_repr=target_name,
            details={'description': f"رفع مسدودیت امنیتی {target_name} و بازنشانی شمارنده تلاش‌های ناموفق سیستم ضد نفوذ (Axes)"},
            severity='warning',
            ip_address=get_client_ip(request)
        )

        return Response({
            'success': True,
            'username': username,
            'ip_address': ip_address,
            'message': f"قفل امنیتی {target_name} با موفقیت بازنشانی شد و کاربر مجدداً مجاز به ورود است."
        })

def audit_entity_kinds() -> dict:
    """نگاشت AuditLog.target_model به نوع سند (ص۱۰ طرح ۲.۶): contribute('audit_entity_kinds', (target_model, kind))."""
    out = {}
    for item in contributions('audit_entity_kinds'):
        try:
            model_name, kind = item
        except (TypeError, ValueError):
            continue
        if model_name and kind:
            out[str(model_name)] = str(kind)
    return out


def denied_audit_ids(user, logs) -> set:
    """
    شناسهٔ ردیف‌هایی از `logs` که قلاب دیدن سندشان برای کاربر منع صریح می‌دهد (entity_view_decision is False).
    برای هر (نوع، شناسه) فقط یک بار تصمیم گرفته می‌شود؛ مدیر کل فعال چیزی را از دست نمی‌دهد. بی‌نظر = نمایش.
    """
    if user is None or (getattr(user, 'is_superuser', False) and getattr(user, 'is_active', False)):
        return set()
    kinds = audit_entity_kinds()
    if not kinds:
        return set()
    decided = {}
    denied = set()
    for log in logs:
        kind = kinds.get(log.target_model or '')
        if not kind or not log.target_object_id:
            continue
        key = (kind, str(log.target_object_id))
        if key not in decided:
            decided[key] = entity_view_decision(user, kind, key[1]) is False
        if decided[key]:
            denied.add(log.pk)
    return denied


class AuditLogViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = AuditLog.objects.all().select_related('user').prefetch_related('user__groups__customrole').order_by('-created_at')
    serializer_class = AuditLogSerializer
    pagination_class = StandardLogPagination

    def list(self, request, *args, **kwargs):
        # ردیف سندی که قلاب دیدن برای کاربر منع می‌دهد پنهان می‌شود (ص۱۰)؛ تصمیم فقط برای ردیف‌های همین صفحه
        queryset = self.filter_queryset(self.get_queryset())
        page = self.paginate_queryset(queryset)
        rows = page if page is not None else list(queryset)
        denied = denied_audit_ids(request.user, rows)
        if denied:
            rows = [log for log in rows if log.pk not in denied]
        serializer = self.get_serializer(rows, many=True)
        if page is not None:
            return self.get_paginated_response(serializer.data)
        return Response(serializer.data)

    def get_object(self):
        # جزئیات، پیش‌نمایش بازگردانی و بازگردانی تکی: سند منع‌شده مثل ناموجود (۴۰۴)
        obj = super().get_object()
        if denied_audit_ids(self.request.user, [obj]):
            from rest_framework.exceptions import NotFound
            raise NotFound()
        return obj

    def get_serializer_class(self):
        if self.action == 'list':
            return AuditLogListSerializer
        return AuditLogSerializer

    def get_permissions(self):
        from ..permissions import any_menu_access
        # مجوز مشاهدهٔ ممیزی ماژول‌ها (contribute('audit_view_perms', ...)) کنار مجوزهای پلتفرم
        return [any_menu_access('perm_sys_logs', *contributions('audit_view_perms'), 'perm_sys_audit_export', 'perm_sys_purge_logs')]

    def get_queryset(self, app_scope=None):
        req = getattr(self, 'request', None)
        user = getattr(req, 'user', None) if req else None
        qs = super().get_queryset()
        params = getattr(req, 'query_params', getattr(req, 'GET', {})) if req else {}
        # فقط export_excel قلمرو بدنهٔ POST را صریح می‌دهد؛ اکشن‌های نوشتنی (revert) همیشه از query string می‌خوانند
        effective_app_scope = app_scope if app_scope is not None else params.get('app_scope')

        # Warehouse restriction: non-superusers without perm_sys_logs only see assigned warehouses
        if user and user.is_authenticated and not (user.is_superuser or user.has_perm('accounts.perm_sys_logs')):
            qs = restrict_audit_for_user(qs, user, effective_app_scope)


        module = params.get('module')
        if module:
            qs = qs.filter(module=module)

        action_param = params.get('action_type') or params.get('action')
        if action_param:
            qs = qs.filter(action=action_param)

        severity = params.get('severity')
        if severity:
            qs = qs.filter(severity=severity)

        warehouse_id = params.get('warehouse')
        if warehouse_id and str(warehouse_id).strip().upper() != 'ALL':
            qs = qs.filter(warehouse_id=warehouse_id)

        user_id = params.get('user')
        if user_id:
            qs = qs.filter(user_id=user_id)

        company_id = params.get('company_id') or params.get('company')
        if not company_id and req and req.META.get('HTTP_X_COMPANY_ID'):
            company_id = req.META.get('HTTP_X_COMPANY_ID')
        if company_id and str(company_id).isdigit():
            c_int = int(company_id)
            if user and user.is_authenticated and not user.is_superuser:
                try:
                    from companies.access import validate_user_company_access
                    validate_user_company_access(user, c_int)
                except Exception:
                    return qs.none()
            from django.db.models import Q
            qs = qs.filter(Q(company_id=c_int) | Q(details__company_id=c_int) | Q(details__company_id=str(c_int)))
        elif user and user.is_authenticated and not user.is_superuser:
            try:
                from companies.access import get_user_allowed_companies
                from django.db.models import Q
                allowed_cids = list(get_user_allowed_companies(user).values_list('id', flat=True))
                qs = qs.filter(Q(company_id__in=allowed_cids) | Q(company_id__isnull=True))
            except Exception:
                pass

        # فیلتر اختصاصی بر اساس قلمرو سامانه (Domain App-Scope Filtering)
        app_scope = params.get('app_scope')
        if app_scope and str(app_scope).lower().strip() != 'all':
            app_scope = str(app_scope).lower().strip()
            if app_scope in ['security', 'incidents']:
                from django.db.models import Q
                qs = qs.filter(
                    Q(details__event__in=['CROSS_APP_DENIED', 'APP_SCOPE_SWITCH', 'ROLE_BREACH_ATTEMPT']) |
                    Q(module__in=['security', 'users']) |
                    Q(action='ROLLBACK') |
                    Q(severity='critical')
                )
            else:
                qs = apply_audit_app_scope(qs, app_scope)

        event_param = params.get('event')
        if event_param:
            qs = qs.filter(details__event=event_param)

        search = params.get('search')
        if search:
            from django.db.models import Q
            qs = qs.filter(
                Q(target_repr__icontains=search) |
                Q(target_object_id__icontains=search) |
                Q(user__username__icontains=search) |
                Q(user__first_name__icontains=search) |
                Q(user__last_name__icontains=search) |
                Q(actor_username__icontains=search) |
                Q(actor_name__icontains=search)
            )

        from_date = _parse_query_date(params.get('from_date'), is_end_of_day=False)
        if from_date:
            qs = qs.filter(created_at__gte=from_date)

        to_date = _parse_query_date(params.get('to_date'), is_end_of_day=True)
        if to_date:
            qs = qs.filter(created_at__lte=to_date)

        return qs

    @action(detail=False, methods=['get'])
    def stats(self, request):
        from django.db.models import Count, Q
        now_time = now()
        today_start = now_time.replace(hour=0, minute=0, second=0, microsecond=0)
        last_24h = now_time - timedelta(hours=24)

        base_qs = self.filter_queryset(self.get_queryset())
        total_logs = base_qs.count()
        logs_24h = base_qs.filter(created_at__gte=last_24h).count()
        critical_24h = base_qs.filter(created_at__gte=last_24h, severity='critical').count()
        rollbacks_24h = base_qs.filter(created_at__gte=last_24h, action='ROLLBACK').count()

        aggregates = base_qs.aggregate(
            total_all_time=Count('id'),
            critical_all_time=Count('id', filter=Q(severity='critical')),
            warning_all_time=Count('id', filter=Q(severity='warning')),
            rollbacks_all_time=Count('id', filter=Q(action='ROLLBACK'))
        )

        warning_24h = base_qs.filter(created_at__gte=last_24h, severity='warning').count()

        # شاخص‌های تفکیک‌شده ممیزی و امنیت
        security_all_time = base_qs.filter(
            Q(details__event__in=['CROSS_APP_DENIED', 'APP_SCOPE_SWITCH', 'ROLE_BREACH_ATTEMPT']) |
            Q(module__in=['security', 'users']) |
            Q(action='ROLLBACK') |
            Q(severity='critical')
        ).count()
        cross_app_denied_count = base_qs.filter(details__event='CROSS_APP_DENIED').count()
        app_switch_count = base_qs.filter(details__event='APP_SCOPE_SWITCH').count()

        warehouse_logs_count = apply_audit_app_scope(base_qs, 'warehouse').count()
        finance_logs_count = apply_audit_app_scope(base_qs, 'finance').count()

        module_breakdown = dict(
            base_qs.filter(created_at__gte=last_24h)
            .values_list('module')
            .annotate(c=Count('id'))
        )

        login_total = UserLoginLog.objects.count()
        audit_total = aggregates['total_all_time'] or 0
        storage_info = build_log_storage_payload(audit_count=audit_total, login_count=login_total)

        return Response({
            'total_all_time': aggregates['total_all_time'] or 0,
            'critical_all_time': aggregates['critical_all_time'] or 0,
            'warning_all_time': aggregates['warning_all_time'] or 0,
            'rollbacks_all_time': aggregates['rollbacks_all_time'] or 0,
            'logs_24h': logs_24h,
            'audits_24h': logs_24h,
            'critical_24h': critical_24h,
            'critical_count': critical_24h,
            'warning_24h': warning_24h,
            'warning_count': warning_24h,
            'rollbacks_24h': rollbacks_24h,
            'security_all_time': security_all_time,
            'cross_app_denied_count': cross_app_denied_count,
            'app_switch_count': app_switch_count,
            'warehouse_logs_count': warehouse_logs_count,
            'finance_logs_count': finance_logs_count,
            'module_breakdown': module_breakdown,
            'storage': storage_info
        })

    def _check_export_permission(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_sys_audit_export') or
            user.has_perm('accounts.perm_sys_purge_logs') or
            user.has_perm('accounts.perm_sys_logs')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها کاربران دارای مجوز خروجی ممیزی (perm_sys_audit_export) یا مدیر سیستم مجاز به دریافت گزارش لاگ‌ها هستند.")

    @action(detail=False, methods=['get', 'post'])
    def export_excel(self, request):
        self._check_export_permission(request)
        data = request.data if request.method == 'POST' and request.data else request.query_params

        effective_app_scope = data.get('app_scope') if request.method == 'POST' and request.data else None
        qs = self.get_queryset(app_scope=effective_app_scope)
        if request.method == 'POST' and request.data:
            module = data.get('module')
            if module:
                qs = qs.filter(module=module)
            action_param = data.get('action_type') or data.get('action')
            if action_param:
                qs = qs.filter(action=action_param)
            severity = data.get('severity')
            if severity:
                qs = qs.filter(severity=severity)
            warehouse_id = data.get('warehouse') or data.get('warehouse_id')
            if warehouse_id and str(warehouse_id) != 'ALL':
                qs = qs.filter(warehouse_id=warehouse_id)
            user_id = data.get('user')
            if user_id:
                qs = qs.filter(user_id=user_id)
            app_scope = data.get('app_scope')
            if app_scope and str(app_scope).lower().strip() != 'all':
                app_scope = str(app_scope).lower().strip()
                if app_scope in ['security', 'incidents']:
                    from django.db.models import Q
                    qs = qs.filter(
                        Q(details__event__in=['CROSS_APP_DENIED', 'APP_SCOPE_SWITCH', 'ROLE_BREACH_ATTEMPT']) |
                        Q(module__in=['security', 'users']) |
                        Q(action='ROLLBACK') |
                        Q(severity='critical')
                    )
                else:
                    qs = apply_audit_app_scope(qs, app_scope)
            event_param = data.get('event')
            if event_param:
                qs = qs.filter(details__event=event_param)
            search = data.get('search')
            if search:
                from django.db.models import Q
                qs = qs.filter(
                    Q(target_repr__icontains=search) |
                    Q(target_object_id__icontains=search) |
                    Q(user__username__icontains=search) |
                    Q(user__first_name__icontains=search) |
                    Q(user__last_name__icontains=search) |
                    Q(actor_username__icontains=search) |
                    Q(actor_name__icontains=search)
                )
            from_date = _parse_query_date(data.get('from_date'), is_end_of_day=False)
            if from_date:
                qs = qs.filter(created_at__gte=from_date)
            to_date = _parse_query_date(data.get('to_date'), is_end_of_day=True)
            if to_date:
                qs = qs.filter(created_at__lte=to_date)

        export_limit = 10000
        qs = list(qs[:export_limit])
        wh_ids = {item.warehouse_id for item in qs if getattr(item, 'warehouse_id', None)}
        wh_map = {}
        if wh_ids:
            from platform_core.registry import resolve_entity_map
            wh_map = resolve_entity_map('warehouse', wh_ids)

        all_cols_def = [
            ('index', 'ردیف', 8),
            ('id', 'شناسه', 10),
            ('user', 'کاربر عامل', 22),
            ('warehouse', 'انبار', 18),
            ('module', 'ماژول', 18),
            ('action', 'نوع عملیات', 16),
            ('severity', 'سطح اهمیت', 14),
            ('target_repr', 'موجودیت هدف', 25),
            ('target_object_id', 'شناسه سند', 14),
            ('changes_summary', 'خلاصه تغییرات', 35),
            ('ip_address', 'آدرس آی‌پی', 18),
            ('created_at', 'تاریخ و زمان', 22),
        ]

        col_alias_map = {
            'user_display': 'user',
            'actor_username': 'user',
            'warehouse_name': 'warehouse',
            'module_display': 'module',
            'action_display': 'action',
            'target': 'target_repr',
            'object_id': 'target_object_id',
            'changes': 'changes_summary',
            'ip': 'ip_address',
            'date': 'created_at',
            'timestamp': 'created_at'
        }

        req_cols = data.get('columns') or request.query_params.get('columns')
        if req_cols:
            if isinstance(req_cols, str):
                req_cols = [c.strip() for c in req_cols.split(',') if c.strip()]
            normalized_cols = set([col_alias_map.get(c, c) for c in req_cols])
            selected_cols = [c for c in all_cols_def if c[0] in normalized_cols or (c[0] == 'index' and ('index' in normalized_cols or 'radif' in normalized_cols))]
            if not selected_cols:
                selected_cols = all_cols_def
        else:
            selected_cols = all_cols_def

        def _get_changes_summary(item):
            if not item.before_state and not item.after_state:
                return "—"
            diffs = []
            before = item.before_state or {}
            after = item.after_state or {}
            all_keys = set(before.keys()) | set(after.keys())
            for k in sorted(all_keys):
                old_v = before.get(k)
                new_v = after.get(k)
                if old_v != new_v:
                    diffs.append(f"{k}: {old_v} ➔ {new_v}")
            return ' | '.join(diffs) if diffs else "بدون تغییر داده‌ای"

        # اگر درخواست CSV باشد
        is_csv = (
            request.query_params.get('format') == 'csv' or
            request.query_params.get('file_format') == 'csv' or
            data.get('file_format') == 'csv' or
            data.get('format') == 'csv' or
            request.path.endswith('export_csv')
        )
        if is_csv:
            import csv
            import io
            output = io.StringIO()
            writer = csv.writer(output)
            # سطر ۱: عناوین فارسی
            writer.writerow([c[1] for c in selected_cols])
            # سطر ۲: نام‌های دیتابیسی
            writer.writerow([c[0] for c in selected_cols])

            for idx, item in enumerate(qs, 1):
                if item.user:
                    u_full = f"{item.user.first_name} {item.user.last_name}".strip()
                    user_display = u_full if u_full else item.user.username
                elif item.actor_name:
                    user_display = f"{item.actor_name} (سابق)"
                elif item.actor_username:
                    user_display = f"{item.actor_username} (سابق)"
                else:
                    user_display = "سیستم"

                wh_name = wh_map.get(item.warehouse_id) or (item.warehouse.name if getattr(item, 'warehouse', None) else "عمومی / سیستم")
                row_vals = []
                for col_key, _, _ in selected_cols:
                    if col_key == 'index':
                        row_vals.append(idx)
                    elif col_key == 'id':
                        row_vals.append(item.id)
                    elif col_key in ('user', 'user_display'):
                        row_vals.append(user_display)
                    elif col_key in ('warehouse', 'warehouse_name'):
                        row_vals.append(wh_name)
                    elif col_key in ('module', 'module_display'):
                        row_vals.append(item.get_module_display() if hasattr(item, 'get_module_display') else item.module)
                    elif col_key in ('action', 'action_display'):
                        row_vals.append(item.get_action_display() if hasattr(item, 'get_action_display') else item.action)
                    elif col_key == 'severity':
                        row_vals.append(item.get_severity_display() if hasattr(item, 'get_severity_display') else item.severity)
                    elif col_key == 'target_repr':
                        row_vals.append(item.target_repr or "—")
                    elif col_key == 'target_object_id':
                        row_vals.append(item.target_object_id or "—")
                    elif col_key == 'changes_summary':
                        row_vals.append(_get_changes_summary(item))
                    elif col_key == 'ip_address':
                        row_vals.append(item.ip_address or "—")
                    elif col_key == 'created_at':
                        row_vals.append(format_shamsi_datetime(item.created_at))
                writer.writerow(row_vals)

            response = HttpResponse(output.getvalue().encode('utf-8-sig'), content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = 'attachment; filename="audit_logs_report.csv"'
            return response

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "گزارش رهگیری تغییرات"
        ws.views.sheetView[0].rightToLeft = True

        header_font = Font(name='Tahoma', size=10, bold=True, color='FFFFFF')
        header_fill = PatternFill('solid', fgColor='4F46E5')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        key_font = Font(name='Consolas', size=9, bold=True, color='475569')
        key_fill = PatternFill('solid', fgColor='F1F5F9')
        key_align = Alignment(horizontal='center', vertical='center')

        data_font = Font(name='Tahoma', size=9)
        data_align_center = Alignment(horizontal='center', vertical='center')
        data_align_right = Alignment(horizontal='right', vertical='center')
        zebra_fill = PatternFill('solid', fgColor='F8FAFC')
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # سطر ۱: عناوین نمایشی فارسی
        ws.row_dimensions[1].height = 28
        for col_idx, (_, persian_title, _) in enumerate(selected_cols, 1):
            cell = ws.cell(row=1, column=col_idx, value=persian_title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # سطر ۲: کلیدهای سیستمی لاتین
        ws.row_dimensions[2].height = 20
        for col_idx, (col_key, _, _) in enumerate(selected_cols, 1):
            cell = ws.cell(row=2, column=col_idx, value=col_key)
            cell.font = key_font
            cell.fill = key_fill
            cell.alignment = key_align
            cell.border = thin_border

        # سطرهای داده از ردیف ۳
        for idx, item in enumerate(qs, 1):
            if item.user:
                u_full = f"{item.user.first_name} {item.user.last_name}".strip()
                user_display = u_full if u_full else item.user.username
            elif item.actor_name:
                user_display = f"{item.actor_name} (سابق)"
            elif item.actor_username:
                user_display = f"{item.actor_username} (سابق)"
            else:
                user_display = "سیستم"

            wh_name = wh_map.get(item.warehouse_id) or (item.warehouse.name if getattr(item, 'warehouse', None) else "عمومی / سیستم")
            row_vals = []
            for col_key, _, _ in selected_cols:
                if col_key == 'index':
                    row_vals.append(idx)
                elif col_key == 'id':
                    row_vals.append(item.id)
                elif col_key in ('user', 'user_display'):
                    row_vals.append(user_display)
                elif col_key in ('warehouse', 'warehouse_name'):
                    row_vals.append(wh_name)
                elif col_key in ('module', 'module_display'):
                    row_vals.append(item.get_module_display() if hasattr(item, 'get_module_display') else item.module)
                elif col_key in ('action', 'action_display'):
                    row_vals.append(item.get_action_display() if hasattr(item, 'get_action_display') else item.action)
                elif col_key == 'severity':
                    row_vals.append(item.get_severity_display() if hasattr(item, 'get_severity_display') else item.severity)
                elif col_key == 'target_repr':
                    row_vals.append(item.target_repr or "—")
                elif col_key == 'target_object_id':
                    row_vals.append(item.target_object_id or "—")
                elif col_key == 'changes_summary':
                    summary_txt = _get_changes_summary(item)
                    row_vals.append(summary_txt[:1000] if len(summary_txt) > 1000 else summary_txt)
                elif col_key == 'ip_address':
                    row_vals.append(item.ip_address or "—")
                elif col_key == 'created_at':
                    row_vals.append(format_shamsi_datetime(item.created_at))

            ws.append(sanitize_excel_row(row_vals))
            row_num = idx + 2
            ws.row_dimensions[row_num].height = 22
            is_zebra = (idx % 2 == 0)

            for col_idx, (col_key, _, _) in enumerate(selected_cols, 1):
                c = ws.cell(row=row_num, column=col_idx)
                c.font = data_font
                c.border = thin_border
                if is_zebra:
                    c.fill = zebra_fill
                if col_key in ('user_display', 'warehouse_name', 'target_repr', 'changes_summary'):
                    c.alignment = data_align_right
                else:
                    c.alignment = data_align_center

        for col_idx, (_, _, width_val) in enumerate(selected_cols, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width_val

        ws.freeze_panes = 'A3'

        import io
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="audit_logs_report.xlsx"'
        return response

    @action(detail=False, methods=['get'])
    def export_csv(self, request):
        return self.export_excel(request)

    @action(detail=False, methods=['post'])
    def purge(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_sys_purge_logs')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها مدیر ارشد سامانه یا کاربران دارای مجوز اختصاصی پاکسازی لاگ‌ها (perm_sys_purge_logs) مجاز به حذف لاگ‌های ممیزی هستند.")

        data = request.data or {}
        dry_run = data.get('dry_run') in (True, 'true', '1', 1)
        
        qs = AuditLog.objects.all()

        from_date = _parse_query_date(data.get('from_date'), is_end_of_day=False)
        if from_date:
            qs = qs.filter(created_at__gte=from_date)

        to_date = _parse_query_date(data.get('to_date'), is_end_of_day=True)
        if to_date:
            qs = qs.filter(created_at__lte=to_date)

        warehouse_id = data.get('warehouse') or data.get('warehouse_id')
        if warehouse_id and str(warehouse_id) != 'ALL':
            qs = qs.filter(warehouse_id=warehouse_id)

        module = data.get('module')
        if module:
            qs = qs.filter(module=module)

        days = data.get('days')
        if days and str(days).isdigit():
            cutoff = now() - timedelta(days=int(days))
            qs = qs.filter(created_at__lt=cutoff)

        count = qs.count()

        if dry_run:
            return Response({
                'success': True,
                'dry_run': True,
                'count': count,
                'message': f"{count} رکورد لاگ ممیزی منطبق با شرایط انتخابی برای پاکسازی است."
            })

        confirm_text = str(data.get('confirm_text', '')).strip()
        if confirm_text != 'PURGE_AUDIT_LOGS_CONFIRM':
            return Response({
                'success': False,
                'error': 'جهت پاکسازی قطعی لاگ‌های ممیزی، عبارت تاییدیه امنیتی PURGE_AUDIT_LOGS_CONFIRM باید به صورت دقیق وارد شود.'
            }, status=status.HTTP_400_BAD_REQUEST)

        if count == 0:
            return Response({
                'success': True,
                'purged_count': 0,
                'message': 'هیچ رکوردی منطبق با فیلترهای انتخابی یافت نشد.'
            })

        batch_size = 5000
        while True:
            batch_ids = list(qs.values_list('id', flat=True)[:batch_size])
            if not batch_ids:
                break
            AuditLog.objects.filter(id__in=batch_ids).delete()
            if len(batch_ids) < batch_size:
                break

        from ..audit_utils import log_audit_event
        desc_parts = []
        if from_date:
            desc_parts.append(f"از: {from_date.strftime('%Y-%m-%d')}")
        if to_date:
            desc_parts.append(f"تا: {to_date.strftime('%Y-%m-%d')}")
        if warehouse_id:
            desc_parts.append(f"انبار: {warehouse_id}")
        if module:
            desc_parts.append(f"ماژول: {module}")
        if days:
            desc_parts.append(f"قدیمی‌تر از {days} روز")
        criteria_str = ' | '.join(desc_parts) if desc_parts else 'کلیه لاگ‌ها'

        log_audit_event(
            user=user,
            module='system',
            action='DELETE',
            severity='critical',
            details={
                'description': f"پاکسازی قطعی {count} رکورد از لاگ‌های ممیزی ({criteria_str})",
                'purged_count': count,
                'filters': {
                    'from_date': str(from_date) if from_date else None,
                    'to_date': str(to_date) if to_date else None,
                    'warehouse': warehouse_id,
                    'module': module,
                    'days': days
                }
            },
            ip_address=getattr(request, 'META', {}).get('REMOTE_ADDR')
        )

        return Response({
            'success': True,
            'purged_count': count,
            'message': f"{count} رکورد لاگ ممیزی با موفقیت پاکسازی شد."
        })

    @action(detail=True, methods=['get'])
    def preview_revert(self, request, pk=None):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_rollback_single') or
            user.has_perm('accounts.perm_rollback_bulk') or
            user.has_perm('accounts.perm_rollback_data') or
            user.has_perm('accounts.perm_restore_deleted')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها کاربران مجاز (دارای مجوز بازگردانی یا مدیر سیستم) مجاز به مشاهده پیش‌نمایش بازگردانی هستند.")
        
        log_entry = self.get_object()
        from ..rollback_service import get_revert_preview
        result = get_revert_preview(log_entry)
        return Response(result)

    @action(detail=True, methods=['post'])
    def revert(self, request, pk=None):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_rollback_single') or
            user.has_perm('accounts.perm_rollback_bulk') or
            user.has_perm('accounts.perm_rollback_data') or
            user.has_perm('accounts.perm_restore_deleted')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها کاربران مجاز (دارای مجوز بازگردانی perm_rollback_single یا مدیر سیستم) مجاز به بازگردانی این رکورد هستند.")
        
        log_entry = self.get_object()
        from ..rollback_service import revert_log_entry
        from ..middleware import get_current_ip
        
        reason = request.data.get('reason')
        ip_address = get_current_ip() or getattr(request, 'META', {}).get('REMOTE_ADDR')
        
        res = revert_log_entry(log_entry, user=request.user, reason=reason, ip_address=ip_address)
        if not res.get('success'):
            return Response(res, status=status.HTTP_400_BAD_REQUEST)
        return Response(res)

    @action(detail=False, methods=['post'])
    def bulk_revert(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_rollback_bulk') or
            user.has_perm('accounts.perm_rollback_data')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها مدیر ارشد سامانه یا کاربران دارای مجوز اختصاصی بازگردانی گروهی (perm_rollback_bulk) مجاز به بازگردانی دسته‌ای رکوردها هستند.")
        
        log_ids = request.data.get('log_ids', [])
        if not log_ids or not isinstance(log_ids, list):
            return Response({'success': False, 'error': 'لیست شناسه‌های لاگ الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            
        reason = request.data.get('reason')
        from ..rollback_service import revert_log_entry
        from ..middleware import get_current_ip
        ip_address = get_current_ip() or getattr(request, 'META', {}).get('REMOTE_ADDR')
        
        success_count = 0
        errors = []
        
        allowed_qs = self.get_queryset()
        for lid in log_ids:
            log_item = allowed_qs.filter(id=lid).first()
            if log_item and denied_audit_ids(request.user, [log_item]):
                log_item = None
            if not log_item:
                errors.append(f"لاگ #{lid} یافت نشد یا در قلمرو انبار مجاز شما قرار ندارد.")
                continue
            res = revert_log_entry(log_item, user=request.user, reason=reason, ip_address=ip_address)
            if res.get('success'):
                success_count += 1
            else:
                errors.append(f"لاگ #{lid}: {res.get('error')}")
                
        return Response({
            'success': success_count > 0,
            'success_count': success_count,
            'errors': errors,
            'message': f"{success_count} مورد با موفقیت بازگردانی شد." + (f" ({len(errors)} خطا)" if errors else "")
        })

    @action(detail=False, methods=['post'])
    def preview_point_in_time_rollback(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_rollback_bulk') or
            user.has_perm('accounts.perm_rollback_data')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها کاربران دارای مجوز بازگردانی گروهی (perm_rollback_bulk) یا مدیر سیستم مجاز به شبیه‌سازی بازگردانی به تاریخ هستند.")

        data = request.data or {}
        target_date_raw = data.get('target_datetime') or data.get('target_date')
        if not target_date_raw:
            return Response({'can_rollback': False, 'message': 'تاریخ و زمان مقصد جهت بازگردانی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        from ..rollback_service import preview_point_in_time_rollback
        target_datetime = _parse_query_date(target_date_raw, is_end_of_day=False)
        if not target_datetime:
            return Response({'can_rollback': False, 'message': 'فرمت تاریخ و زمان نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        warehouse_id = data.get('warehouse') or data.get('warehouse_id')
        module = data.get('module')
        target_model = data.get('target_model')

        res = preview_point_in_time_rollback(
            target_datetime=target_datetime,
            warehouse_id=warehouse_id,
            module=module,
            target_model=target_model
        )
        return Response(res)

    @action(detail=False, methods=['post'])
    def execute_point_in_time_rollback(self, request):
        user = getattr(request, 'user', None)
        if not (user and user.is_authenticated and (
            user.is_superuser or
            user.has_perm('accounts.perm_rollback_bulk') or
            user.has_perm('accounts.perm_rollback_data')
        )):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("تنها مدیر ارشد سامانه یا کاربران دارای مجوز اختصاصی بازگردانی گروهی (perm_rollback_bulk) مجاز به بازگردانی داده‌ها به یک تاریخ مشخص هستند.")

        data = request.data or {}
        target_date_raw = data.get('target_datetime') or data.get('target_date')
        if not target_date_raw:
            return Response({'success': False, 'error': 'تاریخ و زمان مقصد جهت بازگردانی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        from ..rollback_service import execute_point_in_time_rollback
        from ..middleware import get_current_ip
        target_datetime = _parse_query_date(target_date_raw, is_end_of_day=False)
        if not target_datetime:
            return Response({'success': False, 'error': 'فرمت تاریخ و زمان نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        warehouse_id = data.get('warehouse') or data.get('warehouse_id')
        module = data.get('module')
        target_model = data.get('target_model')
        reason = data.get('reason')
        ip_address = get_current_ip() or getattr(request, 'META', {}).get('REMOTE_ADDR')

        res = execute_point_in_time_rollback(
            target_datetime=target_datetime,
            user=user,
            warehouse_id=warehouse_id,
            module=module,
            target_model=target_model,
            reason=reason,
            ip_address=ip_address
        )
        if not res.get('success'):
            return Response(res, status=status.HTTP_400_BAD_REQUEST)
        return Response(res)

__all__ = [
    '_parse_query_date',
    'StandardLogPagination',
    'get_db_total_storage_info',
    'get_table_storage_info',
    'format_bytes_to_human',
    'build_log_storage_payload',
    'format_shamsi_datetime',
    'UserLoginLogViewSet',
    'AuditLogViewSet',
]
