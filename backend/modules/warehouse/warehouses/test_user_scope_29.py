"""
طرح ۲.۹-الف (S3 و S7): انبار شرکت دیگر انتساب نمی‌شود و جزئیات کاربر مشترک فقط انبارهای دامنهٔ مدیر را نشان می‌دهد.
این تست‌ها در ماژول انبار است چون هسته نباید مدل انبار را وارد کند (نگهبان ۲۹).
"""
import io

from openpyxl import load_workbook

from accounts.models import CustomUser
from accounts.test_user_scope_29 import USERS, UserScope29Base
from warehouses.models import Warehouse


class WarehouseScopeTests(UserScope29Base):
    """S3: انبار شرکت دیگر انتساب نمی‌شود؛ انتساب‌های بیرون از دامنه در ویرایش پاک نمی‌شوند."""

    def setUp(self):
        super().setUp()
        self.wh_a = Warehouse.objects.create(name='انبار الف ۲۹', code='S29-WA', company=self.a)
        self.wh_b = Warehouse.objects.create(name='انبار ب ۲۹', code='S29-WB', company=self.b)
        self.wh_legacy = Warehouse.objects.create(name='انبار قدیمی ۲۹', code='S29-WL')

    def test_other_company_warehouse_rejected_even_with_is_superuser_flag(self):
        c = self.as_user(self.manager)
        for body in ({'assigned_warehouses': [self.wh_b.id]},
                     {'assigned_warehouses': [self.wh_b.id], 'is_superuser': True}):
            r = c.patch(f'{USERS}{self.local.id}/', body, format='json')
            self.assertEqual(r.status_code, 400, r.content)
        self.assertFalse(self.local.assigned_warehouses.filter(pk=self.wh_b.pk).exists())

    def test_new_user_without_company_cannot_get_other_company_warehouse(self):
        r = self.as_user(self.manager).post(USERS, {
            'username': 's29_new_wh', 'first_name': 'ن', 'last_name': 'و', 'phone_number': '09171112233',
            'assigned_warehouses': [self.wh_b.id]}, format='json')
        self.assertEqual(r.status_code, 400, r.content)
        self.assertFalse(CustomUser.objects.filter(username='s29_new_wh').exists())

    def test_update_keeps_out_of_scope_assignments(self):
        self.local.assigned_warehouses.set([self.wh_legacy])
        r = self.as_user(self.manager).patch(f'{USERS}{self.local.id}/', {'assigned_warehouses': [self.wh_a.id]}, format='json')
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(set(self.local.assigned_warehouses.values_list('id', flat=True)), {self.wh_a.id, self.wh_legacy.id})
        # نمایش برای مدیر فقط انبار دامنهٔ خودش
        self.assertEqual(r.data['assigned_warehouses'], [self.wh_a.id])


class SharedUserDisplayTests(UserScope29Base):
    """S7: جزئیات و اکسل کاربر مشترک فقط شرکت‌ها و انبارهای خود مدیر را نشان می‌دهد."""

    def setUp(self):
        super().setUp()
        self.wh_a = Warehouse.objects.create(name='انبار الف نمایش', code='S29-DA', company=self.a)
        self.wh_b = Warehouse.objects.create(name='انبار ب نمایش', code='S29-DB', company=self.b)
        self.shared.assigned_warehouses.set([self.wh_a, self.wh_b])
        self.shared.company = self.b.name
        self.shared.save(update_fields=['company'])

    def test_retrieve_hides_other_company_and_warehouse(self):
        r = self.as_user(self.manager).get(f'{USERS}{self.shared.id}/')
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual([a['company_id'] for a in r.data['company_accesses']], [self.a.id])
        self.assertEqual(r.data['company_ids'], [self.a.id])
        self.assertEqual(r.data['assigned_warehouses'], [self.wh_a.id])
        self.assertNotIn(self.b.name, r.content.decode('utf-8'))
        full = self.as_user(self.admin).get(f'{USERS}{self.shared.id}/')
        self.assertEqual(set(full.data['company_ids']), {self.a.id, self.b.id})

    def test_excel_export_hides_other_company_and_warehouse(self):
        r = self.as_user(self.manager).get(f'{USERS}export_excel/', {'ids': str(self.shared.id)})
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        text = ' '.join(str(c) for row in ws.iter_rows(values_only=True) for c in row if c)
        self.assertIn(self.wh_a.name, text)
        self.assertNotIn(self.wh_b.name, text)
        self.assertNotIn(self.b.name, text)
