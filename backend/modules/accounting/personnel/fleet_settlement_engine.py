# -*- coding: utf-8 -*-
"""
Fleet Settlement Engine and Group Bank Payment Diskette Generator.
Calculates monthly gross/net payouts for contracted/rented fleet vehicles,
and exports official Bank Meli group payment Excel files.
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from decimal import Decimal
from django.db.models import Sum, F
from .models import VehicleDriverProfile, VehicleTripLog, MonthlyWorkPeriod


SETTLEMENT_SENSITIVE_KEYS = ('driver_national_code', 'account_owner_national_code', 'account_number', 'sheba_number')


def _settlement_scope(warehouse_id, year_month, period_id, section_id, company_ids):
    """
    د۷ طرح ۴.۶: دامنهٔ تسویه. خروجی (period, year_month, section_id, vehicle_qs, trip_qs).
    period_id یا section_id + year_month (دورهٔ بخش)؛ warehouse_id (قدیمی)؛ بی شناسه: company_ids (None = همه، فقط
    برای نگهبان و فراخوانی درونی؛ نما همیشه شرکت‌های مجاز کاربر را می‌دهد).
    """
    from django.db.models import Q
    from .workflows import fleet_period_trips
    period = None
    vehicle_qs = VehicleDriverProfile.objects.all()
    trip_qs = VehicleTripLog.objects.filter(is_deleted=False)
    if period_id:
        period = MonthlyWorkPeriod.objects.filter(pk=period_id).first()
        if period is None:
            raise ValueError("دورهٔ تسویه یافت نشد.")
        year_month = period.year_month
        section_id = period.section_id
    elif section_id:
        from .attendance_periods import find_period
        period = find_period(None, section_id, year_month)
    if not year_month:
        raise ValueError("پارامتر سال و ماه (year_month) الزامی است.")
    if section_id:
        if period is not None:
            trip_qs = fleet_period_trips(period)
        else:
            trip_qs = trip_qs.filter(section_id=section_id, date_shamsi__startswith=year_month)
        base_vehicle_q = Q(section_id=section_id)
    else:
        trip_qs = trip_qs.filter(date_shamsi__startswith=year_month)
        base_vehicle_q = Q()
        if warehouse_id:
            period = MonthlyWorkPeriod.objects.filter(year_month=year_month, warehouse_id=warehouse_id).first()
            trip_qs = trip_qs.filter(warehouse_id=warehouse_id)
            base_vehicle_q = Q(assigned_warehouse_id=warehouse_id)
        else:
            period = MonthlyWorkPeriod.objects.filter(year_month=year_month).first() if company_ids is None else None
        if company_ids is not None:
            ids = list(company_ids)
            trip_qs = trip_qs.filter(Q(vehicle__company_id__in=ids) | Q(section__project__company_id__in=ids)
                                     | Q(project__company_id__in=ids))
            vehicle_qs = vehicle_qs.filter(Q(company_id__in=ids) | Q(section__project__company_id__in=ids))
    # فقط خودروی مصوب (در همان دامنه) یا خودرویی که در این دوره تردد دارد
    vehicle_qs = vehicle_qs.filter((base_vehicle_q & Q(approval_status='approved'))
                                   | Q(id__in=trip_qs.values('vehicle_id')))
    return period, year_month, section_id, vehicle_qs, trip_qs


def calculate_monthly_fleet_settlement(warehouse_id=None, year_month=None, *, period_id=None, section_id=None,
                                       company_ids=None, viewer=None):
    """
    محاسبه کارکرد و مبالغ تسویه ماهانه ناوگان خودرویی
    - خودروهای استیجاری و ملکی راننده: مشمول محاسبه کرایه و تسویه بانکی
    - خودروهای شرکتی: صرفاً ثبت تردد، بدون مبلغ تسویه (۰ تومان)
    د۷ طرح ۴.۶: دامنه با period_id یا section_id + year_month؛ حساب اصلی با prefetch (بی N+1)؛ viewer داده شود،
    کد ملی و شماره حساب و شبای خودروهایی که viewer دید کامل ندارد پوشیده می‌شود.
    """
    from django.db.models import Prefetch
    from .models import VehicleBankAccount
    period, year_month, section_id, vehicle_qs, trip_qs = _settlement_scope(
        warehouse_id, year_month, period_id, section_id, company_ids)
    is_locked = bool(period and period.status in ('LOCKED', 'FINALIZED'))
    fleet_state = {}
    if period is not None and period.section_id:
        from .workflows import fleet_workflow_state
        fleet_state = fleet_workflow_state(period)
        is_locked = is_locked or fleet_state.get('fleet_locked', False)

    vehicle_qs = vehicle_qs.prefetch_related(Prefetch(
        'bank_accounts', queryset=VehicleBankAccount.objects.filter(is_primary=True, is_active=True).order_by('id'),
        to_attr='primary_accounts'))

    # مپ ترددها بر اساس شناسه خودرو
    trips_by_vehicle = {}
    for t in trip_qs.only('id', 'vehicle_id', 'trip_count', 'unit_rate'):
        trips_by_vehicle.setdefault(t.vehicle_id, []).append(t)

    records = []
    total_trips_sum = 0
    total_payable_sum = Decimal(0)
    contractor_count = 0

    sens_cache = {}
    for v in vehicle_qs:
        v_trips = trips_by_vehicle.get(v.id, [])
        trip_count = sum(t.trip_count for t in v_trips)
        gross_amount = sum((t.trip_count * t.unit_rate) for t in v_trips)
        
        # خودروهای شرکتی بدون مبلغ تسویه هستند
        is_payable = (v.ownership_type in ('contract', 'personal'))
        payable_amount = gross_amount if is_payable else Decimal(0)

        total_trips_sum += trip_count
        if is_payable:
            total_payable_sum += payable_amount
            if trip_count > 0:
                contractor_count += 1

        primary_acc = v.primary_accounts[0] if v.primary_accounts else None
        acc_owner_name = (primary_acc.account_owner_name if primary_acc and primary_acc.account_owner_name else v.driver_name) or ''
        acc_owner_nat_code = (primary_acc.account_owner_national_code if primary_acc and primary_acc.account_owner_national_code else (v.driver_national_code or '')) or ''
        bank_name = (primary_acc.bank_name if primary_acc and primary_acc.bank_name else (v.bank_name or 'بانک ملی')) or ''
        account_number = (primary_acc.account_number if primary_acc and primary_acc.account_number else (v.account_number or '')) or ''
        sheba_number = (primary_acc.sheba_number if primary_acc and primary_acc.sheba_number else (v.sheba_number or '')) or ''

        records.append({
            'vehicle_id': v.id,
            'driver_name': v.driver_name,
            'driver_national_code': v.driver_national_code or '',
            'driver_phone': v.driver_phone or '',
            'account_owner_name': acc_owner_name,
            'account_owner_national_code': acc_owner_nat_code,
            'plate_number': v.plate_number,
            'vehicle_type': v.vehicle_type,
            'vehicle_type_display': v.get_vehicle_type_display(),
            'ownership_type': v.ownership_type,
            'ownership_type_display': v.get_ownership_type_display(),
            'default_service_rate': float(v.default_service_rate or 0),
            'bank_name': bank_name,
            'account_number': account_number,
            'sheba_number': sheba_number,
            'trip_count': trip_count,
            'gross_amount': float(gross_amount),
            'payable_amount': float(payable_amount),
            'is_payable': is_payable,
            'is_active': v.is_active
        })
        if viewer is not None:
            from .fleet_lifecycle_services import can_view_vehicle_sensitive
            from .lifecycle_services import mask_sensitive_value
            if not can_view_vehicle_sensitive(viewer, v, sens_cache):
                for key in SETTLEMENT_SENSITIVE_KEYS:
                    records[-1][key] = mask_sensitive_value(records[-1][key]) if records[-1][key] else ''

    # مرتب‌سازی بر اساس نام راننده
    records.sort(key=lambda r: r['driver_name'])

    summary = {
        'year_month': year_month,
        'warehouse_id': warehouse_id,
        'is_locked': is_locked,
        'period_status': period.status if period else 'OPEN',
        'period_id': period.pk if period else None,
        'section_id': section_id,
        'fleet_workflow_status': fleet_state.get('fleet_workflow_status'),
        'fleet_workflow_step': fleet_state.get('fleet_workflow_step'),
        'fleet_approved': bool(fleet_state.get('fleet_approved', False)),
        'fleet_locked': bool(fleet_state.get('fleet_locked', False)),
        'total_vehicles_count': len(records),
        'active_contractor_count': contractor_count,
        'total_fleet_trips': total_trips_sum,
        'total_payable_settlement': float(total_payable_sum)
    }

    return {
        'summary': summary,
        'records': records
    }


def generate_fleet_bank_meli_excel(settlement_data: dict, year_month_title: str) -> io.BytesIO:
    """
    تولید فایل اکسل پرداخت گروهی بانک ملی برای رانندگان ناوگان
    مطابق ساختار رسمی پورتال بانک ملی (۷ ستونه)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fleet_Bank_Payment"
    ws.views.sheetView[0].rightToLeft = True

    headers = [
        "ردیف",
        "مبلغ (ریال)",
        "شماره حساب ملی / شبا مقصد",
        "نام گیرنده (راننده/پیمانکار)",
        "توضیحات",
        "شناسه واریز (اختیاری)",
        "کد ملی راننده"
    ]

    header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    records = settlement_data.get('records', [])
    row_idx = 1
    for rec in records:
        # فقط رانندگان دارای مبلغ تسویه مثبت
        if not rec.get('is_payable', True):
            continue
        payable_toman = rec.get('payable_amount', 0)
        if payable_toman <= 0:
            continue

        # تبدیل تومان به ریال در خروجی دیسکت بانکی (هر تومان = ۱۰ ریال)
        payable_rial = int(round(payable_toman * 10))

        acc_or_sheba = (rec.get('sheba_number') or rec.get('account_number') or '').strip()
        payee_name = (rec.get('account_owner_name') or rec.get('driver_name') or '').strip()
        desc = f"تسویه کرایه حمل ناوگان {year_month_title} - {rec.get('plate_number', '')}"
        national_code = (rec.get('account_owner_national_code') or rec.get('driver_national_code') or '').strip()

        ws.append([
            row_idx,
            payable_rial,
            acc_or_sheba,
            payee_name,
            desc,
            "",
            national_code
        ])

        current_row = row_idx + 1
        for col_num in range(1, 8):
            c = ws.cell(row=current_row, column=col_num)
            c.border = border_thin
            c.font = Font(name="Tahoma", size=9)
            if col_num in (1, 2, 3, 6, 7):
                c.alignment = align_center
                if col_num == 2:
                    c.number_format = '#,##0'
            else:
                c.alignment = align_right

        row_idx += 1

    # تنظیم عرض ستون‌ها
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 28
    ws.column_dimensions['D'].width = 28
    ws.column_dimensions['E'].width = 38
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 18

    # فریز ردیف اول
    ws.freeze_panes = 'A2'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
