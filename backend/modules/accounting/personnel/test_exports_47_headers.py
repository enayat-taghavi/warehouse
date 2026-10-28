"""
تست‌های هدر دوسربرگه اکسل، فریز پنل و پوشش اطلاعات حساس (۴.۷-ب).
اجرا: set PYTHONUTF8=1&& .\\venv\\Scripts\\python.exe manage.py test personnel.test_exports_47_headers --keepdb
"""
import io
import re
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from rest_framework.test import APIClient

from companies.models import UserCompanyAccess
from personnel.models import (
    CompanyBankAccount, Counterparty, DailyAttendance, ExpenseHead,
    ExpenseInvoice, ExpenseInvoiceLine, InvoicePayment, MonthlyPayrollRecord,
    MonthlyWorkPeriod, PersonnelProfile, PettyCashAccount, PettyCashTransaction,
    VehicleDriverProfile, VehicleTripLog,
)
from personnel.scoped_test_support import grant_section_perms
from personnel.test_fiscal_period_32 import FiscalFixtures

User = get_user_model()
REPORTS = '/api/personnel/reports/'
SHAMSI_DATE_RE = re.compile(r'^1[34]\d\d/\d\d/\d\d')
GREGORIAN_DATE_RE = re.compile(r'^20\d\d')


class Exports47HeaderTests(FiscalFixtures):
    def setUp(self):
        super().setUp()
        self.heads = list(
            ExpenseHead.objects.filter(company=self.company, parent__isnull=False)
            .order_by('code')[:3]
        )

    def make_invoice(self, status='paid', discount=Decimal('100'), vat=Decimal('200'), section=None):
        inv = ExpenseInvoice.objects.create(
            section=section or self.sec,
            counterparty=self.cp,
            invoice_number=f'R47T-{ExpenseInvoice.objects.count()}',
            invoice_date_shamsi='1405/08/01',
            amount=Decimal('0'),
            discount_amount=discount,
            vat_amount=vat,
            description='d',
            created_by=self.admin,
            status=status,
        )
        specs = [(self.heads[0], '2', '500'), (self.heads[1], '1', '400')]
        for pos, (head, qty, price) in enumerate(specs, start=1):
            ExpenseInvoiceLine.objects.create(
                invoice=inv,
                position=pos,
                description='l',
                quantity=Decimal(qty),
                unit_price=Decimal(price),
                expense_head=head,
            )
        subtotal = sum(line.line_amount for line in inv.lines.all())
        inv.subtotal = subtotal
        inv.amount = subtotal - discount + vat
        inv.save()
        return inv

    def member(self, name, codes, section=None, role='employee'):
        u = User.objects.create_user(username=name, password='x')
        UserCompanyAccess.objects.create(user=u, company=self.company, is_default=True)
        grant_section_perms(u, codes, section=section or self.sec, role=role)
        c = APIClient()
        c.force_authenticate(u)
        c.credentials(HTTP_X_COMPANY_ID=str(self.company.id))
        return c

    def assert_sheet_headers_and_dates(self, ws, expected_freeze='A3', check_dates=True):
        """اعتبارسنجی هدر دوسربرگه: سطر ۱ فارسی، سطر ۲ کلیدهای لاتین اسکی، فریز پنل و تاریخ‌های شمسی."""
        self.assertEqual(ws.freeze_panes, expected_freeze)
        max_col = ws.max_column
        self.assertGreaterEqual(max_col, 1)

        row1 = [ws.cell(row=1, column=c).value for c in range(1, max_col + 1)]
        row2 = [ws.cell(row=2, column=c).value for c in range(1, max_col + 1)]

        # سطر ۱ عنوان‌های فارسی دارد (شامل کاراکترهای غیر اسکی)
        self.assertTrue(any(v and not str(v).isascii() for v in row1), f'Row 1 missing non-ascii titles: {row1}')
        # سطر ۲ کلیدهای لاتین سیستمی دارد (اسکی)
        for k in row2:
            if k is not None and str(k).strip():
                self.assertTrue(str(k).isascii(), f'Row 2 key is not ascii: {k}')

        if check_dates and ws.max_row >= 3:
            # شناسایی ستون‌های تاریخ بر اساس کلید لاتین یا عنوان فارسی
            date_col_indices = []
            for col_idx, (t1, t2) in enumerate(zip(row1, row2), start=1):
                key = str(t2 or '').lower()
                title = str(t1 or '')
                if 'date' in key or 'تاریخ' in title or key == 'created_at':
                    date_col_indices.append(col_idx)

            for col_idx in date_col_indices:
                for row_idx in range(3, ws.max_row + 1):
                    val = str(ws.cell(row=row_idx, column=col_idx).value or '').strip()
                    if val and val not in ('-', 'None', ''):
                        self.assertFalse(
                            GREGORIAN_DATE_RE.match(val),
                            f'Gregorian date found in col {col_idx} row {row_idx}: {val}',
                        )
                        self.assertTrue(
                            SHAMSI_DATE_RE.match(val),
                            f'Expected Shamsi date starting with 13xx or 14xx in col {col_idx} row {row_idx}, got: {val}',
                        )

    def test_counterparties_export_headers(self):
        """اکسپورت طرف‌حساب‌ها: وضعیت 200، سطر ۱ فارسی، سطر ۲ لاتین، فریز A3 و تاریخ ثبت شمسی."""
        Counterparty.objects.create(
            company=self.company,
            name='طرف‌حساب تست هدر',
            counterparty_type='contractor',
            national_id='10100000001',
            is_active=True,
        )
        res = self.api.get('/api/personnel/counterparties/export-excel/')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assert_sheet_headers_and_dates(ws, expected_freeze='A3')

    def test_expense_invoices_export_headers(self):
        """اکسپورت فاکتورهای هزینه: وضعیت 200، دو شیت با هدر دو سطری، فریز A3 و تاریخ فاکتور شمسی."""
        self.make_invoice(status='paid')
        res = self.api.get('/api/personnel/expense-invoices/export-excel/')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        for ws in wb.worksheets:
            self.assert_sheet_headers_and_dates(ws, expected_freeze='A3')

    def test_petty_cash_transactions_export_headers(self):
        """اکسپورت گردش تنخواه: وضعیت 200، هدر دوسربرگه، فریز A3 و تاریخ تراکنش شمسی."""
        petty_acc = PettyCashAccount.objects.create(
            company=self.company, custodian=self.admin, ceiling_amount=Decimal('5000000')
        )
        PettyCashTransaction.objects.create(
            account=petty_acc, custodian=self.admin, created_by=self.admin,
            transaction_type='expense', amount=Decimal('15000'),
            transaction_date_shamsi='1405/08/01', title='سند تنخواه آزمایشی',
            status='approved', line_status='approved', section=self.sec,
            expense_head=self.heads[0],
        )
        res = self.api.get('/api/personnel/petty-cash-transactions/export-excel/')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assert_sheet_headers_and_dates(ws, expected_freeze='A3')

    def test_profiles_export_headers(self):
        """اکسپورت پرونده‌های پرسنل: وضعیت 200، شیت Emp_info با فریز A3 و تاریخ‌های تولد و استخدام شمسی."""
        PersonnelProfile.objects.create(
            company=self.company,
            section=self.sec,
            national_code='0010376488',
            first_name='علی',
            last_name='محمدی',
            birth_date='13700101',
            start_date='14050101',
            approval_status='approved',
            is_active=True,
        )
        res = self.api.get('/api/personnel/profiles/export-excel/')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assert_sheet_headers_and_dates(ws, expected_freeze='A3')

    def test_vehicles_export_headers(self):
        """اکسپورت ناوگان و رانندگان: وضعیت 200، هدر دوسربرگه و فریز A3."""
        VehicleDriverProfile.objects.create(
            company=self.company,
            section=self.sec,
            plate_number='12ب345-67',
            vehicle_type='sedan',
            ownership_type='contract',
            driver_name='راننده آزمون',
            driver_national_code='0010376488',
            account_number='1234567890',
            sheba_number='IR120120000000001234567890',
            approval_status='approved',
            is_active=True,
        )
        res = self.api.get('/api/personnel/vehicles/export-excel/')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assert_sheet_headers_and_dates(ws, expected_freeze='A3')

    def test_attendance_export_headers(self):
        """اکسپورت کارکرد ماهانه پرسنل: وضعیت 200، فریز F3 و هدر دوسربرگه."""
        profile = PersonnelProfile.objects.create(
            company=self.company,
            section=self.sec,
            national_code='0010376488',
            first_name='سارا',
            last_name='احمدی',
            approval_status='approved',
            is_active=True,
        )
        DailyAttendance.objects.create(
            personnel=profile,
            section=self.sec,
            project=self.sec.project,
            date_shamsi='1405/08/01',
            status='PRESENT_10H',
            effective_hours=Decimal('10'),
            approval_status='APPROVED',
        )
        res = self.api.get('/api/personnel/attendance/export-monthly-excel/?year_month=1405/08')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assert_sheet_headers_and_dates(ws, expected_freeze='F3')

    def test_trips_export_headers(self):
        """اکسپورت ۳۱ روزه تردد ناوگان: وضعیت 200، هدر دوسربرگه و فریز A3."""
        vehicle = VehicleDriverProfile.objects.create(
            company=self.company,
            section=self.sec,
            plate_number='34ج567-89',
            vehicle_type='pickup',
            ownership_type='contract',
            driver_name='احمد کریمی',
            driver_national_code='0010376488',
            approval_status='approved',
            is_active=True,
        )
        VehicleTripLog.objects.create(
            vehicle=vehicle,
            section=self.sec,
            project=self.sec.project,
            date_shamsi='1405/08/01',
            trip_count=2,
            unit_rate=Decimal('250000'),
            is_deleted=False,
        )
        res = self.api.get('/api/personnel/trips/export-monthly-excel/?year_month=1405/08')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assert_sheet_headers_and_dates(ws, expected_freeze='A3')

    def test_monthly_payroll_export_headers(self):
        """اکسپورت ۵۸ ستونی حقوق ماهانه: وضعیت 200، هدر دوسربرگه و فریز A3."""
        profile = PersonnelProfile.objects.create(
            company=self.company,
            section=self.sec,
            national_code='0010376488',
            first_name='مهران',
            last_name='مدیری',
            approval_status='approved',
            is_active=True,
        )
        period = MonthlyWorkPeriod.objects.create(
            company=self.company,
            section=self.sec,
            project=self.sec.project,
            year_month='1405/08',
            status='READY_TO_PAY',
        )
        MonthlyPayrollRecord.objects.create(
            period=period,
            personnel=profile,
            row_number=1,
            national_code=profile.national_code,
            full_name=profile.full_name,
            base_salary=Decimal('50000000'),
            gross_salary=Decimal('60000000'),
            net_salary=Decimal('55000000'),
        )
        res = self.api.get(f'/api/personnel/monthly-payroll/export-monthly-excel/?period_id={period.id}')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assert_sheet_headers_and_dates(ws, expected_freeze='A3')

    def test_sensitive_columns_masked_without_perm(self):
        """کاربر بدون دسترسی داده‌های حساس، در اکسپورت پرسنل و خودرو کد ملی و شبا را ماسک‌شده می‌بیند."""
        PersonnelProfile.objects.create(
            company=self.company,
            section=self.sec,
            national_code='0010376488',
            first_name='پرسنل',
            last_name='حساس',
            account_number='1234567890',
            approval_status='approved',
            is_active=True,
        )
        VehicleDriverProfile.objects.create(
            company=self.company,
            section=self.sec,
            plate_number='55س111-22',
            vehicle_type='sedan',
            ownership_type='contract',
            driver_name='راننده حساس',
            driver_national_code='0010376488',
            account_number='1234567890',
            sheba_number='IR120120000000001234567890',
            approval_status='approved',
            is_active=True,
        )

        viewer = self.member('r47_exp_viewer', ['view_sys_personnel'])

        # ۱. بررسی اکسپورت پرسنل برای کاربر عادی
        r_p = viewer.get('/api/personnel/profiles/export-excel/')
        self.assertEqual(r_p.status_code, 200)
        wb_p = openpyxl.load_workbook(io.BytesIO(r_p.content))
        ws_p = wb_p.active
        nc_col = next(c for c in range(1, ws_p.max_column + 1) if ws_p.cell(row=2, column=c).value == 'national_code')
        masked_nc = str(ws_p.cell(row=3, column=nc_col).value)
        self.assertTrue(masked_nc.startswith('******'))
        self.assertTrue(masked_nc.endswith('6488'))

        # ۲. بررسی اکسپورت خودرو برای کاربر عادی
        r_v = viewer.get('/api/personnel/vehicles/export-excel/')
        self.assertEqual(r_v.status_code, 200)
        wb_v = openpyxl.load_workbook(io.BytesIO(r_v.content))
        ws_v = wb_v.active
        sh_col = next(c for c in range(1, ws_v.max_column + 1) if ws_v.cell(row=2, column=c).value == 'sheba_number')
        dnc_col = next(c for c in range(1, ws_v.max_column + 1) if ws_v.cell(row=2, column=c).value == 'driver_national_code')
        masked_sheba = str(ws_v.cell(row=3, column=sh_col).value)
        masked_dnc = str(ws_v.cell(row=3, column=dnc_col).value)
        self.assertTrue(masked_sheba.startswith('******'))
        self.assertTrue(masked_sheba.endswith('7890'))
        self.assertTrue(masked_dnc.startswith('******'))
        self.assertTrue(masked_dnc.endswith('6488'))

        # ۳. بررسی عدم ماسک برای مدیر کل / کاربر صاحب دسترسی
        r_p_admin = self.api.get('/api/personnel/profiles/export-excel/')
        ws_p_admin = openpyxl.load_workbook(io.BytesIO(r_p_admin.content)).active
        admin_nc_col = next(c for c in range(1, ws_p_admin.max_column + 1) if ws_p_admin.cell(row=2, column=c).value == 'national_code')
        self.assertEqual(str(ws_p_admin.cell(row=3, column=admin_nc_col).value), '0010376488')

        r_v_admin = self.api.get('/api/personnel/vehicles/export-excel/')
        ws_v_admin = openpyxl.load_workbook(io.BytesIO(r_v_admin.content)).active
        admin_sh_col = next(c for c in range(1, ws_v_admin.max_column + 1) if ws_v_admin.cell(row=2, column=c).value == 'sheba_number')
        admin_dnc_col = next(c for c in range(1, ws_v_admin.max_column + 1) if ws_v_admin.cell(row=2, column=c).value == 'driver_national_code')
        self.assertEqual(str(ws_v_admin.cell(row=3, column=admin_sh_col).value), 'IR120120000000001234567890')
        self.assertEqual(str(ws_v_admin.cell(row=3, column=admin_dnc_col).value), '0010376488')

    def test_report_exports_two_header_rows(self):
        """هر شش گزارش اکسل دارای سربرگ دوسربرگه با فریز A3 هستند و شبا برای حسابدار ماسک می‌شود."""
        bank = CompanyBankAccount.objects.create(
            company=self.company, bank_name='ملت', account_title='حساب شرکت',
            account_number='1234567890', sheba_number='IR120170000000000000000001',
            opening_balance=Decimal('1000000'), opening_balance_date='1405/01/01', is_primary=True,
        )
        inv = self.make_invoice(status='paid')
        InvoicePayment.objects.create(
            invoice=inv, amount=Decimal('2000'), payment_date_shamsi='1405/08/01',
            tracking_code='R47T-BNK1', method='transfer', company_bank_account=bank,
            created_by=self.admin,
        )

        for key in ('expense', 'unpaid-invoices', 'petty-cash', 'bank-payments', 'labor-cost', 'fleet-cost'):
            r = self.api.get(f'{REPORTS}{key}/export-excel/')
            self.assertEqual(r.status_code, 200, (key, r.content[:200]))
            wb = openpyxl.load_workbook(io.BytesIO(r.content))
            for ws in wb.worksheets:
                self.assert_sheet_headers_and_dates(ws, expected_freeze='A3', check_dates=False)

        # بررسی ماسک شبا در خروجی اکسل bank-payments برای حسابدار
        accountant = self.member('r47_acc_exp', ['perm_accounting_reports', 'perm_invoice_record'], section=self.sec)
        r_acc = accountant.get(f'{REPORTS}bank-payments/export-excel/')
        self.assertEqual(r_acc.status_code, 200)
        wb_acc = openpyxl.load_workbook(io.BytesIO(r_acc.content))
        # در شیت خلاصه حساب بانکی
        ws_sum = wb_acc.worksheets[0]
        sheba_col = next(c for c in range(1, ws_sum.max_column + 1) if ws_sum.cell(row=2, column=c).value == 'sheba_number')
        sheba_val = str(ws_sum.cell(row=3, column=sheba_col).value)
        self.assertTrue(sheba_val.startswith('******'))
        self.assertTrue(sheba_val.endswith('0001'))

        # در شیت ریز پرداخت‌ها
        ws_det = wb_acc.worksheets[1]
        det_sheba_col = next(c for c in range(1, ws_det.max_column + 1) if ws_det.cell(row=2, column=c).value == 'sheba_number')
        det_sheba_val = str(ws_det.cell(row=3, column=det_sheba_col).value)
        self.assertTrue(det_sheba_val.startswith('******'))
        self.assertTrue(det_sheba_val.endswith('0001'))

        # برای مدیر کل کامل است
        r_adm = self.api.get(f'{REPORTS}bank-payments/export-excel/')
        wb_adm = openpyxl.load_workbook(io.BytesIO(r_adm.content))
        ws_adm_sum = wb_adm.worksheets[0]
        self.assertEqual(str(ws_adm_sum.cell(row=3, column=sheba_col).value), 'IR120170000000000000000001')
