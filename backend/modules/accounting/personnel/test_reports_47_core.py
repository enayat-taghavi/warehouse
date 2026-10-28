"""
تست‌های سریع هستهٔ گزارش‌های حسابداری (۴.۷-ب): عدد گ۱ برابر سند، مجوز 403، قلمرو 404، ورودی 400، تعداد پرس‌وجوی ثابت
و اکسل دوسربرگه با فریز A3.
اجرا: set PYTHONUTF8=1&& .\\venv\\Scripts\\python.exe manage.py test personnel.test_reports_47_core --keepdb
"""
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.db import connection
from django.test.utils import CaptureQueriesContext
from rest_framework.test import APIClient

from companies.models import UserCompanyAccess
from personnel.accounting_reports import split_discount
from personnel.models import ExpenseHead, ExpenseInvoice, ExpenseInvoiceLine
from personnel.scoped_test_support import grant_section_perms
from personnel.test_fiscal_period_32 import FiscalFixtures

User = get_user_model()
REPORTS = '/api/personnel/reports/'


class Reports47CoreTests(FiscalFixtures):
    def setUp(self):
        super().setUp()
        self.heads = list(ExpenseHead.objects.filter(company=self.company, parent__isnull=False)
                          .order_by('code')[:3])

    def make_invoice(self, status='paid', discount=Decimal('101'), vat=Decimal('290'), section=None):
        inv = ExpenseInvoice.objects.create(
            section=section or self.sec, counterparty=self.cp, invoice_number=f'R47-{ExpenseInvoice.objects.count()}',
            invoice_date_shamsi='1405/08/01', amount=Decimal('0'), discount_amount=discount, vat_amount=vat,
            description='d', created_by=self.admin, status=status)
        specs = [(self.heads[0], '2.5', '1001'), (self.heads[1], '1', '500'), (self.heads[2], '3', '333'),
                 (self.heads[0], '1', '7')]
        for pos, (head, qty, price) in enumerate(specs, start=1):
            ExpenseInvoiceLine.objects.create(invoice=inv, position=pos, description='l', quantity=Decimal(qty),
                                              unit_price=Decimal(price), expense_head=head)
        subtotal = sum(line.line_amount for line in inv.lines.all())
        inv.subtotal = subtotal
        inv.amount = subtotal - discount + vat
        inv.save()
        return inv

    def member(self, name, codes, section=None):
        u = User.objects.create_user(username=name, password='x')
        UserCompanyAccess.objects.create(user=u, company=self.company, is_default=True)
        grant_section_perms(u, codes, section=section or self.sec)
        c = APIClient()
        c.force_authenticate(u)
        return c

    def test_split_discount_matches_document(self):
        inv = self.make_invoice()
        doc = [(head.pk, net) for head, net in inv.head_amounts()]
        self.assertEqual(sum(n for _, n in doc), inv.subtotal - inv.discount_amount)
        self.assertEqual(split_discount([(1, 100), (2, 200)], 0), [(1, Decimal('100')), (2, Decimal('200'))])

    def test_expense_by_head_equals_head_amounts_and_excludes_draft_rejected(self):
        inv = self.make_invoice()
        self.make_invoice(status='draft')
        self.make_invoice(status='rejected')
        r = self.api.get(f'{REPORTS}expense/?group_by=head')
        self.assertEqual(r.status_code, 200, r.content)
        data = r.json()
        got = {row['key']: row['invoice_amount'] for row in data['rows']}
        want = {head.pk: int(net) for head, net in inv.head_amounts()}
        self.assertEqual(got, want)
        self.assertEqual(data['totals']['invoice_amount'], int(inv.amount - inv.vat_amount))
        self.assertEqual(data['totals']['vat_amount'], 290)
        self.assertEqual(data['totals']['invoice_count'], 1)
        for group_by in ('section', 'project', 'counterparty', 'month'):
            r = self.api.get(f'{REPORTS}expense/?group_by={group_by}')
            self.assertEqual(r.status_code, 200, r.content)
            self.assertEqual(r.json()['totals']['total'], int(inv.amount - inv.vat_amount), group_by)

    def test_permission_403_and_scope_404(self):
        no_perm = self.member('r47_noperm', ['perm_invoice_record'])
        self.assertEqual(no_perm.get(f'{REPORTS}expense/').status_code, 403)
        self.assertEqual(no_perm.get(f'{REPORTS}expense/export-excel/').status_code, 403)
        reporter = self.member('r47_rep', ['perm_invoice_record', 'perm_accounting_reports'])
        self.assertEqual(reporter.get(f'{REPORTS}expense/').status_code, 200)
        self.assertEqual(reporter.get(f'{REPORTS}expense/?company_id={self.other.id}').status_code, 404)
        self.assertEqual(reporter.get(f'{REPORTS}expense/?section_id={self.other_sec.id}').status_code, 404)
        self.assertEqual(reporter.get(f'{REPORTS}unpaid-invoices/?project_id={self.other_sec.project_id}')
                         .status_code, 404)
        self.assertEqual(reporter.get(f'{REPORTS}labor-cost/').status_code, 403)
        self.assertEqual(reporter.get(f'{REPORTS}fleet-cost/').status_code, 403)
        self.assertEqual(reporter.get(f'{REPORTS}unknown/').status_code, 404)
        # فاکتور بخش دیگر شرکت دیگر در عدد کاربر نیست
        self.make_invoice(section=self.other_sec)
        mine = self.make_invoice()
        data = reporter.get(f'{REPORTS}expense/?group_by=section').json()
        self.assertEqual([row['key'] for row in data['rows']], [self.sec.id])
        self.assertEqual(data['totals']['invoice_amount'], int(mine.amount - mine.vat_amount))

    def test_bad_input_400(self):
        for qs in ('date_from=1404/01/01&date_to=1405/08/01', 'date_from=1405/13/40', 'company_id=abc',
                   'group_by=xyz', 'date_from=1405/08/05&date_to=1405/08/01'):
            r = self.api.get(f'{REPORTS}expense/?{qs}')
            self.assertEqual(r.status_code, 400, qs)
            self.assertTrue(r.json()['error'])
        self.assertEqual(self.api.get(f'{REPORTS}unpaid-invoices/?page=x').status_code, 400)
        self.assertEqual(self.api.get(f'{REPORTS}bank-payments/?date_from=1400/01/01').status_code, 400)

    def _bulk(self, n):
        invs = ExpenseInvoice.objects.bulk_create([
            ExpenseInvoice(section=self.sec, counterparty=self.cp, invoice_number=f'B{n}-{i}',
                           invoice_date_shamsi='1405/08/01', amount=Decimal('1290'), subtotal=Decimal('1100'),
                           discount_amount=Decimal('100'), vat_amount=Decimal('290'), description='d',
                           status='ready_to_pay', created_by=self.admin) for i in range(n)])
        ExpenseInvoiceLine.objects.bulk_create([
            ExpenseInvoiceLine(invoice=inv, position=pos, description='l', quantity=1, unit_price=amt,
                               line_amount=amt, expense_head=head)
            for inv in invs for pos, (head, amt) in enumerate(((self.heads[0], 700), (self.heads[1], 400)), 1)])

    def _count(self, path):
        with CaptureQueriesContext(connection) as ctx:
            r = self.api.get(path)
        self.assertEqual(r.status_code, 200, r.content)
        return len(ctx.captured_queries), r.json()

    def test_query_count_constant(self):
        self._bulk(30)
        small = {p: self._count(f'{REPORTS}{p}')[0] for p in ('expense/', 'unpaid-invoices/', 'bank-payments/',
                                                             'petty-cash/', 'dashboard-summary/')}
        self._bulk(2970)
        for path, n in small.items():
            big, data = self._count(f'{REPORTS}{path}')
            self.assertEqual(big, n, path)
        _n, data = self._count(f'{REPORTS}expense/')
        self.assertEqual(data['totals']['invoice_count'], 3000)
        self.assertEqual(data['totals']['invoice_amount'], 3000 * 1000)
        _n, unpaid = self._count(f'{REPORTS}unpaid-invoices/?page=2&page_size=100')
        self.assertEqual(unpaid['totals']['remaining'], 3000 * 1290)
        self.assertEqual(unpaid['details']['count'], 3000)
        self.assertEqual(len(unpaid['details']['results']), 100)

    def test_export_excel_two_header_rows_and_a3(self):
        self.make_invoice()
        for key in ('expense', 'unpaid-invoices', 'petty-cash', 'bank-payments', 'labor-cost', 'fleet-cost'):
            r = self.api.get(f'{REPORTS}{key}/export-excel/')
            self.assertEqual(r.status_code, 200, (key, r.content[:200]))
            wb = openpyxl.load_workbook(io.BytesIO(r.content))
            for ws in wb.worksheets:
                self.assertEqual(ws.freeze_panes, 'A3', key)
                self.assertTrue(ws.sheet_view.rightToLeft, key)
                fa = [c.value for c in ws[1]]
                keys = [c.value for c in ws[2]]
                self.assertTrue(all(k and k.isascii() for k in keys), (key, keys))
                self.assertTrue(any(v and not v.isascii() for v in fa), (key, fa))
            self.assertTrue(wb.properties.title)
        self.assertEqual(self.api.get(f'{REPORTS}dashboard-summary/export-excel/').status_code, 404)
