"""
ابزار مشترک اکسل حسابداری (ق۱۰، د۵ طرح ۴.۷): سربرگ دو ردیفه (ردیف ۱ فارسی، ردیف ۲ کلید لاتین)، فریز A3،
راست‌به‌چپ و قالب عددی پول. columns: [(عنوان فارسی، کلید لاتین، نوع)] با نوع 'right' | 'num' | 'money' | هر چیز دیگر.
"""
import io

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_TITLE_FONT = dict(name='B Nazanin', size=11, bold=True, color='FFFFFF')
_KEY_FONT = dict(name='Segoe UI', size=8, italic=True, color='94A3B8')
_TITLE_FILL = '1E293B'
_KEY_FILL = '0F172A'
_BORDER_COLOR = 'E2E8F0'


def _safe(v):
    # متن کاربر که با = + - @ شروع شود فرمول اکسل نشود
    return f"'{v}" if isinstance(v, str) and v[:1] in ('=', '+', '-', '@') else v


def write_standard_sheet(ws, columns, rows):
    """برگهٔ معمولی (نه write_only): سربرگ دو ردیفه، راست‌به‌چپ، ردیف‌های ثابت و قالب عددی پول (ق۱۰)."""
    import openpyxl

    ws.views.sheetView[0].rightToLeft = True
    title_font = Font(**_TITLE_FONT)
    key_font = Font(**_KEY_FONT)
    title_fill = PatternFill(start_color=_TITLE_FILL, end_color=_TITLE_FILL, fill_type='solid')
    key_fill = PatternFill(start_color=_KEY_FILL, end_color=_KEY_FILL, fill_type='solid')
    thin = Side(style='thin', color=_BORDER_COLOR)
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    data_font = Font(name='B Nazanin', size=10)
    data_font_num = Font(name='Segoe UI', size=10)

    ws.append([c[0] for c in columns])
    ws.append([c[1] for c in columns])
    for col_idx in range(1, len(columns) + 1):
        cell_fa = ws.cell(row=1, column=col_idx)
        cell_fa.font = title_font
        cell_fa.fill = title_fill
        cell_fa.alignment = Alignment(horizontal='center', vertical='center')
        cell_en = ws.cell(row=2, column=col_idx)
        cell_en.font = key_font
        cell_en.fill = key_fill
        cell_en.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 16
    ws.freeze_panes = 'A3'

    for r_idx, row in enumerate(rows, start=3):
        ws.append([_safe(v) for v in row])
        ws.row_dimensions[r_idx].height = 20
        for col_idx, (_fa, _key, kind) in enumerate(columns, start=1):
            c = ws.cell(row=r_idx, column=col_idx)
            c.border = thin_border
            c.alignment = Alignment(horizontal='right' if kind == 'right' else 'center', vertical='center')
            c.font = data_font_num if kind in ('num', 'money') else data_font
            if kind == 'money':
                c.number_format = '#,##0'

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)


def write_standard_sheet_stream(ws, columns, rows):
    """
    همان قالب برای برگهٔ write_only (فهرست‌های ریز بزرگ): سبک‌ها روی WriteOnlyCell، فریز A3 و پهنای ستون پیش از
    نخستین ردیف. rows می‌تواند generator باشد.
    """
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.views import Pane

    ws.sheet_view.rightToLeft = True
    ws.sheet_view.pane = Pane(topLeftCell='A3', ySplit=2, state='frozen', activePane='bottomLeft')
    for idx, (fa, key, _kind) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(fa), len(key)) + 8, 60)
    center = Alignment(horizontal='center', vertical='center')
    right = Alignment(horizontal='right', vertical='center')
    thin = Side(style='thin', color=_BORDER_COLOR)
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_font, key_font = Font(**_TITLE_FONT), Font(**_KEY_FONT)
    title_fill = PatternFill(start_color=_TITLE_FILL, end_color=_TITLE_FILL, fill_type='solid')
    key_fill = PatternFill(start_color=_KEY_FILL, end_color=_KEY_FILL, fill_type='solid')
    data_font, data_font_num = Font(name='B Nazanin', size=10), Font(name='Segoe UI', size=10)

    def header(values, font, fill):
        out = []
        for v in values:
            c = WriteOnlyCell(ws, value=v)
            c.font, c.fill, c.alignment = font, fill, center
            out.append(c)
        return out

    ws.append(header([c[0] for c in columns], title_font, title_fill))
    ws.append(header([c[1] for c in columns], key_font, key_fill))
    for row in rows:
        out = []
        for value, (_fa, _key, kind) in zip(row, columns):
            c = WriteOnlyCell(ws, value=_safe(value))
            c.border = thin_border
            c.alignment = right if kind == 'right' else center
            c.font = data_font_num if kind in ('num', 'money') else data_font
            if kind == 'money':
                c.number_format = '#,##0'
            out.append(c)
        ws.append(out)


def build_report_workbook(title, sheets):
    """
    کتاب write_only با یک یا چند برگه: sheets = [(نام برگه، columns، rows)]. عنوان و زمان شمسی در ویژگی‌های فایل
    (title/subject/description)، نه ردیف بالای سربرگ (د۵). خروجی BytesIO.
    """
    import openpyxl
    from common.excel_utils import jalali_now_str

    wb = openpyxl.Workbook(write_only=True)
    stamp = jalali_now_str()
    wb.properties.title = title
    wb.properties.subject = title
    wb.properties.description = f'{title} — تهیه‌شده در {stamp}'
    for name, columns, rows in sheets:
        ws = wb.create_sheet(title=_sheet_name(name))
        write_standard_sheet_stream(ws, columns, rows)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _sheet_name(name):
    for ch in '[]:*?/\\':
        name = name.replace(ch, ' ')
    return name[:31] or 'Sheet'
