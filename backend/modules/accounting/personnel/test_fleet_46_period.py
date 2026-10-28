# -*- coding: utf-8 -*-
"""
Release v1 step 4.6-b2: monthly fleet work on the section period and engine 2.3 (plan 4.6, D5, D6, D7, D10).
Trip -> section period link, full flow with the accountant sheba gate and treasury pay, fleet lock on every trip
write path, return unlocks, no double settlement, personnel attendance lock has no effect, treasury list scope,
disburse_fleet input/scope errors, settlement scope/approval and constant query count.
"""
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from rest_framework.test import APIClient

from companies.models import Company, Project, ProjectSection, UserCompanyAccess
from personnel.attendance_periods import get_or_create_period
from personnel.fleet_settlement_engine import calculate_monthly_fleet_settlement
from personnel.models import VehicleBankAccount, VehicleDriverProfile, VehicleTripLog
from personnel.scoped_test_support import grant_section_perms
from personnel.treasury_engine import TreasuryDisbursementService
from personnel.workflows import FLEET_PERIOD_WORKFLOW, fleet_open_task, is_fleet_locked
from workflow import engine
from workflow.errors import WorkflowError
from workflow.models import WorkflowInstance

User = get_user_model()
YM = '1405/06'
DAY = '1405/06/05'
TRIPS = '/api/personnel/trips/'
TREASURY_API = '/api/personnel/cartable/treasury/'
SETTLE_API = '/api/personnel/fleet-settlement/'
VALID_SHEBA = 'IR820540102680020817909002'


class FleetPeriod46Tests(TestCase):
    def setUp(self):
        self.company, self.project, self.section = self.make_org('F46P')
        self.other_company, self.other_project, self.other_section = self.make_org('F46X')
        self.operator = self.make_user('op', ['view_sys_fleet_attendance'])
        self.supervisor = self.make_user('sup', ['perm_approve_fleet_supervisor', 'view_sys_fleet_attendance'],
                                         role='supervisor')
        self.accountant = self.make_user('acc', ['perm_approve_fleet_finance', 'view_sys_fleet_settlement'],
                                         role='accountant')
        self.manager = self.make_user('mgr', ['perm_approve_fleet_manager'], role='manager')
        self.treasurer = self.make_user('trs', ['perm_treasury_disburse_action'], role='treasury')
        self.rented = self.make_vehicle('11ب111-11', 'contract')
        self.company_car = self.make_vehicle('22ب222-22', 'company')
        self.no_sheba = self.make_vehicle('33ب333-33', 'personal', sheba='')

    # ── helpers ──
    def make_org(self, code):
        company = Company.objects.create(code=f'{code}-CO', name=f'شرکت {code}', is_active=True)
        project = Project.objects.create(code=f'{code}-P', name=f'پروژه {code}', company=company, is_active=True)
        section = ProjectSection.objects.create(project=project, code=f'{code}-S', name=f'بخش {code}', is_active=True)
        return company, project, section

    def make_user(self, name, codes, role='employee', section=None, company=None):
        u = User.objects.create_user(username=f'f46p_{name}', password='x')
        UserCompanyAccess.objects.create(user=u, company=company or self.company, is_default=True)
        grant_section_perms(u, list(codes), section=section or self.section, role=role, exact=True)
        return u

    def make_vehicle(self, plate, ownership, sheba=VALID_SHEBA, section=None, company=None):
        v = VehicleDriverProfile.objects.create(
            plate_number=plate, ownership_type=ownership, driver_name=f'راننده {plate}', driver_phone='09170000000',
            driver_national_code='0012345679', default_service_rate=1000, sheba_number=sheba,
            account_number='0105678901004', bank_name='بانک ملی ایران', section=section or self.section,
            project=(section or self.section).project, company=company or self.company,
            approval_status='approved', is_active=True, valid_until_date='1406/12/29',
            insurance_expiry_date='1406/12/29', inspection_expiry_date='1406/12/29',
            license_expiry_date='1406/12/29',
        )
        if sheba:
            VehicleBankAccount.objects.create(vehicle=v, bank_name='بانک ملی ایران', account_number='0105678901004',
                                              sheba_number=sheba, account_owner_name=v.driver_name,
                                              is_primary=True, is_active=True)
        return v

    def client_for(self, user, company=None):
        c = APIClient()
        c.force_authenticate(user)
        c.credentials(HTTP_X_COMPANY_ID=str((company or self.company).pk))
        return c

    def day_trip(self, vehicle, count=3, user=None, date=DAY):
        return self.client_for(user or self.operator).post(f'{TRIPS}update-day-trip/', {
            'vehicle_id': vehicle.pk, 'date_shamsi': date, 'trip_count': count, 'section_id': self.section.pk,
        }, format='json')

    def record_month(self):
        for v in (self.rented, self.company_car, self.no_sheba):
            r = self.day_trip(v)
            self.assertEqual(r.status_code, 200, r.data)
        return get_or_create_period(None, self.section.pk, YM)[0]

    def submit(self, user=None):
        return self.client_for(user or self.operator).post(
            f'{TRIPS}submit-fleet-period/', {'section_id': self.section.pk, 'year_month': YM}, format='json')

    def to_treasury(self, period):
        engine.act(self.supervisor, fleet_open_task(period.pk))
        self.no_sheba.sheba_number = VALID_SHEBA
        self.no_sheba.save(update_fields=['sheba_number'])
        engine.act(self.accountant, fleet_open_task(period.pk))
        engine.act(self.manager, fleet_open_task(period.pk))
        self.assertEqual(fleet_open_task(period.pk).step_key, 'treasury_pay')

    def pay(self, period, code='FLT-123'):
        return self.client_for(self.treasurer).post(TREASURY_API, {
            'action': 'disburse_fleet', 'period_id': period.pk, 'tracking_code': code}, format='json')

    def assert_locked(self, response):
        self.assertEqual(response.status_code, 400, getattr(response, 'data', None))
        self.assertEqual(response.data.get('code'), 'fleet_locked', response.data)

    # ── D5 ──
    def test_trip_links_to_section_period(self):
        period = self.record_month()
        trip = VehicleTripLog.objects.get(vehicle=self.rented, date_shamsi=DAY, is_deleted=False)
        self.assertEqual(trip.period_id, period.pk)
        self.assertEqual(period.section_id, self.section.pk)

    def test_full_flow_settles_rented_trips(self):
        period = self.record_month()
        r = self.submit()
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['fleet_workflow_step'], 'supervisor_ok')
        self.assertTrue(r.data['fleet_locked'])
        engine.act(self.supervisor, fleet_open_task(period.pk))
        # accountant gate: rented vehicle without a valid sheba blocks with its plate
        with self.assertRaises(WorkflowError) as ctx:
            engine.act(self.accountant, fleet_open_task(period.pk))
        self.assertEqual(ctx.exception.code, 'action_blocked')
        self.assertIn(self.no_sheba.plate_number, ctx.exception.message_fa)
        self.assertNotIn(self.company_car.plate_number, ctx.exception.message_fa)
        # treasury list excludes the period before manager approval
        data = self.client_for(self.treasurer).get(TREASURY_API).data
        self.assertEqual(data['fleet_periods'], [])
        self.assertEqual(data['trips'], [])
        # bank excel before manager approval: 400 fleet_not_approved
        r = self.client_for(self.accountant).get(f'{SETTLE_API}export-bank-excel/', {'period_id': period.pk})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.data['code'], 'fleet_not_approved')
        self.no_sheba.sheba_number = VALID_SHEBA
        self.no_sheba.save(update_fields=['sheba_number'])
        engine.act(self.accountant, fleet_open_task(period.pk))
        engine.act(self.manager, fleet_open_task(period.pk))
        data = self.client_for(self.treasurer).get(TREASURY_API).data
        self.assertEqual([p['id'] for p in data['fleet_periods']], [period.pk])
        self.assertEqual(data['fleet_periods'][0]['total_payable'], 6000)
        self.assertEqual(sorted(t['vehicle_id'] for t in data['trips']), sorted([self.rented.pk, self.no_sheba.pk]))
        r = self.client_for(self.accountant).get(f'{SETTLE_API}export-bank-excel/', {'period_id': period.pk})
        self.assertEqual(r.status_code, 200)
        # pay without tracking code is refused by the generic engine act
        with self.assertRaises(WorkflowError) as ctx:
            engine.act(self.treasurer, fleet_open_task(period.pk))
        self.assertEqual(ctx.exception.code, 'tracking_code_required')
        r = self.pay(period)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['settled_trips_count'], 2)
        settled = VehicleTripLog.objects.filter(period=period, is_settled=True)
        self.assertEqual({t.vehicle_id for t in settled}, {self.rented.pk, self.no_sheba.pk})
        self.assertTrue(all(t.payment_tracking_code == 'FLT-123' for t in settled))
        self.assertFalse(VehicleTripLog.objects.get(vehicle=self.company_car, date_shamsi=DAY).is_settled)
        inst = WorkflowInstance.objects.get(workflow_key=FLEET_PERIOD_WORKFLOW, entity_id=str(period.pk))
        self.assertEqual(inst.status, WorkflowInstance.COMPLETED)
        self.assertTrue(is_fleet_locked(period))
        # second pay: no task, no double settlement
        before = list(settled.values_list('id', 'settled_at'))
        r = self.pay(period, code='FLT-999')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.data['code'], 'not_ready_to_pay')
        self.assertEqual(list(VehicleTripLog.objects.filter(period=period, is_settled=True)
                              .values_list('id', 'settled_at')), before)
        again = TreasuryDisbursementService.settle_fleet_trips([i for i, _ in before], self.treasurer, 'X')
        self.assertEqual(again['settled_trips_count'], 0)

    def test_every_write_path_locked_then_return_unlocks(self):
        period = self.record_month()
        self.assertEqual(self.submit().status_code, 200)
        trip = VehicleTripLog.objects.get(vehicle=self.rented, date_shamsi=DAY, is_deleted=False)
        op = self.client_for(self.operator)
        self.assert_locked(self.day_trip(self.rented, count=5))
        self.assert_locked(op.post(f'{TRIPS}bulk-save/', {
            'date_shamsi': '1405/06/06', 'section_id': self.section.pk,
            'items': [{'vehicle_id': self.rented.pk, 'trip_count': 2}]}, format='json'))
        self.assert_locked(op.post(f'{TRIPS}bulk-save-monthly-grid/', {
            'year_month': YM, 'section_id': self.section.pk,
            'items': [{'vehicle_id': self.rented.pk, 'day': 7, 'trip_count': 2}]}, format='json'))
        self.assert_locked(op.post(TRIPS, {'vehicle': self.rented.pk, 'date_shamsi': '1405/06/08', 'trip_count': 1,
                                           'unit_rate': 1000, 'section': self.section.pk}, format='json'))
        self.assert_locked(op.patch(f'{TRIPS}{trip.pk}/', {'trip_count': 9}, format='json'))
        self.assert_locked(op.delete(f'{TRIPS}{trip.pk}/'))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=2, column=3, value='plate_number')
        ws.cell(row=2, column=7 + 9, value='day_09')
        ws.cell(row=3, column=3, value=self.rented.plate_number)
        ws.cell(row=3, column=7 + 9, value=4)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'fleet.xlsx'
        self.assert_locked(op.post(f'{TRIPS}import-monthly-excel/', {
            'file': buf, 'year_month': YM, 'section_id': self.section.pk}, format='multipart'))
        trip.refresh_from_db()
        self.assertEqual(trip.trip_count, 3)
        self.assertFalse(trip.is_deleted)
        self.assertFalse(VehicleTripLog.objects.filter(date_shamsi__in=['1405/06/06', '1405/06/07', '1405/06/08',
                                                                        '1405/06/09']).exists())
        grid = op.get(f'{TRIPS}monthly-grid/', {'year_month': YM, 'section_id': self.section.pk}).data
        self.assertTrue(grid['is_locked'])
        self.assertTrue(grid['fleet_locked'])
        # resubmit while open: 409
        self.assertEqual(self.submit(self.supervisor).status_code, 409)
        # return for fix unlocks; nothing deleted
        engine.return_for_fix(self.supervisor, fleet_open_task(period.pk), 'اصلاح تعداد')
        self.assertFalse(is_fleet_locked(period))
        r = self.day_trip(self.rented, count=5)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(VehicleTripLog.objects.filter(period=period, is_deleted=False).count(), 3)
        # resubmit after return
        self.assertEqual(self.submit().status_code, 200)

    def test_personnel_attendance_lock_does_not_block_trips(self):
        period = self.record_month()
        period.status = 'SUBMITTED_ACCOUNTANT'
        period.save(update_fields=['status'])
        self.assertFalse(is_fleet_locked(period))
        r = self.day_trip(self.rented, count=7)
        self.assertEqual(r.status_code, 200, r.data)

    def test_submit_rules(self):
        empty = self.client_for(self.operator).post(
            f'{TRIPS}submit-fleet-period/', {'section_id': self.section.pk, 'year_month': '1405/07'}, format='json')
        self.assertEqual(empty.status_code, 400)
        self.assertEqual(empty.data['code'], 'no_fleet_trips')
        self.record_month()
        outsider = self.make_user('out', ['view_sys_fleet_attendance'], section=self.other_section,
                                  company=self.other_company)
        r = self.client_for(outsider, self.other_company).post(
            f'{TRIPS}submit-fleet-period/', {'section_id': self.section.pk, 'year_month': YM}, format='json')
        self.assertEqual(r.status_code, 404)
        bad = self.client_for(self.operator).post(f'{TRIPS}submit-fleet-period/', {'period_id': 'x'}, format='json')
        self.assertEqual(bad.status_code, 400)

    # ── D6 ──
    def test_disburse_fleet_errors(self):
        period = self.record_month()
        trs = self.client_for(self.treasurer)
        trip = VehicleTripLog.objects.filter(period=period).first()
        r = trs.post(TREASURY_API, {'action': 'disburse_fleet', 'trip_ids': [trip.pk], 'tracking_code': 'T'},
                     format='json')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.data['code'], 'fleet_period_required')
        r = trs.post(TREASURY_API, {'action': 'disburse_fleet', 'period_id': 'abc', 'tracking_code': 'T'},
                     format='json')
        self.assertEqual(r.status_code, 400)
        r = trs.post(TREASURY_API, {'action': 'disburse_fleet', 'period_id': 999999, 'tracking_code': 'T'},
                     format='json')
        self.assertEqual(r.status_code, 404)
        other_period = get_or_create_period(None, self.other_section.pk, YM)[0]
        r = trs.post(TREASURY_API, {'action': 'disburse_fleet', 'period_id': other_period.pk, 'tracking_code': 'T'},
                     format='json')
        self.assertEqual(r.status_code, 404)
        # not yet at treasury_pay
        r = self.pay(period)
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.data['code'], 'not_ready_to_pay')
        self.assertFalse(VehicleTripLog.objects.filter(is_settled=True).exists())

    # ── D7 ──
    def test_settlement_scope(self):
        period = self.record_month()
        self.assertEqual(self.client_for(self.operator).get(
            f'{SETTLE_API}calculate/', {'period_id': period.pk}).status_code, 403)
        self.assertEqual(self.client_for(self.operator).get(
            f'{SETTLE_API}calculate/', {'year_month': YM}).status_code, 403)
        other_period = get_or_create_period(None, self.other_section.pk, YM)[0]
        self.assertEqual(self.client_for(self.accountant).get(
            f'{SETTLE_API}calculate/', {'period_id': other_period.pk}).status_code, 404)
        self.assertEqual(self.client_for(self.accountant).get(
            f'{SETTLE_API}calculate/', {'period_id': 'zz'}).status_code, 400)
        r = self.client_for(self.accountant).get(f'{SETTLE_API}calculate/', {'period_id': period.pk})
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['summary']['total_payable_settlement'], 6000.0)
        self.assertFalse(r.data['summary']['fleet_approved'])
        # no ids: only allowed companies
        self.make_vehicle('44ب444-44', 'contract', section=self.other_section, company=self.other_company)
        r = self.client_for(self.accountant).get(f'{SETTLE_API}calculate/', {'year_month': YM})
        self.assertEqual(r.status_code, 200)
        self.assertNotIn('44ب444-44', [rec['plate_number'] for rec in r.data['records']])
        # supervisor-level viewer with the settlement perm but no sensitive view in the vehicle section is masked
        viewer = self.make_user('setl', ['view_sys_fleet_settlement'], section=self.other_section)
        masked = calculate_monthly_fleet_settlement(period_id=period.pk, viewer=viewer)
        self.assertTrue(all(rec['sheba_number'] != VALID_SHEBA for rec in masked['records'] if rec['sheba_number']))
        full = calculate_monthly_fleet_settlement(period_id=period.pk, viewer=self.accountant)
        self.assertIn(VALID_SHEBA, [rec['sheba_number'] for rec in full['records']])

    def test_settlement_constant_queries(self):
        period = get_or_create_period(None, self.section.pk, YM)[0]

        def add(n, start):
            vs = VehicleDriverProfile.objects.bulk_create([VehicleDriverProfile(
                plate_number=f'{start + i}ب', ownership_type='contract', driver_name=f'ر{start + i}',
                default_service_rate=1000, section=self.section, project=self.project, company=self.company,
                approval_status='approved', is_active=True) for i in range(n)])
            VehicleBankAccount.objects.bulk_create([VehicleBankAccount(
                vehicle=v, bank_name='ملی', sheba_number=VALID_SHEBA, is_primary=True, is_active=True) for v in vs])
            VehicleTripLog.objects.bulk_create([VehicleTripLog(
                vehicle=v, section=self.section, period=period, date_shamsi=DAY, trip_count=2, unit_rate=1000,
                total_amount=2000) for v in vs])

        def count():
            with CaptureQueriesContext(connection) as ctx:
                calculate_monthly_fleet_settlement(period_id=period.pk, viewer=self.accountant)
            return len(ctx.captured_queries)

        add(10, 1000)
        count()  # warm per-request caches (content types, permission cache)
        small = count()
        add(290, 2000)
        self.assertEqual(count(), small)
