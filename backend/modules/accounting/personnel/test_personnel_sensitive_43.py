"""
د۶ تا د۸ طرح ۴.۳ (زیرمرحلهٔ ب۲): پوشش دادهٔ حساس پرونده، نوشتن مقدار ستاره‌دار، اکسل، مدرک، صفحه‌بندی و چاپ.
"""
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from rest_framework.test import APIClient

from accounts.models import AuditLog
from companies.models import Company, UserCompanyAccess
from personnel.models import PersonnelBankAccount, PersonnelDocument, PersonnelProfile, Project, ProjectSection
from personnel.scoped_test_support import grant_section_perms

User = get_user_model()
BASE = '/api/personnel/profiles/'
NC = '0010376488'
SHEBA = 'IR820540102680020817909002'


class PersonnelSensitive43Tests(TestCase):
    def setUp(self):
        self.co = Company.objects.create(code='PS43', name='شرکت پرونده', is_active=True)
        self.co2 = Company.objects.create(code='PS43X', name='شرکت دیگر', is_active=True)
        project = Project.objects.create(company=self.co, code='PS43-P', name='پروژه')
        self.s1 = ProjectSection.objects.create(project=project, code='PS43-S1', name='بخش یک')
        self.s2 = ProjectSection.objects.create(project=project, code='PS43-S2', name='بخش دو')
        self.sup = self.user('sup', 'supervisor', self.s1)
        self.acc = self.user('acc', 'accountant', self.s1)
        self.acc2 = self.user('acc2', 'accountant', self.s2)
        self.owner = self.user('owner', 'employee', self.s1)
        project2 = Project.objects.create(company=self.co2, code='PS43-P2', name='پروژه دو')
        self.out = self.user('out', 'accountant', ProjectSection.objects.create(
            project=project2, code='PS43-X1', name='بخش بیرونی'), company=self.co2)
        self.profile = PersonnelProfile.objects.create(
            company=self.co, project=project, section=self.s1, user=self.owner, first_name='علی', last_name='رضایی',
            national_code=NC, id_number='5566', account_number='1234567890', sheba_number=SHEBA,
            card_number='6037991234567890', daily_base_wage=Decimal('5000000'), approval_status='pending_supervisor',
        )
        PersonnelBankAccount.objects.create(personnel=self.profile, bank_name='ملی', account_number='1234567890',
                                            sheba_number=SHEBA, card_number='6037991234567890', is_primary=True)
        PersonnelDocument.objects.create(personnel=self.profile, doc_type='other', uploaded_by=self.acc,
                                         file=SimpleUploadedFile('a.pdf', b'%PDF-1.4'))

    def user(self, name, role, section, company=None):
        u = User.objects.create_user(username=f'ps43_{name}', password='x')
        UserCompanyAccess.objects.create(user=u, company=company or self.co, is_default=True)
        grant_section_perms(u, ['view_sys_personnel'], section=section, role=role, exact=False)
        return u

    def call(self, user, method, url, data=None, company=None, fmt='json'):
        c = APIClient()
        c.force_authenticate(user)
        kwargs = {'HTTP_X_COMPANY_ID': str((company or self.co).pk)}
        if method == 'get':
            return c.get(url, data or {}, **kwargs)
        return getattr(c, method)(url, data or {}, format=fmt, **kwargs)

    def detail(self, user):
        res = self.call(user, 'get', f'{BASE}{self.profile.pk}/')
        self.assertEqual(res.status_code, 200, res.content[:300])
        return res.json()

    def test_supervisor_masked_no_wages(self):
        data = self.detail(self.sup)
        self.assertEqual(data['national_code'], '******6488')
        self.assertTrue(data['sheba_number'].endswith('9002') and '*' in data['sheba_number'])
        self.assertTrue(all('*' in a['sheba_number'] for a in data['bank_accounts']))
        self.assertNotIn('daily_base_wage', data)
        self.assertEqual(data['documents'], [])
        self.assertEqual(data['documents_count'], 1)
        res = self.call(self.sup, 'get', f'{BASE}{self.profile.pk}/bank-accounts/')
        self.assertTrue('*' in res.json()[0]['account_number'])
        docs = self.call(self.sup, 'get', f'{BASE}{self.profile.pk}/documents/')
        self.assertEqual(docs.json(), [])
        self.assertEqual(docs['X-Documents-Count'], '1')

    def test_accountant_same_section_full(self):
        data = self.detail(self.acc)
        self.assertEqual(data['national_code'], NC)
        self.assertEqual(data['bank_accounts'][0]['sheba_number'], SHEBA)
        self.assertIn('daily_base_wage', data)
        self.assertEqual(len(data['documents']), 1)

    def test_accountant_other_section_masked_and_owner_full(self):
        grant_section_perms(self.acc2, ['view_sys_personnel'], section=self.s1, role='employee', exact=False)
        data = self.detail(self.acc2)
        self.assertEqual(data['national_code'], '******6488')
        self.assertEqual(self.detail(self.owner)['national_code'], NC)

    def test_update_with_masked_value_keeps_old(self):
        res = self.call(self.sup, 'patch', f'{BASE}{self.profile.pk}/', {
            'national_code': '******6488', 'sheba_number': '**********************9002', 'father_name': 'حسن',
        })
        self.assertEqual(res.status_code, 200, res.content[:300])
        self.profile.refresh_from_db()
        self.assertEqual(self.profile.national_code, NC)
        self.assertEqual(self.profile.sheba_number, SHEBA)
        self.assertEqual(self.profile.father_name, 'حسن')

    def _excel_row(self, user):
        res = self.call(user, 'get', f'{BASE}export-excel/')
        self.assertEqual(res.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        keys = [c.value for c in ws[2]]
        return dict(zip(keys, [c.value for c in ws[3]]))

    def test_export_excel_masking_and_audit(self):
        row = self._excel_row(self.sup)
        self.assertEqual(row['national_code'], '******6488')
        self.assertIn(row['daily_base_wage'], ('', None))
        self.assertFalse(AuditLog.objects.filter(target_repr='personnel_export_full').exists())
        row = self._excel_row(self.acc)
        self.assertEqual(row['national_code'], NC)
        self.assertEqual(AuditLog.objects.filter(target_repr='personnel_export_full').count(), 1)

    def test_lookup_masked_and_other_company(self):
        res = self.call(self.sup, 'get', f'{BASE}lookup-by-national-code/', {'national_code': NC})
        self.assertEqual(res.json(), {'found': True, 'masked': True,
                                      'personnel': {'first_name': 'علی', 'last_name': 'رضایی'}})
        res = self.call(self.out, 'get', f'{BASE}lookup-by-national-code/', {'national_code': NC}, company=self.co2)
        self.assertEqual(res.json(), {'found': False})
        res = self.call(self.out, 'get', f'{BASE}{self.profile.pk}/documents/', company=self.co2)
        self.assertEqual(res.status_code, 404)

    def test_pagination_and_constant_queries(self):
        for i in range(3):
            PersonnelProfile.objects.create(company=self.co, section=self.s1, first_name='ن', last_name=f'{i}',
                                            national_code=f'00000000{i}{i}', approval_status='draft')
        res = self.call(self.sup, 'get', BASE)
        self.assertIsInstance(res.json(), list)
        self.assertEqual(len(res.json()), 4)
        res = self.call(self.sup, 'get', BASE, {'page': 1, 'page_size': 2})
        self.assertEqual(res.json()['count'], 4)
        self.assertEqual(len(res.json()['results']), 2)
        res = self.call(self.sup, 'get', '/api/personnel/personnel-change-requests/', {'page': 1})
        self.assertIn('results', res.json())
        with CaptureQueriesContext(connection) as small:
            self.call(self.sup, 'get', BASE, {'section_id': self.s1.pk})
        for i in range(4, 9):
            PersonnelProfile.objects.create(company=self.co, section=self.s1, first_name='ن', last_name=f'{i}',
                                            national_code=f'00000000{i}{i}', approval_status='draft')
        with CaptureQueriesContext(connection) as big:
            self.call(self.sup, 'get', BASE, {'section_id': self.s1.pk})
        self.assertEqual(len(small.captured_queries), len(big.captured_queries))

    def test_timeline_metadata_masked(self):
        from personnel.models import WorkflowAuditLog
        WorkflowAuditLog.objects.create(
            content_type='personnel', object_id=self.profile.pk, actor=self.acc, from_status='draft',
            to_status='draft', action='update', metadata={'company_id': self.co.pk,
                                                          'diff': {'national_code': NC, 'father_name': 'حسن'},
                                                          'old_diff': {'sheba_number': {'old': SHEBA}}})
        url = '/api/personnel/workflow-audit-logs/'
        params = {'content_type': 'personnel', 'object_id': self.profile.pk}
        meta = self.call(self.sup, 'get', url, params).json()[0]['metadata']
        self.assertEqual(meta['diff']['national_code'], '******6488')
        self.assertEqual(meta['diff']['father_name'], 'حسن')
        self.assertTrue(meta['old_diff']['sheba_number']['old'].endswith('9002'))
        self.assertNotIn(SHEBA, str(meta))
        meta = self.call(self.acc, 'get', url, params).json()[0]['metadata']
        self.assertEqual(meta['diff']['national_code'], NC)

    def test_print_masking_and_scope(self):
        from personnel.personnel_print import printable_personnel_profile
        doc = printable_personnel_profile(self.sup, self.profile.pk)
        text = ' '.join(doc.paragraphs) + ' '.join(' '.join(r) for r in doc.table_rows)
        from reports.document_pdf import fa_digits
        self.assertNotIn(fa_digits(SHEBA), text)
        self.assertNotIn(fa_digits(NC), text)
        self.assertNotIn('مزد روزانه', text)
        full = printable_personnel_profile(self.acc, self.profile.pk)
        full_text = ' '.join(full.paragraphs) + ' '.join(' '.join(r) for r in full.table_rows)
        self.assertIn(fa_digits(SHEBA), full_text)
        self.assertIn('مزد روزانه', full_text)
        self.assertIsNone(printable_personnel_profile(self.out, self.profile.pk))
        # مدرک‌ها: بی‌دید فقط تعداد
        self.assertFalse(any(p.startswith('مدرک:') for p in doc.paragraphs))
        self.assertEqual(doc.attachments_count, 1)
        self.assertTrue(any(p.startswith('مدرک:') for p in full.paragraphs))

    def test_partial_national_code_search_only_exact_for_plain_users(self):
        res = self.call(self.sup, 'get', BASE, {'search': NC[:7]})
        self.assertEqual(res.json(), [])
        res = self.call(self.sup, 'get', BASE, {'search': NC})
        self.assertEqual(len(res.json()), 1)
