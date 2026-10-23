import re
from platform_core.registry import permission_alias, expand_permission_aliases
import jdatetime
from django.db.models import Q
from common.date_utils import normalize_digits
from common.validators import (
    validate_iranian_national_code,
    validate_iranian_mobile,
    validate_iranian_postal_code,
)


def normalize_job_title(text) -> str:
    """
    DEF-22: یکدست‌سازی عنوان شغل؛ «ي» و «ك» عربی به فارسی، حذف کشیدگی و نویسه‌های نامرئی و فاصله‌های تکراری.
    فایل مرجع عنوان‌ها را با حروف عربی و گاه بی‌فاصله دارد (مثل «مديرپروژه»).
    """
    if text is None:
        return ''
    s = str(text).replace('ي', 'ی').replace('ك', 'ک').replace('ـ', '')
    s = re.sub(r'[​‍‎‏﻿]', '', s)
    s = re.sub(r'\s+', ' ', s)
    return s.strip()


def job_title_match_key(text) -> str:
    """کلید تطبیق عنوان شغل: پس از یکدست‌سازی، فاصله و نیم‌فاصله هم نادیده گرفته می‌شود («مدیرپروژه» = «مدیر پروژه»)."""
    return re.sub(r'[\s‌]+', '', normalize_job_title(text))


def get_month_last_day(year: int, month: int) -> int:
    """محاسبه آخرین روز ماه در گاه‌شماری هجری خورشیدی، با احتساب اسفند کبیسه"""
    if 1 <= month <= 6:
        return 31
    elif 7 <= month <= 11:
        return 30
    elif month == 12:
        try:
            return 30 if jdatetime.date(year, 1, 1).isleap() else 29
        except Exception:
            return 29
    return 30


def parse_shamsi_date(value):
    """
    تبدیل رشتهٔ تاریخ شمسی (YYYY/MM/DD یا YYYY-MM-DD، با ارقام فارسی یا لاتین، با یا بدون صفر پیش‌رو)
    به jdatetime.date. برای ورودی خالی یا نامعتبر None برمی‌گرداند.
    """
    if value is None:
        return None
    raw = normalize_digits(str(value)).strip().replace('-', '/')
    if not raw:
        return None
    if len(raw) == 8 and raw.isdigit() and '/' not in raw:
        raw = f"{raw[:4]}/{raw[4:6]}/{raw[6:8]}"
    parts = raw.split('/')
    if len(parts) != 3 or not all(p.isdigit() for p in parts) or len(parts[0]) != 4:
        return None
    try:
        return jdatetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
    except (ValueError, TypeError):
        return None


def normalize_shamsi_date(value) -> str:
    """
    DEF-20: قالب استاندارد YYYY/MM/DD با صفر پیش‌رو. ورودی خالی رشتهٔ خالی برمی‌گرداند؛
    ورودی نامعتبر ValueError می‌دهد.
    """
    if value is None or not str(value).strip():
        return ''
    parsed = parse_shamsi_date(value)
    if parsed is None:
        raise ValueError('تاریخ باید شمسی و معتبر در قالب YYYY/MM/DD باشد (مثلاً 1405/04/05).')
    return f"{parsed.year:04d}/{parsed.month:02d}/{parsed.day:02d}"


def validate_personnel_validity(personnel, year_month_or_date: str) -> tuple[bool, str]:
    """
    تسک ۲.۱: سرویس اعتبارسنجی تاریخ انقضای شمسی بر مبنای آخرین روز ماه محاسباتی (با اسفند کبیسه).
    اگر valid_until_date خالی باشد، پرونده بدون انقضا (نامحدود) تلقی می‌شود.
    مقایسه با شیء تاریخ انجام می‌شود، نه رشته (DEF-20)؛ چون «1405/4/5» در مقایسهٔ رشته‌ای از «1405/04/15» بزرگ‌تر است.
    در صورت انقضا، (False, پیام خطا) برمی‌گرداند.
    """
    if not personnel:
        return True, ""

    raw_valid_until = getattr(personnel, 'valid_until_date', None)
    if not raw_valid_until or not str(raw_valid_until).strip():
        return True, ""

    valid_until = parse_shamsi_date(raw_valid_until)
    if valid_until is None:
        import logging
        logging.getLogger(__name__).warning(
            "[Validity] valid_until_date نامعتبر برای پرسنل %s: %r — بدون انقضا تلقی شد.",
            getattr(personnel, 'pk', None), raw_valid_until
        )
        return True, ""
    val_label = f"{valid_until.year:04d}/{valid_until.month:02d}/{valid_until.day:02d}"

    target_str = normalize_digits(str(year_month_or_date or '')).strip().replace('-', '/')
    parts = target_str.split('/')
    try:
        if len(parts) == 2:
            # بررسی دوره محاسباتی ماهانه (مانند 1405/04): معیار، آخرین روز ماه است
            y, m = int(parts[0]), int(parts[1])
            effective_check_date = jdatetime.date(y, m, get_month_last_day(y, m))
            if valid_until < effective_check_date:
                msg = (
                    f"اعتبار پرونده پرسنل «{personnel.full_name}» در تاریخ {val_label} به پایان رسیده "
                    f"و برای دوره محاسباتی {target_str} منقضی شده است."
                )
                return False, msg
        elif len(parts) == 3:
            # بررسی تاریخ روزانه (مانند 1405/04/15)
            target_date = parse_shamsi_date(target_str)
            if target_date is not None and valid_until < target_date:
                msg = (
                    f"اعتبار پرونده پرسنل «{personnel.full_name}» در تاریخ {val_label} منقضی شده "
                    f"و ثبت کارکرد برای تاریخ {target_str} مجاز نمی‌باشد."
                )
                return False, msg
    except (ValueError, TypeError):
        pass

    return True, ""


def split_expired_payroll_records(records, year_month: str):
    """
    DEF-20: جداسازی فیش‌های پرسنل منقضی از خروجی دیسکت (تصمیم مصوب ۶.۶).
    خروجی: (فیش‌های معتبر، فهرست پرسنل کنارگذاشته با نام و کد ملی و تاریخ اعتبار)
    """
    valid_records = []
    excluded = []
    for r in records:
        is_valid, _ = validate_personnel_validity(r.personnel, year_month)
        if is_valid:
            valid_records.append(r)
        else:
            excluded.append({
                'personnel_id': r.personnel_id,
                'full_name': r.personnel.full_name,
                'national_code': r.personnel.national_code,
                'valid_until_date': r.personnel.valid_until_date,
            })
    return valid_records, excluded


def check_personnel_history(personnel) -> dict:
    """
    تسک ۲.۳: سرویس استعلام جامع سابقه پرسنل پیش از حذف فیزیکی یا رد قطعی.
    تمام جداول و سوابق را بررسی کرده و جزییات را برمی‌گرداند.
    """
    from personnel.models import (
        DailyAttendance,
        MonthlyPayrollRecord,
        AccountingEvent,
        PersonnelChangeRequest,
        PersonnelSectionAssignment
    )

    pers_id = personnel.id
    
    # ۱. کارکرد روزانه (شامل رکوردهای فعال و حذف‌نرم‌شده)
    attendance_count = DailyAttendance.objects.filter(personnel_id=pers_id).count()

    # ۲. فیش‌های حقوق و دستمزد ماهانه (دیسکت‌های بیمه، مالیات و بانک از همین جدول ساخته می‌شوند)
    payroll_count = MonthlyPayrollRecord.objects.filter(personnel_id=pers_id).count()
    payroll_ids = list(MonthlyPayrollRecord.objects.filter(personnel_id=pers_id).values_list('id', flat=True))

    # ۳. رویدادهای مالی دفتر کل (AccountingEvent با source_model پرسنل یا فیش‌های حقوق)
    acc_events_q = Q(source_model='PersonnelProfile', source_id=str(pers_id))
    if payroll_ids:
        acc_events_q |= Q(source_model='MonthlyPayrollRecord', source_id__in=[str(pid) for pid in payroll_ids])
    accounting_events_count = AccountingEvent.objects.filter(acc_events_q).count()

    # ۴. درخواست‌های تغییرات پرسنل
    change_requests_count = PersonnelChangeRequest.objects.filter(personnel_id=pers_id).count()

    # ۵. انتسابات سازمانی به بخش‌ها
    section_assignments_count = PersonnelSectionAssignment.objects.filter(personnel_id=pers_id).count()

    # ۶. حساب کاربری متصل در سامانه
    has_user_account = bool(personnel.user_id)

    # تجمیع و ارزیابی
    reasons = []
    if attendance_count > 0:
        reasons.append(f"{attendance_count} رکورد کارکرد روزانه (شامل رکوردهای آرشیوی)")
    if payroll_count > 0:
        reasons.append(f"{payroll_count} رکورد فیش حقوق و محاسبات ماهانه")
    if accounting_events_count > 0:
        reasons.append(f"{accounting_events_count} سند یا رویداد مالی ثبت‌شده در دفتر کل")
    if change_requests_count > 0:
        reasons.append(f"{change_requests_count} درخواست تغییرات پرسنلی")
    if section_assignments_count > 0:
        reasons.append(f"{section_assignments_count} انتساب سازمانی به بخش‌ها")
    if has_user_account:
        reasons.append("حساب کاربری فعال متصل به پرسنل در سیستم")

    has_history = len(reasons) > 0

    return {
        'has_history': has_history,
        'can_delete': not has_history,
        'reasons': reasons,
        'details': {
            'attendance_count': attendance_count,
            'payroll_count': payroll_count,
            'accounting_events_count': accounting_events_count,
            'change_requests_count': change_requests_count,
            'section_assignments_count': section_assignments_count,
            'has_user_account': has_user_account,
        }
    }


# ماتریس فیلدهای الزامی بر اساس مرحله گردش کار (تسک ۲.۶ و بخش ۳.۲ طرح مصوب)
REQUIRED_FIELDS_BY_STAGE = {
    'draft': ['first_name', 'last_name', 'national_code', 'section'],
    'submit': [
        'first_name', 'last_name', 'national_code', 'section',
        'father_name', 'birth_date', 'id_number', 'gender', 'job_title',
        'phone_number', 'address', 'postal_code', 'primary_bank_account'
    ],
    'accountant': [
        'first_name', 'last_name', 'national_code', 'section',
        'father_name', 'birth_date', 'id_number', 'gender', 'job_title',
        'phone_number', 'address', 'postal_code', 'primary_bank_account',
        'job_code', 'job_grade', 'contract_type', 'start_date', 'contract_hours',
        'wage_or_salary', 'insurance_type', 'insurance_name', 'exemption_type',
        'tax_payment_type', 'tax_service_location', 'tax_exceptions', 'tax_currency_type'
    ]
}

FIELD_LABELS = {
    'first_name': 'نام',
    'last_name': 'نام خانوادگی',
    'national_code': 'کد ملی معتبر',
    'section': 'بخش/دپارتمان پروژه',
    'father_name': 'نام پدر',
    'birth_date': 'تاریخ تولد',
    'id_number': 'شماره شناسنامه',
    'gender': 'جنسیت',
    'job_title': 'سمت و عنوان شغلی',
    'phone_number': 'شماره همراه (09...)',
    'address': 'نشانی محل سکونت',
    'postal_code': 'کد پستی (۱۰ رقم)',
    'primary_bank_account': 'حساب بانکی اصلی با شبای معتبر',
    'job_code': 'کد ۶ رقمی شغل بیمه',
    'job_grade': 'گروه شغلی',
    'contract_type': 'نوع قرارداد',
    'start_date': 'تاریخ شروع به کار',
    'contract_hours': 'ساعت کار قرارداد شده ماهانه',
    'wage_or_salary': 'دستمزد روزانه یا حقوق ماهانه قرارداد',
    'insurance_type': 'نوع بیمه',
    'insurance_name': 'نام بیمه',
    'exemption_type': 'نوع معافیت مالیاتی',
    'tax_payment_type': 'نوع پرداختی مالیات (WH)',
    'tax_service_location': 'محل خدمت مالیات (WH)',
    'tax_exceptions': 'استثنائات مالیات (WH)',
    'tax_currency_type': 'نوع ارز مالیات (WH)',
}


def get_missing_fields(personnel, stage: str) -> list[dict]:
    """
    تسک ۲.۶: گیت سروری تکمیل مرحله و اعتبارسنجی فیلدهای الزامی.
    stage می‌تواند یکی از مقادیر زیر باشد:
    - 'draft': ذخیره موقت پیش‌نویس
    - 'submit': ارسال کارمند به سرپرست (SoD کامل - فیلدهای دستمزد و کدهای مالیاتی حذف شده‌اند)
    - 'accountant': ارسال حسابدار به مدیر (بررسی کلیه فیلدهای بیمه، مالیات و دستمزد)
    """
    missing = []

    def add_missing(field_name, custom_label=None):
        missing.append({
            'field': field_name,
            'label': custom_label or FIELD_LABELS.get(field_name, field_name)
        })

    # ۱. حداقل‌های پیش‌نویس (draft)
    if not getattr(personnel, 'first_name', None) or not str(personnel.first_name).strip():
        add_missing('first_name')
    if not getattr(personnel, 'last_name', None) or not str(personnel.last_name).strip():
        add_missing('last_name')
    
    ncode = getattr(personnel, 'national_code', None)
    if not ncode:
        add_missing('national_code', 'کد ملی الزامی است')
    else:
        try:
            validate_iranian_national_code(ncode)
        except ValueError as e:
            add_missing('national_code', str(e))

    if not getattr(personnel, 'section_id', None):
        add_missing('section')

    if stage == 'draft':
        return missing

    # ۲. فیلدهای مرحله ارسال کارمند به سرپرست (submit)
    if not getattr(personnel, 'father_name', None) or not str(personnel.father_name).strip():
        add_missing('father_name')
    if not getattr(personnel, 'birth_date', None) or not str(personnel.birth_date).strip():
        add_missing('birth_date')
    if not getattr(personnel, 'id_number', None) or not str(personnel.id_number).strip():
        add_missing('id_number')
    if not getattr(personnel, 'gender', None) or not str(personnel.gender).strip():
        add_missing('gender')
    if not getattr(personnel, 'job_title', None) or not str(personnel.job_title).strip():
        add_missing('job_title')

    phone = getattr(personnel, 'phone_number', None)
    if not phone:
        add_missing('phone_number', 'شماره همراه الزامی است')
    else:
        try:
            validate_iranian_mobile(phone)
        except ValueError as e:
            add_missing('phone_number', str(e))

    if not getattr(personnel, 'address', None) or not str(personnel.address).strip():
        add_missing('address')

    postal = getattr(personnel, 'postal_code', None)
    if not postal:
        add_missing('postal_code', 'کد پستی الزامی است')
    else:
        try:
            validate_iranian_postal_code(postal)
        except ValueError as e:
            add_missing('postal_code', str(e))

    # بررسی حساب اصلی بانکی: DEF-23 — وجود حساب کافی نیست؛ شبای حساب اصلی باید معتبر (Mod 97) باشد
    from personnel.sheba_utils import clean_sheba, validate_sheba
    primary_sheba = None
    if getattr(personnel, 'pk', None) and hasattr(personnel, 'bank_accounts'):
        # .all() از کش prefetch فهرست استفاده می‌کند تا هر سطر یک کوئری جدا نزند (DEF-27)
        primary = next((a for a in personnel.bank_accounts.all() if a.is_primary and a.is_active), None)
        if primary is not None:
            primary_sheba = primary.sheba_number
    if primary_sheba is None:
        # پرونده‌ای که هنوز حساب ندارد: فیلدهای آینه پرسنل
        primary_sheba = getattr(personnel, 'sheba_number', None)
    has_primary_bank = False
    if primary_sheba and str(primary_sheba).strip():
        has_primary_bank = validate_sheba(clean_sheba(primary_sheba))[0]
    if not has_primary_bank:
        add_missing('primary_bank_account')

    if stage in ('submit', 'employee', 'pending_supervisor'):
        return missing

    # ۳. فیلدهای مرحله حسابدار به مدیر (accountant / pending_manager)
    jcode = getattr(personnel, 'job_code', None)
    if not jcode or len(str(jcode).strip()) != 6 or not str(jcode).strip().isdigit():
        add_missing('job_code', 'کد ۶ رقمی شغل تامین اجتماعی الزامی است')

    if not getattr(personnel, 'job_grade', None) or not str(personnel.job_grade).strip():
        add_missing('job_grade')
    if not getattr(personnel, 'contract_type', None):
        add_missing('contract_type')
    if not getattr(personnel, 'start_date', None) or not str(personnel.start_date).strip():
        add_missing('start_date')

    chours = getattr(personnel, 'contract_hours', 0) or 0
    if float(chours) <= 0:
        add_missing('contract_hours', 'ساعت کار ماهانه قرارداد باید بزرگتر از صفر باشد')

    dwage = getattr(personnel, 'daily_base_wage', 0) or 0
    msalary = getattr(personnel, 'contract_base_salary', 0) or 0
    if float(dwage) <= 0 and float(msalary) <= 0:
        add_missing('wage_or_salary', 'حداقل یکی از دستمزد روزانه یا حقوق ماهانه قرارداد باید مشخص باشد')

    if not getattr(personnel, 'insurance_type', None):
        add_missing('insurance_type')
    if not getattr(personnel, 'insurance_name', None) or not str(personnel.insurance_name).strip():
        add_missing('insurance_name')
    if not getattr(personnel, 'exemption_type', None):
        add_missing('exemption_type')

    # کدهای تنظیمی جدول مالیاتی WH
    if not getattr(personnel, 'tax_payment_type', None):
        add_missing('tax_payment_type')
    if not getattr(personnel, 'tax_service_location', None):
        add_missing('tax_service_location')
    if not getattr(personnel, 'tax_exceptions', None):
        add_missing('tax_exceptions')
    if not getattr(personnel, 'tax_currency_type', None):
        add_missing('tax_currency_type')

    return missing


# مدیرِ تاییدکنندهٔ ماژول‌های دیگر (مثلاً نقش سازگاری مدیر) از alias اعلام‌شده می‌آید؛ کد مجوز ماژول دیگر اینجا نیست
MANAGER_APPROVER = permission_alias('manager_approver')

# ── مجوزهای مراحل گردش کار پرسنل: منبع واحد برای تأیید تکی، تأیید گروهی و عبور هوشمند ──
SUPERVISOR_STAGE_PERMS = (
    'perm_approve_personnel_supervisor', 'can_act_as_workshop_supervisor',
    'perm_approve_personnel_manager', MANAGER_APPROVER,
)
FINANCE_STAGE_PERMS = ('perm_approve_personnel_finance',)
MANAGER_STAGE_PERMS = ('perm_approve_personnel_manager', MANAGER_APPROVER, 'can_act_as_company_manager')


def user_has_any_perm(user, codenames) -> bool:
    if not user or not getattr(user, 'is_authenticated', False):
        return False
    if user.is_superuser:
        return True
    return any(user.has_perm(f'accounts.{c}') for c in expand_permission_aliases(codenames))


class PersonnelWorkflowError(Exception):
    """خطای قابل‌نمایش گردش کار پرسنل؛ ویو آن را به پاسخ HTTP تبدیل می‌کند."""

    def __init__(self, message, status_code=400, missing_fields=None, code=None, errors=None):
        super().__init__(message)
        self.message = message
        self.status_code = status_code
        self.missing_fields = missing_fields or []
        self.code = code
        self.errors = errors

    def to_response_data(self) -> dict:
        data = {'error': self.message}
        if self.missing_fields:
            data['missing_fields'] = self.missing_fields
            data['detail'] = f"فیلدهای ناقص: {', '.join(m['label'] for m in self.missing_fields)}"
        if self.errors:
            data['errors'] = self.errors
        if self.code:
            data['code'] = self.code
        return data


# قواعد تأیید هر مرحله (جدول گذار §۳.۱-الف طرح). وضعیت‌های gated_from فقط با گیت مرحلهٔ حسابدار پذیرفته می‌شوند.
APPROVAL_STAGES = {
    'supervisor': {
        'perms': SUPERVISOR_STAGE_PERMS,
        'perm_error': 'شما دسترسی تایید مرحله اول (سرپرست) پرسنل را ندارید.',
        'from': ('pending_supervisor',),
        'gated_from': (),
        'to': 'pending_accountant',
        'action': 'approve_supervisor',
        'actor_role': 'supervisor',
        'status_error': 'پرسنل در وضعیت «{status}» امکان تایید سرپرست ندارد (صرفاً پرونده‌های در انتظار سرپرست قابل تایید هستند).',
        'gate_error': '',
    },
    'finance': {
        'perms': FINANCE_STAGE_PERMS,
        'perm_error': 'شما دسترسی تایید مرحله دوم (مالی) پرسنل را ندارید.',
        'from': (),
        'gated_from': ('pending_accountant', 'supervisor_approved'),
        'to': 'pending_manager',
        'action': 'approve_finance',
        'actor_role': 'accountant',
        'status_error': 'پرسنل در وضعیت «{status}» امکان تایید مالی ندارد.',
        'gate_error': 'فیلدهای الزامی مرحله حسابداری تکمیل نشده است.',
    },
    'manager': {
        'perms': MANAGER_STAGE_PERMS,
        'perm_error': 'شما دسترسی تصویب نهایی (مدیر) پرسنل را ندارید.',
        'from': ('pending_manager', 'accountant_approved'),
        'gated_from': ('pending_supervisor', 'pending_accountant', 'supervisor_approved'),
        'to': 'approved',
        'action': 'approve_manager',
        'actor_role': 'manager',
        'status_error': 'پرسنل در وضعیت «{status}» امکان تصویب نهایی مدیر را ندارد.',
        'gate_error': 'پرونده دارای کسری فیلدهای مرحله مالی است و امکان پرش یا تصویب نهایی بدون تکمیل آن‌ها وجود ندارد.',
    },
}


def approve_personnel_stage(personnel_id, user, stage: str, note: str = '', batch: bool = False):
    """
    DEF-09: سرویس مشترک تأیید تکی و گروهی پرسنل. فهرست سفید وضعیت، مجوز، گیت، لاگ و اعلان یک‌جا اعمال می‌شود.
    بررسی قلمرو با فراخواننده است: فقط شناسه‌ای که از get_queryset کاربر گذشته به این تابع داده می‌شود.
    """
    from django.db import transaction
    from django.utils import timezone
    from personnel.models import PersonnelProfile, WorkflowAuditLog

    rule = APPROVAL_STAGES.get(stage)
    if not rule:
        raise PersonnelWorkflowError('مرحله ارسالی نامعتبر است.')
    if not user_has_any_perm(user, rule['perms']):
        raise PersonnelWorkflowError(rule['perm_error'], status_code=403)

    with transaction.atomic():
        p = PersonnelProfile.objects.select_for_update().filter(pk=personnel_id).first()
        if p is None:
            raise PersonnelWorkflowError('پرونده پرسنلی یافت نشد.', status_code=404)

        old_status = p.approval_status
        if stage == 'manager' and old_status in ('draft', 'revision_required'):
            raise PersonnelWorkflowError('امکان تصویب پرونده در وضعیت پیش‌نویس یا عودت‌داده‌شده برای مدیر وجود ندارد. پرونده ابتدا باید توسط کارمند ارسال و مراحل را طی کند.')
        if old_status not in rule['from'] and old_status not in rule['gated_from']:
            raise PersonnelWorkflowError(rule['status_error'].format(status=p.get_approval_status_display()))
        if old_status in rule['gated_from']:
            missing = get_missing_fields(p, stage='accountant')
            if missing:
                raise PersonnelWorkflowError(rule['gate_error'], missing_fields=missing)

        now = timezone.now()
        p.approval_status = rule['to']
        p.rejection_reason = None
        if stage == 'supervisor':
            p.supervisor_approved_by = user
            p.supervisor_approved_at = now
        elif stage == 'finance':
            p.accountant_approved_by = user
            p.accountant_approved_at = now
        else:
            p.manager_approved_by = user
            p.manager_approved_at = now
            if not p.supervisor_approved_by_id:
                p.supervisor_approved_by = user
                p.supervisor_approved_at = now
            if not p.accountant_approved_by_id:
                p.accountant_approved_by = user
                p.accountant_approved_at = now
            p.is_active = True
        p.save()

        WorkflowAuditLog.objects.create(
            content_type='personnel',
            object_id=p.id,
            actor=user,
            from_status=old_status,
            to_status=rule['to'],
            action=rule['action'],
            reason=note,
            metadata={'approval_note': note, 'actor_role': rule['actor_role'], 'batch': batch, 'company_id': p.company_id}
        )
        dispatch_workflow_notification(p, stage=rule['to'], actor=user, note=note)
    return p


def auto_pass_personnel(personnel, user, source: str = 'create', draft_on_incomplete: bool = False) -> dict:
    """
    DEF-14: عبور هوشمند پرسنل با گیت (جایگزین process_creation_with_auto_pass برای پرسنل؛ موتور مشترک دست نمی‌خورد).
    - عبور سرپرست فقط با گیت submit؛ عبور حسابدار و مدیر فقط با گیت accountant.
    - اگر نقصی باشد پرونده در بالاترین مرحلهٔ مجاز می‌ماند؛ با draft_on_incomplete=True (درون‌ریزی) پیش‌نویس می‌ماند.
    - هر عبور با action='auto_pass' در WorkflowAuditLog ثبت می‌شود.
    خروجی: {'status', 'passed_stages', 'missing_fields'}
    """
    from django.db import transaction
    from django.utils import timezone
    from personnel.models import WorkflowAuditLog

    can_mgr = user_has_any_perm(user, MANAGER_STAGE_PERMS)
    can_fin = user_has_any_perm(user, FINANCE_STAGE_PERMS)
    can_sup = user_has_any_perm(user, SUPERVISOR_STAGE_PERMS)
    if not (can_mgr or can_fin or can_sup):
        return {'status': personnel.approval_status, 'passed_stages': [], 'missing_fields': []}

    submit_missing = get_missing_fields(personnel, stage='submit')
    target = None
    if can_mgr or can_fin:
        acc_missing = get_missing_fields(personnel, stage='accountant')
        missing_out = acc_missing
        if not acc_missing:
            target = 'approved' if can_mgr else 'pending_manager'
        elif not submit_missing and not draft_on_incomplete:
            target = 'pending_accountant'
    else:
        missing_out = submit_missing
        if not submit_missing:
            target = 'pending_accountant'

    if target is None:
        return {'status': personnel.approval_status, 'passed_stages': [], 'missing_fields': missing_out}

    now = timezone.now()
    old_status = personnel.approval_status
    passed = ['supervisor']
    with transaction.atomic():
        personnel.supervisor_approved_by = user
        personnel.supervisor_approved_at = now
        if target in ('pending_manager', 'approved'):
            personnel.accountant_approved_by = user
            personnel.accountant_approved_at = now
            passed.append('accountant')
        if target == 'approved':
            personnel.manager_approved_by = user
            personnel.manager_approved_at = now
            personnel.is_active = True
            passed.append('manager')
        personnel.is_auto_passed = True
        personnel.auto_passed_by = user
        personnel.auto_passed_at = now
        personnel.approval_status = target
        personnel.rejection_reason = None
        personnel.save()

        WorkflowAuditLog.objects.create(
            content_type='personnel',
            object_id=personnel.id,
            actor=user,
            from_status=old_status,
            to_status=target,
            action='auto_pass',
            reason='عبور هوشمند مراحل پیشین هنگام ثبت توسط نقش بالاتر',
            metadata={
                'passed_stages': passed,
                'source': source,
                'company_id': personnel.company_id,
                'missing_fields': missing_out,
            }
        )
        dispatch_workflow_notification(personnel, stage=target, actor=user)
    return {'status': target, 'passed_stages': passed, 'missing_fields': missing_out}


# وضعیت‌هایی که حساب بانکی در آن‌ها برای همه قفل است و تغییر فقط از مسیر درخواست تغییرات است
BANK_ACCOUNT_LOCKED_STATUSES = ('pending_manager', 'accountant_approved', 'manager_approved', 'approved', 'ready_to_pay', 'paid')


def check_bank_account_edit_allowed(user, personnel):
    """
    DEF-10: گارد وضعیت و نقش برای تغییر حساب‌های بانکی پرسنل.
    کارمند فقط در draft و revision_required، حسابدار در pending_accountant، و سوپریوزر جز در وضعیت‌های قفل.
    خروجی: (مجاز؟, پیام خطا)
    """
    st = personnel.approval_status
    if st in BANK_ACCOUNT_LOCKED_STATUSES:
        if st == 'pending_manager':
            return False, 'پرونده در انتظار تصویب مدیر است و حساب‌های بانکی آن قفل است؛ در صورت نیاز، مدیر پرونده را به حسابدار عودت دهد.'
        return False, 'ویرایش یا تغییر حساب‌های بانکی برای پرونده مصوب مجاز نیست؛ لطفاً از طریق ثبت درخواست تغییرات اقدام نمایید.'
    if user and user.is_superuser:
        return True, ''
    if st in ('draft', 'revision_required'):
        return True, ''
    if st == 'pending_accountant' and user_has_any_perm(user, FINANCE_STAGE_PERMS):
        return True, ''
    return False, f'در وضعیت «{personnel.get_approval_status_display()}» مجاز به تغییر حساب‌های بانکی این پرونده نیستید.'


# ── DEF-21: دستمزد و حقوق فقط برای نقش‌های مالی (تصمیم‌های مصوب ۶.۵ و ۶.۹) ──
WAGE_VIEW_PERMS = FINANCE_STAGE_PERMS + MANAGER_STAGE_PERMS + ('view_sys_payroll',)
WAGE_FIELDS = (
    'daily_base_wage', 'daily_seniority_bonus', 'base_daily_rate', 'contract_base_salary',
    'housing_allowance', 'food_allowance', 'spouse_allowance', 'weather_bonus',
    'asaluyeh_parsian_bonus', 'remote_hardship_bonus', 'market_attraction_bonus', 'transport_allowance',
    'effective_daily_rate', 'hourly_rate', 'real_hourly_rate',
)


def user_can_view_wages(user) -> bool:
    return user_has_any_perm(user, WAGE_VIEW_PERMS)


# ── درخواست تغییرات پروندهٔ مصوب (تسک ۳.۴، DEF-16 و DEF-06) ──
# فیلدهایی که هرگز از مسیر درخواست تغییرات عوض نمی‌شوند (هویت کلیدی، قلمرو و فیلدهای گردش کار)
CHANGE_REQUEST_EXCLUDED_FIELDS = {
    'id', 'pk', 'national_code', 'section', 'company', 'approval_status', 'rejection_reason',
    'has_pending_changes', 'is_active', 'created_at', 'updated_at', 'created_by', 'bank_accounts', 'documents',
    # قلمرو و پیوندهای ساختاری از مسیر درخواست تغییرات جابه‌جا نمی‌شوند
    'user', 'assigned_warehouse_id',
}
# فهرست سفید نقش غیرمالی: مشخصات فردی، تماس و بانکی (قفل فیلدهای مالیاتی، بیمه‌ای و دستمزد)
EMPLOYEE_CHANGE_REQUEST_FIELDS = {
    'first_name', 'last_name', 'father_name', 'gender', 'id_number', 'id_series', 'id_serial',
    'birth_date', 'birth_place', 'issue_place', 'issue_date', 'marital_status', 'children_count',
    'phone_number', 'address', 'postal_code', 'bank_name', 'account_number', 'sheba_number',
    'card_number', 'notes', 'valid_until_date', 'recommender',
}
PENDING_CHANGE_REQUEST_STATUSES = ('draft', 'pending_supervisor', 'pending_accountant', 'pending_manager', 'supervisor_approved', 'accountant_approved', 'manager_approved', 'revision_required')


def _diff_repr(value):
    """نمایش یکدست مقدار برای مقایسهٔ diff: خالی و None یکی‌اند و کلید خارجی با شناسه مقایسه می‌شود."""
    if value is None or value == '':
        return None
    if hasattr(value, 'pk') and not isinstance(value, (str, int, float)):
        return value.pk
    return str(value)


def compute_profile_diff(instance, validated_data, excluded=CHANGE_REQUEST_EXCLUDED_FIELDS):
    """
    DEF-16: ابتدا diff محاسبه می‌شود؛ فقط کلیدهای واقعاً تغییرکرده برمی‌گردند.
    خروجی: (diff جدید، مقادیر قبلی)
    """
    diff, old_diff = {}, {}
    for key, new_value in validated_data.items():
        if key in excluded:
            continue
        old_rep = _diff_repr(getattr(instance, key, None))
        new_rep = _diff_repr(new_value)
        if old_rep != new_rep:
            diff[key] = new_rep
            old_diff[key] = old_rep
    return diff, old_diff


def change_request_fields_allowed_for(user):
    """فیلدهای مجاز درخواست تغییرات برای این کاربر؛ None یعنی همهٔ فیلدهای غیرمستثنی (نقش مالی)."""
    if user_can_view_wages(user):
        return None
    return EMPLOYEE_CHANGE_REQUEST_FIELDS


def log_change_request_event(cr, actor, action: str, reason: str = '', extra: dict = None):
    """
    DEF-19: لاگ اکشن‌های درخواست تغییرات روی تایملاین همان پرونده (content_type='personnel').
    Diff Viewer با change_request_id همین لاگ‌ها را پیدا می‌کند.
    """
    from personnel.models import WorkflowAuditLog
    personnel = cr.personnel
    metadata = {'change_request_id': cr.id, 'change_request_status': cr.status, 'company_id': personnel.company_id}
    if extra:
        metadata.update(extra)
    WorkflowAuditLog.objects.create(
        content_type='personnel',
        object_id=personnel.id,
        actor=actor if getattr(actor, 'is_authenticated', False) else None,
        from_status=personnel.approval_status,
        to_status=personnel.approval_status,
        action=action,
        reason=reason or '',
        metadata=metadata,
    )


def refresh_has_pending_changes(personnel, exclude_cr_id=None):
    from personnel.models import PersonnelChangeRequest
    pending = PersonnelChangeRequest.objects.filter(personnel=personnel, status__in=PENDING_CHANGE_REQUEST_STATUSES)
    if exclude_cr_id:
        pending = pending.exclude(pk=exclude_cr_id)
    has_pending = pending.exists()
    if personnel.has_pending_changes != has_pending:
        personnel.has_pending_changes = has_pending
        personnel.save(update_fields=['has_pending_changes'])


def apply_personnel_change_request(cr, request=None):
    """
    DEF-06: منبع واحد اعمال درخواست تغییرات روی پرونده (هم برای approve-manager ویو و هم موتور گردش کار).
    - کلیدهای مستثنی پیش از اعتبارسنجی حذف می‌شوند و فهرست سفید نقشِ درخواست‌دهنده دوباره سنجیده می‌شود.
    - اعتبارسنجی با PersonnelProfileSerializer و context؛ کلید خارجی (مثل project) با شناسه پذیرفته می‌شود.
    - اگر نامعتبر باشد PersonnelWorkflowError می‌دهد و هیچ setattr خامی انجام نمی‌شود؛ درخواست معلق می‌ماند.
    """
    from personnel.serializers import PersonnelProfileSerializer

    personnel = cr.personnel
    proposed = dict(cr.proposed_changes or {})
    clean_data = {k: v for k, v in proposed.items() if k not in CHANGE_REQUEST_EXCLUDED_FIELDS}

    allowed = change_request_fields_allowed_for(cr.requested_by) if cr.requested_by_id else None
    if allowed is not None:
        locked = sorted(k for k in clean_data if k not in allowed)
        if locked:
            raise PersonnelWorkflowError(
                f"درخواست تغییرات شامل فیلدهای قفل برای نقش درخواست‌دهنده است: {', '.join(locked)}"
            )

    context = {'request': request} if request is not None else {}
    context['apply_change_request'] = True
    serializer = PersonnelProfileSerializer(personnel, data=clean_data, partial=True, context=context)
    if not serializer.is_valid():
        raise PersonnelWorkflowError(
            'تغییرات پیشنهادی معتبر نیست و اعمال نشد؛ درخواست معلق ماند.',
            code='invalid_change_request',
            errors=serializer.errors,
        )
    serializer.save()
    return personnel


# ── DEF-05 و DEF-04: برآورد سروری هزینهٔ ماهانهٔ پرسنل با همان موتور حقوق ──
COST_ESTIMATE_OVERRIDE_FIELDS = {
    'job_grade': 'str', 'daily_base_wage': 'decimal', 'daily_seniority_bonus': 'decimal',
    'base_years_experience': 'int', 'contract_hours': 'decimal', 'contract_base_salary': 'decimal',
    'marital_status': 'str', 'children_count': 'int', 'include_in_insurance': 'bool',
}


def current_shamsi_year_month() -> str:
    today = jdatetime.date.today()
    return f"{today.year:04d}/{today.month:02d}"


def _coerce_override(kind, value):
    from decimal import Decimal, InvalidOperation
    if kind == 'bool':
        return value if isinstance(value, bool) else str(value).strip().lower() in ('1', 'true', 'yes', 'بله')
    raw = normalize_digits(str(value if value is not None else '')).replace(',', '').strip()
    if kind == 'str':
        return raw
    if raw == '':
        raw = '0'
    try:
        return int(Decimal(raw)) if kind == 'int' else Decimal(raw)
    except (InvalidOperation, ValueError):
        raise ValueError(f'مقدار «{value}» عددی معتبر نیست.')


def estimate_monthly_personnel_cost(personnel, year_month: str = None, overrides: dict = None, cache: dict = None) -> dict:
    """
    DEF-05: برآورد هزینهٔ یک ماه کامل استاندارد با PayrollCalculationEngine.calculate_single_employee،
    تا ماشین‌حساب حسابدار و ویجت مدیر یک منبع داشته باشند: حق تأهل فقط برای متأهل، حق اولاد برابر
    تعداد فرزند × حق اولاد ماهانه، و پایهٔ بیمه شامل همهٔ مزایای مشمول بیمه با سقف قانونی.
    overrides مقادیر فرم ذخیره‌نشده است و روی یک کپی ذخیره‌نشده اعمال می‌شود؛ پرونده تغییر نمی‌کند.
    """
    import copy
    from types import SimpleNamespace
    from personnel.payroll_engine import PayrollCalculationEngine, get_effective_payroll_settings

    year_month = year_month or current_shamsi_year_month()
    parts = normalize_digits(str(year_month)).replace('-', '/').split('/')
    if len(parts) < 2 or not parts[0].isdigit() or not parts[1].isdigit() or not (1 <= int(parts[1]) <= 12):
        raise ValueError('ماه محاسباتی باید در قالب YYYY/MM باشد.')
    year_month = f"{int(parts[0]):04d}/{int(parts[1]):02d}"
    month_num = int(parts[1])
    # همان طول ماه موتور حقوق (calculate_period)
    month_days = 31 if month_num <= 6 else (30 if month_num <= 11 else 29)

    est = copy.copy(personnel)
    est.id = None  # بدون کارکرد ثبت‌شده: موتور ماه کامل استاندارد را فرض می‌کند
    for key, value in (overrides or {}).items():
        kind = COST_ESTIMATE_OVERRIDE_FIELDS.get(key)
        if kind:
            setattr(est, key, _coerce_override(kind, value))
    if est.marital_status not in ('single', 'married'):
        est.marital_status = 'single'

    cache = cache if cache is not None else {}
    cache_key = (year_month, est.project_id, est.company_id)
    if cache_key not in cache:
        settings = get_effective_payroll_settings(year_month, est.project_id, est.company_id)
        grades = {}
        if settings:
            for jg in settings.job_grades.all():
                grades[str(jg.grade_number)] = (jg.daily_base_wage, jg.daily_seniority_bonus)
        cache[cache_key] = (settings, grades)
    settings, grades = cache[cache_key]

    data = PayrollCalculationEngine().calculate_single_employee(
        personnel=est,
        period=SimpleNamespace(year_month=year_month),
        settings=settings,
        job_grades=grades,
        month_days=month_days,
    )
    std_hours = float(settings.standard_daily_hours) if settings and settings.standard_daily_hours else 10.0
    gross = int(data['gross_salary'])
    employer_insurance = int(data['employer_insurance'])
    unemployment_insurance = int(data['unemployment_insurance'])
    return {
        'year_month': year_month,
        'settings_id': settings.id if settings else None,
        'job_grade': data['job_grade'],
        'base_daily_rate': int(data['base_daily_rate']),
        'insurance_base_hourly_rate': round(float(data['base_daily_rate']) / std_hours, 2) if std_hours else None,
        'real_hourly_rate': est.real_hourly_rate,
        'insurance_days': int(data['insurance_days']),
        'gross_salary': gross,
        'base_salary': int(data['base_salary']),
        'marital_allowance': int(data['marital_allowance']),
        'child_allowance': int(data['child_allowance']),
        'housing_allowance': int(data['housing_allowance']),
        'food_allowance': int(data['food_allowance']),
        'total_insurable_salary_allowances': int(data['total_insurable_salary_allowances']),
        'worker_insurance': int(data['worker_insurance']),
        'employer_insurance': employer_insurance,
        'unemployment_insurance': unemployment_insurance,
        'net_salary': int(data['net_salary']),
        'employer_monthly_cost': gross + employer_insurance + unemployment_insurance,
        'employer_insurance_rate': float(settings.employer_insurance_rate) if settings else 20.0,
        'unemployment_insurance_rate': float(settings.unemployment_insurance_rate) if settings else 3.0,
        'worker_insurance_rate': float(settings.worker_insurance_rate) if settings else 7.0,
    }


def _snapshot_value(value):
    import datetime
    from decimal import Decimal
    from django.db.models.fields.files import FieldFile

    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    if isinstance(value, FieldFile):
        return value.name or None
    return str(value)


def build_personnel_snapshot(personnel) -> dict:
    """
    DEF-11: عکس کامل پرونده پیش از حذف یا Purge (شرط ۳ بخش ۳.۱-ب طرح):
    همهٔ فیلدهای پرونده، همهٔ حساب‌های بانکی با همهٔ فیلدها، فهرست مدارک و مسیر attachment.
    کلیدهای خلاصهٔ قبلی برای سازگاری با نمایش تایملاین حفظ شده‌اند.
    """
    profile = {f.attname: _snapshot_value(getattr(personnel, f.attname)) for f in personnel._meta.concrete_fields}
    bank_accounts = [
        {f.attname: _snapshot_value(getattr(acc, f.attname)) for f in acc._meta.concrete_fields}
        for acc in personnel.bank_accounts.all().order_by('id')
    ]
    documents = []
    for doc in personnel.documents.all().select_related('uploaded_by').order_by('id'):
        documents.append({
            'id': doc.id,
            'doc_type': doc.doc_type,
            'file': doc.file.name if doc.file else None,
            'uploaded_by_id': doc.uploaded_by_id,
            'uploaded_by_username': doc.uploaded_by.username if doc.uploaded_by else None,
            'created_at': _snapshot_value(doc.created_at),
        })
    return {
        'snapshot_version': 2,
        'id': personnel.id,
        'full_name': personnel.full_name,
        'national_code': personnel.national_code,
        'section_id': personnel.section_id,
        'company_id': personnel.company_id,
        'approval_status': personnel.approval_status,
        'profile': profile,
        'bank_accounts': bank_accounts,
        'documents': documents,
        'attachment': personnel.attachment.name if personnel.attachment else None,
    }


def _workflow_recipient_ids(stage: str, company_id, section_id, created_by_id, perms_by_stage: dict = None) -> set:
    """
    DEF-12: گیرندگان اعلان هر مرحله در یک کوئری.
    UserCompanyAccess فیلد is_active ندارد؛ دسترسی شرکت همان قاعدهٔ get_user_allowed_companies است:
    دسترسی کامل صریح (workspace_full) یا انتساب فعال به خود شرکت، یکی از پروژه‌ها یا یکی از بخش‌های آن.
    """
    from django.contrib.auth import get_user_model
    from django.db.models import Q

    if stage in ('revision_required', 'approved', 'rejected', 'deleted', 'purged', 'end_collaboration'):
        return {created_by_id} if created_by_id else set()

    if stage == 'reopened':
        recipients = {created_by_id} if created_by_id else set()
        sup_ids = _workflow_recipient_ids('pending_supervisor', company_id, section_id, created_by_id, perms_by_stage)
        return recipients | sup_ids

    if perms_by_stage is None:
        perms_by_stage = {
            'pending_supervisor': ('perm_approve_personnel_supervisor', 'can_act_as_workshop_supervisor'),
            'pending_accountant': ('perm_approve_personnel_finance',),
            'pending_manager': ('perm_approve_personnel_manager', MANAGER_APPROVER, 'can_act_as_company_manager'),
        }
    perm_codes = expand_permission_aliases(perms_by_stage.get(stage) or ())
    if not perm_codes:
        return set()

    User = get_user_model()
    has_perm_q = (
        Q(is_superuser=True)
        | Q(user_permissions__codename__in=perm_codes)
        | Q(groups__permissions__codename__in=perm_codes)
    )
    users = User.objects.filter(has_perm_q, is_active=True)
    if company_id:
        # عضو شرکت فقط از منبع‌های خودکار (طرح مجوزهای بخش‌به‌بخش، ق۲)
        from platform_core.registry import contributions
        # انتساب سازمانی در هر سطح (شرکت، پروژه یا بخش) از سازندهٔ هسته (build_section_member_q) می‌آید
        in_company_q = Q(is_superuser=True)
        for builder in contributions('company_member_q_builders'):
            builder_q = builder({company_id})
            if builder_q is not None:
                in_company_q |= builder_q
        users = users.filter(in_company_q)
    else:
        users = users.filter(is_superuser=True)
    candidate_ids = set(users.values_list('id', flat=True).distinct())

    # سرپرست: اگر کاربر انتساب سازمانی دارد، فقط کسی اعلان می‌گیرد که بخشِ پرونده را پوشش می‌دهد
    # (انتساب مستقیم بخش یا انتساب سطح پروژه یا شرکتِ آن؛ طرح ۲.۱)
    if stage in ('pending_supervisor', 'change_request_created') and section_id and candidate_ids:
        from personnel.models import SectionAssignment
        from companies.org_roles import assignments_covering_section
        superuser_ids = set(User.objects.filter(id__in=candidate_ids, is_superuser=True).values_list('id', flat=True))
        users_with_any_section = set(
            SectionAssignment.objects.filter(user_id__in=candidate_ids, is_active=True).values_list('user_id', flat=True)
        )
        users_in_this_section = set(
            assignments_covering_section(section_id).filter(user_id__in=candidate_ids).values_list('user_id', flat=True)
        )
        candidate_ids = {
            uid for uid in candidate_ids
            if uid in superuser_ids or uid in users_in_this_section or uid not in users_with_any_section
        }
    return candidate_ids


WORKFLOW_NOTIFICATION_TEXTS = {
    'pending_supervisor': ('پرونده پرسنل در انتظار تایید سرپرست', 'پرونده جدید «{name}» جهت بررسی و تایید سرپرستی ارسال گردید.'),
    'pending_accountant': ('پرونده پرسنل در انتظار بررسی حسابدار', 'پرونده پرسنل «{name}» به تایید سرپرست رسید و جهت تکمیل اطلاعات مالی به کارتابل حسابداری ارسال شد.'),
    'pending_manager': ('پرونده پرسنل در انتظار تصویب مدیر', 'پرونده پرسنل «{name}» به تایید مالی رسید و جهت تصویب نهایی در کارتابل مدیر قرار گرفت.'),
    'approved': ('تصویب نهایی پرونده پرسنل', 'پرونده پرسنل «{name}» با موفقیت به تصویب مدیر رسید و فعال گردید.'),
    'revision_required': ('عودت پرونده پرسنل جهت بازنگری', 'پرونده پرسنل «{name}» جهت بازنگری و اصلاح عودت داده شد{reason}.'),
    'rejected': ('رد پرونده پرسنل', 'پرونده پرسنل «{name}» رد شد{reason}.'),
    'purged': ('پاک‌سازی پرونده پرسنل', 'پرونده پرسنل «{name}» به دلیل رد قطعی و عدم سابقه پاک‌سازی شد.'),
    'end_collaboration': ('قطع همکاری پرسنل', 'همکاری پرسنل «{name}» پایان یافت و پرونده غیرفعال شد{reason}.'),
}
RETURN_NOTIFICATION_TEXT = ('عودت پرونده پرسنل جهت بازنگری', 'پرونده پرسنل «{name}» جهت بازنگری به کارتابل شما عودت داده شد{reason}.')


def dispatch_workflow_notification(personnel, stage: str, actor=None, reason=None, note=None, is_return=False):
    """
    تسک ۴.۱ و ۴.۲: ارسال اعلان‌های هدفمند گردش کار پرسنل و انتشار به وب‌سوکت کانال اختصاصی کاربر (user_{id}).
    - stage همان وضعیت مقصد است؛ گیرندگان کاربران دارای مجوز همان مرحله در شرکت پرونده‌اند.
    - is_return=True یعنی عودت: گیرنده بر اساس مقصد تعیین می‌شود (کارمند، سرپرست یا حسابدار)، نه همیشه ثبت‌کننده (DEF-12).
    - مقادیر پرونده همین حالا خوانده می‌شوند تا پس از حذف (Purge) هم شناسهٔ درست ثبت شود.
    - اجرای نهایی پس از commit موفق تراکنش (transaction.on_commit).
    """
    if not personnel:
        return

    from django.db import transaction

    pers_id = getattr(personnel, 'id', None)
    cid = getattr(personnel, 'company_id', None)
    if not cid and getattr(personnel, 'section_id', None) and personnel.section and personnel.section.project:
        cid = personnel.section.project.company_id
    sec_id = getattr(personnel, 'section_id', None)
    created_by_id = getattr(personnel, 'created_by_id', None)
    full_name = getattr(personnel, 'full_name', '') or 'پرسنل'
    approval_status = getattr(personnel, 'approval_status', '')
    actor_id = getattr(actor, 'id', None)

    def _execute():
        import logging
        try:
            from notifications.models import Notification
            from channels.layers import get_channel_layer
            from asgiref.sync import async_to_sync

            if is_return and stage != 'revision_required':
                title_tpl, message_tpl = RETURN_NOTIFICATION_TEXT
            elif stage in WORKFLOW_NOTIFICATION_TEXTS:
                title_tpl, message_tpl = WORKFLOW_NOTIFICATION_TEXTS[stage]
            elif stage == 'deleted':
                title_tpl, message_tpl = WORKFLOW_NOTIFICATION_TEXTS['purged']
            else:
                return
            reason_txt = f' (دلیل: {reason})' if reason else ''
            title = title_tpl
            message = message_tpl.format(name=full_name, reason=reason_txt)

            target_user_ids = _workflow_recipient_ids(stage, cid, sec_id, created_by_id)
            # حذف اقدام‌کننده از گیرندگان
            if actor_id:
                target_user_ids.discard(actor_id)
            if not target_user_ids:
                return

            channel_layer = get_channel_layer()
            pers_id_str = str(pers_id) if pers_id is not None else ''

            for uid in sorted(target_user_ids):
                notif = Notification.objects.create(
                    recipient_id=uid,
                    title=title,
                    message=message,
                    notification_type='personnel_workflow',
                    target_model='PersonnelProfile',
                    target_id=pers_id_str,
                )
                if channel_layer:
                    async_to_sync(channel_layer.group_send)(
                        f'user_{uid}',
                        {
                            'type': 'send_notification',
                            'type_str': 'personnel_workflow',
                            'notification_id': notif.id,
                            'title': title,
                            'message': message,
                            'target_model': 'PersonnelProfile',
                            'target_id': pers_id_str,
                            'personnel_id': pers_id,
                            'company_id': cid,
                            'approval_status': approval_status,
                            'created_at': notif.created_at.isoformat(),
                        }
                    )
        except Exception:
            logging.getLogger(__name__).exception(
                "[Notification] Error dispatching workflow notification (stage=%s, personnel=%s)", stage, pers_id
            )

    transaction.on_commit(_execute)


def get_personnel_missing_primary_sheba(queryset=None) -> list[dict]:
    """
    DEF-23: واکشی فقط‌خواندنی پرونده‌های پرسنل که حساب اصلی آن‌ها شبای معتبر (Mod 97) ندارد.
    ستون‌ها: شرکت، بخش، نام، کد ملی، وضعیت، شماره حساب، شبای فعلی و راهنمای تکمیل.
    کاملاً فقط‌خواندنی بدون هیچ‌گونه ویرایش یا حذف در دیتابیس.
    """
    from personnel.models import PersonnelProfile
    from personnel.sheba_utils import clean_sheba, validate_sheba

    if queryset is None:
        qs = PersonnelProfile.objects.all()
    else:
        qs = queryset

    qs = qs.select_related('company', 'section', 'section__project', 'section__project__company').prefetch_related('bank_accounts')

    results = []
    for p in qs:
        # حساب اصلی فعال در صورت وجود چند حسابی
        primary = next((a for a in p.bank_accounts.all() if a.is_primary and a.is_active), None)
        if primary is not None:
            account_no = primary.account_number or ''
            sheba_no = primary.sheba_number or ''
        else:
            account_no = getattr(p, 'account_number', '') or ''
            sheba_no = getattr(p, 'sheba_number', '') or ''

        # بررسی اعتبار شبا (ISO 7064 Mod 97)
        is_valid_sheba = False
        cleaned = clean_sheba(sheba_no)
        if cleaned:
            is_valid_sheba = validate_sheba(cleaned)[0]

        if not is_valid_sheba:
            comp_name = '-'
            if p.company:
                comp_name = p.company.name
            elif p.section and p.section.project and p.section.project.company:
                comp_name = p.section.project.company.name

            sec_name = p.section.name if p.section else '-'
            status_display = p.get_approval_status_display() if hasattr(p, 'get_approval_status_display') else p.approval_status

            results.append({
                'personnel_id': p.id,
                'company_name': comp_name,
                'section_name': sec_name,
                'full_name': p.full_name or '',
                'national_code': p.national_code or '',
                'approval_status': p.approval_status or '',
                'approval_status_display': status_display or '',
                'is_active': p.is_active,
                'is_active_display': 'فعال' if p.is_active else 'غیرفعال',
                'account_number': account_no,
                'sheba_number': sheba_no,
                'completion_guide': 'درخواست تغییرات (سرپرست/حسابدار/مدیر)',
            })

    return results


def generate_missing_primary_sheba_excel(queryset=None) -> bytes:
    """
    DEF-23: تولید فایل اکسل گزارش پرونده‌های دارای حساب اصلی بدون شبای معتبر.
    رعایت معماری ۲ سطری هدر: سطر ۱ عناوین فارسی، سطر ۲ کلیدهای دیتابیس در قلم خاکستری، freeze_panes = 'A3'.
    """
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    data = get_personnel_missing_primary_sheba(queryset=queryset)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "فاقد شبا"
    ws.views.sheetView[0].rightToLeft = True
    ws.freeze_panes = 'A3'

    row1_titles = [
        "ردیف",
        "شرکت",
        "بخش / پروژه",
        "نام و نام خانوادگی",
        "کد ملی",
        "وضعیت پرونده",
        "وضعیت فعالیت",
        "شماره حساب",
        "شماره شبا (فعلی)",
        "مسیر قانونی تکمیل اطلاعات"
    ]

    row2_keys = [
        "row_number",
        "company_name",
        "section_name",
        "full_name",
        "national_code",
        "approval_status",
        "is_active",
        "account_number",
        "sheba_number",
        "completion_guide"
    ]

    header_font_r1 = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")
    header_fill_r1 = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font_r2 = Font(name="Courier New", size=9, bold=False, color="94A3B8")
    header_fill_r2 = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(row1_titles)
    for col_idx in range(1, len(row1_titles) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font_r1
        cell.fill = header_fill_r1
        cell.alignment = align_center
        cell.border = border_thin

    ws.append(row2_keys)
    for col_idx in range(1, len(row2_keys) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font_r2
        cell.fill = header_fill_r2
        cell.alignment = align_center
        cell.border = border_thin

    data_font = Font(name="B Nazanin", size=10)
    for idx, row in enumerate(data, start=1):
        r_num = idx + 2
        values = [
            idx,
            row['company_name'],
            row['section_name'],
            row['full_name'],
            row['national_code'],
            row['approval_status_display'],
            row['is_active_display'],
            row['account_number'],
            row['sheba_number'],
            row['completion_guide']
        ]
        ws.append(values)
        for col_idx in range(1, len(values) + 1):
            cell = ws.cell(row=r_num, column=col_idx)
            cell.font = data_font
            cell.border = border_thin
            if col_idx in (1, 5, 6, 7, 8, 9):
                cell.alignment = align_center
            else:
                cell.alignment = align_right

    # تنظیم عرض ستون‌ها
    col_widths = {
        'A': 8,
        'B': 24,
        'C': 22,
        'D': 24,
        'E': 16,
        'F': 18,
        'G': 14,
        'H': 22,
        'I': 28,
        'J': 35,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

