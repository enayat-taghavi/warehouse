# -*- coding: utf-8 -*-
"""
آزمون‌های نقش «راهبر تشکیلات» و مجوزهای طرف‌حساب حسابدار
مرجع: Documents/Organization_Admin_And_Menu_Architecture/implementation_plan_org_admin_menu.md §۷
از گام ۲ طرح Project_Section_To_Core راهبر تشکیلات دارندهٔ سطح «نوشتن» بخش هستهٔ projects است
(جای مجوز حذف‌شدهٔ perm_manage_projects_sections) و API پروژه و بخش زیر /api/org/ است.
"""
from companies.testing import grant_membership
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from rest_framework.test import APIClient

from accounts.middleware import get_user_valid_roles_for_app, resolve_effective_role
from personnel.models import (
    Company,
    Counterparty,
    Project,
    ProjectSection,
    UserCompanyAccess,
    SectionAssignment,
)

User = get_user_model()

PROJECTS_WRITE = {'projects': 'write'}


class OrgAdminAndCounterpartyPermissionsTestCase(TestCase):

    def setUp(self):
        self.client = APIClient()
        self.company = Company.objects.create(
            code='ORG_TEST',
            name='شرکت آزمون راهبر',
            national_id='14001234567',
            economic_code='411123456789',
            is_active=True
        )
        self.other_company = Company.objects.create(
            code='ORG_OTHER',
            name='شرکت دیگر',
            national_id='14007654321',
            economic_code='411987654321',
            is_active=True
        )
        self.project = Project.objects.create(company=self.company, code='PRJ_ORG', name='پروژه آزمون')
        self.section = ProjectSection.objects.create(project=self.project, code='SEC_ORG', name='بخش آزمون')
        self.counterparty = Counterparty.objects.create(company=self.company, name='طرف‌حساب آزمون', counterparty_type='other')

        self.org_admin = self._user('org_admin_user', [], levels=PROJECTS_WRITE)
        self.accountant = self._user('acc_user', ['can_act_as_accountant', 'view_sys_payroll'])
        self.finance_approver = self._user('fin_user', ['perm_approve_personnel_finance'])
        self.plain = self._user('plain_user', [])

    def _user(self, username, codenames, levels=None):
        u = User.objects.create_user(username=username, password='password123', email=f'{username}@test.com')
        UserCompanyAccess.objects.create(user=u, company=self.company, is_default=True, section_levels=levels or {})
        grant_membership(u, self.company)
        for codename in codenames:
            u.user_permissions.add(Permission.objects.get(codename=codename, content_type__app_label='accounts'))
        return User.objects.get(pk=u.pk)

    # ─── نقش‌های معتبر ───
    def test_org_admin_is_valid_role(self):
        self.assertEqual(get_user_valid_roles_for_app(self.org_admin, 'accounting'), ['org_admin'])

    def test_org_admin_does_not_get_operator_fallback(self):
        u = self._user('org_admin_personnel', ['view_sys_personnel'], levels=PROJECTS_WRITE)
        self.assertEqual(get_user_valid_roles_for_app(u, 'accounting'), ['org_admin'])

    def test_org_admin_first_so_default_role_is_unchanged(self):
        u = self._user('org_admin_accountant', ['can_act_as_accountant'], levels=PROJECTS_WRITE)
        roles = get_user_valid_roles_for_app(u, 'accounting')
        self.assertEqual(roles, ['org_admin', 'accountant'])
        self.assertEqual(resolve_effective_role(u, '', 'accounting'), 'accountant')
        self.assertEqual(resolve_effective_role(u, 'org_admin', 'accounting'), 'org_admin')

    def test_superuser_can_act_as_org_admin(self):
        su = User.objects.create_superuser(username='su_org', password='password123', email='su@test.com')
        self.assertIn('org_admin', get_user_valid_roles_for_app(su, 'accounting'))

    # ─── ساختار پروژه‌ها و بخش‌ها ───
    def test_org_admin_can_write_project_structure(self):
        self.client.force_authenticate(user=self.org_admin)
        res = self.client.post('/api/org/projects/', {'company': self.company.id, 'code': 'PRJ_NEW', 'name': 'پروژه تازه'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)

        res = self.client.post('/api/org/project-sections/', {'project': self.project.id, 'code': 'SEC_NEW', 'name': 'بخش تازه'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)

        res = self.client.post('/api/org/section-assignments/', {'user': self.plain.id, 'section': self.section.id, 'role': 'employee'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)

    def test_plain_user_cannot_write_project_structure(self):
        self.client.force_authenticate(user=self.plain)
        res = self.client.post('/api/org/projects/', {'company': self.company.id, 'code': 'PRJ_X', 'name': 'x'}, format='json')
        self.assertEqual(res.status_code, 403)

    def test_org_admin_cannot_grant_company_access(self):
        self.client.force_authenticate(user=self.org_admin)
        res = self.client.post('/api/personnel/user-company-access/', {'user': self.plain.id, 'company': self.other_company.id}, format='json')
        self.assertEqual(res.status_code, 403)

    # ─── فهرست کاربران برای انتساب ───
    def test_org_admin_can_list_but_not_edit_users(self):
        self.client.force_authenticate(user=self.org_admin)
        self.assertEqual(self.client.get('/api/auth/users/').status_code, 200)
        res = self.client.patch(f'/api/auth/users/{self.plain.id}/', {'first_name': 'تغییر'}, format='json')
        self.assertEqual(res.status_code, 403)
        self.assertEqual(self.client.get(f'/api/auth/users/{self.plain.id}/').status_code, 403)

    # ─── طرف‌حساب‌ها ───
    def test_accountant_can_edit_and_delete_counterparty(self):
        self.client.force_authenticate(user=self.accountant)
        res = self.client.patch(f'/api/personnel/counterparties/{self.counterparty.id}/', {'phone': '09120000000'}, format='json')
        self.assertEqual(res.status_code, 200, res.content)
        res = self.client.delete(f'/api/personnel/counterparties/{self.counterparty.id}/')
        self.assertEqual(res.status_code, 204, res.content)

    def test_finance_approver_can_edit_counterparty(self):
        self.client.force_authenticate(user=self.finance_approver)
        res = self.client.patch(f'/api/personnel/counterparties/{self.counterparty.id}/', {'phone': '09120000001'}, format='json')
        self.assertEqual(res.status_code, 200, res.content)

    def test_plain_user_cannot_edit_counterparty(self):
        self.client.force_authenticate(user=self.plain)
        res = self.client.patch(f'/api/personnel/counterparties/{self.counterparty.id}/', {'phone': '09120000002'}, format='json')
        self.assertEqual(res.status_code, 403)

    def test_org_admin_cannot_edit_counterparty(self):
        self.client.force_authenticate(user=self.org_admin)
        res = self.client.patch(f'/api/personnel/counterparties/{self.counterparty.id}/', {'phone': '09120000003'}, format='json')
        self.assertEqual(res.status_code, 403)


class OrgAdminAndAccountantTenantScopeTestCase(TestCase):
    """
    یافته‌های بازبینی تفکیک داده (۹ مهر ۱۴۰۵) و تصمیم کاربر: هر نویسندهٔ غیرسوپریوزر (راهبر تشکیلات، حسابدار و مدیر)
    فقط درون شرکت‌های مجاز خودش می‌نویسد، به رکورد سراسری (بدون شرکت) دست نمی‌زند، فقط اعضای همان شرکت را منتسب می‌کند
    و درون‌ریزی اکسلش به شرکت فعالِ مجاز محدود است.
    """

    def setUp(self):
        self.client = APIClient()
        self.company = Company.objects.create(
            code='SCOPE_A', name='شرکت الف', national_id='14001111111', economic_code='411111111111', is_active=True
        )
        self.other_company = Company.objects.create(
            code='SCOPE_B', name='شرکت ب', national_id='14002222222', economic_code='411222222222', is_active=True
        )
        self.project = Project.objects.create(company=self.company, code='PRJ_A', name='پروژه الف')
        self.section = ProjectSection.objects.create(project=self.project, code='SEC_A', name='بخش الف')
        self.other_project = Project.objects.create(company=self.other_company, code='PRJ_B', name='پروژه ب')
        self.other_section = ProjectSection.objects.create(project=self.other_project, code='SEC_B', name='بخش ب')
        # پروژه و بخش سراسری (بی‌شرکت) دیگر وجود ندارد (طرح ۲.۱، ت۲)
        self.counterparty = Counterparty.objects.create(company=self.company, name='طرف‌حساب الف', counterparty_type='other')
        self.global_counterparty = Counterparty.objects.create(company=None, name='طرف‌حساب سراسری', counterparty_type='other')

        self.org_admin = self._user('scope_org_admin', [], self.company, levels=PROJECTS_WRITE)
        self.accountant = self._user('scope_accountant', ['can_act_as_accountant', 'view_sys_payroll'], self.company)
        self.manager = self._user('scope_manager', ['perm_approve_personnel_manager'], self.company, levels=PROJECTS_WRITE)
        self.member = self._user('scope_member', [], self.company)
        self.outsider = self._user('scope_outsider', [], self.other_company)

    def _user(self, username, codenames, company, levels=None):
        u = User.objects.create_user(username=username, password='password123', email=f'{username}@test.com')
        UserCompanyAccess.objects.create(user=u, company=company, is_default=True, section_levels=levels or {})
        grant_membership(u, company)
        for codename in codenames:
            u.user_permissions.add(Permission.objects.get(codename=codename, content_type__app_label='accounts'))
        return User.objects.get(pk=u.pk)

    def _as(self, user):
        self.client.force_authenticate(user=user)

    @staticmethod
    def _rows(res):
        body = res.json()
        return body.get('results', body) if isinstance(body, dict) else body

    @staticmethod
    def _xlsx(rows):
        """فایل اکسل با دو سطر سرآیند (قالب موتور org_excel_engine) و سطرهای داده از ردیف ۳"""
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['header'])
        ws.append(['keys'])
        for row in rows:
            ws.append(list(row))
        buf = io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile('import.xlsx', buf.getvalue())

    def _import(self, url, rows, company):
        return self.client.post(url, {'file': self._xlsx(rows)}, format='multipart', HTTP_X_COMPANY_ID=str(company.id))

    # ─── راهبر: پروژه ───
    def test_org_admin_project_writes_stay_inside_own_company(self):
        self._as(self.org_admin)
        url = '/api/org/projects/'
        self.assertEqual(self.client.post(url, {'company': self.other_company.id, 'code': 'PRJ_X1', 'name': 'x'}, format='json').status_code, 403)
        # پروژهٔ بی‌شرکت ممکن نیست: شرکت فیلد الزامی است (طرح ۲.۱، ت۲)
        self.assertEqual(self.client.post(url, {'code': 'PRJ_X2', 'name': 'x'}, format='json').status_code, 400)
        self.assertEqual(self.client.patch(f'{url}{self.other_project.id}/', {'name': 'تغییر'}, format='json').status_code, 404)
        self.assertEqual(self.client.delete(f'{url}{self.other_project.id}/').status_code, 404)
        self.assertEqual(self.client.patch(f'{url}{self.project.id}/', {'company': None}, format='json').status_code, 400)
        self.assertEqual(self.client.patch(f'{url}{self.project.id}/', {'company': self.other_company.id}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'{url}{self.project.id}/', {'name': 'نام تازه'}, format='json').status_code, 200)
        self.other_project.refresh_from_db()
        self.assertEqual(self.other_project.name, 'پروژه ب')

    # ─── راهبر: بخش ───
    def test_org_admin_section_writes_stay_inside_own_company(self):
        self._as(self.org_admin)
        url = '/api/org/project-sections/'
        self.assertEqual(self.client.post(url, {'project': self.other_project.id, 'code': 'S1', 'name': 'x'}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'{url}{self.section.id}/', {'project': self.other_project.id}, format='json').status_code, 403)
        self.assertEqual(self.client.post(url, {'project': self.project.id, 'code': 'S3', 'name': 'بخش مجاز'}, format='json').status_code, 201)

    # ─── راهبر: انتساب (عضویت بخش خودش دسترسی شرکت می‌دهد) ───
    def test_org_admin_cannot_escalate_into_other_company_via_assignment(self):
        from companies.access import get_user_allowed_companies
        self._as(self.org_admin)
        base = '/api/org/section-assignments/'
        res = self.client.post(base, {'user': self.org_admin.id, 'section': self.other_section.id, 'role': 'manager'}, format='json')
        self.assertEqual(res.status_code, 403)
        res = self.client.post(f'{base}bulk-assign/', {'section_id': self.other_section.id, 'user_ids': [self.org_admin.id], 'role': 'manager'}, format='json')
        self.assertEqual(res.status_code, 403)
        self.assertFalse(SectionAssignment.objects.filter(user=self.org_admin, section=self.other_section).exists())
        self.assertNotIn(self.other_company.id, set(get_user_allowed_companies(self.org_admin).values_list('id', flat=True)))

    def test_org_admin_assigns_only_company_members_and_never_to_foreign_sections(self):
        self._as(self.org_admin)
        base = '/api/org/section-assignments/'
        self.assertEqual(self.client.post(base, {'user': self.outsider.id, 'section': self.section.id, 'role': 'employee'}, format='json').status_code, 403)
        res = self.client.post(f'{base}bulk-assign/', {'section_id': self.section.id, 'user_ids': [self.member.id, self.outsider.id], 'role': 'employee'}, format='json')
        self.assertEqual(res.status_code, 403)
        self.assertFalse(SectionAssignment.objects.filter(section=self.section).exists())
        self.assertEqual(self.client.post(base, {'user': self.member.id, 'section': self.other_section.id, 'role': 'employee'}, format='json').status_code, 403)
        self.assertEqual(self.client.post(base, {'user': self.member.id, 'section': self.section.id, 'role': 'employee'}, format='json').status_code, 201)

    def test_org_admin_cannot_touch_other_company_assignments(self):
        foreign = SectionAssignment.objects.create(user=self.outsider, section=self.other_section, role='employee', is_active=True)
        own = SectionAssignment.objects.create(user=self.member, section=self.section, role='employee', is_active=True)
        self._as(self.org_admin)
        base = '/api/org/section-assignments/'
        # انتساب‌های شرکت دیگر اصلاً در فهرست او نیستند
        self.assertEqual(self.client.delete(f'{base}{foreign.id}/').status_code, 404)
        self.assertEqual(self.client.patch(f'{base}{foreign.id}/', {'role': 'manager'}, format='json').status_code, 404)
        listed = {row['id'] for row in self._rows(self.client.get(base))}
        self.assertIn(own.id, listed)
        self.assertNotIn(foreign.id, listed)
        self.assertEqual(self.client.patch(f'{base}{own.id}/', {'section': self.other_section.id}, format='json').status_code, 403)
        foreign.refresh_from_db()
        self.assertTrue(foreign.is_active)
        self.assertEqual(foreign.role, 'employee')

    def test_org_admin_cannot_clone_assignments_across_companies(self):
        SectionAssignment.objects.create(user=self.outsider, section=self.other_section, role='employee', is_active=True)
        self._as(self.org_admin)
        res = self.client.post(
            f'/api/org/project-sections/{self.section.id}/clone-assignments/',
            {'source_section_id': self.other_section.id}, format='json'
        )
        self.assertEqual(res.status_code, 403)
        self.assertFalse(SectionAssignment.objects.filter(user=self.outsider, section=self.section).exists())

    # ─── درون‌ریزی اکسل شرکت‌محور ───
    def test_project_import_writes_only_inside_active_company(self):
        self._as(self.org_admin)
        url = '/api/org/projects/import-excel/'
        # کد PRJ_B مال شرکت ب است؛ در شرکت الف پروژهٔ تازه ساخته می‌شود و پروژهٔ شرکت ب دست نمی‌خورد
        res = self._import(url, [('PRJ_A', 'نام تازه الف'), ('PRJ_B', 'پروژهٔ تازه در الف'), ('PRJ_G', 'پروژهٔ تازه دیگر')], self.company)
        self.assertEqual(res.status_code, 200, res.content)
        self.project.refresh_from_db()
        self.other_project.refresh_from_db()
        self.assertEqual(self.project.name, 'نام تازه الف')
        self.assertEqual(self.other_project.name, 'پروژه ب')
        self.assertTrue(Project.objects.filter(company=self.company, code='PRJ_B').exists())
        self.assertTrue(Project.objects.filter(company=self.company, code='PRJ_G').exists())
        # شرکتی که کاربر به آن دسترسی ندارد
        self.assertEqual(self._import(url, [('PRJ_X', 'x')], self.other_company).status_code, 403)

    def test_section_import_finds_parent_project_only_in_active_company(self):
        self._as(self.org_admin)
        res = self._import('/api/org/project-sections/import-excel/', [
            ('PRJ_A', '', 'SEC_NEW', 'بخش تازه'),
            ('PRJ_B', '', 'SEC_LEAK', 'نشت به شرکت ب'),
            ('PRJ_G', '', 'SEC_LEAK2', 'نشت به سراسری'),
        ], self.company)
        self.assertEqual(res.status_code, 400)  # دو ردیف خطا دارند؛ ردیف معتبر ثبت می‌شود
        self.assertTrue(ProjectSection.objects.filter(project=self.project, code='SEC_NEW').exists())
        self.assertFalse(ProjectSection.objects.filter(code__in=['SEC_LEAK', 'SEC_LEAK2']).exists())

    def test_assignment_import_only_for_members_and_sections_of_active_company(self):
        self._as(self.org_admin)
        res = self._import('/api/org/section-assignments/import-excel/', [
            (self.member.username, '', 'PRJ_A', 'SEC_A', 'employee', '', 'بله'),
            (self.outsider.username, '', 'PRJ_A', 'SEC_A', 'employee', '', 'بله'),
            (self.org_admin.username, '', 'PRJ_B', 'SEC_B', 'manager', '', 'بله'),
        ], self.company)
        self.assertEqual(res.status_code, 400)
        self.assertTrue(SectionAssignment.objects.filter(user=self.member, section=self.section).exists())
        self.assertFalse(SectionAssignment.objects.filter(user=self.outsider).exists())
        self.assertFalse(SectionAssignment.objects.filter(section=self.other_section).exists())

    def test_counterparty_import_never_touches_global_or_other_company_records(self):
        other_cp = Counterparty.objects.create(company=self.other_company, name='طرف‌حساب ب', counterparty_type='other', phone='0211')
        self._as(self.accountant)
        url = '/api/personnel/counterparties/import-excel/'
        res = self._import(url, [
            ('طرف‌حساب سراسری', 'other', '', '09129999999', '', '', '', '', 'بله'),
            ('طرف‌حساب ب', 'other', '', '09128888888', '', '', '', '', 'بله'),
            ('طرف‌حساب الف', 'other', '', '09127777777', '', '', '', '', 'بله'),
        ], self.company)
        self.assertEqual(res.status_code, 200, res.content)
        self.global_counterparty.refresh_from_db()
        other_cp.refresh_from_db()
        self.counterparty.refresh_from_db()
        self.assertNotEqual(self.global_counterparty.phone, '09129999999')
        self.assertEqual(other_cp.phone, '0211')
        self.assertEqual(self.counterparty.phone, '09127777777')
        self.assertEqual(Counterparty.objects.filter(company=self.company).count(), 3)
        self.assertEqual(self._import(url, [('x', 'other')], self.other_company).status_code, 403)

    # ─── خروجی اکسل ───
    def test_excel_export_contains_only_permitted_companies(self):
        import io
        import openpyxl
        self._as(self.org_admin)
        res = self.client.get('/api/org/projects/export-excel/')
        self.assertEqual(res.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        codes = {ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)}
        self.assertEqual(codes, {'PRJ_A'})

    # ─── نویسندهٔ قدیمی (مدیر) هم محدود به شرکت خودش است ───
    def test_legacy_manager_is_company_scoped_too(self):
        self._as(self.manager)
        base = '/api/org/section-assignments/'
        self.assertEqual(self.client.post(base, {'user': self.manager.id, 'section': self.other_section.id, 'role': 'manager'}, format='json').status_code, 403)
        res = self.client.post(f'{base}bulk-assign/', {'section_id': self.other_section.id, 'user_ids': [self.manager.id], 'role': 'manager'}, format='json')
        self.assertEqual(res.status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.other_project.id}/', {'name': 'تغییر'}, format='json').status_code, 404)
        self.assertEqual(self.client.patch(f'/api/personnel/counterparties/{self.global_counterparty.id}/', {'phone': '1'}, format='json').status_code, 403)
        self.assertEqual(self._import('/api/org/projects/import-excel/', [('PRJ_X', 'x')], self.other_company).status_code, 403)

    def test_company_access_grants_stay_inside_permitted_companies(self):
        newcomer = User.objects.create_user(username='scope_newcomer', password='password123', email='n@test.com')
        self._as(self.manager)
        url = '/api/personnel/user-company-access/'
        # بدون «نوشتن» دسترسی کاربران هیچ اعطایی ممکن نیست (طرح مجوزهای بخش‌به‌بخش)
        res = self.client.post(url, {'user': newcomer.id, 'company': self.company.id, 'section_levels': {}}, format='json')
        self.assertEqual(res.status_code, 403)
        UserCompanyAccess.objects.filter(user=self.manager, company=self.company).update(section_levels={'user_access': 'write'})
        # خود را به شرکت دیگر راه دادن
        res = self.client.post(url, {'user': self.manager.id, 'company': self.other_company.id, 'section_levels': {}}, format='json')
        self.assertEqual(res.status_code, 403)
        # اعطای پیوند در شرکت خودش مجاز است
        res = self.client.post(url, {'user': newcomer.id, 'company': self.company.id, 'section_levels': {}}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        listed_companies = {row['company'] for row in self._rows(self.client.get(url))}
        self.assertNotIn(self.other_company.id, listed_companies)
        outsider_access = UserCompanyAccess.objects.get(user=self.outsider, company=self.other_company)
        self.assertEqual(self.client.delete(f'{url}{outsider_access.id}/').status_code, 404)

    def test_user_list_company_param_is_validated_for_user_managers(self):
        admin_like = self._user('scope_user_admin', ['view_sys_users'], self.company)
        self._as(admin_like)
        self.assertEqual(self.client.get(f'/api/auth/users/?company_id={self.other_company.id}').status_code, 403)
        ids = {row['id'] for row in self._rows(self.client.get('/api/auth/users/'))}
        self.assertNotIn(self.outsider.id, ids)

    # ─── فهرست کاربران راهبر ───
    def test_org_admin_user_directory_is_scoped_and_minimal(self):
        derived_member = User.objects.create_user(username='scope_derived', password='password123', email='d@test.com')
        SectionAssignment.objects.create(user=derived_member, section=self.section, role='employee', is_active=True)
        self._as(self.org_admin)
        res = self.client.get('/api/auth/users/')
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        rows = body.get('results', body) if isinstance(body, dict) else body
        ids = {row['id'] for row in rows}
        self.assertIn(self.member.id, ids)
        self.assertIn(derived_member.id, ids)
        self.assertNotIn(self.outsider.id, ids)
        for row in rows:
            self.assertEqual(set(row.keys()), {'id', 'username', 'first_name', 'last_name', 'is_active'})
        self.assertEqual(self.client.get(f'/api/auth/users/?company_id={self.other_company.id}').status_code, 403)

    # ─── حسابدار: طرف‌حساب ───
    def test_accountant_counterparty_writes_stay_inside_own_company(self):
        self._as(self.accountant)
        url = '/api/personnel/counterparties/'
        self.assertEqual(self.client.patch(f'{url}{self.global_counterparty.id}/', {'phone': '09120000009'}, format='json').status_code, 403)
        self.assertEqual(self.client.delete(f'{url}{self.global_counterparty.id}/').status_code, 403)
        self.assertEqual(self.client.patch(f'{url}{self.counterparty.id}/', {'company': self.other_company.id}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'{url}{self.counterparty.id}/', {'company': None}, format='json').status_code, 403)
        self.assertEqual(self.client.patch(f'{url}{self.counterparty.id}/', {'phone': '09120000010'}, format='json').status_code, 200)
        self.assertTrue(Counterparty.objects.filter(pk=self.global_counterparty.pk).exists())

    def test_counterparty_create_never_lands_in_foreign_company_or_global(self):
        """شرکت غیرمجاز نادیده گرفته می‌شود و رد نمی‌شود (تا رکورد صف آفلاین حذف نشود)؛ بی‌هدر هم سراسری ساخته نمی‌شود"""
        self._as(self.accountant)
        url = '/api/personnel/counterparties/'
        res = self.client.post(url, {'name': 'تزریق به شرکت ب', 'counterparty_type': 'other', 'company': self.other_company.id}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(Counterparty.objects.get(name='تزریق به شرکت ب').company_id, self.company.id)
        # بازپخش صف آفلاین بدون هدر شرکت و بدون شرکت در بدنه
        res = self.client.post(url, {'name': 'ثبت آفلاین', 'counterparty_type': 'other'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(Counterparty.objects.get(name='ثبت آفلاین').company_id, self.company.id)
        res = self.client.post(url, {'name': 'هدر شرکت ب', 'counterparty_type': 'other'}, format='json', HTTP_X_COMPANY_ID=str(self.other_company.id))
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(Counterparty.objects.get(name='هدر شرکت ب').company_id, self.company.id)
        self.assertFalse(Counterparty.objects.filter(company=self.other_company).exists())

    # ─── یافته‌های بازبینی دوم ───
    def test_reactivating_a_users_only_assignment_is_allowed(self):
        own = SectionAssignment.objects.create(user=self.member, section=self.section, role='employee', is_active=True)
        UserCompanyAccess.objects.filter(user=self.member).delete()  # تنها عضویت او همین انتساب است
        self._as(self.org_admin)
        base = '/api/org/section-assignments/'
        self.assertEqual(self.client.delete(f'{base}{own.id}/').status_code, 204)
        res = self.client.patch(f'{base}{own.id}/', {'is_active': True}, format='json')
        self.assertEqual(res.status_code, 200, res.content)
        self.client.delete(f'{base}{own.id}/')
        res = self.client.post(base, {'user': self.member.id, 'section': self.section.id, 'role': 'employee'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        own.refresh_from_db()
        self.assertTrue(own.is_active)

    def test_sections_and_projects_cannot_move_between_companies(self):
        UserCompanyAccess.objects.create(user=self.org_admin, company=self.other_company)
        grant_membership(self.org_admin, self.other_company)
        self._as(self.org_admin)
        res = self.client.patch(f'/api/org/project-sections/{self.section.id}/', {'project': self.other_project.id}, format='json')
        self.assertEqual(res.status_code, 403)
        res = self.client.patch(f'/api/org/projects/{self.project.id}/', {'company': self.other_company.id}, format='json')
        self.assertEqual(res.status_code, 403)
        self.section.refresh_from_db()
        self.project.refresh_from_db()
        self.assertEqual(self.section.project_id, self.project.id)
        self.assertEqual(self.project.company_id, self.company.id)

    def test_org_admin_cannot_assign_themselves_but_manager_still_can(self):
        base = '/api/org/section-assignments/'
        self._as(self.org_admin)
        self.assertEqual(self.client.post(base, {'user': self.org_admin.id, 'section': self.section.id, 'role': 'manager'}, format='json').status_code, 403)
        res = self.client.post(f'{base}bulk-assign/', {'section_id': self.section.id, 'user_ids': [self.member.id, self.org_admin.id], 'role': 'manager'}, format='json')
        self.assertEqual(res.status_code, 403)
        self.assertFalse(SectionAssignment.objects.filter(user=self.org_admin).exists())
        self._as(self.manager)
        self.assertEqual(self.client.post(base, {'user': self.manager.id, 'section': self.section.id, 'role': 'manager'}, format='json').status_code, 201)

    def test_user_edit_cannot_grant_foreign_company_access(self):
        editor = self._user('scope_user_editor', ['perm_usr_edit'], self.company)
        both = self._user('scope_both', [], self.company)
        UserCompanyAccess.objects.create(user=both, company=self.other_company)
        grant_membership(both, self.other_company)
        self._as(editor)
        res = self.client.patch(f'/api/auth/users/{editor.id}/', {'company_ids': [self.company.id, self.other_company.id]}, format='json')
        self.assertEqual(res.status_code, 403)
        self.assertFalse(UserCompanyAccess.objects.filter(user=editor, company=self.other_company).exists())
        # کاربر مشترک: دسترسی او به شرکتِ بیرون از قلمرو ویرایشگر پاک نمی‌شود
        res = self.client.patch(f'/api/auth/users/{both.id}/', {'company_ids': [self.company.id]}, format='json')
        self.assertEqual(res.status_code, 200, res.content)
        self.assertTrue(UserCompanyAccess.objects.filter(user=both, company=self.other_company).exists())

    def test_org_admin_directory_search_ignores_personal_fields(self):
        self.member.national_code = '0012345678'
        self.member.save(update_fields=['national_code'])
        self._as(self.org_admin)
        ids = {row['id'] for row in self._rows(self.client.get('/api/auth/users/?search=00123'))}
        self.assertNotIn(self.member.id, ids)
        ids = {row['id'] for row in self._rows(self.client.get(f'/api/auth/users/?search={self.member.username}'))}
        self.assertIn(self.member.id, ids)

    def test_superuser_import_uses_active_company(self):
        su = User.objects.create_superuser(username='scope_su', password='password123', email='su2@test.com')
        self._as(su)
        res = self._import('/api/org/projects/import-excel/', [('PRJ_SU', 'پروژهٔ سوپریوزر')], self.company)
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(Project.objects.get(code='PRJ_SU').company_id, self.company.id)

    # ─── نویسندهٔ قدیمی: رفتار عادی در شرکت خودش دست نخورده است ───
    def test_legacy_structure_manager_still_writes_in_own_company(self):
        self._as(self.manager)
        self.assertEqual(self.client.patch(f'/api/personnel/counterparties/{self.counterparty.id}/', {'phone': '09120000011'}, format='json').status_code, 200)
        self.assertEqual(self.client.patch(f'/api/org/projects/{self.project.id}/', {'name': 'ویرایش مدیر'}, format='json').status_code, 200)
        res = self.client.post('/api/org/section-assignments/', {'user': self.member.id, 'section': self.section.id, 'role': 'employee'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
