# -*- coding: utf-8 -*-
"""
ViewSets for counterparties, expense invoices, petty cash accounts/transactions, and workflow audit logs.
"""

import io
import openpyxl
from decimal import Decimal
from rest_framework import viewsets, permissions, status, parsers
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied, ValidationError
from django.http import HttpResponse
from django.db.models import Q, Sum
from django.utils import timezone
from django.apps import apps
from django.contrib.auth import get_user_model
from django.shortcuts import get_object_or_404
from common.date_utils import normalize_digits
from accounts.views import OptionalPageNumberPagination
from companies.access import (
    get_user_allowed_companies,
    get_request_company_id,
    validate_user_company_access,
    ensure_scoped_write,
)
from ..models import (
    Counterparty,
    PartyBankAccount,
    ExpenseInvoice,
    PettyCashAccount,
    PettyCashTransaction,
    WorkflowAuditLog,
    ProjectSection,
    PersonnelProfile,
    PersonnelChangeRequest,
)
from ..serializers import (
    CounterpartySerializer,
    PartyBankAccountSerializer,
    ExpenseInvoiceSerializer,
    PettyCashAccountSerializer,
    PettyCashTransactionSerializer,
    WorkflowAuditLogSerializer,
)
from ..permissions import (
    ensure_counterparty_manager,
    IsCounterpartyManagerOrReadOnly,
    requires_company_scope,
)
from .scope_helpers import (
    broadcast_org_structure_updated,
    readable_company_ids,
    resolve_import_company,
)
from companies.org_roles import user_has_section_role
from accounts.scoped_permissions import can, section_q, sections_where
from ..excel_utils import write_standard_sheet
from ..counterparty_services import (
    create_bank_account,
    create_counterparty,
    delete_bank_account,
    remove_counterparty,
    update_bank_account,
    update_counterparty,
)

# مجوزهای بخشی ت۱۳ طرح ۲.۲: «چه کاری» از مجوز و «کجا» از انتساب (ق۶). فهرست‌ها با sections_where/section_q
# و کار روی هر ردیف با can(user, code, section) سنجیده می‌شود، نه has_perm (ت۸).
INVOICE_PERM = 'perm_invoice_record'
PETTY_CASH_MANAGE_PERM = 'perm_petty_cash_manage'

# فاکتور روی موتور گردش کار هسته (ت۱۷ طرح ۲.۳): نشانی و ورودی اکشن‌ها همان است، وضعیت در خود فاکتور (ت۲)
from django.db import transaction
from django.http import Http404
from workflow import engine as wf_engine
from workflow.errors import WorkflowError
from ..workflows import (INVOICE_ENTITY, INVOICE_WORKFLOW, can_view_invoice, can_view_petty_statement,
                         can_view_petty_tx)
from ..workflow_engine import notify_document_changed
from ..fiscal import ensure_period_open
from ..invoice_services import (duplicate_queryset, ensure_invoice_number, invoice_duplicates,
                                void_invoice_number)
from ..invoice_payments import (PAYMENTS_EXIST_CODE, PAYMENTS_EXIST_MESSAGE, RECEIPT_FIELD, add_payment,
                                attach_receipt, invoice_remaining, read_receipt,
                                void_payment)
from ..models import InvoicePayment
from ..serializers import InvoicePaymentSerializer
from django.db.models import Prefetch

INVOICE_PAYMENTS_PREFETCH = Prefetch(
    'payments',
    queryset=InvoicePayment.objects.select_related('company_bank_account', 'party_bank_account__party',
                                                   'created_by', 'voided_by'))
from ..models import normalize_invoice_number

INVOICE_EDITABLE_STATUSES = ('draft', 'revision_required')


class InvoicePagePagination(OptionalPageNumberPagination):
    """د۶ طرح ۴.۱: فقط با پارامتر page صفحه‌بندی می‌شود (page_size تنها فهرست کامل را عوض نمی‌کند)."""

    def paginate_queryset(self, queryset, request, view=None):
        if 'page' not in request.query_params:
            return None
        return super().paginate_queryset(queryset, request, view)


def _record_submit_reason(invoice, reason):
    """دلیل ارسال فاکتور تکراری در ردیف «ارسال» سابقهٔ گردش (موتور هسته دلیل ارسال نمی‌گیرد)."""
    from workflow.models import WorkflowStep
    wf_instance = wf_engine.instance_of(INVOICE_WORKFLOW, invoice.pk)
    if wf_instance is None:
        return
    step = (WorkflowStep.objects.filter(instance=wf_instance, action=WorkflowStep.SUBMIT)
            .order_by('-at', '-pk').first())
    if step is not None and not step.reason:
        step.reason = f'دلیل ارسال فاکتور تکراری: {reason}'[:2000]
        step.save(update_fields=['reason'])


def _invoice_wf_call(fn, raise_drf=False):
    """اجرای یک فراخوان موتور؛ خطای موتور = پاسخ {code, error} با وضعیت همان خطا (مثلاً 409 task_already_done)."""
    try:
        fn()
    except WorkflowError as exc:
        if raise_drf:
            from rest_framework.exceptions import APIException
            err = APIException(exc.message_fa)
            err.status_code = exc.status
            raise err
        return Response({'code': exc.code, 'error': exc.message_fa, 'detail': exc.message_fa}, status=exc.status)
    return None


def _ensure_doc_period_open(section, value, field):
    """ص۵ طرح ۳.۲: سند تاریخ‌دارِ دورهٔ مالی بستهٔ شرکتِ بخش ثبت، ویرایش، حذف یا در گردش جابه‌جا نمی‌شود (400)."""
    company_id = getattr(getattr(section, 'project', None), 'company_id', None)
    ensure_period_open(company_id, value, field)


def _ensure_invoice_open(invoice):
    _ensure_doc_period_open(invoice.section, invoice.invoice_date_shamsi, 'invoice_date_shamsi')


def _ensure_petty_cash_open(tx):
    # شرکت ردیف از حساب سطح شرکت، وگرنه از بخش (ردیف قدیمی)؛ ص۱۱ طرح ۳.۶
    ensure_period_open(tx.wf_company_id, tx.transaction_date_shamsi, 'transaction_date_shamsi')


def _section_id(value):
    try:
        return int(getattr(value, 'pk', value))
    except (TypeError, ValueError):
        return None


def _ensure_own_petty_cash_section(user, section):
    """
    تنخواه خودِ کاربر فقط در بخشی که انتساب فعال (هر نقشی، در هر سطح) آن را می‌پوشاند یا مدیر تنخواه آن است.
    """
    if user.is_superuser:
        return
    if section is None or not (
        user_has_section_role(user, section) or can(user, PETTY_CASH_MANAGE_PERM, section)
    ):
        raise PermissionDenied("شما در این بخش انتساب فعال ندارید و ثبت تنخواه در آن مجاز نیست.")


def _ensure_petty_cash_manager(user, section, message="شما مجوز مدیریت تنخواه این بخش را ندارید."):
    if not user.is_superuser and not can(user, PETTY_CASH_MANAGE_PERM, section):
        raise PermissionDenied(message)


# ─── تنخواه سطح شرکت (طرح ۳.۶) ─────────────────────────────────────────────────────────────────────
from accounts.scoped_permissions import can_at_company_level, codes_in_company
from workflow.models import WorkflowInstance
from ..models import PettyCashStatement
from ..serializers import PettyCashStatementSerializer
from ..fiscal import to_shamsi_str
from ..expense_heads import PETTY_CASH_FIXED_CATEGORY
from ..workflows import (
    PETTY_CHARGE_ENTITY, PETTY_CHARGE_WORKFLOW, PETTY_MANAGER_PERM, PETTY_REVIEW_PERM,
    PETTY_STATEMENT_ENTITY, PETTY_STATEMENT_LINE_REJECT_STEPS, PETTY_STATEMENT_WORKFLOW, PETTY_COMPANY_VIEW_PERMS,
    is_petty_statement_locked, petty_company_allowed, petty_fund_level, recompute_statement,
)
from .scope_helpers import resolve_write_company
from rest_framework.exceptions import APIException
from ..models import PettyCashDebtSettlement
from ..serializers import PettyCashDebtSettlementSerializer
from ..petty_settlements import (
    SettlementError, create_settlement, complete_charge_offsets, payroll_deduction_available, void_settlement,
)
from ..workflows import (
    petty_allocation_category, petty_charge_exceeds_ceiling, petty_fund_level_breakdown,
)
from ..models import PettyCashReconciliation
from ..serializers import PettyCashReconciliationSerializer
from ..serializers_domains.treasury_audit import petty_bank_full_view
from ..counterparty_services import mask_tail
from ..petty_lifecycle import (
    PettyLifecycleError, close_account, create_reconciliation, deputy_account_q, ensure_account_open,
    ensure_period_unlocked, ensure_statement_number, handover_account, is_petty_deputy,
)


class _PettyApiError(APIException):
    """خطای تنخواه با همان شکل {code, error, detail} پاسخ‌های اکشن‌ها، از داخل perform_*."""

    def __init__(self, code, message, status_code=status.HTTP_400_BAD_REQUEST):
        self.status_code = status_code
        super().__init__({'code': code, 'error': message, 'detail': message})


def _lifecycle(fn, *args, **kwargs):
    """خطای چرخهٔ عمر تنخواه (د۵ تا د۸) با همان شکل {code, error, detail}."""
    try:
        return fn(*args, **kwargs)
    except PettyLifecycleError as exc:
        raise _PettyApiError(exc.code, exc.message_fa, exc.status)


def _ensure_account_open(account):
    """د۶: حساب بسته فقط خواندنی است (400 account_closed)."""
    _lifecycle(ensure_account_open, account)


def _ensure_period_unlocked(account_id, day, tx_type='expense'):
    """د۵: هزینه با تاریخ تا آخرین تطبیق حساب ثبت یا ویرایش نمی‌شود (400 period_locked؛ صف آفلاین «ردشده» می‌ماند)."""
    _lifecycle(ensure_period_unlocked, account_id, day, tx_type)


def _shamsi_param(request, name):
    raw = request.query_params.get(name)
    return to_shamsi_str(raw) if raw else None


def _petty_offset_subquery():
    """جمع کسر بدهی (در انتظار یا قطعی) هر شارژ برای annotate روی فهرست ردیف‌ها (یک کوئری، بی N+1)."""
    from django.db.models import DecimalField, OuterRef, Subquery
    from django.db.models.functions import Coalesce as _Coalesce
    money = DecimalField(max_digits=15, decimal_places=0)
    agg = (PettyCashDebtSettlement.objects
           .filter(charge=OuterRef('pk'), method='charge_offset', status__in=['pending', 'done'])
           .order_by().values('charge').annotate(s=Sum('amount')).values('s'))
    return _Coalesce(Subquery(agg, output_field=money), Decimal('0'), output_field=money)


def _manager_approved_current_round(tx):
    """
    آیا مدیر در دور فعلی گردش این شارژ (از آخرین ارسال) تأیید کرده است؟ مرز دور زمان ارسال است، نه کار بررسی
    حسابدار، چون بررسی برای حسابدارِ ثبت‌کننده‌ای که تنها حسابدار است رد می‌شود و کاری باز نمی‌کند.
    """
    from workflow.models import WorkflowTask
    inst = wf_engine.instance_of(PETTY_CHARGE_WORKFLOW, tx.pk)
    if inst is None:
        return False
    return WorkflowTask.objects.filter(
        instance=inst, track=WorkflowTask.TRACK_STEP, step_key='manager_ok', outcome=WorkflowTask.ACTED,
        opened_at__gte=inst.submitted_at).exists()


def _petty_account_has_activity(account):
    """د۲: اولین گردش حساب = هر ردیف یا صورت‌حساب."""
    return account.transactions.exists() or account.statements.exists()


def _log_opening_balance(user, account, old_amount, old_date, reason='', had_activity=False):
    first = not old_amount and not old_date
    _log_petty_audit(user, 'pettycashaccount', account.pk,
                     'petty_opening_balance_set' if first else 'petty_opening_balance_changed',
                     str(int(old_amount or 0)), str(int(account.opening_balance or 0)), reason=reason,
                     company_id=account.company_id, custodian_id=account.custodian_id,
                     before={'opening_balance': str(int(old_amount or 0)), 'opening_balance_date': old_date},
                     after={'opening_balance': str(int(account.opening_balance or 0)),
                            'opening_balance_date': account.opening_balance_date},
                     had_activity=had_activity)

# ص۱۰: اسناد قدیمی در این وضعیت‌ها فقط با اکشن قدیمی حسابدار بسته می‌شوند؛ سند تازه هرگز به آن‌ها نمی‌رود
LEGACY_PETTY_PENDING = ('pending_supervisor', 'pending_accountant')
PETTY_RETIRED_MSG = 'این مسیر قدیمی تنخواه کنار رفته است؛ از درخواست شارژ یا صورت‌حساب تنخواه استفاده کنید.'


def _petty_retired():
    return Response({'code': 'petty_cash_path_retired', 'error': PETTY_RETIRED_MSG, 'detail': PETTY_RETIRED_MSG},
                    status=status.HTTP_410_GONE)


def _petty_viewer_company_ids(user):
    """شرکت‌های مجاز کاربر که در آن‌ها تنخواه دیگران را در سطح شرکت می‌بیند (مدیریت تنخواه یا مجوز مرحله‌ها)."""
    allowed = get_user_allowed_companies(user).values_list('id', flat=True)
    return [cid for cid in allowed
            if any(can_at_company_level(user, code, cid) for code in PETTY_COMPANY_VIEW_PERMS)]


def _petty_tx_company_q(company_ids):
    return Q(account__company_id__in=company_ids) | Q(account__isnull=True, section__project__company_id__in=company_ids)


def scope_invoices_for_user(user, qs, company_id=None):
    """
    قلمرو فهرست فاکتور (و گزارش‌های ۴.۷): شرکت داده‌شده (اعتبارسنجی‌شده) و فقط بخش‌هایی که کاربر در آن‌ها
    «ثبت و دیدن فاکتور» دارد (ت۱۳)؛ خروجی اکسل هم همین queryset است.
    """
    if company_id:
        qs = qs.filter(section__project__company_id=company_id)
    return qs.filter(section_q(user, INVOICE_PERM))


def scope_petty_transactions_for_user(user, qs, company_id=None):
    """
    قلمرو اسناد تنخواه (فهرست و گزارش‌های ۴.۷): همیشه فقط شرکت‌های مجاز کاربر. در آن‌ها: اسناد خودش، همهٔ اسناد
    شرکتی که در آن انتساب سطح شرکت با مدیریت تنخواه یا مجوز مرحله‌ها دارد (ص۷ طرح ۳.۶)، و اسناد قدیمی بخش‌هایی که
    مدیر تنخواه آن‌هاست. پارامتر custodian_id قلمرو را باز نمی‌کند.
    """
    if company_id:
        qs = qs.filter(_petty_tx_company_q([company_id]))
    if not user.is_superuser:
        allowed = list(get_user_allowed_companies(user).values_list('id', flat=True))
        qs = qs.filter(_petty_tx_company_q(allowed)).filter(
            Q(custodian=user)
            | _petty_tx_company_q(_petty_viewer_company_ids(user))
            | Q(section_id__in=sections_where(user, PETTY_CASH_MANAGE_PERM).values('id'))
            | deputy_account_q(user, prefix='account__')  # د۸: حساب صاحب کار در بازهٔ جانشینی
        )
    return qs


def scope_petty_accounts_for_user(user, qs, company_id=None):
    """
    قلمرو حساب‌های تنخواه (ص۱ طرح ۳.۶): همیشه فقط شرکت‌های مجاز کاربر؛ در آن‌ها حساب خودش، همهٔ حساب‌های شرکتی که
    در آن انتساب سطح شرکت با مدیریت تنخواه یا مجوز مرحله‌ها دارد، و حساب قدیمی بخش‌هایی که مدیر تنخواه آن‌هاست.
    """
    if company_id:
        qs = qs.filter(company_id=company_id)
    if not user.is_superuser:
        qs = qs.filter(company_id__in=get_user_allowed_companies(user).values('id')).filter(
            Q(custodian=user)
            | Q(company_id__in=_petty_viewer_company_ids(user))
            | Q(section_id__in=sections_where(user, PETTY_CASH_MANAGE_PERM).values('id'))
            | deputy_account_q(user)
        )
    return qs


def _ensure_petty_custodian(user, company_id):
    """ص۱: ثبت هزینه و درخواست تنخواه فقط با «مدیریت تنخواه» در همان شرکت (هر سطح انتساب) و فقط برای حساب خود."""
    if not petty_company_allowed(user, company_id):
        raise PermissionDenied("شما به این شرکت دسترسی ندارید.")
    if not user.is_superuser and PETTY_CASH_MANAGE_PERM not in codes_in_company(user, company_id):
        if not PettyCashAccount.objects.filter(company_id=company_id, custodian=user, is_active=True).exists():
            raise PermissionDenied("شما در این شرکت مجوز ثبت تنخواه ندارید.")


def _own_petty_account(user, company_id, create=False):
    account = PettyCashAccount.objects.filter(company_id=company_id, custodian=user, is_active=True).first()
    if account is None and create:
        # د۶: حساب بسته دوباره باز یا خودکار جایگزین نمی‌شود؛ حساب تازه را مدیریت تنخواه می‌سازد
        if PettyCashAccount.objects.filter(company_id=company_id, custodian=user, closed_at__isnull=False).exists():
            raise _PettyApiError('account_closed', "حساب تنخواه شما در این شرکت بسته شده است؛ "
                                                   "برای حساب تازه با مدیریت تنخواه هماهنگ کنید.")
        from django.db import IntegrityError
        try:
            with transaction.atomic():
                # سقف تعیین‌نشده (۰): هر شارژ تا تعیین سقف به دست مدیر، به تأیید مدیر می‌رسد (ص۳)
                account = PettyCashAccount.objects.create(company_id=company_id, custodian=user, ceiling_amount=0)
        except IntegrityError:
            account = PettyCashAccount.objects.filter(company_id=company_id, custodian=user, is_active=True).first()
    return account


def _charge_in_progress(tx):
    inst = wf_engine.instance_of(PETTY_CHARGE_WORKFLOW, tx.pk)
    return inst is not None and inst.status == WorkflowInstance.IN_PROGRESS


def _ensure_petty_tx_editable(tx):
    if tx.status != 'draft' or tx.statement_id or _charge_in_progress(tx):
        raise PermissionDenied("این سند تنخواه در گردش کار، در صورت‌حساب یا بسته‌شده است و تغییر نمی‌کند.")


def _petty_wf_error(exc):
    return Response({'code': exc.code, 'error': exc.message_fa, 'detail': exc.message_fa}, status=exc.status)


def _log_petty_audit(user, content_type, object_id, action_name, from_status='', to_status='', reason='', **metadata):
    WorkflowAuditLog.objects.create(content_type=content_type, object_id=object_id, actor=user,
                                    from_status=from_status, to_status=to_status, action=action_name,
                                    reason=reason, metadata=metadata)

def _tab_id(request):
    return request.headers.get('X-Client-Tab-Id') or (request.data.get('client_tab_id') if hasattr(request.data, 'get') else None)


def _broadcast_counterparty(request, action_name, cp_id, name, company_id):
    # طرف حساب سراسری دیگر نیست (ص۲ طرح ۳.۱): رویداد همیشه به شرکت خود طرف می‌رود
    broadcast_org_structure_updated('counterparty', action_name, cp_id, name, company_id=company_id,
                                    global_scope=False, client_tab_id=_tab_id(request), sender_id=request.user.id)


class CounterpartyViewSet(viewsets.ModelViewSet):
    """
    مدیریت طرف‌حساب‌های مالی (رانندگان، تعمیرگاه‌ها، جایگاه سوخت و پیمانکاران) روی طرف بیرونی هسته (ص۵ طرح ۳.۱).
    """
    queryset = Counterparty.objects.select_related('party', 'party__company').prefetch_related('party__bank_accounts')
    serializer_class = CounterpartySerializer
    permission_classes = [permissions.IsAuthenticated, IsCounterpartyManagerOrReadOnly]
    pagination_class = OptionalPageNumberPagination

    def get_permissions(self):
        if self.action in ['list', 'retrieve', 'create']:
            return [permissions.IsAuthenticated()]
        return [permissions.IsAuthenticated(), IsCounterpartyManagerOrReadOnly()]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if not user.is_authenticated:
            return qs.none()

        # ایزولاسیون چندمستاجری طرف‌حساب‌ها بر مبنای شرکت فعال؛ هر طرف حساب یک شرکت دارد (ق۶)
        company_param = get_request_company_id(self.request)
        try:
            cid = int(company_param) if company_param else None
        except (ValueError, TypeError):
            cid = None
        if not user.is_superuser:
            allowed_cids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if cid is not None:
                if cid not in allowed_cids:
                    raise PermissionDenied("شما به طرف‌حساب‌های این شرکت دسترسی ندارید.")
                qs = qs.filter(party__company_id=cid)
            else:
                qs = qs.filter(party__company_id__in=allowed_cids)
        elif cid is not None:
            qs = qs.filter(party__company_id=cid)

        c_type = self.request.query_params.get('counterparty_type')
        if c_type:
            qs = qs.filter(counterparty_type=c_type)
        kind = self.request.query_params.get('kind')
        if kind:
            qs = qs.filter(party__kind=kind)
        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() == 'true')
        search = self.request.query_params.get('search')
        if search:
            search_clean = normalize_digits(str(search)).strip()
            qs = qs.filter(Q(party__name__icontains=search) | Q(party__national_id__icontains=search_clean)
                           | Q(party__phone__icontains=search_clean) | Q(party__mobile__icontains=search_clean)
                           | Q(account_code__icontains=search_clean))
        return qs.order_by('party__name', 'id')

    def _resolve_company(self, explicit_company):
        """
        شرکت طرف تازه (ص۵): بدنه (اگر مجاز) ← سرآیند شرکت فعال (اگر مجاز) ← تنها شرکت مجاز ← اولین شرکت مجاز.
        شرکت غیرمجاز نادیده گرفته می‌شود و رد نمی‌شود، چون پاسخ 4xx رکورد صف آفلاین را حذف می‌کند.
        کاربر بی هیچ شرکت مجاز 403 و مدیر کل سامانهٔ بی‌شرکت 400 می‌گیرد. خروجی: (company_id, fallback).
        """
        user = self.request.user
        try:
            req_cid = get_request_company_id(self.request)
            req_cid = int(req_cid) if req_cid else None
        except (ValueError, TypeError):
            req_cid = None
        if not requires_company_scope(user):
            company_id = explicit_company.id if explicit_company else req_cid
            if not company_id:
                raise ValidationError({'company': "شرکت طرف‌حساب را انتخاب کنید."})
            return company_id, False
        allowed_cids = sorted(get_user_allowed_companies(user).values_list('id', flat=True))
        if not allowed_cids:
            raise PermissionDenied("شما به هیچ شرکتی دسترسی ندارید و نمی‌توانید طرف‌حساب ثبت کنید.")
        candidates = [explicit_company.id if explicit_company else None, req_cid]
        if len(allowed_cids) == 1:
            candidates.append(allowed_cids[0])
        company_id = next((c for c in candidates if c is not None and c in allowed_cids), None)
        if company_id is None:
            return allowed_cids[0], True
        return company_id, False

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        company_id, fallback = self._resolve_company(serializer.validated_data.get('company'))
        party_data, role_data, _flat, bank_rows = serializer.split_validated()
        details = {'company_fallback': True} if fallback else None
        cp, existing = create_counterparty(company_id, party_data, role_data, bank_rows, request.user, details=details)
        cp = Counterparty.objects.select_related('party', 'party__company').prefetch_related(
            'party__bank_accounts').get(pk=cp.pk)
        if not existing:
            _broadcast_counterparty(request, 'create', cp.id, cp.name, cp.company_id)
        data = dict(self.get_serializer(cp).data)
        data['existing'] = existing
        return Response(data, status=status.HTTP_200_OK if existing else status.HTTP_201_CREATED)

    def perform_update(self, serializer):
        user = self.request.user
        instance = serializer.instance
        if requires_company_scope(user):
            ensure_scoped_write(user, instance.company_id, "فقط طرف‌حساب‌های شرکت‌های مجاز خود را می‌توانید ویرایش کنید.")
        ensure_counterparty_manager(user, instance.company_id)
        party_data, role_data, flat_bank, _rows = serializer.split_validated()
        new_nid = party_data.get('national_id')
        if new_nid and new_nid != instance.party.national_id:
            from contacts.services import find_party
            other = find_party(instance.company_id, new_nid)
            if other is not None and other.pk != instance.party_id:
                raise ValidationError({'national_id': "طرف دیگری با همین کد یا شناسهٔ ملی در این شرکت هست."})
        update_counterparty(instance, party_data, role_data, flat_bank, user)
        serializer.instance = Counterparty.objects.select_related('party', 'party__company').prefetch_related(
            'party__bank_accounts').get(pk=instance.pk)
        _broadcast_counterparty(self.request, 'update', instance.id, instance.name, instance.company_id)

    def perform_destroy(self, instance):
        if requires_company_scope(self.request.user):
            ensure_scoped_write(self.request.user, instance.company_id, "فقط طرف‌حساب‌های شرکت‌های مجاز خود را می‌توانید حذف کنید.")
        ensure_counterparty_manager(self.request.user, instance.company_id)
        instance_id, instance_name, company_id = instance.id, instance.name, instance.company_id
        # ص۵: طرف حساب دارای فاکتور یا تراکنش تنخواه فقط غیرفعال می‌شود
        result = remove_counterparty(instance, self.request.user)
        _broadcast_counterparty(self.request, 'update' if result == 'deactivated' else 'delete',
                                instance_id, instance_name, company_id)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        from ..org_excel_engine import export_counterparties_excel
        section_id = request.query_params.get('section_id')
        return export_counterparties_excel(section_id, company_ids=readable_company_ids(request.user), user=request.user)

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        from ..org_excel_engine import download_counterparties_template
        return download_counterparties_template()

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request):
        from ..org_excel_engine import import_counterparties_from_excel
        # غیرسوپریوزر فقط طرف‌حساب‌های شرکت فعالِ مجاز خودش را می‌سازد یا به‌روز می‌کند
        company_id = resolve_import_company(request)
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایلی ارسال نشده است'}, status=status.HTTP_400_BAD_REQUEST)
        dry_run = request.data.get('dry_run') in ['true', 'True', True]
        if company_id is None:
            return Response({'error': 'شرکت مقصد درون‌ریزی را انتخاب کنید.'}, status=status.HTTP_400_BAD_REQUEST)
        res = import_counterparties_from_excel(file_obj, dry_run=dry_run, company_id=company_id, user=request.user)
        if not dry_run:
            _broadcast_counterparty(request, 'import', None, 'Bulk Import', company_id)
        return Response(res, status=status.HTTP_200_OK if res['success'] else status.HTTP_400_BAD_REQUEST)

class PartyBankAccountViewSet(viewsets.ModelViewSet):
    """
    حساب‌های بانکی طرف حساب (ص۵ طرح ۳.۱): `/api/personnel/party-bank-accounts/?party=<id>`.
    خواندن برای اعضای شرکت طرف (با پوشش ص۴)؛ نوشتن با حق ویرایش طرف حساب و فقط در شرکت‌های مجاز.
    """
    serializer_class = PartyBankAccountSerializer
    permission_classes = [permissions.IsAuthenticated, IsCounterpartyManagerOrReadOnly]
    pagination_class = OptionalPageNumberPagination

    def get_queryset(self):
        user = self.request.user
        qs = PartyBankAccount.objects.select_related('party').filter(party__counterparty__isnull=False)
        if not user.is_superuser:
            qs = qs.filter(party__company_id__in=get_user_allowed_companies(user).values('id'))
        party_id = self.request.query_params.get('party')
        if party_id:
            try:
                qs = qs.filter(party_id=int(party_id))
            except (ValueError, TypeError):
                return qs.none()
        return qs.order_by('-is_primary', '-created_at')

    def _ensure_party_scope(self, party):
        if requires_company_scope(self.request.user):
            ensure_scoped_write(self.request.user, party.company_id, "فقط حساب بانکی طرف‌حساب‌های شرکت‌های مجاز خود را می‌توانید تغییر دهید.")
        ensure_counterparty_manager(self.request.user, party.company_id)

    def perform_create(self, serializer):
        party = serializer.validated_data.get('party')
        if party is None or not Counterparty.objects.filter(party=party).exists():
            raise ValidationError({'party': "طرف‌حساب را انتخاب کنید."})
        self._ensure_party_scope(party)
        data = {k: v for k, v in serializer.validated_data.items() if k != 'party'}
        serializer.instance = create_bank_account(party, data, self.request.user)

    def perform_update(self, serializer):
        self._ensure_party_scope(serializer.instance.party)
        data = {k: v for k, v in serializer.validated_data.items() if k != 'party'}
        serializer.instance = update_bank_account(serializer.instance, data, self.request.user)

    def perform_destroy(self, instance):
        self._ensure_party_scope(instance.party)
        delete_bank_account(instance, self.request.user)


class ExpenseInvoiceViewSet(viewsets.ModelViewSet):
    """
    مدیریت فاکتورهای هزینه‌ای بخش‌ها
    """
    queryset = (ExpenseInvoice.objects.all()
                .select_related('section', 'section__project', 'counterparty__party', 'created_by', 'paid_by')
                .prefetch_related('lines__expense_head__parent', INVOICE_PAYMENTS_PREFETCH))
    serializer_class = ExpenseInvoiceSerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    # د۶ طرح ۴.۱: صفحه‌بندی فقط با ?page=؛ بی آن همان فهرست کامل قبلی
    pagination_class = InvoicePagePagination

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        qs = scope_invoices_for_user(user, qs, cid)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(section_id=section_id)
        project_id = self.request.query_params.get('project_id')
        if project_id:
            qs = qs.filter(section__project_id=project_id)
        inv_status = self.request.query_params.get('status')
        if inv_status:
            qs = qs.filter(status=inv_status)
        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(Q(invoice_number__icontains=search) | Q(number__icontains=search)
                           | Q(counterparty__party__name__icontains=search) | Q(description__icontains=search))
        return qs.order_by('-invoice_date_shamsi', '-created_at')

    def perform_create(self, serializer):
        user = self.request.user
        section = serializer.validated_data.get('section')
        if section and section.project and section.project.company and not user.is_superuser:
            allowed_ids = set(get_user_allowed_companies(user).values_list('id', flat=True))
            if section.project.company_id not in allowed_ids:
                raise PermissionDenied("شما مجاز به ثبت فاکتور برای این شرکت نیستید.")
        if not user.is_superuser and (section is None or not can(user, INVOICE_PERM, section)):
            raise PermissionDenied("شما مجوز ثبت فاکتور در این بخش را ندارید.")
        _ensure_doc_period_open(section, serializer.validated_data.get('invoice_date_shamsi'), 'invoice_date_shamsi')
        # Guardian G2: default to draft status
        req_status = self.request.data.get('status')
        if not user.is_superuser:
            instance = serializer.save(created_by=user, status='draft')
        else:
            instance = serializer.save(created_by=user, status=req_status or 'draft')
        if req_status and req_status != 'draft':
            # ثبت مستقیم با وضعیت غیر پیش‌نویس (فقط مدیر کل) = ارسال به گردش موتور هسته
            instance.status = 'draft'
            instance.save(update_fields=['status', 'updated_at'])
            _invoice_wf_call(lambda: wf_engine.submit(INVOICE_WORKFLOW, instance, user), raise_drf=True)
            ensure_invoice_number(instance, user)
        notify_document_changed(instance, 'created')

    def perform_update(self, serializer):
        user = self.request.user
        instance = self.get_object()
        new_section = serializer.validated_data.get('section')
        if (new_section is not None and new_section.pk != instance.section_id and not user.is_superuser
                and not can(user, INVOICE_PERM, new_section)):
            raise PermissionDenied("شما مجوز ثبت فاکتور در بخش مقصد را ندارید.")
        new_status = serializer.validated_data.pop('status', None)
        if new_status and new_status != instance.status:
            # ت۲ طرح ۲.۳: وضعیت فاکتور فقط از راه گردش کار (ارسال، تأیید، عودت، رد، پرداخت، لغو) عوض می‌شود
            raise PermissionDenied("وضعیت فاکتور فقط از راه گردش کار تغییر می‌کند.")
        if instance.status not in INVOICE_EDITABLE_STATUSES and not user.is_superuser:
            raise PermissionDenied("فاکتور در جریان گردش کار یا بسته‌شده قابل ویرایش نیست.")
        if (instance.status in ('ready_to_pay', 'partially_paid', 'paid') or instance.paid_amount
                or instance.payments.exists()):
            # در مرحلهٔ پرداخت حتی مدیر کل ویرایش نمی‌کند؛ ذخیرهٔ کامل ردیف قسط همزمان را بازنویسی می‌کرد
            raise PermissionDenied("برای این فاکتور قسط پرداخت ثبت شده است و قابل ویرایش نیست.")
        # سند نه از دورهٔ بسته بیرون می‌رود و نه به آن وارد می‌شود
        _ensure_invoice_open(instance)
        _ensure_doc_period_open(new_section or instance.section,
                                serializer.validated_data.get('invoice_date_shamsi', instance.invoice_date_shamsi),
                                'invoice_date_shamsi')
        notify_document_changed(serializer.save(), 'updated')

    def perform_destroy(self, instance):
        _ensure_invoice_open(instance)
        if instance.status not in INVOICE_EDITABLE_STATUSES and not self.request.user.is_superuser:
            # مثل ویرایش: فاکتور در جریان گردش کار یا بسته‌شده حذف نمی‌شود (ردیف‌هایش CASCADE است)
            raise PermissionDenied("فاکتور در جریان گردش کار یا بسته‌شده قابل حذف نیست.")
        if instance.payments.exists():
            # قسط حذف نمی‌شود (د۵)؛ پس فاکتور قسط‌دار هم حذف نمی‌شود (FK آن PROTECT است)
            raise PermissionDenied("برای این فاکتور قسط پرداخت ثبت شده است و حذف نمی‌شود.")
        # in a transaction, so the event goes out only after the delete commits
        with transaction.atomic():
            notify_document_changed(instance, 'deleted')
            # پیش‌نویس شماره‌دار (پس از لغو): شماره باطل می‌شود، نه آزاد (د۴ طرح ۴.۱)
            void_invoice_number(instance, self.request.user)
            instance.delete()

    def _duplicate_payload(self, matches_qs):
        """{count, hidden_count, matches}: جزئیات فقط برای فاکتورهای دامنهٔ دیدن خود کاربر؛ بقیه فقط شمرده می‌شوند."""
        user = self.request.user
        visible_qs = matches_qs if user.is_superuser else matches_qs.filter(section_q(user, INVOICE_PERM))
        total = matches_qs.count()
        visible_count = visible_qs.count()
        visible = list(visible_qs.select_related('section').order_by('-invoice_date_shamsi', '-created_at')[:20])
        return {
            'count': total,
            'hidden_count': max(0, total - visible_count),
            'matches': [{
                'id': inv.pk, 'number': inv.number, 'invoice_number': inv.invoice_number,
                'invoice_date_shamsi': inv.invoice_date_shamsi, 'amount': inv.amount, 'status': inv.status,
                'status_display': inv.get_status_display(), 'section_name': inv.section.name if inv.section else '',
            } for inv in visible],
        }

    @action(detail=False, methods=['get'], url_path='duplicates')
    def duplicates(self, request):
        """هشدار فاکتور تکراری در فرم (د۳ طرح ۴.۱): ?counterparty=&invoice_number=&exclude="""
        user = request.user
        cid = validate_user_company_access(user, get_request_company_id(request))
        cp_id = str(request.query_params.get('counterparty') or '').strip()
        norm = normalize_invoice_number(request.query_params.get('invoice_number') or '')
        exclude = str(request.query_params.get('exclude') or '').strip()
        if not cp_id.isdigit() or (exclude and not exclude.isdigit()):
            return Response({'code': 'invalid_params', 'error': 'طرف‌حساب یا شناسهٔ فاکتور نامعتبر است.'},
                            status=status.HTTP_400_BAD_REQUEST)
        cp_qs = Counterparty.objects.select_related('party').filter(pk=int(cp_id))
        if cid:
            cp_qs = cp_qs.filter(party__company_id=cid)
        elif not user.is_superuser:
            cp_qs = cp_qs.filter(party__company_id__in=get_user_allowed_companies(user).values_list('id', flat=True))
        counterparty = cp_qs.first()
        if counterparty is None:
            return Response({'code': 'counterparty_not_found', 'error': 'طرف‌حساب در شرکت شما یافت نشد.'},
                            status=status.HTTP_400_BAD_REQUEST)
        if not norm:
            return Response({'count': 0, 'hidden_count': 0, 'matches': []})
        matches = duplicate_queryset(counterparty.company_id, counterparty.pk, norm, int(exclude) if exclude else None)
        return Response(self._duplicate_payload(matches))

    def _workflow_invoice(self, read_only=False):
        """
        فاکتور برای اقدام گردش کار: از فهرست بخش‌های «ثبت و دیدن فاکتور» کاربر، یا از راه کار باز همین فاکتور
        (ت۱۱: دریافت‌کننده یا جانشینی که مجوز فهرست آن بخش را ندارد). نادیدنی = 404.
        read_only (فقط دیدن جزئیات): شرکت‌کنندهٔ گردش و مدیر/حسابدار دامنه هم (قلاب can_view، رسید پرداخت).
        """
        pk = self.kwargs.get(self.lookup_url_kwarg or self.lookup_field)
        obj = self.filter_queryset(self.get_queryset()).filter(pk=pk).first()
        if obj is None and (wf_engine.task_grants_view(self.request.user, INVOICE_ENTITY, pk)
                            or (read_only and can_view_invoice(self.request.user, pk))):
            obj = ExpenseInvoice.objects.select_related('section', 'section__project', 'counterparty__party').filter(pk=pk).first()
            if obj is not None and not self._invoice_company_visible(obj):
                obj = None
        if obj is None:
            raise Http404
        return obj

    def _invoice_company_visible(self, obj):
        """راه کار باز فقط در شرکت فعال درخواست و با عضویت زندهٔ کاربر در شرکت فاکتور (ردیف کهنهٔ دریافت‌کننده نه)."""
        company_id = obj.section.project.company_id
        active = get_request_company_id(self.request)
        if active and str(active) != str(company_id):
            return False
        try:
            validate_user_company_access(self.request.user, company_id)
        except PermissionDenied:
            return False
        return True

    def retrieve(self, request, *args, **kwargs):
        return Response(self.get_serializer(self._workflow_invoice(read_only=True)).data)

    def _act_on_open_task(self, request, instance, fn):
        task = wf_engine.open_task(INVOICE_WORKFLOW, instance.pk)
        if task is None:
            return Response({'code': 'task_already_done', 'error': 'این فاکتور کار بازی در گردش ندارد.'},
                            status=status.HTTP_409_CONFLICT)
        revision = request.data.get('revision')
        revision = int(revision) if str(revision or '').isdigit() else None
        return _invoice_wf_call(lambda: fn(task, revision))

    @action(detail=True, methods=['post'], url_path='submit')
    def submit_invoice(self, request, pk=None):
        instance = self.get_object()
        _ensure_invoice_open(instance)
        if instance.status not in INVOICE_EDITABLE_STATUSES:
            return Response({'error': 'فقط فاکتورهای پیش‌نویس یا نیازمند بازنگری قابل ارسال هستند.'}, status=status.HTTP_400_BAD_REQUEST)
        # ثبت‌کننده باید در بخش فاکتور «ثبت و دیدن فاکتور» داشته باشد (get_object همین را سنجیده است)
        if not request.user.is_superuser and not can(request.user, INVOICE_PERM, instance.section):
            raise PermissionDenied("شما مجوز ثبت فاکتور در این بخش را ندارید.")
        # د۳: فاکتور تکراری فقط با دلیل ارسال می‌شود؛ ثبت و ویرایش هرگز به‌خاطر تکراری رد نمی‌شود
        matches = invoice_duplicates(instance)
        reason = str(request.data.get('duplicate_reason') or '').strip()
        is_duplicate = matches.exists()
        if is_duplicate and not reason:
            return Response({'code': 'duplicate_reason_required',
                             'error': 'این شمارهٔ فاکتور برای همین طرف‌حساب قبلاً ثبت شده است؛ دلیل ارسال را بنویسید.',
                             **self._duplicate_payload(matches)}, status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            if is_duplicate:
                instance.possible_duplicate = True
                instance.duplicate_reason = reason[:2000]
                instance.save(update_fields=['possible_duplicate', 'duplicate_reason', 'updated_at'])
            err = _invoice_wf_call(lambda: wf_engine.submit(INVOICE_WORKFLOW, instance, request.user))
            if err is not None:
                transaction.set_rollback(True)
            else:
                if is_duplicate:
                    _record_submit_reason(instance, instance.duplicate_reason)
                # د۴: شمارهٔ داخلی در اولین ارسال و در همین تراکنش؛ ارسال دوباره همان شماره را نگه می‌دارد
                ensure_invoice_number(instance, request.user)
        if err is not None:
            instance.refresh_from_db()
            return err
        return Response({'message': 'فاکتور با موفقیت به گردش کار ارسال شد.', 'status': instance.status,
                         'number': instance.number})

    @action(detail=True, methods=['post'], url_path='approve')
    def approve_invoice(self, request, pk=None):
        instance = self._workflow_invoice()
        _ensure_invoice_open(instance)
        notes = (request.data.get('notes') or '').strip()
        err = self._act_on_open_task(request, instance, lambda task, rev: wf_engine.act(
            request.user, task, reason=notes, revision=rev, obj=instance))
        if err is not None:
            return err
        return Response({'message': 'تایید فاکتور با موفقیت ثبت شد.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='reject')
    def reject_invoice(self, request, pk=None):
        instance = self._workflow_invoice()
        _ensure_invoice_open(instance)
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'درج دلیل رد فاکتور الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        err = self._act_on_open_task(request, instance, lambda task, rev: wf_engine.reject(
            request.user, task, reason, obj=instance, revision=rev))
        if err is not None:
            return err
        instance.rejection_reason = reason
        instance.save(update_fields=['rejection_reason', 'updated_at'])
        return Response({'message': 'فاکتور هزینه رد گردید.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='request-revision')
    def request_revision_invoice(self, request, pk=None):
        instance = self._workflow_invoice()
        _ensure_invoice_open(instance)
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'درج دلیل بازنگری الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if instance.paid_amount > 0:
            return self._payments_exist_response()
        err = self._act_on_open_task(request, instance, lambda task, rev: wf_engine.return_for_fix(
            request.user, task, reason, obj=instance, revision=rev))
        if err is not None:
            return err
        instance.rejection_reason = reason
        instance.save(update_fields=['rejection_reason', 'updated_at'])
        return Response({'message': 'فاکتور با ثبت اشکال جهت بازنگری ارجاع شد.', 'status': instance.status})

    @staticmethod
    def _payments_exist_response():
        return Response({'code': PAYMENTS_EXIST_CODE, 'error': PAYMENTS_EXIST_MESSAGE, 'detail': PAYMENTS_EXIST_MESSAGE},
                        status=status.HTTP_409_CONFLICT)

    @staticmethod
    def _revision(request):
        revision = request.data.get('revision')
        return int(revision) if str(revision or '').isdigit() else None

    def _payment_result(self, request, payment, invoice, message, code=status.HTTP_200_OK):
        invoice.refresh_from_db()
        return Response({
            'message': message,
            'status': invoice.status,
            'status_display': invoice.get_status_display(),
            'amount': str(invoice.amount),
            'paid_amount': str(invoice.paid_amount),
            'remaining': str(invoice_remaining(invoice)),
            'payment': (InvoicePaymentSerializer(payment, context=self.get_serializer_context()).data
                        if payment is not None else None),
        }, status=code)

    def _payment_call(self, fn):
        """خطای موتور/قسط = {code, error, detail} با وضعیت همان خطا؛ خطای فیلد = 400 معمول DRF."""
        try:
            return fn(), None
        except WorkflowError as exc:
            return None, Response({'code': exc.code, 'error': exc.message_fa, 'detail': exc.message_fa},
                                  status=exc.status)

    @action(detail=True, methods=['post'], url_path='pay')
    def pay_invoice(self, request, pk=None):
        """سازگاری تا ۴.۱-ج: پرداخت کل مانده در یک قسط با همان فیلدهای قسط (tracking_code یا payment_ref، notes)."""
        instance = self._workflow_invoice()
        _ensure_invoice_open(instance)
        result, err = self._payment_call(lambda: add_payment(
            request.user, instance.pk, request.data, revision=self._revision(request), legacy=True,
            receipt=request.FILES.get(RECEIPT_FIELD)))
        if err is not None:
            return err
        payment, invoice = result
        return self._payment_result(request, payment, invoice, 'تسویه فاکتور با موفقیت در خزانه‌داری ثبت گردید.')

    @action(detail=True, methods=['get', 'post'], url_path='payments')
    def payments(self, request, pk=None):
        """GET فهرست قسط‌ها (باطل‌شده‌ها هم)؛ POST ثبت یک قسط (د۵ طرح ۴.۱)."""
        instance = self._workflow_invoice()
        if request.method == 'GET':
            rows = (instance.payments.select_related('company_bank_account', 'party_bank_account__party',
                                                     'created_by', 'voided_by').order_by('created_at', 'pk'))
            return Response({
                'amount': str(instance.amount), 'paid_amount': str(instance.paid_amount),
                'remaining': str(invoice_remaining(instance)), 'status': instance.status,
                'results': InvoicePaymentSerializer(rows, many=True, context=self.get_serializer_context()).data,
            })
        _ensure_invoice_open(instance)
        result, err = self._payment_call(lambda: add_payment(
            request.user, instance.pk, request.data, revision=self._revision(request),
            receipt=request.FILES.get(RECEIPT_FIELD)))
        if err is not None:
            return err
        payment, invoice = result
        message = ('قسط ثبت شد و فاکتور کامل پرداخت شد.' if invoice.status == 'paid'
                   else 'قسط پرداخت ثبت شد.')
        return self._payment_result(request, payment, invoice, message, code=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path=r'payments/(?P<pid>\d+)/void')
    def void_payment(self, request, pk=None, pid=None):
        instance = self._workflow_invoice()
        _ensure_invoice_open(instance)
        result, err = self._payment_call(lambda: void_payment(
            request.user, instance.pk, int(pid), request.data.get('reason')))
        if err is not None:
            return err
        payment, invoice = result
        return self._payment_result(request, payment, invoice, 'قسط باطل شد و مانده برگشت.')

    @action(detail=True, methods=['post'], url_path='cancel')
    def cancel_invoice(self, request, pk=None):
        """لغو گردش به دست ثبت‌کننده، پیش از اقدام نفر دیگر (ت۸)؛ فاکتور به پیش‌نویس برمی‌گردد."""
        instance = self.get_object()
        _ensure_invoice_open(instance)
        if instance.paid_amount > 0:
            return self._payments_exist_response()
        wf_instance = wf_engine.instance_of(INVOICE_WORKFLOW, instance.pk)
        if wf_instance is None:
            return Response({'code': 'cancel_not_allowed', 'error': 'این فاکتور در گردش کار نیست.'},
                            status=status.HTTP_409_CONFLICT)
        err = _invoice_wf_call(lambda: wf_engine.cancel(request.user, wf_instance,
                                                        reason=(request.data.get('reason') or '').strip(),
                                                        obj=instance))
        if err is not None:
            return err
        return Response({'message': 'گردش فاکتور لغو شد و فاکتور به پیش‌نویس برگشت.', 'status': instance.status})

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی اکسل استاندارد دو ردیفهٔ سربرگ (ق۱۰، د۸ طرح ۴.۱): برگهٔ فاکتورها و برگهٔ «ردیف‌ها». دادهٔ بانکی نمی‌آید.
        همان queryset دامنه‌دار فهرست؛ ردیف‌ها از prefetch (بی N+1).
        """
        import openpyxl

        qs = self.get_queryset()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "فاکتورهای هزینه"
        lines_ws = wb.create_sheet("ردیف‌ها")

        invoice_cols = [
            ('ردیف', 'row_num', 'num'),
            ('شمارهٔ داخلی', 'number', 'num'),
            ('شمارهٔ فاکتور فروشنده', 'invoice_number', 'num'),
            ('تاریخ فاکتور (شمسی)', 'invoice_date_shamsi', 'num'),
            ('پروژه', 'project_name', 'text'),
            ('بخش', 'section_name', 'text'),
            ('طرف‌حساب', 'counterparty_name', 'right'),
            ('سرفصل‌های هزینه', 'expense_heads', 'text'),
            ('جمع ردیف‌ها (ریال)', 'subtotal', 'money'),
            ('تخفیف (ریال)', 'discount_amount', 'money'),
            ('مالیات بر ارزش افزوده (ریال)', 'vat_amount', 'money'),
            ('مبلغ قابل پرداخت (ریال)', 'amount', 'money'),
            ('پرداخت‌شده (ریال)', 'paid_amount', 'money'),
            ('مانده (ریال)', 'remaining', 'money'),
            ('وضعیت', 'status_display', 'text'),
            ('احتمال تکراری', 'possible_duplicate', 'text'),
            ('ثبت‌کننده', 'created_by_name', 'text'),
            ('شرح هزینه', 'description', 'right'),
        ]
        line_cols = [
            ('ردیف', 'row_num', 'num'),
            ('شمارهٔ داخلی', 'number', 'num'),
            ('شمارهٔ فاکتور فروشنده', 'invoice_number', 'num'),
            ('شمارهٔ ردیف', 'position', 'num'),
            ('شرح ردیف', 'description', 'right'),
            ('سرفصل هزینه', 'expense_head', 'text'),
            ('تعداد / مقدار', 'quantity', 'num'),
            ('واحد', 'unit', 'text'),
            ('فی (ریال)', 'unit_price', 'money'),
            ('مبلغ ردیف (ریال)', 'line_amount', 'money'),
        ]

        invoice_rows = []
        line_rows = []
        for idx, inv in enumerate(qs, start=1):
            created_name = (f"{inv.created_by.first_name} {inv.created_by.last_name}".strip() or inv.created_by.username
                            if inv.created_by else 'سیستم')
            lines = list(inv.lines.all())
            invoice_rows.append([
                idx,
                inv.number or '',
                inv.invoice_number,
                inv.invoice_date_shamsi,
                inv.section.project.name if inv.section and inv.section.project else '',
                inv.section.name if inv.section else '',
                inv.counterparty.name if inv.counterparty else '',
                '، '.join(dict.fromkeys(line.expense_head.title for line in lines if line.expense_head)),
                inv.subtotal,
                inv.discount_amount,
                inv.vat_amount,
                inv.amount,
                inv.paid_amount,
                invoice_remaining(inv),
                inv.get_status_display(),
                'تکراری' if inv.possible_duplicate else '',
                created_name,
                inv.description or '',
            ])
            for line in lines:
                line_rows.append([
                    len(line_rows) + 1,
                    inv.number or '',
                    inv.invoice_number,
                    line.position,
                    line.description or '',
                    line.expense_head.title if line.expense_head else '',
                    line.quantity.normalize() if line.quantity is not None else '',
                    line.unit or '',
                    line.unit_price,
                    line.line_amount,
                ])

        _write_standard_sheet(ws, invoice_cols, invoice_rows)
        _write_standard_sheet(lines_ws, line_cols, line_rows)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"expense_invoices_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# سربرگ دو ردیفهٔ مشترک به personnel/excel_utils.py رفت (د۵ طرح ۴.۷)
_write_standard_sheet = write_standard_sheet

# ==============================================================================
# 0.2 ویوست‌های مدیریت تنخواه‌گردان (Petty Cash Management ViewSets)
# ==============================================================================

class PettyCashAccountViewSet(viewsets.ModelViewSet):
    """
    مدیریت و تعریف حساب‌ها / سقف‌های تنخواه‌گردان کارمندان در بخش‌های پروژه
    """
    queryset = PettyCashAccount.objects.all().select_related('section', 'section__project', 'custodian')
    serializer_class = PettyCashAccountSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = PettyCashAccount.annotate_debt(super().get_queryset().select_related('company'))
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        qs = scope_petty_accounts_for_user(user, qs, cid)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(section_id=section_id)
        custodian_id = self.request.query_params.get('custodian_id')
        if custodian_id:
            qs = qs.filter(custodian_id=custodian_id)
        return qs.order_by('company', 'custodian')

    @staticmethod
    def _ensure_bank_editor(user, company_id, vd, instance=None):
        """د۱۱: کارت و شبا را فقط مدیریت تنخواه در سطح شرکت عوض می‌کند؛ برمی‌گرداند آیا عوض شده است."""
        changed = any(name in vd and (vd[name] or None) != (getattr(instance, name, None) or None)
                      for name in ('card_or_account_number', 'sheba_number'))
        if changed and not user.is_superuser and not can_at_company_level(user, PETTY_CASH_MANAGE_PERM, company_id):
            raise PermissionDenied("کارت و شبای تنخواه‌گردان را فقط مدیریت تنخواه در سطح شرکت عوض می‌کند.")
        return changed

    @staticmethod
    def _log_bank_change(user, account, before):
        after = {name: mask_tail(getattr(account, name, None)) for name in ('card_or_account_number', 'sheba_number')}
        _log_petty_audit(user, 'pettycashaccount', account.pk, 'petty_bank_info_changed', '', 'changed',
                         company_id=account.company_id, custodian_id=account.custodian_id,
                         before={k: mask_tail(v) for k, v in before.items()}, after=after)

    def _ensure_lifecycle_actor(self, user, account):
        """د۶ و د۷: مدیریت تنخواه یا مدیر در سطح شرکت؛ نه خود تنخواه‌گردان و نه جانشین او (د۸)."""
        if user.is_superuser:
            return
        if account.custodian_id == user.pk or is_petty_deputy(user, account) or not self._can_manage(user, account):
            raise PermissionDenied("بستن و تحویل حساب تنخواه فقط با مدیریت تنخواه یا مدیر در سطح همین شرکت ممکن است.")

    @action(detail=True, methods=['post'], url_path='close')
    def close(self, request, pk=None):
        """د۶: بستن حساب؛ مانده > ۰ با سند عودت (شمارهٔ پیگیری لازم). حساب بسته دوباره باز نمی‌شود."""
        account = self.get_object()
        self._ensure_lifecycle_actor(request.user, account)
        account, refund = _lifecycle(close_account, request.user, account, request.data.get('reference'),
                                     request.data.get('note'))
        if refund is not None:
            notify_document_changed(refund, 'created')
        data = self.get_serializer(self.get_queryset().filter(pk=account.pk).first() or account).data
        data['refund_id'] = refund.pk if refund else None
        return Response(data)

    @action(detail=True, methods=['post'], url_path='handover')
    def handover(self, request, pk=None):
        """د۷: تحویل به تنخواه‌گردان تازه در یک تراکنش (عودت و بستن حساب قدیمی، حساب تازه با همان سقف و مانده)."""
        account = self.get_object()
        self._ensure_lifecycle_actor(request.user, account)
        raw = request.data.get('new_custodian') or request.data.get('new_custodian_id')
        new_id = _section_id(raw)
        new_custodian = get_user_model().objects.filter(pk=new_id).first() if new_id else None
        old, new = _lifecycle(handover_account, request.user, account, new_custodian,
                              request.data.get('reference'), request.data.get('note'))
        qs = PettyCashAccount.annotate_debt(PettyCashAccount.objects.select_related('company', 'custodian', 'section'))
        return Response({
            'old_account': self.get_serializer(qs.filter(pk=old.pk).first()).data,
            'new_account': self.get_serializer(qs.filter(pk=new.pk).first()).data,
        }, status=status.HTTP_201_CREATED)

    def _can_manage(self, user, account_or_company_id, section_id=None):
        company_id = getattr(account_or_company_id, 'company_id', account_or_company_id)
        return (can_at_company_level(user, PETTY_CASH_MANAGE_PERM, company_id)
                or can_at_company_level(user, PETTY_MANAGER_PERM, company_id)
                or bool(section_id and can(user, PETTY_CASH_MANAGE_PERM, section_id)))

    def perform_create(self, serializer):
        user = self.request.user
        vd = serializer.validated_data
        section = vd.pop('section', None)
        # کلاینت قدیمی بخش می‌فرستد: شرکت از بخش؛ وگرنه شرکت فعال درخواست. حساب تازه بخش نمی‌گیرد (ص۹).
        company_id = section.project.company_id if section is not None else resolve_write_company(self.request)
        if not petty_company_allowed(user, company_id):
            raise PermissionDenied("شما مجاز به تعریف حساب تنخواه برای این شرکت نیستید.")
        if not self._can_manage(user, company_id):
            raise PermissionDenied("تعریف حساب تنخواه فقط با مجوز مدیریت تنخواه در سطح شرکت ممکن است.")
        custodian = vd.get('custodian')
        if custodian is None or not (custodian.is_superuser or
                                     get_user_allowed_companies(custodian).filter(id=company_id).exists()):
            raise ValidationError({'custodian': "تنخواه‌گردان باید عضو همین شرکت باشد."})
        if vd.get('is_active', True) and PettyCashAccount.objects.filter(
                company_id=company_id, custodian=custodian, is_active=True).exists():
            raise ValidationError({'custodian': "این شخص در این شرکت از قبل یک حساب تنخواه فعال دارد."})
        ceiling = vd.get('ceiling_amount') or Decimal('0')
        if ceiling and not can_at_company_level(user, PETTY_MANAGER_PERM, company_id):
            raise PermissionDenied("سقف حساب تنخواه را فقط مدیر تعیین می‌کند.")
        bank_changed = self._ensure_bank_editor(user, company_id, vd)
        with transaction.atomic():
            account = serializer.save(company_id=company_id, section=None)
            if bank_changed:
                self._log_bank_change(user, account, {})
            if ceiling:
                _log_petty_audit(user, 'pettycashaccount', account.pk, 'ceiling_change', '0', str(int(ceiling)),
                                 company_id=company_id, custodian_id=account.custodian_id)
            # د۲: پیش از اولین گردش، سازندهٔ حساب موجودی اولیه را آزادانه تعیین می‌کند
            if account.opening_balance or account.opening_balance_date:
                _log_opening_balance(user, account, None, None,
                                     reason=str(self.request.data.get('reason') or '').strip())

    def perform_update(self, serializer):
        user = self.request.user
        instance = serializer.instance
        vd = serializer.validated_data
        if not self._can_manage(user, instance, instance.section_id):
            raise PermissionDenied("شما مجوز ویرایش حساب تنخواه این شرکت را ندارید.")
        # د۶: حساب بسته فقط خواندنی است و دوباره باز نمی‌شود
        _ensure_account_open(instance)
        bank_before = {name: getattr(instance, name, None) for name in ('card_or_account_number', 'sheba_number')}
        bank_changed = self._ensure_bank_editor(user, instance.company_id, vd, instance)
        vd.pop('section', None)
        if 'custodian' in vd and vd['custodian'].pk != instance.custodian_id:
            raise ValidationError({'custodian': "تنخواه‌گردان حساب عوض نمی‌شود؛ حساب تازه تعریف کنید."})
        old_ceiling = instance.ceiling_amount
        new_ceiling = vd.get('ceiling_amount', old_ceiling)
        if new_ceiling != old_ceiling and not can_at_company_level(user, PETTY_MANAGER_PERM, instance.company_id):
            # ص۳: سقف را فقط مدیر (مجوز تأیید پرداخت در سطح شرکت) تعیین یا عوض می‌کند، نه حسابدار
            raise PermissionDenied("سقف حساب تنخواه را فقط مدیر تعیین یا عوض می‌کند.")
        if vd.get('is_active') and not instance.is_active and PettyCashAccount.objects.filter(
                company_id=instance.company_id, custodian_id=instance.custodian_id, is_active=True).exists():
            raise ValidationError({'custodian': "این شخص در این شرکت از قبل یک حساب تنخواه فعال دارد."})
        old_opening, old_opening_date = instance.opening_balance, instance.opening_balance_date
        opening_changed = (vd.get('opening_balance', old_opening) != old_opening
                           or vd.get('opening_balance_date', old_opening_date) != old_opening_date)
        reason = str(self.request.data.get('reason') or '').strip()
        with transaction.atomic():
            had_activity = False
            if opening_changed:
                # د۲: پس از اولین گردش فقط مدیر (مجوز تأیید پرداخت در سطح شرکت) و با دلیل
                PettyCashAccount.objects.select_for_update().filter(pk=instance.pk).order_by().first()
                had_activity = _petty_account_has_activity(instance)
                if had_activity:
                    if not can_at_company_level(user, PETTY_MANAGER_PERM, instance.company_id):
                        raise _PettyApiError('opening_balance_locked',
                                             "پس از اولین گردش حساب، موجودی اولیه را فقط مدیر عوض می‌کند.",
                                             status.HTTP_403_FORBIDDEN)
                    if not reason:
                        raise _PettyApiError('reason_required', "دلیل تغییر موجودی اولیه الزامی است.")
            account = serializer.save()
            if new_ceiling != old_ceiling:
                _log_petty_audit(user, 'pettycashaccount', account.pk, 'ceiling_change',
                                 str(int(old_ceiling)), str(int(new_ceiling)),
                                 company_id=account.company_id, custodian_id=account.custodian_id)
            if opening_changed:
                _log_opening_balance(user, account, old_opening, old_opening_date, reason=reason,
                                     had_activity=had_activity)
            if bank_changed:
                self._log_bank_change(user, account, bank_before)

    def perform_destroy(self, instance):
        if not self._can_manage(self.request.user, instance, instance.section_id):
            raise PermissionDenied("شما مجوز حذف حساب تنخواه این شرکت را ندارید.")
        _ensure_account_open(instance)
        if (instance.transactions.exists() or instance.statements.exists()
                or instance.reconciliations.exists() or instance.debt_settlements.exists()):
            # هیچ سند تنخواهی بی‌حساب نمی‌ماند؛ حساب دارای سند فقط غیرفعال می‌شود
            raise ValidationError({'error': "این حساب سند تنخواه دارد و حذف نمی‌شود؛ آن را غیرفعال کنید."})
        instance.delete()


class PettyCashDebtSettlementViewSet(viewsets.ReadOnlyModelViewSet):
    """
    تسویهٔ بدهی تنخواه‌گردان (د۴ طرح ۴.۲). فهرست: تنخواه‌گردان تسویه‌های حساب خودش، حسابدار و بیننده‌های سطح شرکت
    تسویه‌های شرکت. ثبت و ابطال فقط با `perm_invoice_review` در سطح همان شرکت؛ تنخواه‌گردان روی حساب خودش نه.
    حذف وجود ندارد (ابطال به‌جای حذف).
    """
    queryset = PettyCashDebtSettlement.objects.all().select_related('account', 'created_by', 'charge')
    serializer_class = PettyCashDebtSettlementSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(company_id=cid)
        params = self.request.query_params
        for param, field in (('account_id', 'account_id'), ('status', 'status'), ('method', 'method'),
                             ('charge_id', 'charge_id')):
            value = params.get(param)
            if value:
                qs = qs.filter(**{field: value})
        if not user.is_superuser:
            qs = qs.filter(company_id__in=get_user_allowed_companies(user).values('id')).filter(
                Q(account__custodian=user) | Q(company_id__in=_petty_viewer_company_ids(user)))
        return qs.order_by('-created_at', '-id')

    @staticmethod
    def _error(exc):
        return Response({'code': exc.code, 'error': exc.message_fa, 'detail': exc.message_fa}, status=exc.status)

    @staticmethod
    def _ensure_reviewer(user, account):
        """حسابدار همان شرکت (سطح شرکت)؛ کسی بدهی خودش را تسویه نمی‌کند."""
        if user.is_superuser:
            return
        if account.custodian_id == user.pk or is_petty_deputy(user, account) \
                or not can_at_company_level(user, PETTY_REVIEW_PERM, account.company_id):
            raise PermissionDenied("ثبت و ابطال تسویهٔ بدهی تنخواه فقط با مجوز حسابدار در سطح همین شرکت ممکن است.")

    def create(self, request, *args, **kwargs):
        user = request.user
        raw = request.data.get('account') or request.data.get('account_id')
        account_id = _section_id(raw)
        account = PettyCashAccount.objects.filter(pk=account_id).first() if account_id else None
        # حساب بیرون از شرکت‌های مجاز کاربر یا شرکت فعال دیگر: 404 (وجودش لو نمی‌رود)
        active = get_request_company_id(request)
        if account is None or not petty_company_allowed(user, account.company_id) \
                or (active and str(active) != str(account.company_id)):
            raise Http404
        self._ensure_reviewer(user, account)
        _ensure_account_open(account)
        charge = request.data.get('charge') or request.data.get('charge_id')
        try:
            settlement = create_settlement(
                user, account, request.data.get('amount'), str(request.data.get('method') or '').strip(),
                reference=request.data.get('reference'), charge=_section_id(charge) if charge else None,
                note=request.data.get('note'))
        except SettlementError as exc:
            return self._error(exc)
        return Response(self.get_serializer(settlement).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='void')
    def void(self, request, pk=None):
        settlement = self.get_object()
        self._ensure_reviewer(request.user, settlement.account)
        _ensure_account_open(settlement.account)
        try:
            settlement = void_settlement(request.user, settlement, request.data.get('reason'))
        except SettlementError as exc:
            return self._error(exc)
        return Response(self.get_serializer(settlement).data)

class PettyCashReconciliationViewSet(viewsets.ReadOnlyModelViewSet):
    """
    تطبیق و بستن دورهٔ تنخواه (د۵ طرح ۴.۲): فهرست با همان قلمرو تسویه‌ها؛ ثبت فقط با حسابدار سطح شرکت (نه
    تنخواه‌گردان و نه جانشینش). تطبیق ویرایش یا حذف نمی‌شود؛ اشتباه با تطبیق تازه درست می‌شود.
    """
    queryset = PettyCashReconciliation.objects.all().select_related('account', 'created_by')
    serializer_class = PettyCashReconciliationSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(company_id=cid)
        account_id = self.request.query_params.get('account_id')
        if account_id:
            qs = qs.filter(account_id=account_id)
        if not user.is_superuser:
            qs = qs.filter(company_id__in=get_user_allowed_companies(user).values('id')).filter(
                Q(account__custodian=user) | Q(company_id__in=_petty_viewer_company_ids(user)))
        return qs.order_by('-period_end_date', '-id')

    def create(self, request, *args, **kwargs):
        user = request.user
        account_id = _section_id(request.data.get('account') or request.data.get('account_id'))
        account = PettyCashAccount.objects.filter(pk=account_id).first() if account_id else None
        active = get_request_company_id(request)
        if account is None or not petty_company_allowed(user, account.company_id) \
                or (active and str(active) != str(account.company_id)):
            raise Http404
        PettyCashDebtSettlementViewSet._ensure_reviewer(user, account)
        row = _lifecycle(create_reconciliation, user, account, request.data.get('period_end_date'),
                         request.data.get('counted_cash'), request.data.get('note'))
        return Response(self.get_serializer(row).data, status=status.HTTP_201_CREATED)


class PettyCashTransactionViewSet(viewsets.ModelViewSet):
    """
    مدیریت اسناد، هزینه‌کردها، درخواست‌های شارژ و تسویه تنخواه‌گردان
    """
    queryset = PettyCashTransaction.objects.all().select_related(
        'section', 'section__project', 'custodian', 'counterparty__party', 'created_by', 'account',
        'expense_head__parent'
    )
    serializer_class = PettyCashTransactionSerializer
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]
    # د۱۳: فقط با page صفحه‌بندی می‌شود؛ کلاینت قدیمی و صف آفلاین همان فهرست کامل را می‌گیرند
    pagination_class = InvoicePagePagination

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        qs = scope_petty_transactions_for_user(user, qs, cid)
        section_id = self.request.query_params.get('section_id')
        if section_id:
            qs = qs.filter(section_id=section_id)
        custodian_id = self.request.query_params.get('custodian_id')
        if custodian_id:
            qs = qs.filter(custodian_id=custodian_id)
        statement_id = self.request.query_params.get('statement_id')
        if statement_id:
            qs = qs.filter(statement_id=statement_id)
        line_status = self.request.query_params.get('line_status')
        if line_status:
            qs = qs.filter(line_status=line_status)
        tx_type = self.request.query_params.get('transaction_type')
        if tx_type:
            qs = qs.filter(transaction_type=tx_type)
        status_val = self.request.query_params.get('status')
        if status_val:
            qs = qs.filter(status=status_val)
        date_from = _shamsi_param(self.request, 'date_from')
        if date_from:
            qs = qs.filter(transaction_date_shamsi__gte=date_from)
        date_to = _shamsi_param(self.request, 'date_to')
        if date_to:
            qs = qs.filter(transaction_date_shamsi__lte=date_to)
        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(
                Q(title__icontains=search) |
                Q(receipt_number__icontains=search) |
                Q(category__icontains=search) |
                Q(counterparty__party__name__icontains=search) |
                Q(description__icontains=search) |
                Q(statement__document_number__icontains=search)
            )

        return qs.annotate(offset_total_amount=_petty_offset_subquery()).order_by('-transaction_date_shamsi', '-created_at')

    def _petty_tx_object(self, read_only=False):
        """
        سند برای اقدام گردش: از فهرست قلمرو کاربر، یا از راه کار باز همین سند در شرکت فعال (مثل فاکتور).
        read_only (فقط دیدن جزئیات): شرکت‌کنندهٔ گردش و مدیر/حسابدار سطح شرکت هم (قلاب can_view).
        """
        pk = self.kwargs.get(self.lookup_url_kwarg or self.lookup_field)
        obj = self.filter_queryset(self.get_queryset()).filter(pk=pk).first()
        if obj is None and (wf_engine.task_grants_view(self.request.user, PETTY_CHARGE_ENTITY, pk)
                            or (read_only and can_view_petty_tx(self.request.user, pk))):
            obj = PettyCashTransaction.objects.select_related('account', 'section__project', 'custodian').filter(pk=pk).first()
            active = get_request_company_id(self.request)
            if obj is not None and (not petty_company_allowed(self.request.user, obj.wf_company_id)
                                    or (active and str(active) != str(obj.wf_company_id))):
                obj = None
        if obj is None:
            raise Http404
        return obj

    def retrieve(self, request, *args, **kwargs):
        return Response(self.get_serializer(self._petty_tx_object(read_only=True)).data)

    @staticmethod
    def _ensure_same_company(company_id, section, project):
        if section is not None and section.project.company_id != company_id:
            raise ValidationError({'section': "بخش انتخاب‌شده از شرکت این حساب تنخواه نیست."})
        if project is not None and project.company_id != company_id:
            raise ValidationError({'project': "پروژهٔ انتخاب‌شده از شرکت این حساب تنخواه نیست."})

    def perform_create(self, serializer):
        """ص۱ طرح ۳.۶: تنخواه‌گردان فقط روی حساب خودش ثبت می‌کند؛ بخش و پروژه اختیاری ولی از شرکت همان حساب."""
        user = self.request.user
        vd = serializer.validated_data
        section, project, account = vd.get('section'), vd.get('project'), vd.get('account')
        # د۸: جانشین فعال روی حساب صاحب کار ثبت می‌کند (custodian همان صاحب حساب، created_by جانشین)
        deputy_account = None
        if not user.is_superuser:
            target = account
            if target is None and vd.get('custodian') is not None and vd['custodian'].pk != user.pk:
                target = PettyCashAccount.objects.filter(
                    custodian=vd['custodian'], is_active=True,
                    company_id=resolve_write_company(self.request)).first()
            if target is not None and target.custodian_id != user.pk and is_petty_deputy(user, target):
                deputy_account = account = target
        custodian = vd.get('custodian') or (deputy_account.custodian if deputy_account else user)
        if custodian.pk != user.pk and not user.is_superuser and deputy_account is None:
            raise PermissionDenied("ثبت سند تنخواه فقط روی حساب تنخواه خود شما ممکن است.")
        if account is not None:
            _ensure_account_open(account)
            if account.custodian_id != custodian.pk or not account.is_active:
                raise ValidationError({'account': "حساب تنخواه انتخاب‌شده حساب فعال خود شما نیست."})
            company_id = account.company_id
        elif section is not None:
            company_id = section.project.company_id
        elif project is not None:
            company_id = project.company_id
        else:
            company_id = resolve_write_company(self.request)
        if deputy_account is None:
            _ensure_petty_custodian(user, company_id)
        self._ensure_same_company(company_id, section, project)
        ensure_period_open(company_id, vd.get('transaction_date_shamsi'), 'transaction_date_shamsi')
        if account is None:
            account = _own_petty_account(custodian, company_id, create=True)
        if section is None and account and account.section_id:
            section = account.section
        with transaction.atomic():
            # قفل حساب: بستن، تحویل یا تطبیق هم‌زمان میان سنجش و ذخیره جا نمی‌گیرد
            if account is not None:
                account = PettyCashAccount.objects.select_for_update().get(pk=account.pk)
                _ensure_account_open(account)
            _ensure_period_unlocked(account.pk if account else None, vd.get('transaction_date_shamsi'),
                                    vd.get('transaction_type') or 'expense')
            # ص۱۰: سند تازه پیش‌نویس می‌ماند؛ شارژ با request-charge و هزینه با صورت‌حساب به گردش موتور می‌رود
            instance = serializer.save(created_by=user, custodian=custodian, account=account, section=section, status='draft')
        notify_document_changed(instance, 'created')

    def update(self, request, *args, **kwargs):
        # د۵ و د۶: حساب بسته و ردیف دورهٔ تطبیق‌شده پیش از اعتبارسنجی بدنه با کد روشن رد می‌شوند (صف آفلاین)
        instance = self.get_object()
        _ensure_account_open(instance.account)
        _ensure_period_unlocked(instance.account_id, instance.transaction_date_shamsi, instance.transaction_type)
        return super().update(request, *args, **kwargs)

    def perform_update(self, serializer):
        user = self.request.user
        instance = serializer.instance
        vd = serializer.validated_data
        deputy = (not user.is_superuser and instance.custodian_id != user.pk
                  and is_petty_deputy(user, instance.account))
        if not user.is_superuser and instance.custodian_id != user.pk and not deputy:
            raise PermissionDenied("فقط تنخواه‌گردان سند خودش را ویرایش می‌کند.")
        _ensure_account_open(instance.account)
        # د۵: ردیف دورهٔ تطبیق‌شده و جابه‌جایی به آن دوره هر دو 400 period_locked
        _ensure_period_unlocked(instance.account_id, instance.transaction_date_shamsi, instance.transaction_type)
        _ensure_period_unlocked(instance.account_id, vd.get('transaction_date_shamsi', instance.transaction_date_shamsi),
                                vd.get('transaction_type', instance.transaction_type))
        new_status = vd.pop('status', None)
        if new_status and new_status != instance.status:
            raise PermissionDenied("وضعیت سند تنخواه فقط از راه گردش کار تغییر می‌کند.")
        for name in ('custodian', 'account'):
            if name in vd and getattr(vd[name], 'pk', None) != getattr(instance, f'{name}_id'):
                raise ValidationError({name: "تنخواه‌گردان و حساب سند تنخواه عوض نمی‌شود."})
        _ensure_petty_tx_editable(instance)
        _ensure_petty_cash_open(instance)
        company_id = instance.wf_company_id
        if not user.is_superuser and not deputy:
            _ensure_petty_custodian(user, company_id)
        self._ensure_same_company(company_id, vd.get('section'), vd.get('project'))
        ensure_period_open(company_id, vd.get('transaction_date_shamsi', instance.transaction_date_shamsi),
                           'transaction_date_shamsi')
        with transaction.atomic():
            # قفل حساب و سنجش دوباره: بستن یا تطبیق هم‌زمان میان سنجش و ذخیره جا نمی‌گیرد
            if instance.account_id:
                _ensure_account_open(PettyCashAccount.objects.select_for_update().get(pk=instance.account_id))
                _ensure_period_unlocked(instance.account_id,
                                        vd.get('transaction_date_shamsi', instance.transaction_date_shamsi),
                                        vd.get('transaction_type', instance.transaction_type))
            saved = serializer.save()
        notify_document_changed(saved, 'updated')

    def perform_destroy(self, instance):
        user = self.request.user
        if not user.is_superuser and instance.custodian_id != user.pk:
            raise PermissionDenied("فقط تنخواه‌گردان سند خودش را حذف می‌کند.")
        _ensure_account_open(instance.account)
        _ensure_period_unlocked(instance.account_id, instance.transaction_date_shamsi, instance.transaction_type)
        _ensure_petty_tx_editable(instance)
        _ensure_petty_cash_open(instance)
        if not user.is_superuser:
            _ensure_petty_custodian(user, instance.wf_company_id)
        if instance.debt_offsets.exists():
            # د۴: شارژی که کسر بدهی رویش ثبت شده (حتی باطل‌شده) حذف نمی‌شود؛ سابقهٔ تسویه می‌ماند
            raise _PettyApiError('has_settlements', "روی این شارژ کسر بدهی ثبت شده و حذف نمی‌شود.")
        # in a transaction, so the event goes out only after the delete commits
        with transaction.atomic():
            notify_document_changed(instance, 'deleted')
            instance.delete()

    @action(detail=True, methods=['post'], url_path='submit')
    def submit_tx(self, request, pk=None):
        instance = self.get_object()
        _ensure_petty_cash_open(instance)
        # ص۱۰: ارسال قدیمی به سرپرست کنار رفته است
        return _petty_retired()

    @action(detail=True, methods=['post'], url_path='approve')
    def approve_tx(self, request, pk=None):
        instance = self.get_object()
        _ensure_petty_cash_open(instance)
        if instance.status not in LEGACY_PETTY_PENDING:
            return _petty_retired()
        # Plan 2.2 (T13): approving petty cash is petty-cash management in the record's own section
        _ensure_petty_cash_manager(request.user, instance.section_id)
        from ..workflow_engine import advance_workflow_step
        advance_workflow_step(instance, request.user, notes=request.data.get('notes'), is_financial=True)
        return Response({'message': 'تایید تراکنش تنخواه با موفقیت ثبت شد.', 'status': instance.status})

    @action(detail=True, methods=['post'], url_path='reject')
    def reject_tx(self, request, pk=None):
        instance = self.get_object()
        _ensure_petty_cash_open(instance)
        if instance.status not in LEGACY_PETTY_PENDING:
            return _petty_retired()
        _ensure_petty_cash_manager(request.user, instance.section_id)
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'درج دلیل رد الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        from ..workflow_engine import reject_workflow
        reject_workflow(instance, request.user, reason=reason)
        return Response({'message': 'تراکنش تنخواه رد گردید.', 'status': instance.status})

    @action(detail=False, methods=['get'], url_path='my-balance')
    def my_balance(self, request):
        """
        ماندهٔ تنخواه حساب سطح شرکت (ص۱ و ص۵ طرح ۳.۶): سقف، شارژهای دریافتی، هزینه‌های تأییدشده، در جریان و بدهی.
        شرکت: section_id (کلاینت قدیمی) یا company_id یا شرکت فعال درخواست.
        """
        user = request.user
        section_id = request.query_params.get('section_id')
        if section_id:
            sid = _section_id(section_id)
            section = ProjectSection.objects.select_related('project').filter(pk=sid).first() if sid else None
            if section is None:
                return Response({'error': 'شناسه بخش نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
            company_id = section.project.company_id
        else:
            raw = request.query_params.get('company_id') or get_request_company_id(request)
            try:
                company_id = int(raw) if raw else None
            except (TypeError, ValueError):
                company_id = None
            if company_id is None:
                allowed = list(get_user_allowed_companies(user).values_list('id', flat=True)[:2])
                company_id = allowed[0] if len(allowed) == 1 else None
            if company_id is None:
                return Response({'error': 'شرکت فعال مشخص نیست.'}, status=status.HTTP_400_BAD_REQUEST)
        if not petty_company_allowed(user, company_id):
            raise PermissionDenied("شما به این شرکت دسترسی ندارید.")

        custodian = user
        target_user_id = request.query_params.get('custodian_id')
        if target_user_id and str(target_user_id) != str(user.pk):
            if user.is_superuser or company_id in _petty_viewer_company_ids(user):
                custodian = get_user_model().objects.filter(id=target_user_id).first() or user
            else:
                principal_account = PettyCashAccount.objects.filter(
                    company_id=company_id, custodian_id=_section_id(target_user_id), is_active=True).first()
                if principal_account is not None and is_petty_deputy(user, principal_account):
                    custodian = principal_account.custodian

        account = PettyCashAccount.objects.filter(company_id=company_id, custodian=custodian, is_active=True).first()
        ceiling = account.ceiling_amount if account else Decimal('0')
        rows = PettyCashTransaction.objects.filter(account=account) if account else PettyCashTransaction.objects.none()
        pending_settlement = rows.filter(
            transaction_type='expense', status__in=['draft', 'pending_supervisor', 'pending_accountant'],
        ).exclude(line_status='rejected').aggregate(total=Sum('amount'))['total'] or Decimal('0')
        # د۱ طرح ۴.۲: سطح تنخواه و بدهی با اجزایشان
        level = petty_fund_level_breakdown(account)
        allocations_sum = level['received_total'] - level['offset_total']
        spent_sum = level['spent_total']
        current_balance = level['fund_level']
        debt_parts = account.debt_breakdown() if account else dict.fromkeys(
            ('rejected_total', 'shortage_total', 'settled_total', 'debt'), Decimal('0'))
        debt = debt_parts['debt']
        pending_debt_settlements = PettyCashDebtSettlement.objects.filter(account=account, status='pending') \
            .aggregate(total=Sum('amount'))['total'] or Decimal('0') if account else Decimal('0')

        utilization_rate = 0.0
        if ceiling > Decimal('0'):
            utilization_rate = float(min(Decimal('100.0'), (spent_sum / ceiling) * Decimal('100.0')))
        custodian_full_name = f"{custodian.first_name} {custodian.last_name}".strip() or custodian.username
        # د۱۱: کارت و شبای کامل فقط برای خود تنخواه‌گردان، خزانه‌دار و مدیریت تنخواه سطح شرکت
        full_bank = custodian.pk == user.pk or petty_bank_full_view(user, company_id)
        card = account.card_or_account_number if account else ''
        sheba = account.sheba_number if account else ''
        if not full_bank:
            card, sheba = mask_tail(card), mask_tail(sheba)

        return Response({
            'account_id': account.id if account else None,
            'company_id': company_id,
            'ceiling_amount': int(ceiling),
            'ceiling_amount_tomans': int(ceiling // 10),
            'total_allocated': int(allocations_sum),
            'total_allocated_tomans': int(allocations_sum // 10),
            'total_spent': int(spent_sum),
            'total_spent_tomans': int(spent_sum // 10),
            'current_balance': int(current_balance),
            'current_balance_tomans': int(current_balance // 10),
            'pending_settlement': int(pending_settlement),
            'pending_settlement_tomans': int(pending_settlement // 10),
            'debt': int(debt),
            'debt_tomans': int(debt // 10),
            'opening_balance': int(level['opening_balance']),
            'opening_balance_date': account.opening_balance_date if account else None,
            'refund_total': int(level['refund_total']),
            'rejected_total': int(debt_parts['rejected_total']),
            'shortage_total': int(debt_parts['shortage_total']),
            'settled_total': int(debt_parts['settled_total']),
            'pending_debt_settlements': int(pending_debt_settlements),
            'payroll_deduction_available': payroll_deduction_available(),
            'utilization_rate': round(utilization_rate, 1),
            'card_or_account_number': card,
            'sheba_number': sheba,
            'custodian_name': custodian_full_name,
            'custodian_id': custodian.id
        })

    @action(detail=False, methods=['post'], url_path='request-replenishment')
    def request_replenishment(self, request):
        """ص۱۰: مسیر قدیمی شارژ کنار رفته است؛ جایش request-charge است."""
        return _petty_retired()

    def _submit_charge(self, tx, user):
        try:
            with transaction.atomic():
                wf_engine.submit(PETTY_CHARGE_WORKFLOW, tx, user)
        except WorkflowError as exc:
            return exc
        return None

    @action(detail=False, methods=['post'], url_path='request-charge')
    def request_charge(self, request):
        """
        ص۲ طرح ۳.۶: درخواست شارژ تنخواه‌گردان روی حساب فعال خودش در شرکت؛ سند allocation ساخته و به گردش
        petty_cash_charge فرستاده می‌شود (حسابدار ← مدیر اگر از سقف بگذرد ← خزانه‌دار).
        """
        user = request.user
        try:
            amount = Decimal(str(request.data.get('amount')))
            if amount <= 0 or amount != amount.to_integral_value():
                raise ValueError()
        except Exception:
            return Response({'error': 'مبلغ باید عددی صحیح و مثبت باشد.'}, status=status.HTTP_400_BAD_REQUEST)
        company_id = resolve_write_company(request)
        principal_account = None
        raw_custodian = request.data.get('custodian') or request.data.get('custodian_id')
        if raw_custodian and str(raw_custodian) != str(user.pk):
            # د۸: جانشین برای حساب فعال صاحب کار شارژ می‌خواهد
            principal_account = PettyCashAccount.objects.filter(
                company_id=company_id, custodian_id=_section_id(raw_custodian), is_active=True).first()
            if principal_account is None or not is_petty_deputy(user, principal_account):
                raise PermissionDenied("درخواست شارژ فقط برای حساب خود شما یا صاحب کاری که جانشین اوست ممکن است.")
        else:
            _ensure_petty_custodian(user, company_id)
        today = to_shamsi_str(timezone.now())
        ensure_period_open(company_id, today, 'transaction_date_shamsi')
        with transaction.atomic():
            account = principal_account or _own_petty_account(user, company_id, create=True)
            tx = PettyCashTransaction.objects.create(
                account=account, custodian=account.custodian, created_by=user, transaction_type='allocation',
                amount=amount,
                transaction_date_shamsi=today,
                title=(str(request.data.get('title') or '').strip() or 'درخواست شارژ تنخواه')[:200],
                category=petty_allocation_category('charge'), source='charge',
                description=str(request.data.get('description') or ''), status='draft',
            )
            err = self._submit_charge(tx, user)
            if err is not None:
                transaction.set_rollback(True)
        if err is not None:
            return _petty_wf_error(err)
        notify_document_changed(tx, 'created')
        return Response(self.get_serializer(tx).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='send-charge')
    def send_charge(self, request, pk=None):
        """ارسال دوبارهٔ درخواست شارژ پیش‌نویس (پس از عودت یا لغو) به دست خود تنخواه‌گردان."""
        tx = self.get_object()
        deputy = tx.custodian_id != request.user.pk and is_petty_deputy(request.user, tx.account)
        if (tx.custodian_id != request.user.pk and not deputy) or tx.transaction_type != 'allocation':
            raise PermissionDenied("فقط تنخواه‌گردان درخواست شارژ خودش را می‌فرستد.")
        _ensure_account_open(tx.account)
        _ensure_petty_tx_editable(tx)
        _ensure_petty_cash_open(tx)
        if not deputy:
            _ensure_petty_custodian(request.user, tx.wf_company_id)
        err = self._submit_charge(tx, request.user)
        if err is not None:
            return _petty_wf_error(err)
        return Response(self.get_serializer(tx).data)

    @action(detail=True, methods=['post'], url_path='pay-charge')
    def pay_charge(self, request, pk=None):
        """واریز خزانه‌دار (مرحلهٔ EXECUTE) با شمارهٔ پیگیری؛ شارژ تأیید و به ماندهٔ حساب افزوده می‌شود."""
        tx = self._petty_tx_object()
        _ensure_petty_cash_open(tx)
        tracking = str(request.data.get('tracking_code') or request.data.get('payment_ref') or '').strip()
        if not tracking:
            return Response({'error': 'درج شماره پیگیری یا شماره سند بانکی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        receipt = request.FILES.get(RECEIPT_FIELD)
        receipt_meta = read_receipt(receipt)
        task = wf_engine.open_task(PETTY_CHARGE_WORKFLOW, tx.pk)
        if task is None or task.step_key != 'treasury_pay':
            return Response({'code': 'not_payable', 'error': 'درخواست شارژ هنوز به مرحلهٔ واریز نرسیده است.'},
                            status=status.HTTP_409_CONFLICT)
        revision = request.data.get('revision')
        revision = int(revision) if str(revision or '').isdigit() else None
        err = None
        with transaction.atomic():
            # قفل حساب: دو واریز هم‌زمان سقف را با هم نمی‌شکنند
            if tx.account_id:
                PettyCashAccount.objects.select_for_update().filter(pk=tx.account_id).order_by().first()
            # د۱۰: مرحلهٔ مدیر در این دور رد شده بود ولی حالا (سطح تازه + شارژهای در جریان + خالص این شارژ) از سقف
            # بیشتر است: واریز رد می‌شود تا شارژ عودت و دوباره (این بار با مدیر) فرستاده شود
            if not _manager_approved_current_round(tx) and petty_charge_exceeds_ceiling(tx):
                msg = 'با وضع فعلی تنخواه این شارژ از سقف حساب بیشتر است؛ آن را عودت دهید تا به تأیید مدیر برسد.'
                return Response({'code': 'ceiling_exceeded_at_pay', 'error': msg, 'detail': msg},
                                status=status.HTTP_409_CONFLICT)
            tx.receipt_number = tracking[:100]
            tx.save(update_fields=['receipt_number', 'updated_at'])
            tx._payment_receipt_checked = True  # قلاب پایان گردش فقط با این پرچم شارژ را می‌بندد
            try:
                wf_engine.act(request.user, task, reason=str(request.data.get('notes') or '').strip(),
                              revision=revision, obj=tx)
                # د۴: با واریز، کسرهای در انتظار این شارژ قطعی می‌شوند (خزانه‌دار مبلغ خالص را واریز کرده است)
                complete_charge_offsets(tx, request.user)
                attach_receipt(request.user, PETTY_CHARGE_ENTITY, tx.pk, receipt, receipt_meta)
            except WorkflowError as exc:
                err = exc
                transaction.set_rollback(True)
        if err is not None:
            tx.refresh_from_db()
            return _petty_wf_error(err)
        tx = self.get_queryset().filter(pk=tx.pk).first() or tx
        return Response(self.get_serializer(tx).data)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی اکسل رسمی ۲ ردیفه گردش و تراکنش‌های تنخواه‌گردان مطابق استانداردهای سیستم
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        qs = self.get_queryset()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "گردش تنخواه‌گردان"
        ws.views.sheetView[0].rightToLeft = True

        headers_fa = [
            'ردیف', 'نوع تراکنش', 'شماره رسید/فاکتور', 'تاریخ (شمسی)', 'پروژه', 'بخش',
            'تنخواه‌دار', 'طرف‌حساب/فروشنده', 'سرفصل هزینه', 'عنوان / شرح مختصر',
            'مبلغ (ریال)', 'وضعیت سند', 'ثبت‌کننده', 'توضیحات تفصیلی'
        ]
        headers_en = [
            'row_num', 'transaction_type_display', 'receipt_number', 'transaction_date_shamsi',
            'project_name', 'section_name', 'custodian_name', 'counterparty_name',
            'category', 'title', 'amount', 'status_display', 'created_by_name', 'description'
        ]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(headers_fa) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.freeze_panes = 'A3'

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, tx in enumerate(qs, start=1):
            curr_row = idx + 2
            cust_name = f"{tx.custodian.first_name} {tx.custodian.last_name}".strip() or tx.custodian.username
            cp_name = tx.counterparty.name if tx.counterparty else '-'
            creator_name = f"{tx.created_by.first_name} {tx.created_by.last_name}".strip() if tx.created_by else '-'
            row_data = [
                idx,
                tx.get_transaction_type_display(),
                tx.receipt_number or '-',
                tx.transaction_date_shamsi,
                tx.section.project.name if tx.section else '-',
                tx.section.name if tx.section else '-',
                cust_name,
                cp_name,
                tx.category,
                tx.title,
                float(tx.amount),
                tx.get_status_display(),
                creator_name,
                tx.description or '-'
            ]
            ws.append(row_data)

            for col_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal='center' if col_idx not in [8, 10, 14] else 'right', vertical='center')
                if col_idx in [1, 3, 4, 11]:
                    c.font = data_font_num
                else:
                    c.font = data_font
                if col_idx == 11:
                    c.number_format = '#,##0'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"petty_cash_transactions_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

class PettyCashStatementViewSet(viewsets.ReadOnlyModelViewSet):
    """
    صورت‌حساب تنخواه (ص۴ طرح ۳.۶): تنخواه‌گردان هزینه‌های باز خودش را در یک صورت‌حساب می‌فرستد ← بررسی حسابدار ←
    تأیید مدیر (سقف شرکت ردش می‌کند) ← شارژ دوباره به دست خزانه‌دار به مبلغ ردیف‌های تأییدشده.
    بررسی، تأیید و عودت کل با دلیل از اقدام‌های عمومی «کارهای من» موتور؛ رد ردیف و پرداخت از همین‌جا.
    """
    serializer_class = PettyCashStatementSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = InvoicePagePagination  # د۱۳

    def get_queryset(self):
        from django.db.models import Count
        user = self.request.user
        qs = PettyCashStatement.objects.select_related('company', 'account', 'custodian', 'created_by') \
            .annotate(lines_count=Count('lines'))
        cid = validate_user_company_access(user, get_request_company_id(self.request))
        if cid:
            qs = qs.filter(company_id=cid)
        st = self.request.query_params.get('status')
        if st:
            qs = qs.filter(status=st)
        custodian_id = self.request.query_params.get('custodian_id')
        if custodian_id:
            qs = qs.filter(custodian_id=custodian_id)
        account_id = self.request.query_params.get('account_id')
        if account_id:
            qs = qs.filter(account_id=account_id)
        date_from = _shamsi_param(self.request, 'date_from')
        if date_from:
            qs = qs.filter(statement_date_shamsi__gte=date_from)
        date_to = _shamsi_param(self.request, 'date_to')
        if date_to:
            qs = qs.filter(statement_date_shamsi__lte=date_to)
        search = (self.request.query_params.get('search') or '').strip()
        if search:
            term = Q(document_number__icontains=search) | Q(custodian__first_name__icontains=search) \
                | Q(custodian__last_name__icontains=search) | Q(custodian__username__icontains=search)
            if search.isdigit():
                term |= Q(pk=int(search))
            qs = qs.filter(term)
        if not user.is_superuser:
            qs = qs.filter(company_id__in=get_user_allowed_companies(user).values('id')).filter(
                Q(custodian=user) | Q(company_id__in=_petty_viewer_company_ids(user))
                | deputy_account_q(user, prefix='account__'))
        return qs.order_by('-statement_date_shamsi', '-created_at')

    def _statement_object(self, read_only=False):
        pk = self.kwargs.get(self.lookup_url_kwarg or self.lookup_field)
        obj = self.get_queryset().filter(pk=pk).first()
        if obj is None and (wf_engine.task_grants_view(self.request.user, PETTY_STATEMENT_ENTITY, pk)
                            or (read_only and can_view_petty_statement(self.request.user, pk))):
            obj = PettyCashStatement.objects.select_related('company', 'account', 'custodian').filter(pk=pk).first()
            active = get_request_company_id(self.request)
            if obj is not None and (not petty_company_allowed(self.request.user, obj.company_id)
                                    or (active and str(active) != str(obj.company_id))):
                obj = None
        if obj is None:
            raise Http404
        return obj

    def retrieve(self, request, *args, **kwargs):
        return Response(self.get_serializer(self._statement_object(read_only=True)).data)

    @staticmethod
    def _submit(statement, user):
        try:
            with transaction.atomic():
                # د۹: شمارهٔ ۲.۸ در اولین ارسال، در همان تراکنش؛ ارسال دوباره همان شماره را نگه می‌دارد
                ensure_statement_number(statement, user)
                wf_engine.submit(PETTY_STATEMENT_WORKFLOW, statement, user)
        except WorkflowError as exc:
            return exc
        return None

    def create(self, request, *args, **kwargs):
        """ساخت صورت‌حساب از ردیف‌های هزینهٔ باز خود (یک حساب، یک شرکت) و فرستادن آن به گردش."""
        user = request.user
        raw_ids = request.data.get('line_ids') or []
        try:
            ids = sorted({int(x) for x in raw_ids})
        except (TypeError, ValueError):
            ids = []
        if not ids:
            return Response({'error': 'دست‌کم یک ردیف هزینه برای صورت‌حساب لازم است.'}, status=status.HTTP_400_BAD_REQUEST)
        date = to_shamsi_str(request.data.get('statement_date_shamsi')) or to_shamsi_str(timezone.now())
        err = None
        with transaction.atomic():
            lines = list(PettyCashTransaction.objects.select_for_update()
                         .filter(pk__in=ids, transaction_type='expense'))
            owner_id = lines[0].custodian_id if lines else None
            # د۸: جانشین فعال برای ردیف‌های حساب صاحب کار صورت‌حساب می‌سازد
            deputy = bool(lines) and owner_id != user.pk and all(l.custodian_id == owner_id for l in lines) \
                and is_petty_deputy(user, lines[0].account)
            if not deputy:
                lines = [l for l in lines if l.custodian_id == user.pk]
            if len(lines) != len(ids):
                return Response({'error': 'برخی ردیف‌ها پیدا نشدند یا هزینهٔ خود شما نیستند.'},
                                status=status.HTTP_400_BAD_REQUEST)
            if any(l.statement_id or l.line_status != 'open' or l.status != 'draft' for l in lines):
                return Response({'error': 'برخی ردیف‌ها از قبل در صورت‌حساب یا گردش دیگری هستند.'},
                                status=status.HTTP_400_BAD_REQUEST)
            account_ids = {l.account_id for l in lines}
            account = lines[0].account
            _ensure_account_open(account)
            if len(account_ids) != 1 or account is None or not account.is_active or not account.company_id:
                return Response({'error': 'همهٔ ردیف‌ها باید از یک حساب تنخواه فعال باشند.'},
                                status=status.HTTP_400_BAD_REQUEST)
            company_id = account.company_id
            if not deputy:
                _ensure_petty_custodian(user, company_id)
            for line in lines:
                ensure_period_open(company_id, line.transaction_date_shamsi, 'transaction_date_shamsi')
            ensure_period_open(company_id, date, 'statement_date_shamsi')
            statement = PettyCashStatement.objects.create(
                company_id=company_id, account=account, custodian_id=account.custodian_id, statement_date_shamsi=date,
                created_by=user, status='draft')
            PettyCashTransaction.objects.filter(pk__in=ids).update(statement=statement)
            recompute_statement(statement)
            err = self._submit(statement, user)
            if err is not None:
                transaction.set_rollback(True)
        if err is not None:
            return _petty_wf_error(err)
        return Response(self.get_serializer(statement).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='send')
    def send(self, request, pk=None):
        """ارسال دوبارهٔ صورت‌حساب پیش‌نویس (پس از لغو) به دست خود تنخواه‌گردان."""
        statement = self._statement_object()
        deputy = statement.custodian_id != request.user.pk and is_petty_deputy(request.user, statement.account)
        if statement.custodian_id != request.user.pk and not deputy:
            raise PermissionDenied("فقط تنخواه‌گردان صورت‌حساب خودش را می‌فرستد.")
        _ensure_account_open(statement.account)
        if statement.status != 'draft' or not statement.lines.exists():
            return Response({'error': 'فقط صورت‌حساب پیش‌نویس دارای ردیف فرستاده می‌شود.'}, status=status.HTTP_400_BAD_REQUEST)
        if not deputy:
            _ensure_petty_custodian(request.user, statement.company_id)
        ensure_period_open(statement.company_id, statement.statement_date_shamsi, 'statement_date_shamsi')
        recompute_statement(statement)
        err = self._submit(statement, request.user)
        if err is not None:
            return _petty_wf_error(err)
        return Response(self.get_serializer(statement).data)

    @action(detail=True, methods=['post'], url_path='reject-line')
    def reject_line(self, request, pk=None):
        """
        رد یک ردیف با دلیل (ص۴) به دست حسابدار یا مدیری که مرحلهٔ جاری را دارد (دریافت‌کننده یا مجوز همان مرحله در
        سطح شرکت). ردیف ردشده شارژ دوباره نمی‌گیرد و بدهی تنخواه‌گردان می‌شود (ص۵). هیچ ردیفی حذف نمی‌شود.
        """
        user = request.user
        statement = self._statement_object()
        reason = str(request.data.get('reason') or '').strip()
        if not reason:
            return Response({'error': 'درج دلیل رد ردیف الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        if is_petty_statement_locked(statement.pk):
            return Response({'code': 'entity_locked', 'error': 'این صورت‌حساب قفل است.'}, status=status.HTTP_409_CONFLICT)
        task = wf_engine.open_task(PETTY_STATEMENT_WORKFLOW, statement.pk)
        if task is None or task.step_key not in PETTY_STATEMENT_LINE_REJECT_STEPS:
            return Response({'code': 'not_reviewable', 'error': 'صورت‌حساب در مرحلهٔ بررسی یا تأیید نیست.'},
                            status=status.HTTP_409_CONFLICT)
        step_perm = PETTY_REVIEW_PERM if task.step_key == 'accountant_review' else PETTY_MANAGER_PERM
        holder = wf_engine.task_grants_view(user, PETTY_STATEMENT_ENTITY, statement.pk) \
            or can_at_company_level(user, step_perm, statement.company_id)
        if not petty_company_allowed(user, statement.company_id) or not holder \
                or (statement.custodian_id == user.pk and not user.is_superuser):
            raise PermissionDenied("رد ردیف فقط کار حسابدار یا مدیر مرحلهٔ جاری است.")
        try:
            line_id = int(request.data.get('line_id'))
        except (TypeError, ValueError):
            return Response({'error': 'شناسهٔ ردیف نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            line = statement.lines.select_for_update().filter(pk=line_id).first()
            if line is None:
                raise Http404
            if line.line_status == 'rejected':
                return Response({'error': 'این ردیف از قبل رد شده است.'}, status=status.HTTP_400_BAD_REQUEST)
            old = line.line_status
            line.line_status = 'rejected'
            line.line_reject_reason = reason
            line.save(update_fields=['line_status', 'line_reject_reason', 'updated_at'])
            recompute_statement(statement)
            _log_petty_audit(user, 'pettycashstatement', statement.pk, 'line_reject', old, 'rejected', reason,
                             company_id=statement.company_id, line_id=line.pk, amount=str(line.amount),
                             step=task.step_key)
        notify_document_changed(statement)
        return Response(self.get_serializer(statement).data)

    @action(detail=True, methods=['post'], url_path='pay')
    def pay(self, request, pk=None):
        """شارژ دوباره به دست خزانه‌دار (EXECUTE) با شمارهٔ پیگیری: ردیف‌های باز تأیید و شارژ تازه ساخته می‌شود."""
        statement = self._statement_object()
        tracking = str(request.data.get('tracking_code') or request.data.get('payment_ref') or '').strip()
        if not tracking:
            return Response({'error': 'درج شماره پیگیری یا شماره سند بانکی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        receipt = request.FILES.get(RECEIPT_FIELD)
        receipt_meta = read_receipt(receipt)
        task = wf_engine.open_task(PETTY_STATEMENT_WORKFLOW, statement.pk)
        if task is None or task.step_key != 'treasury_pay':
            return Response({'code': 'not_payable', 'error': 'صورت‌حساب هنوز به مرحلهٔ شارژ دوباره نرسیده است.'},
                            status=status.HTTP_409_CONFLICT)
        revision = request.data.get('revision')
        revision = int(revision) if str(revision or '').isdigit() else None
        statement._payment_ref = tracking[:100]
        statement._paid_by = request.user
        err = None
        with transaction.atomic():
            try:
                wf_engine.act(request.user, task, reason=str(request.data.get('notes') or '').strip(),
                              revision=revision, obj=statement)
                attach_receipt(request.user, PETTY_STATEMENT_ENTITY, statement.pk, receipt, receipt_meta)
            except WorkflowError as exc:
                err = exc
                transaction.set_rollback(True)
        if err is not None:
            statement.refresh_from_db()
            return _petty_wf_error(err)
        return Response(self.get_serializer(statement).data)


class WorkflowAuditLogViewSet(viewsets.ReadOnlyModelViewSet):
    """
    مشاهده تاریخچه ممیزی گردش کار و رویدادهای تایید/رد با ایزولاسیون شرکتی
    """
    queryset = WorkflowAuditLog.objects.all().select_related('actor')
    serializer_class = WorkflowAuditLogSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    # برچسب‌های قدیمی موتور گردش کار ('personnelprofile') هم لاگ پرونده شمرده می‌شوند
    PERSONNEL_LOG_TYPES = ('personnel', 'personnelprofile')
    # مسیر رسیدن به شرکت برای مدل‌های دیگری که موتور گردش کار برایشان لاگ می‌نویسد
    COMPANY_PATH_CANDIDATES = (
        ('company', 'company_id'),
        ('section', 'section__project__company_id'),
        ('project', 'project__company_id'),
        ('personnel', 'personnel__company_id'),
        ('vehicle', 'vehicle__company_id'),
        ('period', 'period__company_id'),
    )

    def _company_scope_q(self, base_qs, allowed_ids):
        """
        DEF-19: محدودیت قطعی لاگ‌ها به شرکت‌های مجاز کاربر، با یا بدون سربرگ شرکت.
        لاگ پروندهٔ حذف‌شده یا Purgeشده با metadata.company_id سنجیده می‌شود.
        """
        personnel_ids = PersonnelProfile.objects.filter(
            Q(company_id__in=allowed_ids) | Q(section__project__company_id__in=allowed_ids)
        ).values('id')
        scope = Q(content_type__in=self.PERSONNEL_LOG_TYPES) & (
            Q(object_id__in=personnel_ids) | Q(metadata__company_id__in=allowed_ids)
        )
        scope |= Q(
            content_type='personnelchangerequest',
            object_id__in=PersonnelChangeRequest.objects.filter(personnel_id__in=personnel_ids).values('id'),
        )

        other_types = (
            base_qs.exclude(content_type__in=self.PERSONNEL_LOG_TYPES + ('personnelchangerequest',))
            .values_list('content_type', flat=True).distinct()
        )
        personnel_models = {m.__name__.lower(): m for m in apps.get_app_config('personnel').get_models()}
        for ct in other_types:
            model = personnel_models.get(ct)
            if model is None:
                continue
            field_names = {f.name for f in model._meta.get_fields()}
            path = next((p for name, p in self.COMPANY_PATH_CANDIDATES if name in field_names), None)
            if path is None:
                continue
            scope |= Q(content_type=ct, object_id__in=model.objects.filter(**{f'{path}__in': allowed_ids}).values('pk'))
        return scope

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        cid = validate_user_company_access(user, get_request_company_id(self.request))

        content_type = (self.request.query_params.get('content_type') or '').strip().lower()
        object_id = self.request.query_params.get('object_id')
        oid = None
        if object_id:
            try:
                oid = int(object_id)
            except (ValueError, TypeError):
                return qs.none()

        if content_type in self.PERSONNEL_LOG_TYPES:
            qs = qs.filter(content_type__in=self.PERSONNEL_LOG_TYPES)
            if oid is not None:
                qs = qs.filter(object_id=oid)
        elif content_type == 'personnelchangerequest':
            # Diff Viewer: لاگ اکشن‌های درخواست تغییرات روی تایملاین پرونده و با change_request_id ثبت می‌شود
            if oid is not None:
                qs = qs.filter(
                    Q(content_type='personnelchangerequest', object_id=oid) |
                    Q(content_type__in=self.PERSONNEL_LOG_TYPES, metadata__change_request_id=oid)
                )
            else:
                qs = qs.filter(
                    Q(content_type='personnelchangerequest') |
                    Q(content_type__in=self.PERSONNEL_LOG_TYPES, metadata__has_key='change_request_id')
                )
        elif content_type:
            qs = qs.filter(content_type=content_type)
            if oid is not None:
                qs = qs.filter(object_id=oid)
        elif oid is not None:
            qs = qs.filter(object_id=oid)

        # ایزولاسیون شرکتی همیشگی (DEF-19): بدون سربرگ هم به شرکت‌های مجاز کاربر محدود می‌شود
        if not user.is_superuser:
            allowed_ids = [cid] if cid else list(get_user_allowed_companies(user).values_list('id', flat=True))
            if not allowed_ids:
                return qs.none()
            qs = qs.filter(self._company_scope_q(qs, allowed_ids))

        return qs.order_by('-created_at')

__all__ = [
    'CounterpartyViewSet',
    'ExpenseInvoiceViewSet',
    'PettyCashAccountViewSet',
    'PettyCashDebtSettlementViewSet',
    'PettyCashTransactionViewSet',
    'PettyCashStatementViewSet',
    'WorkflowAuditLogViewSet',
]
