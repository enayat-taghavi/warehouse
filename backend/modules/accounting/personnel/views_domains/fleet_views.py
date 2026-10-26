# -*- coding: utf-8 -*-
"""
ViewSets for vehicle & driver profiles and vehicle change requests.
"""

import io
from platform_core.registry import has_permission_alias
import openpyxl
from decimal import Decimal
from rest_framework import viewsets, permissions, status, serializers
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied
from django.http import HttpResponse
from django.db import transaction
from django.db.models import Count, Q
from django.utils import timezone
from common.date_utils import normalize_digits
from personnel.permissions import (
    CanApproveFleetManager,
    CanApproveFleetFinance,
)
from companies.access import (
    get_request_company_id,
    validate_user_company_access,
)
from ..models import (
    VehicleDriverProfile,
    VehicleChangeRequest,
    VehicleTripLog,
    WorkflowAuditLog,
    ProjectSection,
)
from companies.org_roles import user_has_section_role, user_sections
from ..serializers import (
    VehicleDriverProfileSerializer,
    VehicleChangeRequestSerializer,
    normalize_plate,
)
from ..permissions import (
    requires_company_scope,
)
from .scope_helpers import (
    broadcast_vehicle_update,
    ensure_record_write_scope,
    is_vehicle_global_auditor,
    log_vehicle_audit,
    readable_company_ids,
    resolve_import_company,
    scope_vehicle_queryset,
)

class VehicleDriverProfileViewSet(viewsets.ModelViewSet):
    queryset = VehicleDriverProfile.objects.all().select_related(
        'user',
        'supervisor_approved_by',
        'accountant_approved_by',
        'manager_approved_by',
        'treasury_paid_by',
        'revision_requested_by',
        'auto_passed_by',
        'created_by',
        'section',
        'project',
    ).prefetch_related('change_requests', 'bank_accounts', 'documents')
    serializer_class = VehicleDriverProfileSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None



    def get_queryset(self):
        qs = super().get_queryset()
        qs = scope_vehicle_queryset(self.request, qs)

        warehouse_id = self.request.query_params.get('warehouse_id')
        if warehouse_id:
            qs = qs.filter(Q(assigned_warehouse_id=warehouse_id) | Q(assigned_warehouse_id__isnull=True))
            
        section_id = self.request.query_params.get('section_id')
        if section_id:
            if not is_vehicle_global_auditor(self.request.user) and self.request.user and self.request.user.is_authenticated:
                if str(section_id).isdigit() and not user_has_section_role(self.request.user, int(section_id)):
                    return qs.none()
            qs = qs.filter(section_id=section_id)
        project_id = self.request.query_params.get('project_id')
        if project_id:
            qs = qs.filter(section__project_id=project_id)

        is_active = self.request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() in ['true', '1'])

        approval_status = self.request.query_params.get('approval_status')
        if approval_status:
            if ',' in approval_status:
                status_list = [s.strip() for s in approval_status.split(',') if s.strip()]
                qs = qs.filter(approval_status__in=status_list)
            else:
                qs = qs.filter(approval_status=approval_status)

        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(
                Q(driver_name__icontains=search) |
                Q(plate_number__icontains=search) |
                Q(driver_national_code__icontains=search) |
                Q(driver_phone__icontains=search)
            )
        return qs

    def create(self, request, *args, **kwargs):
        data = request.data.copy() if hasattr(request.data, 'copy') else dict(request.data)
        section_id = data.get('section')
        if not data.get('company') and section_id:
            try:
                sec = ProjectSection.objects.filter(id=int(section_id)).select_related('project').first()
                if sec and sec.project and sec.project.company_id:
                    data['company'] = sec.project.company_id
            except (ValueError, TypeError):
                pass
        if not data.get('company'):
            cid = get_request_company_id(request)
            if cid:
                valid_cid = validate_user_company_access(request.user, cid)
                if valid_cid:
                    data['company'] = valid_cid
        serializer = self.get_serializer(data=data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        instance = serializer.instance
        if instance:
            instance.refresh_from_db()
            serializer = self.get_serializer(instance)
        headers = self.get_success_headers(serializer.data)
        return Response(serializer.data, status=status.HTTP_201_CREATED, headers=headers)

    def perform_create(self, serializer):
        user = self.request.user
        is_mgr = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or has_permission_alias(user, 'manager_approver')
        
        section = serializer.validated_data.get('section')
        # گارد انتساب بخش: کاربر عادی باید انتساب فعال به بخش انتخاب‌شده داشته باشد
        if not is_mgr and section and user and user.is_authenticated:
            if not user_has_section_role(user, section):
                raise PermissionDenied("شما انتساب فعال در این بخش برای معرفی خودرو جدید ندارید.")

        # سنجش قلمرو شرکتی نوشتن (Task 0.3):
        ensure_record_write_scope(
            user,
            company=serializer.validated_data.get('company'),
            section=section,
            project=serializer.validated_data.get('project'),
            label='ناوگان',
        )

        extra_kwargs = {'created_by': user}
        if section and hasattr(section, 'project') and section.project:
            if not serializer.validated_data.get('project'):
                extra_kwargs['project'] = section.project
            if not serializer.validated_data.get('company') and section.project.company_id:
                extra_kwargs['company_id'] = section.project.company_id

        # تسک ۳.۰ و ب-۱۴: ایجاد اولیه همیشه در وضعیت draft، و عبور فقط با auto_pass_vehicle
        extra_kwargs['approval_status'] = 'draft'
        instance = serializer.save(**extra_kwargs)
        from ..fleet_lifecycle_services import auto_pass_vehicle
        auto_pass_vehicle(instance, user, source='create')
        broadcast_vehicle_update(instance, action_type='created', message=f'پرونده خودرو «{instance.plate_number}» با راننده «{instance.driver_name}» ثبت شد.', sender_id=user.id if user else None)

    def _ensure_update_scope(self, instance, validated_data):
        """جابه‌جایی شرکت، بخش یا پروژهٔ پرونده فقط درون شرکت‌های مجاز کاربر (پرونده به شرکت بیگانه منتقل نمی‌شود)."""
        def _changed(field):
            return field in validated_data and getattr(validated_data[field], 'pk', None) != getattr(instance, f'{field}_id')

        if not any(_changed(f) for f in ('company', 'section', 'project')):
            return
        ensure_record_write_scope(
            self.request.user,
            company=validated_data.get('company', instance.company),
            section=validated_data.get('section', instance.section),
            project=validated_data['project'] if _changed('project') else None,
            label='ناوگان',
        )

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        # ب-۱ (Task 0.1): عبور از قلمرو و واکشی شیء پیش از قفل ردیف درون تراکنش
        target = self.get_object()
        user = request.user
        is_mgr = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or has_permission_alias(user, 'manager_approver')
        is_fin = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_finance')

        from ..fleet_lifecycle_services import (
            VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS,
            compute_profile_diff,
            log_vehicle_change_request_event,
            get_vehicle_missing_fields,
            dispatch_vehicle_workflow_notification,
            refresh_vehicle_has_pending_changes,
        )

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)

            # پروندهٔ ردشده بایگانی است و پیش از هر ویرایش باید با اکشن reopen بازگشایی شود (تسک ۱۱.۱، ن-۱، الگوی پرسنل)
            if instance.approval_status == 'rejected' and not user.is_superuser:
                return Response(
                    {'error': 'این پرونده رد شده است. برای ویرایش و ارسال دوباره، ابتدا آن را بازگشایی کنید.', 'code': 'reopen_required'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # ۱. پرونده در انتظار تصویب مدیر: برای همه جز سوپریوزر قفل است (تسک ۳.۷)
            if instance.approval_status in ['pending_manager', 'accountant_approved']:
                if not user.is_superuser:
                    return Response(
                        {'error': f'این پرونده در وضعیت «{instance.get_approval_status_display()}» قرار دارد و اطلاعات آن قفل است؛ در صورت نیاز، مدیر پرونده را به حسابدار عودت دهد.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )

            # ۲. پرونده در انتظار بررسی حسابدار (pending_accountant):
            # ویرایش مستقیم فقط برای حسابدار یا سوپریوزر مجاز است و به صورت accountant_edit لاگ می‌شود (تسک ۳.۷)
            if instance.approval_status in ['pending_accountant', 'supervisor_approved']:
                if not (user.is_superuser or is_fin):
                    return Response(
                        {'error': f'این پرونده در وضعیت «{instance.get_approval_status_display()}» قرار دارد و اطلاعات آن تا پایان بررسی حسابداری قابل تغییر مستقیم توسط کارمند نیست.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
                serializer = self.get_serializer(instance, data=request.data, partial=partial)
                serializer.is_valid(raise_exception=True)
                self._ensure_update_scope(instance, serializer.validated_data)
                diff, old_diff = compute_profile_diff(instance, serializer.validated_data, excluded=VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS)
                updated_instance = serializer.save()
                if diff:
                    WorkflowAuditLog.objects.create(
                        content_type='vehicle',
                        object_id=updated_instance.id,
                        actor=user,
                        from_status=updated_instance.approval_status,
                        to_status=updated_instance.approval_status,
                        action='accountant_edit',
                        reason='ویرایش مستقیم اطلاعات پرونده خودرو توسط حسابدار',
                        metadata={'diff': diff, 'old_diff': old_diff, 'company_id': updated_instance.company_id}
                    )
                broadcast_vehicle_update(updated_instance, action_type='updated', message=f'اطلاعات مالی خودرو «{updated_instance.plate_number}» توسط حسابدار به‌روزرسانی شد.', sender_id=user.id if user else None)
                return Response(self.get_serializer(updated_instance).data)

            # ۳. پرونده مصوب: سوپریوزر ویرایش مستقیم دارد (با لاگ superuser_edit)؛ سایر کاربران به کارتابل درخواست تغییرات هدایت می‌شوند
            if instance.approval_status in ['approved', 'manager_approved']:
                if user.is_superuser:
                    serializer = self.get_serializer(instance, data=request.data, partial=partial)
                    serializer.is_valid(raise_exception=True)
                    self._ensure_update_scope(instance, serializer.validated_data)
                    diff, old_diff = compute_profile_diff(instance, serializer.validated_data, excluded=VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS)
                    updated_instance = serializer.save()
                    if diff:
                        WorkflowAuditLog.objects.create(
                            content_type='vehicle',
                            object_id=updated_instance.id,
                            actor=user,
                            from_status=updated_instance.approval_status,
                            to_status=updated_instance.approval_status,
                            action='superuser_edit',
                            reason='ویرایش مستقیم پرونده مصوب خودرو توسط سوپریوزر',
                            metadata={'diff': diff, 'old_diff': old_diff, 'company_id': updated_instance.company_id}
                        )
                    broadcast_vehicle_update(updated_instance, action_type='updated', message=f'پرونده مصوب خودرو «{updated_instance.plate_number}» توسط سوپریوزر ویرایش شد.', sender_id=user.id if user else None)
                    return Response(self.get_serializer(updated_instance).data)

                # جلوگیری از ایجاد درخواست تغییرات تکراری تا زمان تعیین تکلیف قبلی (CONCUR-01)
                has_pending = VehicleChangeRequest.objects.filter(
                    vehicle=instance,
                    status__in=['pending_supervisor', 'pending_accountant', 'pending_manager', 'supervisor_approved', 'accountant_approved', 'manager_approved']
                ).exists()
                if has_pending:
                    return Response(
                        {'error': 'این خودرو دارای درخواست تغییرات در حال بررسی است. تا زمان تعیین تکلیف درخواست قبلی امکان ثبت تغییرات جدید وجود ندارد.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )

                from decimal import Decimal, InvalidOperation
                from ..fleet_lifecycle_services import vehicle_change_request_fields_allowed_for

                allowed_change_fields = vehicle_change_request_fields_allowed_for(user)
                # تسک ۱۱.۳ (ن-۳، ط-۲): تلاش صریح نقش غیرمالی برای تغییر نرخ خودروی مصوب رد می‌شود
                if allowed_change_fields is not None and 'default_service_rate' in request.data:
                    raw_new_rate = request.data.get('default_service_rate')
                    if raw_new_rate not in (None, '', 0, '0'):
                        try:
                            new_rate_dec = Decimal(str(raw_new_rate))
                            old_rate_dec = Decimal(str(instance.default_service_rate or 0))
                            if new_rate_dec > 0 and new_rate_dec != old_rate_dec:
                                return Response(
                                    {'error': 'تغییر نرخ پایه سرویس ناوگان مصوب منحصراً در صلاحیت نقش حسابداری ناوگان است.', 'code': 'forbidden_rate_change'},
                                    status=status.HTTP_403_FORBIDDEN
                                )
                        except (InvalidOperation, ValueError):
                            pass

                serializer = self.get_serializer(instance, data=request.data, partial=partial)
                serializer.is_valid(raise_exception=True)
                self._ensure_update_scope(instance, serializer.validated_data)
                
                diff, old_diff = compute_profile_diff(instance, serializer.validated_data, excluded=VEHICLE_CHANGE_REQUEST_EXCLUDED_FIELDS)
                if allowed_change_fields is not None:
                    diff = {k: v for k, v in diff.items() if k in allowed_change_fields}
                    old_diff = {k: v for k, v in old_diff.items() if k in diff}

                if diff:
                    cr = VehicleChangeRequest.objects.create(
                        vehicle=instance,
                        requested_by=user,
                        proposed_changes=diff,
                        previous_values=old_diff,
                        status='pending_supervisor'
                    )
                    instance.has_pending_changes = True
                    instance.save(update_fields=['has_pending_changes'])
                    log_vehicle_change_request_event(cr, user, 'change_request_created', reason='ثبت درخواست تغییرات پرونده مصوب خودرو')
                    # تسک ۱۱.۱۰ (ن-۱۰): ارسال اعلان به سرپرست
                    dispatch_vehicle_workflow_notification(instance, stage='change_request_created', actor=user)
                    return Response({
                        'message': 'درخواست تغییرات خودرو با موفقیت ثبت شد و جهت بررسی در کارتابل سرپرست قرار گرفت. اطلاعات قبلی تا زمان تصویب نهایی معتبر باقی می‌ماند.',
                        'change_request_id': cr.id,
                        'data': self.get_serializer(instance).data
                    }, status=status.HTTP_202_ACCEPTED)
                return Response(self.get_serializer(instance).data)

            # ۴. ویرایش رکوردهای پیش‌نویس یا عودت‌داده‌شده:
            serializer = self.get_serializer(instance, data=request.data, partial=partial)
            serializer.is_valid(raise_exception=True)
            self._ensure_update_scope(instance, serializer.validated_data)

            extra_kwargs = {}
            section = serializer.validated_data.get('section', instance.section)
            if section and hasattr(section, 'project') and section.project:
                if not serializer.validated_data.get('project') and not instance.project_id:
                    extra_kwargs['project'] = section.project

            # تسک ۲.۷، ب-۱۳ و تسک ۱۱.۱ (ن-۱، ط-۱): پل سازگاری PATCH برای وضعیت pending_supervisor
            is_submit_attempt = request.data.get('approval_status') == 'pending_supervisor'
            can_bridge_transition = instance.approval_status in ('draft', 'revision_required')
            updated_instance = serializer.save(**extra_kwargs)

            submitted = False
            missing = []
            if is_submit_attempt:
                if can_bridge_transition:
                    missing = get_vehicle_missing_fields(updated_instance, stage='submit')
                    if not missing:
                        old_st = updated_instance.approval_status
                        updated_instance.approval_status = 'pending_supervisor'
                        updated_instance.rejection_reason = None
                        updated_instance.save(update_fields=['approval_status', 'rejection_reason'])
                        submitted = True
                        WorkflowAuditLog.objects.create(
                            content_type='vehicle',
                            object_id=updated_instance.id,
                            actor=user,
                            from_status=old_st,
                            to_status='pending_supervisor',
                            action='submit',
                            reason='ارسال پرونده خودرو به سرپرست جهت بررسی',
                            metadata={'company_id': updated_instance.company_id}
                        )
                        dispatch_vehicle_workflow_notification(updated_instance, stage='pending_supervisor', actor=user)
                else:
                    submitted = False

            broadcast_vehicle_update(updated_instance, action_type='updated', message=f'پرونده خودرو «{updated_instance.plate_number}» به‌روزرسانی شد.', sender_id=user.id if user else None)
            res_data = self.get_serializer(updated_instance).data
            if is_submit_attempt:
                res_data['submitted'] = submitted
                if not submitted:
                    if can_bridge_transition:
                        res_data['missing_fields'] = missing
                        res_data['detail'] = f"تغییرات ذخیره شد، اما به دلیل نقص اطلاعات ارسال به سرپرست انجام نشد: {', '.join(m['label'] for m in missing)}"
                    else:
                        res_data['detail'] = f"تغییرات ذخیره شد، اما گذار وضعیت از «{instance.get_approval_status_display()}» به «در انتظار سرپرست» مجاز نیست."
            return Response(res_data)

    @staticmethod
    def _user_may_delete(user, instance) -> bool:
        if not user or not user.is_authenticated:
            return False
        if user.is_superuser or user.has_perm('accounts.perm_delete_vehicle_nomovement'):
            return True
        if instance.approval_status not in ('draft', 'revision_required'):
            return False
        return VehicleDriverProfileViewSet._is_creator_or_co_section(user, instance)

    @staticmethod
    def _is_creator_or_co_section(user, instance) -> bool:
        if instance.created_by_id == user.id:
            return True
        return bool(instance.section_id) and user_has_section_role(user, instance.section_id)

    def destroy(self, request, *args, **kwargs):
        """
        تسک ۰.۶، ۱.۲، ۲.۳، ۳.۹، ب-۹ و ب-۱۰: حذف خودرو با قفل ردیف و سنجش سابقه درون تراکنش:
        - اگر خودرو دارای سابقه کارکرد یا تردد، رویداد مالی یا درخواست تغییرات باشد،
          حذف فیزیکی مسدود است و ۴۰۰ با code='history_exists' و پیشنهاد قطع همکاری برمی‌گردد.
        - اگر فاقد سابقه باشد، ثبت‌کننده یا هم‌بخش در draft/revision_required یا دارنده perm_delete_vehicle_nomovement مجاز به حذف است.
        - عکسبرداری کامل پرونده (اسنپ‌شات) در WorkflowAuditLog ثبت می‌گردد.
        """
        from ..fleet_lifecycle_services import check_vehicle_history, build_vehicle_snapshot, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            history = check_vehicle_history(instance)
            if history['has_history']:
                return Response({
                    'error': 'این خودرو دارای سابقه کارکرد، رویداد مالی یا درخواست تغییرات در سامانه است و حذف فیزیکی آن مجاز نمی‌باشد. لطفاً از طریق ثبت تاریخ پایان همکاری و غیرفعال‌سازی پرونده اقدام به قطع همکاری نمایید.',
                    'code': 'history_exists',
                    'suggested_action': 'end_collaboration',
                    'history_reasons': history['reasons'],
                    'history_details': history['details']
                }, status=status.HTTP_400_BAD_REQUEST)

            if not self._user_may_delete(user, instance):
                if self._is_creator_or_co_section(user, instance):
                    return Response(
                        {'error': 'فقط پرونده‌های در وضعیت پیش‌نویس یا عودت‌داده‌شده قابل حذف هستند.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
                raise PermissionDenied("شما مجوز حذف این پرونده خودرو را ندارید.")

            snapshot = build_vehicle_snapshot(instance)
            plate_number = instance.plate_number
            cid = instance.company_id
            target_id = instance.id
            approval_status = instance.approval_status
            broadcast_vehicle_update(instance, action_type='deleted', message=f'خودرو «{plate_number}» حذف شد.', sender_id=user.id if user else None)
            dispatch_vehicle_workflow_notification(instance, stage='deleted', actor=user)
            self.perform_destroy(instance)
            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=target_id,
                actor=user,
                from_status=approval_status,
                to_status='deleted',
                action='delete',
                reason='حذف فیزیکی خودرو فاقد سابقه',
                metadata={'snapshot': snapshot, 'company_id': cid, 'plate_number': plate_number}
            )
            return Response({'message': f'خودرو «{plate_number}» با موفقیت از سامانه حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get'], url_path='delete_impact')
    def delete_impact(self, request, pk=None):
        """
        تسک ۵.۵: استعلام اثرات حذف آبشاری و سوابق ناوگان منطبق بر قرارداد smart-delete-modal و check_vehicle_history
        """
        instance = self.get_object()
        from django.db.models.deletion import Collector, ProtectedError, RestrictedError
        from django.db import router
        from ..fleet_lifecycle_services import check_vehicle_history

        using = router.db_for_write(instance.__class__, instance=instance)
        collector = Collector(using=using)
        impact_summary = []

        try:
            collector.collect([instance])
            for model, instances in collector.data.items():
                if model == instance.__class__:
                    continue
                impact_summary.append({
                    'model_name': str(model._meta.verbose_name or model.__name__),
                    'count': len(instances)
                })
        except (ProtectedError, RestrictedError) as pe:
            prot_counts = {}
            for obj in getattr(pe, 'protected_objects', []):
                m = obj.__class__
                prot_counts[m] = prot_counts.get(m, 0) + 1
            for m, cnt in prot_counts.items():
                impact_summary.append({
                    'model_name': str(m._meta.verbose_name or m.__name__),
                    'count': cnt
                })
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        impact_summary.sort(key=lambda x: x['count'], reverse=True)

        history = check_vehicle_history(instance)
        user = request.user
        can_delete = bool(history['can_delete'] and self._user_may_delete(user, instance))

        return Response({
            'impact': impact_summary,
            'total_affected': sum(item['count'] for item in impact_summary),
            'history': history,
            'can_delete': can_delete,
            'has_history': history['has_history'],
            'reasons': history['reasons'],
            'blocking_reasons': history['reasons'],
            'history_details': history.get('details', {})
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get', 'post', 'put', 'patch', 'delete'], url_path='bank-accounts')
    def bank_accounts(self, request, pk=None):
        """
        تسک ۵.۱: مدیریت حساب‌های بانکی تودرتوی ناوگان با رعایت تفکیک وظایف (SoD).
        صاحب حساب همیشه همین خودرو است (vehicle فقط‌خواندنی)، و تغییر حساب با گارد وضعیت و نقش انجام می‌شود.
        """
        from ..fleet_lifecycle_services import check_vehicle_bank_account_edit_allowed
        from ..lifecycle_services import _snapshot_value
        from ..serializers import VehicleBankAccountSerializer
        instance = self.get_object()

        if request.method == 'GET':
            accounts = instance.bank_accounts.all()
            serializer = VehicleBankAccountSerializer(accounts, many=True)
            return Response(serializer.data, status=status.HTTP_200_OK)

        allowed, deny_msg = check_vehicle_bank_account_edit_allowed(request.user, instance)
        if not allowed:
            return Response({'error': deny_msg}, status=status.HTTP_400_BAD_REQUEST)

        def _log_bank_change(op, account_id, diff=None):
            in_draft = instance.approval_status in ('draft', 'revision_required')
            if in_draft and op != 'delete':
                return
            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=request.user,
                from_status=instance.approval_status,
                to_status=instance.approval_status,
                action='bank_account_deleted' if in_draft else 'accountant_edit',
                reason=request.data.get('note', '') or 'تغییر حساب بانکی خودرو',
                metadata={'bank_account': {'op': op, 'id': account_id, 'diff': diff or {}}, 'company_id': instance.company_id}
            )

        if request.method == 'POST':
            serializer = VehicleBankAccountSerializer(data=request.data)
            serializer.is_valid(raise_exception=True)
            with transaction.atomic():
                acc = serializer.save(vehicle=instance)
                _log_bank_change('create', acc.id, {k: str(v) for k, v in serializer.validated_data.items()})
            broadcast_vehicle_update(instance, action_type='updated', message=f'حساب بانکی جدید برای خودرو «{instance.plate_number}» افزوده شد.', sender_id=request.user.id)
            return Response(VehicleBankAccountSerializer(acc).data, status=status.HTTP_201_CREATED)

        account_id = request.data.get('id') or request.query_params.get('account_id') or request.query_params.get('id')
        if not account_id:
            return Response({'error': 'شناسه حساب بانکی (id) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        acc = instance.bank_accounts.filter(pk=account_id).first()
        if not acc:
            return Response({'error': 'حساب بانکی موردنظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)

        if request.method in ['PUT', 'PATCH']:
            serializer = VehicleBankAccountSerializer(acc, data=request.data, partial=True)
            serializer.is_valid(raise_exception=True)
            diff = {
                k: {'old': str(getattr(acc, k)) if getattr(acc, k) is not None else None, 'new': str(v) if v is not None else None}
                for k, v in serializer.validated_data.items() if getattr(acc, k) != v
            }
            with transaction.atomic():
                updated_acc = serializer.save()
                if diff:
                    _log_bank_change('update', updated_acc.id, diff)
            broadcast_vehicle_update(instance, action_type='updated', message=f'حساب بانکی خودرو «{instance.plate_number}» ویرایش شد.', sender_id=request.user.id)
            return Response(VehicleBankAccountSerializer(updated_acc).data, status=status.HTTP_200_OK)

        if request.method == 'DELETE':
            with transaction.atomic():
                account_snapshot = {
                    f.attname: _snapshot_value(getattr(acc, f.attname))
                    for f in acc._meta.concrete_fields
                }
                _log_bank_change('delete', acc.id, account_snapshot)
                acc.delete()
            broadcast_vehicle_update(instance, action_type='updated', message=f'حساب بانکی خودرو «{instance.plate_number}» حذف شد.', sender_id=request.user.id)
            return Response({'message': 'حساب بانکی با موفقیت حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get', 'post', 'delete'], url_path='documents')
    def documents(self, request, pk=None):
        """
        تسک ۵.۲ و ب-۳۰: مدیریت و بارگذاری مدارک چندگانه ناوگان و خودرو با اعتبارسنجی نوع، پسوند و حجم فایل.
        در حذف مدرک، رکورد دیتابیس حذف می‌شود ولی فایل فیزیکی جهت ممیزی حفظ می‌گردد.
        """
        from ..models import VehicleDocument
        from ..serializers import VehicleDocumentSerializer, validate_uploaded_document
        from ..fleet_lifecycle_services import check_vehicle_bank_account_edit_allowed
        instance = self.get_object()

        if request.method == 'GET':
            docs = instance.documents.all().order_by('-created_at')
            return Response(VehicleDocumentSerializer(docs, many=True).data, status=status.HTTP_200_OK)

        elif request.method == 'POST':
            file_obj = request.FILES.get('file') or request.FILES.get('attachment')
            if not file_obj:
                return Response({'error': 'فایل مدرک الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            doc_type = request.data.get('doc_type') or 'other'
            if doc_type not in dict(VehicleDocument.DOC_TYPE_CHOICES):
                return Response({'error': f'نوع مدرک «{doc_type}» نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)
            try:
                validate_uploaded_document(
                    file_obj,
                    max_size_bytes=10 * 1024 * 1024,
                    allowed_exts=('.pdf', '.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.heic', '.heif'),
                )
            except serializers.ValidationError as e:
                return Response({'error': ' '.join(str(d) for d in e.detail)}, status=status.HTTP_400_BAD_REQUEST)

            from ..fleet_lifecycle_services import normalize_shamsi_date
            title = (request.data.get('title') or '').strip() or None
            raw_issue_date = (request.data.get('issue_date') or '').strip()
            raw_expiry_date = (request.data.get('expiry_date') or '').strip()
            # تاریخ نامعتبر ۴۰۰ روشن می‌گیرد، نه خطای ۵۰۰ (normalize_shamsi_date برای ورودی نامعتبر ValueError می‌دهد)
            try:
                issue_date = normalize_shamsi_date(raw_issue_date) if raw_issue_date else None
                expiry_date = normalize_shamsi_date(raw_expiry_date) if raw_expiry_date else None
            except ValueError as e:
                return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

            with transaction.atomic():
                doc = VehicleDocument.objects.create(
                    vehicle=instance,
                    file=file_obj,
                    doc_type=doc_type,
                    title=title,
                    issue_date=issue_date,
                    expiry_date=expiry_date,
                    uploaded_by=request.user if request.user.is_authenticated else None
                )
                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=request.user if request.user.is_authenticated else None,
                    from_status=instance.approval_status,
                    to_status=instance.approval_status,
                    action='document_uploaded',
                    reason=request.data.get('note', '') or 'بارگذاری مدرک پیوست خودرو',
                    metadata={
                        'document': {
                            'id': doc.id,
                            'doc_type': doc.doc_type,
                            'title': doc.title,
                            'file_name': doc.file.name if doc.file else None,
                            'issue_date': doc.issue_date,
                            'expiry_date': doc.expiry_date,
                            'uploaded_by_id': doc.uploaded_by_id,
                        },
                        'company_id': instance.company_id
                    }
                )
            broadcast_vehicle_update(instance, action_type='updated', message=f'مدرک جدید برای خودرو «{instance.plate_number}» بارگذاری شد.', sender_id=request.user.id)
            return Response(VehicleDocumentSerializer(doc).data, status=status.HTTP_201_CREATED)

        elif request.method == 'DELETE':
            doc_id = request.data.get('doc_id') or request.query_params.get('doc_id') or request.query_params.get('id')
            if not doc_id or not str(doc_id).isdigit():
                return Response({'error': 'شناسه مدرک الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
            allowed, deny_msg = check_vehicle_bank_account_edit_allowed(request.user, instance)
            if not allowed:
                return Response({'error': deny_msg.replace('حساب‌های بانکی', 'مدارک')}, status=status.HTTP_400_BAD_REQUEST)
            doc = instance.documents.filter(id=int(doc_id)).select_related('uploaded_by').first()
            if not doc:
                return Response({'error': 'مدرک مورد نظر یافت نشد.'}, status=status.HTTP_404_NOT_FOUND)
            with transaction.atomic():
                doc_snapshot = {
                    'id': doc.id,
                    'doc_type': doc.doc_type,
                    'title': doc.title,
                    'file_name': doc.file.name if doc.file else None,
                    'issue_date': doc.issue_date,
                    'expiry_date': doc.expiry_date,
                    'uploaded_by_id': doc.uploaded_by_id,
                }
                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=request.user,
                    from_status=instance.approval_status,
                    to_status=instance.approval_status,
                    action='document_deleted',
                    reason=request.data.get('note', '') or 'حذف مدرک پیوست خودرو',
                    metadata={'document': doc_snapshot, 'company_id': instance.company_id}
                )
                doc.delete()
            broadcast_vehicle_update(instance, action_type='updated', message=f'مدرک خودرو «{instance.plate_number}» حذف شد.', sender_id=request.user.id)
            return Response({'message': 'مدرک با موفقیت حذف گردید.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='submit')
    def submit(self, request, pk=None):
        """
        تسک ۲.۷: اکشن صریح ارسال پرونده خودرو به سرپرست جهت بررسی و تایید
        """
        from ..fleet_lifecycle_services import get_vehicle_missing_fields, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user
        is_mgr = user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or has_permission_alias(user, 'manager_approver')
        if not is_mgr and target.created_by_id != user.id:
            if not (target.section_id and user_has_section_role(user, target.section_id)):
                raise PermissionDenied("شما دسترسی ارسال این پرونده به سرپرست را ندارید.")

        if target.approval_status not in ('draft', 'revision_required'):
            return Response(
                {'error': f'ارسال به سرپرست فقط برای پرونده‌های پیش‌نویس یا نیازمند بازنگری مجاز است (وضعیت فعلی: {target.get_approval_status_display()}).'},
                status=status.HTTP_400_BAD_REQUEST
            )

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            missing = get_vehicle_missing_fields(instance, stage='submit')
            if missing:
                return Response({
                    'error': 'فیلدهای الزامی پرونده خودرو جهت ارسال به سرپرست تکمیل نشده است.',
                    'missing_fields': missing,
                    'detail': f"فیلدهای ناقص: {', '.join(m['label'] for m in missing)}"
                }, status=status.HTTP_400_BAD_REQUEST)

            old_status = instance.approval_status
            instance.approval_status = 'pending_supervisor'
            instance.rejection_reason = None
            instance.save(update_fields=['approval_status', 'rejection_reason'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status=old_status,
                to_status='pending_supervisor',
                action='submit',
                reason=(request.data.get('note') or 'ارسال پرونده خودرو به سرپرست').strip(),
                metadata={'company_id': instance.company_id}
            )
            dispatch_vehicle_workflow_notification(instance, stage='pending_supervisor', actor=user)

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» به سرپرست ارسال شد.', sender_id=user.id if user else None)
        return Response({
            'message': 'پرونده خودرو با موفقیت جهت بررسی و تایید به کارتابل سرپرست ارسال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='approve-supervisor')
    def approve_supervisor(self, request, pk=None):
        """
        تسک ۳.۴: تایید مرحله اول (سرپرست) پرونده خودرو با استفاده از ماژول متمرکز گردش کار
        """
        from ..fleet_lifecycle_services import approve_vehicle_stage, VehicleWorkflowError
        target = self.get_object()
        note = (request.data.get('note') or request.data.get('approval_note') or '').strip()
        try:
            instance = approve_vehicle_stage(target.id, request.user, stage='supervisor', note=note)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» به تایید سرپرست رسید و به حسابداری ارسال شد.', sender_id=request.user.id if request.user else None)
        return Response({
            'message': 'تایید سرپرست با موفقیت ثبت شد و پرونده خودرو به حسابداری ارسال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='approve-finance')
    def approve_finance(self, request, pk=None):
        """
        تسک ۳.۴: تایید مرحله دوم (مالی) پرونده خودرو همراه با گیت الزامی نرخ و شماره شبا
        """
        from ..fleet_lifecycle_services import approve_vehicle_stage, VehicleWorkflowError
        target = self.get_object()
        note = (request.data.get('note') or request.data.get('approval_note') or '').strip()
        try:
            instance = approve_vehicle_stage(target.id, request.user, stage='finance', note=note)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» به تایید مالی رسید و به مدیر ارسال شد.', sender_id=request.user.id if request.user else None)
        return Response({
            'message': 'تایید مالی با موفقیت ثبت شد و پرونده خودرو جهت تصویب نهایی به کارتابل مدیر ارسال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='approve-manager')
    def approve_manager(self, request, pk=None):
        """
        تسک ۳.۴: تصویب نهایی مدیر و فعال‌سازی پرونده خودرو
        """
        from ..fleet_lifecycle_services import approve_vehicle_stage, VehicleWorkflowError
        target = self.get_object()
        note = (request.data.get('note') or request.data.get('approval_note') or '').strip()
        try:
            instance = approve_vehicle_stage(target.id, request.user, stage='manager', note=note)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        broadcast_vehicle_update(instance, action_type='updated', message=f'خودرو «{instance.plate_number}» با راننده «{instance.driver_name}» تصویب و فعال شد.', sender_id=request.user.id if request.user else None)
        return Response({
            'message': 'تصویب نهایی مدیر با موفقیت ثبت شد و خودرو فعال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=False, methods=['post'], url_path='batch-approve')
    def batch_approve(self, request):
        """
        تسک ۳.۶ و DEF-09: تایید گروهی پرونده‌های خودرو با تراکنش مستقل برای هر رکورد
        """
        from ..fleet_lifecycle_services import approve_vehicle_stage, FLEET_APPROVAL_STAGES, user_has_any_perm, VehicleWorkflowError

        stage = request.data.get('stage')
        ids = request.data.get('ids', [])
        note = (request.data.get('note') or 'تایید گروهی').strip()

        if not stage or stage not in FLEET_APPROVAL_STAGES:
            return Response({'error': 'مرحله ارسالی نامعتبر است.'}, status=status.HTTP_400_BAD_REQUEST)

        rule = FLEET_APPROVAL_STAGES[stage]
        if not user_has_any_perm(request.user, rule['perms']):
            raise PermissionDenied(rule['perm_error'])

        if not ids or not isinstance(ids, list):
            return Response({'error': 'لیست شناسه‌های خودرو الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        allowed_qs = self.get_queryset().filter(id__in=ids)
        allowed_ids = set(allowed_qs.values_list('id', flat=True))

        success_ids = []
        errors = []

        for v_id in ids:
            if v_id not in allowed_ids:
                errors.append({'id': v_id, 'error': 'خودرو یافت نشد یا در قلمرو دسترسی شما قرار ندارد.'})
                continue
            try:
                approved_v = approve_vehicle_stage(v_id, request.user, stage=stage, note=note, batch=True)
                success_ids.append(v_id)
                broadcast_vehicle_update(
                    approved_v,
                    action_type='updated',
                    message=f'خودرو «{approved_v.plate_number}» به صورت گروهی در مرحله {stage} تایید شد.',
                    sender_id=request.user.id if request.user else None
                )
            except VehicleWorkflowError as e:
                errors.append({'id': v_id, 'error': e.message, 'missing_fields': e.missing_fields})
            except Exception as e:
                errors.append({'id': v_id, 'error': str(e)})

        return Response({
            'success_count': len(success_ids),
            'failed_count': len(errors),
            'success_ids': success_ids,
            'errors': errors,
        })

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        """
        تسک ۳.۱ و ب-۱۸: رد پرونده با فهرست سفید وضعیت‌های در جریان، قفل ردیف، و Purge مشروط در صورت عدم سابقه
        """
        from ..fleet_lifecycle_services import (
            user_has_any_perm,
            FLEET_MANAGER_STAGE_PERMS,
            FLEET_SUPERVISOR_STAGE_PERMS,
            FLEET_FINANCE_STAGE_PERMS,
            check_vehicle_history,
            build_vehicle_snapshot,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        can_mgr = user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS + FLEET_SUPERVISOR_STAGE_PERMS)
        can_fin = user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS)
        if not (can_mgr or can_fin):
            raise PermissionDenied("شما دسترسی لازم برای رد این پرونده را ندارید.")

        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت دلیل رد الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        target = self.get_object()
        valid_statuses = [
            'pending_supervisor', 'pending_accountant', 'pending_manager',
            'supervisor_approved', 'accountant_approved', 'revision_required'
        ]
        if target.approval_status not in valid_statuses:
            return Response(
                {'error': f'امکان رد پرونده در وضعیت «{target.get_approval_status_display()}» وجود ندارد.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            if instance.approval_status not in valid_statuses:
                return Response(
                    {'error': f'وضعیت پرونده تغییر یافته است (وضعیت فعلی: {instance.get_approval_status_display()}).'},
                    status=status.HTTP_409_CONFLICT
                )

            history = check_vehicle_history(instance)
            is_never_approved = instance.manager_approved_at is None
            purge_statuses = ('revision_required', 'pending_supervisor', 'pending_accountant', 'pending_manager')
            can_purge = is_never_approved and history['can_delete'] and instance.approval_status in purge_statuses

            old_status = instance.approval_status
            cid = instance.company_id
            plate = instance.plate_number

            if can_purge:
                snapshot = build_vehicle_snapshot(instance)
                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=user,
                    from_status=old_status,
                    to_status='purged',
                    action='purge',
                    reason=reason,
                    metadata={'snapshot': snapshot, 'company_id': cid}
                )
                instance.delete()
                dispatch_vehicle_workflow_notification(target, stage='purged', actor=user, reason=reason)
                broadcast_vehicle_update(target, action_type='deleted', message=f'پرونده خودرو «{plate}» به دلیل رد قطعی و عدم سابقه پاک‌سازی شد.', sender_id=user.id)
                return Response({
                    'message': f'پرونده خودرو «{plate}» به دلیل رد قطعی و عدم وجود هرگونه سابقه، از سامانه پاک‌سازی شد.',
                    'purged': True
                })
            else:
                instance.approval_status = 'rejected'
                instance.rejection_reason = reason
                instance.save(update_fields=['approval_status', 'rejection_reason'])

                WorkflowAuditLog.objects.create(
                    content_type='vehicle',
                    object_id=instance.id,
                    actor=user,
                    from_status=old_status,
                    to_status='rejected',
                    action='reject',
                    reason=reason,
                    metadata={'company_id': cid}
                )
                dispatch_vehicle_workflow_notification(instance, stage='rejected', actor=user, reason=reason)
                broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{plate}» رد شد: {reason}', sender_id=user.id)
                return Response({
                    'message': f'پرونده خودرو «{plate}» رد شد.',
                    'data': self.get_serializer(instance).data,
                    'purged': False
                })

    @action(detail=True, methods=['post'], url_path='request-revision')
    def request_revision(self, request, pk=None):
        """
        تسک ۳.۲ و ب-۱۹ و DEF-13: عودت پرونده بر پایه وضعیت فعلی:
        - در pending_supervisor: عودت به کارمند (revision_required)
        - در pending_accountant: عودت به کارمند (revision_required) یا سرپرست (pending_supervisor)
        - در pending_manager: عودت به حسابدار (pending_accountant)
        """
        from ..fleet_lifecycle_services import (
            user_has_any_perm,
            FLEET_MANAGER_STAGE_PERMS,
            FLEET_SUPERVISOR_STAGE_PERMS,
            FLEET_FINANCE_STAGE_PERMS,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        target = self.get_object()

        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت توضیحات و دلیل عودت الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            old_status = instance.approval_status

            if old_status == 'pending_supervisor':
                if not user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS):
                    raise PermissionDenied("شما دسترسی سرپرستی برای عودت این پرونده را ندارید.")
                target_status = 'revision_required'

            elif old_status in ('pending_accountant', 'supervisor_approved'):
                if not user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS + FLEET_MANAGER_STAGE_PERMS):
                    raise PermissionDenied("شما دسترسی مالی برای عودت این پرونده را ندارید.")
                target_choice = request.data.get('target_stage', 'employee')
                if target_choice in ('supervisor', 'pending_supervisor'):
                    target_status = 'pending_supervisor'
                else:
                    target_status = 'revision_required'

            elif old_status in ('pending_manager', 'accountant_approved'):
                if not user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS):
                    raise PermissionDenied("شما دسترسی مدیریت برای عودت این پرونده را ندارید.")
                target_status = 'pending_accountant'

            else:
                return Response(
                    {'error': f'امکان عودت پرونده در وضعیت «{instance.get_approval_status_display()}» وجود ندارد.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            instance.approval_status = target_status
            instance.rejection_reason = reason
            instance.save(update_fields=['approval_status', 'rejection_reason'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status=old_status,
                to_status=target_status,
                action='request_revision',
                reason=reason,
                metadata={'company_id': instance.company_id, 'target_stage': target_status}
            )
            dispatch_vehicle_workflow_notification(
                instance, stage=target_status, actor=user, reason=reason, is_return=True
            )

        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» عودت داده شد: {reason}', sender_id=user.id)
        return Response({
            'message': f'پرونده خودرو با موفقیت به وضعیت «{instance.get_approval_status_display()}» عودت داده شد.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['get'], url_path='rate-benchmark')
    def rate_benchmark(self, request, pk=None):
        """
        تسک ۹.۳ و ۱۱.۲ (ن-۲): مقایسهٔ نرخ با میانگین هم‌تیپ‌های پروژه (سروری و محدود به قلمرو).
        گارد نقش: فقط برای کاربران مجاز به مشاهده نرخ‌های ناوگان (مالی و مدیریت).
        """
        user = request.user
        from ..serializers import user_can_view_fleet_rates
        if not user_can_view_fleet_rates(user):
            raise PermissionDenied("مشاهده مقایسه و شاخص نرخ ناوگان فقط برای نقش‌های مدیریت و مالی مجاز است.")

        from django.db.models import Avg, Min, Max, Count
        target = self.get_object()

        # کوئری هم‌تیپ‌ها در همان پروژه یا بخش با رعایت قلمرو شرکتی
        qs = self.get_queryset().filter(
            vehicle_type=target.vehicle_type,
            default_service_rate__gt=0
        )
        if target.section_id:
            if target.section and target.section.project_id:
                qs = qs.filter(section__project_id=target.section.project_id)
            else:
                qs = qs.filter(section_id=target.section_id)

        stats = qs.aggregate(
            count=Count('id'),
            avg_rate=Avg('default_service_rate'),
            min_rate=Min('default_service_rate'),
            max_rate=Max('default_service_rate')
        )

        return Response({
            'vehicle_id': target.id,
            'vehicle_type': target.vehicle_type,
            'vehicle_type_display': target.get_vehicle_type_display() if hasattr(target, 'get_vehicle_type_display') else target.vehicle_type,
            'project_name': target.section.project.name if (target.section and target.section.project) else (target.section.name if target.section else 'عمومی'),
            'count': stats['count'] or 0,
            'avg_rate': round(stats['avg_rate']) if stats['avg_rate'] else 0,
            'min_rate': stats['min_rate'] or 0,
            'max_rate': stats['max_rate'] or 0,
            'current_rate': target.default_service_rate or 0
        })

    @action(detail=True, methods=['get'], url_path='cost-estimate')
    def cost_estimate(self, request, pk=None):
        """
        تسک ۱۰.۱: ویجت بار مالی ماهانه (سروری، محدود به قلمرو) در کارتابل مدیر.
        برآورد بار مالی ماهانه و سالانه خودرو بر مبنای نرخ روزانه، نوع مالکیت و انحراف از میانگین هم‌تیپ‌ها در پروژه.
        """
        user = request.user
        from ..fleet_lifecycle_services import user_has_any_perm, FLEET_MANAGER_STAGE_PERMS, FLEET_FINANCE_STAGE_PERMS
        if not (user.is_superuser or user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS + FLEET_FINANCE_STAGE_PERMS)):
            raise PermissionDenied("مشاهده برآورد مالی و نرخ ناوگان فقط برای نقش‌های مدیریت و مالی مجاز است.")

        target = self.get_object()
        daily_rate = float(target.default_service_rate or 0)
        is_payable = target.ownership_type in ('contract', 'personal')

        # مبنای کاری ماهانه: ۲۶ روز کاری استاندارد (و ۳۰ روز کامل تقویمی)
        standard_days = 26
        full_month_days = 30

        monthly_cost_standard = round(daily_rate * standard_days) if is_payable else 0
        monthly_cost_full = round(daily_rate * full_month_days) if is_payable else 0
        yearly_cost_standard = monthly_cost_standard * 12

        # مقایسه با میانگین هم‌تیپ‌ها در قلمرو پروژه
        from django.db.models import Avg, Count
        qs = self.get_queryset().filter(
            vehicle_type=target.vehicle_type,
            default_service_rate__gt=0
        )
        if target.section_id:
            if target.section and target.section.project_id:
                qs = qs.filter(section__project_id=target.section.project_id)
            else:
                qs = qs.filter(section_id=target.section_id)

        stats = qs.aggregate(
            count=Count('id'),
            avg_rate=Avg('default_service_rate')
        )
        avg_rate = round(stats['avg_rate']) if stats['avg_rate'] else 0
        avg_monthly_cost = round(avg_rate * standard_days) if is_payable else 0
        deviation_percentage = round(((daily_rate - avg_rate) / avg_rate) * 100) if (avg_rate > 0 and is_payable) else 0

        return Response({
            'vehicle_id': target.id,
            'plate_number': target.plate_number,
            'driver_name': target.driver_name,
            'vehicle_type': target.vehicle_type,
            'vehicle_type_display': target.get_vehicle_type_display() if hasattr(target, 'get_vehicle_type_display') else target.vehicle_type,
            'ownership_type': target.ownership_type,
            'ownership_type_display': target.get_ownership_type_display() if hasattr(target, 'get_ownership_type_display') else target.ownership_type,
            'is_payable': is_payable,
            'daily_rate': daily_rate,
            'standard_days': standard_days,
            'full_month_days': full_month_days,
            'monthly_cost_standard': monthly_cost_standard,
            'monthly_cost_full': monthly_cost_full,
            'yearly_cost_standard': yearly_cost_standard,
            'avg_daily_rate': avg_rate,
            'avg_monthly_cost': avg_monthly_cost,
            'deviation_percentage': deviation_percentage,
            'benchmark_count': stats['count'] or 0
        })

    @action(detail=True, methods=['post'], url_path='reopen')
    def reopen(self, request, pk=None):
        """
        تسک ۳.۲-ب: بازگشایی پرونده ردشده (rejected -> draft) با قفل و بررسی وضعیت هم‌زمان
        """
        from ..fleet_lifecycle_services import user_has_any_perm, FLEET_SUPERVISOR_STAGE_PERMS, FLEET_MANAGER_STAGE_PERMS, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user

        is_auth = (
            user.is_superuser
            or user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS + FLEET_MANAGER_STAGE_PERMS)
            or target.created_by_id == user.id
            or (target.section_id and user_has_section_role(user, target.section_id))
        )
        if not is_auth:
            raise PermissionDenied("شما دسترسی بازگشایی این پرونده را ندارید.")

        reason = (request.data.get('reason') or 'بازگشایی مجدد پرونده ردشده جهت اصلاح').strip()
        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)
            if instance.approval_status != 'rejected':
                return Response(
                    {'error': f'فقط پرونده‌های در وضعیت ردشده امکان بازگشایی دارند (وضعیت فعلی: {instance.get_approval_status_display()}).'},
                    status=status.HTTP_409_CONFLICT
                )

            instance.approval_status = 'draft'
            instance.rejection_reason = None
            instance.save(update_fields=['approval_status', 'rejection_reason'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status='rejected',
                to_status='draft',
                action='reopen',
                reason=reason,
                metadata={'company_id': instance.company_id}
            )

        dispatch_vehicle_workflow_notification(instance, stage='reopened', actor=user, reason=reason)
        broadcast_vehicle_update(instance, action_type='updated', message=f'پرونده خودرو «{instance.plate_number}» بازگشایی شد و به پیش‌نویس بازگشت.', sender_id=user.id)
        return Response({
            'message': 'پرونده خودرو با موفقیت بازگشایی شد و به وضعیت پیش‌نویس بازگشت.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=True, methods=['post'], url_path='end-collaboration')
    def end_collaboration(self, request, pk=None):
        """
        تسک ۳.۱-ب: قطع همکاری خودرو؛ فقط برای خودروهای مصوب و فعال با مجوز مدیریت و end_date معتبر
        """
        from ..fleet_lifecycle_services import (
            user_has_any_perm,
            FLEET_MANAGER_STAGE_PERMS,
            normalize_shamsi_date,
            parse_shamsi_date,
            _clean_shamsi_date_str,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        if not user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS):
            raise PermissionDenied("فقط مدیر مجاز به ثبت پایان همکاری خودرو است.")

        target = self.get_object()
        if target.approval_status != 'approved':
            return Response(
                {'error': f'قطع همکاری فقط برای خودروهای مصوب امکان‌پذیر است (وضعیت فعلی: {target.get_approval_status_display()}).'},
                status=status.HTTP_400_BAD_REQUEST
            )

        reason = (request.data.get('reason') or '').strip()
        if not reason:
            return Response({'error': 'ثبت علت قطع همکاری الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        raw_end_date = request.data.get('end_date')
        if not raw_end_date or not str(raw_end_date).strip():
            return Response({'error': 'ثبت تاریخ پایان همکاری (end_date) الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            end_date = normalize_shamsi_date(raw_end_date)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        end_d_obj = parse_shamsi_date(end_date)

        with transaction.atomic():
            instance = VehicleDriverProfile.objects.select_for_update().get(pk=target.pk)

            # تسک ۱۱.۶ (ن-۶): شرط فعال بودن و عدم وجود تاریخ پایان همکاری درون قفل
            if not instance.is_active or instance.end_date:
                return Response(
                    {'error': f'همکاری با این خودرو قبلاً در تاریخ «{instance.end_date}» پایان یافته است و امکان ثبت مجدد وجود ندارد.', 'code': 'already_ended'},
                    status=status.HTTP_409_CONFLICT
                )

            # تسک ۱۱.۶ (ن-۶): یافتن آخرین تردد بر پایه حداکثر شیء تاریخ شمسی (DEF-20)
            trip_dates = list(VehicleTripLog.objects.filter(vehicle=instance, is_deleted=False).values_list('date_shamsi', flat=True).distinct())
            max_trip_d = None
            max_trip_str = None
            for ds in trip_dates:
                if ds:
                    parsed = parse_shamsi_date(_clean_shamsi_date_str(ds))
                    if parsed and (max_trip_d is None or parsed > max_trip_d):
                        max_trip_d = parsed
                        max_trip_str = ds

            if max_trip_d and end_d_obj and end_d_obj < max_trip_d:
                return Response({
                    'error': f'تاریخ پایان همکاری ({end_date}) نمی‌تواند قبل از تاریخ آخرین تردد ثبت‌شده ({max_trip_str}) باشد.'
                }, status=status.HTTP_400_BAD_REQUEST)

            instance.is_active = False
            instance.end_date = end_date
            instance.save(update_fields=['is_active', 'end_date'])

            WorkflowAuditLog.objects.create(
                content_type='vehicle',
                object_id=instance.id,
                actor=user,
                from_status=instance.approval_status,
                to_status=instance.approval_status,
                action='end_collaboration',
                reason=reason,
                metadata={'end_date': end_date, 'company_id': instance.company_id}
            )
            dispatch_vehicle_workflow_notification(instance, stage='end_collaboration', actor=user, reason=reason)

        broadcast_vehicle_update(instance, action_type='updated', message=f'همکاری با خودرو «{instance.plate_number}» در تاریخ {end_date} پایان یافت.', sender_id=user.id)
        return Response({
            'message': f'همکاری با خودرو «{instance.plate_number}» پایان یافت و پرونده غیرفعال گردید.',
            'data': self.get_serializer(instance).data
        })

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request):
        """
        دانلود قالب اکسل خام ناوگان و رانندگان با ساختار ۲ سطری استاندارد شرکت
        """
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Vehicles'
        ws.sheet_view.rightToLeft = True

        columns_fa = [
            'ردیف', 'شماره پلاک', 'نام راننده', 'کد ملی راننده', 'شماره همراه راننده',
            'نام مالک', 'کد ملی مالک', 'شماره همراه مالک',
            'نوع خودرو', 'نوع مالکیت', 'نرخ پایه سرویس (ریال)', 'شماره شبا', 'نام بانک', 'شماره حساب'
        ]
        columns_en = [
            'row_num', 'plate_number', 'driver_name', 'driver_national_code', 'driver_phone',
            'owner_name', 'owner_national_code', 'owner_phone',
            'vehicle_type', 'ownership_type', 'default_service_rate', 'sheba_number', 'bank_name', 'account_number'
        ]

        f_title = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        f_key = Font(name='Segoe UI', size=9, bold=True, color='CBD5E1')
        fill_title = PatternFill(start_color='0284C7', end_color='0284C7', fill_type='solid') # Sky-600
        fill_key = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 20
        ws.freeze_panes = 'A3'

        for col_idx, (c_fa, c_en) in enumerate(zip(columns_fa, columns_en), 1):
            cell_fa = ws.cell(row=1, column=col_idx, value=c_fa)
            cell_fa.font = f_title
            cell_fa.fill = fill_title
            cell_fa.alignment = align_center
            cell_fa.border = thin_border

            cell_en = ws.cell(row=2, column=col_idx, value=c_en)
            cell_en.font = f_key
            cell_en.fill = fill_key
            cell_en.alignment = align_center
            cell_en.border = thin_border

        # ردیف نمونه راهنما
        sample_row = [
            1, '12 الف 345 ایران 63', 'رضا اکبری', '0010376488', '09123456789',
            'رضا اکبری', '0010376488', '09123456789',
            'وانت نیسان', 'استیجاری', 4500000, 'IR120170000000123456789012', 'بانک ملی ایران', '0101234567001'
        ]
        ws.row_dimensions[3].height = 22
        for col_idx, val in enumerate(sample_row, 1):
            c = ws.cell(row=3, column=col_idx, value=val)
            c.alignment = align_center
            c.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="fleet_vehicles_template.xlsx"'
        return response

    @action(detail=False, methods=['post', 'POST', 'get', 'GET'], url_path='import-excel')
    def import_excel(self, request):
        """
        درون‌ریزی مستقیم شیت ناوگان از فایل اکسل شرکت (Upsert بر مبنای شماره پلاک خودرو)
        پشتیبانی از resolve_import_company، ensure_record_write_scope، اعتبارسنجی کد ملی،
        ایجاد حساب در VehicleBankAccount و عبور هوشمند با auto_pass_vehicle(..., draft_on_incomplete=True).
        """
        import openpyxl
        from decimal import Decimal
        from django.db import transaction
        from common.validators import validate_iranian_national_code
        from ..sheba_utils import clean_sheba, validate_sheba, get_bank_from_sheba
        from ..fleet_lifecycle_services import normalize_plate, auto_pass_vehicle
        from ..serializers import user_can_view_fleet_rates
        from personnel.models import VehicleBankAccount

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'فایل اکسل الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        
        section_id = request.data.get('section_id') or request.query_params.get('section_id')
        dry_run = str(request.data.get('dry_run') or request.query_params.get('dry_run', 'false')).lower() in ['true', '1']
        update_existing = str(request.data.get('update_existing') or request.query_params.get('update_existing', 'true')).lower() in ['true', '1']
        
        user = request.user
        is_mgr = bool(user and (user.is_superuser or user.has_perm('accounts.perm_approve_fleet_manager') or has_permission_alias(user, 'manager_approver')))
        can_view_rate = user_can_view_fleet_rates(user)

        # تعیین بخش پیش‌فرض
        target_section = None
        if section_id:
            try:
                target_section = ProjectSection.objects.filter(id=int(section_id)).select_related('project').first()
            except (ValueError, TypeError):
                pass
        if not target_section and not is_mgr and user and user.is_authenticated:
            # نخستین بخش تحت پوشش کاربر؛ انتساب سطح شرکت یا پروژه هم بخش‌های زیرمجموعه را می‌دهد (طرح ۲.۱)
            target_section = user_sections(user).select_related('project').order_by('id').first()

        # تعیین شرکت مقصد با resolve_import_company
        import_company_id = resolve_import_company(request)
        if target_section and target_section.project_id:
            sec_company_id = target_section.project.company_id
            if import_company_id and sec_company_id and import_company_id != sec_company_id:
                return Response({'error': 'بخش انتخاب‌شده با شرکت هدف تطابق ندارد.'}, status=status.HTTP_400_BAD_REQUEST)
            import_company_id = sec_company_id or import_company_id

        if requires_company_scope(user) and not import_company_id:
            return Response({'error': 'تعیین شرکت جهت درون‌ریزی الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            ensure_record_write_scope(user, company=import_company_id, section=target_section, label='ناوگان')
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_403_FORBIDDEN)

        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
        except Exception as e:
            return Response({
                'success': False,
                'dry_run': dry_run,
                'summary': {'total_rows': 0, 'created': 0, 'updated': 0, 'skipped': 0, 'valid_count': 0, 'error_count': 1},
                'errors': [{'row': 0, 'field': 'فایل اکسل', 'message': f'خطا در باز کردن فایل اکسل: {str(e)}'}]
            }, status=status.HTTP_400_BAD_REQUEST)
        
        target_sheet = None
        for name in wb.sheetnames:
            if name.strip().lower() in ['vehicles', 'fleet', 'ناوگان', 'خودرو', 'رانندگان']:
                target_sheet = wb[name]
                break
        if not target_sheet:
            target_sheet = wb.active

        # مپینگ نام ستون‌ها
        vtype_map = {
            'وانت نیسان': 'nissan', 'نیسان': 'nissan', 'nissan': 'nissan',
            'خاور': 'khavar', 'کامیونت': 'khavar', 'khavar': 'khavar',
            'تریلی': 'trailer', 'کشنده': 'trailer', 'trailer': 'trailer',
            'وانت بار': 'pickup', 'وانت': 'pickup', 'pickup': 'pickup',
            'کامیون': 'truck', 'تک': 'truck', 'جفت': 'truck', 'truck': 'truck',
            'لیفتراک': 'forklift', 'forklift': 'forklift',
            'سواری': 'sedan', 'sedan': 'sedan',
            'سایر': 'other', 'other': 'other'
        }
        own_map = {
            'استیجاری': 'contract', 'پیمانکاری': 'contract', 'contract': 'contract',
            'شرکتی': 'company', 'ملکی پروژه': 'company', 'company': 'company',
            'ملکی': 'personal', 'ملکی راننده': 'personal', 'personal': 'personal'
        }

        scoped_vehicles_qs = scope_vehicle_queryset(request, VehicleDriverProfile.objects.all())
        created_count = 0
        updated_count = 0
        skipped_count = 0
        errors = []
        warnings = []
        parsed_vehicles = []
        seen_file_plates = {}

        # بررسی سرستون‌های ردیف ۲ جهت نگاشت منعطف ستون‌ها
        col_map = {}
        for c in range(1, min(target_sheet.max_column + 1, 20)):
            val = str(target_sheet.cell(row=2, column=c).value or '').strip().lower()
            if val:
                col_map[val] = c

        has_owner_cols = 'owner_name' in col_map or 'owner_national_code' in col_map

        for row_idx in range(3, target_sheet.max_row + 1):
            plate_raw = target_sheet.cell(row=row_idx, column=col_map.get('plate_number', 2)).value
            if not plate_raw or not str(plate_raw).strip():
                continue

            plate = normalize_plate(str(plate_raw).strip())
            if not plate:
                continue

            if plate in seen_file_plates:
                errors.append({
                    'row': row_idx,
                    'field': 'شماره پلاک',
                    'message': f'شماره پلاک «{plate}» در ردیف {seen_file_plates[plate]} نیز تکرار شده است (پلاک تکراری در فایل اکسل).'
                })
                continue
            seen_file_plates[plate] = row_idx

            driver_name = str(target_sheet.cell(row=row_idx, column=col_map.get('driver_name', 3)).value or '').strip()
            if not driver_name:
                errors.append({'row': row_idx, 'field': 'نام راننده', 'message': f'نام راننده برای پلاک {plate} الزامی است.'})
                continue

            nat_code_raw = target_sheet.cell(row=row_idx, column=col_map.get('driver_national_code', 4)).value
            nat_code = str(nat_code_raw).strip() if nat_code_raw else ''
            if nat_code:
                clean_nat = ''.join(c for c in normalize_digits(nat_code) if c.isdigit())
                clean_nat = clean_nat.zfill(10)
                try:
                    nat_code = validate_iranian_national_code(clean_nat)
                except ValueError as e:
                    errors.append({
                        'row': row_idx,
                        'field': 'کد ملی راننده',
                        'message': f'کد ملی راننده نامعتبر است ({nat_code_raw}): {str(e)}'
                    })
                    continue
            else:
                nat_code = None

            phone_raw = str(target_sheet.cell(row=row_idx, column=col_map.get('driver_phone', 5)).value or '').strip()
            phone = normalize_digits(phone_raw) if phone_raw else None
            
            if has_owner_cols:
                owner_name = str(target_sheet.cell(row=row_idx, column=col_map.get('owner_name', 6)).value or '').strip() or None
                owner_nat_raw = target_sheet.cell(row=row_idx, column=col_map.get('owner_national_code', 7)).value
                owner_nat = str(owner_nat_raw).strip() if owner_nat_raw else ''
                if owner_nat:
                    clean_o_nat = ''.join(c for c in normalize_digits(owner_nat) if c.isdigit())
                    clean_o_nat = clean_o_nat.zfill(10)
                    try:
                        owner_nat = validate_iranian_national_code(clean_o_nat)
                    except ValueError as e:
                        errors.append({
                            'row': row_idx,
                            'field': 'کد ملی مالک',
                            'message': f'کد ملی مالک نامعتبر است ({owner_nat_raw}): {str(e)}'
                        })
                        continue
                else:
                    owner_nat = None
                owner_phone_raw = str(target_sheet.cell(row=row_idx, column=col_map.get('owner_phone', 8)).value or '').strip()
                owner_phone = normalize_digits(owner_phone_raw) if owner_phone_raw else None
                is_driver_owner = bool(not owner_nat or owner_nat == nat_code)

                vtype_col = col_map.get('vehicle_type', 9)
                own_col = col_map.get('ownership_type', 10)
                rate_col = col_map.get('default_service_rate', 11)
                sheba_col = col_map.get('sheba_number', 12)
                bank_col = col_map.get('bank_name', 13)
                acc_col = col_map.get('account_number', 14)
            else:
                owner_name = driver_name
                owner_nat = nat_code
                owner_phone = phone
                is_driver_owner = True

                vtype_col = col_map.get('vehicle_type', 6)
                own_col = col_map.get('ownership_type', 7)
                rate_col = col_map.get('default_service_rate', 8)
                sheba_col = col_map.get('sheba_number', 9)
                bank_col = col_map.get('bank_name', 10)
                acc_col = col_map.get('account_number', 11)

            vtype_raw = str(target_sheet.cell(row=row_idx, column=vtype_col).value or '').strip().lower()
            vehicle_type = vtype_map.get(vtype_raw)
            if not vehicle_type:
                errors.append({
                    'row': row_idx,
                    'field': 'نوع خودرو',
                    'message': f'نوع خودرو «{vtype_raw}» نامعتبر است. انواع معتبر: نیسان، خاور، تریلی، وانت بار، کامیون، لیفتراک، سواری، سایر.'
                })
                continue

            own_raw = str(target_sheet.cell(row=row_idx, column=own_col).value or '').strip().lower()
            ownership_type = own_map.get(own_raw)
            if not ownership_type:
                errors.append({
                    'row': row_idx,
                    'field': 'نوع مالکیت',
                    'message': f'نوع مالکیت «{own_raw}» نامعتبر است. انواع معتبر: استیجاری، شرکتی، ملکی.'
                })
                continue

            # سنجش پلاک درون شرکت جاری و قلمرو مجاز (یکتایی پلاک به‌ازای شرکت)
            existing_in_company = VehicleDriverProfile.objects.filter(company_id=import_company_id, plate_number=plate).first() if import_company_id else None
            if not existing_in_company and plate_raw:
                existing_in_company = VehicleDriverProfile.objects.filter(company_id=import_company_id, plate_number=str(plate_raw).strip()).first() if import_company_id else None

            # اگر خودروی موجود در وضعیت قفل (مصوب، در انتظار و ...) باشد، اکسل مجاز به تغییر آن نیست
            # طبق مصوبه ط-۴: این رکوردها نباید کل ایمپورت را با خطای ۴۰۰ متوقف کنند، بلکه با هشدار کنار گذاشته می‌شوند
            if existing_in_company and existing_in_company.approval_status not in ['draft', 'revision_required']:
                skipped_count += 1
                status_lbl = existing_in_company.get_approval_status_display() if hasattr(existing_in_company, 'get_approval_status_display') else existing_in_company.approval_status
                warnings.append(
                    f'ردیف {row_idx}: خودروی پلاک «{plate_raw}» در وضعیت «{status_lbl}» قرار دارد و تغییر آن از طریق اکسل مجاز نیست (نادیده گرفته شد).'
                )
                continue

            from ..fleet_lifecycle_services import is_valid_national_plate, ROAD_VEHICLE_TYPES
            if vehicle_type in ROAD_VEHICLE_TYPES and not is_valid_national_plate(plate):
                errors.append({
                    'row': row_idx,
                    'field': 'شماره پلاک',
                    'message': f'شماره پلاک «{plate_raw}» منطبق بر الگوی استاندارد پلاک ملی ایران نیست (مثال: ۱۲ب۳۴۵ایران۶۷ یا ۱۲-ب-۳۴۵-۶۷).'
                })
                continue

            if existing_in_company and not scoped_vehicles_qs.filter(id=existing_in_company.id).exists():
                errors.append({
                    'row': row_idx,
                    'field': 'شماره پلاک',
                    'message': f'شماره پلاک «{plate}» در این شرکت ثبت شده و شما دسترسی ویرایش آن را ندارید.'
                })
                continue

            if can_view_rate:
                rate_raw = target_sheet.cell(row=row_idx, column=rate_col).value
                try:
                    rate = Decimal(str(rate_raw or 0).replace(',', '').strip())
                    if rate < 0:
                        rate = Decimal(0)
                except Exception:
                    rate = Decimal(0)
            else:
                rate = Decimal(0)

            sheba_raw = str(target_sheet.cell(row=row_idx, column=sheba_col).value or '').strip()
            sheba = clean_sheba(sheba_raw) if sheba_raw else None
            
            bank_name = str(target_sheet.cell(row=row_idx, column=bank_col).value or '').strip() or None
            if sheba and not bank_name:
                b_info = get_bank_from_sheba(sheba)
                if b_info:
                    bank_name = b_info.get('name')

            account_no = str(target_sheet.cell(row=row_idx, column=acc_col).value or '').strip() or None

            parsed_vehicles.append({
                'row_idx': row_idx,
                'plate_number': plate,
                'driver_name': driver_name,
                'driver_national_code': nat_code,
                'driver_phone': phone,
                'is_driver_owner': is_driver_owner,
                'owner_name': owner_name,
                'owner_national_code': owner_nat,
                'owner_phone': owner_phone,
                'vehicle_type': vehicle_type,
                'ownership_type': ownership_type,
                'default_service_rate': rate,
                'sheba_number': sheba,
                'bank_name': bank_name,
                'account_number': account_no,
            })

        if errors:
            return Response({
                'success': False,
                'dry_run': dry_run,
                'summary': {
                    'total_rows': len(parsed_vehicles) + len(errors),
                    'created': 0, 'updated': 0, 'skipped': len(errors),
                    'valid_count': len(parsed_vehicles), 'error_count': len(errors)
                },
                'errors': errors
            }, status=status.HTTP_400_BAD_REQUEST)

        if not dry_run:
            from ..fleet_lifecycle_services import build_vehicle_snapshot
            from ..lifecycle_services import compute_profile_diff
            with transaction.atomic():
                for item in parsed_vehicles:
                    plate = item['plate_number']
                    q_exist = scoped_vehicles_qs.filter(plate_number=plate)
                    if import_company_id:
                        q_exist = q_exist.filter(company_id=import_company_id)
                    existing = q_exist.first()

                    if existing:
                        if update_existing:
                            # تسک ۱۱.۸ و ط-۳: به‌روزرسانی فقط در draft و revision_required
                            if existing.approval_status not in ['draft', 'revision_required']:
                                skipped_count += 1
                                warnings.append(
                                    f'ردیف {item["row_idx"]}: خودروی پلاک «{plate}» در وضعیت «{existing.get_approval_status_display() or existing.approval_status}» قرار دارد و تغییر آن از طریق اکسل مجاز نیست (نادیده گرفته شد).'
                                )
                                continue

                            old_snap = build_vehicle_snapshot(existing)

                            existing.driver_name = item['driver_name']
                            if item['driver_national_code']:
                                existing.driver_national_code = item['driver_national_code']
                            if item['driver_phone']:
                                existing.driver_phone = item['driver_phone']
                            existing.is_driver_owner = item['is_driver_owner']
                            if item['owner_name']:
                                existing.owner_name = item['owner_name']
                            if item['owner_national_code']:
                                existing.owner_national_code = item['owner_national_code']
                            if item['owner_phone']:
                                existing.owner_phone = item['owner_phone']
                            existing.vehicle_type = item['vehicle_type']
                            existing.ownership_type = item['ownership_type']
                            if can_view_rate and item['default_service_rate'] is not None:
                                existing.default_service_rate = item['default_service_rate']
                            if item['sheba_number']:
                                existing.sheba_number = item['sheba_number']
                            if item['bank_name']:
                                existing.bank_name = item['bank_name']
                            if item['account_number']:
                                existing.account_number = item['account_number']
                            if target_section and not existing.section_id:
                                existing.section = target_section
                                if target_section.project:
                                    existing.project = target_section.project
                            
                            # تسک ۱۱.۸ و ط-۵: اعمال _skip_bank_sync=True جهت جلوگیری از تکرار لاجیک مدل
                            existing._skip_bank_sync = True
                            existing.save()

                            # به‌روزرسانی یا ایجاد حساب اصلی در VehicleBankAccount
                            if item['sheba_number'] or item['account_number']:
                                pri_acc = VehicleBankAccount.objects.filter(vehicle=existing, is_primary=True).first()
                                if pri_acc:
                                    if item['bank_name']:
                                        pri_acc.bank_name = item['bank_name']
                                    if item['account_number']:
                                        pri_acc.account_number = item['account_number']
                                    if item['sheba_number']:
                                        pri_acc.sheba_number = item['sheba_number']
                                    pri_acc.save()
                                else:
                                    VehicleBankAccount.objects.create(
                                        vehicle=existing,
                                        bank_name=item['bank_name'] or '',
                                        account_number=item['account_number'] or '',
                                        sheba_number=item['sheba_number'] or '',
                                        account_owner_name=existing.driver_name if existing.is_driver_owner else (existing.owner_name or existing.driver_name),
                                        account_owner_national_code=existing.driver_national_code if existing.is_driver_owner else (existing.owner_national_code or existing.driver_national_code),
                                        payee_type='driver' if existing.is_driver_owner else 'owner',
                                        is_primary=True,
                                        is_active=True
                                    )

                            new_snap = build_vehicle_snapshot(existing)
                            diff = {}
                            for k, new_val in new_snap.items():
                                old_val = old_snap.get(k)
                                if old_val != new_val:
                                    diff[k] = {'old': old_val, 'new': new_val}

                            WorkflowAuditLog.objects.create(
                                content_type='vehicle',
                                object_id=existing.id,
                                actor=user if user and user.is_authenticated else None,
                                from_status=existing.approval_status,
                                to_status=existing.approval_status,
                                action='import_update',
                                reason=f'به‌روزرسانی مشخصات خودرو از طریق فایل اکسل (ردیف {item["row_idx"]})',
                                metadata={'company_id': existing.company_id, 'row': item['row_idx'], 'diff': diff}
                            )
                            updated_count += 1
                        else:
                            skipped_count += 1
                    else:
                        # تسک ۱۱.۸ و ط-۵: ایجاد اولیه بدون فیلدهای تخت بانک تا حساب اصلی مضاعف ساخته نشود
                        new_v = VehicleDriverProfile.objects.create(
                            company_id=import_company_id,
                            plate_number=plate,
                            driver_name=item['driver_name'],
                            driver_national_code=item['driver_national_code'],
                            driver_phone=item['driver_phone'],
                            is_driver_owner=item['is_driver_owner'],
                            owner_name=item['owner_name'],
                            owner_national_code=item['owner_national_code'],
                            owner_phone=item['owner_phone'],
                            vehicle_type=item['vehicle_type'],
                            ownership_type=item['ownership_type'],
                            default_service_rate=item['default_service_rate'],
                            sheba_number='',
                            bank_name='',
                            account_number='',
                            section=target_section,
                            project=target_section.project if (target_section and target_section.project) else None,
                            created_by=user if user and user.is_authenticated else None,
                            approval_status='draft',
                            is_active=True
                        )
                        # ایجاد صریح رکورد اصلی در VehicleBankAccount و سپس تنظیم فیلدهای تخت با _skip_bank_sync=True
                        if item['sheba_number'] or item['account_number']:
                            VehicleBankAccount.objects.create(
                                vehicle=new_v,
                                bank_name=item['bank_name'] or '',
                                account_number=item['account_number'] or '',
                                sheba_number=item['sheba_number'] or '',
                                account_owner_name=item['driver_name'] if item['is_driver_owner'] else (item['owner_name'] or item['driver_name']),
                                account_owner_national_code=item['driver_national_code'] if item['is_driver_owner'] else (item['owner_national_code'] or item['driver_national_code']),
                                payee_type='driver' if item['is_driver_owner'] else 'owner',
                                is_primary=True,
                                is_active=True
                            )
                            new_v.sheba_number = item['sheba_number'] or ''
                            new_v.bank_name = item['bank_name'] or ''
                            new_v.account_number = item['account_number'] or ''
                            new_v._skip_bank_sync = True
                            new_v.save(update_fields=['sheba_number', 'bank_name', 'account_number'])

                        # اجرای auto_pass_vehicle با draft_on_incomplete=True
                        auto_pass_vehicle(new_v, user, source='import_excel', draft_on_incomplete=True)
                        created_count += 1
        else:
            for item in parsed_vehicles:
                q_chk = scoped_vehicles_qs.filter(plate_number=item['plate_number'])
                if import_company_id:
                    q_chk = q_chk.filter(company_id=import_company_id)
                existing = q_chk.first()
                if existing:
                    if update_existing:
                        if existing.approval_status not in ['draft', 'revision_required']:
                            skipped_count += 1
                            warnings.append(
                                f'ردیف {item["row_idx"]}: خودروی پلاک «{item["plate_number"]}» در وضعیت «{existing.get_approval_status_display() or existing.approval_status}» قرار دارد و از طریق اکسل قابل ویرایش نیست (نادیده گرفته خواهد شد).'
                            )
                        else:
                            updated_count += 1
                    else:
                        skipped_count += 1
                else:
                    created_count += 1

        return Response({
            'success': True,
            'dry_run': dry_run,
            'summary': {
                'total_rows': len(parsed_vehicles),
                'created': created_count,
                'updated': updated_count,
                'skipped': skipped_count,
                'valid_count': len(parsed_vehicles),
                'error_count': 0
            },
            'created_count': created_count,
            'updated_count': updated_count,
            'warnings': warnings,
            'errors': []
        })

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """
        خروجی رسمی ۲ ردیفه اطلاعات ناوگان و رانندگان (عنوان فارسی + کلید لاتین + فریز A3)
        محدود به شرکت‌های مجاز کاربر (readable_company_ids) و تفکیک وظایف مالی (SoD) برای ستون نرخ.
        """
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse
        from django.utils import timezone
        from ..serializers import user_can_view_fleet_rates

        qs = self.filter_queryset(self.get_queryset())
        cids = readable_company_ids(request.user)
        if cids is not None:
            qs = qs.filter(company_id__in=cids)

        can_view_rate = user_can_view_fleet_rates(request.user)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ناوگان و رانندگان"
        ws.views.sheetView[0].rightToLeft = True

        # تعریف ستون‌ها به همراه کلید، جهت متن، فونت عددی، و فرمت عددی
        cols_spec = [
            ('ردیف', 'row_num', 'center', True, None),
            ('شماره پلاک', 'plate_number', 'center', True, None),
            ('نوع خودرو', 'vehicle_type_display', 'center', False, None),
            ('نوع مالکیت', 'ownership_type_display', 'center', False, None),
            ('نام راننده', 'driver_name', 'right', False, None),
            ('کد ملی راننده', 'driver_national_code', 'center', True, None),
            ('شماره همراه راننده', 'driver_phone', 'center', True, None),
            ('مالک شخص راننده است', 'is_driver_owner', 'center', False, None),
            ('نام مالک', 'owner_name', 'right', False, None),
            ('کد ملی مالک', 'owner_national_code', 'center', True, None),
            ('شماره همراه مالک', 'owner_phone', 'center', True, None),
        ]
        if can_view_rate:
            cols_spec.append(('نرخ پایه سرویس (ریال)', 'default_service_rate', 'center', True, '#,##0'))
        cols_spec.extend([
            ('شماره حساب', 'account_number', 'center', True, None),
            ('شماره شبا', 'sheba_number', 'center', True, None),
            ('نام بانک', 'bank_name', 'center', False, None),
            ('وضعیت تایید پرونده', 'approval_status_display', 'center', False, None),
            ('وضعیت فعال', 'is_active', 'center', False, None),
        ])

        headers_fa = [col[0] for col in cols_spec]
        headers_en = [col[1] for col in cols_spec]

        title_font = Font(name='B Nazanin', size=11, bold=True, color='FFFFFF')
        key_font = Font(name='Segoe UI', size=8, italic=True, color='94A3B8')
        title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        key_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(headers_fa)
        ws.append(headers_en)

        for col_idx in range(1, len(headers_fa) + 1):
            cell_fa = ws.cell(row=1, column=col_idx)
            cell_fa.font = title_font
            cell_fa.fill = title_fill
            cell_fa.alignment = Alignment(horizontal='center', vertical='center')

            cell_en = ws.cell(row=2, column=col_idx)
            cell_en.font = key_font
            cell_en.fill = key_fill
            cell_en.alignment = Alignment(horizontal='center', vertical='center')

        ws.freeze_panes = 'A3'

        data_font = Font(name='B Nazanin', size=10)
        data_font_num = Font(name='Segoe UI', size=10)

        for idx, v in enumerate(qs, start=1):
            curr_row = idx + 2
            is_owner = v.is_driver_owner if v.is_driver_owner is not None else True
            val_map = {
                'row_num': idx,
                'plate_number': v.plate_number,
                'vehicle_type_display': v.get_vehicle_type_display() if hasattr(v, 'get_vehicle_type_display') else v.vehicle_type,
                'ownership_type_display': v.get_ownership_type_display() if hasattr(v, 'get_ownership_type_display') else (v.ownership_type or '-'),
                'driver_name': v.driver_name,
                'driver_national_code': v.driver_national_code or '-',
                'driver_phone': v.driver_phone or '-',
                'is_driver_owner': 'بله' if is_owner else 'خیر',
                'owner_name': v.owner_name or (v.driver_name if is_owner else '-'),
                'owner_national_code': v.owner_national_code or (v.driver_national_code if is_owner else '-'),
                'owner_phone': v.owner_phone or (v.driver_phone if is_owner else '-'),
                'default_service_rate': float(v.default_service_rate or 0),
                'account_number': v.account_number or '-',
                'sheba_number': v.sheba_number or '-',
                'bank_name': v.bank_name or '-',
                'approval_status_display': v.get_approval_status_display() if hasattr(v, 'get_approval_status_display') else str(v.approval_status),
                'is_active': 'فعال' if v.is_active else 'غیرفعال',
            }
            row_data = [val_map[col[1]] for col in cols_spec]
            ws.append(row_data)

            for col_idx, col_spec in enumerate(cols_spec, start=1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal=col_spec[2], vertical='center')
                c.font = data_font_num if col_spec[3] else data_font
                if col_spec[4]:
                    c.number_format = col_spec[4]

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"fleet_vehicles_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

class VehicleChangeRequestViewSet(viewsets.ModelViewSet):
    queryset = VehicleChangeRequest.objects.all().select_related(
        'vehicle', 'requested_by', 'manager_reviewed_by', 'accountant_reviewed_by', 'supervisor_reviewed_by'
    )
    serializer_class = VehicleChangeRequestSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        qs = scope_vehicle_queryset(self.request, qs, prefix='vehicle__')

        status_filter = self.request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        vehicle_id = self.request.query_params.get('vehicle_id')
        if vehicle_id:
            qs = qs.filter(vehicle_id=vehicle_id)
        return qs

    def create(self, request, *args, **kwargs):
        """
        تسک ۳.۵ و ب-۲۳: ثبت مستقیم درخواست تغییرات از این مسیر مسدود است (۴۰۵).
        پیشنهاد تغییرات صرفاً از طریق اکشن اختصاصی پرونده خودرو انجام می‌شود.
        """
        return Response(
            {'error': 'ایجاد مستقیم درخواست تغییرات از این مسیر مجاز نیست. لطفاً از طریق اکشن اختصاصی پرونده خودرو اقدام فرمایید.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def update(self, request, *args, **kwargs):
        """تسک ۳.۵ و ب-۲۳: ویرایش مستقیم درخواست تغییرات مجاز نیست (۴۰۵)."""
        return Response(
            {'error': 'ویرایش مستقیم درخواست تغییرات ناوگان مجاز نیست.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def partial_update(self, request, *args, **kwargs):
        """تسک ۳.۵ و ب-۲۳: ویرایش مستقیم درخواست تغییرات مجاز نیست (۴۰۵)."""
        return Response(
            {'error': 'ویرایش مستقیم درخواست تغییرات ناوگان مجاز نیست.'},
            status=status.HTTP_405_METHOD_NOT_ALLOWED
        )

    def destroy(self, request, *args, **kwargs):
        """
        تسک ۳.۵ و ب-۲۳: حذف فیزیکی رکورد درخواست تغییرات منحصراً برای سوپریوزر (۴۰۳ برای بقیه).
        """
        if not request.user.is_superuser:
            raise PermissionDenied("حذف فیزیکی درخواست تغییرات ناوگان منحصراً در اختیار مدیر ارشد سیستم (سوپریوزر) است.")
        from ..fleet_lifecycle_services import log_vehicle_change_request_event, refresh_vehicle_has_pending_changes
        cr = self.get_object()
        vehicle = cr.vehicle
        with transaction.atomic():
            log_vehicle_change_request_event(cr, request.user, 'change_request_deleted', reason='حذف توسط سوپریوزر')
            response = super().destroy(request, *args, **kwargs)
            refresh_vehicle_has_pending_changes(vehicle)
            return response

    @action(detail=True, methods=['post'], url_path='cancel')
    def cancel(self, request, pk=None):
        """
        تسک ۳.۵ و ب-۲۳: لغو درخواست تغییرات فقط برای ثبت‌کننده (یا سوپریوزر)
        و فقط تا پیش از تایید سرپرست (draft, pending_supervisor, revision_required).
        """
        from ..fleet_lifecycle_services import log_vehicle_change_request_event, refresh_vehicle_has_pending_changes, dispatch_vehicle_workflow_notification
        target = self.get_object()
        user = request.user
        if not (user.is_superuser or target.requested_by_id == user.id):
            raise PermissionDenied("تنها کاربر ثبت‌کننده درخواست یا مدیر سیستم مجاز به انصراف و لغو درخواست تغییرات است.")

        if target.status not in ('draft', 'pending_supervisor', 'revision_required'):
            return Response(
                {'error': f'درخواست در وضعیت «{target.get_status_display()}» قابل لغو نیست؛ ابطال فقط تا پیش از تایید سرپرست ممکن است.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        reason = (request.data.get('reason') or 'لغو شده توسط ثبت‌کننده').strip()
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status not in ('draft', 'pending_supervisor', 'revision_required'):
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'cancelled'
            cr.rejection_reason = reason
            cr.save(update_fields=['status', 'rejection_reason'])
            refresh_vehicle_has_pending_changes(cr.vehicle, exclude_cr_id=cr.id)
            log_vehicle_change_request_event(cr, user, 'change_request_cancelled', reason=reason)
            log_vehicle_audit(user, 'CANCEL_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_cancelled', actor=user, reason=reason, recipient_ids=[cr.requested_by_id])
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» لغو شد.', sender_id=user.id if user else None)

        return Response({
            'message': 'درخواست تغییرات با موفقیت لغو شد.',
            'data': self.get_serializer(cr).data
        })

    @action(detail=True, methods=['post'], url_path='approve-supervisor')
    def approve_supervisor(self, request, pk=None):
        from ..fleet_lifecycle_services import log_vehicle_change_request_event, FLEET_SUPERVISOR_STAGE_PERMS, user_has_any_perm, dispatch_vehicle_workflow_notification
        user = request.user
        if not (user.is_superuser or user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS)):
            raise PermissionDenied("شما دسترسی بررسی مرحله اول (سرپرست) درخواست تغییرات ناوگان را ندارید.")
        
        target = self.get_object()
        # تسک ۱۱.۵ و ن-۵: فقط وضعیت pending_supervisor مورد تایید سرپرست است و تغییر وضعیت ۴۰۹ می‌گیرد
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status != 'pending_supervisor':
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'pending_accountant'
            cr.supervisor_reviewed_by = user
            cr.supervisor_reviewed_at = timezone.now()
            cr.save(update_fields=['status', 'supervisor_reviewed_by', 'supervisor_reviewed_at'])
            log_vehicle_change_request_event(cr, user, 'change_request_supervisor_approved')
            log_vehicle_audit(user, 'APPROVE_SUPERVISOR_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_supervisor_approved', actor=user)
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» به تایید سرپرست رسید.', sender_id=user.id if user else None)
        return Response({'message': 'درخواست تغییرات خودرو به تایید سرپرست رسید و به کارتابل حسابداری ارسال شد.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='approve-finance')
    def approve_finance(self, request, pk=None):
        from ..fleet_lifecycle_services import log_vehicle_change_request_event, FLEET_FINANCE_STAGE_PERMS, user_has_any_perm, dispatch_vehicle_workflow_notification
        user = request.user
        # ب-۸: فقط perm_approve_fleet_finance (بی نشانهٔ نقش حسابدار)
        if not (user.is_superuser or user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS)):
            raise PermissionDenied("شما دسترسی تایید مرحله دوم (مالی) درخواست تغییرات ناوگان را ندارید.")
        
        target = self.get_object()
        valid_statuses = ['pending_accountant', 'supervisor_approved']
        if target.status not in valid_statuses:
            return Response({'error': f'درخواست ابتدا باید به تایید سرپرست برسد (وضعیت فعلی: {target.get_status_display()}).'}, status=status.HTTP_400_BAD_REQUEST)
        
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status not in valid_statuses:
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'pending_manager'
            cr.accountant_reviewed_by = user
            cr.accountant_reviewed_at = timezone.now()
            cr.save(update_fields=['status', 'accountant_reviewed_by', 'accountant_reviewed_at'])
            log_vehicle_change_request_event(cr, user, 'change_request_finance_approved')
            log_vehicle_audit(user, 'APPROVE_FINANCE_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_finance_approved', actor=user)
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» به تایید مالی رسید.', sender_id=user.id if user else None)
        return Response({'message': 'درخواست تغییرات خودرو به تایید مالی رسید و به کارتابل مدیر ارسال شد.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='approve-manager')
    def approve_manager(self, request, pk=None):
        from ..fleet_lifecycle_services import (
            log_vehicle_change_request_event,
            apply_vehicle_change_request,
            refresh_vehicle_has_pending_changes,
            FLEET_MANAGER_STAGE_PERMS,
            user_has_any_perm,
            VehicleWorkflowError,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        if not (user.is_superuser or user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS)):
            raise PermissionDenied("شما دسترسی تصویب نهایی درخواست تغییرات ناوگان را ندارید.")
        
        target = self.get_object()
        valid_statuses = ['pending_manager', 'accountant_approved']
        if user.is_superuser:
            valid_statuses += ['pending_supervisor', 'pending_accountant', 'draft', 'manager_approved']
        
        if target.status not in valid_statuses:
            return Response({'error': f'درخواست در وضعیت «{target.get_status_display()}» قابل تصویب مدیر نیست.'}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            with transaction.atomic():
                cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
                if cr.status not in valid_statuses:
                    return Response(
                        {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                        status=status.HTTP_409_CONFLICT
                    )
                # اعمال امن تغییرات بر روی پرونده خودرو با سریالایزر بدون هیچگونه setattr خام (تسک ۳.۵ و ب-۲۳)
                vehicle = apply_vehicle_change_request(cr, request=request)

                cr.status = 'approved'
                cr.manager_reviewed_by = user
                cr.manager_reviewed_at = timezone.now()
                if not cr.supervisor_reviewed_by:
                    cr.supervisor_reviewed_by = user
                    cr.supervisor_reviewed_at = timezone.now()
                if not cr.accountant_reviewed_by:
                    cr.accountant_reviewed_by = user
                    cr.accountant_reviewed_at = timezone.now()
                cr.save()

                refresh_vehicle_has_pending_changes(vehicle, exclude_cr_id=cr.id)
                log_vehicle_change_request_event(cr, user, 'change_request_applied')
                log_vehicle_audit(user, 'APPROVE_MANAGER_CHANGE_REQ', vehicle, {'change_request_id': cr.id, 'changes': cr.proposed_changes})
                dispatch_vehicle_workflow_notification(vehicle, stage='change_request_manager_approved', actor=user, recipient_ids=[cr.requested_by_id])
                broadcast_vehicle_update(vehicle, action_type='updated', message=f'تغییرات خودرو «{vehicle.plate_number}» توسط مدیر تایید و اعمال گردید.', sender_id=user.id if user else None)
        except VehicleWorkflowError as e:
            return Response(e.to_response_data(), status=e.status_code)

        return Response({'message': 'درخواست تغییرات خودرو توسط مدیر تصویب و اعمال گردید.', 'data': self.get_serializer(cr).data})

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        from ..fleet_lifecycle_services import (
            log_vehicle_change_request_event,
            refresh_vehicle_has_pending_changes,
            FLEET_SUPERVISOR_STAGE_PERMS,
            FLEET_FINANCE_STAGE_PERMS,
            FLEET_MANAGER_STAGE_PERMS,
            user_has_any_perm,
            dispatch_vehicle_workflow_notification,
        )
        user = request.user
        can_op = (
            user.is_superuser
            or user_has_any_perm(user, FLEET_SUPERVISOR_STAGE_PERMS)
            or user_has_any_perm(user, FLEET_FINANCE_STAGE_PERMS)
            or user_has_any_perm(user, FLEET_MANAGER_STAGE_PERMS)
        )
        if not can_op:
            raise PermissionDenied("دسترسی رد این درخواست را ندارید.")
        
        reason = request.data.get('reason', '').strip()
        if not reason:
            return Response({'error': 'ثبت دلیل رد الزامی است.'}, status=status.HTTP_400_BAD_REQUEST)
        
        target = self.get_object()
        active_statuses = ['draft', 'pending_supervisor', 'pending_accountant', 'pending_manager', 'revision_required', 'supervisor_approved', 'accountant_approved']
        if target.status not in active_statuses:
            return Response({'error': f'درخواست در وضعیت «{target.get_status_display()}» قابل رد نیست.'}, status=status.HTTP_400_BAD_REQUEST)
        
        with transaction.atomic():
            cr = VehicleChangeRequest.objects.select_for_update().get(pk=target.pk)
            if cr.status not in active_statuses:
                return Response(
                    {'error': f'وضعیت درخواست توسط کاربر دیگری تغییر یافته است (وضعیت فعلی: {cr.get_status_display()}).', 'code': 'conflict'},
                    status=status.HTTP_409_CONFLICT
                )
            cr.status = 'rejected'
            cr.rejection_reason = reason
            cr.save(update_fields=['status', 'rejection_reason'])
            refresh_vehicle_has_pending_changes(cr.vehicle, exclude_cr_id=cr.id)
            log_vehicle_change_request_event(cr, user, 'change_request_rejected', reason=reason)
            log_vehicle_audit(user, 'REJECT_CHANGE_REQ', cr.vehicle, {'change_request_id': cr.id, 'reason': reason})
            dispatch_vehicle_workflow_notification(cr.vehicle, stage='change_request_rejected', actor=user, reason=reason, recipient_ids=[cr.requested_by_id])
            broadcast_vehicle_update(cr.vehicle, action_type='updated', message=f'درخواست تغییرات خودرو «{cr.vehicle.plate_number}» رد شد: {reason}', sender_id=user.id if user else None)

        return Response({'message': 'درخواست تغییرات خودرو رد شد.', 'data': self.get_serializer(cr).data})

__all__ = [
    'VehicleDriverProfileViewSet',
    'VehicleChangeRequestViewSet',
]
